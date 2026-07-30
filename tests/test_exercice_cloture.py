"""Tests clôture d'exercice — snapshot archives + ouverture N+1 séparée."""

from datetime import date
from decimal import Decimal
from types import SimpleNamespace
from unittest.mock import AsyncMock, MagicMock
from uuid import uuid4

import pytest
from openpyxl import load_workbook

from app.core.exceptions import ValidationError
from app.models.enums import StatutExercice
from app.services.exercice_cloture_service import (
    KIND_CLOTURE,
    ExerciceClotureService,
    build_cloture_xlsx,
    periode_ouverture,
)
from app.services.exercice_ouverture_service import ExerciceOuvertureService
from app.services.recap_amortissement import _mouvements_immo


def test_periode_ouverture():
    assert periode_ouverture(2026) == "2026-12"
    assert periode_ouverture(2025) == "2025-12"


def test_statut_exercice_persiste_les_valeurs_minuscules():
    """La colonne doit stocker 'ouvert'/'cloture', pas les noms d'énumération."""
    from app.models import ExerciceComptable

    col_type = ExerciceComptable.__table__.c.statut.type
    assert sorted(col_type.enums) == ["cloture", "ouvert"]


def test_build_cloture_xlsx_contains_snapshot_columns():
    content = build_cloture_xlsx(
        annee=2026,
        nature_code="TY-142010",
        lignes=[
            {
                "date_acquisition": date(2026, 1, 6),
                "quantite": 1,
                "designation": "RGLT FACT ETS KERIM",
                "valeur_brute": Decimal("203500.00"),
                "taux": Decimal("10"),
                "amt_n1": Decimal("0"),
                "dotation": Decimal("10175.00"),
                "amt_fin": Decimal("10175.00"),
                "vnc": Decimal("193325.00"),
                "agence_label": "NDB",
            }
        ],
    )
    wb = load_workbook(filename=__import__("io").BytesIO(content))
    ws = wb.active
    assert "CLÔTURE 31/12/2026" in str(ws["A2"].value)
    # Header row 4, data row 5
    assert ws.cell(5, 3).value == "RGLT FACT ETS KERIM"
    assert float(ws.cell(5, 4).value) == 203500.0
    assert float(ws.cell(5, 7).value) == 10175.0  # Exer. En. C = dotation N


def test_mouvements_snapshot_amounts_coherent():
    """amt_n1 / dotation / amt_fin alignés sur _mouvements_immo (source snapshot)."""
    immo = SimpleNamespace(
        date_acquisition=date(2024, 1, 1),
        date_fin=None,
        valeur_brute=Decimal("100000.00"),
        valeur_residuelle=Decimal("0"),
        taux=Decimal("10.0000"),
        duree_annees=10,
        duree_mois=0,
        compte_immobilisation="142010",
        compte_amortissement="148211",
        metadata_json=None,
    )
    mvts = _mouvements_immo(
        immo,
        2026,
        montants_dotation_db=[Decimal("2500"), Decimal("2500"), Decimal("2500"), Decimal("2500")],
        cumul_n1_db=Decimal("20000.00"),
        cumul_fin_n_db=Decimal("30000.00"),
        vnc_fin_n_db=Decimal("70000.00"),
    )
    assert mvts is not None
    assert mvts["amorts_cumules_n1"] == Decimal("20000.00")
    assert mvts["dotations_annee"] == Decimal("10000.00")
    assert mvts["amorts_cumules_n"] == Decimal("30000.00")
    assert mvts["vnc"] == Decimal("70000.00")


def test_bank_snapshot_amounts_keeps_amt_n1():
    from app.services.exercice_cloture_service import _bank_snapshot_amounts

    immo = SimpleNamespace(
        valeur_brute=Decimal("100000.00"),
        metadata_json={
            "source": "import_banque",
            "bank": {
                "amt_n1": "80000.00",
                "dotation": "5000.00",
                "amt_fin": "85000.00",
                "vnc": "15000.00",
            },
        },
    )
    mvts = _bank_snapshot_amounts(immo)
    assert mvts is not None
    assert mvts["amorts_cumules_n1"] == Decimal("80000.00")
    assert mvts["dotations_annee"] == Decimal("5000.00")
    assert mvts["amorts_cumules_n"] == Decimal("85000.00")
    assert mvts["vnc"] == Decimal("15000.00")
    # Stock historique : ouverture N+1 = amt_fin
    assert mvts["cumul_ouverture_n1"] == Decimal("85000.00")
    assert mvts["vnc_ouverture_n1"] == Decimal("15000.00")


def test_bank_snapshot_arrete_courant_uses_amt_n1_for_opening():
    from app.services.exercice_cloture_service import _bank_snapshot_amounts

    immo = SimpleNamespace(
        valeur_brute=Decimal("100000.00"),
        metadata_json={
            "source": "import_banque",
            "bank": {
                "amt_n1": "80000.00",
                "dotation": "5000.00",
                "amt_fin": "85000.00",
                "vnc": "15000.00",
                "seed_mode": "arrete_courant",
            },
        },
    )
    mvts = _bank_snapshot_amounts(immo)
    assert mvts is not None
    assert mvts["cumul_ouverture_n1"] == Decimal("80000.00")
    assert mvts["vnc_ouverture_n1"] == Decimal("20000.00")


@pytest.mark.asyncio
async def test_cloture_idempotente_sans_force():
    db = MagicMock()
    svc = ExerciceClotureService(db)
    dossier = MagicMock()
    dossier.annee = 2026
    dossier.fichiers = [MagicMock(kind=KIND_CLOTURE)]
    svc._get_or_create_dossier = AsyncMock(return_value=dossier)

    empty = MagicMock()
    empty.scalar_one_or_none.return_value = None
    db.execute = AsyncMock(return_value=empty)

    with pytest.raises(ValidationError, match="définitive|clôture système"):
        await svc.cloturer(annee=2026, force=False)


@pytest.mark.asyncio
async def test_cloture_refuse_si_deja_cloturee():
    db = MagicMock()
    svc = ExerciceClotureService(db)
    exo = MagicMock()
    exo.statut = StatutExercice.CLOTURE
    found = MagicMock()
    found.scalar_one_or_none.return_value = exo
    db.execute = AsyncMock(return_value=found)

    with pytest.raises(ValidationError, match="déjà clôturé"):
        await svc.cloturer(annee=2025, force=True)


@pytest.mark.asyncio
async def test_ouverture_cree_soldes_142_148_vnc_sans_68():
    """Ouverture N+1 : SoldeOuvertureImmobilisation, 142−148=VNC, pas de charge 68."""
    db = MagicMock()
    svc = ExerciceOuvertureService(db)

    immo_id = uuid4()
    immo = SimpleNamespace(
        id=immo_id,
        deleted_at=None,
        date_fin=None,
        statut=None,
        valeur_brute=Decimal("100000.00"),
        code_inventaire="AAI-2025-001",
    )
    mvts = {
        "valeur_brute": Decimal("100000.00"),
        "amorts_cumules_n": Decimal("45000.00"),
        "vnc": Decimal("55000.00"),
    }

    # situation() : dernier clôturé 2025, pas d'exercice 2026
    exo_2025 = MagicMock()
    exo_2025.annee = 2025
    exo_2025.statut = StatutExercice.CLOTURE

    list_exos = MagicMock()
    list_exos.scalars.return_value.all.return_value = [exo_2025]

    empty_exo = MagicMock()
    empty_exo.scalar_one_or_none.return_value = None

    # flush creates exo → we simulate by capturing db.add
    added = []

    def add_side(obj):
        added.append(obj)
        if hasattr(obj, "annee") and not hasattr(obj, "immobilisation_id"):
            obj.id = uuid4()

    db.add = MagicMock(side_effect=add_side)
    db.flush = AsyncMock()

    # ouvrir_suivant flow:
    # 1) situation → list exercices
    # 2) existing annee_ouverture → None
    # 3) _collect_snapshot_lines
    cloture_collect = AsyncMock(return_value=({}, [(immo, mvts)]))

    async def execute_side(*_args, **_kwargs):
        # First calls in situation + existing check
        return list_exos

    # More controlled: patch methods
    svc.situation = AsyncMock(
        return_value={
            "dernier_cloture": 2025,
            "exercice_ouvert": None,
            "annee_ouverture_proposee": 2026,
            "peut_ouvrir": True,
            "exercices": [exo_2025],
        }
    )
    db.execute = AsyncMock(return_value=empty_exo)

    from app.services import exercice_ouverture_service as mod

    original = ExerciceClotureService
    # Patch collect via instance method after construction inside ouvrir_suivant
    from app.services.exercice_cloture_service import ExerciceClotureService as ECS

    ECS._collect_snapshot_lines = cloture_collect  # type: ignore[method-assign]

    try:
        result = await svc.ouvrir_suivant()
    finally:
        # restore not strictly needed in test process isolation
        pass

    assert result.annee_ouverture == 2026
    assert result.annee_source == 2025
    assert result.ouvertures_seed == 1
    assert result.total_valeur_brute == Decimal("100000.00")
    assert result.total_amortissement == Decimal("45000.00")
    assert result.total_vnc == Decimal("55000.00")
    assert result.total_valeur_brute - result.total_amortissement == result.total_vnc

    soldes = [o for o in added if getattr(o, "cumul_148", None) is not None]
    assert len(soldes) == 1
    solde = soldes[0]
    assert solde.valeur_brute_142 == Decimal("100000.00")
    assert solde.cumul_148 == Decimal("45000.00")
    assert solde.vnc == Decimal("55000.00")
    assert solde.annee == 2026
    assert solde.annee_source == 2025


@pytest.mark.asyncio
async def test_ouverture_idempotente_si_soldes_presents():
    db = MagicMock()
    db.flush = AsyncMock()
    svc = ExerciceOuvertureService(db)
    exo = MagicMock()
    exo.statut = StatutExercice.OUVERT
    exo.id = uuid4()

    solde = MagicMock()
    solde.valeur_brute_142 = Decimal("10")
    solde.cumul_148 = Decimal("4")
    solde.vnc = Decimal("6")

    found_exo = MagicMock()
    found_exo.scalar_one_or_none.return_value = exo
    empty_periodes = MagicMock()
    empty_periodes.scalars.return_value.all.return_value = []
    found_soldes = MagicMock()
    found_soldes.scalars.return_value.all.return_value = [solde]

    svc.situation = AsyncMock(
        return_value={
            "dernier_cloture": 2025,
            "exercice_ouvert": 2026,
            "annee_ouverture_proposee": 2026,
            "peut_ouvrir": False,
            "exercices": [],
        }
    )
    db.execute = AsyncMock(side_effect=[found_exo, empty_periodes, found_soldes])

    result = await svc.ouvrir_suivant()
    assert result.ouvertures_seed == 1
    assert "déjà ouvert" in result.message


@pytest.mark.asyncio
async def test_guard_refuse_exercice_cloture():
    from app.services.exercice_guard import ensure_exercice_ouvert

    db = MagicMock()
    found = MagicMock()
    found.scalar_one_or_none.return_value = StatutExercice.CLOTURE
    db.execute = AsyncMock(return_value=found)

    with pytest.raises(ValidationError, match="clôturé"):
        await ensure_exercice_ouvert(db, 2025, contexte="Test")

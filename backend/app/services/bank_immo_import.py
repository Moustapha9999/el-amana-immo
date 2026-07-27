"""Import du tableau d'amortissement banque (classeur multi-feuilles IMMO)."""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any
from uuid import UUID

from sqlalchemy import delete, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.exceptions import ValidationError
from app.models import Agence, Amortissement, CategorieImmobilisation, Immobilisation, PieceJointe
from app.models.enums import StatutImmobilisation
from app.services.immobilisation_defaults import apply_categorie_defaults, validate_immobilisation

# Feuilles de détail → code catégorie El Amana
SHEET_CATEGORY_MAP: dict[str, str] = {
    "aai": "TY-142010",
    "logiciel": "TY-147530",
    "frias immob": "TY-147050",
    "frias emmission emprt": "TY-147030",
    "mat bureaudv": "TY-142060",
    "matinfo": "TY-142041",
    "mat transp": "TY-142050",
    "matexhisto": "TY-142097",
    "autre immo": "TY-142080",
    "autres immo corp": "TY-142160",
}

SKIP_DESIGNATION_RE = re.compile(
    r"^(solde\b|s/t\b|total\b|sous[- ]?total\b|désignation|designation|date\b|"
    r"(?:report|rapport)s?\s+de\s+solde|"
    r"(?:report|rapport)s?\s+exercice|"
    r"(?:report|rapport)s?\s+d[eé]cembre|"
    r"(?:report|rapport)\s+de\s+l[' ]|"
    r"(?:report|rapport)s?\s+\d{1,2}[/.\-])",
    re.IGNORECASE,
)
OPENING_STOCK_RE = re.compile(r"^(s/t\b|solde\b|total\b)", re.IGNORECASE)
# Report historique initial (ex. REPORT 2003) — pas les reports/rapports de solde annuels
REPORT_HISTORIQUE_RE = re.compile(r"^(?:report|rapport)\s+(19|20)\d{2}\s*$", re.IGNORECASE)
REPORT_RE = re.compile(r"^(?:report|rapport)\b", re.IGNORECASE)
AGENCE_ALIASES: dict[str, str] = {
    "SIEGE CENTRAL": "00001",
    "SIÈGE CENTRAL": "00001",
    "AGENCE CENTRALE": "00001",
    "AGENCE CENTRALE PARTICULIERS": "00001",
}

# Stock Excel clôturé (ex. Solde 12/12/2025) :
#   2024-12 = colonne « Amt cumulés fin ex. préc. » (amt_n1)
#   2025-12 = colonne « Montant amt fin exercice » (amt_fin) + dotation Excel
PERIODE_STOCK_N1 = "2024-12"
PERIODE_OUVERTURE = "2025-12"
PERIODE_ARRETE = "2026-06"
# Feuilles sans colonnes Dotation / Fin exercice (ex. Logiciel 147530) :
# acquisitions < 2026 → 12 mois pleins, arrêté au 31/12/2026 ;
# acquisitions 2026 → prorata jusqu'au trimestre comptabilisé (30/06/2026).
PERIODE_ARRETE_ANNUELLE = "2026-12"
EXERCICE_CALCUL = 2026
# Année des colonnes Dotation / Fin du tableau Excel importé en stock_ouverture
EXERCICE_STOCK_EXCEL = int(PERIODE_OUVERTURE[:4])
# Arrêté prorata Excel DAYS360(..., 30/06/2026)
DATE_ARRETE_PRORATA = date(2026, 6, 30)
# Alias historique (évite de casser les imports de scripts)
DATE_ARRETE_PRORATA_EXCL = DATE_ARRETE_PRORATA



@dataclass
class ParsedBankRow:
    sheet: str
    categorie_code: str
    row_number: int
    date_acquisition: date
    quantite: int
    designation: str
    valeur_brute: Decimal
    taux: Decimal | None
    amt_n1: Decimal
    dotation: Decimal
    amt_fin: Decimal
    vnc: Decimal
    agence_label: str | None
    is_report: bool = False
    is_negative: bool = False
    # Feuille sans colonnes Dotation / Fin exercice : dotation à recalculer
    calc_dotation: bool = False


@dataclass
class BankImportTotaux:
    compte: str
    created: int = 0
    valeur_brute: Decimal = field(default_factory=lambda: Decimal("0.00"))
    amt_fin: Decimal = field(default_factory=lambda: Decimal("0.00"))
    vnc: Decimal = field(default_factory=lambda: Decimal("0.00"))


@dataclass
class BankImportResult:
    created: int
    amortissements_created: int
    errors: list[str]
    totaux_par_compte: list[dict[str, Any]]
    reports_created: int
    negatives: int


# Taux dominant → catégorie (feuilles renommées Feuil1 / Sheet1)
TAUX_CATEGORY_MAP: dict[str, str] = {
    "4": "TY-142020",  # Constructions
    "4.0": "TY-142020",
    "20": "TY-142041",  # Matériel informatique
    "20.0": "TY-142041",
    "25": "TY-142050",  # Transport
    "25.0": "TY-142050",
    "33": "TY-147050",  # Frais immobilisés
    "33.0": "TY-147050",
    "33.33": "TY-147050",
}

COMPTE_CATEGORY_MAP: dict[str, str] = {
    "142010": "TY-142010",
    "142020": "TY-142020",
    "142041": "TY-142041",
    "142050": "TY-142050",
    "142060": "TY-142060",
    "142080": "TY-142080",
    "142097": "TY-142097",
    "142160": "TY-142160",
    "147030": "TY-147030",
    "147050": "TY-147050",
    "147530": "TY-147530",
}


def resolve_categorie_code(sheet_name: str) -> str | None:
    key = sheet_name.strip().lower()
    key = re.sub(r"\s+", " ", key)
    if key in SHEET_CATEGORY_MAP:
        return SHEET_CATEGORY_MAP[key]
    if key.startswith("construct"):
        return "TY-142020"
    if key.startswith("matinfo"):
        return "TY-142041"
    if "frias" in key and "empr" in key:
        return "TY-147030"
    if "frias" in key:
        return "TY-147050"
    if key.startswith("autre") and "corp" in key:
        return "TY-142160"
    if key.startswith("autre"):
        return "TY-142080"
    return None


def _infer_categorie_from_grid(grid: list[list[Any]]) -> str | None:
    """Déduit la catégorie d'une feuille Feuil1 à partir du compte ou du taux dominant."""
    # Cherche un numéro de compte dans les 20 premières lignes / cellules texte
    for row in grid[:20]:
        for cell in row or []:
            if cell is None:
                continue
            text = str(cell)
            for compte, cat in COMPTE_CATEGORY_MAP.items():
                if compte in text:
                    return cat

    # Taux dominant sur les lignes d'actifs
    from collections import Counter

    counts: Counter[str] = Counter()
    header_indices = [
        i for i, row in enumerate(grid) if _detect_columns(list(row) if row else []) is not None
    ]
    if not header_indices:
        return None
    for h_idx, header_row in enumerate(header_indices):
        cols = _detect_columns(list(grid[header_row]))
        if cols is None:
            continue
        end = header_indices[h_idx + 1] if h_idx + 1 < len(header_indices) else len(grid)
        for r in range(header_row + 1, end):
            raw = list(grid[r]) if grid[r] else []
            des = str(_cell(raw, cols["designation"]) or "").strip()
            if not des or SKIP_DESIGNATION_RE.search(des):
                continue
            if REPORT_RE.search(des) and not REPORT_HISTORIQUE_RE.search(des):
                continue
            taux = parse_taux(_cell(raw, cols["taux"]))
            if taux is None:
                continue
            key = str(taux.normalize() if hasattr(taux, "normalize") else taux)
            # Normalise 4.00 → 4
            try:
                key = str(Decimal(key).normalize())
            except Exception:
                pass
            counts[key] += 1
    if not counts:
        return None
    dominant, _ = counts.most_common(1)[0]
    return TAUX_CATEGORY_MAP.get(dominant) or TAUX_CATEGORY_MAP.get(f"{float(dominant):g}")


def _q(value: Decimal) -> Decimal:
    return value.quantize(Decimal("0.01"))


def parse_amount(value: Any) -> Decimal | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return _q(Decimal(str(value)))
    s = str(value).strip().replace("\xa0", " ").replace(" ", "")
    if not s or s in {"-", "—", "–"}:
        return None
    s = s.replace(",", ".")
    s = re.sub(r"[^0-9.\-]", "", s)
    if not s or s in {".", "-", "-."}:
        return None
    try:
        return _q(Decimal(s))
    except (InvalidOperation, ValueError):
        return None


def parse_taux(value: Any) -> Decimal | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        d = Decimal(str(value))
        if Decimal("0") < d <= Decimal("1"):
            d = d * Decimal("100")
        return _q(d)
    s = str(value).strip().replace("%", "").replace(",", ".")
    if not s or s == "-":
        return None
    try:
        d = Decimal(s)
        if Decimal("0") < d <= Decimal("1"):
            d = d * Decimal("100")
        return _q(d)
    except (InvalidOperation, ValueError):
        return None


def parse_date(value: Any, *, datemode: int = 0) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        try:
            import xlrd

            dt = xlrd.xldate_as_datetime(float(value), datemode)
            return dt.date()
        except Exception:
            # Excel serial fallback (1900 system)
            try:
                base = date(1899, 12, 30)
                return base + timedelta(days=int(value))
            except Exception:
                return None
    s = str(value).strip()
    if not s:
        return None
    # Corrige typos type 01/01/20120
    m = re.match(r"^(\d{1,2})/(\d{1,2})/(\d{2,5})$", s)
    if m:
        day, month, year_s = int(m.group(1)), int(m.group(2)), m.group(3)
        year = int(year_s)
        if len(year_s) == 2:
            year = 2000 + year if year < 70 else 1900 + year
        elif len(year_s) == 3 or year > 2100:
            # 212 → 2012 ; 20120 → 2020
            yy = int(year_s[-2:])
            year = 2000 + yy if yy < 70 else 1900 + yy
        try:
            return date(year, month, day)
        except ValueError:
            return None
    try:
        return date.fromisoformat(s[:10])
    except ValueError:
        return None


def _norm_header(value: Any) -> str:
    s = str(value or "").strip().lower()
    s = s.replace("é", "e").replace("è", "e").replace("ê", "e").replace("à", "a")
    s = re.sub(r"\s+", " ", s)
    return s


def _detect_columns(header_cells: list[Any]) -> dict[str, int] | None:
    norms = [_norm_header(c) for c in header_cells]
    if not any(n == "date" or n.startswith("date ") for n in norms):
        return None

    def find(*preds: Any) -> int | None:
        for i, n in enumerate(norms):
            for pred in preds:
                if callable(pred) and pred(n):
                    return i
                if isinstance(pred, str) and pred in n:
                    return i
        return None

    col_date = find(lambda n: n == "date" or n.startswith("date "))
    col_des = find(lambda n: "design" in n)
    col_qte = find(lambda n: n.startswith("qt") or n == "qte" or "quant" in n)
    col_vb = find(
        lambda n: "acquisition" in n or n.startswith("valeur d"),
        "valeur",
    )
    col_taux = find(lambda n: "taux" in n)
    col_n1 = find(
        lambda n: (
            "prec" in n
            or "preced" in n
            or "debut" in n
            or "cumul" in n
            or "fin exr" in n  # ex. Fin Exr.2025
            or "exr.20" in n
        )
    )
    col_dot = find(lambda n: "dotation" in n or "en. c" in n or "en cours" in n)
    col_fin = find(
        lambda n: (
            "fin exercice" in n
            or "fin . ex" in n
            or "amrt fin" in n
            or "montant amt" in n
            or n.startswith("amt fin")
        )
        and "net" not in n
        and "comptable" not in n
    )
    col_vnc = find(lambda n: "net" in n or "comptable" in n)
    col_agence = find(lambda n: "agence" in n)

    if col_date is None or col_vb is None:
        return None
    # Feuilles sans en-tête Désignation (ex. frais emprunt) → col 2 par défaut
    if col_des is None:
        col_des = 2 if col_vb != 2 else 1

    # La feuille fournit-elle réellement les colonnes Dotation / Fin exercice ?
    # (avant application des positions par défaut ci-dessous)
    has_bank_dotation = col_dot is not None or col_fin is not None

    # Fallback positions if headers are incomplete
    if col_n1 is None and col_taux is not None:
        col_n1 = col_taux + 1
    if col_dot is None and col_n1 is not None:
        col_dot = col_n1 + 1
    if col_fin is None and col_dot is not None:
        col_fin = col_dot + 1
    if col_vnc is None and col_fin is not None:
        col_vnc = col_fin + 1

    return {
        "date": col_date,
        "qte": col_qte if col_qte is not None else 1,
        "designation": col_des,
        "vb": col_vb,
        "taux": col_taux if col_taux is not None else col_vb + 1,
        "n1": col_n1 if col_n1 is not None else col_vb + 2,
        "dotation": col_dot if col_dot is not None else col_vb + 3,
        "fin": col_fin if col_fin is not None else col_vb + 4,
        "vnc": col_vnc if col_vnc is not None else col_vb + 5,
        "agence": col_agence if col_agence is not None else -1,
        "has_dotation": 1 if has_bank_dotation else 0,
    }


def _cell(row: list[Any], idx: int) -> Any:
    if idx < 0 or idx >= len(row):
        return None
    return row[idx]


def _skip_row_label(raw: list[Any], cols: dict[str, Any]) -> str | None:
    """Libellé Solde / Total / Report à ignorer, même s'il est hors colonne Désignation.

    Sur certaines feuilles frais (ex. FRIAS emmission EMPRT), le report de solde
    a le texte dans la colonne Qté et une désignation vide — sinon il devient
    « Ligne sans libellé » et double-compte la VB.
    """
    for key in ("designation", "qte", "date"):
        val = _cell(raw, cols[key])
        if val is None:
            continue
        s = str(val).strip()
        if not s:
            continue
        if SKIP_DESIGNATION_RE.search(s):
            return s
        if REPORT_RE.search(s) and not REPORT_HISTORIQUE_RE.search(s):
            return s
    return None


def parse_bank_workbook(content: bytes, filename: str = "import.xls") -> list[ParsedBankRow]:
    """Parse le classeur banque (.xls ou .xlsx) en lignes exploitables."""
    name_l = (filename or "").lower()
    if name_l.endswith(".xlsx"):
        return _parse_xlsx(content)
    return _parse_xls(content)


def _resolve_sheet_category(sheet_name: str, grid: list[list[Any]]) -> str | None:
    cat = resolve_categorie_code(sheet_name)
    if cat is not None:
        return cat
    # Feuille générique (Feuil1, Sheet1…) avec structure tableau banque
    has_header = any(_detect_columns(list(row) if row else []) for row in grid[:25])
    if not has_header:
        return None
    return _infer_categorie_from_grid(grid)


def _parse_xls(content: bytes) -> list[ParsedBankRow]:
    try:
        import xlrd
    except ImportError as exc:
        raise ValidationError("Le module xlrd est requis pour lire les fichiers .xls") from exc

    book = xlrd.open_workbook(file_contents=content)
    rows: list[ParsedBankRow] = []
    for sheet in book.sheets():
        grid = [
            [sheet.cell_value(r, c) for c in range(sheet.ncols)]
            for r in range(sheet.nrows)
        ]
        cat = _resolve_sheet_category(sheet.name, grid)
        if cat is None:
            continue
        rows.extend(_parse_sheet_grid(sheet.name, cat, grid, datemode=book.datemode))
    return rows


def _parse_xlsx(content: bytes) -> list[ParsedBankRow]:
    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(content), read_only=True, data_only=True)
    rows: list[ParsedBankRow] = []
    for sheet in wb.worksheets:
        grid = [list(row) for row in sheet.iter_rows(values_only=True)]
        cat = _resolve_sheet_category(sheet.title, grid)
        if cat is None:
            continue
        rows.extend(_parse_sheet_grid(sheet.title, cat, grid, datemode=0))
    return rows


def _parse_sheet_grid(
    sheet_name: str,
    categorie_code: str,
    grid: list[list[Any]],
    *,
    datemode: int,
) -> list[ParsedBankRow]:
    # Plusieurs blocs d'en-tête possibles dans une même feuille
    header_indices: list[int] = []
    for i, row in enumerate(grid):
        cols = _detect_columns(list(row) if row else [])
        if cols is not None:
            header_indices.append(i)

    if not header_indices:
        return []

    out: list[ParsedBankRow] = []
    anon = 0
    # Dernière ligne "Solde ..." de la feuille = totaux de référence du tableau
    last_solde: tuple[Decimal, Decimal, Decimal | None] | None = None
    for h_idx, header_row in enumerate(header_indices):
        cols = _detect_columns(list(grid[header_row]))
        if cols is None:
            continue
        calc_mode = cols["has_dotation"] == 0
        end = header_indices[h_idx + 1] if h_idx + 1 < len(header_indices) else len(grid)
        # Un seul REPORT YYYY par bloc d'en-tête (ex. REPORT 2003).
        # MatexHisto a souvent 2 tableaux (comptes 1420970002 & 1420970015) :
        # chaque bloc a son propre stock d'ouverture — ne pas bloquer le 2e.
        # Dans un même bloc, les REPORT suivants (REPORT 2006…) restent ignorés.
        report_historique_pris = False
        block_has_rows = False
        for r in range(header_row + 1, end):
            raw = list(grid[r]) if grid[r] else []
            if not raw or all(c is None or str(c).strip() == "" for c in raw):
                continue

            des_raw = _cell(raw, cols["designation"])
            des = str(des_raw).strip() if des_raw is not None else ""
            vb = parse_amount(_cell(raw, cols["vb"]))

            if not des and vb is None:
                continue

            # Reports de solde / exercice / date → ignorer (évite le double comptage).
            # Exception : premier S/T / Solde / Total du bloc = stock d'ouverture
            # (ex. matinfo 31/12/05 S/T avant les acquisitions 2006).
            skip_label = _skip_row_label(raw, cols)
            opening_stock = False
            if skip_label:
                if (
                    not report_historique_pris
                    and not block_has_rows
                    and vb is not None
                    and OPENING_STOCK_RE.search(skip_label)
                ):
                    opening_stock = True
                    des = skip_label
                else:
                    des_l = skip_label.lower()
                    if (
                        des_l.startswith("solde")
                        or des_l.startswith("total")
                        or des_l.startswith("s/t")
                    ) and vb is not None:
                        s_n1 = parse_amount(_cell(raw, cols["n1"]))
                        s_fin = parse_amount(_cell(raw, cols["fin"]))
                        if s_n1 is not None and (s_fin is not None or calc_mode):
                            last_solde = (
                                _q(vb),
                                _q(s_n1),
                                _q(s_fin) if s_fin is not None else None,
                            )
                    continue

            is_report = bool(
                opening_stock or (des and REPORT_HISTORIQUE_RE.search(des))
            )
            if is_report:
                if report_historique_pris:
                    continue
                report_historique_pris = True
            # Total de fin d'année sans libellé (ex. 31/12/08 montant seul)
            if not des and vb is not None:
                if block_has_rows and parse_taux(_cell(raw, cols["taux"])) is None:
                    s_n1 = parse_amount(_cell(raw, cols["n1"]))
                    s_fin = parse_amount(_cell(raw, cols["fin"]))
                    if s_n1 is not None and (s_fin is not None or calc_mode):
                        last_solde = (
                            _q(vb),
                            _q(s_n1),
                            _q(s_fin) if s_fin is not None else None,
                        )
                    continue
                anon += 1
                des = f"Ligne sans libellé {anon}"
            if not des or vb is None:
                continue

            d_acq = parse_date(_cell(raw, cols["date"]), datemode=datemode)
            if d_acq is None:
                if is_report:
                    m_year = re.search(r"(19|20)\d{2}", des)
                    year = int(m_year.group(0)) if m_year else 2003
                    d_acq = date(year, 1, 1)
                elif des.startswith("Ligne sans libellé"):
                    d_acq = date(2006, 1, 1)
                else:
                    continue

            # REPORT 2003 daté 01/01/15 (saisie banque) : ancrer sur l'année du libellé
            # pour l'afficher en tête d'historique, pas au milieu de 2015.
            if is_report and not opening_stock:
                m_rep = re.search(r"(?:report|rapport)\s+((?:19|20)\d{2})\s*$", des, re.I)
                if m_rep:
                    report_year = int(m_rep.group(1))
                    if d_acq.year != report_year:
                        d_acq = date(report_year, 1, 1)

            if opening_stock:
                des = f"REPORT {d_acq.year}"

            qte_val = parse_amount(_cell(raw, cols["qte"]))
            quantite = int(qte_val) if qte_val is not None and qte_val > 0 else 1

            taux = parse_taux(_cell(raw, cols["taux"]))
            amt_n1 = parse_amount(_cell(raw, cols["n1"])) or Decimal("0.00")
            dotation = parse_amount(_cell(raw, cols["dotation"])) or Decimal("0.00")
            amt_fin = parse_amount(_cell(raw, cols["fin"]))
            if amt_fin is None:
                amt_fin = _q(amt_n1 + dotation)
            vnc = parse_amount(_cell(raw, cols["vnc"]))
            if vnc is None:
                vnc = _q(vb - amt_fin)

            agence_label = None
            if cols["agence"] >= 0:
                ag_raw = _cell(raw, cols["agence"])
                if ag_raw is not None and str(ag_raw).strip():
                    agence_label = str(ag_raw).strip()

            if is_report and not des.upper().startswith("REPORT"):
                des = f"Report historique — {des}"

            out.append(
                ParsedBankRow(
                    sheet=sheet_name,
                    categorie_code=categorie_code,
                    row_number=r + 1,
                    date_acquisition=d_acq,
                    quantite=max(1, quantite),
                    designation=des[:255],
                    valeur_brute=vb,
                    taux=taux,
                    amt_n1=_q(amt_n1),
                    dotation=_q(dotation),
                    amt_fin=_q(amt_fin),
                    vnc=_q(vnc),
                    agence_label=agence_label,
                    is_report=is_report,
                    is_negative=vb < 0,
                    calc_dotation=calc_mode,
                )
            )
            block_has_rows = True
    _reconcile_with_solde(out, last_solde)
    return out


def _reconcile_with_solde(
    rows: list[ParsedBankRow], solde: tuple[Decimal, Decimal, Decimal | None] | None
) -> None:
    """Absorbe les écarts d'arrondi entre les lignes et la ligne Solde du tableau.

    Les totaux Excel sont calculés en pleine précision alors que les cellules
    lignes sont arrondies au centime : quelques centimes d'écart peuvent
    apparaître. On les reporte sur la ligne au cumul le plus élevé afin que les
    totaux plateforme soient identiques au tableau banque.

    ``s_fin`` peut être absent (feuilles sans colonne Fin exercice, dotation
    recalculée) : on ne réconcilie alors que la VB et le cumul N-1.
    """
    if not rows or solde is None:
        return
    s_vb, s_n1, s_fin = solde
    tolerance = Decimal("1.00")
    zero = Decimal("0.00")
    d_vb = _q(sum((r.valeur_brute for r in rows), zero) - s_vb)
    d_n1 = _q(sum((r.amt_n1 for r in rows), zero) - s_n1)
    d_fin = _q(sum((r.amt_fin for r in rows), zero) - s_fin) if s_fin is not None else zero
    if d_vb == 0 and d_n1 == 0 and d_fin == 0:
        return
    # Écart trop important : probable problème de parsing, ne pas le masquer
    if abs(d_vb) > tolerance or abs(d_n1) > tolerance or abs(d_fin) > tolerance:
        return
    if d_vb:
        cible_vb = max(rows, key=lambda r: abs(r.valeur_brute))
        cible_vb.valeur_brute = _q(cible_vb.valeur_brute - d_vb)
        cible_vb.vnc = _q(cible_vb.valeur_brute - cible_vb.amt_fin)
    cible = max(rows, key=lambda r: abs(r.amt_fin))
    cible.amt_n1 = _q(cible.amt_n1 - d_n1)
    if s_fin is not None:
        cible.amt_fin = _q(cible.amt_fin - d_fin)
        cible.dotation = _q(cible.amt_fin - cible.amt_n1)
    else:
        # Dotation recalculée plus tard : le cumul fin suit le cumul N-1
        cible.amt_fin = _q(cible.amt_fin - d_n1)
    cible.vnc = _q(cible.valeur_brute - cible.amt_fin)


# Lignes de mouvement (pas une acquisition) : pas de nouvelle dotation calculée.
# Le cumul N-1 banque est figé ; VNC = VB − cumul N-1.
SKIP_DOTATION_DES_RE = re.compile(
    r"\b(nivellement|reclassement|annulation|annuation)\b",
    re.IGNORECASE,
)


def compute_dotation_exercice(
    vb: Decimal,
    amt_n1: Decimal,
    taux_pct: Decimal | None,
    date_acquisition: date,
    *,
    annee: int = EXERCICE_CALCUL,
    designation: str | None = None,
) -> Decimal:
    """Dotation de l'exercice ``annee`` pour une feuille sans colonne Dotation.

    Règles banque :
      • VB ≤ 0 (reclassement / annulation) → pas de dotation ;
      • nivellement / reclassement (libellé) → pas de dotation (transfert) ;
      • acquisition antérieure à ``annee`` → 12 mois pleins (VB × taux) ;
      • acquisition de ``annee`` → prorata 30/360 jusqu'au trimestre
        comptabilisé (arrêté du 30/06/2026) ;
      • plafonnée au restant à amortir (VB − cumul N-1).
    """
    if taux_pct is None or taux_pct == 0 or vb <= 0:
        return Decimal("0.00")
    if designation and SKIP_DOTATION_DES_RE.search(designation):
        return Decimal("0.00")
    from app.services.amortissement_engine import days_360

    taux = Decimal(taux_pct) / Decimal("100")
    if date_acquisition.year < annee:
        duree = Decimal("1")
    else:
        jours = days_360(max(date_acquisition, date(annee, 1, 1)), DATE_ARRETE_PRORATA)
        if jours <= 0:
            return Decimal("0.00")
        duree = Decimal(jours) / Decimal("360")

    dotation = _q(vb * taux * duree)
    restant = _q(vb - amt_n1)
    if restant <= 0:
        return Decimal("0.00")
    return min(dotation, restant)


def _prefix_from_category(code: str) -> str:
    from app.services.code_inventaire import prefix_from_categorie_code

    return prefix_from_categorie_code(code)


def _resolve_agence_id(
    label: str | None,
    agences_by_code: dict[str, UUID],
    agences_by_libelle: dict[str, UUID],
) -> UUID | None:
    if not label:
        return None
    key = label.strip().upper()
    code = AGENCE_ALIASES.get(key)
    if code and code in agences_by_code:
        return agences_by_code[code]
    if key in agences_by_code:
        return agences_by_code[key]
    lib = label.strip().upper()
    if lib in agences_by_libelle:
        return agences_by_libelle[lib]
    for libelle, aid in agences_by_libelle.items():
        if lib in libelle or libelle in lib:
            return aid
    return None


class BankImmoImportService:
    def __init__(self, db: AsyncSession):
        self.db = db

    def _import_banque_filter(self):
        """Biens issus de l'import tableau banque (metadata source)."""
        return (
            Immobilisation.deleted_at.is_(None),
            Immobilisation.metadata_json.contains({"source": "import_banque"}),
        )

    async def count_import_banque(self) -> int:
        result = await self.db.execute(
            select(Immobilisation.id).where(*self._import_banque_filter())
        )
        return len(list(result.scalars().all()))

    async def purge_import_banque(self) -> int:
        """Supprime définitivement les biens (et amortissements) de l'import banque.

        Suppression physique pour libérer les codes inventaire et permettre un ré-import.
        """
        result = await self.db.execute(
            select(Immobilisation.id).where(*self._import_banque_filter())
        )
        ids = list(result.scalars().all())
        if not ids:
            return 0

        from app.models import EcritureComptable, InventaireScan

        await self.db.execute(delete(Amortissement).where(Amortissement.immobilisation_id.in_(ids)))
        await self.db.execute(delete(PieceJointe).where(PieceJointe.immobilisation_id.in_(ids)))
        await self.db.execute(
            delete(EcritureComptable).where(EcritureComptable.immobilisation_id.in_(ids))
        )
        await self.db.execute(
            delete(InventaireScan).where(InventaireScan.immobilisation_id.in_(ids))
        )
        await self.db.execute(delete(Immobilisation).where(Immobilisation.id.in_(ids)))
        await self.db.flush()
        return len(ids)

    async def import_from_bytes(self, content: bytes, filename: str = "import.xls") -> BankImportResult:
        parsed = parse_bank_workbook(content, filename)
        if not parsed:
            raise ValidationError(
                "Aucune feuille de détail reconnue (aai, logiciel, matinfo, …)."
            )

        categories = {
            c.code: c
            for c in (
                await self.db.execute(
                    select(CategorieImmobilisation).where(CategorieImmobilisation.deleted_at.is_(None))
                )
            ).scalars().all()
        }
        agences = list(
            (await self.db.execute(select(Agence).where(Agence.deleted_at.is_(None)))).scalars().all()
        )
        agences_by_code = {a.code: a.id for a in agences}
        agences_by_libelle = {a.libelle.strip().upper(): a.id for a in agences}

        from app.services.code_inventaire import format_code_inventaire, load_max_sequences_map

        existing_codes = set(
            (
                await self.db.execute(select(Immobilisation.code_inventaire))
            ).scalars().all()
        )
        # Compteurs par (préfixe nature, année d'acquisition) → AAI-2026-001
        counters: dict[tuple[str, int], int] = await load_max_sequences_map(self.db)
        totaux: dict[str, BankImportTotaux] = {}
        errors: list[str] = []
        created = 0
        amort_created = 0
        reports_created = 0
        negatives = 0

        for row in parsed:
            categorie = categories.get(row.categorie_code)
            if categorie is None:
                errors.append(
                    f"{row.sheet} L{row.row_number}: catégorie {row.categorie_code} absente du référentiel"
                )
                continue

            prefix = _prefix_from_category(row.categorie_code)
            year = row.date_acquisition.year
            key = (prefix, year)
            counters[key] = counters.get(key, 0) + 1
            code = format_code_inventaire(prefix, year, counters[key])
            while code in existing_codes:
                counters[key] += 1
                code = format_code_inventaire(prefix, year, counters[key])
            existing_codes.add(code)

            agence_id = _resolve_agence_id(row.agence_label, agences_by_code, agences_by_libelle)
            from app.services.nature_immo_referentiel import normalize_taux_for_categorie

            taux_raw = row.taux if row.taux is not None else categorie.taux_lineaire_defaut
            taux = normalize_taux_for_categorie(categorie.code, taux_raw)

            if row.calc_dotation:
                # Feuille sans colonnes Dotation / Fin exercice (ex. Logiciel 147530) :
                # 12 mois pour les acquisitions < 2026, prorata jusqu'au trimestre
                # comptabilisé (30/06/2026) pour celles de 2026, puis cumul et VNC recalculés.
                row.dotation = compute_dotation_exercice(
                    row.valeur_brute,
                    row.amt_n1,
                    taux,
                    row.date_acquisition,
                    designation=row.designation,
                )
                row.amt_fin = _q(row.amt_n1 + row.dotation)
                row.vnc = _q(row.valeur_brute - row.amt_fin)

            # Tableau déjà à l'arrêté N (ex. 30/06/2026) s'il contient des acq. N.
            # Sinon : stock historique clôturé → amt_fin = ouverture N.
            arrete_courant = any(
                r.date_acquisition.year >= EXERCICE_CALCUL for r in parsed
            )
            seed_mode = "arrete_courant" if arrete_courant else "stock_ouverture"

            try:
                immo = Immobilisation(
                    code_inventaire=code,
                    designation=row.designation,
                    quantite=row.quantite,
                    categorie_id=categorie.id,
                    agence_id=agence_id,
                    date_acquisition=row.date_acquisition,
                    date_comptabilisation=row.date_acquisition,
                    date_mise_en_service=row.date_acquisition,
                    valeur_brute=row.valeur_brute,
                    valeur_residuelle=Decimal("0.00"),
                    taux=taux,
                    statut=StatutImmobilisation.EN_SERVICE,
                    metadata_json={
                        "source": "import_banque",
                        "fichier": filename,
                        "feuille": row.sheet,
                        "is_report": row.is_report,
                        "dotation_calculee": row.calc_dotation,
                        "bank": {
                            "amt_n1": str(row.amt_n1),
                            "dotation": str(row.dotation),
                            "amt_fin": str(row.amt_fin),
                            "vnc": str(row.vnc),
                            "seed_mode": seed_mode,
                        },
                    },
                )
                apply_categorie_defaults(
                    immo,
                    categorie,
                    override_comptes=True,
                    preserve_taux=taux is not None,
                )
                if taux is not None:
                    immo.taux = taux
                # Validation assouplie pour VB négative (reclassements banque)
                if row.valeur_brute >= 0:
                    validate_immobilisation(immo, categorie)
                else:
                    if immo.date_comptabilisation is None:
                        immo.date_comptabilisation = immo.date_acquisition
                    immo.periodicite = "trimestriel"
                    negatives += 1

                immo.qr_code_data = f"IMMO:{immo.code_inventaire}"
                immo.barcode_data = immo.code_inventaire
                self.db.add(immo)
                await self.db.flush()

                # Seed amortissements banque.
                # stock_ouverture : conserver l'exercice Excel (N-1 + dotation → fin),
                #   puis ouverture N+1 = amt_fin (PERIODE_OUVERTURE).
                # arrete_courant : cumul ouverture = amt_n1, + période d'arrêté N.
                if seed_mode == "stock_ouverture":
                    if row.calc_dotation:
                        dotation_excel = _q(row.amt_fin - row.amt_n1)
                    else:
                        dotation_excel = _q(row.dotation)
                    self.db.add(
                        Amortissement(
                            immobilisation_id=immo.id,
                            periode=PERIODE_STOCK_N1,
                            montant=Decimal("0.00"),
                            cumul=row.amt_n1,
                            vnc=_q(row.valeur_brute - row.amt_n1),
                            valide=True,
                            annule=False,
                            simule=False,
                        )
                    )
                    amort_created += 1
                    self.db.add(
                        Amortissement(
                            immobilisation_id=immo.id,
                            periode=PERIODE_OUVERTURE,
                            montant=dotation_excel,
                            cumul=row.amt_fin,
                            vnc=row.vnc,
                            valide=True,
                            annule=False,
                            simule=False,
                        )
                    )
                    amort_created += 1
                else:
                    self.db.add(
                        Amortissement(
                            immobilisation_id=immo.id,
                            periode=PERIODE_OUVERTURE,
                            montant=Decimal("0.00"),
                            cumul=row.amt_n1,
                            vnc=_q(row.valeur_brute - row.amt_n1),
                            valide=True,
                            annule=False,
                            simule=False,
                        )
                    )
                    amort_created += 1

                    # Période de la dotation : 31/12/2026 (12 mois) si feuille sans
                    # colonnes Dotation et acquisition < 2026 ; sinon arrêté 30/06/2026.
                    if row.calc_dotation and row.date_acquisition.year < EXERCICE_CALCUL:
                        periode_arrete = PERIODE_ARRETE_ANNUELLE
                    else:
                        periode_arrete = PERIODE_ARRETE

                    # Conserver la colonne « Exer. En C » pour le Total 68.
                    if row.calc_dotation:
                        dotation_arrete = _q(row.amt_fin - row.amt_n1)
                    else:
                        dotation_arrete = _q(row.dotation)
                    if dotation_arrete != 0 or row.dotation != 0:
                        self.db.add(
                            Amortissement(
                                immobilisation_id=immo.id,
                                periode=periode_arrete,
                                montant=dotation_arrete,
                                cumul=row.amt_fin,
                                vnc=row.vnc,
                                valide=True,
                                annule=False,
                                simule=False,
                            )
                        )
                        amort_created += 1

                created += 1
                if row.is_report:
                    reports_created += 1

                compte = (immo.compte_immobilisation or categorie.compte_immobilisation or "").strip()
                bucket = totaux.setdefault(compte, BankImportTotaux(compte=compte))
                bucket.created += 1
                bucket.valeur_brute = _q(bucket.valeur_brute + row.valeur_brute)
                bucket.amt_fin = _q(bucket.amt_fin + row.amt_fin)
                bucket.vnc = _q(bucket.vnc + row.vnc)
            except Exception as exc:
                errors.append(f"{row.sheet} L{row.row_number}: {exc}")

        return BankImportResult(
            created=created,
            amortissements_created=amort_created,
            errors=errors,
            totaux_par_compte=[
                {
                    "compte": t.compte,
                    "created": t.created,
                    "valeur_brute": float(t.valeur_brute),
                    "amt_fin": float(t.amt_fin),
                    "vnc": float(t.vnc),
                }
                for t in sorted(totaux.values(), key=lambda x: x.compte)
            ],
            reports_created=reports_created,
            negatives=negatives,
        )


async def repair_stock_ouverture_amorts(db: AsyncSession) -> dict[str, int]:
    """Répare les imports stock : 2024-12 = amt_n1, 2025-12 = amt_fin + dotation Excel.

    Ancien seed : uniquement 2025-12 avec cumul=amt_fin et montant=0 → Comptes 2025
    affichait Amt N-1 = amt_fin (ex. 9 000 au lieu de 8 325, écart 675).
    """
    result = await db.execute(
        select(Immobilisation).where(
            Immobilisation.deleted_at.is_(None),
            Immobilisation.metadata_json.contains({"source": "import_banque"}),
        )
    )
    immos = list(result.scalars().all())
    inserted = 0
    updated = 0
    skipped = 0

    for immo in immos:
        meta = immo.metadata_json if isinstance(immo.metadata_json, dict) else {}
        bank = meta.get("bank") if isinstance(meta, dict) else None
        if not isinstance(bank, dict):
            skipped += 1
            continue
        if str(bank.get("seed_mode") or "stock_ouverture") == "arrete_courant":
            skipped += 1
            continue
        try:
            amt_n1 = _q(Decimal(str(bank.get("amt_n1") or "0")))
            amt_fin = _q(Decimal(str(bank.get("amt_fin") or "0")))
            if bank.get("dotation") is not None:
                dotation = _q(Decimal(str(bank.get("dotation"))))
            else:
                dotation = _q(amt_fin - amt_n1)
            vb = _q(Decimal(immo.valeur_brute or 0))
            vnc_n1 = _q(vb - amt_n1)
            vnc_fin = _q(vb - amt_fin)
            if bank.get("vnc") is not None:
                vnc_fin = _q(Decimal(str(bank.get("vnc"))))
        except Exception:  # noqa: BLE001
            skipped += 1
            continue

        am_rows = (
            await db.execute(
                select(Amortissement).where(
                    Amortissement.immobilisation_id == immo.id,
                    Amortissement.annule.is_(False),
                )
            )
        ).scalars().all()
        by_per = {str(a.periode): a for a in am_rows}

        # Ne pas toucher un import déjà en mode arrêté (période 2026-06 / 2026-12 seed).
        if PERIODE_ARRETE in by_per or PERIODE_ARRETE_ANNUELLE in by_per:
            skipped += 1
            continue

        n1_row = by_per.get(PERIODE_STOCK_N1)
        if n1_row is None:
            db.add(
                Amortissement(
                    immobilisation_id=immo.id,
                    periode=PERIODE_STOCK_N1,
                    montant=Decimal("0.00"),
                    cumul=amt_n1,
                    vnc=vnc_n1,
                    valide=True,
                    annule=False,
                    simule=False,
                )
            )
            inserted += 1
        else:
            if n1_row.cumul != amt_n1 or n1_row.montant != 0:
                n1_row.montant = Decimal("0.00")
                n1_row.cumul = amt_n1
                n1_row.vnc = vnc_n1
                n1_row.valide = True
                updated += 1

        ouv = by_per.get(PERIODE_OUVERTURE)
        if ouv is None:
            db.add(
                Amortissement(
                    immobilisation_id=immo.id,
                    periode=PERIODE_OUVERTURE,
                    montant=dotation,
                    cumul=amt_fin,
                    vnc=vnc_fin,
                    valide=True,
                    annule=False,
                    simule=False,
                )
            )
            inserted += 1
        else:
            if ouv.cumul != amt_fin or ouv.montant != dotation:
                ouv.montant = dotation
                ouv.cumul = amt_fin
                ouv.vnc = vnc_fin
                ouv.valide = True
                updated += 1

        # Persister seed_mode pour les prochains traitements
        if not bank.get("seed_mode"):
            bank["seed_mode"] = "stock_ouverture"
            meta["bank"] = bank
            immo.metadata_json = dict(meta)

    await db.flush()
    return {"inserted": inserted, "updated": updated, "skipped": skipped, "immos": len(immos)}

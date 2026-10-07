"""Formation & Sensibilisation — import de l'historique Excel (remplace le suivi Excel).

Deux temps, rien n'est écrit avant confirmation :
1. ``analyser`` : détection de la feuille / ligne d'en-tête / colonnes, lecture des lignes,
   éclatement des cellules multi-valeurs (« LBC FT/Procédure… », « A / B »), rapprochement
   avec les référentiels et les employés existants, regroupement en sessions, anomalies.
   L'analyse est conservée dans ``formation_imports.analyse`` (traçabilité).
2. ``confirmer`` : l'utilisateur a validé les correspondances, les créations de valeurs et le
   découpage Nom / Prénom (jamais déduit automatiquement de façon « fiable » : il est proposé,
   marqué « à vérifier » et doit être confirmé). Sessions historiques créées CLÔTURÉES.
"""

from __future__ import annotations

import hashlib
import re
import uuid
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta
from difflib import SequenceMatcher
from io import BytesIO
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.core.exceptions import AppError
from app.models import (
    FormationEmploye,
    FormationEntite,
    FormationImport,
    FormationParticipant,
    FormationReferentiel,
    FormationSession,
    FormationSessionFormateur,
    FormationSessionTheme,
    User,
)
from app.services.formation_service import (
    Ctx,
    FormationService,
    Refs,
    cle,
    cle_identite,
    conflit,
    maintenant,
    propre,
)

TAILLE_MAX = 10 * 1024 * 1024
DOMAINES_IMPORT = ("THEME", "LIEU", "FORMATEUR", "FONCTION", "ENTITE")

SYNONYMES: dict[str, list[str]] = {
    "date": ["date", "date formation", "date de formation", "date session", "date de la formation"],
    "theme": ["theme", "themes", "sujet", "intitule formation", "theme formation", "formation"],
    "lieu": ["lieu", "ville", "localite", "lieu de formation"],
    "formateur": ["formateur", "formateurs", "animateur", "intervenant", "formateur s"],
    "participant": ["nom du participant", "participant", "nom et prenom", "nom prenom", "nom complet",
                    "participants", "employe", "collaborateur", "nom des participants", "prenom et nom"],
    "nom": ["nom", "nom de famille"],
    "prenom": ["prenom", "prenoms"],
    "fonction": ["fonction", "poste", "emploi", "fonction du participant"],
    "entite": ["entite", "agence", "service", "structure", "entite agence"],
    "perimetre": ["perimetre"],
    "presence": ["presence", "present", "statut presence", "emargement", "present absent"],
}
LIBELLES_CHAMPS = {
    "date": "Date", "theme": "Thème", "lieu": "Lieu", "formateur": "Formateur", "participant": "Nom du participant",
    "nom": "Nom", "prenom": "Prénom", "fonction": "Fonction", "entite": "Entité", "perimetre": "Périmètre",
    "presence": "Présence",
}
VALEURS_VIDES = {"", "0", "neant", "xxx", "ref", "n a", "na", "nd", "aucun", "aucune", "valeur"}
SEP_THEMES = re.compile(r"\s*[,;/+\n]\s*")
SEP_FORMATEURS = re.compile(r"\s*(?:[,;/&+\n]|\bet\b)\s*", re.I)


def _texte(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return propre(str(v))


def _vide(v: str) -> bool:
    c = cle(v)
    return not c or c in VALEURS_VIDES or str(v).strip().startswith("#")


def _date(v: Any) -> date | None:
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, (int, float)) and 20000 < float(v) < 80000:
        try:
            return from_excel(v).date()
        except Exception:
            return None
    s = str(v).strip()
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%d/%m/%y", "%d.%m.%Y", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def _presence(v: Any) -> tuple[str | None, bool]:
    """(valeur, reconnue)."""
    c = cle(_texte(v))
    if not c:
        return None, True
    if c in ("present", "presente", "p", "oui", "o", "x", "1", "yes", "ok"):
        return "PRESENT", True
    if c in ("absent", "absente", "a", "non", "n", "0", "no"):
        return "ABSENT", True
    return None, False


def _detecter_colonnes(entetes: list[str]) -> dict[str, int]:
    colonnes: dict[str, int] = {}
    cles = [cle(h) for h in entetes]
    for champ, syns in SYNONYMES.items():
        for i, c in enumerate(cles):
            if c and c in syns and i not in colonnes.values():
                colonnes[champ] = i
                break
    for champ, syns in SYNONYMES.items():
        if champ in colonnes:
            continue
        for i, c in enumerate(cles):
            if c and i not in colonnes.values() and any(c.startswith(s) for s in syns if len(s) > 4):
                colonnes[champ] = i
                break
    # « Nom » seul + « Prénom » : pas de colonne nom complet.
    if "participant" in colonnes and "nom" in colonnes and colonnes["nom"] == colonnes["participant"]:
        colonnes.pop("nom")
    return colonnes


def _score(c: str, ref_cle: str) -> float:
    if c == ref_cle:
        return 1.0
    a, b = set(c.split()), set(ref_cle.split())
    if b and b <= a:
        return 0.86 + 0.1 * len(b) / max(len(a), 1)
    if a and a <= b:
        return 0.84 + 0.1 * len(a) / max(len(b), 1)
    return SequenceMatcher(None, c, ref_cle).ratio()


def _proposer(c: str, candidats: list[dict]) -> dict:
    meilleur, score = None, 0.0
    for r in candidats:
        s = _score(c, r["cle"])
        if s > score or (s == score and meilleur and r["actif"] and not meilleur["actif"]):
            meilleur, score = r, s
    if meilleur and score >= 0.82:
        return {"action": "MAPPER", "cible_id": meilleur["id"], "cible": meilleur["libelle"],
                "score": round(score, 2), "a_verifier": score < 1.0}
    return {"action": "CREER", "cible_id": None, "cible": None, "score": round(score, 2), "a_verifier": True,
            "suggestion": meilleur["libelle"] if meilleur and score >= 0.6 else None,
            "suggestion_id": meilleur["id"] if meilleur and score >= 0.6 else None}


class FormationImportService:
    def __init__(self, db: AsyncSession, ctx: Ctx):
        self.db = db
        self.ctx = ctx
        self.svc = FormationService(db, ctx)

    # ---------------------------------------------------------------- analyse
    async def analyser(self, contenu: bytes, nom_fichier: str) -> dict:
        self.ctx.exiger("formation.import.execute")
        if not contenu:
            raise AppError("Fichier vide", 422, code="FICHIER_VIDE")
        if len(contenu) > TAILLE_MAX:
            raise AppError("Fichier trop volumineux (10 Mo maximum)", 413, code="FICHIER_TROP_VOLUMINEUX")
        if not nom_fichier.lower().endswith((".xlsx", ".xlsm")):
            raise AppError("Format non pris en charge : enregistrez le fichier au format .xlsx", 422,
                           code="FORMAT_NON_SUPPORTE")
        try:
            wb = load_workbook(BytesIO(contenu), data_only=True, read_only=True)
        except Exception as exc:
            raise AppError("Fichier Excel illisible ou corrompu", 422, code="FICHIER_ILLISIBLE") from exc

        meilleure = None
        for ws in wb.worksheets:
            lignes = []
            for row in ws.iter_rows(values_only=True):
                lignes.append(list(row))
                if len(lignes) >= 5000:
                    break
            for idx, row in enumerate(lignes[:25]):
                entetes = [_texte(v) for v in row]
                cols = _detecter_colonnes(entetes)
                score = len(cols) + (3 if ("participant" in cols or "nom" in cols) else 0) + (2 if "date" in cols else 0)
                if meilleure is None or score > meilleure[0]:
                    meilleure = (score, ws.title, idx, entetes, cols, lignes)
        feuilles = [ws.title for ws in wb.worksheets]
        wb.close()
        if not meilleure or meilleure[0] < 6:
            raise AppError("Colonnes non reconnues : le fichier doit contenir au moins Date, Thème, Lieu et "
                           "Nom du participant", 422, code="COLONNES_NON_RECONNUES")
        _score_, feuille, idx_entete, entetes, cols, lignes = meilleure
        manquantes = [LIBELLES_CHAMPS[c] for c in ("date", "theme", "lieu") if c not in cols]
        if "participant" not in cols and "nom" not in cols:
            manquantes.append("Nom du participant")
        if manquantes:
            raise AppError(f"Colonnes obligatoires absentes : {', '.join(manquantes)}", 422,
                           code="COLONNES_MANQUANTES")

        refs = await Refs.charger(self.db)
        par_domaine: dict[str, list[dict]] = defaultdict(list)
        for r in refs.refs.values():
            par_domaine[r.domaine].append({"id": str(r.id), "libelle": r.libelle, "cle": r.cle, "actif": r.actif})
        entites = [{"id": str(e.id), "libelle": e.libelle, "cle": e.cle, "actif": e.actif,
                    "perimetre_id": str(e.perimetre_id)} for e in refs.entites.values()]
        perimetres = {r["cle"]: r for r in par_domaine["PERIMETRE"]}

        def val(row: list, champ: str) -> Any:
            i = cols.get(champ)
            return row[i] if i is not None and i < len(row) else None

        anomalies: list[dict] = []
        parsed: list[dict] = []
        lignes_formules = 0
        today = date.today()
        for n, row in enumerate(lignes[idx_entete + 1:], start=idx_entete + 2):
            if not any(_texte(v) for v in row):
                continue
            participant = _texte(val(row, "participant"))
            nom_col, prenom_col = _texte(val(row, "nom")), _texte(val(row, "prenom"))
            if not participant and (nom_col or prenom_col):
                participant = propre(f"{prenom_col} {nom_col}")
            d = _date(val(row, "date"))
            themes = [t for t in SEP_THEMES.split(_texte(val(row, "theme"))) if not _vide(t)]
            lieu = _texte(val(row, "lieu"))
            if not (participant or d or themes or _texte(val(row, "date")) or not _vide(lieu)):
                lignes_formules += 1
                continue
            formateurs = [f for f in SEP_FORMATEURS.split(_texte(val(row, "formateur"))) if not _vide(f)]
            presence, presence_ok = _presence(val(row, "presence")) if "presence" in cols else ("PRESENT", True)
            ligne = {
                "ligne": n, "date": d.isoformat() if d else None, "date_brut": _texte(val(row, "date")),
                "themes": themes, "lieu": "" if _vide(lieu) else lieu, "formateurs": formateurs,
                "participant": participant, "nom": nom_col or None, "prenom": prenom_col or None,
                "fonction": "" if _vide(_texte(val(row, "fonction"))) else _texte(val(row, "fonction")),
                "entite": "" if _vide(_texte(val(row, "entite"))) else _texte(val(row, "entite")),
                "perimetre": "" if _vide(_texte(val(row, "perimetre"))) else _texte(val(row, "perimetre")),
                "perimetre_brut": _texte(val(row, "perimetre")),
                "presence": presence if d and d <= today else None,
                "bloquante": False,
            }
            erreurs = []
            if not participant:
                erreurs.append("Nom du participant manquant")
            if not d:
                erreurs.append(f"Date absente ou invalide ({ligne['date_brut'] or 'vide'})")
            if not themes:
                erreurs.append("Thème manquant")
            if not ligne["lieu"]:
                erreurs.append("Lieu manquant")
            for e in erreurs:
                anomalies.append({"ligne": n, "niveau": "bloquant", "message": e})
            ligne["bloquante"] = bool(erreurs)
            if d and (d.year < 2000 or d > today + timedelta(days=730)):
                anomalies.append({"ligne": n, "niveau": "alerte", "message": f"Date inhabituelle : {d:%d/%m/%Y}"})
            if not formateurs and not erreurs:
                anomalies.append({"ligne": n, "niveau": "alerte", "message": "Formateur non renseigné"})
            if ligne["perimetre_brut"] and not ligne["perimetre"]:
                anomalies.append({"ligne": n, "niveau": "alerte",
                                  "message": f"Périmètre invalide « {ligne['perimetre_brut']} » (déduit de l'entité)"})
            if _texte(val(row, "entite")) and not ligne["entite"]:
                anomalies.append({"ligne": n, "niveau": "alerte",
                                  "message": f"Entité fictive « {_texte(val(row, 'entite'))} » ignorée"})
            if not presence_ok:
                anomalies.append({"ligne": n, "niveau": "alerte",
                                  "message": f"Présence non reconnue « {_texte(val(row, 'presence'))} » : laissée non saisie"})
            parsed.append(ligne)

        valides = [l for l in parsed if not l["bloquante"]]
        # Valeurs distinctes à rapprocher.
        occur: dict[str, Counter] = {d: Counter() for d in DOMAINES_IMPORT}
        brut: dict[str, dict[str, str]] = {d: {} for d in DOMAINES_IMPORT}
        per_entite: dict[str, Counter] = defaultdict(Counter)
        for l in valides:
            for dom, valeurs in (("THEME", l["themes"]), ("LIEU", [l["lieu"]]), ("FORMATEUR", l["formateurs"]),
                                 ("FONCTION", [l["fonction"]] if l["fonction"] else []),
                                 ("ENTITE", [l["entite"]] if l["entite"] else [])):
                for v in valeurs:
                    c = cle(v)
                    occur[dom][c] += 1
                    brut[dom].setdefault(c, propre(v))
            if l["entite"] and l["perimetre"]:
                per_entite[cle(l["entite"])][cle(l["perimetre"])] += 1
        valeurs: dict[str, list[dict]] = {}
        for dom in DOMAINES_IMPORT:
            cands = entites if dom == "ENTITE" else par_domaine[dom]
            items = []
            for c, nb in occur[dom].most_common():
                prop = _proposer(c, cands)
                item = {"cle": c, "brut": brut[dom][c], "occurrences": nb, "libelle": brut[dom][c], **prop}
                if dom == "ENTITE" and prop["action"] == "CREER":
                    per = per_entite.get(c)
                    p_cle = per.most_common(1)[0][0] if per else None
                    p_ref = perimetres.get(p_cle) if p_cle else None
                    if not p_ref and p_cle:
                        p_prop = _proposer(p_cle, par_domaine["PERIMETRE"])
                        p_ref = {"id": p_prop["cible_id"], "libelle": p_prop["cible"]} if p_prop["cible_id"] else None
                    if not p_ref:
                        p_ref = perimetres.get(c)
                        if not p_ref:
                            p_prop = _proposer(c, par_domaine["PERIMETRE"])
                            p_ref = {"id": p_prop["cible_id"], "libelle": p_prop["cible"]} if p_prop["cible_id"] else None
                    item["perimetre_id"] = p_ref["id"] if p_ref else None
                    item["perimetre"] = p_ref["libelle"] if p_ref else None
                    item["perimetre_brut"] = p_cle
                    if not p_ref:
                        anomalies.append({"ligne": None, "niveau": "alerte",
                                          "message": f"Entité nouvelle « {item['brut']} » : périmètre à choisir"})
                items.append(item)
            valeurs[dom] = items

        # Personnes.
        groupes: dict[str, list[dict]] = defaultdict(list)
        for l in valides:
            groupes[cle_identite(l["participant"])].append(l)
        existants = {}
        if groupes:
            for e in (await self.db.scalars(select(FormationEmploye).where(
                    FormationEmploye.cle_identite.in_(list(groupes))))).unique().all():
                existants.setdefault(e.cle_identite, e)
        personnes = []
        for ci, ls in groupes.items():
            variantes = Counter(l["participant"] for l in ls)
            principal = variantes.most_common(1)[0][0]
            recent = max(ls, key=lambda l: l["date"] or "")
            e = existants.get(ci)
            p = {
                "cle": ci, "brut": principal, "variantes": sorted(variantes), "occurrences": len(ls),
                "mots": principal.split(), "fonction": recent["fonction"], "entite": recent["entite"],
                "nom_colonne": recent.get("nom"), "prenom_colonne": recent.get("prenom"),
                "action": "EXISTANT" if e else "NOUVEAU",
                "employe_id": str(e.id) if e else None,
                "employe": (f"{e.prenom} {e.nom}" if e.prenom else e.nom) if e else None,
                "candidats": [],
            }
            if not e:
                proches = await self.svc.doublons(principal)
                p["candidats"] = [{"id": c["id"], "nom_complet": c["nom_complet"], "entite": c["entite"]}
                                  for c in proches[:5]]
                if p["candidats"]:
                    anomalies.append({"ligne": None, "niveau": "alerte",
                                      "message": f"« {principal} » ressemble à un employé existant : à vérifier"})
            if len(variantes) > 1:
                anomalies.append({"ligne": None, "niveau": "info",
                                  "message": f"Orthographes multiples regroupées : {', '.join(sorted(variantes))}"})
            personnes.append(p)
        personnes.sort(key=lambda p: p["brut"].lower())

        # Sessions.
        sessions: dict[str, dict] = {}
        for l in valides:
            k = self._cle_session(l)
            s = sessions.setdefault(k, {"cle": k, "date": l["date"], "themes": l["themes"], "lieu": l["lieu"],
                                        "formateurs": l["formateurs"], "participants": [], "lignes": []})
            ci = cle_identite(l["participant"])
            if ci in s["participants"]:
                anomalies.append({"ligne": l["ligne"], "niveau": "info",
                                  "message": f"Doublon : « {l['participant']} » déjà présent dans cette session"})
                continue
            s["participants"].append(ci)
            s["lignes"].append(l["ligne"])
        liste_sessions = sorted(sessions.values(), key=lambda s: s["date"] or "")
        for s in liste_sessions:
            s["nb_participants"] = len(s["participants"])

        sha = hashlib.sha256(contenu).hexdigest()
        precedent = await self.db.scalar(select(FormationImport).where(
            FormationImport.fichier_sha256 == sha, FormationImport.statut == "IMPORTE"
        ).order_by(FormationImport.importe_le.desc()))
        analyse = {
            "fichier": nom_fichier, "feuille": feuille, "feuilles": feuilles, "ligne_entete": idx_entete + 1,
            "colonnes": [{"champ": c, "libelle": LIBELLES_CHAMPS[c], "index": i, "entete": entetes[i]}
                         for c, i in sorted(cols.items(), key=lambda x: x[1])],
            "colonnes_ignorees": [h for i, h in enumerate(entetes) if h and i not in cols.values()],
            "presence_colonne": "presence" in cols,
            "lignes": parsed, "valeurs": valeurs, "personnes": personnes, "sessions": liste_sessions,
            "anomalies": anomalies,
            "stats": {
                "lignes_total": len(parsed), "lignes_valides": len(valides), "lignes_vides": lignes_formules,
                "lignes_ignorees": len(parsed) - len(valides), "sessions": len(liste_sessions),
                "personnes": len(personnes), "employes_existants": sum(1 for p in personnes if p["action"] == "EXISTANT"),
                "nouveaux_employes": sum(1 for p in personnes if p["action"] == "NOUVEAU"),
                "valeurs_a_creer": sum(1 for d in valeurs.values() for v in d if v["action"] == "CREER"),
                "bloquants": sum(1 for a in anomalies if a["niveau"] == "bloquant"),
                "alertes": sum(1 for a in anomalies if a["niveau"] == "alerte"),
            },
            "deja_importe": {"le": precedent.importe_le.isoformat() if precedent and precedent.importe_le else None,
                             "id": str(precedent.id)} if precedent else None,
        }
        imp = FormationImport(fichier_nom=nom_fichier[:255], fichier_sha256=sha, statut="ANALYSE",
                              nb_lignes=len(parsed), analyse=analyse, created_by_id=self.ctx.user.id)
        self.db.add(imp)
        await self.db.flush()
        await self.svc.audit("formation.import.analyse", "formation_import", imp.id,
                             after={"fichier": nom_fichier, **analyse["stats"]})
        return self._import_dict(imp, avec_analyse=True)

    @staticmethod
    def _cle_session(l: dict) -> str:
        return "|".join([
            l["date"] or "", ",".join(sorted(cle(t) for t in l["themes"])), cle(l["lieu"]),
            ",".join(sorted(cle(f) for f in l["formateurs"])),
        ])

    def _import_dict(self, imp: FormationImport, *, avec_analyse: bool = False, auteur: str | None = None) -> dict:
        out = {
            "id": str(imp.id), "fichier_nom": imp.fichier_nom, "statut": imp.statut, "nb_lignes": imp.nb_lignes,
            "created_at": imp.created_at.isoformat() if imp.created_at else None,
            "importe_le": imp.importe_le.isoformat() if imp.importe_le else None,
            "stats": (imp.analyse or {}).get("stats"), "resultat": imp.resultat, "auteur": auteur,
        }
        if avec_analyse:
            out["analyse"] = imp.analyse
        return out

    async def lister(self) -> list[dict]:
        rows = (await self.db.execute(
            select(FormationImport, User.full_name).outerjoin(User, User.id == FormationImport.created_by_id)
            .order_by(FormationImport.created_at.desc()).limit(100))).all()
        return [self._import_dict(i, auteur=n) for i, n in rows]

    async def detail(self, iid: uuid.UUID) -> dict:
        imp = await self.db.get(FormationImport, iid)
        if not imp:
            raise AppError("Import introuvable", 404, code="NOT_FOUND")
        return self._import_dict(imp, avec_analyse=True)

    async def abandonner(self, iid: uuid.UUID) -> dict:
        self.ctx.exiger("formation.import.execute")
        imp = await self.db.get(FormationImport, iid)
        if not imp:
            raise AppError("Import introuvable", 404, code="NOT_FOUND")
        if imp.statut != "ANALYSE":
            raise conflit("Cet import n'est plus en attente", "IMPORT_TERMINE")
        imp.statut = "ABANDONNE"
        await self.db.flush()
        await self.svc.audit("formation.import.abandon", "formation_import", imp.id, after={"fichier": imp.fichier_nom})
        return self._import_dict(imp)

    # ------------------------------------------------------------- confirmation
    async def confirmer(self, iid: uuid.UUID, payload: dict) -> dict:
        self.ctx.exiger("formation.import.execute")
        imp = await self.db.scalar(select(FormationImport).where(FormationImport.id == iid).with_for_update())
        if not imp:
            raise AppError("Import introuvable", 404, code="NOT_FOUND")
        if imp.statut != "ANALYSE":
            raise conflit("Cet import a déjà été confirmé ou abandonné", "IMPORT_TERMINE")
        a = imp.analyse or {}
        if a.get("deja_importe") and not payload.get("forcer"):
            raise conflit("Ce fichier a déjà été importé. Confirmez pour l'importer de nouveau "
                          "(les participations existantes ne sont pas dupliquées).", "FICHIER_DEJA_IMPORTE")
        choix_personnes: dict[str, dict] = payload.get("personnes") or {}
        nouveaux = [p for p in a.get("personnes", [])
                    if (choix_personnes.get(p["cle"]) or {}).get("action", p["action"]) == "NOUVEAU"]
        if nouveaux and not payload.get("verification_noms"):
            raise AppError("Vérifiez le découpage Nom / Prénom des nouveaux employés puis cochez la confirmation",
                           422, code="VERIFICATION_NOMS_REQUISE")

        user_id = self.ctx.user.id
        stats = Counter()
        # 1. Valeurs de référentiel.
        mapping: dict[str, dict[str, uuid.UUID | None]] = {d: {} for d in DOMAINES_IMPORT}
        decisions = payload.get("decisions") or {}
        for dom in DOMAINES_IMPORT:
            for item in a.get("valeurs", {}).get(dom, []):
                choix = {**item, **((decisions.get(dom) or {}).get(item["cle"]) or {})}
                action = choix.get("action")
                if action == "IGNORER":
                    mapping[dom][item["cle"]] = None
                elif action == "MAPPER":
                    if not choix.get("cible_id"):
                        raise AppError(f"Correspondance manquante pour « {item['brut']} »", 422, code="MAPPING_INCOMPLET")
                    mapping[dom][item["cle"]] = uuid.UUID(str(choix["cible_id"]))
                else:
                    libelle = propre(choix.get("libelle")) or item["brut"]
                    mapping[dom][item["cle"]] = await self._creer_valeur(dom, libelle, choix)
                    stats[f"crees_{dom.lower()}"] += 1

        # 2. Employés.
        entites = {e.id: e for e in (await self.db.scalars(select(FormationEntite))).unique().all()}
        employes: dict[str, uuid.UUID] = {}
        for p in a.get("personnes", []):
            choix = {**p, **(choix_personnes.get(p["cle"]) or {})}
            if choix.get("action") == "EXISTANT":
                if not choix.get("employe_id"):
                    raise AppError(f"Employé existant non choisi pour « {p['brut']} »", 422, code="EMPLOYE_NON_CHOISI")
                e = await self.db.get(FormationEmploye, uuid.UUID(str(choix["employe_id"])))
                if not e:
                    raise AppError(f"Employé introuvable pour « {p['brut']} »", 422, code="EMPLOYE_INCONNU")
                employes[p["cle"]] = e.id
                stats["employes_existants"] += 1
                continue
            nom = propre(choix.get("nom")) or p["brut"]
            prenom = propre(choix.get("prenom")) or None
            ci = cle_identite(nom, prenom)
            deja = await self.db.scalar(select(FormationEmploye).where(FormationEmploye.cle_identite == ci).limit(1))
            if deja:
                employes[p["cle"]] = deja.id
                stats["employes_existants"] += 1
                continue
            fonction_id = mapping["FONCTION"].get(cle(p.get("fonction"))) if p.get("fonction") else None
            entite_id = mapping["ENTITE"].get(cle(p.get("entite"))) if p.get("entite") else None
            e = FormationEmploye(nom=nom, prenom=prenom, cle_identite=ci, fonction_id=fonction_id, entite_id=entite_id,
                                 actif=True, source="IMPORT", import_id=imp.id, created_by_id=user_id,
                                 updated_by_id=user_id)
            self.db.add(e)
            await self.db.flush()
            employes[p["cle"]] = e.id
            stats["employes_crees"] += 1

        # 3. Sessions + participations.
        lignes = {l["ligne"]: l for l in a.get("lignes", [])}
        today = date.today()
        now = maintenant()
        for s in a.get("sessions", []):
            l0 = lignes[s["lignes"][0]] if s["lignes"] else None
            if not l0:
                continue
            themes = list(dict.fromkeys(t for t in (mapping["THEME"].get(cle(x)) for x in s["themes"]) if t))
            lieu_id = mapping["LIEU"].get(cle(s["lieu"]))
            if not themes or not lieu_id:
                stats["sessions_ignorees"] += 1
                continue
            formateurs = list(dict.fromkeys(f for f in (mapping["FORMATEUR"].get(cle(x)) for x in s["formateurs"]) if f))
            d = date.fromisoformat(s["date"])
            session = await self._session_existante(d, lieu_id, set(themes))
            if session is None:
                session = FormationSession(
                    reference=await self.svc._reference(d.year), date_session=d, lieu_id=lieu_id,
                    statut="PLANIFIEE", source="IMPORT", import_id=imp.id, revision=1,
                    observations=f"Import Excel « {imp.fichier_nom} »", created_by_id=user_id, updated_by_id=user_id)
                session.themes = [FormationSessionTheme(theme_id=t, ordre=i) for i, t in enumerate(themes)]
                session.formateurs = [FormationSessionFormateur(formateur_id=f, ordre=i) for i, f in enumerate(formateurs)]
                session.participants = []
                self.db.add(session)
                await self.db.flush()
                stats["sessions_creees"] += 1
            else:
                stats["sessions_fusionnees"] += 1
            deja = {p.employe_id for p in session.participants}
            ordre = max((p.ordre for p in session.participants), default=0)
            for ligne_no in s["lignes"]:
                l = lignes[ligne_no]
                eid = employes.get(cle_identite(l["participant"]))
                if not eid or eid in deja:
                    stats["participations_existantes"] += 1 if eid else 0
                    continue
                e = await self.db.get(FormationEmploye, eid)
                ent_id = mapping["ENTITE"].get(cle(l["entite"])) if l["entite"] else e.entite_id
                ent = entites.get(ent_id) if ent_id else None
                if ent is None and ent_id:
                    ent = await self.db.get(FormationEntite, ent_id)
                fonction_id = mapping["FONCTION"].get(cle(l["fonction"])) if l["fonction"] else e.fonction_id
                presence = l.get("presence") if d <= today else None
                ordre += 1
                session.participants.append(FormationParticipant(
                    session_id=session.id, employe_id=eid, presence=presence, fonction_id=fonction_id,
                    entite_id=ent_id, perimetre_id=ent.perimetre_id if ent else None, ordre=ordre,
                    presence_saisie_le=now if presence else None, presence_saisie_par_id=user_id if presence else None))
                deja.add(eid)
                stats["participations_creees"] += 1
            await self.db.flush()
            complet = bool(session.participants) and all(p.presence for p in session.participants)
            if session.statut in ("PLANIFIEE", "REALISEE"):
                if complet and d <= today:
                    session.statut, session.cloturee_le, session.cloturee_par_id = "CLOTUREE", now, user_id
                    session.presences_saisies_le, session.presences_saisies_par_id = now, user_id
                else:
                    session.statut = "PLANIFIEE"
            await self.db.flush()

        imp.statut = "IMPORTE"
        imp.importe_le = now
        imp.importe_par_id = user_id
        imp.resultat = dict(stats)
        await self.db.flush()
        await self.svc.audit("formation.import.confirm", "formation_import", imp.id,
                             after={"fichier": imp.fichier_nom, **dict(stats)})
        return self._import_dict(imp)

    async def _creer_valeur(self, dom: str, libelle: str, choix: dict) -> uuid.UUID:
        c = cle(libelle)
        if dom == "ENTITE":
            existant = await self.db.scalar(select(FormationEntite).where(FormationEntite.cle == c))
            if existant:
                return existant.id
            if not choix.get("perimetre_id"):
                raise AppError(f"Périmètre obligatoire pour la nouvelle entité « {libelle} »", 422,
                               code="PERIMETRE_OBLIGATOIRE")
            e = FormationEntite(libelle=libelle, cle=c, perimetre_id=uuid.UUID(str(choix["perimetre_id"])),
                                lieu_id=uuid.UUID(str(choix["lieu_id"])) if choix.get("lieu_id") else None,
                                actif=True, ordre=999, created_by_id=self.ctx.user.id, updated_by_id=self.ctx.user.id)
            self.db.add(e)
            await self.db.flush()
            await self.svc.audit("formation.entite.create", "formation_entite", e.id,
                                 after={"libelle": libelle, "source": "import"})
            return e.id
        existant = await self.db.scalar(select(FormationReferentiel).where(
            FormationReferentiel.domaine == dom, FormationReferentiel.cle == c))
        if existant:
            return existant.id
        r = FormationReferentiel(domaine=dom, libelle=libelle, cle=c, ordre=999, actif=True,
                                 created_by_id=self.ctx.user.id, updated_by_id=self.ctx.user.id)
        self.db.add(r)
        await self.db.flush()
        await self.svc.audit("formation.referentiel.create", "formation_referentiel", r.id,
                             after={"domaine": dom, "libelle": libelle, "source": "import"})
        return r.id

    async def _session_existante(self, d: date, lieu_id: uuid.UUID, themes: set[uuid.UUID]) -> FormationSession | None:
        cands = (await self.db.scalars(
            select(FormationSession).options(selectinload(FormationSession.participants))
            .where(FormationSession.date_session == d, FormationSession.lieu_id == lieu_id,
                   FormationSession.statut != "ANNULEE"))).unique().all()
        for s in cands:
            if {t.theme_id for t in s.themes} == themes:
                return s
        return None

"""Déclaration SIC : construction des lignes PP / PM, anomalies, fichier Excel et contre-vérification. Pur.

Aucune valeur n'est inventée : une donnée obligatoire absente ou invalide produit une anomalie BLOQUANTE et la
cellule reste vide. Tous les candidats sont exportés ; les lignes incomplètes sont listées dans le rapport
d'anomalies.
"""

from __future__ import annotations

import io
import re
from collections import Counter
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill

from app.data.clientele_sic import (
    COLONNES,
    FEUILLES,
    SECTEURS_GENERIQUES,
    decouper_nom,
    norm,
    normaliser_telephone,
)
from app.services.clientele.bcm_calcul import ClientBcm, forme_client

BLOQUANTE, A_CORRIGER, AVERTISSEMENT = "BLOQUANTE", "A_CORRIGER", "AVERTISSEMENT"
_RANG = {BLOQUANTE: 0, A_CORRIGER: 1, AVERTISSEMENT: 2}
FORMAT_DATE_EXCEL = "mm/dd/yyyy hh:mm:ss"
LONGUEUR_NOM_ORION = 25


@dataclass
class CandidatSic:
    racine: str
    profil: str | None = None
    categorie: str | None = None
    statut_resident: str | None = None
    nni: str | None = None
    nif: str | None = None
    nom_complet: str | None = None
    prenoms: str | None = None
    secteur_orion: str | None = None
    date_entree: date | None = None
    agence_id: str | None = None
    agence_code: str | None = None
    nb_comptes: int = 0
    nb_agences: int = 0
    eer_sexe: str | None = None
    eer_telephone: str | None = None
    kyc_genre: str | None = None
    kyc_telephone: str | None = None
    kyc_csp: str | None = None
    kyc_secteur: str | None = None


@dataclass
class ContexteSic:
    """``regles`` : valeurs courantes des règles SIC (code → valeur) ; ``regles_bcm`` : règles de forme du mensuel ;
    ``codes_ref`` : code BCM → usage (PP / PM / AUTRE) de la version de référentiel utilisée."""

    regles: dict[str, Any]
    regles_bcm: dict[str, Any] = field(default_factory=dict)
    codes_ref: dict[str, str] = field(default_factory=dict)
    agences: dict[str, str] = field(default_factory=dict)
    secteurs: dict[str, tuple[str, str]] = field(default_factory=dict)


def _anomalie(gravite: str, champ: str, valeur: Any, motif: str, action: str) -> dict:
    return {"gravite": gravite, "champ": champ, "valeur": None if valeur in (None, "") else str(valeur),
            "motif": motif, "action": action}


def _txt(v: Any) -> str:
    return " ".join(str(v or "").split())


def formater_date(d: date, fmt: str) -> str | datetime:
    if fmt == "TEXTE_MM_JJ_AAAA_HEURE":
        return d.strftime("%m/%d/%Y 00:00:00")
    if fmt == "DATE_EXCEL":
        return datetime(d.year, d.month, d.day)
    return d.strftime("%d/%m/%Y")


def _source_complement(c: CandidatSic, ctx: ContexteSic, champ: str, valeurs: dict[str, str | None],
                       anomalies: list[dict]) -> str | None:
    """Première source renseignée selon l'ordre ``sic.sources_complement`` ; divergence signalée."""
    ordre = ctx.regles.get("sic.sources_complement") or ["EER", "KYC"]
    presentes = [(s, valeurs.get(s)) for s in ordre if valeurs.get(s)]
    if not presentes:
        return None
    if len({v for _, v in presentes}) > 1:
        detail = " ; ".join(f"{s} = {v}" for s, v in presentes)
        anomalies.append(_anomalie(A_CORRIGER, champ, detail, f"Sources divergentes ({detail}) : {presentes[0][0]} "
                                   "retenu selon l'ordre de priorité", "Mettre en cohérence EER et KYC"))
    return presentes[0][1]


def _commun_debut(c: CandidatSic, type_sic: str, ctx: ContexteSic, anomalies: list[dict]) -> str:
    racine = _txt(c.racine)
    if not re.fullmatch(r"\d{6}", racine):
        anomalies.append(_anomalie(BLOQUANTE, "CodeClient", racine, "Racine client ORION hors format (6 chiffres)",
                                   "Vérifier la racine dans ORION"))
    if c.profil not in ("PP", "PM"):
        anomalies.append(_anomalie(BLOQUANTE, "Profil", c.profil, "Profil PP / PM non identifié (agent économique "
                                   "et catégorie juridique insuffisants)", "Compléter la fiche client dans ORION"))
    elif ctx.regles_bcm.get("forme.correspondance") is not None:
        forme = forme_client(ClientBcm(racine=racine, d_entree=None, a_ouvert=False, categorie=c.categorie,
                                       profil=c.profil), ctx.regles_bcm)
        if forme in ("PP", "PM") and forme != type_sic:
            anomalies.append(_anomalie(BLOQUANTE, "Profil", c.categorie, f"Confusion PP / PM : profil {c.profil}, "
                                       f"catégorie juridique rattachée à {forme}",
                                       "Corriger l'agent économique ou la catégorie juridique dans ORION"))
        elif forme == "CJ":
            anomalies.append(_anomalie(A_CORRIGER, "Profil", c.categorie, "Construction juridique : rattachement au "
                                       "fichier PM à confirmer", "Confirmer le traitement SIC de cette forme"))
    return racine


def _commun_fin(c: CandidatSic, ctx: ContexteSic, anomalies: list[dict]) -> tuple[Any, str | None]:
    d = None
    if c.date_entree is None:
        anomalies.append(_anomalie(BLOQUANTE, "DateEntreeEnRelation", None, "Aucune date d'ouverture de compte",
                                   "Vérifier les comptes ORION du client"))
    else:
        d = formater_date(c.date_entree, ctx.regles.get("sic.format_date") or "TEXTE_JJ_MM_AAAA")
    agence = None
    if not c.agence_id:
        anomalies.append(_anomalie(BLOQUANTE, "MoughataaAgence", None, "Agence du premier compte inconnue",
                                   "Rattacher le compte à une agence"))
    else:
        agence = ctx.agences.get(str(c.agence_id))
        if not agence:
            anomalies.append(_anomalie(BLOQUANTE, "MoughataaAgence", c.agence_code,
                                       f"Valeur SIC non paramétrée pour l'agence {c.agence_code or '—'}",
                                       "Renseigner la correspondance agence dans les référentiels BCM"))
    if c.nb_agences > 1:
        anomalies.append(_anomalie(AVERTISSEMENT, "MoughataaAgence", c.agence_code,
                                   f"Comptes dans {c.nb_agences} agences : agence du premier compte retenue",
                                   "Aucune (information)"))
    return d, agence


def _code_ref(code: str | None, usage: str, champ: str, ctx: ContexteSic, anomalies: list[dict]) -> str | None:
    code = _txt(code)
    if not code:
        return None
    if code not in ctx.codes_ref:
        anomalies.append(_anomalie(BLOQUANTE, champ, code, "Code absent de la version active du référentiel BCM",
                                   "Corriger le code ou importer le référentiel à jour"))
        return None
    if ctx.codes_ref[code] != usage:
        anomalies.append(_anomalie(BLOQUANTE, champ, code, f"Code réservé aux {ctx.codes_ref[code]} dans le "
                                   f"référentiel BCM (attendu : {usage})", "Choisir un code de la bonne catégorie"))
        return None
    return code


def ligne_pp(c: CandidatSic, ctx: ContexteSic) -> tuple[list, list[dict]]:
    an: list[dict] = []
    racine = _commun_debut(c, "PP", ctx, an)
    residence = (ctx.regles.get("sic.residence") or {}).get(norm(c.statut_resident))
    if not residence:
        an.append(_anomalie(BLOQUANTE, "StatutResidentiel", c.statut_resident, "Statut de résidence ORION absent "
                            "ou sans correspondance", "Renseigner le statut de résidence dans ORION"))
    nni = _txt(c.nni)
    if not nni:
        an.append(_anomalie(BLOQUANTE, "NNI", None, "NNI absent", "Renseigner le NNI dans ORION"))
    elif not re.fullmatch(ctx.regles.get("sic.nni_format") or r".+", nni):
        an.append(_anomalie(A_CORRIGER, "NNI", nni, "NNI hors format attendu (exporté tel quel)",
                            "Vérifier la pièce d'identité"))
    nom, prenom, echec = decouper_nom(c.nom_complet, c.prenoms,
                                      en_tete=ctx.regles.get("sic.nom_prenom") == "PRENOM_EN_FIN_OU_EN_TETE")
    if echec:
        champ = "Prenom" if "Prénom absent" in echec else "Nom"
        an.append(_anomalie(BLOQUANTE, champ, _txt(c.nom_complet), echec,
                            "Corriger le nom / prénom dans ORION (aucun découpage approximatif)"))
    genre_map = ctx.regles.get("sic.genre") or {}
    eer_genre = genre_map.get(norm(c.eer_sexe)) if c.eer_sexe else None
    kyc_genre = genre_map.get(norm(c.kyc_genre)) if c.kyc_genre else None
    if c.eer_sexe and not eer_genre:
        an.append(_anomalie(A_CORRIGER, "Genre", c.eer_sexe, "Sexe EER sans correspondance H / F",
                            "Corriger le dossier EER"))
    genre = _source_complement(c, ctx, "Genre", {"EER": eer_genre, "KYC": kyc_genre}, an)
    if not genre:
        an.append(_anomalie(BLOQUANTE, "Genre", None, "Genre absent (ni EER validé, ni complément KYC)",
                            "Saisir le genre dans les compléments SIC"))
    csp = _code_ref(c.kyc_csp, "PP", "CodeCategorieSocioprofessionnelle", ctx, an)
    if not _txt(c.kyc_csp):
        an.append(_anomalie(BLOQUANTE, "CodeCategorieSocioprofessionnelle", None,
                            "Catégorie socioprofessionnelle BCM absente du KYC",
                            "Saisir le code CSP dans les compléments SIC"))
    tel = _source_complement(c, ctx, "NumTel", {"EER": normaliser_telephone(c.eer_telephone) or None,
                                                "KYC": normaliser_telephone(c.kyc_telephone) or None}, an)
    if not tel:
        an.append(_anomalie(BLOQUANTE, "NumTel", None, "Téléphone absent (ni EER validé, ni complément KYC)",
                            "Saisir le téléphone dans les compléments SIC"))
    elif not re.fullmatch(ctx.regles.get("sic.telephone_format") or r".+", tel):
        an.append(_anomalie(BLOQUANTE, "NumTel", tel, "Numéro de téléphone hors format", "Corriger le numéro"))
        tel = None
    d, agence = _commun_fin(c, ctx, an)
    return [racine, residence, nni or None, prenom, nom, genre, csp, tel, d, agence], an


def ligne_pm(c: CandidatSic, ctx: ContexteSic) -> tuple[list, list[dict]]:
    an: list[dict] = []
    racine = _commun_debut(c, "PM", ctx, an)
    nif = _txt(c.nif)
    if not nif and ctx.regles.get("sic.nif_obligatoire", True):
        an.append(_anomalie(BLOQUANTE, "NIF", None, "NIF absent (le NNI n'est jamais utilisé à sa place)",
                            "Renseigner le NIF dans ORION"))
    raison = _txt(c.nom_complet)
    if not raison:
        an.append(_anomalie(BLOQUANTE, "RaisonSociale", None, "Raison sociale absente", "Compléter ORION"))
    elif len(raison) == LONGUEUR_NOM_ORION:
        an.append(_anomalie(AVERTISSEMENT, "RaisonSociale", raison, f"Raison sociale de {LONGUEUR_NOM_ORION} "
                            "caractères : possiblement tronquée par ORION", "Vérifier la raison sociale complète"))
    secteur = None
    if _txt(c.kyc_secteur):
        secteur = _code_ref(c.kyc_secteur, "PM", "CodeSecteurActivite", ctx, an)
    else:
        lib = norm(c.secteur_orion)
        if not lib or lib in SECTEURS_GENERIQUES:
            an.append(_anomalie(BLOQUANTE, "CodeSecteurActivite", c.secteur_orion or None,
                                "Secteur ORION absent ou générique : aucun rattachement automatique",
                                "Saisir le code secteur BCM dans les compléments SIC"))
        elif lib not in ctx.secteurs:
            an.append(_anomalie(BLOQUANTE, "CodeSecteurActivite", c.secteur_orion,
                                "Secteur ORION sans correspondance BCM",
                                "Ajouter la correspondance secteur ou saisir le code dans les compléments SIC"))
        else:
            code, statut = ctx.secteurs[lib]
            secteur = _code_ref(code, "PM", "CodeSecteurActivite", ctx, an)
            if secteur and statut != "CONFIRMEE":
                an.append(_anomalie(A_CORRIGER, "CodeSecteurActivite", f"{c.secteur_orion} → {code}",
                                    "Code issu d'une correspondance secteur à valider",
                                    "Confirmer la correspondance dans les référentiels BCM"))
    d, agence = _commun_fin(c, ctx, an)
    return [racine, nif or None, raison or None, secteur, d, agence], an


def gravite_max(anomalies: list[dict]) -> str | None:
    return min((a["gravite"] for a in anomalies), key=_RANG.__getitem__, default=None)


def construire(candidats: list[CandidatSic], type_sic: str, ctx: ContexteSic) -> list[dict]:
    """Une entrée par candidat (toujours EXPORTEE) : racine, valeurs, anomalies, gravité maximale."""
    fn = ligne_pp if type_sic == "PP" else ligne_pm
    lignes = []
    for c in candidats:
        valeurs, an = fn(c, ctx)
        lignes.append({"racine": _txt(c.racine), "valeurs": valeurs, "anomalies": an})
    racines = Counter(lg["racine"] for lg in lignes)
    i_id = 2 if type_sic == "PP" else 1
    champ_id = COLONNES[type_sic][i_id]
    ids = Counter(lg["valeurs"][i_id] for lg in lignes if lg["valeurs"][i_id])
    for lg in lignes:
        if racines[lg["racine"]] > 1:
            lg["anomalies"].append(_anomalie(BLOQUANTE, "CodeClient", lg["racine"], "Racine client en double",
                                             "Dédoublonner la racine dans ORION"))
        ident = lg["valeurs"][i_id]
        if ident and ids[ident] > 1:
            lg["anomalies"].append(_anomalie(A_CORRIGER, champ_id, ident,
                                             f"{champ_id} partagé par {ids[ident]} racines clients",
                                             "Vérifier s'il s'agit du même client (doublon de racine)"))
    for lg in lignes:
        lg["gravite_max"] = gravite_max(lg["anomalies"])
        lg["statut"] = "EXPORTEE"
        lg["valeurs"] = [v.isoformat() if isinstance(v, datetime) else v for v in lg["valeurs"]]
    return lignes


# --------------------------------------------------------------------------- fichier Excel
def _valeur_cellule(colonne: str, v: Any, fmt_date: str) -> Any:
    if colonne == "DateEntreeEnRelation" and fmt_date == "DATE_EXCEL" and v:
        return datetime.fromisoformat(v) if isinstance(v, str) else v
    return None if v is None else str(v)


def ecrire_xlsx(type_sic: str, lignes: list[list], fmt_date: str) -> bytes:
    """Feuille unique ``pp`` / ``pm`` : en-têtes du modèle puis une ligne par client exporté. Identifiants en texte."""
    wb = Workbook()
    ws = wb.active
    ws.title = FEUILLES[type_sic]
    colonnes = COLONNES[type_sic]
    for j, nom in enumerate(colonnes, start=1):
        cell = ws.cell(row=1, column=j, value=nom)
        cell.number_format = "@"
        cell.font = Font(bold=True)
    for i, valeurs in enumerate(lignes, start=2):
        for j, (nom, v) in enumerate(zip(colonnes, valeurs), start=1):
            cell = ws.cell(row=i, column=j, value=_valeur_cellule(nom, v, fmt_date))
            date_excel = nom == "DateEntreeEnRelation" and fmt_date == "DATE_EXCEL"
            cell.number_format = FORMAT_DATE_EXCEL if date_excel else "@"
    for j, nom in enumerate(colonnes, start=1):
        ws.column_dimensions[ws.cell(row=1, column=j).column_letter].width = max(14, len(nom) + 4)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


_MOTIF_DATE = {
    "TEXTE_JJ_MM_AAAA": r"\d{2}/\d{2}/\d{4}",
    "TEXTE_MM_JJ_AAAA_HEURE": r"\d{2}/\d{2}/\d{4} 00:00:00",
}


def verifier_xlsx(contenu: bytes, type_sic: str, attendu: list[list], fmt_date: str, codes_ref: dict[str, str],
                  regles: dict[str, Any]) -> dict:
    """Relit le fichier généré : feuille, en-têtes, ordre, types, nombre de lignes, doublons, valeurs autorisées.

    Seuls les contrôles de structure sont bloquants ; cellules vides et valeurs hors référentiel sont signalées.
    """
    controles: list[dict] = []

    def ok(code: str, libelle: str, reussi: bool, detail: str = "", bloquant: bool = True) -> None:
        controles.append({"code": code, "libelle": libelle, "ok": bool(reussi), "detail": detail,
                          "bloquant": bloquant})

    def hors(valeurs: list, admis) -> int:
        return sum(1 for v in valeurs if v not in (None, "") and not admis(v))

    wb = load_workbook(io.BytesIO(contenu))
    colonnes = COLONNES[type_sic]
    ok("feuille", f"Feuille unique « {FEUILLES[type_sic]} »", wb.sheetnames == [FEUILLES[type_sic]],
       ", ".join(wb.sheetnames))
    ws = wb[wb.sheetnames[0]]
    entetes = tuple(ws.cell(row=1, column=j).value for j in range(1, ws.max_column + 1))
    ok("entetes", f"{len(colonnes)} colonnes du modèle, dans l'ordre", entetes == colonnes,
       "" if entetes == colonnes else f"Lu : {', '.join(str(x) for x in entetes)}")
    lignes = [list(r) for r in ws.iter_rows(min_row=2, max_col=len(colonnes), values_only=True)]
    ok("nb_lignes", "Nombre de lignes égal aux clients exportés", len(lignes) == len(attendu),
       f"{len(lignes)} lu(s) / {len(attendu)} attendu(s)")
    texte_ko = 0
    for row in ws.iter_rows(min_row=2, max_col=len(colonnes)):
        for nom, cell in zip(colonnes, row):
            if cell.value is None:
                continue
            if nom == "DateEntreeEnRelation" and fmt_date == "DATE_EXCEL":
                texte_ko += not isinstance(cell.value, datetime)
            else:
                texte_ko += not (cell.data_type == "s" and cell.number_format == "@")
    ok("types", "Identifiants et codes au format texte (zéros initiaux conservés)", texte_ko == 0,
       f"{texte_ko} cellule(s) hors type attendu" if texte_ko else "")
    racines = [str(r[0]) if r and r[0] is not None else "" for r in lignes]
    doublons = [k for k, n in Counter(racines).items() if n > 1]
    ok("doublons", "Aucun CodeClient en double", not doublons, ", ".join(doublons[:10]))
    cc_ko = sum(1 for r in racines if not re.fullmatch(r"\d{6}", r))
    ok("code_client", "CodeClient à 6 chiffres", cc_ko == 0, f"{cc_ko} hors format" if cc_ko else "", bloquant=False)
    vides = sum(1 for r in lignes for v in r if v in (None, ""))
    incompletes = sum(1 for r in lignes if any(v in (None, "") for v in r))
    ok("obligatoires", "Aucune cellule vide", vides == 0,
       f"{vides} cellule(s) vide(s) sur {incompletes} ligne(s), à compléter (voir rapport d'anomalies)" if vides
       else "", bloquant=False)
    idx = {nom: j for j, nom in enumerate(colonnes)}

    def colonne(nom: str) -> list:
        return [r[idx[nom]] for r in lignes]

    d_ko = 0
    if fmt_date in _MOTIF_DATE:
        d_ko = hors(colonne("DateEntreeEnRelation"), lambda v: re.fullmatch(_MOTIF_DATE[fmt_date], str(v)))
    ok("dates", "DateEntreeEnRelation au format paramétré", d_ko == 0, f"{d_ko} date(s) hors format" if d_ko else "",
       bloquant=False)
    if type_sic == "PP":
        resid = set((regles.get("sic.residence") or {}).values()) or {"MR", "ETR"}
        n = hors(colonne("StatutResidentiel"), lambda v: v in resid)
        ok("residence", "StatutResidentiel autorisé", n == 0, f"{n} hors liste" if n else "", bloquant=False)
        n = hors(colonne("Genre"), lambda v: v in ("H", "F"))
        ok("genre", "Genre H / F", n == 0, f"{n} hors H / F" if n else "", bloquant=False)
        n = hors(colonne("CodeCategorieSocioprofessionnelle"), lambda v: codes_ref.get(str(v)) == "PP")
        ok("csp", "Code CSP présent dans le référentiel (PP)", n == 0, f"{n} hors référentiel" if n else "",
           bloquant=False)
    else:
        n = hors(colonne("CodeSecteurActivite"), lambda v: codes_ref.get(str(v)) == "PM")
        ok("secteur", "Code secteur présent dans le référentiel (PM)", n == 0, f"{n} hors référentiel" if n else "",
           bloquant=False)
    relu = [[v.isoformat() if isinstance(v, datetime) else v for v in r] for r in lignes]
    ok("contenu", "Contenu identique aux lignes contrôlées", relu == attendu)
    return {"ok": all(c["ok"] for c in controles if c["bloquant"]), "controles": controles}


# --------------------------------------------------------------------------- rapport d'anomalies
_ENTETES_RAPPORT = ("Racine client", "Ligne", "Gravité", "Champ", "Valeur observée", "Motif", "Action attendue")
_FOND_ENTETE = PatternFill("solid", fgColor="E2E3E5")


def rapport_anomalies_xlsx(entete: list[tuple[str, Any]], lignes: list[dict]) -> bytes:
    """Synthèse + une ligne par anomalie. ``lignes`` : racine, statut, anomalies. Racines en texte."""
    wb = Workbook(write_only=True)
    gras = Font(bold=True)

    def cellule(ws, v, *, police=None, fond=None):
        c = WriteOnlyCell(ws, value=v)
        if not isinstance(v, (int, float)):
            c.number_format = "@"
        if police:
            c.font = police
        if fond:
            c.fill = fond
        return c

    ws = wb.create_sheet("Synthèse")
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 60
    for k, v in entete:
        ws.append([cellule(ws, k, police=gras),
                   cellule(ws, v if isinstance(v, (int, float)) else ("" if v is None else str(v)))])
    wa = wb.create_sheet("Anomalies")
    for j, w in enumerate((14, 10, 16, 32, 30, 70, 60)):
        wa.column_dimensions["ABCDEFG"[j]].width = w
    wa.freeze_panes = "A2"
    wa.append([cellule(wa, h, police=gras, fond=_FOND_ENTETE) for h in _ENTETES_RAPPORT])
    for lg in lignes:
        if lg["statut"] == "EXCLUE":
            statut = "Exclue"
        elif any(a["gravite"] == BLOQUANTE for a in lg["anomalies"]):
            statut = "Incomplète"
        else:
            statut = "Exportée"
        for a in lg["anomalies"]:
            wa.append([str(lg["racine"]), statut, a["gravite"], a["champ"], a.get("valeur") or "", a["motif"],
                       a["action"]])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

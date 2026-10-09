"""Lots de clients : un fichier Excel avec une colonne Racine client (+ colonnes libres).

Deux usages :
- ``LISTE``     : quelques racines reçues de l'extérieur → évaluation moteur, classement, export ;
- ``SITUATION`` : classeur Conformité « Situation des comptes PP et PM » → reprise de la
  classe de risque historique (source ``EXCEL_CONFORMITE``).

Le fichier et le résultat d'analyse sont conservés sur disque (pas de table) le temps du
traitement. Rien n'est écrit dans la classification sans confirmation explicite. La racine
n'est jamais modifiée ; une racine absente du référentiel n'est jamais créée ici.
"""

from __future__ import annotations

import hashlib
import json
import time
import uuid
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import get_settings
from app.core.exceptions import AppError
from app.models import ClienteleClassification, ClienteleClient
from app.services.clientele.classification import NIVEAUX, ClienteleClassificationService
from app.services.clientele.consolidation import normaliser_entete
from app.services.clientele.service import Ctx, ClienteleService
from app.services.reporting_export import build_styled_workbook_multi

TAILLE_MAX = 40 * 1024 * 1024
LIGNES_MAX = 80_000
LIMITE_EVALUATION = 2_000
APERCU = 300
DUREE_VIE_S = 7 * 24 * 3600
TYPES = ("LISTE", "SITUATION")

ENTETES_RACINE = ("RACINE_CLIENT", "RACINE CLIENT", "RACINE", "CLIENT", "CODE CLIENT", "NUMERO CLIENT")
ENTETES_NIVEAU = (
    "NIVEAU", "CLASSE RISQUE LBC FT", "CLASSE DE RISQUE", "CLASSE RISQUE", "NIVEAU RISQUE",
    "NIVEAU DE RISQUE", "RISQUE LBC FT", "CLASSIFICATION",
)
ENTETES_MOTIF = ("MOTIF", "MOTIF_CLASSEMENT", "MOTIF CLASSEMENT", "MOTIF RISQUE", "MOTIF_RISQUE", "OBSERVATION")
ENTETES_NOM = ("RAISON_SOCIAL", "RAISON SOCIALE", "NOM", "NOM CLIENT", "INTITULE")


def _norm(v: Any) -> str:
    return normaliser_entete(str(v)) if v is not None else ""


def normaliser_racine(v: Any) -> str | None:
    if v is None:
        return None
    if isinstance(v, float):
        if not v.is_integer():
            return None
        v = int(v)
    s = str(v).strip()
    if not s or not s.isdigit() or len(s) > 6:
        return None
    return s.zfill(6)


def normaliser_niveau(v: Any) -> str | None:
    s = _norm(v).replace("RISQUE", "").strip()
    if not s:
        return None
    if s.startswith("INTERDIT"):
        return "INTERDIT"
    if s.startswith("ELEV") or s in ("HAUT", "FORT"):
        return "ELEVE"
    if s.startswith("MOY"):
        return "MOYEN"
    if s.startswith("FAIBL") or s == "BAS":
        return "FAIBLE"
    return None


def _cherche(entetes: list[str], candidats: tuple[str, ...]) -> str | None:
    norm = {_norm(e): e for e in entetes if e}
    for c in candidats:
        if _norm(c) in norm:
            return norm[_norm(c)]
    return None


def _entetes(row: tuple) -> list[str]:
    out = []
    for i, c in enumerate(row):
        s = str(c).strip() if c is not None else ""
        out.append(s or f"(colonne {get_column_letter(i + 1)})")
    return out


def _detecter(wb) -> tuple[str, int, list[str]]:
    meilleur: tuple[int, str, int, list[str]] | None = None
    for ws in wb.worksheets:
        for i, row in enumerate(ws.iter_rows(min_row=1, max_row=15, values_only=True), start=1):
            entetes = _entetes(row)
            score = sum(1 for e in entetes if e and _norm(e) in {_norm(x) for x in
                                                                ENTETES_RACINE + ENTETES_NIVEAU + ENTETES_NOM})
            if _cherche(entetes, ENTETES_RACINE):
                score += 5
            if score and (meilleur is None or score > meilleur[0]):
                meilleur = (score, ws.title, i, entetes)
    if not meilleur:
        ws = wb.worksheets[0]
        row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
        return ws.title, 1, _entetes(row)
    return meilleur[1], meilleur[2], meilleur[3]


class ClienteleLotService:
    def __init__(self, db: AsyncSession, ctx: Ctx):
        self.db = db
        self.ctx = ctx
        self.svc = ClienteleService(db, ctx)

    # ------------------------------------------------------------------ stockage
    def _dossier(self) -> Path:
        p = Path(get_settings().upload_dir) / "clientele" / "lots"
        p.mkdir(parents=True, exist_ok=True)
        return p

    def _xlsx(self, lid: uuid.UUID) -> Path:
        return self._dossier() / f"{lid}.xlsx"

    def _json(self, lid: uuid.UUID) -> Path:
        return self._dossier() / f"{lid}.json"

    def _charger(self, lid: uuid.UUID) -> dict:
        p = self._json(lid)
        if not p.exists():
            raise AppError("Lot introuvable ou expiré ; déposez de nouveau le fichier", 404, code="NOT_FOUND")
        lot = json.loads(p.read_text(encoding="utf-8"))
        if lot.get("created_by") not in (None, str(self.ctx.user.id)) and not self.ctx.peut("clientele.admin"):
            raise AppError("Lot introuvable ou expiré ; déposez de nouveau le fichier", 404, code="NOT_FOUND")
        return lot

    def _sauver(self, lot: dict) -> None:
        self._json(uuid.UUID(lot["id"])).write_text(json.dumps(lot, ensure_ascii=False), encoding="utf-8")

    def _purger(self) -> None:
        limite = time.time() - DUREE_VIE_S
        for p in self._dossier().glob("*"):
            try:
                if p.stat().st_mtime < limite:
                    p.unlink()
            except OSError:
                pass

    # ------------------------------------------------------------------ analyse
    async def analyser(self, contenu: bytes, nom_fichier: str, *, type_lot: str = "LISTE",
                       colonne_racine: str | None = None, colonne_niveau: str | None = None,
                       colonne_motif: str | None = None, lid: uuid.UUID | None = None) -> dict:
        if not self.ctx.peut("clientele.import.execute", "clientele.classif.execute"):
            self.ctx.exiger("clientele.classif.execute")
        type_lot = type_lot if type_lot in TYPES else "LISTE"
        if not contenu:
            raise AppError("Fichier vide", 422, code="FICHIER_VIDE")
        if len(contenu) > TAILLE_MAX:
            raise AppError("Fichier trop volumineux (40 Mo maximum)", 413, code="FICHIER_TROP_VOLUMINEUX")
        if not nom_fichier.lower().endswith((".xlsx", ".xlsm")):
            raise AppError("Format non pris en charge : fichier .xlsx attendu", 422, code="FORMAT_NON_SUPPORTE")
        try:
            wb = load_workbook(BytesIO(contenu), read_only=True, data_only=True)
        except Exception as exc:
            raise AppError("Fichier Excel illisible ou corrompu", 422, code="FICHIER_ILLISIBLE") from exc
        try:
            feuille, ligne_entete, entetes = _detecter(wb)
            col_r = colonne_racine if colonne_racine in entetes else _cherche(entetes, ENTETES_RACINE)
            col_n = colonne_niveau if colonne_niveau in entetes else _cherche(entetes, ENTETES_NIVEAU)
            col_m = colonne_motif if colonne_motif in entetes else _cherche(entetes, ENTETES_MOTIF)
            col_nom = _cherche(entetes, ENTETES_NOM)
            if colonne_niveau == "":
                col_n = None
            if not col_r:
                lid = lid or uuid.uuid4()
                self._xlsx(lid).write_bytes(contenu)
                lot = {
                    "id": str(lid), "fichier_nom": nom_fichier[:255], "type": type_lot,
                    "created_by": str(self.ctx.user.id), "feuille": feuille, "ligne_entete": ligne_entete,
                    "entetes": [e for e in entetes if e],
                    "colonne_racine": None, "colonne_niveau": col_n, "colonne_motif": col_m,
                    "stats": {}, "lignes": [], "evaluation": None,
                    "message": "Colonne Racine client non reconnue : choisissez-la dans la liste.",
                }
                self._sauver(lot)
                return self._public(lot)
            ir = entetes.index(col_r)
            i_n = entetes.index(col_n) if col_n else None
            i_m = entetes.index(col_m) if col_m else None
            i_nom = entetes.index(col_nom) if col_nom else None
            autres = [(i, e) for i, e in enumerate(entetes) if e and i not in (ir, i_n, i_m, i_nom)][:6]
            lignes: list[dict] = []
            ws = wb[feuille]
            for numero, row in enumerate(ws.iter_rows(min_row=ligne_entete + 1, values_only=True),
                                         start=ligne_entete + 1):
                if not row or all(c is None or str(c).strip() == "" for c in row):
                    continue
                if len(lignes) >= LIGNES_MAX:
                    raise AppError(f"Plus de {LIGNES_MAX} lignes : découpez le fichier", 422, code="TROP_DE_LIGNES")
                cell = lambda i: row[i] if i is not None and i < len(row) else None  # noqa: E731
                brut = cell(ir)
                lignes.append({
                    "ligne": numero,
                    "brut": None if brut is None else str(brut)[:40],
                    "racine": normaliser_racine(brut),
                    "niveau_fichier_brut": None if cell(i_n) is None else str(cell(i_n))[:40],
                    "niveau_fichier": normaliser_niveau(cell(i_n)) if i_n is not None else None,
                    "motif_fichier": None if cell(i_m) is None else str(cell(i_m)).strip()[:500] or None,
                    "nom_fichier": None if cell(i_nom) is None else str(cell(i_nom)).strip()[:120],
                    "colonnes": {e: (None if cell(i) is None else str(cell(i))[:80]) for i, e in autres},
                })
        finally:
            wb.close()

        await self._controler(lignes)
        lid = lid or uuid.uuid4()
        self._purger()
        self._xlsx(lid).write_bytes(contenu)
        lot = {
            "id": str(lid), "fichier_nom": nom_fichier[:255], "type": type_lot,
            "sha256": hashlib.sha256(contenu).hexdigest(), "created_by": str(self.ctx.user.id),
            "feuille": feuille, "ligne_entete": ligne_entete, "entetes": [e for e in entetes if e],
            "colonne_racine": col_r, "colonne_niveau": col_n, "colonne_motif": col_m,
            "colonnes_affichees": [e for _, e in autres],
            "stats": self._stats(lignes), "lignes": lignes, "evaluation": None,
        }
        self._sauver(lot)
        await self.svc.audit("clientele.lot.analyse", "clientele_lot", lid,
                             after={"fichier": nom_fichier, "type": type_lot, **lot["stats"]})
        return self._public(lot)

    async def recartographier(self, lid: uuid.UUID, payload: dict) -> dict:
        lot = self._charger(lid)
        p = self._xlsx(lid)
        if not p.exists():
            raise AppError("Fichier du lot introuvable ; déposez-le de nouveau", 409, code="FICHIER_ABSENT")
        return await self.analyser(
            p.read_bytes(), lot["fichier_nom"], type_lot=payload.get("type") or lot["type"],
            colonne_racine=payload.get("colonne_racine"), colonne_niveau=payload.get("colonne_niveau"),
            colonne_motif=payload.get("colonne_motif"), lid=lid)

    async def _controler(self, lignes: list[dict]) -> None:
        racines = sorted({x["racine"] for x in lignes if x["racine"]})
        connus: dict[str, str] = {}
        classes: dict[str, dict] = {}
        for i in range(0, len(racines), 5000):
            lot = racines[i:i + 5000]
            for r, nom in (await self.db.execute(
                    select(ClienteleClient.racine_client, ClienteleClient.raison_sociale)
                    .where(ClienteleClient.racine_client.in_(lot)))).all():
                connus[r] = nom
            for c in (await self.db.scalars(
                    select(ClienteleClassification).where(ClienteleClassification.racine_client.in_(lot)))).all():
                classes[c.racine_client] = {"niveau": c.niveau, "source": c.source}
        vues: set[str] = set()
        for x in lignes:
            r = x["racine"]
            if not r:
                x["statut"] = "INVALIDE"
            elif r in vues:
                x["statut"] = "DOUBLON"
            elif r not in connus:
                x["statut"] = "INCONNUE"
            else:
                x["statut"] = "OK"
                vues.add(r)
            x["nom"] = connus.get(r) if r else None
            x["classe_actuelle"] = classes.get(r, {}).get("niveau") if r else None
            x["source_actuelle"] = classes.get(r, {}).get("source") if r else None
            if x["niveau_fichier_brut"] and not x["niveau_fichier"]:
                x["niveau_fichier_invalide"] = True

    @staticmethod
    def _stats(lignes: list[dict]) -> dict:
        def n(pred) -> int:
            return sum(1 for x in lignes if pred(x))
        return {
            "lignes": len(lignes),
            "racines_ok": n(lambda x: x["statut"] == "OK"),
            "inconnues": n(lambda x: x["statut"] == "INCONNUE"),
            "invalides": n(lambda x: x["statut"] == "INVALIDE"),
            "doublons": n(lambda x: x["statut"] == "DOUBLON"),
            "avec_niveau": n(lambda x: x["statut"] == "OK" and x["niveau_fichier"]),
            "niveau_illisible": n(lambda x: x.get("niveau_fichier_invalide")),
            "deja_classes": n(lambda x: x["statut"] == "OK" and x["classe_actuelle"]),
            "manuels": n(lambda x: x["statut"] == "OK" and x["source_actuelle"] in ("MANUEL", "MANUEL_SURCHARGE")),
        }

    def _public(self, lot: dict, *, statut: str | None = None, page: int = 1, taille: int = APERCU) -> dict:
        lignes = lot.get("lignes") or []
        if statut:
            lignes = [x for x in lignes if x.get("statut") == statut]
        debut = (max(1, page) - 1) * taille
        out = {k: v for k, v in lot.items() if k not in ("lignes", "evaluation")}
        out["total_filtre"] = len(lignes)
        out["page"] = page
        out["lignes"] = lignes[debut:debut + taille]
        ev = lot.get("evaluation")
        out["evaluation"] = {k: v for k, v in ev.items() if k != "resultats"} if ev else None
        out["limite_evaluation"] = LIMITE_EVALUATION
        return out

    async def detail(self, lid: uuid.UUID, *, statut: str | None = None, page: int = 1) -> dict:
        return self._public(self._charger(lid), statut=statut, page=page)

    async def supprimer(self, lid: uuid.UUID) -> None:
        self._charger(lid)
        for p in (self._xlsx(lid), self._json(lid)):
            if p.exists():
                p.unlink()

    # ------------------------------------------------------------------ opérations
    async def evaluer(self, lid: uuid.UUID) -> dict:
        self.ctx.exiger("clientele.classif.view")
        lot = self._charger(lid)
        racines = [x["racine"] for x in lot["lignes"] if x.get("statut") == "OK"]
        if not racines:
            raise AppError("Aucune racine connue du référentiel dans ce fichier", 422, code="AUCUNE_RACINE")
        if len(racines) > LIMITE_EVALUATION:
            raise AppError(
                f"{len(racines)} racines : l'évaluation moteur est limitée à {LIMITE_EVALUATION} par lot. "
                "Découpez le fichier.", 422, code="LOT_TROP_GRAND")
        classif = ClienteleClassificationService(self.db, self.ctx)
        resultats: dict[str, dict] = {}
        compteurs: dict[str, int] = {}
        for r in racines:
            ev = await classif.evaluer_racine(r, persister=False)
            niveau = ev.get("niveau_final") or "NON_CLASSE"
            compteurs[niveau] = compteurs.get(niveau, 0) + 1
            resultats[r] = {
                "score": ev.get("score_total"), "niveau": ev.get("niveau_final"),
                "statut": ev.get("statut"), "motif": ev.get("motif_genere") or ev.get("motif_principal"),
                "classable": ev.get("statut") in ("EVALUE", "BLOQUANT") and bool(ev.get("niveau_final")),
            }
        for x in lot["lignes"]:
            res = resultats.get(x["racine"]) if x.get("statut") == "OK" else None
            x["moteur"] = res
        lot["evaluation"] = {
            "nb": len(racines), "par_niveau": compteurs,
            "classables": sum(1 for v in resultats.values() if v["classable"]),
            "a_arbitrer": sum(1 for v in resultats.values() if v["statut"] == "A_ARBITRER"),
            "resultats": resultats,
        }
        self._sauver(lot)
        await self.svc.audit("clientele.lot.evaluer", "clientele_lot", lid,
                             after={"nb": len(racines), "par_niveau": compteurs})
        return self._public(lot)

    async def classer(self, lid: uuid.UUID, payload: dict) -> dict:
        self.ctx.exiger("clientele.classif.execute")
        lot = self._charger(lid)
        origine = payload.get("origine")
        motif = (payload.get("motif") or "").strip()
        forcer = bool(payload.get("forcer"))
        if origine not in ("MOTEUR", "FICHIER"):
            raise AppError("Origine attendue : MOTEUR ou FICHIER", 422, code="ORIGINE_INVALIDE")
        if not motif:
            raise AppError("Motif de classement obligatoire", 422, code="MOTIF_OBLIGATOIRE")
        classif = ClienteleClassificationService(self.db, self.ctx)
        ecrits = ignores_manuels = inchanges = non_classables = 0
        for x in lot["lignes"]:
            if x.get("statut") != "OK":
                continue
            r = x["racine"]
            actuelle = await self.db.get(ClienteleClassification, r)
            if actuelle and actuelle.source in ("MANUEL", "MANUEL_SURCHARGE") and not forcer:
                ignores_manuels += 1
                continue
            if origine == "FICHIER":
                niveau = x.get("niveau_fichier")
                if niveau not in NIVEAUX:
                    non_classables += 1
                    continue
                if actuelle and actuelle.niveau == niveau:
                    inchanges += 1
                    continue
                source = "EXCEL_CONFORMITE" if lot["type"] == "SITUATION" else "EXCEL"
                motifs = [{"critere": source, "motif": x.get("motif_fichier") or motif, "niveau": niveau,
                           "fichier": lot["fichier_nom"], "ligne": x["ligne"]}]
                await classif._enregistrer(
                    r, niveau, source=source, version_id=None, motifs=motifs,
                    motif_risque=x.get("motif_fichier"), motif_classement=motif)
            else:
                if not lot.get("evaluation"):
                    raise AppError("Lancez d'abord l'évaluation moteur du lot", 409, code="EVALUATION_ABSENTE")
                ev = await classif.evaluer_racine(r, persister=True)
                niveau = ev.get("niveau_final")
                if ev.get("statut") not in ("EVALUE", "BLOQUANT") or niveau not in NIVEAUX:
                    non_classables += 1
                    continue
                if actuelle and actuelle.niveau == niveau and actuelle.source == "MOTEUR":
                    inchanges += 1
                    continue
                motifs = [{
                    "critere": lg["critere"], "libelle": lg["libelle"], "famille": lg.get("famille"),
                    "valeur": lg.get("valeur"), "poids": lg.get("poids"), "niveau": lg.get("niveau_retenu"),
                    "motif": lg.get("motif"),
                } for lg in ev.get("lignes", []) if lg.get("contribue_au_score") or lg.get("blocking_propose")]
                motifs.append({"critere": "SCORE", "poids": ev.get("score_total"),
                               "version_regles": ev.get("version_regles"),
                               "evaluation_id": ev.get("evaluation_id")})
                await classif._enregistrer(
                    r, niveau, source="MOTEUR", version_id=None, motifs=motifs,
                    motif_risque=ev.get("motif_genere"), motif_classement=motif)
            ecrits += 1
        bilan = {"origine": origine, "ecrits": ecrits, "inchanges": inchanges,
                 "manuels_preserves": ignores_manuels, "non_classables": non_classables}
        await self.svc.audit("clientele.lot.classer", "clientele_lot", lid,
                             after={**bilan, "fichier": lot["fichier_nom"], "motif": motif})
        await self._controler(lot["lignes"])
        lot["stats"] = self._stats(lot["lignes"])
        lot["dernier_classement"] = bilan
        self._sauver(lot)
        return {**self._public(lot), "bilan": bilan}

    async def exporter(self, lid: uuid.UUID) -> bytes:
        lot = self._charger(lid)
        extra = lot.get("colonnes_affichees") or []
        lignes = []
        for x in lot["lignes"]:
            m = x.get("moteur") or {}
            lignes.append([x["ligne"], x.get("brut"), x.get("racine"), x.get("statut"),
                           x.get("nom") or x.get("nom_fichier"),
                           *[(x.get("colonnes") or {}).get(e) for e in extra],
                           x.get("niveau_fichier") or x.get("niveau_fichier_brut"),
                           x.get("classe_actuelle"), x.get("source_actuelle"),
                           m.get("score"), m.get("niveau"), m.get("statut"), m.get("motif")])
        entetes = ["Ligne", "Racine (fichier)", "Racine", "Statut", "Nom", *extra,
                   "Niveau fichier", "Classe actuelle", "Source actuelle",
                   "Score moteur", "Niveau moteur", "Statut moteur", "Motif moteur"]
        notes = [
            ["Fichier source", lot["fichier_nom"]],
            ["Moteur SCORE (CDC 1.0)", "Résultat de simulation tant qu'il n'est pas appliqué."],
            ["Statut INCONNUE", "Racine absente du référentiel clients (jamais créée ici)."],
        ]
        return build_styled_workbook_multi(
            report_title="Lot de classification",
            sheets=[("Lot", f"Lot — {lot['fichier_nom']}", entetes, lignes),
                    ("Lire", "Notes de lecture", ["Élément", "Explication"], notes)],
            subtitle=f"{len(lignes)} ligne(s) · fichier {lot['fichier_nom']}",
        )

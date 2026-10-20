"""Classification CLIENT (racine) : moteur de scoring CDC 1.0, décisions, historique.

Niveaux : FAIBLE / MOYEN / ELEVE / INTERDIT. La racine n'est jamais modifiable ici.

Politique de décision (directive du 10/10/2026) :
- le moteur de scoring (``scoring.evaluer``) propose ; seul un résultat EVALUE / BLOQUANT est admissible ;
- la classe d'un fichier reste une donnée source, distincte du niveau calculé ;
- une classification manuelle n'est jamais remplacée par un import ou le moteur sans forçage
  (permission ``clientele.classif.force`` + confirmation + motif) ; sinon un réexamen est ouvert ;
- un blocage INTERDIT du moteur ne s'abaisse qu'avec la même permission de forçage.

``appliquer_moteur`` est l'ancien moteur à règles saisies en base (versions), distinct du scoring.
"""

from __future__ import annotations

import uuid
from datetime import date
from io import BytesIO
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import func, or_, select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import aliased, selectinload

from app.core.config import get_settings
from app.core.exceptions import AppError
from app.models import (
    ClienteleAlerte,
    ClienteleClassifCritere,
    ClienteleClassifDivergence,
    ClienteleClassifDivergenceDecision,
    ClienteleClassifEvaluation,
    ClienteleClassifEvaluationLigne,
    ClienteleClassifHistorique,
    ClienteleClassifLot,
    ClienteleClassifLotLigne,
    ClienteleClassifNiveau,
    ClienteleClassifReexamen,
    ClienteleClassifRegle,
    ClienteleClassifVersion,
    ClienteleClassification,
    ClienteleClient,
    ClienteleCompte,
    ClienteleSituation,
    EerDossier,
    User,
)
from app.services.clientele.analystes import ClienteleAnalysteService
from app.services.clientele.journal import journaliser
from app.services.clientele.scoring import (
    SOURCES_LIBELLES, DonneesClient, Evaluation, codes_liste_interdiction, divergences_dict, evaluer,
    matrice_maitre_dict, nommer, valeurs_maitre_dict,
)
from app.services.clientele.service import Ctx, ClienteleService, maintenant
from app.services.reporting_export import build_styled_pdf, build_styled_workbook

NIVEAUX = ("FAIBLE", "MOYEN", "ELEVE", "INTERDIT")
SOURCES_MANUELLES = ("MANUEL", "MANUEL_SURCHARGE")
SOURCES_FICHIER = ("EXCEL", "EXCEL_CONFORMITE")
TAILLE_PAQUET_MASSE = 300


def autorisation_classement_masse() -> dict:
    s = get_settings()
    reference = (s.clientele_classement_masse_validation or "").strip()
    actif = bool(s.clientele_classement_masse_actif and reference)
    return {
        "actif": actif, "reference_validation": reference or None,
        "message": None if actif else (
            "Classement de masse suspendu côté serveur : les règles du moteur SCORE (liste ORION, comptes "
            "fermés, conflits Matrice / Fiche, complétude, doubles comptages) attendent la validation de la "
            "Conformité. La simulation reste disponible (lecture seule) ; le classement individuel motivé "
            "depuis la fiche client reste possible."),
    }


def exiger_classement_masse_autorise() -> str:
    """Verrou serveur de tout classement moteur en masse. Retourne la référence de validation (audit)."""
    a = autorisation_classement_masse()
    if not a["actif"]:
        raise AppError(a["message"], 409, code="CLASSEMENT_MASSE_SUSPENDU")
    return a["reference_validation"]
MODE_PAR_SOURCE = {
    "MOTEUR": "MOTEUR", "EXCEL": "FICHIER", "EXCEL_CONFORMITE": "FICHIER",
    "MANUEL": "MANUEL", "MANUEL_SURCHARGE": "MANUEL",
}
CHAMPS_TRACE = (
    "score_calcule", "niveau_propose", "classe_source_fichier", "mode_decision", "motif_moteur",
    "motif_source", "motif_decision_manuelle", "evaluation_id", "lot_id", "ligne_source",
    "date_maj_metier", "analyste_kyc", "date_reference", "version_moteur", "version_regles",
    "force", "blocage_leve",
)
STATUTS_DIVERGENCE = ("A_ARBITRER", "VALIDEE", "INVALIDE_SOURCE", "ALIGNEE")


def trace_evaluation(ev: dict) -> dict:
    """Champs de traçabilité issus d'un résultat moteur (dict ``Evaluation.to_dict`` ou ligne de lot)."""
    ref = ev.get("date_reference")
    return {
        "score_calcule": ev.get("score_total"),
        "niveau_propose": ev.get("niveau_final"),
        "motif_moteur": ev.get("motif_genere") or ev.get("motif_principal"),
        "evaluation_id": uuid.UUID(str(ev["evaluation_id"])) if ev.get("evaluation_id") else None,
        "date_reference": date.fromisoformat(ref) if isinstance(ref, str) else ref,
        "version_moteur": ev.get("version_moteur"),
        "version_regles": ev.get("version_regles"),
    }
CHAMPS_CLIENT = frozenset({
    "raison_sociale", "prenoms", "nationalite", "statut_resident", "agent_economique",
    "situation_juridique", "categorie_juridique", "secteur_activite", "famille_secteur_activite",
    "type_identifiant", "nni", "nif", "rcs", "type_client",
})
CHAMPS_COMPTE = frozenset({
    "etat_compte", "liste_interdiction", "conformite_compte", "devise", "ncg",
})
OPERATEURS = ("EGAL", "DIFFERENT", "IN", "NOT_IN", "CONTIENT", "VIDE", "NON_VIDE")


def _val(v: Any) -> str | None:
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def _match(operateur: str, observe: Any, attendu: Any) -> bool:
    obs = _val(observe)
    if operateur == "VIDE":
        return obs is None
    if operateur == "NON_VIDE":
        return obs is not None
    if operateur == "EGAL":
        return obs is not None and obs.upper() == str(attendu).strip().upper()
    if operateur == "DIFFERENT":
        return obs is None or obs.upper() != str(attendu).strip().upper()
    if operateur == "CONTIENT":
        return obs is not None and str(attendu).strip().upper() in obs.upper()
    valeurs = attendu if isinstance(attendu, (list, tuple)) else [attendu]
    cible = {str(x).strip().upper() for x in valeurs if x is not None and str(x).strip()}
    if operateur == "IN":
        return obs is not None and obs.upper() in cible
    if operateur == "NOT_IN":
        return obs is None or obs.upper() not in cible
    return False


def _habiller_modele(ws, largeurs: list[int]) -> None:
    """Style seulement : l'en-tête reste en ligne 1, lue telle quelle par ``analyser_excel``."""
    entete = PatternFill("solid", fgColor="1E3A5F")
    bord = Border(bottom=Side(style="thin", color="CBD5E1"))
    for cell in ws[1]:
        cell.fill = entete
        cell.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Calibri", size=10, color="0F172A")
            cell.border = bord
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 22
    for i, largeur in enumerate(largeurs, start=1):
        ws.column_dimensions[get_column_letter(i)].width = largeur
    ws.freeze_panes = "A2"


def _date_fr(iso: str | None, *, heure: bool = False) -> str | None:
    if not iso:
        return None
    jour = f"{iso[8:10]}/{iso[5:7]}/{iso[0:4]}"
    return f"{jour} {iso[11:16]}" if heure and len(iso) >= 16 else jour


def _xlsx(titre: str, entetes: list[str], lignes: list[list]) -> bytes:
    """Export tabulaire : racines et identifiants en texte (zéros de tête conservés)."""
    return build_styled_workbook(
        sheet_title=titre, report_title=titre, headers=entetes,
        rows=[[" " + v if isinstance(v, str) and v.startswith("=") else v for v in lg] for lg in lignes])


def _pdf(titre: str, entetes: list[str], lignes: list[list], *, sous_titre: str | None = None) -> bytes:
    return build_styled_pdf(report_title=titre, headers=entetes, rows=lignes, subtitle=sous_titre,
                            landscape_mode=len(entetes) > 6)


class ClienteleClassificationService:
    def __init__(self, db: AsyncSession, ctx: Ctx):
        self.db = db
        self.ctx = ctx
        self.svc = ClienteleService(db, ctx)

    async def referentiel(self) -> dict:
        self.ctx.exiger("clientele.classif.view")
        niveaux = list((await self.db.scalars(
            select(ClienteleClassifNiveau).order_by(ClienteleClassifNiveau.rang))).all())
        criteres = list((await self.db.scalars(
            select(ClienteleClassifCritere).order_by(ClienteleClassifCritere.code))).all())
        versions = list((await self.db.scalars(
            select(ClienteleClassifVersion).order_by(ClienteleClassifVersion.numero.desc()))).all())
        return {
            "niveaux": [{"code": n.code, "libelle": n.libelle, "rang": n.rang, "actif": n.actif} for n in niveaux],
            "criteres": [{"id": str(c.id), "code": c.code, "libelle": c.libelle,
                          "description": c.description, "champ_defaut": c.champ_defaut,
                          "actif": c.actif} for c in criteres],
            "versions": [self._version_dict(v, avec_regles=False) for v in versions],
            "operateurs": list(OPERATEURS),
            "champs_client": sorted(CHAMPS_CLIENT),
            "champs_compte": sorted(CHAMPS_COMPTE),
            "avertissement": (
                "Moteur SCORE uniquement (CDC 1.0). Les valeurs du référentiel maître ne sont pas ACTIVE. "
                "INTERDIT = blocage métier, le score reste calculé. Conflit Matrice des risques LBC/FT BEA / "
                "Fiche de scoring non arbitré : exclu de la somme. Donnée absente ≠ Faible. 0 critère → NON CLASSÉ."
            ),
            "matrice": matrice_maitre_dict()["meta"],
        }

    async def version_detail(self, vid: uuid.UUID) -> dict:
        self.ctx.exiger("clientele.classif.view")
        v = await self.db.scalar(
            select(ClienteleClassifVersion).where(ClienteleClassifVersion.id == vid)
            .options(selectinload(ClienteleClassifVersion.regles).selectinload(ClienteleClassifRegle.critere)))
        if not v:
            raise AppError("Version introuvable", 404, code="NOT_FOUND")
        return self._version_dict(v, avec_regles=True)

    async def creer_regle(self, vid: uuid.UUID, payload: dict) -> dict:
        self.ctx.exiger("clientele.classif.admin")
        v = await self.db.get(ClienteleClassifVersion, vid)
        if not v:
            raise AppError("Version introuvable", 404, code="NOT_FOUND")
        if v.statut == "ARCHIVEE":
            raise AppError("Version archivée", 409, code="VERSION_ARCHIVEE")
        r = self._regle_depuis(vid, payload)
        self.db.add(r)
        await self.db.flush()
        await self.svc.audit("clientele.classif.regle.create", "clientele_classif_regle", r.id,
                             after={"version": v.numero, "niveau": r.niveau_cible})
        return await self.version_detail(vid)

    async def modifier_regle(self, rid: uuid.UUID, payload: dict) -> dict:
        self.ctx.exiger("clientele.classif.admin")
        r = await self.db.get(ClienteleClassifRegle, rid)
        if not r:
            raise AppError("Règle introuvable", 404, code="NOT_FOUND")
        v = await self.db.get(ClienteleClassifVersion, r.version_id)
        if v and v.statut == "ARCHIVEE":
            raise AppError("Version archivée", 409, code="VERSION_ARCHIVEE")
        self._appliquer_payload(r, payload)
        await self.db.flush()
        await self.svc.audit("clientele.classif.regle.update", "clientele_classif_regle", rid)
        return await self.version_detail(r.version_id)

    async def supprimer_regle(self, rid: uuid.UUID) -> dict:
        self.ctx.exiger("clientele.classif.admin")
        r = await self.db.get(ClienteleClassifRegle, rid)
        if not r:
            raise AppError("Règle introuvable", 404, code="NOT_FOUND")
        vid = r.version_id
        await self.db.delete(r)
        await self.db.flush()
        return await self.version_detail(vid)

    async def activer_version(self, vid: uuid.UUID) -> dict:
        self.ctx.exiger("clientele.classif.admin")
        v = await self.db.get(ClienteleClassifVersion, vid)
        if not v:
            raise AppError("Version introuvable", 404, code="NOT_FOUND")
        actives = list((await self.db.scalars(
            select(ClienteleClassifVersion).where(ClienteleClassifVersion.statut == "ACTIVE"))).all())
        for autre in actives:
            if autre.id != v.id:
                autre.statut = "ARCHIVEE"
                autre.date_fin = date.today()
        v.statut = "ACTIVE"
        v.date_fin = None
        await self.db.flush()
        await self.svc.audit("clientele.classif.version.activer", "clientele_classif_version", vid,
                             after={"numero": v.numero})
        return await self.version_detail(vid)

    async def creer_version(self, payload: dict) -> dict:
        self.ctx.exiger("clientele.classif.admin")
        dernier = await self.db.scalar(select(func.max(ClienteleClassifVersion.numero))) or 0
        v = ClienteleClassifVersion(
            numero=int(dernier) + 1,
            libelle=(payload.get("libelle") or f"Version {int(dernier) + 1}")[:160],
            mode=payload.get("mode") if payload.get("mode") in ("MAX_NIVEAU", "SCORE") else "MAX_NIVEAU",
            seuils=payload.get("seuils"),
            date_effet=date.fromisoformat(payload["date_effet"]) if payload.get("date_effet") else date.today(),
            statut="BROUILLON",
            created_by_id=self.ctx.user.id,
        )
        self.db.add(v)
        await self.db.flush()
        source_id = payload.get("copier_version_id")
        if source_id:
            src = await self.db.scalar(
                select(ClienteleClassifVersion).where(ClienteleClassifVersion.id == uuid.UUID(str(source_id)))
                .options(selectinload(ClienteleClassifVersion.regles)))
            if src:
                for r in src.regles:
                    self.db.add(ClienteleClassifRegle(
                        version_id=v.id, critere_id=r.critere_id, priorite=r.priorite, poids=r.poids,
                        niveau_cible=r.niveau_cible, operateur=r.operateur, champ_source=r.champ_source,
                        portee=r.portee, valeur=r.valeur, motif=r.motif, actif=r.actif))
        await self.db.flush()
        return await self.version_detail(v.id)

    async def lister(self, *, niveau: str | None = None, source: str | None = None,
                     q: str | None = None, analyste: str | None = None,
                     page: int = 1, taille: int = 50) -> dict:
        self.ctx.exiger("clientele.classif.view")
        filtres = []
        if niveau:
            filtres.append(ClienteleClassification.niveau == niveau)
        if source == "MANUELLES":
            filtres.append(ClienteleClassification.source.in_(SOURCES_MANUELLES))
        elif source:
            filtres.append(ClienteleClassification.source == source)
        if analyste:
            filtres.append(ClienteleClassification.analyste_kyc.ilike(f"%{analyste.strip()}%"))
        stmt = (select(ClienteleClassification, ClienteleClient.raison_sociale)
                .join(ClienteleClient, ClienteleClient.racine_client == ClienteleClassification.racine_client))
        if filtres:
            stmt = stmt.where(*filtres)
        if q:
            like = f"%{q.strip()}%"
            stmt = stmt.where(or_(ClienteleClassification.racine_client.like(like),
                                  ClienteleClient.raison_sociale.ilike(like)))
        page, taille = max(1, page), min(max(1, taille), 200)
        total = await self.db.scalar(select(func.count()).select_from(stmt.subquery())) or 0
        ordre = ((ClienteleClassification.classifie_le.desc(), ClienteleClassification.racine_client)
                 if source == "MANUELLES" else
                 (ClienteleClassification.niveau.desc(), ClienteleClassification.racine_client))
        rows = (await self.db.execute(
            stmt.order_by(*ordre).offset((page - 1) * taille).limit(taille)
        )).all()
        return {
            "total": int(total), "page": page, "taille": taille,
            "items": [{**self._classif_dict(c), "raison_sociale": nom} for c, nom in rows],
        }

    async def rechercher(self, q: str, *, page: int = 1, taille: int = 25) -> dict:
        """Recherche par identifiant stable (racine, compte, RIB, NNI, NIF) ou par nom."""
        self.ctx.exiger("clientele.classif.view")
        q = (q or "").strip()
        if len(q) < 2:
            raise AppError("Saisissez au moins 2 caractères (racine, compte, NNI, NIF ou nom)", 422,
                           code="RECHERCHE_TROP_COURTE")
        C = ClienteleClient
        criteres = [C.raison_sociale.ilike(f"%{q}%"), C.prenoms.ilike(f"%{q}%")]
        if q.isdigit():
            criteres += [C.nni == q, C.nif == q, C.racine_client.in_(
                select(ClienteleCompte.racine_client).where(or_(ClienteleCompte.compte == q,
                                                                ClienteleCompte.rib == q)))]
            if len(q) <= 6:
                criteres.append(C.racine_client == q.zfill(6))
        nb_comptes = (select(func.count()).select_from(ClienteleCompte)
                      .where(ClienteleCompte.racine_client == C.racine_client).scalar_subquery())
        stmt = (select(C, ClienteleClassification, nb_comptes.label("nb_comptes"))
                .outerjoin(ClienteleClassification, ClienteleClassification.racine_client == C.racine_client)
                .where(or_(*criteres)))
        page, taille = max(1, page), min(max(1, taille), 100)
        total = await self.db.scalar(select(func.count()).select_from(stmt.subquery())) or 0
        rows = (await self.db.execute(stmt.order_by(C.racine_client).offset((page - 1) * taille)
                                      .limit(taille))).all()
        return {"total": int(total), "page": page, "taille": taille, "items": [{
            "racine_client": c.racine_client, "raison_sociale": c.raison_sociale, "prenoms": c.prenoms,
            "nni": c.nni, "nif": c.nif, "source_fiche": c.source, "nb_comptes": int(n or 0),
            "classification": ({"niveau": k.niveau, "source": k.source, "mode_decision": k.mode_decision,
                                "classifie_le": k.classifie_le.isoformat() if k.classifie_le else None}
                               if k else None),
        } for c, k, n in rows]}

    async def analystes(self) -> list[str]:
        """Analystes KYC actifs du référentiel (seules valeurs acceptées pour une décision)."""
        self.ctx.exiger("clientele.classif.view")
        return await ClienteleAnalysteService(self.db, self.ctx).noms_actifs()

    async def historique(self, racine: str) -> list[dict]:
        self.ctx.exiger("clientele.classif.view")
        racine = self._racine(racine)
        rows = list((await self.db.execute(
            select(ClienteleClassifHistorique, User.full_name)
            .outerjoin(User, User.id == ClienteleClassifHistorique.created_by_id)
            .where(ClienteleClassifHistorique.racine_client == racine)
            .order_by(ClienteleClassifHistorique.created_at.desc())
        )).all())
        return [{
            "id": str(h.id), "ancienne_classe": h.ancienne_classe, "nouvelle_classe": h.nouvelle_classe,
            "motif_risque": h.motif_risque, "motif_classement": h.motif_classement,
            "source": h.source, "version_id": str(h.version_id) if h.version_id else None,
            "detail": h.detail, "utilisateur": nom,
            "date": h.created_at.isoformat() if h.created_at else None,
            **self._trace_dict(h),
        } for h, nom in rows]

    def _exiger_forcage(self, confirmation: bool, motif: str | None) -> None:
        self.ctx.exiger("clientele.classif.force")
        if not confirmation:
            raise AppError("Forçage : confirmation explicite requise", 422, code="CONFIRMATION_REQUISE")
        if not (motif or "").strip():
            raise AppError("Forçage : motif obligatoire", 422, code="MOTIF_OBLIGATOIRE")

    async def decider_manuel(self, racine: str, niveau: str, motif: str, *, motif_risque: str | None = None,
                             forcer: bool = False, confirmation: bool = False,
                             trace: dict | None = None, type_action: str = "MODIFICATION",
                             groupe_id: uuid.UUID | None = None, detail_journal: dict | None = None) -> str:
        """Décision humaine. Motif obligatoire. INTERDIT moteur : abaissement seulement par forçage.

        Journalisée (``type_action`` : MODIFICATION ou ACCEPTATION d'une proposition) ; origine IMPORT
        si la décision porte sur une ligne de lot, SAISIE sinon."""
        self.ctx.exiger("clientele.classif.execute")
        niveau = (niveau or "").upper()
        if niveau not in NIVEAUX:
            raise AppError("Niveau invalide (FAIBLE, MOYEN, ELEVE, INTERDIT)", 422, code="NIVEAU_INVALIDE")
        motif = (motif or "").strip()
        if not motif:
            raise AppError("Motif de décision obligatoire", 422, code="MOTIF_OBLIGATOIRE")
        trace = dict(trace or {})
        precedente = await self.db.get(ClienteleClassification, racine)
        avant = (precedente.niveau, precedente.source) if precedente else (None, None)
        if "niveau_propose" not in trace:
            ev = await self.evaluer_racine(racine, persister=True)
            trace.update(trace_evaluation(ev))
            trace["_statut_moteur"] = ev.get("statut")
        blocage_leve = False
        if trace.pop("_statut_moteur", None) == "BLOQUANT" and niveau != "INTERDIT":
            if not forcer:
                raise AppError(
                    "Le moteur bloque ce client (critère INTERDIT). Un niveau inférieur exige un forçage "
                    "autorisé, confirmé et motivé.", 409, code="BLOCAGE_INTERDIT")
            self._exiger_forcage(confirmation, motif)
            blocage_leve = True
        await self._enregistrer(
            racine, niveau, source="MANUEL", version_id=None,
            motifs=[{"critere": "MANUEL", "motif": motif, "niveau": niveau}],
            motif_risque=motif_risque, motif_classement=motif,
            trace={**trace, "mode_decision": "MANUEL", "motif_decision_manuelle": motif,
                   "force": blocage_leve, "blocage_leve": blocage_leve})
        await self._clore_reexamens(racine, f"Décision manuelle {niveau} : {motif}")
        await journaliser(
            self.db, self.ctx, racine, type_action, origine="IMPORT" if trace.get("lot_id") else "SAISIE",
            champ="classification", avant=avant[0], apres=niveau, motif=motif, analyste=trace.get("analyste_kyc"),
            lot_id=trace.get("lot_id"), ligne=trace.get("ligne_source"), groupe_id=groupe_id,
            detail={"source_precedente": avant[1], "niveau_moteur": trace.get("niveau_propose"),
                    "classe_fichier": trace.get("classe_source_fichier"), "force": blocage_leve,
                    **(detail_journal or {})})
        await self.svc.audit(
            "clientele.classif.manuel.blocage_leve" if blocage_leve else "clientele.classif.manuel",
            "clientele_classification", racine,
            after={"niveau": niveau, "niveau_propose": trace.get("niveau_propose"),
                   "score": trace.get("score_calcule"), "lot_id": str(trace.get("lot_id") or "") or None,
                   "ligne": trace.get("ligne_source"), "type_action": type_action,
                   "groupe_id": str(groupe_id) if groupe_id else None})
        return "ECRIT"

    async def modifier_individuel(self, racine: str, payload: dict) -> dict:
        self.ctx.exiger("clientele.classif.execute")
        if (payload.get("racine_client") or "").strip() not in ("", racine.strip()):
            raise AppError("La racine client n'est pas modifiable depuis la classification", 422,
                           code="RACINE_IMMUTABLE")
        racine = self._racine(racine)
        await self._client_existe(racine)
        analyste = await ClienteleAnalysteService(self.db, self.ctx).exiger(payload.get("analyste_kyc"))
        await self.decider_manuel(
            racine, payload.get("niveau") or "", payload.get("motif_classement") or "",
            motif_risque=(payload.get("motif_risque") or "").strip() or None,
            forcer=bool(payload.get("forcer")), confirmation=bool(payload.get("confirmation")),
            trace={"analyste_kyc": analyste})
        return await self.courante(racine)

    async def ecrire_resultat_moteur(self, racine: str, ev: dict, motif: str, *, forcer: bool = False,
                                     confirmation: bool = False, trace: dict | None = None,
                                     preserver_fichier: bool = False) -> str:
        """Enregistre un résultat moteur déjà calculé. Retourne ECRIT / INCHANGE / MANUEL_PRESERVE /
        FICHIER_PRESERVE / NON_ADMISSIBLE. Ne remplace jamais un manuel sans forçage ; ouvre un réexamen
        à la place. ``preserver_fichier`` : idem pour une classe reprise d'un fichier (classement de masse)."""
        self.ctx.exiger("clientele.classif.execute")
        motif = (motif or "").strip()
        if not motif:
            raise AppError("Motif de classement obligatoire", 422, code="MOTIF_OBLIGATOIRE")
        niveau = ev.get("niveau_final")
        if ev.get("statut") not in ("EVALUE", "BLOQUANT") or niveau not in NIVEAUX:
            return "NON_ADMISSIBLE"
        trace = {**trace_evaluation(ev), **(trace or {})}
        actuelle = await self.db.get(ClienteleClassification, racine)
        force = False
        if actuelle and actuelle.source in SOURCES_MANUELLES:
            if not forcer:
                if actuelle.niveau != niveau:
                    await self.ouvrir_reexamen(
                        racine, origine="LOT" if trace.get("lot_id") else "MOTEUR", source_suggestion="MOTEUR",
                        niveau_suggere=niveau, actuelle=actuelle, lot_id=trace.get("lot_id"),
                        ligne=trace.get("ligne_source"),
                        motif=f"Le moteur propose {niveau} (score {ev.get('score_total')}) ; "
                              f"décision manuelle {actuelle.niveau} conservée.")
                return "MANUEL_PRESERVE"
            self._exiger_forcage(confirmation, motif)
            force = True
        elif actuelle and preserver_fichier and actuelle.source in SOURCES_FICHIER:
            if actuelle.niveau != niveau:
                await self.ouvrir_reexamen(
                    racine, origine="MOTEUR", source_suggestion="MOTEUR", niveau_suggere=niveau,
                    actuelle=actuelle,
                    motif=f"Le moteur propose {niveau} (score {ev.get('score_total')}) ; "
                          f"classe du fichier {actuelle.niveau} conservée.")
            return "FICHIER_PRESERVE"
        elif actuelle and actuelle.niveau == niveau and actuelle.source == "MOTEUR":
            return "INCHANGE"
        motifs = [{
            "critere": lg["critere"], "libelle": lg["libelle"], "famille": lg.get("famille"),
            "valeur": lg.get("valeur"), "poids": lg.get("poids"), "niveau": lg.get("niveau_retenu"),
            "motif": lg.get("motif"),
        } for lg in ev.get("lignes", []) if lg.get("contribue_au_score") or lg.get("blocking_propose")]
        motifs.append({"critere": "SCORE", "poids": ev.get("score_total"),
                       "version_regles": ev.get("version_regles"), "evaluation_id": ev.get("evaluation_id")})
        await self._enregistrer(
            racine, niveau, source="MOTEUR", version_id=None, motifs=motifs,
            motif_risque=ev.get("motif_genere"), motif_classement=motif,
            trace={**trace, "mode_decision": "MOTEUR", "force": force})
        if force:
            await self.svc.audit("clientele.classif.forcage", "clientele_classification", racine,
                                 before={"niveau": actuelle.niveau, "source": actuelle.source},
                                 after={"niveau": niveau, "mode": "MOTEUR", "motif": motif})
        return "ECRIT"

    async def classer_par_moteur(self, racine: str, payload: dict) -> dict:
        """Fiche client : évaluer avec le moteur de scoring puis enregistrer si admissible."""
        self.ctx.exiger("clientele.classif.execute")
        racine = self._racine(racine)
        ev = await self.evaluer_racine(racine, persister=True)
        avant = (ev.get("classification_courante") or {}).get("niveau")
        analyste = None
        if (payload.get("analyste_kyc") or "").strip():
            analyste = await ClienteleAnalysteService(self.db, self.ctx).exiger(payload.get("analyste_kyc"))
        resultat = await self.ecrire_resultat_moteur(
            racine, ev, payload.get("motif") or "", forcer=bool(payload.get("forcer")),
            confirmation=bool(payload.get("confirmation")), trace={"analyste_kyc": analyste})
        if resultat == "ECRIT":
            await journaliser(
                self.db, self.ctx, racine, "ACCEPTATION", origine="SAISIE", champ="classification",
                avant=avant, apres=ev.get("niveau_final"), motif=payload.get("motif"), analyste=analyste,
                detail={"proposition": "MOTEUR", "score": ev.get("score_total"),
                        "evaluation_id": ev.get("evaluation_id")})
        await self.svc.audit("clientele.classif.moteur_scoring", "clientele_classification", racine,
                             after={"resultat": resultat, "niveau_propose": ev.get("niveau_final"),
                                    "statut": ev.get("statut"), "score": ev.get("score_total")})
        return {"resultat": resultat, "evaluation": ev, **(await self.courante(racine))}

    def _perimetre_masse(self, payload: dict):
        q = select(ClienteleClient.racine_client)
        profil = (payload.get("profil") or "").strip().upper() or None
        agence = (payload.get("code_agence") or "").strip() or None
        if profil and profil not in ("PP", "PM", "NON_IDENTIFIE"):
            raise AppError("Profil attendu : PP, PM ou NON_IDENTIFIE", 422, code="PROFIL_INVALIDE")
        if profil or agence:
            q = q.join(ClienteleSituation, ClienteleSituation.racine_client == ClienteleClient.racine_client)
            if profil:
                q = q.where(ClienteleSituation.profil_derive == profil)
            if agence:
                q = q.where(ClienteleSituation.code_agence == agence)
        return q, {"profil": profil, "code_agence": agence}

    async def classer_masse_moteur(self, payload: dict) -> dict:
        """Classement de masse par le moteur SCORE, par paquets (curseur ``apres`` = dernière racine traitée).

        Résultats admissibles seulement. Classements MOTEUR mis à jour ; décisions manuelles et classes
        reprises d'un fichier jamais écrasées (réexamen ouvert si le moteur propose un autre niveau)."""
        self.ctx.exiger("clientele.classif.execute")
        exiger_classement_masse_autorise()
        motif = (payload.get("motif") or "").strip()
        if not motif:
            raise AppError("Motif de classement obligatoire", 422, code="MOTIF_OBLIGATOIRE")
        q, filtres = self._perimetre_masse(payload)
        apres = (payload.get("apres") or "").strip()
        taille = min(max(int(payload.get("taille") or TAILLE_PAQUET_MASSE), 1), 500)
        total = None
        if not apres:
            total = int(await self.db.scalar(select(func.count()).select_from(q.subquery())) or 0)
        racines = list((await self.db.scalars(
            q.where(ClienteleClient.racine_client > apres).order_by(ClienteleClient.racine_client)
            .limit(taille))).all())
        bilan, _ = await self.traiter_paquet_moteur(racines, motif, filtres)
        prochain = racines[-1] if len(racines) == taille else None
        return {"total": total, "traites": len(racines), "bilan": bilan, "prochain": prochain,
                "filtres": filtres}

    async def traiter_paquet_moteur(self, racines: list[str], motif: str, filtres: dict, *,
                                    isoler_erreurs: bool = False) -> tuple[dict, list[dict]]:
        """Évalue et enregistre un paquet de racines (mêmes règles que la fiche). Un audit par paquet.

        ``isoler_erreurs`` : chaque racine dans un point de sauvegarde ; une erreur technique est comptée
        (``ERREUR``) sans annuler les autres racines du paquet."""
        validation = exiger_classement_masse_autorise()
        bilan = {"ECRIT": 0, "INCHANGE": 0, "MANUEL_PRESERVE": 0, "FICHIER_PRESERVE": 0, "NON_ADMISSIBLE": 0}
        erreurs: list[dict] = []
        if isoler_erreurs:
            bilan["ERREUR"] = 0
        for r in racines:
            if not isoler_erreurs:
                ev = await self.evaluer_racine(r, persister=True)
                bilan[await self.ecrire_resultat_moteur(r, ev, motif, preserver_fichier=True)] += 1
                continue
            try:
                async with self.db.begin_nested():
                    ev = await self.evaluer_racine(r, persister=True)
                    resultat = await self.ecrire_resultat_moteur(r, ev, motif, preserver_fichier=True)
                bilan[resultat] += 1
            except Exception as exc:  # noqa: BLE001 — compté et tracé, le paquet continue
                bilan["ERREUR"] += 1
                detail = getattr(exc, "message", None) or getattr(exc, "detail", None) or exc.__class__.__name__
                erreurs.append({"racine": r, "message": str(detail)[:300]})
        await self.svc.audit("clientele.classif.moteur_masse", "clientele_classification", None,
                             after={**filtres, "motif": motif, "de": racines[0] if racines else None,
                                    "a": racines[-1] if racines else None, "validation": validation, **bilan})
        return bilan, erreurs

    async def ouvrir_reexamen(self, racine: str, *, origine: str, source_suggestion: str, niveau_suggere: str,
                              actuelle: ClienteleClassification | None, motif: str,
                              lot_id: uuid.UUID | None = None, ligne: int | None = None) -> bool:
        """Un seul signal OUVERT par (racine, source, niveau suggéré) : un import répété ne le duplique pas."""
        deja = await self.db.scalar(select(ClienteleClassifReexamen.id).where(
            ClienteleClassifReexamen.racine_client == racine,
            ClienteleClassifReexamen.source_suggestion == source_suggestion,
            ClienteleClassifReexamen.niveau_suggere == niveau_suggere,
            ClienteleClassifReexamen.statut == "OUVERT"))
        if deja:
            return False
        self.db.add(ClienteleClassifReexamen(
            racine_client=racine, origine=origine, lot_id=lot_id, ligne_source=ligne,
            niveau_actuel=actuelle.niveau if actuelle else None,
            source_actuelle=actuelle.source if actuelle else None,
            niveau_suggere=niveau_suggere, source_suggestion=source_suggestion, motif=motif[:2000],
            statut="OUVERT", created_by_id=self.ctx.user.id))
        await self.db.flush()
        return True

    async def _clore_reexamens(self, racine: str, decision: str) -> None:
        for r in (await self.db.scalars(select(ClienteleClassifReexamen).where(
                ClienteleClassifReexamen.racine_client == racine,
                ClienteleClassifReexamen.statut == "OUVERT"))).all():
            r.statut, r.decision = "TRAITE", decision[:2000]
            r.traite_par_id, r.traite_le = self.ctx.user.id, maintenant()
        await self.db.flush()

    async def lister_reexamens(self, *, statut: str | None = "OUVERT", racine: str | None = None,
                               page: int = 1, taille: int = 50) -> dict:
        self.ctx.exiger("clientele.classif.view")
        q = select(ClienteleClassifReexamen, ClienteleClient.raison_sociale).join(
            ClienteleClient, ClienteleClient.racine_client == ClienteleClassifReexamen.racine_client)
        if statut:
            q = q.where(ClienteleClassifReexamen.statut == statut)
        if racine:
            q = q.where(ClienteleClassifReexamen.racine_client == self._racine(racine))
        total = await self.db.scalar(select(func.count()).select_from(q.subquery())) or 0
        rows = (await self.db.execute(q.order_by(ClienteleClassifReexamen.created_at.desc())
                                      .offset((max(1, page) - 1) * taille).limit(min(taille, 200)))).all()
        return {"total": int(total), "page": page, "taille": taille, "items": [{
            "id": str(r.id), "racine_client": r.racine_client, "raison_sociale": nom, "origine": r.origine,
            "lot_id": str(r.lot_id) if r.lot_id else None, "ligne_source": r.ligne_source,
            "niveau_actuel": r.niveau_actuel, "source_actuelle": r.source_actuelle,
            "niveau_suggere": r.niveau_suggere, "source_suggestion": r.source_suggestion,
            "motif": r.motif, "statut": r.statut, "decision": r.decision,
            "created_at": r.created_at.isoformat() if r.created_at else None,
            "traite_le": r.traite_le.isoformat() if r.traite_le else None,
        } for r, nom in rows]}

    async def traiter_reexamen(self, rid: uuid.UUID, payload: dict) -> dict:
        """Clôture d'un réexamen. Changer le niveau passe par une décision manuelle (motif, historique)."""
        self.ctx.exiger("clientele.classif.arbitrer")
        r = await self.db.get(ClienteleClassifReexamen, rid)
        if not r:
            raise AppError("Réexamen introuvable", 404, code="NOT_FOUND")
        if r.statut != "OUVERT":
            raise AppError("Réexamen déjà traité", 409, code="REEXAMEN_TRAITE")
        decision = (payload.get("decision") or "").strip()
        if not decision:
            raise AppError("Décision motivée obligatoire", 422, code="MOTIF_OBLIGATOIRE")
        niveau = (payload.get("niveau") or "").upper() or None
        if niveau:
            await self.decider_manuel(r.racine_client, niveau, decision,
                                      forcer=bool(payload.get("forcer")),
                                      confirmation=bool(payload.get("confirmation")))
        if r.statut == "OUVERT":
            r.statut = "CLOS"
            r.decision = f"Décision manuelle conservée : {decision}"[:2000]
            r.traite_par_id, r.traite_le = self.ctx.user.id, maintenant()
        await self.db.flush()
        await self.svc.audit("clientele.classif.reexamen", "clientele_classif_reexamen", rid,
                             after={"statut": r.statut, "niveau": niveau, "decision": decision})
        return {"id": str(r.id), "statut": r.statut, "decision": r.decision}

    async def arbitrer_divergence(self, did: uuid.UUID, payload: dict) -> dict:
        self.ctx.exiger("clientele.classif.arbitrer")
        d = await self.db.get(ClienteleClassifDivergence, did)
        if not d:
            raise AppError("Divergence introuvable (référentiel non chargé en base ?)", 404, code="NOT_FOUND")
        statut = (payload.get("statut") or "").upper()
        if statut not in STATUTS_DIVERGENCE:
            raise AppError("Statut attendu : " + ", ".join(STATUTS_DIVERGENCE), 422, code="STATUT_INVALIDE")
        justification = (payload.get("justification") or "").strip()
        if not justification:
            raise AppError("Justification obligatoire", 422, code="MOTIF_OBLIGATOIRE")
        decision = (payload.get("decision") or "").strip() or None
        avant = d.statut
        self.db.add(ClienteleClassifDivergenceDecision(
            divergence_id=d.id, ancien_statut=avant, nouveau_statut=statut, decision=decision,
            justification=justification, created_by_id=self.ctx.user.id))
        d.statut, d.decision, d.justification = statut, decision, justification
        d.validee_par_id, d.validee_le = self.ctx.user.id, maintenant()
        await self.db.flush()
        await self.svc.audit("clientele.classif.divergence.arbitrer", "clientele_classif_divergence", did,
                             before={"statut": avant}, after={"statut": statut, "decision": decision})
        return {**(await self.divergence_historique(did)), "statut": statut}

    async def divergence_historique(self, did: uuid.UUID) -> dict:
        self.ctx.exiger("clientele.classif.view")
        rows = (await self.db.execute(
            select(ClienteleClassifDivergenceDecision, User.full_name)
            .outerjoin(User, User.id == ClienteleClassifDivergenceDecision.created_by_id)
            .where(ClienteleClassifDivergenceDecision.divergence_id == did)
            .order_by(ClienteleClassifDivergenceDecision.created_at.desc()))).all()
        return {"id": str(did), "decisions": [{
            "ancien_statut": x.ancien_statut, "nouveau_statut": x.nouveau_statut, "decision": x.decision,
            "justification": x.justification, "utilisateur": nom,
            "date": x.created_at.isoformat() if x.created_at else None,
        } for x, nom in rows]}

    async def _client_existe(self, racine: str) -> None:
        if not await self.db.scalar(select(ClienteleClient.racine_client)
                                    .where(ClienteleClient.racine_client == racine)):
            raise AppError("Client introuvable", 404, code="NOT_FOUND")

    async def courante(self, racine: str) -> dict:
        self.ctx.exiger("clientele.classif.view")
        racine = self._racine(racine)
        c = await self.db.get(ClienteleClassification, racine)
        calcul = await self.dernier_calcul(racine)
        return {"racine_client": racine, "classification": self._classif_dict(c) if c else None,
                "dernier_calcul": calcul, "ecart": self.ecart(c, calcul),
                "historique": await self.historique(racine)}

    async def dernier_calcul(self, racine: str) -> dict | None:
        """Dernière évaluation du moteur pour la racine (photographie), distincte de la classe retenue."""
        e = await self.db.scalar(
            select(ClienteleClassifEvaluation).where(ClienteleClassifEvaluation.racine_client == racine)
            .order_by(ClienteleClassifEvaluation.created_at.desc()).limit(1))
        return self._calcul_dict(e)

    @staticmethod
    def _calcul_dict(e: ClienteleClassifEvaluation | None) -> dict | None:
        if not e:
            return None
        return {
            "evaluation_id": str(e.id), "niveau": e.niveau_final, "statut": e.statut,
            "admissible": e.statut in ("EVALUE", "BLOQUANT") and e.niveau_final in NIVEAUX,
            "score": e.score_total, "version_moteur": e.version_moteur, "version_regles": e.version_regles,
            "date_calcul": e.created_at.isoformat() if e.created_at else None,
            "date_reference": e.date_reference.isoformat() if e.date_reference else None,
            "motif": e.motif_principal,
            "lot_id": str(e.lot_id) if e.lot_id else None, "ligne_source": e.ligne_source,
        }

    @staticmethod
    def ecart(retenue: ClienteleClassification | dict | None, calcul: dict | None) -> bool:
        """Vrai si le dernier calcul admissible du moteur diffère de la classe retenue."""
        if not retenue or not calcul or not calcul.get("admissible"):
            return False
        niveau = retenue.get("niveau") if isinstance(retenue, dict) else retenue.niveau
        return niveau != calcul.get("niveau")

    def _requete_resultats(self, *, ecart: bool | None = None, niveau: str | None = None,
                           statut: str | None = None, source: str | None = None, q: str | None = None):
        E = ClienteleClassifEvaluation
        dern = (select(E).distinct(E.racine_client)
                .order_by(E.racine_client, E.created_at.desc()).subquery())
        e = aliased(E, dern)
        K = ClienteleClassification
        stmt = (select(e, K, ClienteleClient.raison_sociale, ClienteleSituation.profil_derive)
                .join(ClienteleClient, ClienteleClient.racine_client == e.racine_client)
                .outerjoin(K, K.racine_client == e.racine_client)
                .outerjoin(ClienteleSituation, ClienteleSituation.racine_client == e.racine_client))
        admissible = e.statut.in_(("EVALUE", "BLOQUANT")) & e.niveau_final.in_(NIVEAUX)
        if ecart is True:
            stmt = stmt.where(admissible, K.niveau.is_not(None), K.niveau != e.niveau_final)
        elif ecart is False:
            stmt = stmt.where(or_(~admissible, K.niveau.is_(None), K.niveau == e.niveau_final))
        if niveau:
            stmt = stmt.where(e.niveau_final == niveau)
        if statut:
            stmt = stmt.where(e.statut == statut)
        if source == "MANUELLES":
            stmt = stmt.where(K.source.in_(SOURCES_MANUELLES))
        elif source:
            stmt = stmt.where(K.source == source)
        if q:
            like = f"%{q.strip()}%"
            stmt = stmt.where(or_(e.racine_client.like(like), ClienteleClient.raison_sociale.ilike(like)))
        return stmt, e

    def _resultat_dict(self, e, k, nom, profil) -> dict:
        calcul = self._calcul_dict(e)
        return {
            "racine_client": e.racine_client, "raison_sociale": nom, "profil": profil,
            "calcul": calcul, "retenue": self._classif_dict(k) if k else None, "ecart": self.ecart(k, calcul),
        }

    async def resultats(self, *, ecart: bool | None = None, niveau: str | None = None,
                        statut: str | None = None, source: str | None = None, q: str | None = None,
                        page: int = 1, taille: int = 50) -> dict:
        """Résultats du moteur : dernier calcul par racine, face à la classe retenue (écart signalé)."""
        self.ctx.exiger("clientele.classif.view")
        stmt, e = self._requete_resultats(ecart=ecart, niveau=niveau, statut=statut, source=source, q=q)
        page, taille = max(1, page), min(max(1, taille), 200)
        total = await self.db.scalar(select(func.count()).select_from(stmt.subquery())) or 0
        rows = (await self.db.execute(
            stmt.order_by(e.created_at.desc(), e.racine_client).offset((page - 1) * taille).limit(taille))).all()
        return {"total": int(total), "page": page, "taille": taille,
                "items": [self._resultat_dict(*r) for r in rows]}

    async def resultats_xlsx(self, **filtres) -> bytes:
        self.ctx.exiger("clientele.export")
        stmt, e = self._requete_resultats(**filtres)
        rows = (await self.db.execute(stmt.order_by(e.racine_client).limit(100_000))).all()
        lignes = []
        for r in rows:
            d = self._resultat_dict(*r)
            c, k = d["calcul"] or {}, d["retenue"] or {}
            lignes.append([
                d["racine_client"], d["profil"], d["raison_sociale"], c.get("niveau"), c.get("statut"),
                c.get("score"), c.get("version_moteur"), _date_fr(c.get("date_calcul")), k.get("niveau"),
                k.get("source"), k.get("analyste_kyc"), k.get("motif_decision_manuelle"),
                "Oui" if d["ecart"] else "Non",
            ])
        await self.svc.audit("clientele.classif.resultats.export", "clientele_classif_evaluation", None,
                             after={"lignes": len(lignes), **{k: v for k, v in filtres.items() if v is not None}})
        return _xlsx("Résultats moteur", [
            "Racine", "PP/PM", "Identité", "Classe moteur", "Statut moteur", "Score", "Version moteur",
            "Date du calcul", "Classe retenue", "Source", "Analyste KYC", "Motif manuel", "Écart",
        ], lignes)

    def _requete_historique(self, *, q: str | None = None, source: str | None = None,
                            analyste: str | None = None, lot_id: uuid.UUID | None = None,
                            du: date | None = None, au: date | None = None, profil: str | None = None):
        H = ClienteleClassifHistorique
        stmt = (select(H, User.full_name, ClienteleClient.raison_sociale, ClienteleSituation.profil_derive)
                .outerjoin(User, User.id == H.created_by_id)
                .outerjoin(ClienteleClient, ClienteleClient.racine_client == H.racine_client)
                .outerjoin(ClienteleSituation, ClienteleSituation.racine_client == H.racine_client))
        if profil in ("PP", "PM"):
            stmt = stmt.where(ClienteleSituation.profil_derive == profil)
        if q:
            like = f"%{q.strip()}%"
            stmt = stmt.where(or_(H.racine_client.like(like), ClienteleClient.raison_sociale.ilike(like)))
        if source == "MANUELLES":
            stmt = stmt.where(H.source.in_(SOURCES_MANUELLES))
        elif source:
            stmt = stmt.where(H.source == source)
        if analyste:
            stmt = stmt.where(H.analyste_kyc.ilike(f"%{analyste.strip()}%"))
        if lot_id:
            stmt = stmt.where(H.lot_id == lot_id)
        if du:
            stmt = stmt.where(func.date(H.created_at) >= du)
        if au:
            stmt = stmt.where(func.date(H.created_at) <= au)
        return stmt

    def _historique_dict(self, h, utilisateur, nom, profil) -> dict:
        return {
            "id": str(h.id), "racine_client": h.racine_client, "raison_sociale": nom, "profil": profil,
            "ancienne_classe": h.ancienne_classe, "nouvelle_classe": h.nouvelle_classe,
            "motif_risque": h.motif_risque, "motif_classement": h.motif_classement, "source": h.source,
            "utilisateur": utilisateur, "date": h.created_at.isoformat() if h.created_at else None,
            **self._trace_dict(h),
        }

    async def historique_global(self, *, q: str | None = None, source: str | None = None,
                                analyste: str | None = None, lot_id: uuid.UUID | None = None,
                                du: date | None = None, au: date | None = None, profil: str | None = None,
                                page: int = 1, taille: int = 50) -> dict:
        """Historique de toutes les classifications (table en ajout seul)."""
        self.ctx.exiger("clientele.classif.view")
        stmt = self._requete_historique(q=q, source=source, analyste=analyste, lot_id=lot_id, du=du, au=au,
                                        profil=profil)
        page, taille = max(1, page), min(max(1, taille), 200)
        total = await self.db.scalar(select(func.count()).select_from(stmt.subquery())) or 0
        rows = (await self.db.execute(
            stmt.order_by(ClienteleClassifHistorique.created_at.desc())
            .offset((page - 1) * taille).limit(taille))).all()
        return {"total": int(total), "page": page, "taille": taille,
                "items": [self._historique_dict(*r) for r in rows]}

    async def historique_xlsx(self, **filtres) -> bytes:
        self.ctx.exiger("clientele.export")
        stmt = self._requete_historique(**filtres)
        rows = (await self.db.execute(
            stmt.order_by(ClienteleClassifHistorique.created_at.desc()).limit(100_000))).all()
        lignes = []
        for r in rows:
            d = self._historique_dict(*r)
            lignes.append([
                d["racine_client"], d["profil"], d["raison_sociale"], d["ancienne_classe"], d["nouvelle_classe"],
                d["niveau_propose"], d["score_calcule"], d["version_moteur"], d["source"],
                d["mode_decision"] or MODE_PAR_SOURCE.get(d["source"] or ""),
                d["motif_decision_manuelle"] or d["motif_classement"],
                d["analyste_kyc"], d["utilisateur"], _date_fr(d["date"], heure=True),
                d["lot_id"], d["ligne_source"],
            ])
        await self.svc.audit("clientele.classif.historique.export", "clientele_classif_historique", None,
                             after={"lignes": len(lignes), "format": "xlsx",
                                    "filtres": sorted(k for k, v in filtres.items() if v)})
        return _xlsx("Historique classifications", [
            "Racine", "PP/PM", "Identité", "Ancienne classe", "Nouvelle classe", "Classe moteur", "Score",
            "Version moteur", "Source", "Mode", "Motif", "Analyste KYC", "Utilisateur", "Date", "Lot",
            "Ligne du fichier",
        ], lignes)

    async def historique_compter(self, **filtres) -> int:
        self.ctx.exiger("clientele.classif.view")
        stmt = self._requete_historique(**filtres)
        return int(await self.db.scalar(select(func.count()).select_from(stmt.subquery())) or 0)

    async def historique_pdf(self, **filtres) -> bytes:
        """Version imprimable : colonnes de décision seulement (les traces techniques restent dans l'Excel)."""
        self.ctx.exiger("clientele.export")
        total = await self.historique_compter(**filtres)
        if total > 5_000:
            raise AppError(f"{total} lignes : le PDF est limité à 5 000 lignes. Utilisez l'export Excel ou "
                           "affinez les filtres.", 422, code="EXPORT_TROP_VOLUMINEUX")
        stmt = self._requete_historique(**filtres)
        rows = (await self.db.execute(stmt.order_by(ClienteleClassifHistorique.created_at.desc()))).all()
        lignes = []
        for r in rows:
            d = self._historique_dict(*r)
            lignes.append([
                _date_fr(d["date"], heure=True), d["racine_client"], d["profil"], d["raison_sociale"],
                d["ancienne_classe"], d["nouvelle_classe"], d["niveau_propose"],
                d["mode_decision"] or MODE_PAR_SOURCE.get(d["source"] or ""),
                d["motif_decision_manuelle"] or d["motif_classement"], d["analyste_kyc"], d["utilisateur"],
            ])
        await self.svc.audit("clientele.classif.historique.export", "clientele_classif_historique", None,
                             after={"lignes": len(lignes), "format": "pdf",
                                    "filtres": sorted(k for k, v in filtres.items() if v)})
        profil = filtres.get("profil")
        return _pdf(f"Historique des classifications{' ' + profil if profil in ('PP', 'PM') else ''}", [
            "Date", "Racine", "PP/PM", "Identité", "Ancienne", "Nouvelle", "Classe moteur", "Mode", "Motif",
            "Analyste KYC", "Utilisateur",
        ], lignes)

    async def appliquer_moteur(self, *, racine: str | None = None, forcer: bool = False) -> dict:
        """Ancien moteur à règles. Une décision manuelle n'est jamais remplacée ici, même avec ``forcer`` :
        le remplacement d'un manuel passe par le forçage tracé (``ecrire_resultat_moteur``)."""
        self.ctx.exiger("clientele.classif.execute")
        if not racine:
            exiger_classement_masse_autorise()
        version = await self._version_active()
        regles = list((await self.db.scalars(
            select(ClienteleClassifRegle)
            .where(ClienteleClassifRegle.version_id == version.id, ClienteleClassifRegle.actif.is_(True))
            .options(selectinload(ClienteleClassifRegle.critere))
        )).all())
        if not regles:
            raise AppError(
                "Aucune règle active : le moteur ne classe pas tant que le métier n'a pas validé "
                "et saisi les règles (matrice / seuils).",
                422, code="AUCUNE_REGLE")
        rangs = {n.code: n.rang for n in (await self.db.scalars(select(ClienteleClassifNiveau))).all()}
        q = select(ClienteleClient)
        if racine:
            q = q.where(ClienteleClient.racine_client == self._racine(racine))
        clients = list((await self.db.scalars(q)).all())
        traites = ignores = changes = 0
        for client in clients:
            actuelle = await self.db.get(ClienteleClassification, client.racine_client)
            if actuelle and actuelle.source in SOURCES_MANUELLES:
                ignores += 1
                continue
            comptes = list((await self.db.scalars(
                select(ClienteleCompte).where(ClienteleCompte.racine_client == client.racine_client)
            )).all())
            resultat = self._evaluer(client, comptes, regles, rangs, version.mode, version.seuils)
            traites += 1
            if not resultat:
                continue
            if actuelle and actuelle.niveau == resultat["niveau"] and actuelle.version_id == version.id:
                continue
            await self._enregistrer(
                client.racine_client, resultat["niveau"], source="MOTEUR", version_id=version.id,
                motifs=resultat["motifs"], motif_risque=resultat["motif_risque"],
                motif_classement=resultat["motif_classement"])
            changes += 1
        await self.svc.audit("clientele.classif.moteur", "clientele_classif_version", version.id,
                             after={"traites": traites, "changes": changes, "ignores": ignores})
        return {"version": version.numero, "traites": traites, "modifies": changes,
                "manuels_preserves": ignores}

    async def modele_excel(self) -> bytes:
        self.ctx.exiger("clientele.classif.execute")
        wb = Workbook()
        ws = wb.active
        ws.title = "Classification"
        ws.append(["RACINE_CLIENT", "NIVEAU", "MOTIF_RISQUE", "MOTIF_CLASSEMENT"])
        ws.append(["000001", "ELEVE", "PPE (exemple)", "Saisie manuelle — règles à valider"])
        ws2 = wb.create_sheet("Niveaux")
        ws2.append(["Code", "Libellé"])
        for code in NIVEAUX:
            ws2.append([code, code])
        note = wb.create_sheet("Lire")
        note.append(["Notes de lecture"])
        note.append(["La colonne RACINE_CLIENT identifie le client. Elle n'est jamais modifiée."])
        note.append(["Niveaux autorisés : FAIBLE, MOYEN, ELEVE, INTERDIT."])
        note.append(["Les règles métier (matrice, listes, seuils) ne sont pas déduites de ce fichier."])
        _habiller_modele(ws, [18, 14, 36, 44])
        ws["A2"].number_format = "@"
        _habiller_modele(ws2, [16, 24])
        _habiller_modele(note, [90])
        buf = BytesIO()
        wb.save(buf)
        return buf.getvalue()

    async def analyser_excel(self, contenu: bytes) -> dict:
        self.ctx.exiger("clientele.classif.execute")
        try:
            wb = load_workbook(BytesIO(contenu), read_only=True, data_only=True)
        except Exception as exc:
            raise AppError("Fichier Excel illisible", 422, code="FICHIER_ILLISIBLE") from exc
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        wb.close()
        if not rows:
            raise AppError("Feuille vide", 422, code="FICHIER_VIDE")
        entetes = [str(c).strip().upper() if c else "" for c in rows[0]]
        idx = {h: i for i, h in enumerate(entetes)}
        if "RACINE_CLIENT" not in idx or "NIVEAU" not in idx:
            raise AppError("Colonnes obligatoires : RACINE_CLIENT, NIVEAU", 422, code="COLONNES_MANQUANTES")
        if any("RACINE" in h and h != "RACINE_CLIENT" for h in entetes):
            raise AppError("Aucune colonne ne peut modifier la racine", 422, code="RACINE_IMMUTABLE")
        valides: list[dict] = []
        rejets: list[dict] = []
        vues: set[str] = set()
        for numero, row in enumerate(rows[1:], start=2):
            if not row or all(c is None or str(c).strip() == "" for c in row):
                continue
            racine = str(row[idx["RACINE_CLIENT"]] or "").strip().zfill(6)[:6]
            niveau = str(row[idx["NIVEAU"]] or "").strip().upper().replace("É", "E")
            if niveau == "ELEVEE":
                niveau = "ELEVE"
            motif_r = str(row[idx["MOTIF_RISQUE"]]).strip() if "MOTIF_RISQUE" in idx and row[idx["MOTIF_RISQUE"]] else None
            motif_c = str(row[idx["MOTIF_CLASSEMENT"]]).strip() if "MOTIF_CLASSEMENT" in idx and row[idx["MOTIF_CLASSEMENT"]] else None
            if len(racine) != 6 or not racine.isdigit():
                rejets.append({"ligne": numero, "motif": "racine invalide"})
                continue
            if niveau not in NIVEAUX:
                rejets.append({"ligne": numero, "racine": racine, "motif": "niveau invalide"})
                continue
            if racine in vues:
                rejets.append({"ligne": numero, "racine": racine, "motif": "doublon dans le fichier"})
                continue
            if not await self.db.scalar(select(ClienteleClient.racine_client)
                                        .where(ClienteleClient.racine_client == racine)):
                rejets.append({"ligne": numero, "racine": racine, "motif": "client absent du référentiel"})
                continue
            vues.add(racine)
            valides.append({"racine_client": racine, "niveau": niveau,
                            "motif_risque": motif_r, "motif_classement": motif_c or "Import Excel contrôlé"})
        return {"nb_valides": len(valides), "nb_rejets": len(rejets), "valides": valides[:200],
                "rejets": rejets[:200], "lignes": valides}

    async def confirmer_excel(self, lignes: list[dict]) -> dict:
        self.ctx.exiger("clientele.classif.execute")
        if not lignes:
            raise AppError("Aucune ligne à importer", 422, code="AUCUNE_LIGNE")
        n = preserves = 0
        for ligne in lignes:
            racine = self._racine(ligne["racine_client"])
            niveau = ligne["niveau"]
            if niveau not in NIVEAUX:
                continue
            actuelle = await self.db.get(ClienteleClassification, racine)
            if actuelle and actuelle.source in SOURCES_MANUELLES:
                preserves += 1
                if actuelle.niveau != niveau:
                    await self.ouvrir_reexamen(
                        racine, origine="EXCEL", source_suggestion="FICHIER", niveau_suggere=niveau,
                        actuelle=actuelle, motif=f"Import Excel : {niveau} ; décision manuelle "
                                                 f"{actuelle.niveau} conservée.")
                continue
            await self._enregistrer(
                racine, niveau, source="EXCEL", version_id=None,
                motifs=[{"critere": "EXCEL", "motif": ligne.get("motif_classement"), "niveau": niveau}],
                motif_risque=ligne.get("motif_risque"),
                motif_classement=ligne.get("motif_classement") or "Import Excel contrôlé",
                trace={"classe_source_fichier": niveau, "motif_source": ligne.get("motif_risque")})
            n += 1
        await self.svc.audit("clientele.classif.excel", "clientele_classification", None,
                             after={"nb": n, "manuels_preserves": preserves})
        return {"modifies": n, "manuels_preserves": preserves}

    def _evaluer(self, client: ClienteleClient, comptes: list[ClienteleCompte],
                 regles: list[ClienteleClassifRegle], rangs: dict[str, int],
                 mode: str, seuils: dict | None) -> dict | None:
        matches: list[tuple[ClienteleClassifRegle, ClienteleClassifCritere | None]] = []
        for r in sorted(regles, key=lambda x: x.priorite):
            if self._regle_matche(r, client, comptes):
                matches.append((r, r.critere if hasattr(r, "critere") else None))
        if not matches:
            return None
        if any(r.niveau_cible == "INTERDIT" for r, _ in matches):
            retenues = [(r, c) for r, c in matches if r.niveau_cible == "INTERDIT"]
            niveau = "INTERDIT"
        elif mode == "SCORE":
            score = sum(r.poids for r, _ in matches)
            niveau = self._niveau_depuis_score(score, seuils, rangs)
            retenues = matches
        else:
            niveau = max((r.niveau_cible for r, _ in matches), key=lambda n: rangs.get(n, 0))
            retenues = [(r, c) for r, c in matches if r.niveau_cible == niveau] or matches
        motifs = [{
            "critere": (c.code if c else r.champ_source),
            "libelle": (c.libelle if c else r.champ_source),
            "motif": r.motif, "niveau": r.niveau_cible, "poids": r.poids,
        } for r, c in retenues]
        return {
            "niveau": niveau,
            "motifs": motifs,
            "motif_risque": " ; ".join(m["libelle"] for m in motifs),
            "motif_classement": " ; ".join(m["motif"] for m in motifs),
        }

    def _regle_matche(self, r: ClienteleClassifRegle, client: ClienteleClient,
                      comptes: list[ClienteleCompte]) -> bool:
        if r.portee == "COMPTE" or r.champ_source in CHAMPS_COMPTE:
            if r.operateur in ("VIDE",) and not comptes:
                return True
            return any(_match(r.operateur, getattr(c, r.champ_source, None), r.valeur) for c in comptes)
        if r.champ_source == "date_ouverture":
            dates = [c.date_ouverture for c in comptes if c.date_ouverture]
            observe = min(dates).isoformat() if dates else None
            return _match(r.operateur, observe, r.valeur)
        if r.champ_source not in CHAMPS_CLIENT:
            return False
        return _match(r.operateur, getattr(client, r.champ_source, None), r.valeur)

    @staticmethod
    def _niveau_depuis_score(score: int, seuils: dict | None, rangs: dict[str, int]) -> str:
        s = seuils or {}
        try:
            if score >= int(s.get("INTERDIT", 10**9)):
                return "INTERDIT"
            if score >= int(s.get("ELEVE", 10**9)):
                return "ELEVE"
            if score >= int(s.get("MOYEN", 10**9)):
                return "MOYEN"
        except (TypeError, ValueError):
            pass
        return "FAIBLE" if "FAIBLE" in rangs else "MOYEN"

    async def _enregistrer(self, racine: str, niveau: str, *, source: str,
                           version_id: uuid.UUID | None, motifs: list, motif_risque: str | None,
                           motif_classement: str | None, trace: dict | None = None) -> None:
        """Écrit l'historique puis la classification courante. Les champs de trace sont tous réécrits :
        une décision ne garde jamais le score ou le lot d'une décision antérieure."""
        trace = {k: v for k, v in (trace or {}).items() if k in CHAMPS_TRACE}
        trace.setdefault("mode_decision", MODE_PAR_SOURCE.get(source))
        valeurs = {k: trace.get(k, False if k in ("force", "blocage_leve") else None) for k in CHAMPS_TRACE}
        actuelle = await self.db.get(ClienteleClassification, racine)
        ancienne = actuelle.niveau if actuelle else None
        self.db.add(ClienteleClassifHistorique(
            racine_client=racine, ancienne_classe=ancienne, nouvelle_classe=niveau,
            motif_risque=motif_risque, motif_classement=motif_classement, source=source,
            version_id=version_id, detail={"motifs": motifs, "ancienne_source": actuelle.source if actuelle else None},
            created_by_id=self.ctx.user.id, **valeurs))
        now = maintenant()
        if actuelle:
            actuelle.niveau = niveau
            actuelle.version_id = version_id
            actuelle.motifs = motifs
            actuelle.source = source
            actuelle.motif_risque = motif_risque
            actuelle.motif_classement = motif_classement
            actuelle.classifie_le = now
            actuelle.classifie_par_id = self.ctx.user.id
            actuelle.updated_at = now
            for k, v in valeurs.items():
                setattr(actuelle, k, v)
        else:
            self.db.add(ClienteleClassification(
                racine_client=racine, niveau=niveau, version_id=version_id, motifs=motifs,
                source=source, motif_risque=motif_risque, motif_classement=motif_classement,
                classifie_le=now, classifie_par_id=self.ctx.user.id, updated_at=now, **valeurs))
        await self.db.flush()

    async def _version_active(self) -> ClienteleClassifVersion:
        v = await self.db.scalar(
            select(ClienteleClassifVersion).where(ClienteleClassifVersion.statut == "ACTIVE")
            .order_by(ClienteleClassifVersion.date_effet.desc()).limit(1))
        if not v:
            raise AppError("Aucune version ACTIVE de règles", 422, code="AUCUNE_VERSION")
        return v

    def _regle_depuis(self, vid: uuid.UUID, payload: dict) -> ClienteleClassifRegle:
        r = ClienteleClassifRegle(version_id=vid, motif="règle")
        self._appliquer_payload(r, payload)
        return r

    def _appliquer_payload(self, r: ClienteleClassifRegle, payload: dict) -> None:
        critere_id = payload.get("critere_id")
        if not critere_id:
            raise AppError("critere_id obligatoire", 422, code="CRITERE_OBLIGATOIRE")
        niveau = (payload.get("niveau_cible") or "").upper()
        if niveau not in NIVEAUX:
            raise AppError("niveau_cible invalide", 422, code="NIVEAU_INVALIDE")
        op = (payload.get("operateur") or "").upper()
        if op not in OPERATEURS:
            raise AppError("opérateur invalide", 422, code="OPERATEUR_INVALIDE")
        champ = payload.get("champ_source") or "nationalite"
        portee = payload.get("portee") if payload.get("portee") in ("CLIENT", "COMPTE") else (
            "COMPTE" if champ in CHAMPS_COMPTE else "CLIENT")
        r.critere_id = uuid.UUID(str(critere_id))
        r.niveau_cible = niveau
        r.operateur = op
        r.champ_source = str(champ)[:40]
        r.portee = portee
        r.valeur = payload.get("valeur")
        r.motif = (payload.get("motif") or "").strip() or f"{champ} {op}"
        r.priorite = int(payload.get("priorite") or 100)
        r.poids = int(payload.get("poids") or 0)
        r.actif = bool(payload.get("actif", True))

    @staticmethod
    def _racine(valeur: str) -> str:
        racine = (valeur or "").strip()
        if len(racine) != 6 or not racine.isdigit():
            raise AppError("Racine client invalide (6 chiffres)", 422, code="RACINE_INVALIDE")
        return racine

    @staticmethod
    def _trace_dict(o: Any) -> dict:
        out = {}
        for k in CHAMPS_TRACE:
            v = getattr(o, k, None)
            out[k] = v.isoformat() if isinstance(v, date) else (str(v) if isinstance(v, uuid.UUID) else v)
        return out

    @classmethod
    def _classif_dict(cls, c: ClienteleClassification | None) -> dict | None:
        if not c:
            return None
        return {
            "racine_client": c.racine_client, "niveau": c.niveau, "niveau_retenu": c.niveau,
            "source": c.source, "motifs": c.motifs or [],
            "motif_risque": c.motif_risque, "motif_classement": c.motif_classement,
            "version_id": str(c.version_id) if c.version_id else None,
            "classifie_le": c.classifie_le.isoformat() if c.classifie_le else None,
            **cls._trace_dict(c),
        }

    async def matrice_maitre(self) -> dict:
        self.ctx.exiger("clientele.classif.view")
        return matrice_maitre_dict()

    async def valeurs_maitres(self, *, dimension: str | None = None, statut: str | None = None,
                              page: int = 1, taille: int = 80) -> dict:
        self.ctx.exiger("clientele.classif.view")
        out = valeurs_maitre_dict(dimension=dimension, statut=statut, page=page, taille=taille)
        try:
            from app.models.clientele import ClienteleClassifValeur
            q = select(func.count()).select_from(ClienteleClassifValeur)
            n = await self.db.scalar(q)
            out["en_base"] = int(n or 0)
        except Exception:
            out["en_base"] = None
        return out

    async def divergences(self, *, domaine: str | None = None, page: int = 1, taille: int = 80) -> dict:
        self.ctx.exiger("clientele.classif.view")
        q = select(ClienteleClassifDivergence).order_by(
            ClienteleClassifDivergence.domaine, ClienteleClassifDivergence.cle)
        if domaine:
            q = q.where(ClienteleClassifDivergence.domaine == domaine)
        try:
            total = await self.db.scalar(select(func.count()).select_from(q.subquery())) or 0
            rows = list((await self.db.scalars(q.offset((page - 1) * taille).limit(min(taille, 200)))).all())
            if rows:
                return {
                    "total": int(total), "page": page, "taille": taille,
                    "sources_libelles": SOURCES_LIBELLES,
                    "items": [{
                        "id": str(d.id), "domaine": d.domaine, "cle": d.cle,
                        "source_a": d.source_a, "valeur_a": d.valeur_a,
                        "source_a_libelle": SOURCES_LIBELLES.get(d.source_a or "", d.source_a),
                        "source_b": d.source_b, "valeur_b": d.valeur_b,
                        "source_b_libelle": SOURCES_LIBELLES.get(d.source_b or "", d.source_b),
                        "statut": d.statut, "note": nommer(d.note), "decision": nommer(d.decision),
                    } for d in rows],
                }
        except Exception:
            pass
        brut = divergences_dict()
        items = brut["items"]
        if domaine:
            items = [x for x in items if x.get("domaine") == domaine]
        debut = (page - 1) * taille
        return {"total": len(items), "page": page, "taille": taille,
                "sources_libelles": SOURCES_LIBELLES,
                "items": [
                    {**x, "source_a_libelle": SOURCES_LIBELLES.get(x.get("source_a") or "", x.get("source_a")),
                     "source_b_libelle": SOURCES_LIBELLES.get(x.get("source_b") or "", x.get("source_b"))}
                    for x in items[debut:debut + taille]
                ]}

    async def evaluer_racine(self, racine: str, *, persister: bool = True, complements: dict | None = None,
                             lot_id: uuid.UUID | None = None, ligne: int | None = None) -> dict:
        """Évalue sans écrire la classification retenue (snapshot seulement).

        ``complements`` : données de la ligne de lot (date d'ouverture, Entreprise ouverture récente).
        Sans compléments, la dernière valeur connue d'un lot est relue : individuel et lot concordent.
        """
        self.ctx.exiger("clientele.classif.view")
        racine = self._racine(racine)
        client = await self.db.scalar(
            select(ClienteleClient).where(ClienteleClient.racine_client == racine))
        if not client:
            raise AppError("Client introuvable", 404, code="NOT_FOUND")
        comptes = list((await self.db.scalars(
            select(ClienteleCompte).where(ClienteleCompte.racine_client == racine))).all())
        donnees = await self._collecter(client, comptes, complements)
        ev = evaluer(donnees)
        payload = ev.to_dict()
        payload["donnees_source"] = self._donnees_dict(donnees)
        if persister:
            payload["evaluation_id"] = str(await self._photographier(ev, payload["donnees_source"],
                                                                     lot_id=lot_id, ligne=ligne))
        payload["classification_courante"] = self._classif_dict(
            await self.db.get(ClienteleClassification, racine))
        return payload

    async def _photographier(self, ev: Evaluation, donnees: dict, *, lot_id: uuid.UUID | None,
                             ligne: int | None) -> uuid.UUID:
        rec = ClienteleClassifEvaluation(
            racine_client=ev.racine,
            version_moteur=ev.version_moteur,
            version_regles=ev.version_regles,
            mode=ev.mode,
            score_total=ev.score_total,
            nb_evalues=ev.nb_evalues,
            niveau_score=ev.niveau_score,
            niveau_score_v4=ev.niveau_score_v4,
            niveau_max=ev.niveau_max,
            niveau_final=ev.niveau_final,
            statut=ev.statut,
            coherence=ev.coherence,
            motif_principal=ev.motif_principal,
            date_reference=date.fromisoformat(ev.date_reference) if ev.date_reference else None,
            source_date_reference=(ev.source_date_reference or "")[:20] or None,
            lot_id=lot_id,
            ligne_source=ligne,
            detail={
                "autres_facteurs": ev.autres_facteurs,
                "blocking_proposes": ev.blocking_proposes,
                "divergences": ev.divergences,
                "avertissement": ev.avertissement,
                "scores_familles": ev.scores_familles,
                "motif_genere": ev.motif_genere,
                "donnees_source": donnees,
            },
            created_by_id=self.ctx.user.id,
        )
        self.db.add(rec)
        await self.db.flush()
        champs = ("critere", "libelle", "etat", "valeur", "source_donnee", "poids", "niveau_matrice",
                  "niveau_v4", "niveau_retenu", "type_decision", "motif", "divergence", "blocking_propose",
                  "statut_regle", "famille", "contribue_au_score")
        for lg in ev.lignes:
            d = lg.to_dict()
            self.db.add(ClienteleClassifEvaluationLigne(evaluation_id=rec.id, **{k: d[k] for k in champs}))
        await self.db.flush()
        return rec.id

    @staticmethod
    def _donnees_dict(d: DonneesClient) -> dict:
        out = {}
        for k, v in vars(d).items():
            out[k] = v.isoformat() if isinstance(v, date) else v
        return out

    async def backtester(self, *, limite: int = 50) -> dict:
        """Compare le moteur proposition vs classification déjà stockée. Pas d'écriture de classe."""
        self.ctx.exiger("clientele.classif.execute")
        limite = max(1, min(int(limite), 500))
        racines = list((await self.db.scalars(
            select(ClienteleClassification.racine_client).order_by(
                ClienteleClassification.racine_client).limit(limite)
        )).all())
        if not racines:
            racines = list((await self.db.scalars(
                select(ClienteleClient.racine_client).order_by(ClienteleClient.racine_client).limit(limite)
            )).all())
        compteurs: dict[str, int] = {}
        echantillon: list[dict] = []
        for racine in racines:
            ev = await self.evaluer_racine(racine, persister=False)
            hist = (ev.get("classification_courante") or {}).get("niveau")
            nouveau = ev.get("niveau_final") or "NON_CLASSE"
            cle = f"{hist or 'ABSENT'}→{nouveau}"
            compteurs[cle] = compteurs.get(cle, 0) + 1
            if hist and nouveau != "NON_CLASSE" and hist != nouveau and len(echantillon) < 30:
                echantillon.append({
                    "racine_client": racine, "historique": hist, "moteur": nouveau,
                    "score": ev.get("score_total"), "coherence": ev.get("coherence"),
                    "motif": ev.get("motif_principal"), "statut": ev.get("statut"),
                })
        return {
            "limite": limite, "nb": len(racines), "compteurs": compteurs,
            "echantillon_ecarts": echantillon,
            "avertissement": (
                "Backtest sur copie de test uniquement. Moteur SCORE CDC 1.0. "
                "Aucune classification n'est écrasée."
            ),
        }

    async def _complements_connus(self, racine: str) -> dict:
        """Dernière date d'ouverture / valeur « Entreprise ouverture récente » apportée par un lot."""
        out: dict = {}
        L = ClienteleClassifLotLigne
        recence = func.coalesce(L.date_ouverture_saisie_le, ClienteleClassifLot.created_at)
        base = (select(L).join(ClienteleClassifLot, ClienteleClassifLot.id == L.lot_id)
                .where(L.racine_client == racine, L.statut == "OK"))
        lg = await self.db.scalar(base.where(L.date_ouverture_entreprise.is_not(None))
                                  .order_by(recence.desc(), L.ligne.desc()).limit(1))
        if lg:
            out["date_ouverture_entreprise"] = lg.date_ouverture_entreprise
            out["source_date_ouverture"] = lg.source_date_ouverture or "FICHIER"
        lg = await self.db.scalar(base.where(L.entreprise_recente_brut.is_not(None))
                                  .order_by(ClienteleClassifLot.created_at.desc(), L.ligne.desc()).limit(1))
        if lg:
            out["entreprise_recente"] = lg.entreprise_recente_brut
        return out

    async def _collecter(self, client: ClienteleClient, comptes: list[ClienteleCompte],
                         complements: dict | None = None) -> DonneesClient:
        if complements is None:
            complements = await self._complements_connus(client.racine_client)
        listes = codes_liste_interdiction(c.liste_interdiction for c in comptes)
        ppe = procuration = pays = origine = eer_op = date_creation_eer = None
        eer_present = False
        try:
            dossier = await self.db.scalar(
                select(EerDossier)
                .where(EerDossier.racine_client == client.racine_client, EerDossier.deleted_at.is_(None))
                .options(selectinload(EerDossier.client), selectinload(EerDossier.parties))
                .order_by(EerDossier.date_eer.desc(), EerDossier.created_at.desc())
                .limit(1))
            if dossier:
                eer_present = True
                eer_op = dossier.operation_type
                origine = dossier.origine_fonds
                if dossier.ppe_dossier:
                    ppe = True
                parties = list(dossier.parties or [])
                for p in parties:
                    if p.role == "CLIENT" and p.ppe is not None:
                        ppe = p.ppe
                    if p.role == "MANDATAIRE" or p.forme_mandat == "PROCURATION":
                        procuration = True
                if procuration is None and parties:
                    procuration = False
                if getattr(dossier, "client", None) and getattr(dossier.client, "pays_residence", None):
                    pays = dossier.client.pays_residence
                morale = getattr(getattr(dossier, "client", None), "morale", None)
                if morale is not None and morale.date_creation:
                    date_creation_eer = morale.date_creation
        except Exception:
            eer_present = False
        filtrage = None
        try:
            conf = await self.db.scalar(
                select(ClienteleAlerte.id).where(
                    ClienteleAlerte.racine_client == client.racine_client,
                    ClienteleAlerte.statut == "CONFIRMEE").limit(1))
            if conf:
                filtrage = True
        except Exception:
            filtrage = None
        date_ouv = complements.get("date_ouverture_entreprise")
        source_date = complements.get("source_date_ouverture") if date_ouv else None
        if not date_ouv and date_creation_eer:
            date_ouv, source_date = date_creation_eer, "EER"
        declare = complements.get("entreprise_recente")
        return DonneesClient(
            date_creation_pm=date_ouv,
            source_date_creation=source_date,
            entreprise_recente_declaree=declare,
            source_entreprise_recente="DECLARATIF_KYC" if declare not in (None, "") else None,
            date_reference=client.date_extraction,
            source_date_reference="EXTRACTION_ORION" if client.date_extraction else None,
            racine=client.racine_client,
            nationalite=client.nationalite,
            statut_resident=client.statut_resident,
            categorie_juridique=client.categorie_juridique,
            situation_juridique=client.situation_juridique,
            secteur_activite=client.secteur_activite,
            famille_secteur=client.famille_secteur_activite,
            agent_economique=client.agent_economique,
            type_client=client.type_client,
            date_naissance=client.date_naissance,
            liste_interdiction=", ".join(listes) if listes else ("" if comptes else None),
            ppe=ppe,
            procuration=procuration,
            pays_residence=pays,
            origine_fonds=origine,
            filtrage_confirme=filtrage,
            eer_present=eer_present,
            eer_operation=eer_op,
        )

    def _version_dict(self, v: ClienteleClassifVersion, *, avec_regles: bool) -> dict:
        out = {
            "id": str(v.id), "numero": v.numero, "libelle": v.libelle, "mode": v.mode,
            "seuils": v.seuils, "date_effet": v.date_effet.isoformat() if v.date_effet else None,
            "date_fin": v.date_fin.isoformat() if v.date_fin else None, "statut": v.statut,
        }
        if avec_regles:
            out["regles"] = [{
                "id": str(r.id), "critere_id": str(r.critere_id),
                "critere": r.critere.code if r.critere else None,
                "critere_libelle": r.critere.libelle if r.critere else None,
                "priorite": r.priorite, "poids": r.poids, "niveau_cible": r.niveau_cible,
                "operateur": r.operateur, "champ_source": r.champ_source, "portee": r.portee,
                "valeur": r.valeur, "motif": r.motif, "actif": r.actif,
            } for r in v.regles]
        return out

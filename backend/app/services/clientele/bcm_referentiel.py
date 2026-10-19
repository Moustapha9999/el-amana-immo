"""Référentiels BCM : versions du référentiel officiel (feuille ``ref``), correspondances agences et secteurs,
compléments KYC SIC (genre, téléphone, codes CSP / secteur BCM).

Une version n'est jamais modifiée après import ; une version utilisée par une déclaration n'est jamais supprimée.
Les compléments KYC ne touchent ni la classification, ni les comptes, ni les données ORION.
"""

from __future__ import annotations

import hashlib
import io
import re
import uuid
from datetime import datetime, timezone

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from sqlalchemy import func, select, text
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.exceptions import AppError
from app.data.clientele_sic import SECTEURS_GENERIQUES, lire_referentiel, norm, normaliser_telephone
from app.models import (
    Agence,
    ClienteleBcmRefCode,
    ClienteleBcmRefVersion,
    ClienteleDeclarationSic,
    ClienteleKyc,
    ClienteleSicAgence,
    ClienteleSicSecteur,
)
from app.services.clientele.service import Ctx, ClienteleService

SOURCE_COMPLEMENT = "SIC_COMPLEMENT"
COLONNES_COMPLEMENT = ("CodeClient", "Genre", "NumTel", "CodeCategorieSocioprofessionnelle", "CodeSecteurActivite")
MAX_LIGNES_COMPLEMENT = 50_000


def _maintenant() -> datetime:
    return datetime.now(timezone.utc)


def _masquer_tel(v: str | None) -> str | None:
    return None if not v else ("•" * max(0, len(v) - 2)) + v[-2:]


def _version_dict(v: ClienteleBcmRefVersion, utilisee: int = 0) -> dict:
    return {
        "id": str(v.id), "libelle": v.libelle, "fichier_nom": v.fichier_nom, "sha256": v.sha256,
        "feuille": v.feuille, "statut": v.statut, "nb_codes": v.nb_codes, "nb_pp": v.nb_pp, "nb_pm": v.nb_pm,
        "anomalies": v.anomalies or [], "observations": v.observations, "importe_par": v.importe_par_nom,
        "importe_le": v.created_at.isoformat() if v.created_at else None,
        "activee_le": v.activee_le.isoformat() if v.activee_le else None,
        "archivee_le": v.archivee_le.isoformat() if v.archivee_le else None,
        "nb_declarations": utilisee,
        "doublons": sum(1 for a in v.anomalies or [] if "dupliqué" in a.get("motif", "")),
    }


class ClienteleBcmReferentielService:
    def __init__(self, db: AsyncSession, ctx: Ctx):
        self.db = db
        self.ctx = ctx
        self.svc = ClienteleService(db, ctx)

    def _nom(self) -> str:
        return self.ctx.user.full_name or self.ctx.user.email

    def _lire(self) -> None:
        self.ctx.exiger("clientele.bcm.view", "clientele.bcm.referentiel")

    def _gerer(self) -> None:
        self.ctx.exiger("clientele.bcm.referentiel")

    # ------------------------------------------------------------------ versions
    async def version_active(self) -> ClienteleBcmRefVersion | None:
        V = ClienteleBcmRefVersion
        return await self.db.scalar(select(V).where(V.statut == "ACTIVE"))

    async def codes_version(self, vid: uuid.UUID) -> dict[str, str]:
        C = ClienteleBcmRefCode
        rows = await self.db.execute(select(C.code, C.usage).where(C.version_id == vid, C.actif.is_(True)))
        return {code: usage for code, usage in rows.all()}

    async def versions(self) -> list[dict]:
        self._lire()
        V, D = ClienteleBcmRefVersion, ClienteleDeclarationSic
        usages = dict((await self.db.execute(
            select(D.ref_version_id, func.count()).where(D.ref_version_id.is_not(None)).group_by(D.ref_version_id)
        )).all())
        rows = (await self.db.scalars(select(V).order_by(V.created_at.desc()))).all()
        return [_version_dict(v, usages.get(v.id, 0)) for v in rows]

    async def _version(self, vid: uuid.UUID) -> ClienteleBcmRefVersion:
        v = await self.db.get(ClienteleBcmRefVersion, vid)
        if not v:
            raise AppError("Version du référentiel introuvable", 404, code="NOT_FOUND")
        return v

    async def importer(self, contenu: bytes, fichier_nom: str, libelle: str | None, observations: str | None) -> dict:
        self._gerer()
        if not contenu:
            raise AppError("Fichier vide", 422, code="FICHIER_VIDE")
        sha = hashlib.sha256(contenu).hexdigest()
        V = ClienteleBcmRefVersion
        deja = await self.db.scalar(select(V).where(V.sha256 == sha))
        if deja:
            raise AppError(f"Fichier déjà importé (version « {deja.libelle} », {deja.statut})", 409,
                           code="DEJA_IMPORTE")
        try:
            wb = load_workbook(io.BytesIO(contenu), read_only=True, data_only=True)
        except Exception as e:
            raise AppError("Fichier Excel illisible", 422, code="FICHIER_ILLISIBLE") from e
        feuille = next((n for n in wb.sheetnames if norm(n) == "REF"), None)
        if not feuille:
            raise AppError("Feuille « ref » absente du fichier", 422, code="FEUILLE_ABSENTE")
        lignes = [tuple(r) for r in wb[feuille].iter_rows(values_only=True)]
        wb.close()
        codes, anomalies = lire_referentiel(lignes)
        if not codes:
            premiers = " ; ".join(f"ligne {a['ligne']} : {a['motif']}" for a in anomalies[:5])
            raise AppError(f"Aucun code exploitable dans la feuille ref ({premiers})", 422, code="REFERENTIEL_VIDE")
        v = V(libelle=(libelle or "").strip()[:120] or f"Référentiel BCM du {_maintenant():%d/%m/%Y}",
              fichier_nom=fichier_nom[:255], sha256=sha, feuille=feuille, statut="BROUILLON",
              nb_codes=len(codes), nb_pp=sum(c["usage"] == "PP" for c in codes),
              nb_pm=sum(c["usage"] == "PM" for c in codes), anomalies=anomalies,
              observations=(observations or "").strip() or None,
              importe_par_id=self.ctx.user.id, importe_par_nom=self._nom())
        self.db.add(v)
        await self.db.flush()
        self.db.add_all([ClienteleBcmRefCode(version_id=v.id, ligne=c["ligne"], code=c["code"],
                                             libelle=c["libelle"], categorie=c["categorie"] or None,
                                             usage=c["usage"]) for c in codes])
        await self.svc.audit("bcm.referentiel.importer", "clientele_bcm_ref_version", v.id,
                             after={"libelle": v.libelle, "fichier": v.fichier_nom, "sha256": sha,
                                    "nb_codes": v.nb_codes, "nb_anomalies": len(anomalies)})
        await self.db.commit()
        return _version_dict(v)

    async def activer(self, vid: uuid.UUID, motif: str) -> dict:
        self._gerer()
        motif = (motif or "").strip()
        if not motif:
            raise AppError("Motif d'activation obligatoire", 422, code="MOTIF_OBLIGATOIRE")
        v = await self._version(vid)
        if v.statut == "ACTIVE":
            raise AppError("Version déjà active", 409, code="DEJA_ACTIVE")
        if v.statut != "BROUILLON":
            raise AppError("Une version archivée ne se réactive pas : réimporter le fichier si nécessaire", 409,
                           code="VERSION_ARCHIVEE")
        doublons = [a for a in v.anomalies or [] if "dupliqué" in a.get("motif", "")]
        if doublons:
            codes = ", ".join(str(a["code"]) for a in doublons[:10])
            raise AppError(f"Activation impossible : {len(doublons)} code(s) dupliqué(s) dans le fichier ({codes})",
                           409, code="DOUBLONS")
        ancienne = await self.version_active()
        maintenant = _maintenant()
        if ancienne:
            ancienne.statut = "ARCHIVEE"
            ancienne.archivee_le = maintenant
            await self.db.flush()
        v.statut = "ACTIVE"
        v.activee_le = maintenant
        v.activee_par_id = self.ctx.user.id
        await self.db.flush()
        await self.svc.audit("bcm.referentiel.activer", "clientele_bcm_ref_version", v.id,
                             before={"active": str(ancienne.id) if ancienne else None},
                             after={"active": str(v.id), "libelle": v.libelle, "motif": motif})
        await self.db.commit()
        return _version_dict(v)

    async def supprimer(self, vid: uuid.UUID) -> None:
        self._gerer()
        v = await self._version(vid)
        if v.statut != "BROUILLON":
            raise AppError("Seule une version brouillon (jamais activée) peut être supprimée", 409,
                           code="VERSION_UTILISEE")
        D = ClienteleDeclarationSic
        if await self.db.scalar(select(func.count()).select_from(D).where(D.ref_version_id == vid)):
            raise AppError("Version utilisée par une déclaration SIC", 409, code="VERSION_UTILISEE")
        await self.svc.audit("bcm.referentiel.supprimer", "clientele_bcm_ref_version", v.id,
                             before={"libelle": v.libelle, "sha256": v.sha256})
        await self.db.delete(v)
        await self.db.commit()

    async def codes(self, vid: uuid.UUID | None, q: str | None, usage: str | None, page: int, taille: int) -> dict:
        self._lire()
        v = await self._version(vid) if vid else await self.version_active()
        if not v:
            return {"version": None, "total": 0, "items": []}
        C = ClienteleBcmRefCode
        stmt = select(C).where(C.version_id == v.id)
        if usage in ("PP", "PM", "AUTRE"):
            stmt = stmt.where(C.usage == usage)
        if q and q.strip():
            motif = f"%{q.strip()}%"
            stmt = stmt.where(C.code.ilike(motif) | C.libelle.ilike(motif) | C.categorie.ilike(motif))
        total = await self.db.scalar(select(func.count()).select_from(stmt.subquery()))
        taille = max(1, min(taille, 500))
        rows = (await self.db.scalars(stmt.order_by(C.ligne).offset((max(page, 1) - 1) * taille).limit(taille))).all()
        return {"version": _version_dict(v), "total": total or 0,
                "items": [{"ligne": c.ligne, "code": c.code, "libelle": c.libelle, "categorie": c.categorie,
                           "usage": c.usage} for c in rows]}

    async def comparer(self, a: uuid.UUID, b: uuid.UUID) -> dict:
        self._lire()
        va, vb = await self._version(a), await self._version(b)
        C = ClienteleBcmRefCode

        async def charger(vid: uuid.UUID) -> dict[str, tuple]:
            rows = await self.db.execute(select(C.code, C.libelle, C.categorie, C.usage).where(C.version_id == vid))
            return {r[0]: (r[1], r[2], r[3]) for r in rows.all()}

        ca, cb = await charger(va.id), await charger(vb.id)
        champs = ("libelle", "categorie", "usage")
        ajoutes = [{"code": k, **dict(zip(champs, cb[k]))} for k in sorted(cb.keys() - ca.keys(), key=_cle)]
        retires = [{"code": k, **dict(zip(champs, ca[k]))} for k in sorted(ca.keys() - cb.keys(), key=_cle)]
        modifies = []
        for k in sorted(ca.keys() & cb.keys(), key=_cle):
            diff = {ch: {"avant": x, "apres": y} for ch, x, y in zip(champs, ca[k], cb[k]) if x != y}
            if diff:
                modifies.append({"code": k, "ecarts": diff})
        return {"a": _version_dict(va), "b": _version_dict(vb), "ajoutes": ajoutes, "retires": retires,
                "modifies": modifies}

    # ------------------------------------------------------------------ agences → MoughataaAgence
    async def agences(self) -> list[dict]:
        self._lire()
        rows = (await self.db.execute(text("""
            SELECT a.id, a.code, a.libelle, a.ville, m.valeur, m.observation, m.maj_par_nom, m.maj_le,
                   (SELECT count(*) FROM clientele_situation s WHERE s.code_agence = a.code) AS nb_clients
            FROM agences a LEFT JOIN clientele_sic_agences m ON m.agence_id = a.id
            WHERE a.deleted_at IS NULL
            ORDER BY a.code
        """))).all()
        return [{"agence_id": str(r[0]), "code": r[1], "libelle": r[2], "ville": r[3], "valeur": r[4],
                 "observation": r[5], "maj_par": r[6], "maj_le": r[7].isoformat() if r[7] else None,
                 "nb_clients": r[8]} for r in rows]

    async def definir_agence(self, agence_id: uuid.UUID, valeur: str, observation: str | None) -> list[dict]:
        self._gerer()
        valeur = " ".join((valeur or "").split())
        if not valeur:
            raise AppError("Valeur MoughataaAgence obligatoire", 422, code="VALEUR_OBLIGATOIRE")
        if not await self.db.get(Agence, agence_id):
            raise AppError("Agence introuvable", 404, code="NOT_FOUND")
        M = ClienteleSicAgence
        m = await self.db.scalar(select(M).where(M.agence_id == agence_id))
        avant = {"valeur": m.valeur} if m else None
        if not m:
            m = M(agence_id=agence_id, valeur=valeur)
            self.db.add(m)
        m.valeur = valeur[:120]
        m.observation = (observation or "").strip() or None
        m.maj_par_id, m.maj_par_nom, m.maj_le = self.ctx.user.id, self._nom(), _maintenant()
        await self.db.flush()
        await self.svc.audit("sic.agence.definir", "clientele_sic_agence", agence_id, before=avant,
                             after={"valeur": m.valeur, "observation": m.observation})
        await self.db.commit()
        return await self.agences()

    async def retirer_agence(self, agence_id: uuid.UUID) -> list[dict]:
        self._gerer()
        M = ClienteleSicAgence
        m = await self.db.scalar(select(M).where(M.agence_id == agence_id))
        if m:
            await self.svc.audit("sic.agence.retirer", "clientele_sic_agence", agence_id, before={"valeur": m.valeur})
            await self.db.delete(m)
            await self.db.commit()
        return await self.agences()

    async def table_agences(self) -> dict[str, str]:
        M = ClienteleSicAgence
        return {str(a): v for a, v in (await self.db.execute(select(M.agence_id, M.valeur))).all()}

    # ------------------------------------------------------------------ secteurs ORION → code BCM
    async def secteurs(self) -> dict:
        self._lire()
        observes = (await self.db.execute(text("""
            SELECT coalesce(nullif(trim(secteur_activite), ''), '(vide)') AS lib, count(*)
            FROM clientele_situation WHERE profil_derive = 'PM' GROUP BY 1 ORDER BY 2 DESC
        """))).all()
        S = ClienteleSicSecteur
        table = {s.libelle_orion: s for s in (await self.db.scalars(select(S))).all()}
        active = await self.version_active()
        codes = await self.codes_version(active.id) if active else {}
        items = []
        vus = set()
        for lib, n in observes:
            cle = norm(lib)
            s = table.get(cle)
            vus.add(cle)
            items.append(self._secteur_dict(lib, cle, n, s, codes))
        for cle, s in table.items():
            if cle not in vus:
                items.append(self._secteur_dict(cle, cle, 0, s, codes))
        return {"items": items, "version_active": _version_dict(active) if active else None}

    @staticmethod
    def _secteur_dict(lib: str, cle: str, n: int, s: ClienteleSicSecteur | None, codes: dict[str, str]) -> dict:
        generique = cle in SECTEURS_GENERIQUES or lib == "(vide)"
        return {"libelle_orion": lib, "cle": cle, "nb_clients_pm": n, "generique": generique,
                "code_bcm": s.code_bcm if s else None, "statut": s.statut if s else None,
                "observation": s.observation if s else None, "maj_par": s.maj_par_nom if s else None,
                "maj_le": s.maj_le.isoformat() if s and s.maj_le else None,
                "code_valide": (codes.get(s.code_bcm) == "PM") if s else None}

    async def definir_secteur(self, libelle: str, code_bcm: str, statut: str, observation: str | None) -> dict:
        self._gerer()
        cle = norm(libelle)
        if not cle or cle == "(VIDE)":
            raise AppError("Libellé ORION obligatoire", 422, code="LIBELLE_OBLIGATOIRE")
        if cle in SECTEURS_GENERIQUES:
            raise AppError(f"« {cle} » est un secteur générique : jamais rattaché automatiquement. Renseigner le "
                           "code BCM client par client dans les compléments SIC.", 422, code="SECTEUR_GENERIQUE")
        statut = (statut or "A_VALIDER").upper()
        if statut not in ("CONFIRMEE", "A_VALIDER"):
            raise AppError("Statut attendu : CONFIRMEE ou A_VALIDER", 422, code="STATUT_INVALIDE")
        code = (code_bcm or "").strip()
        active = await self.version_active()
        if not active:
            raise AppError("Aucune version active du référentiel BCM : importer et activer le référentiel", 409,
                           code="REFERENTIEL_ABSENT")
        if (await self.codes_version(active.id)).get(code) != "PM":
            raise AppError(f"Code {code or '—'} absent du référentiel actif ou réservé aux PP", 422,
                           code="CODE_INVALIDE")
        S = ClienteleSicSecteur
        s = await self.db.scalar(select(S).where(S.libelle_orion == cle))
        avant = {"code_bcm": s.code_bcm, "statut": s.statut} if s else None
        if not s:
            s = S(libelle_orion=cle, code_bcm=code)
            self.db.add(s)
        s.code_bcm, s.statut = code, statut
        s.observation = (observation or "").strip() or None
        s.maj_par_id, s.maj_par_nom, s.maj_le = self.ctx.user.id, self._nom(), _maintenant()
        await self.db.flush()
        await self.svc.audit("sic.secteur.definir", "clientele_sic_secteur", None, before=avant,
                             after={"libelle_orion": cle, "code_bcm": code, "statut": statut})
        await self.db.commit()
        return await self.secteurs()

    async def retirer_secteur(self, libelle: str) -> dict:
        self._gerer()
        S = ClienteleSicSecteur
        s = await self.db.scalar(select(S).where(S.libelle_orion == norm(libelle)))
        if s:
            await self.svc.audit("sic.secteur.retirer", "clientele_sic_secteur", None,
                                 before={"libelle_orion": s.libelle_orion, "code_bcm": s.code_bcm})
            await self.db.delete(s)
            await self.db.commit()
        return await self.secteurs()

    async def table_secteurs(self) -> dict[str, tuple[str, str]]:
        S = ClienteleSicSecteur
        return {lib: (code, st) for lib, code, st in
                (await self.db.execute(select(S.libelle_orion, S.code_bcm, S.statut))).all()}

    # ------------------------------------------------------------------ compléments KYC SIC
    async def complements(self, racine: str | None, page: int, taille: int) -> dict:
        self._lire()
        K = ClienteleKyc
        stmt = select(K).where((K.genre.is_not(None)) | (K.telephone.is_not(None)) | (K.code_csp_bcm.is_not(None))
                               | (K.code_secteur_bcm.is_not(None)))
        if racine and racine.strip():
            stmt = stmt.where(K.racine_client.like(f"{racine.strip()}%"))
        total = await self.db.scalar(select(func.count()).select_from(stmt.subquery()))
        taille = max(1, min(taille, 200))
        rows = (await self.db.scalars(stmt.order_by(K.racine_client).offset((max(page, 1) - 1) * taille)
                                      .limit(taille))).all()
        return {"total": total or 0, "items": [self._complement_dict(k) for k in rows]}

    @staticmethod
    def _complement_dict(k: ClienteleKyc) -> dict:
        return {"racine": k.racine_client, "genre": k.genre, "telephone": k.telephone, "code_csp_bcm": k.code_csp_bcm,
                "code_secteur_bcm": k.code_secteur_bcm, "sic_source": k.sic_source,
                "sic_maj_le": k.sic_maj_le.isoformat() if k.sic_maj_le else None}

    async def _controler_complement(self, valeurs: dict, codes: dict[str, str]) -> tuple[dict, list[str]]:
        erreurs: list[str] = []
        propre: dict = {}
        genre = norm(valeurs.get("genre"))
        if genre:
            genre = {"M": "H", "MASCULIN": "H", "HOMME": "H", "FEMININ": "F", "FEMME": "F"}.get(genre, genre)
            if genre not in ("H", "F"):
                erreurs.append(f"Genre « {valeurs.get('genre')} » invalide (H / F)")
            else:
                propre["genre"] = genre
        tel = normaliser_telephone(valeurs.get("telephone"))
        if tel:
            if not re.fullmatch(r"\+?\d{6,20}", tel):
                erreurs.append(f"Téléphone « {valeurs.get('telephone')} » invalide")
            else:
                propre["telephone"] = tel
        for champ, usage, libelle in (("code_csp_bcm", "PP", "Code CSP"), ("code_secteur_bcm", "PM", "Code secteur")):
            brut = valeurs.get(champ)
            code = str(int(brut)) if isinstance(brut, float) and brut.is_integer() else str(brut or "").strip()
            if code:
                if codes.get(code) != usage:
                    erreurs.append(f"{libelle} {code} absent du référentiel actif ou hors catégorie {usage}")
                else:
                    propre[champ] = code
        return propre, erreurs

    async def _appliquer_complement(self, racine: str, propre: dict, source: str, motif: str | None,
                                    ligne: int | None = None) -> None:
        K = ClienteleKyc
        k = await self.db.get(K, racine)
        avant = None
        if not k:
            k = K(racine_client=racine, source=SOURCE_COMPLEMENT, ligne_source=ligne, maj_par_id=self.ctx.user.id)
            self.db.add(k)
        else:
            avant = {"genre": k.genre, "telephone": _masquer_tel(k.telephone), "code_csp_bcm": k.code_csp_bcm,
                     "code_secteur_bcm": k.code_secteur_bcm}
        for champ, v in propre.items():
            setattr(k, champ, v)
        k.sic_source, k.sic_maj_par_id, k.sic_maj_le = source, self.ctx.user.id, _maintenant()
        await self.db.flush()
        apres = {ch: (_masquer_tel(v) if ch == "telephone" else v) for ch, v in propre.items()}
        await self.svc.audit("sic.complement.definir", "clientele_kyc", racine, before=avant,
                             after={**apres, "source": source, "motif": motif})

    async def definir_complement(self, racine: str, valeurs: dict, motif: str) -> dict:
        self._gerer()
        motif = (motif or "").strip()
        if not motif:
            raise AppError("Motif obligatoire", 422, code="MOTIF_OBLIGATOIRE")
        racine = (racine or "").strip()
        if not await self.db.scalar(text("SELECT 1 FROM clientele_clients WHERE racine_client = :r"), {"r": racine}):
            raise AppError("Racine client inconnue", 404, code="NOT_FOUND")
        active = await self.version_active()
        propre, erreurs = await self._controler_complement(valeurs, await self.codes_version(active.id)
                                                           if active else {})
        if erreurs:
            raise AppError(" ; ".join(erreurs), 422, code="VALEUR_INVALIDE")
        if not propre:
            raise AppError("Aucune valeur à enregistrer", 422, code="VALEUR_OBLIGATOIRE")
        await self._appliquer_complement(racine, propre, "SAISIE", motif)
        await self.db.commit()
        return self._complement_dict(await self.db.get(ClienteleKyc, racine))

    async def importer_complements(self, contenu: bytes, fichier_nom: str, appliquer: bool) -> dict:
        """Analyse (``appliquer=False``) puis application. Les cellules vides ne remplacent rien."""
        self._gerer()
        try:
            wb = load_workbook(io.BytesIO(contenu), read_only=True, data_only=True)
        except Exception as e:
            raise AppError("Fichier Excel illisible", 422, code="FICHIER_ILLISIBLE") from e
        ws = wb[wb.sheetnames[0]]
        rows = list(ws.iter_rows(values_only=True))
        wb.close()
        if not rows:
            raise AppError("Fichier vide", 422, code="FICHIER_VIDE")
        entete = [norm(x) for x in rows[0]]
        idx = {}
        for col in COLONNES_COMPLEMENT:
            if norm(col) in entete:
                idx[col] = entete.index(norm(col))
        if "CodeClient" not in idx:
            raise AppError(f"Colonnes attendues : {', '.join(COLONNES_COMPLEMENT)}", 422, code="ENTETES")
        if len(rows) - 1 > MAX_LIGNES_COMPLEMENT:
            raise AppError(f"Fichier limité à {MAX_LIGNES_COMPLEMENT} lignes", 422, code="FICHIER_TROP_GROS")
        active = await self.version_active()
        codes = await self.codes_version(active.id) if active else {}
        existants = set((await self.db.scalars(text("SELECT racine_client FROM clientele_clients"))).all())
        champs = {"Genre": "genre", "NumTel": "telephone", "CodeCategorieSocioprofessionnelle": "code_csp_bcm",
                  "CodeSecteurActivite": "code_secteur_bcm"}
        retenues, rejets, vus = [], [], {}
        for n, row in enumerate(rows[1:], start=2):
            brut = row[idx["CodeClient"]] if idx["CodeClient"] < len(row) else None
            if brut is None and not any(row):
                continue
            racine = _racine(brut)
            if racine not in existants:
                rejets.append({"ligne": n, "racine": racine, "motif": "Racine client inconnue"})
                continue
            if racine in vus:
                rejets.append({"ligne": n, "racine": racine, "motif": f"Racine en double (ligne {vus[racine]})"})
                continue
            vus[racine] = n
            valeurs = {champ: row[idx[col]] for col, champ in champs.items() if col in idx and idx[col] < len(row)}
            propre, erreurs = await self._controler_complement(valeurs, codes)
            if erreurs:
                rejets.append({"ligne": n, "racine": racine, "motif": " ; ".join(erreurs)})
            elif propre:
                retenues.append((n, racine, propre))
        resultat = {"fichier": fichier_nom, "lignes": len(rows) - 1, "retenues": len(retenues),
                    "rejets": rejets[:500], "nb_rejets": len(rejets), "applique": False}
        if appliquer and retenues:
            for n, racine, propre in retenues:
                await self._appliquer_complement(racine, propre, "IMPORT", f"Import {fichier_nom}", ligne=n)
            await self.svc.audit("sic.complement.importer", "clientele_kyc", None,
                                 after={"fichier": fichier_nom, "retenues": len(retenues), "rejets": len(rejets)})
            await self.db.commit()
            resultat["applique"] = True
        return resultat

    @staticmethod
    def modele_complements() -> bytes:
        wb = Workbook()
        ws = wb.active
        ws.title = "complements"
        for j, h in enumerate(COLONNES_COMPLEMENT, start=1):
            c = ws.cell(row=1, column=j, value=h)
            c.font = Font(bold=True)
            c.number_format = "@"
            ws.column_dimensions[c.column_letter].width = len(h) + 6
            for i in range(2, 502):
                ws.cell(row=i, column=j).number_format = "@"
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()


def _racine(brut) -> str:
    """Racine lue dans Excel : une cellule numérique a perdu ses zéros initiaux, on les restitue."""
    if isinstance(brut, (int, float)) and float(brut).is_integer():
        return str(int(brut)).zfill(6)
    s = str(brut or "").strip()
    return s.zfill(6) if s.isdigit() and len(s) < 6 else s


def _cle(code: str) -> tuple[int, str]:
    return (int(code) if code.isdigit() else 10**9, code)

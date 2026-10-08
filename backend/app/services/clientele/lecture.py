"""Lecture d'un classeur État Compte ORION (openpyxl, read_only)."""

from __future__ import annotations

from dataclasses import dataclass, field
from io import BytesIO

from openpyxl import load_workbook

from app.services.clientele.consolidation import (
    COLONNES_OBLIGATOIRES,
    Anomalie,
    Consolidation,
    LigneInvalide,
    LigneOrion,
    cartographier,
    consolider,
    lire_ligne,
)

LIGNES_ENTETE_MAX = 15
APERCU_MAX = 20


@dataclass
class AnalyseClasseur:
    feuille: str
    ligne_entete: int
    mapping: dict[str, str]
    colonnes_ignorees: list[str]
    colonnes_manquantes: list[str]
    consolidation: Consolidation
    lignes: list[LigneOrion] = field(default_factory=list)
    rejets: list[Anomalie] = field(default_factory=list)
    nb_vides: int = 0
    nb_lues: int = 0
    apercu: list[dict] = field(default_factory=list)


def _vide(row: tuple) -> bool:
    return all(v is None or (isinstance(v, str) and not str(v).strip()) for v in row)


def _score(mapping: dict[str, str]) -> int:
    return sum(3 if c in COLONNES_OBLIGATOIRES else 1 for c in mapping)


def _detecter(wb) -> tuple[str, int, dict[str, str], list[str]]:
    meilleure: tuple | None = None
    for ws in wb.worksheets:
        for i, row in enumerate(ws.iter_rows(values_only=True, max_row=LIGNES_ENTETE_MAX), 1):
            entetes = ["" if v is None else str(v).strip() for v in row]
            if not any(entetes):
                continue
            mapping, ignores = cartographier(entetes)
            score = _score(mapping)
            if meilleure is None or score > meilleure[0]:
                meilleure = (score, ws.title, i, mapping, ignores)
    if not meilleure or meilleure[0] < 6:
        raise ValueError("COLONNES_NON_RECONNUES")
    return meilleure[1], meilleure[2], meilleure[3], meilleure[4]


def detecter_feuille(contenu: bytes) -> tuple[str, int, dict[str, str], list[str]]:
    wb = load_workbook(BytesIO(contenu), read_only=True, data_only=True)
    try:
        return _detecter(wb)
    finally:
        wb.close()


def _indices(entetes: list[str], mapping: dict[str, str]) -> dict[str, int]:
    """Champ canonique → index de colonne (dernière occurrence d'un en-tête dupliqué)."""
    position = {e: i for i, e in enumerate(entetes)}
    return {champ: position[entete] for champ, entete in mapping.items() if entete in position}


def lire_classeur(contenu: bytes, mapping: dict[str, str] | None = None) -> AnalyseClasseur:
    """Lit le classeur, cartographie les colonnes, consolide par racine client."""
    wb = load_workbook(BytesIO(contenu), read_only=True, data_only=True)
    try:
        feuille, ligne_entete, auto, ignores_auto = _detecter(wb)
        mapping = dict(mapping or auto)
        manquantes = [c for c in COLONNES_OBLIGATOIRES if c not in mapping]

        ws = next(s for s in wb.worksheets if s.title == feuille)
        entetes: list[str] = []
        indices: dict[str, int] = {}
        lignes_orion: list[LigneOrion] = []
        rejets: list[Anomalie] = []
        apercu: list[dict] = []
        vides = lues = 0
        for numero, row in enumerate(ws.iter_rows(values_only=True), 1):
            if numero < ligne_entete:
                continue
            if numero == ligne_entete:
                entetes = ["" if v is None else str(v).strip() for v in row]
                indices = _indices(entetes, mapping)
                continue
            if _vide(row):
                vides += 1
                continue
            lues += 1
            if manquantes:
                continue
            n = len(row)
            canon = {champ: (row[i] if i < n else None) for champ, i in indices.items()}
            for champ in mapping:
                canon.setdefault(champ, None)
            try:
                ligne = lire_ligne(canon, numero)
            except LigneInvalide as exc:
                rejets.append(Anomalie(numero, None, "LIGNE_INVALIDE", " ; ".join(exc.motifs), True))
                continue
            lignes_orion.append(ligne)
            if len(apercu) < APERCU_MAX:
                apercu.append({
                    "ligne": numero,
                    "racine_client": ligne.client.racine_client,
                    "raison_sociale": ligne.client.raison_sociale,
                    "compte": ligne.compte.compte,
                    "rib": ligne.compte.rib,
                    "etat_compte": ligne.compte.etat_compte,
                    "code_agence": ligne.compte.code_agence,
                    "devise": ligne.compte.devise,
                })
        ignores = [e for e in entetes if e and e not in mapping.values()]
    finally:
        wb.close()

    if manquantes:
        consolidation = Consolidation()
        consolidation.nb_lignes = lues
    else:
        consolidation = consolider(lignes_orion)
    return AnalyseClasseur(
        feuille=feuille, ligne_entete=ligne_entete, mapping=mapping,
        colonnes_ignorees=ignores if entetes else ignores_auto,
        colonnes_manquantes=manquantes, consolidation=consolidation, lignes=lignes_orion,
        rejets=rejets, nb_vides=vides, nb_lues=lues, apercu=apercu,
    )

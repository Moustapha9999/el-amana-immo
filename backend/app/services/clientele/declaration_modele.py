"""Export Excel de la Déclaration mensuelle BCM au format du modèle officiel BEA (août 2026).

Feuille « Déclaration mensuelle » : mêmes positions (A1:E44), intitulés et largeurs que le modèle.
Les valeurs viennent exclusivement des cellules calculées de la déclaration ; une cellule
``A_CONFIGURER`` est écrite « À CONFIGURER », jamais un nombre. Feuille « Contrôle » : traçabilité.
"""

from __future__ import annotations

from datetime import datetime
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

A_CONFIGURER = "À CONFIGURER"
NON_CALCULABLE = "Non calculable"
TITRE = "TABLEAU DE DECLARATION MENSUELLE BEA à LA BCM"
FEUILLE = "Déclaration mensuelle"
LARGEURS = {"A": 162.14, "B": 44.29, "C": 24.86, "D": 26.57, "E": 7.86}

COLS_JURIDIQUE = ("Personnes physiques", "Personnes morales", "Constructions juridiques", "TOTAL")
COLS_RISQUE = ("Clients à risque élevé", "Clients à risque moyen", "Clients à faible risque", "TOTAL")
COL_NOMBRE = "Le nombre à la fin du mois de déclaration"

# (tableau, ligne de titre, ligne de sous-titre ou None, ligne d'en-tête, en-têtes, colonnes moteur,
#  lignes : (code ligne moteur, intitulé du modèle)). Intitulés repris mot pour mot du modèle.
BLOCS: tuple[dict[str, Any], ...] = (
    {
        "tableau": "T1A", "titre": (2, "Tableau: Mise à jour des données clients"),
        "sous_titre": (3, "Selon la forme juridique"), "entete": 4,
        "colonnes": COLS_JURIDIQUE, "codes": ("pp", "pm", "cj", "total"),
        "lignes": (
            ("stock_m1", "Le nombre à la fin du mois précédent le mois de déclaration (nbre de client jusqu'au "
                         "mois prècedent le mois de déclaration actif+inactif)"),
            ("cible", "Nombre cible de clients à mettre à jour leurs données pendant le mois de référence(les "
                      "clients actifs jusqu'au mois de declaration )"),
            ("maj_mois", "Nombre de clients dont les données ont été mises à jour au cours du mois considéré (les "
                         "EER du mois de declaration )"),
            ("maj_5ans", "Nombre de clients dont les données ont été mises à jour au cours des cinq dernières "
                         "années (tous les clients actifs a partir de ( 5 ans avant l'annee courante )"),
        ),
    },
    {
        "tableau": "T1B", "titre": None, "sous_titre": (10, "Selon le degré de risque"), "entete": 11,
        "colonnes": COLS_RISQUE, "codes": ("eleve", "moyen", "faible", "total"),
        "lignes": (
            ("stock_m1", "Le nombre à la fin du mois précédent le mois de déclaration (nbre de client jusqu'au "
                         "mois prècedent le mois de déclaration actif+inactif, par degree de risque )"),
            ("cible", "Nombre cible de clients à mettre à jour leurs données pendant le mois de référence(les "
                      "clients actifs jusqu'au mois de declaration , par degree de risque )"),
            ("maj_mois", "Nombre de clients dont les données ont été mises à jour au cours du mois considéré (les "
                         "EER du mois de declaration, par degree de risque  )"),
            ("maj_5ans", "Nombre de clients dont les données ont été mises à jour au cours des cinq dernières "
                         "années (tous les clients actifs a partir de ( 5 ans avant l'annee courante, par degree "
                         "de risque )"),
        ),
    },
    {
        "tableau": "T2", "titre": (19, "Tableau: Opérations inhabituelles"), "sous_titre": None, "entete": 20,
        "colonnes": (COL_NOMBRE,), "codes": ("valeur",),
        "lignes": (
            ("extraites", "Nombre d’opérations inhabituelles extraites du système de l’institution assujettie "
                          "(Alertes declaree du debut d'anne jusqu'au mois de declaration)"),
            ("analysees", "Nombre d’opérations analysées à partir de ces Opérations (toutes les operations en "
                          "cours d'analyse , cloturer, et declarer )"),
            ("suivi", "Nombre d’opérations pour lesquelles l’examen a donné lieu à un suivi continu (tous les "
                      "operations en cours d'analyse, et en attente de reponse)"),
            ("enregistrees", "Nombre d’opérations enregistrées à partir  de ces processus"),
            ("umef_mois", "Nombre de déclaration de soupçon envoyées à l’UMEF sur la base de l’examen de ces "
                          "opérations (nombre de declaration effectuer durant le mois de declaration )"),
        ),
    },
    {
        "tableau": "T3", "titre": (29, "Tableau: Opérations  suspectes"), "sous_titre": None, "entete": 30,
        "colonnes": (COL_NOMBRE,), "codes": ("valeur",),
        "lignes": (
            ("recues", "Nombre de Transactions suspectes reçues par le Département de la conformité de la part "
                       "des employés de l’institution assujettie"),
            ("analysees", "Nombre d’opérations analysées à partir de ces Opérations"),
            ("suivi", "Nombre d’opérations pour lesquelles l’examen a donné lieu à un suivi continu"),
            ("enregistrees", "Nombre d’opérations enregistrées à partir  de ces processus"),
            ("umef", "Nombre de déclaration de soupçon envoyées à l’UMEF sur la base de l’examen de ces "
                     "opérations"),
        ),
    },
    {
        "tableau": "T4",
        "titre": (39, "Tableau: Classification des clients selon les risques de blanchiment d’argent et de "
                      "financement du Terrorisme"),
        "sous_titre": None, "entete": 40,
        "colonnes": (COL_NOMBRE, "Ratio %"), "codes": ("valeur", "ratio"),
        "lignes": (
            ("eleve", "Clients à risque élevé (tous les clients actifs / inactifs, jusqu'au mois de declaration "
                      "inclut) avec niveau de risque eleve"),
            ("moyen", "Clients à risque moyen (tous les clients actifs / inactifs, jusqu'au mois de declaration "
                      "inclut) avec niveau de risque moyen"),
            ("faible", "Clients à faible risque (tous les clients actifs / inactifs, jusqu'au mois de declaration "
                       "inclut) avec niveau de risque faible"),
            ("total", "Total clients"),
        ),
    },
)

_FIN = Side(style="thin")
_BORDURE = Border(left=_FIN, right=_FIN, top=_FIN, bottom=_FIN)
_GRAS = Font(name="Calibri", size=11, bold=True)
_NORMAL = Font(name="Calibri", size=11)
_FOND_ENTETE = PatternFill("solid", fgColor="FFCCFFFF")
_FORMAT_NOMBRE = "#,##0"
_FORMAT_RATIO = "0.00"


def _cellules_par_code(cellules: dict | None) -> dict[tuple[str, str], dict[str, dict]]:
    return {(tab["code"], ligne["code"]): ligne.get("cellules") or {}
            for tab in (cellules or {}).get("tableaux") or [] for ligne in tab.get("lignes") or []}


def _valeur(cell: dict | None) -> int | float | str:
    if not cell or cell.get("statut") == "A_CONFIGURER":
        return A_CONFIGURER
    if cell.get("valeur") is None:
        return NON_CALCULABLE
    return cell["valeur"]


def positions() -> dict[str, tuple[str, str, str]]:
    """Référence Excel → (tableau, ligne, colonne moteur) de chaque cellule chiffrée du modèle."""
    out: dict[str, tuple[str, str, str]] = {}
    for bloc in BLOCS:
        for i, (code_ligne, _lib) in enumerate(bloc["lignes"]):
            for j, code_col in enumerate(bloc["codes"]):
                out[f"{get_column_letter(j + 2)}{bloc['entete'] + 1 + i}"] = (bloc["tableau"], code_ligne, code_col)
    return out


def _feuille_declaration(ws, cellules: dict | None) -> list[list]:
    ws.title = FEUILLE
    for col, largeur in LARGEURS.items():
        ws.column_dimensions[col].width = largeur
    ws["A1"] = TITRE
    ws["A1"].font = _NORMAL
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.merge_cells("A1:E1")
    index = _cellules_par_code(cellules)
    a_configurer: list[list] = []
    for bloc in BLOCS:
        for repere in (bloc["titre"], bloc["sous_titre"]):
            if repere:
                ws.cell(row=repere[0], column=1, value=repere[1]).font = _GRAS
        r0 = bloc["entete"]
        for j, libelle in enumerate(("Déclarations", *bloc["colonnes"])):
            c = ws.cell(row=r0, column=j + 1, value=libelle)
            c.font, c.fill, c.border = _GRAS, _FOND_ENTETE, _BORDURE
        for i, (code_ligne, libelle) in enumerate(bloc["lignes"]):
            r = r0 + 1 + i
            c = ws.cell(row=r, column=1, value=libelle)
            c.font, c.border = _NORMAL, _BORDURE
            source = index.get((bloc["tableau"], code_ligne), {})
            for j, code_col in enumerate(bloc["codes"]):
                cell = source.get(code_col)
                v = _valeur(cell)
                c = ws.cell(row=r, column=j + 2, value=v)
                c.font, c.border = _NORMAL, _BORDURE
                if isinstance(v, str):
                    c.alignment = Alignment(horizontal="center")
                    if v == A_CONFIGURER:
                        a_configurer.append([c.coordinate, bloc["tableau"], libelle,
                                             bloc["colonnes"][j], (cell or {}).get("reserve")])
                else:
                    c.number_format = _FORMAT_RATIO if code_col == "ratio" else _FORMAT_NOMBRE
    return a_configurer


def _horodatage(v: Any) -> str | None:
    if isinstance(v, datetime):
        return v.strftime("%d/%m/%Y %H:%M")
    return v


def _feuille_controle(ws, meta: dict, a_configurer: list[list], controles: list[dict]) -> None:
    ws.title = "Contrôle"
    for col, largeur in {"A": 34, "B": 22, "C": 70, "D": 30, "E": 70, "F": 12, "G": 12}.items():
        ws.column_dimensions[col].width = largeur
    r = 1
    ws.cell(row=r, column=1, value="Document de préparation — vérification avant transmission. "
                                   "Aucune transmission automatique à la BCM.").font = _GRAS
    r += 2
    for libelle, valeur in (
        ("Mois de déclaration", meta.get("periode")),
        ("Statut de la déclaration", meta.get("statut")),
        ("Généré le", _horodatage(meta.get("genere_le"))),
        ("Généré par", meta.get("genere_par")),
        ("Calculée le", _horodatage(meta.get("calculee_le"))),
        ("Validée le", _horodatage(meta.get("validee_le"))),
        ("Version moteur", meta.get("moteur_version")),
        ("Version grille", meta.get("grille_version")),
    ):
        ws.cell(row=r, column=1, value=libelle).font = _GRAS
        ws.cell(row=r, column=2, value=valeur)
        r += 1
    r += 1
    ws.cell(row=r, column=1, value=f"Cellules « {A_CONFIGURER} » (définition métier ou source absente)").font = _GRAS
    r += 1
    for j, h in enumerate(("Cellule", "Tableau", "Ligne", "Colonne", "Réserve")):
        c = ws.cell(row=r, column=j + 1, value=h)
        c.font, c.fill, c.border = _GRAS, _FOND_ENTETE, _BORDURE
    for ligne in a_configurer:
        r += 1
        for j, v in enumerate(ligne):
            ws.cell(row=r, column=j + 1, value=v).border = _BORDURE
    r += 2
    ws.cell(row=r, column=1, value="Contrôles de cohérence").font = _GRAS
    r += 1
    for j, h in enumerate(("Code", "État", "Libellé", "Attendu", "Message", "Observé", "Bloquant")):
        c = ws.cell(row=r, column=j + 1, value=h)
        c.font, c.fill, c.border = _GRAS, _FOND_ENTETE, _BORDURE
    for ctrl in controles:
        r += 1
        valeurs = (ctrl.get("code"), ctrl.get("etat") or ctrl.get("niveau"), ctrl.get("libelle"),
                   ctrl.get("attendu"), ctrl.get("message"), ctrl.get("observe"),
                   "oui" if ctrl.get("bloquant") else "non")
        for j, v in enumerate(valeurs):
            ws.cell(row=r, column=j + 1, value=v).border = _BORDURE


def classeur_modele(cellules: dict | None, controles: list[dict] | None, meta: dict) -> bytes:
    wb = Workbook()
    a_configurer = _feuille_declaration(wb.active, cellules)
    _feuille_controle(wb.create_sheet(), meta, a_configurer, controles or [])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()

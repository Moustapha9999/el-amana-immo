"""Vue 360° clientèle — lecture seule (docs/conformite/clientele-vue360.md).

Présente les résultats SCORE tels qu'enregistrés : aucune écriture sur clientele_classifications,
l'historique ou les évaluations, aucun recalcul. Les contrôles de cohérence sont des diagnostics
calculés à part, jamais présentés comme un motif du moteur.

Unités : un client = une racine ORION (clientele_situation, COUNT DISTINCT) ; une alerte = une ligne
clientele_alertes ; un événement = une ligne d'historique. Ces unités ne s'additionnent pas.
Fuseau : Africa/Nouakchott = UTC+0 toute l'année ; les bornes de période sont donc des jours UTC.
"""

from __future__ import annotations

import io
import re
from collections import defaultdict
from dataclasses import dataclass
from datetime import UTC, date, datetime, time, timedelta
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import String, column, func, select, text
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.exceptions import AppError
from app.data.clientele_classif_matrice import SEUILS_CDC, critere_par_code
from app.models import (
    AuditLog,
    ClienteleClassifEvaluation,
    ClienteleClassifEvaluationLigne,
    ClienteleClassifReexamen,
    User,
)
from app.services.clientele.classification import ClienteleClassificationService
from app.services.clientele.filtrage import GROUPES
from app.services.clientele.service import Ctx, ClienteleService, maintenant
from app.services.clientele.situation import ClienteleSituationService

FUSEAU = "Africa/Nouakchott (UTC+0)"
NIVEAUX = ("FAIBLE", "MOYEN", "ELEVE", "INTERDIT")
PROFILS = ("PP", "PM", "NON_IDENTIFIE")
PERIODES = ("jour", "7j", "30j", "trimestre", "annee", "perso")
PERMS_AUDIT = ("clientele.admin", "clientele.classif.admin", "clientele.filtrage.decide")
LIMITE_EXPORT = 50_000


@dataclass(frozen=True)
class Anomalie:
    code: str
    colonne: str
    libelle: str
    definition: str
    famille: str


ANOMALIES = (
    Anomalie("NON_EVALUE", "f_non_evalue", "Sans classification enregistrée",
             "Aucune ligne dans clientele_classifications pour la racine. Jamais compté en Faible.", "PROFILAGE"),
    Anomalie("A_ARBITRER", "f_a_arbitrer", "Évaluation bloquée (conflit de référentiel)",
             "Sans classification et dernière évaluation SCORE au statut A_ARBITRER : "
             "aucun profil exploitable tant que le référentiel n'est pas arbitré.", "PROFILAGE"),
    Anomalie("CUMUL_CRITERES_ELEVES", "f_cumul", "Plusieurs critères à 100 000 points",
             "Au moins deux critères ≥ 100 000 dans les motifs enregistrés. Le niveau reste ÉLEVÉ "
             "(seuil CDC 100 000), mais le motif affiché peut n'en citer qu'un : à vérifier.", "COHERENCE"),
    Anomalie("NIVEAU_SCORE_INCOHERENT", "f_niveau_score", "Niveau différent des seuils CDC",
             "Classification MOTEUR dont le niveau ne correspond pas aux seuils CDC appliqués au score "
             "enregistré (INTERDIT exclu : blocage métier).", "COHERENCE"),
    Anomalie("LISTE_INTERDICTION_COMPTEE", "f_liste_comptee", "Liste d'interdiction ORION comptée",
             "Le score enregistré inclut la Liste d'interdiction ORION à 100 000 alors que le référentiel "
             "courant classe ce critère hors score (SUSPENDU) : résultat antérieur à la suspension.", "COHERENCE"),
    Anomalie("DIVERGENCE_EVALUATION", "f_eval_diverge", "Classe retenue ≠ dernière évaluation",
             "Le niveau retenu diffère du niveau final de la dernière évaluation photographiée "
             "(par exemple décision manuelle ou évaluation postérieure non enregistrée).", "COHERENCE"),
    Anomalie("REGLES_MODIFIEES_SANS_VERSION", "f_regles_homonymes", "Scores différents, même version de règles",
             "Au moins deux résultats MOTEUR historisés avec le même libellé version_regles et des scores "
             "différents : les règles ont changé sans changement de numéro de version.", "COHERENCE"),
    Anomalie("REEXAMEN_OUVERT", "f_reexamen", "Réexamen ouvert",
             "Une source suggère de revoir la décision (clientele_classif_reexamens au statut OUVERT).", "COHERENCE"),
    Anomalie("SANS_IDENTIFIANT", "f_sans_identifiant", "Sans identifiant exploitable",
             "NNI, NIF et identifiant ORION vides.", "DONNEES"),
    Anomalie("SANS_NATIONALITE", "f_sans_nationalite", "Nationalité vide",
             "Champ NATIONALITE vide dans l'extraction ORION.", "DONNEES"),
    Anomalie("PROFIL_NON_IDENTIFIE", "f_profil_inconnu", "Profil PP / PM non identifié",
             "L'heuristique ORION (agent économique puis catégorie juridique) ne permet pas de conclure.", "DONNEES"),
)
_ANOMALIE = {a.code: a for a in ANOMALIES}
_COHERENCE = tuple(a.colonne for a in ANOMALIES if a.famille == "COHERENCE")
_DONNEES = tuple(a.colonne for a in ANOMALIES if a.famille == "DONNEES")

FILTRAGE = {
    "SANS": "al_total = 0",
    "AVEC": "al_total > 0",
    "A_EXAMINER": "al_a_examiner > 0",
    "EN_ANALYSE": "al_en_analyse > 0",
    "CONFIRMEE": "al_confirmees > 0",
    "ECARTEE": "al_ecartees > 0",
    "CLOTUREE": "al_cloturees > 0",
}
PROFILAGE = {
    "EXPLOITABLE": "NOT f_non_evalue",
    "NON_PROFILE": "f_non_evalue",
    "A_ARBITRER": "f_a_arbitrer",
}

_BASE = f"""
WITH der_eval AS (
    SELECT DISTINCT ON (racine_client) racine_client, statut AS eval_statut, niveau_final AS eval_niveau,
           score_total AS eval_score, created_at AS eval_le
    FROM clientele_classif_evaluations
    ORDER BY racine_client, created_at DESC
),
homonymes AS (
    SELECT DISTINCT racine_client FROM clientele_classif_historique
    WHERE source = 'MOTEUR' AND score_calcule IS NOT NULL AND version_regles IS NOT NULL
    GROUP BY racine_client, version_regles HAVING count(DISTINCT score_calcule) > 1
),
alertes AS (
    SELECT racine_client,
           count(*) FILTER (WHERE statut IN ('NOUVELLE', 'A_ANALYSER')) AS al_a_examiner,
           count(*) FILTER (WHERE statut = 'EN_INVESTIGATION') AS al_en_analyse,
           count(*) FILTER (WHERE statut = 'CONFIRMEE') AS al_confirmees,
           count(*) FILTER (WHERE statut IN ('FAUX_POSITIF', 'REJETEE')) AS al_ecartees,
           count(*) FILTER (WHERE statut = 'CLOTUREE') AS al_cloturees,
           count(*) AS al_total
    FROM clientele_alertes GROUP BY racine_client
),
reex AS (
    SELECT racine_client, count(*) AS nb FROM clientele_classif_reexamens
    WHERE statut = 'OUVERT' GROUP BY racine_client
),
crit AS (
    SELECT c.racine_client,
           count(*) FILTER (WHERE m->>'critere' <> 'SCORE' AND (m->>'poids') ~ '^[0-9]+$'
                            AND (m->>'poids')::bigint >= :poids_eleve) AS nb_eleves,
           bool_or(m->>'critere' = 'LISTE_INTERDICTION' AND (m->>'poids') ~ '^[0-9]+$'
                   AND (m->>'poids')::bigint >= :poids_eleve) AS liste_comptee
    FROM clientele_classifications c
    CROSS JOIN LATERAL jsonb_array_elements(
        CASE WHEN jsonb_typeof(c.motifs) = 'array' THEN c.motifs ELSE '[]'::jsonb END) m
    GROUP BY c.racine_client
),
f AS (
    SELECT s.racine_client, s.nom_client, s.prenoms, s.profil_derive, s.code_agence, s.agence,
           s.nb_comptes, s.etat_client, s.date_extraction, s.nni, s.nif,
           c.niveau, c.source, c.mode_decision, c.score_calcule, c.niveau_propose,
           c.motif_risque, c.motif_moteur, c.motif_source, c.motif_decision_manuelle, c.motif_classement,
           c.classifie_le, c.version_regles,
           e.eval_statut, e.eval_niveau, e.eval_score, e.eval_le,
           coalesce(a.al_a_examiner, 0) AS al_a_examiner, coalesce(a.al_en_analyse, 0) AS al_en_analyse,
           coalesce(a.al_confirmees, 0) AS al_confirmees, coalesce(a.al_ecartees, 0) AS al_ecartees,
           coalesce(a.al_cloturees, 0) AS al_cloturees, coalesce(a.al_total, 0) AS al_total,
           coalesce(r.nb, 0) AS reexamens,
           (c.racine_client IS NULL) AS f_non_evalue,
           (c.racine_client IS NULL AND e.eval_statut = 'A_ARBITRER') AS f_a_arbitrer,
           (coalesce(k.nb_eleves, 0) >= 2) AS f_cumul,
           (c.source = 'MOTEUR' AND c.score_calcule IS NOT NULL AND c.niveau <> 'INTERDIT'
            AND c.niveau <> CASE WHEN c.score_calcule >= :seuil_eleve THEN 'ELEVE'
                                 WHEN c.score_calcule >= :seuil_moyen THEN 'MOYEN' ELSE 'FAIBLE' END
           ) IS TRUE AS f_niveau_score,
           (CAST(:liste_hors_score AS boolean) AND coalesce(k.liste_comptee, false)) AS f_liste_comptee,
           (c.niveau IS NOT NULL AND e.eval_niveau IS NOT NULL AND c.niveau <> e.eval_niveau) AS f_eval_diverge,
           (h.racine_client IS NOT NULL) AS f_regles_homonymes,
           (coalesce(r.nb, 0) > 0) AS f_reexamen,
           (coalesce(s.nni, '') = '' AND coalesce(s.nif, '') = ''
            AND coalesce(s.identifiant_orion, '') = '') AS f_sans_identifiant,
           (coalesce(s.nationalite, '') = '') AS f_sans_nationalite,
           (s.profil_derive = 'NON_IDENTIFIE') AS f_profil_inconnu
    FROM clientele_situation s
    LEFT JOIN clientele_classifications c ON c.racine_client = s.racine_client
    LEFT JOIN der_eval e ON e.racine_client = s.racine_client
    LEFT JOIN alertes a ON a.racine_client = s.racine_client
    LEFT JOIN reex r ON r.racine_client = s.racine_client
    LEFT JOIN crit k ON k.racine_client = s.racine_client
    LEFT JOIN homonymes h ON h.racine_client = s.racine_client
),
g AS (
    SELECT f.*,
           ({' OR '.join(_DONNEES)}) AS f_incomplet,
           ({' OR '.join(_COHERENCE)}) AS f_a_verifier
    FROM f
),
pop AS (SELECT * FROM g WHERE {{where}})
"""
# Population sans diagnostics : suffit quand les filtres ne portent que sur le référentiel.
_BASE_LEGERE = """
WITH pop AS (
    SELECT racine_client, nom_client, profil_derive, code_agence, nni, nif
    FROM clientele_situation WHERE {where}
)
"""
_FILTRES_LEGERS = frozenset({"profil", "agence", "q", "periode", "du", "au"})


@dataclass(frozen=True)
class Periode:
    code: str
    debut: date
    fin: date
    prec_debut: date
    prec_fin: date

    @property
    def jours(self) -> int:
        return (self.fin - self.debut).days + 1

    def bornes(self, precedente: bool = False) -> tuple[datetime, datetime]:
        d, f = (self.prec_debut, self.prec_fin) if precedente else (self.debut, self.fin)
        return datetime.combine(d, time.min, tzinfo=UTC), datetime.combine(f + timedelta(days=1), time.min, tzinfo=UTC)

    def dict(self) -> dict:
        return {"code": self.code, "du": self.debut.isoformat(), "au": self.fin.isoformat(), "jours": self.jours,
                "precedente": {"du": self.prec_debut.isoformat(), "au": self.prec_fin.isoformat()}}


def resoudre_periode(code: str | None, du: date | None = None, au: date | None = None,
                     aujourd_hui: date | None = None) -> Periode:
    """Période de référence et période précédente de même durée, immédiatement antérieure."""
    j = aujourd_hui or maintenant().date()
    code = code if code in PERIODES else "30j"
    if code == "perso":
        if not du or not au:
            raise AppError("Période personnalisée : dates de début et de fin obligatoires", 422,
                           code="PERIODE_INCOMPLETE")
        if au < du:
            raise AppError("Période personnalisée : fin antérieure au début", 422, code="PERIODE_INVALIDE")
        if (au - du).days > 3660:
            raise AppError("Période limitée à 10 ans", 422, code="PERIODE_TROP_LONGUE")
        debut, fin = du, au
    elif code == "jour":
        debut = fin = j
    elif code == "7j":
        debut, fin = j - timedelta(days=6), j
    elif code == "trimestre":
        debut, fin = date(j.year, 3 * ((j.month - 1) // 3) + 1, 1), j
    elif code == "annee":
        debut, fin = date(j.year, 1, 1), j
    else:
        debut, fin = j - timedelta(days=29), j
    duree = (fin - debut).days + 1
    return Periode(code, debut, fin, debut - timedelta(days=duree), debut - timedelta(days=1))


def _pas(p: Periode) -> str:
    return "day" if p.jours <= 62 else ("week" if p.jours <= 370 else "month")


def _iso(v: Any) -> str | None:
    return v.isoformat() if isinstance(v, (date, datetime)) else None


_DATES_SYNTHESE = ("dernier_classement", "derniere_evaluation", "stock_au")


def _cumuler(groupes: list[dict]) -> defaultdict:
    """Somme des compteurs par groupe ; maximum pour les dates. Population vide → compteurs à 0."""
    total: defaultdict = defaultdict(int, dict.fromkeys(_DATES_SYNTHESE))
    for g in groupes:
        for cle, v in g.items():
            if cle in ("profil", "niveau_groupe"):
                continue
            if cle in _DATES_SYNTHESE:
                if v is not None and (total[cle] is None or v > total[cle]):
                    total[cle] = v
            else:
                total[cle] += int(v or 0)
    return total


def motif_enregistre(r: dict) -> tuple[str | None, str | None]:
    """Motif tel qu'enregistré et champ d'origine — jamais reconstitué."""
    for champ in ("motif_decision_manuelle", "motif_moteur", "motif_source", "motif_risque", "motif_classement"):
        if r.get(champ):
            return r[champ], champ
    return None, None


class ClienteleVue360Service:
    def __init__(self, db: AsyncSession, ctx: Ctx):
        self.db = db
        self.ctx = ctx
        self.svc = ClienteleService(db, ctx)

    @property
    def voit_classif(self) -> bool:
        return self.ctx.peut("clientele.classif.view")

    # ------------------------------------------------------------------ filtres
    def _where(self, filtres: dict, periode: Periode | None = None) -> tuple[str, dict]:
        """Fragments SQL constants uniquement ; toutes les valeurs passent en paramètres liés.

        Sans clientele.classif.view, les filtres qui révèleraient un niveau de risque sont ignorés."""
        if not self.voit_classif:
            anomalie = _ANOMALIE.get(filtres.get("anomalie") or "")
            filtres = {k: v for k, v in filtres.items()
                       if k not in ("niveau", "profilage", "coherence", "classe_periode")
                       and not (k == "anomalie" and (not anomalie or anomalie.famille != "DONNEES"))}
        clauses, params = ["TRUE"], {}
        profil = filtres.get("profil")
        if profil in PROFILS:
            clauses.append("profil_derive = :profil")
            params["profil"] = profil
        if filtres.get("agence"):
            clauses.append("code_agence = :agence")
            params["agence"] = str(filtres["agence"]).strip()
        niveau = filtres.get("niveau")
        if niveau == "NON_EVALUE":
            clauses.append("niveau IS NULL")
        elif niveau in NIVEAUX:
            clauses.append("niveau = :niveau")
            params["niveau"] = niveau
        if filtres.get("profilage") in PROFILAGE:
            clauses.append(PROFILAGE[filtres["profilage"]])
        if filtres.get("filtrage") in FILTRAGE:
            clauses.append(FILTRAGE[filtres["filtrage"]])
        if filtres.get("anomalie") in _ANOMALIE:
            clauses.append(_ANOMALIE[filtres["anomalie"]].colonne)
        coherence = filtres.get("coherence")
        if coherence == "A_VERIFIER":
            clauses.append("f_a_verifier")
        elif coherence == "OK":
            clauses.append("NOT f_a_verifier AND NOT f_non_evalue")
        completude = filtres.get("completude")
        if completude == "INCOMPLET":
            clauses.append("f_incomplet")
        elif completude == "COMPLET":
            clauses.append("NOT f_incomplet")
        q = (filtres.get("q") or "").strip()
        if q:
            clauses.append("(racine_client LIKE :q_like OR nom_client ILIKE :q_like OR nni = :q OR nif = :q "
                           "OR racine_client IN (SELECT racine_client FROM clientele_comptes "
                           "WHERE compte = :q OR rib = :q))")
            params["q"], params["q_like"] = q, f"%{q}%"
        if filtres.get("classe_periode") and periode:
            debut, fin = periode.bornes()
            clauses.append("classifie_le >= :cp_debut AND classifie_le < :cp_fin")
            params["cp_debut"], params["cp_fin"] = debut, fin
        return " AND ".join(clauses), params

    def _sql(self, filtres: dict, corps: str, periode: Periode | None = None, *, leger: bool = False,
             **extra: Any) -> tuple[Any, dict]:
        where, params = self._where(filtres, periode)
        if leger and set(filtres) <= _FILTRES_LEGERS:
            return text(_BASE_LEGERE.replace("{where}", where) + corps), {**params, **extra}
        liste = critere_par_code("LISTE_INTERDICTION") or {}
        params.update({
            "poids_eleve": 100_000,
            "seuil_eleve": int(SEUILS_CDC["ELEVE_MIN"]),
            "seuil_moyen": int(SEUILS_CDC["FAIBLE_MAX"]),
            "liste_hors_score": liste.get("statut") == "SUSPENDU",
            **extra,
        })
        return text(_BASE.replace("{where}", where) + corps), params

    def population(self, filtres: dict) -> Any | None:
        """Sous-requête des racines de la population (filtres de diagnostic), ou None si aucun."""
        lourds = {k: v for k, v in filtres.items() if k not in _FILTRES_LEGERS}
        if not lourds:
            return None
        stmt, params = self._sql(lourds, "SELECT racine_client FROM pop")
        return stmt.bindparams(**params).columns(column("racine_client", String))

    async def _lignes(self, filtres: dict, corps: str, periode: Periode | None = None, *, leger: bool = False,
                      **extra: Any) -> list[dict]:
        """``leger`` : le corps n'utilise que racine_client de ``pop``."""
        stmt, params = self._sql(filtres, corps, periode, leger=leger, **extra)
        return [dict(r) for r in (await self.db.execute(stmt, params)).mappings().all()]

    # ------------------------------------------------------------------ synthèse
    async def synthese(self, filtres: dict) -> dict:
        self.ctx.exiger("clientele.filtrage.view")
        p = resoudre_periode(filtres.get("periode"), filtres.get("du"), filtres.get("au"))
        anomalies_sql = ", ".join(f"count(*) FILTER (WHERE {a.colonne}) AS {a.colonne}" for a in ANOMALIES)
        niveaux_sql = ", ".join(f"count(*) FILTER (WHERE niveau = '{n}') AS n_{n.lower()}" for n in NIVEAUX)
        groupes = await self._lignes(filtres, f"""
            SELECT profil_derive AS profil, coalesce(niveau, 'NON_EVALUE') AS niveau_groupe,
                   count(*) AS clients,
                   count(*) FILTER (WHERE profil_derive = 'PP') AS pp,
                   count(*) FILTER (WHERE profil_derive = 'PM') AS pm,
                   count(*) FILTER (WHERE profil_derive = 'NON_IDENTIFIE') AS non_identifies,
                   count(*) FILTER (WHERE NOT f_non_evalue) AS profiles,
                   {niveaux_sql},
                   count(*) FILTER (WHERE f_a_verifier) AS a_verifier,
                   count(*) FILTER (WHERE f_incomplet) AS incomplets,
                   count(*) FILTER (WHERE al_total > 0) AS clients_avec_correspondance,
                   count(*) FILTER (WHERE al_a_examiner > 0) AS clients_a_examiner,
                   coalesce(sum(al_a_examiner), 0) AS al_a_examiner,
                   coalesce(sum(al_en_analyse), 0) AS al_en_analyse,
                   coalesce(sum(al_confirmees), 0) AS al_confirmees,
                   coalesce(sum(al_ecartees), 0) AS al_ecartees,
                   coalesce(sum(al_cloturees), 0) AS al_cloturees,
                   coalesce(sum(reexamens), 0) AS reexamens,
                   {anomalies_sql},
                   max(classifie_le) AS dernier_classement, max(eval_le) AS derniere_evaluation,
                   max(date_extraction) AS stock_au
            FROM pop GROUP BY 1, 2 ORDER BY 1, 2""", p)
        k = _cumuler(groupes)
        clients = int(k["clients"])
        risque_profil = [{"profil": g["profil"], "niveau": g["niveau_groupe"], "clients": g["clients"]}
                         for g in groupes]

        voit = self.voit_classif
        out: dict[str, Any] = {
            "reference": {
                "stock_au": _iso(k["stock_au"]), "periode": p.dict(), "fuseau": FUSEAU,
                "genere_le": maintenant().isoformat(),
                "note": "Les indicateurs de stock (clients, risque, profilage) décrivent l'état courant "
                        "au stock indiqué ; seuls les indicateurs d'activité suivent la période.",
            },
            "autorisations": {"classification": voit, "audit": self.ctx.peut(*PERMS_AUDIT)},
            "clients": {"total": clients, "pp": int(k["pp"]), "pm": int(k["pm"]),
                        "non_identifies": int(k["non_identifies"])},
            "profilage": {
                "exploitables": int(k["profiles"]), "non_profiles": int(k["f_non_evalue"]),
                "a_arbitrer": int(k["f_a_arbitrer"]), "a_verifier": int(k["a_verifier"]),
                "incomplets": int(k["incomplets"]),
                "dernier_classement": _iso(k["dernier_classement"]),
                "derniere_evaluation": _iso(k["derniere_evaluation"]),
            } if voit else None,
            "risque": ({
                **{n: int(k[f"n_{n.lower()}"]) for n in NIVEAUX}, "NON_EVALUE": int(k["f_non_evalue"]),
            } if voit else None),
            "risque_par_profil": [{**r, "clients": int(r["clients"])} for r in risque_profil] if voit else None,
            "filtrage": {
                "clients_avec_correspondance": int(k["clients_avec_correspondance"]),
                "clients_sans_correspondance": clients - int(k["clients_avec_correspondance"]),
                "clients_a_examiner": int(k["clients_a_examiner"]),
                "alertes": {"A_EXAMINER": int(k["al_a_examiner"]), "EN_ANALYSE": int(k["al_en_analyse"]),
                            "CONFIRMEE": int(k["al_confirmees"]), "ECARTEE": int(k["al_ecartees"]),
                            "CLOTUREE": int(k["al_cloturees"])},
            },
            "reexamens_ouverts": int(k["reexamens"]) if voit else None,
            "anomalies": [{
                "code": a.code, "libelle": a.libelle, "definition": a.definition, "famille": a.famille,
                "clients": int(k[a.colonne]),
            } for a in ANOMALIES if voit or a.famille == "DONNEES"],
        }
        out["alertes"] = await self._alertes_agregats(filtres, p)
        out["activite"] = await self._activite(filtres, p)
        out["historique"] = await self._historique(filtres, p)
        out["sources_filtrage"] = await self._sources_filtrage(filtres)
        return out

    async def _alertes_agregats(self, filtres: dict, p: Periode) -> dict:
        rows = await self._lignes(filtres, """
            SELECT al.statut, al.motif, (al.assignee_id IS NULL) AS libre,
                   CASE WHEN al.statut NOT IN ('NOUVELLE', 'A_ANALYSER', 'EN_INVESTIGATION') THEN NULL
                        WHEN al.created_at > now() - interval '7 days' THEN '0_7'
                        WHEN al.created_at > now() - interval '30 days' THEN '8_30'
                        WHEN al.created_at > now() - interval '90 days' THEN '31_90'
                        ELSE '90_PLUS' END AS age,
                   count(*) AS n, min(al.created_at) FILTER (
                       WHERE al.statut IN ('NOUVELLE', 'A_ANALYSER', 'EN_INVESTIGATION')) AS plus_ancienne
            FROM clientele_alertes al JOIN pop USING (racine_client)
            GROUP BY 1, 2, 3, 4""", p, leger=True)
        par_statut: dict[str, int] = {}
        par_motif: dict[str, int] = {}
        par_source = {"ORION": 0, "LISTE": 0}
        anciennete = {"0_7": 0, "8_30": 0, "31_90": 0, "90_PLUS": 0}
        ouvertes_libres = 0
        plus_ancienne = None
        for r in rows:
            n = int(r["n"])
            par_statut[r["statut"]] = par_statut.get(r["statut"], 0) + n
            par_motif[r["motif"]] = par_motif.get(r["motif"], 0) + n
            par_source["ORION" if r["motif"] == "LISTE_INTERDICTION" else "LISTE"] += n
            if r["age"]:
                anciennete[r["age"]] += n
                if r["libre"]:
                    ouvertes_libres += n
                if r["plus_ancienne"] and (plus_ancienne is None or r["plus_ancienne"] < plus_ancienne):
                    plus_ancienne = r["plus_ancienne"]
        ouvertes = sum(par_statut.get(s, 0) for s in GROUPES["OUVERTES"])
        return {
            "total": sum(par_statut.values()), "ouvertes": ouvertes, "ouvertes_non_attribuees": ouvertes_libres,
            "en_attente_decision": par_statut.get("EN_INVESTIGATION", 0),
            "par_statut": par_statut, "par_groupe": {g: sum(par_statut.get(s, 0) for s in sts)
                                                     for g, sts in GROUPES.items()},
            "par_motif": par_motif, "par_source": par_source, "anciennete_ouvertes": anciennete,
            "plus_ancienne_ouverte": _iso(plus_ancienne),
            "priorite": {"statut": "NON_DEFINI", "message": "Aucun niveau de priorité n'est enregistré sur les "
                         "alertes ; le score affiché est un indice de similarité, pas une priorité."},
            "delais": {"statut": "NON_DEFINI", "message": "Aucun délai métier (SLA) n'est paramétré : "
                       "les alertes en retard ne peuvent pas être calculées. L'ancienneté est affichée."},
        }

    async def _activite(self, filtres: dict, p: Periode) -> dict:
        d1, f1 = p.bornes()
        d0, f0 = p.bornes(precedente=True)
        voit = self.voit_classif
        corps = """
            SELECT
              (SELECT count(*) FROM clientele_alertes al JOIN pop USING (racine_client)
                WHERE al.created_at >= :d AND al.created_at < :f) AS alertes_creees,
              (SELECT count(*) FROM clientele_alertes al JOIN pop USING (racine_client)
                WHERE al.created_at >= :d0 AND al.created_at < :f0) AS alertes_creees_prec,
              (SELECT count(*) FROM clientele_alerte_evenements ev
                 JOIN clientele_alertes al ON al.id = ev.alerte_id JOIN pop USING (racine_client)
                WHERE ev.decision IS NOT NULL AND ev.created_at >= :d AND ev.created_at < :f) AS decisions,
              (SELECT count(*) FROM clientele_alerte_evenements ev
                 JOIN clientele_alertes al ON al.id = ev.alerte_id JOIN pop USING (racine_client)
                WHERE ev.decision IS NOT NULL AND ev.created_at >= :d0 AND ev.created_at < :f0) AS decisions_prec,
              (SELECT count(*) FROM clientele_classif_historique h JOIN pop USING (racine_client)
                WHERE h.created_at >= :d AND h.created_at < :f) AS classifs,
              (SELECT count(*) FROM clientele_classif_historique h JOIN pop USING (racine_client)
                WHERE h.created_at >= :d0 AND h.created_at < :f0) AS classifs_prec,
              (SELECT count(*) FROM clientele_classif_historique h JOIN pop USING (racine_client)
                WHERE h.created_at >= :d AND h.created_at < :f AND h.ancienne_classe IS NULL) AS premieres,
              (SELECT count(*) FROM clientele_classif_historique h JOIN pop USING (racine_client)
                WHERE h.created_at >= :d0 AND h.created_at < :f0 AND h.ancienne_classe IS NULL) AS premieres_prec,
              (SELECT count(*) FROM clientele_classif_historique h JOIN pop USING (racine_client)
                WHERE h.created_at >= :d AND h.created_at < :f
                  AND h.ancienne_classe IS NOT NULL AND h.ancienne_classe <> h.nouvelle_classe) AS changements,
              (SELECT count(*) FROM clientele_classif_historique h JOIN pop USING (racine_client)
                WHERE h.created_at >= :d0 AND h.created_at < :f0
                  AND h.ancienne_classe IS NOT NULL AND h.ancienne_classe <> h.nouvelle_classe) AS changements_prec
        """
        r = (await self._lignes(filtres, corps, p, leger=True, d=d1, f=f1, d0=d0, f0=f0))[0]

        def paire(cle: str) -> dict:
            return {"valeur": int(r[cle]), "precedente": int(r[f"{cle}_prec"])}

        return {
            "alertes_creees": paire("alertes_creees"),
            "decisions_alertes": paire("decisions"),
            "classifications_enregistrees": paire("classifs") if voit else None,
            "premieres_classifications": paire("premieres") if voit else None,
            "changements_de_classe": paire("changements") if voit else None,
            "note": "Une variation du nombre de classifications enregistrées reflète des traitements "
                    "(classements, reclassements), pas une évolution du risque réel de la clientèle.",
        }

    async def _historique(self, filtres: dict, p: Periode) -> dict:
        d1, f1 = p.bornes()
        pas = _pas(p)
        series: dict[str, list[dict]] = {}
        if self.voit_classif:
            rows = await self._lignes(filtres, f"""
                SELECT date_trunc('{pas}', h.created_at AT TIME ZONE 'UTC')::date AS jour, h.source,
                       count(*) AS ecrits,
                       count(*) FILTER (WHERE h.ancienne_classe IS NULL) AS premieres,
                       count(*) FILTER (WHERE h.ancienne_classe IS NOT NULL
                                        AND h.ancienne_classe <> h.nouvelle_classe) AS changements
                FROM clientele_classif_historique h JOIN pop USING (racine_client)
                WHERE h.created_at >= :d AND h.created_at < :f
                GROUP BY 1, 2 ORDER BY 1, 2""", p, leger=True, d=d1, f=f1)
            series["classifications"] = [{"jour": _iso(r["jour"]), "source": r["source"],
                                          "ecrits": int(r["ecrits"]), "premieres": int(r["premieres"]),
                                          "changements": int(r["changements"])} for r in rows]
        rows = await self._lignes(filtres, f"""
            SELECT jour, sum(creees) AS creees, sum(decisions) AS decisions FROM (
                SELECT date_trunc('{pas}', al.created_at AT TIME ZONE 'UTC')::date AS jour,
                       count(*) AS creees, 0 AS decisions
                FROM clientele_alertes al JOIN pop USING (racine_client)
                WHERE al.created_at >= :d AND al.created_at < :f GROUP BY 1
                UNION ALL
                SELECT date_trunc('{pas}', ev.created_at AT TIME ZONE 'UTC')::date, 0, count(*)
                FROM clientele_alerte_evenements ev
                JOIN clientele_alertes al ON al.id = ev.alerte_id JOIN pop USING (racine_client)
                WHERE ev.decision IS NOT NULL AND ev.created_at >= :d AND ev.created_at < :f GROUP BY 1
            ) x GROUP BY jour ORDER BY jour""", p, leger=True, d=d1, f=f1)
        series["alertes"] = [{"jour": _iso(r["jour"]), "creees": int(r["creees"]),
                              "decisions": int(r["decisions"])} for r in rows]
        return {"pas": pas, **series}

    async def _sources_filtrage(self, filtres: dict) -> dict:
        scan = (await self.db.execute(
            select(AuditLog.created_at, AuditLog.after_data, User.full_name)
            .outerjoin(User, User.id == AuditLog.user_id)
            .where(AuditLog.action == "clientele.filtrage.scan")
            .order_by(AuditLog.created_at.desc()).limit(1))).first()
        listes = (await self.db.execute(text("""
            SELECT l.id, l.code, l.libelle, l.source, l.actif, l.created_at,
                   count(e.id) FILTER (WHERE e.actif) AS entrees, max(e.created_at) AS derniere_entree
            FROM clientele_filtrage_listes l LEFT JOIN clientele_filtrage_entrees e ON e.liste_id = l.id
            GROUP BY l.id ORDER BY l.code"""))).mappings().all()
        codes = await self._lignes(filtres, """
            SELECT code, count(DISTINCT racine_client) AS clients, count(*) AS comptes FROM (
                SELECT co.racine_client, regexp_replace(trim(tok), '\\s+', ' ', 'g') AS code
                FROM clientele_comptes co JOIN pop USING (racine_client)
                CROSS JOIN LATERAL regexp_split_to_table(coalesce(co.liste_interdiction, ''), ',') tok
            ) x WHERE code <> '' GROUP BY code ORDER BY clients DESC LIMIT 12""", leger=True)
        extraction = await self.db.scalar(text("SELECT max(date_extraction) FROM clientele_situation"))
        nouveaux = None
        if scan:
            nouveaux = await self.db.scalar(text(
                "SELECT count(*) FROM clientele_clients WHERE created_at > :le"), {"le": scan[0]})
        return {
            "dernier_scan": {"le": _iso(scan[0]), "par": scan[2],
                             "alertes_creees": (scan[1] or {}).get("crees")} if scan else None,
            "clients_crees_depuis_scan": int(nouveaux) if nouveaux is not None else None,
            "orion": {"date_extraction": _iso(extraction),
                      "codes": [{"code": c["code"], "clients": int(c["clients"]), "comptes": int(c["comptes"])}
                                for c in codes],
                      "note": "Codes bruts du champ LISTE_INTERDICTION d'ORION (tous comptes). Leur signification "
                              "(DI, OI…) n'est pas documentée dans BEA DIGITAL : à confirmer par la Conformité."},
            "listes": [{
                "id": str(row["id"]), "code": row["code"], "libelle": row["libelle"], "source": row["source"],
                "actif": row["actif"], "entrees": int(row["entrees"]), "creee_le": _iso(row["created_at"]),
                "derniere_entree": _iso(row["derniere_entree"]),
                "preuve_import": False,
            } for row in listes],
            "note_listes": "Aucun journal d'import externe n'existe pour les listes : seule la date de la "
                           "dernière entrée saisie est connue. L'exhaustivité et la fraîcheur ne sont pas prouvées.",
        }

    # ------------------------------------------------------------------ clients
    async def clients(self, filtres: dict, page: int = 1, taille: int = 50, tri: str | None = None) -> dict:
        self.ctx.exiger("clientele.filtrage.view")
        p = resoudre_periode(filtres.get("periode"), filtres.get("du"), filtres.get("au"))
        page, taille = max(1, page), min(max(1, taille), 200)
        ordre = {
            "score": "score_calcule DESC NULLS LAST, racine_client",
            "classement": "classifie_le DESC NULLS LAST, racine_client",
            "nom": "nom_client, racine_client",
            "alertes": "al_a_examiner + al_en_analyse DESC, racine_client",
        }.get(tri or "", "racine_client")
        total = (await self._lignes(filtres, "SELECT count(*) AS n FROM pop", p))[0]["n"]
        rows = await self._lignes(filtres, f"SELECT * FROM pop ORDER BY {ordre} LIMIT :lim OFFSET :off", p,
                                  lim=taille, off=(page - 1) * taille)
        return {"total": int(total), "page": page, "taille": taille, "periode": p.dict(),
                "items": [self._ligne_client(r) for r in rows]}

    def _ligne_client(self, r: dict) -> dict:
        voit = self.voit_classif
        motif, origine = motif_enregistre(r) if voit else (None, None)
        anomalies = [a.code for a in ANOMALIES if r.get(a.colonne) and (voit or a.famille == "DONNEES")]
        if not voit:
            coherence = None
        elif r["f_non_evalue"]:
            coherence = "NON_EVALUE"
        else:
            coherence = "A_VERIFIER" if r["f_a_verifier"] else "OK"
        return {
            "racine_client": r["racine_client"], "nom": r["nom_client"], "prenoms": r["prenoms"],
            "profil": r["profil_derive"], "code_agence": r["code_agence"], "agence": r["agence"],
            "nb_comptes": int(r["nb_comptes"] or 0), "etat_client": r["etat_client"],
            "classification": {
                "niveau": r["niveau"], "source": r["source"], "mode_decision": r["mode_decision"],
                "score_enregistre": r["score_calcule"], "niveau_propose": r["niveau_propose"],
                "motif_enregistre": motif, "motif_champ": origine,
                "classifie_le": _iso(r["classifie_le"]), "version_regles": r["version_regles"],
                "derniere_evaluation": {"statut": r["eval_statut"], "niveau": r["eval_niveau"],
                                        "score": r["eval_score"], "le": _iso(r["eval_le"])},
            } if voit else None,
            "completude": "INCOMPLET" if r["f_incomplet"] else "COMPLET",
            "coherence": coherence,
            "anomalies": anomalies,
            "alertes": {"a_examiner": int(r["al_a_examiner"]), "en_analyse": int(r["al_en_analyse"]),
                        "confirmees": int(r["al_confirmees"]), "ecartees": int(r["al_ecartees"]),
                        "cloturees": int(r["al_cloturees"]), "total": int(r["al_total"])},
            "reexamens": int(r["reexamens"]) if voit else None,
        }

    # ------------------------------------------------------------------ fiche
    async def fiche(self, racine: str) -> dict:
        self.ctx.exiger("clientele.filtrage.view")
        racine = (racine or "").strip()
        if not re.fullmatch(r"[0-9]{6}", racine):
            raise AppError("Racine client invalide (6 chiffres)", 422, code="RACINE_INVALIDE")
        base = await ClienteleSituationService(self.db, self.ctx).fiche(racine)
        ligne = await self._lignes({}, "SELECT * FROM pop WHERE racine_client = :r", None, r=racine)
        voit = self.voit_classif
        out: dict[str, Any] = {
            "client": base["client"], "situation": base["situation"], "comptes": base["comptes"],
            "dernier_import": base["dernier_import"],
            "synthese": self._ligne_client(ligne[0]) if ligne else None,
            "diagnostics": [{"code": a.code, "libelle": a.libelle, "definition": a.definition, "famille": a.famille}
                            for a in ANOMALIES if ligne and ligne[0].get(a.colonne)
                            and (voit or a.famille == "DONNEES")],
            "autorisations": {"classification": voit},
        }
        if voit:
            classif = ClienteleClassificationService(self.db, self.ctx)
            out["classification"] = base["classification"]
            out["historique_classification"] = await classif.historique(racine)
            out["derniere_evaluation"] = await self._derniere_evaluation(racine)
            out["reexamens"] = [{
                "id": str(x.id), "statut": x.statut, "origine": x.origine, "niveau_actuel": x.niveau_actuel,
                "niveau_suggere": x.niveau_suggere, "source_suggestion": x.source_suggestion, "motif": x.motif,
                "decision": x.decision, "cree_le": _iso(x.created_at), "traite_le": _iso(x.traite_le),
            } for x in (await self.db.scalars(
                select(ClienteleClassifReexamen).where(ClienteleClassifReexamen.racine_client == racine)
                .order_by(ClienteleClassifReexamen.created_at.desc()))).all()]
        alertes = (await self.db.execute(text("""
            SELECT al.id, al.statut, al.motif, al.score, al.precedent_faux_positif, al.created_at, al.updated_at,
                   al.cloturee_le, u.full_name AS assignee,
                   (SELECT max(ev.created_at) FROM clientele_alerte_evenements ev WHERE ev.alerte_id = al.id)
                       AS derniere_action
            FROM clientele_alertes al LEFT JOIN users u ON u.id = al.assignee_id
            WHERE al.racine_client = :r ORDER BY al.created_at DESC"""), {"r": racine})).mappings().all()
        out["alertes"] = [{
            "id": str(a["id"]), "statut": a["statut"], "motif": a["motif"], "similarite": a["score"],
            "source": "ORION" if a["motif"] == "LISTE_INTERDICTION" else "LISTE",
            "precedent_faux_positif": a["precedent_faux_positif"], "assignee": a["assignee"],
            "cree_le": _iso(a["created_at"]), "derniere_action": _iso(a["derniere_action"]),
            "cloturee_le": _iso(a["cloturee_le"]),
        } for a in alertes]
        ids = [racine, *(a["id"] for a in out["alertes"])]
        journal = (await self.db.execute(
            select(AuditLog.action, AuditLog.entity, AuditLog.entity_id, AuditLog.created_at, User.full_name)
            .outerjoin(User, User.id == AuditLog.user_id)
            .where(AuditLog.module_code == "clientele", AuditLog.entity_id.in_(ids))
            .order_by(AuditLog.created_at.desc()).limit(100))).all()
        out["journal"] = [{"action": a, "objet": e, "objet_id": eid, "le": _iso(le), "utilisateur": nom}
                          for a, e, eid, le, nom in journal]
        return out

    async def _derniere_evaluation(self, racine: str) -> dict | None:
        ev = await self.db.scalar(
            select(ClienteleClassifEvaluation).where(ClienteleClassifEvaluation.racine_client == racine)
            .order_by(ClienteleClassifEvaluation.created_at.desc()).limit(1))
        if not ev:
            return None
        lignes = (await self.db.scalars(
            select(ClienteleClassifEvaluationLigne).where(ClienteleClassifEvaluationLigne.evaluation_id == ev.id)
            .order_by(ClienteleClassifEvaluationLigne.contribue_au_score.desc(),
                      ClienteleClassifEvaluationLigne.poids.desc().nulls_last()))).all()
        return {
            "id": str(ev.id), "le": _iso(ev.created_at), "score_total": ev.score_total,
            "niveau_final": ev.niveau_final, "statut": ev.statut, "coherence": ev.coherence,
            "motif_principal": ev.motif_principal, "version_moteur": ev.version_moteur,
            "version_regles": ev.version_regles, "date_reference": _iso(ev.date_reference),
            "lignes": [{
                "critere": lg.critere, "libelle": lg.libelle, "etat": lg.etat, "valeur": lg.valeur,
                "poids": lg.poids, "niveau_retenu": lg.niveau_retenu, "statut_regle": lg.statut_regle,
                "contribue_au_score": lg.contribue_au_score, "blocking_propose": lg.blocking_propose,
                "motif": lg.motif,
            } for lg in lignes],
            "note": "Photographie enregistrée au moment de l'évaluation : rien n'est recalculé à l'ouverture.",
        }

    # ------------------------------------------------------------------ audit
    async def journal(self, *, action: str | None = None, objet_id: str | None = None,
                      du: date | None = None, au: date | None = None, page: int = 1, taille: int = 50) -> dict:
        self.ctx.exiger("clientele.filtrage.view")
        if not self.ctx.peut(*PERMS_AUDIT):
            raise AppError("Journal d'audit réservé aux profils de supervision", 403, code="PERMISSION_DENIED")
        page, taille = max(1, page), min(max(1, taille), 200)
        q = (select(AuditLog.id, AuditLog.action, AuditLog.entity, AuditLog.entity_id, AuditLog.created_at,
                    User.full_name)
             .outerjoin(User, User.id == AuditLog.user_id)
             .where(AuditLog.module_code == "clientele"))
        if action:
            q = q.where(AuditLog.action == action)
        if objet_id:
            q = q.where(AuditLog.entity_id == objet_id.strip())
        if du:
            q = q.where(AuditLog.created_at >= datetime.combine(du, time.min, tzinfo=UTC))
        if au:
            q = q.where(AuditLog.created_at < datetime.combine(au + timedelta(days=1), time.min, tzinfo=UTC))
        total = await self.db.scalar(select(func.count()).select_from(q.subquery())) or 0
        rows = (await self.db.execute(
            q.order_by(AuditLog.created_at.desc()).offset((page - 1) * taille).limit(taille))).all()
        actions = (await self.db.execute(text(
            "SELECT action, count(*) FROM audit_logs WHERE module_code = 'clientele' "
            "GROUP BY action ORDER BY action"))).all()
        return {
            "total": int(total), "page": page, "taille": taille,
            "actions": [{"action": a, "n": int(n)} for a, n in actions],
            "items": [{"id": str(i), "action": a, "objet": e, "objet_id": eid, "le": _iso(le), "utilisateur": nom}
                      for i, a, e, eid, le, nom in rows],
            "note": "Lecture seule. Les valeurs avant / après ne sont pas exposées ici (données personnelles).",
        }

    # ------------------------------------------------------------------ export
    async def export_xlsx(self, filtres: dict) -> tuple[bytes, int]:
        self.ctx.exiger("clientele.filtrage.view")
        self.ctx.exiger("clientele.export")
        s = await self.synthese(filtres)
        p = resoudre_periode(filtres.get("periode"), filtres.get("du"), filtres.get("au"))
        rows = await self._lignes(filtres, "SELECT * FROM pop ORDER BY racine_client LIMIT :lim", p,
                                  lim=LIMITE_EXPORT + 1)
        tronque = len(rows) > LIMITE_EXPORT
        rows = rows[:LIMITE_EXPORT]
        wb = Workbook()
        ws = wb.active
        ws.title = "Synthèse"
        gras = Font(bold=True)
        ref = s["reference"]
        entete = [
            ("Vue 360° clientèle — BEA DIGITAL", ""),
            ("Généré le", ref["genere_le"]),
            ("Généré par", self.ctx.user.full_name),
            ("Stock au", ref["stock_au"] or "aucun stock chargé"),
            ("Période d'activité", f"{ref['periode']['du']} → {ref['periode']['au']} ({FUSEAU})"),
            ("Filtres", ", ".join(f"{k}={v}" for k, v in filtres.items() if v not in (None, "", False)) or "aucun"),
            ("Limites", ref["note"]),
            ("", ""),
            ("Clients distincts (racines)", s["clients"]["total"]),
            ("PP", s["clients"]["pp"]), ("PM", s["clients"]["pm"]),
            ("Profil non identifié", s["clients"]["non_identifies"]),
        ]
        if s["profilage"]:
            entete += [("Clients profilés (classification enregistrée)", s["profilage"]["exploitables"]),
                       ("Clients non profilés", s["profilage"]["non_profiles"]),
                       ("Clients à vérifier (cohérence)", s["profilage"]["a_verifier"])]
            entete += [(f"Risque {k}", v) for k, v in s["risque"].items()]
        entete += [("Clients avec correspondance de filtrage", s["filtrage"]["clients_avec_correspondance"]),
                   ("Alertes ouvertes", s["alertes"]["ouvertes"]),
                   ("Alertes ouvertes non attribuées", s["alertes"]["ouvertes_non_attribuees"]),
                   ("Priorité des alertes", s["alertes"]["priorite"]["message"]),
                   ("Délais / retards", s["alertes"]["delais"]["message"])]
        if tronque:
            entete.append(("Attention", f"Liste des clients tronquée à {LIMITE_EXPORT} lignes"))
        for row in entete:
            ws.append(list(row))
        ws["A1"].font = gras
        ws.column_dimensions["A"].width = 48
        ws.column_dimensions["B"].width = 90

        wa = wb.create_sheet("Anomalies")
        wa.append(["Code", "Libellé", "Famille", "Clients", "Définition (diagnostic, pas un motif moteur)"])
        for a in s["anomalies"]:
            wa.append([a["code"], a["libelle"], a["famille"], a["clients"], a["definition"]])

        wc = wb.create_sheet("Clients")
        voit = self.voit_classif
        cols = ["Racine", "Nom", "Profil", "Agence", "Comptes", "Complétude", "Anomalies",
                "Alertes à examiner", "Alertes en analyse", "Alertes confirmées"]
        if voit:
            cols[5:5] = ["Niveau enregistré", "Score enregistré", "Mode", "Motif enregistré", "Champ du motif",
                         "Classé le", "Version règles", "Cohérence"]
        wc.append(cols)
        for r in rows:
            c = self._ligne_client(r)
            ligne = [c["racine_client"], c["nom"], c["profil"], c["code_agence"], c["nb_comptes"],
                     c["completude"], " ".join(c["anomalies"]), c["alertes"]["a_examiner"],
                     c["alertes"]["en_analyse"], c["alertes"]["confirmees"]]
            if voit:
                k = c["classification"]
                ligne[5:5] = [k["niveau"] or "NON ÉVALUÉ", k["score_enregistre"], k["mode_decision"],
                              k["motif_enregistre"], k["motif_champ"], k["classifie_le"], k["version_regles"],
                              c["coherence"]]
            wc.append(ligne)
        for w in (wa, wc):
            for cell in w[1]:
                cell.font = gras
        buf = io.BytesIO()
        wb.save(buf)
        await self.svc.audit("clientele.vue360.export", "clientele_vue360", None,
                             after={"filtres": {k: str(v) for k, v in filtres.items() if v not in (None, "")},
                                    "lignes": len(rows), "tronque": tronque})
        return buf.getvalue(), len(rows)


def filtres_depuis(**kw: Any) -> dict:
    return {k: v for k, v in kw.items() if v not in (None, "")}

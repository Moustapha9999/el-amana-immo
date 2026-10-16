"""Import Excel de la liste officielle des directions, départements et services dans l'organigramme CORE.

Une ligne = une entité feuille (la colonne la plus fine renseignée) avec son chemin ; les niveaux
intermédiaires sont retrouvés ou créés. Rapprochement par libellé normalisé (``org_cle``), sigle ou
alias, **sous le parent** : deux services homonymes de départements différents restent distincts.

Règles :
- aperçu sans écriture (CREER / EXISTANT / A_CONFIRMER / AMBIGU / ERREUR par ligne) ;
- application transactionnelle : refusée tant qu'une ligne est ambiguë ou en erreur (pas d'import
  partiel) ; crée les entités manquantes (source ``LISTE_BANQUE``), confirme les entités existantes
  « à confirmer » ; ne supprime ni ne désactive jamais rien ;
- les entités actives absentes du fichier sont listées pour décision humaine ;
- chaque création / confirmation est auditée (module ``core``).
"""

from __future__ import annotations

import hashlib
import io
import re
import unicodedata
import uuid
from dataclasses import dataclass, field

from fastapi import Request, UploadFile
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.exceptions import AppError
from app.models import User
from app.models.organisation import OrgEntite
from app.services import org_referentiel as org

SOURCE = "LISTE_BANQUE"
FEUILLE = "Organigramme"
TAILLE_MAX = 2 * 1024 * 1024
MAX_LIGNES = 2000
NIVEAUX = (("direction", "Direction", ("DIRECTION_GENERALE", "DIRECTION")),
           ("departement", "Département", ("DEPARTEMENT",)),
           ("service", "Service", ("SERVICE", "SECTION", "CELLULE")))
TYPES_SERVICE = ("SERVICE", "SECTION", "CELLULE")
# Niveaux suivis par la liste de la banque (les comités / conseils ne sont pas concernés).
TYPES_SUIVIS = ("DIRECTION", "DEPARTEMENT", *TYPES_SERVICE)
COLONNES = [
    ("direction", "Direction (libellé officiel)."),
    ("departement", "Département (libellé officiel). Vide si le service dépend directement de la direction."),
    ("service", "Service, section ou cellule (libellé officiel). Vide pour une ligne décrivant une direction ou un "
                "département."),
    ("type_service", "SERVICE (défaut), SECTION ou CELLULE."),
    ("code", "Code de l'entité de la ligne (facultatif, unique ; généré sinon). Jamais modifié sur l'existant."),
    ("sigle", "Sigle de l'entité de la ligne (facultatif ; complété s'il est vide sur l'existant)."),
]
_ALIAS = {
    "direction": ("direction", "directions"),
    "departement": ("departement", "departements"),
    "service": ("service", "services", "service section cellule"),
    "type_service": ("type service", "type", "niveau"),
    "code": ("code",),
    "sigle": ("sigle",),
}


def _erreur(message: str, code: str, status: int = 422) -> AppError:
    return AppError(message, status, code=code)


def _texte(v) -> str:
    return org._propre(str(v)) if v is not None else ""


def _cle_entete(v) -> str:
    t = unicodedata.normalize("NFKD", str(v or "")).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", " ", t).strip()


# --------------------------------------------------------------------------- modèle
def modele() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = FEUILLE
    for i, (code, _aide) in enumerate(COLONNES, start=1):
        c = ws.cell(row=1, column=i, value=code)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
        ws.column_dimensions[c.column_letter].width = 42 if code in ("direction", "departement", "service") else 16
    ws.freeze_panes = "A2"
    aide = wb.create_sheet("Consignes")
    aide.column_dimensions["A"].width = 18
    aide.column_dimensions["B"].width = 100
    aide.append(["Import de la liste officielle des directions, départements et services"])
    aide["A1"].font = Font(bold=True, size=13)
    aide.append([])
    for texte in (
        "Une ligne par entité : renseignez son chemin complet (direction, puis département, puis service).",
        "Les niveaux intermédiaires absents de l'organigramme sont créés automatiquement.",
        "Rapprochement sur le libellé (casse, accents et ponctuation ignorés), le sigle ou un alias, sous le même "
        "parent : deux services homonymes de départements différents restent distincts.",
        "Rien n'est supprimé ni désactivé : les entités absentes du fichier sont seulement listées.",
        "Aucune écriture tant qu'une ligne est ambiguë ou en erreur : corrigez le fichier puis relancez.",
    ):
        aide.append(["", texte])
    aide.append([])
    aide.append(["Colonne", "Règle"])
    for code, regle in COLONNES:
        aide.append([code, regle])
    for row in aide.iter_rows():
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def lire_classeur(contenu: bytes) -> list[tuple[int, dict]]:
    try:
        wb = load_workbook(io.BytesIO(contenu), read_only=True, data_only=True)
    except Exception:
        raise _erreur("Fichier Excel illisible : utilisez le modèle (.xlsx)", "FICHIER_ILLISIBLE") from None
    try:
        ws = wb[FEUILLE] if FEUILLE in wb.sheetnames else wb.worksheets[0]
        lignes = ws.iter_rows(values_only=True)
        colonnes: dict[str, int] = {}
        numero = 0
        for numero, valeurs in enumerate(lignes, start=1):
            cles = [_cle_entete(v) for v in valeurs]
            colonnes = {champ: cles.index(a) for champ, alias in _ALIAS.items() for a in alias if a in cles}
            if {"direction", "service"} <= set(colonnes) or numero >= 10:
                break
        if not {"direction", "departement", "service"} <= set(colonnes):
            raise _erreur("En-têtes introuvables : direction, departement, service (téléchargez le modèle)",
                          "ENTETES_MANQUANTES")
        out: list[tuple[int, dict]] = []
        for rang, valeurs in enumerate(lignes, start=numero + 1):
            brut = {c: _texte(valeurs[i]) if i < len(valeurs) else "" for c, i in colonnes.items()}
            if not any(brut.values()):
                continue
            out.append((rang, brut))
            if len(out) > MAX_LIGNES:
                raise _erreur(f"{MAX_LIGNES} lignes au plus", "TROP_DE_LIGNES")
    finally:
        wb.close()
    if not out:
        raise _erreur("Le fichier ne contient aucune ligne", "FICHIER_VIDE")
    return out


# --------------------------------------------------------------------------- plan
@dataclass
class Noeud:
    """Entité du fichier (chemin normalisé) : existante (``entite``) ou à créer."""

    chemin: tuple[str, ...]
    libelle: str
    type_entite: str
    parent: Noeud | None
    entite: OrgEntite | None = None
    statut: str = "CREER"
    message: str | None = None
    code: str | None = None
    sigle: str | None = None
    lignes: list[int] = field(default_factory=list)


class Planificateur:
    def __init__(self, arbre: org.Arbre, cles: dict[str, str]):
        self.arbre = arbre
        self.cles = cles
        self.noeuds: dict[tuple[str, ...], Noeud] = {}
        actives_dg = [e for e in arbre.entites.values() if e.type_entite == "DIRECTION_GENERALE" and e.is_active]
        self.racine = actives_dg[0] if len(actives_dg) == 1 else None

    def cle(self, texte: str) -> str:
        return self.cles.get(texte, "")

    def _candidats(self, parent_id: uuid.UUID | None, libelle: str, types: tuple[str, ...],
                   *, partout: bool = False) -> list[OrgEntite]:
        c = self.cle(libelle)
        sigle = libelle.replace(" ", "").upper()
        if partout:
            pool = [e for e in self.arbre.entites.values() if e.type_entite in types]
        else:
            pool = [self.arbre.entites[i] for i in self.arbre.enfants.get(parent_id, []) if i in self.arbre.entites]
            pool = [e for e in pool if e.type_entite in types or e.type_entite in TYPES_SUIVIS]
        exacts = [e for e in pool if e.cle == c]
        if exacts:
            return exacts
        return [e for e in pool if c in (e.alias or []) or (e.sigle and e.sigle.replace(" ", "").upper() == sigle)]

    def noeud(self, chemin: tuple[str, ...], libelle: str, types: tuple[str, ...], type_cree: str,
              parent: Noeud | None, rang: int) -> Noeud:
        if chemin in self.noeuds:
            n = self.noeuds[chemin]
            n.lignes.append(rang)
            return n
        n = Noeud(chemin=chemin, libelle=libelle, type_entite=type_cree, parent=parent, lignes=[rang])
        self.noeuds[chemin] = n
        if parent is not None and parent.statut in ("AMBIGU", "ERREUR"):
            n.statut, n.message = "ERREUR", "parent non résolu"
            return n
        if parent is not None and parent.entite is None:
            return n
        if parent is None:
            candidats = self._candidats(None, libelle, types, partout=True)
        else:
            candidats = self._candidats(parent.entite.id, libelle, types)
        if len(candidats) > 1:
            n.statut = "AMBIGU"
            n.message = "plusieurs entités correspondent : " + " | ".join(
                self.arbre.libelle_chemin(e.id) or e.libelle for e in candidats[:4])
        elif candidats:
            e = candidats[0]
            n.entite, n.type_entite = e, e.type_entite
            if not e.is_active:
                n.statut, n.message = "ERREUR", "entité désactivée dans l'organigramme : réactivez-la d'abord"
            elif e.a_confirmer:
                n.statut = "A_CONFIRMER"
            else:
                n.statut = "EXISTANT"
        elif parent is None and self.racine is None:
            n.statut, n.message = "ERREUR", "direction inconnue et direction générale introuvable pour la rattacher"
        return n

    def ligne(self, rang: int, brut: dict) -> dict:
        valeurs = [(champ, nom, types, brut.get(champ, "")) for champ, nom, types in NIVEAUX]
        renseignes = [v for v in valeurs if v[3]]
        resultat = {"ligne": rang, "direction": brut.get("direction") or None,
                    "departement": brut.get("departement") or None, "service": brut.get("service") or None}
        if not renseignes:
            return {**resultat, "statut": "ERREUR", "message": "aucun libellé"}
        if not brut.get("direction") and not brut.get("departement"):
            return {**resultat, "statut": "ERREUR", "message": "indiquez au moins la direction ou le département"}
        type_service = (brut.get("type_service") or "SERVICE").upper()
        if type_service not in TYPES_SERVICE:
            return {**resultat, "statut": "ERREUR", "message": "type_service : SERVICE, SECTION ou CELLULE"}
        parent: Noeud | None = None
        chemin: tuple[str, ...] = ()
        n: Noeud | None = None
        for champ, _nom, types, libelle in renseignes:
            if not self.cle(libelle):
                return {**resultat, "statut": "ERREUR", "message": f"{champ} illisible"}
            chemin = (*chemin, f"{champ}:{self.cle(libelle)}")
            type_cree = {"direction": "DIRECTION", "departement": "DEPARTEMENT"}.get(champ, type_service)
            if parent is None and champ != "direction":
                # Département sans direction : rapproché dans tout l'organigramme (unique exigé).
                candidats = self._candidats(None, libelle, types, partout=True)
                n = self.noeuds.get(chemin) or Noeud(chemin=chemin, libelle=libelle, type_entite=type_cree,
                                                      parent=None, lignes=[])
                self.noeuds[chemin] = n
                n.lignes.append(rang)
                if len(candidats) == 1:
                    n.entite, n.type_entite = candidats[0], candidats[0].type_entite
                    n.statut = "A_CONFIRMER" if candidats[0].a_confirmer else "EXISTANT"
                else:
                    n.statut = "AMBIGU" if candidats else "ERREUR"
                    n.message = ("plusieurs départements portent ce libellé : précisez la direction" if candidats
                                 else "département inconnu : précisez sa direction")
            else:
                n = self.noeud(chemin, libelle, types, type_cree, parent, rang)
            parent = n
        if brut.get("code"):
            n.code = org._code_propre(brut["code"])
        if brut.get("sigle"):
            n.sigle = brut["sigle"][:40]
        chemin_lib = self.arbre.libelle_chemin(n.entite.id) if n.entite else " › ".join(
            x for x in (brut.get("direction"), brut.get("departement"), brut.get("service")) if x)
        return {**resultat, "statut": n.statut, "message": n.message, "type_entite": n.type_entite,
                "entite_id": str(n.entite.id) if n.entite else None, "chemin": chemin_lib,
                "_noeud": n}


async def _cles(db: AsyncSession, textes: set[str]) -> dict[str, str]:
    out: dict[str, str] = {}
    for t in textes:
        if t:
            out[t] = await db.scalar(select(func.org_cle(t))) or ""
    return out


async def _plan(db: AsyncSession, contenu: bytes) -> tuple[Planificateur, list[dict]]:
    brutes = lire_classeur(contenu)
    textes = {v for _r, b in brutes for k, v in b.items() if k in ("direction", "departement", "service")}
    p = Planificateur(await org.charger_arbre(db), await _cles(db, textes))
    lignes = [p.ligne(rang, brut) for rang, brut in brutes]
    for lg in lignes:
        n = lg.get("_noeud")
        if n is not None:
            lg["statut"], lg["message"] = n.statut, n.message
    return p, lignes


def _absentes(p: Planificateur) -> list[dict]:
    vus = {n.entite.id for n in p.noeuds.values() if n.entite is not None}
    for n in p.noeuds.values():
        if n.entite is not None:
            vus |= {e.id for e in p.arbre.chemin(n.entite.id)}
    return [{"id": str(e.id), "libelle": e.libelle, "type_entite": e.type_entite,
             "chemin": p.arbre.libelle_chemin(e.id), "a_confirmer": e.a_confirmer}
            for e in sorted(p.arbre.entites.values(), key=lambda x: p.arbre.libelle_chemin(x.id) or "")
            if e.is_active and e.type_entite in TYPES_SUIVIS and e.id not in vus]


def _synthese(p: Planificateur, lignes: list[dict], empreinte: str) -> dict:
    compte: dict[str, int] = {}
    for n in p.noeuds.values():
        compte[n.statut] = compte.get(n.statut, 0) + 1
    bloquantes = [lg for lg in lignes if lg["statut"] in ("AMBIGU", "ERREUR")]
    return {
        "empreinte": empreinte,
        "lignes": [{k: v for k, v in lg.items() if k != "_noeud"} for lg in lignes],
        "entites": {"a_creer": compte.get("CREER", 0), "existantes": compte.get("EXISTANT", 0),
                    "a_confirmer": compte.get("A_CONFIRMER", 0), "ambigues": compte.get("AMBIGU", 0),
                    "erreurs": compte.get("ERREUR", 0)},
        "nb_lignes": len(lignes), "nb_lignes_bloquantes": len(bloquantes),
        "applicable": not bloquantes,
        "absentes_du_fichier": _absentes(p),
    }


async def _lire(fichier: UploadFile) -> bytes:
    nom = (fichier.filename or "").lower()
    if not nom.endswith(".xlsx"):
        raise _erreur("Format attendu : classeur Excel .xlsx (téléchargez le modèle)", "FORMAT_IMPORT")
    contenu = await fichier.read(TAILLE_MAX + 1)
    if len(contenu) > TAILLE_MAX:
        raise _erreur(f"Fichier trop volumineux ({TAILLE_MAX // (1024 * 1024)} Mo au plus)", "FICHIER_TROP_GROS", 413)
    return contenu


async def apercu(db: AsyncSession, fichier: UploadFile) -> dict:
    contenu = await _lire(fichier)
    p, lignes = await _plan(db, contenu)
    return _synthese(p, lignes, hashlib.sha256(contenu).hexdigest())


async def appliquer(db: AsyncSession, actor: User, fichier: UploadFile, empreinte_attendue: str | None,
                    request: Request | None = None) -> dict:
    """Rejoue l'aperçu sur l'état courant, puis écrit tout ou rien."""
    contenu = await _lire(fichier)
    empreinte = hashlib.sha256(contenu).hexdigest()
    if empreinte_attendue and empreinte_attendue != empreinte:
        raise _erreur("Le fichier diffère de celui de l'aperçu : relancez l'aperçu", "FICHIER_MODIFIE", 409)
    p, lignes = await _plan(db, contenu)
    synthese = _synthese(p, lignes, empreinte)
    if not synthese["applicable"]:
        raise _erreur(f"{synthese['nb_lignes_bloquantes']} ligne(s) ambiguë(s) ou en erreur : corrigez le fichier",
                      "IMPORT_BLOQUE")
    crees, confirmees = [], []
    for n in sorted(p.noeuds.values(), key=lambda x: len(x.chemin)):
        if n.statut == "CREER":
            parent = n.parent.entite if n.parent else p.racine
            code = n.code
            if code and await db.scalar(select(OrgEntite.id).where(func.upper(OrgEntite.code) == code)):
                raise _erreur(f"Le code {code} est déjà utilisé", "CODE_EXISTANT", 409)
            code = code or await org._code_auto(db, parent, n.type_entite)
            ordre = (await db.scalar(select(func.coalesce(func.max(OrgEntite.ordre), 0)).where(
                OrgEntite.parent_id == parent.id)) or 0) + 1
            e = OrgEntite(code=code, libelle=n.libelle, cle=p.cle(n.libelle), type_entite=n.type_entite,
                          parent_id=parent.id,
                          sigle=n.sigle, ordre=ordre, alias=[], a_confirmer=False, source=SOURCE,
                          metadonnees={"import_empreinte": empreinte}, is_active=True,
                          created_by_id=actor.id, updated_by_id=actor.id)
            db.add(e)
            await org._flush(db)
            n.entite = e
            crees.append(e)
        elif n.statut == "A_CONFIRMER" and n.entite is not None:
            avant = {"a_confirmer": True, "note_confirmation": n.entite.note_confirmation}
            n.entite.a_confirmer, n.entite.note_confirmation = False, None
            if n.sigle and not n.entite.sigle:
                n.entite.sigle = n.sigle
            n.entite.updated_by_id = actor.id
            await org._auditer(db, actor, request, "org.entite.confirm", "org_entite", n.entite.id, before=avant,
                               after={"a_confirmer": False, "source_confirmation": SOURCE})
            confirmees.append(n.entite)
        elif n.statut == "EXISTANT" and n.entite is not None and n.sigle and not n.entite.sigle:
            n.entite.sigle = n.sigle
    await org._flush(db)
    arbre = await org.charger_arbre(db)
    for e in crees:
        await org._auditer(db, actor, request, "org.entite.create", "org_entite", e.id,
                           after={**org.entite_dict(arbre.entites[e.id], arbre), "source": SOURCE})
    await org._auditer(db, actor, request, "org.import.apply", "org_entite", None, after={
        "empreinte": empreinte, "lignes": len(lignes), "creees": len(crees), "confirmees": len(confirmees)})
    return {"empreinte": empreinte, "creees": [org.entite_resume(arbre.entites[e.id], arbre) for e in crees],
            "confirmees": [org.entite_resume(arbre.entites[e.id], arbre) for e in confirmees],
            "absentes_du_fichier": synthese["absentes_du_fichier"]}

"""Centre de reporting Achats & Appro — données métier réelles (pas de mock)."""

from __future__ import annotations

import csv
import io
import uuid
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from typing import Any, Callable, Sequence
from zoneinfo import ZoneInfo

from fastapi import HTTPException, status
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import Select, func, or_, select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.data.el_amana_referentiel import BANQUE_EL_AMANA
from app.models.auth import Agence, User
from app.models.mg_achats import (
    MgAchatBl,
    MgAchatComparaison,
    MgAchatConsultation,
    MgAchatDemande,
    MgAchatDevis,
    MgAchatFacture,
    MgAchatPaiement,
    MgAchatReception,
)
from app.models.mg_ops import MgBonCommande
from app.models.organisation import Fournisseur
from app.services.reporting_export import (
    build_styled_pdf,
    build_styled_workbook,
    export_now,
    format_export_datetime,
    format_period_label,
    resolve_bea_logo_path,
)

_TZ = ZoneInfo("Africa/Nouakchott")
EXPORT_MAX_ROWS = 20_000
DEPT_LABEL = "Département Moyens Généraux"


@dataclass(frozen=True)
class ReportColumn:
    key: str
    label: str
    kind: str = "text"  # text | date | number | money | bool


@dataclass
class ReportDef:
    key: str
    label: str
    description: str
    icon: str
    columns: list[ReportColumn]
    filters: list[str]
    csv_enabled: bool = True
    group: str = "objets"  # objets | pilotage | analyses


REPORTS: dict[str, ReportDef] = {
    "demandes": ReportDef(
        key="demandes",
        label="Demandes d'achat",
        description="Demandes internes, priorités et circuit de validation.",
        icon="assignment",
        columns=[
            ReportColumn("id", "ID"),
            ReportColumn("reference", "Référence"),
            ReportColumn("date_demande", "Date", "date"),
            ReportColumn("statut", "Statut"),
            ReportColumn("priorite", "Priorité"),
            ReportColumn("type_achat", "Type"),
            ReportColumn("demandeur_nom", "Demandeur"),
            ReportColumn("agence", "Agence"),
            ReportColumn("projet", "Projet"),
            ReportColumn("motif", "Motif"),
            ReportColumn("montant_estime", "Montant estimé", "money"),
        ],
        filters=["q", "statut", "agence_id", "date_debut", "date_fin", "priorite", "type_achat"],
    ),
    "fournisseurs": ReportDef(
        key="fournisseurs",
        label="Fournisseurs",
        description="Référentiel fournisseurs actifs et inactifs.",
        icon="storefront",
        columns=[
            ReportColumn("id", "ID"),
            ReportColumn("code", "Code"),
            ReportColumn("raison_sociale", "Raison sociale"),
            ReportColumn("type_fournisseur", "Type"),
            ReportColumn("ville", "Ville"),
            ReportColumn("pays", "Pays"),
            ReportColumn("telephone", "Téléphone"),
            ReportColumn("email", "E-mail"),
            ReportColumn("nif", "NIF"),
            ReportColumn("statut", "Statut"),
            ReportColumn("nb_devis", "Nb devis", "number"),
            ReportColumn("nb_bons", "Nb BC", "number"),
            ReportColumn("nb_factures", "Nb factures", "number"),
        ],
        filters=["q", "statut", "type_fournisseur", "ville", "pays"],
    ),
    "consultations": ReportDef(
        key="consultations",
        label="Consultations",
        description="Consultations fournisseurs et dates limites.",
        icon="forum",
        columns=[
            ReportColumn("id", "ID"),
            ReportColumn("reference", "Référence"),
            ReportColumn("date_consultation", "Date", "date"),
            ReportColumn("date_limite", "Limite", "date"),
            ReportColumn("statut", "Statut"),
            ReportColumn("objet", "Objet"),
            ReportColumn("agence", "Agence"),
            ReportColumn("nb_fournisseurs", "Fournisseurs", "number"),
        ],
        filters=["q", "statut", "agence_id", "date_debut", "date_fin"],
    ),
    "devis": ReportDef(
        key="devis",
        label="Devis",
        description="Offres fournisseurs reçues et validité.",
        icon="request_quote",
        columns=[
            ReportColumn("id", "ID"),
            ReportColumn("reference", "Référence"),
            ReportColumn("date_devis", "Date", "date"),
            ReportColumn("date_validite", "Validité", "date"),
            ReportColumn("statut", "Statut"),
            ReportColumn("fournisseur", "Fournisseur"),
            ReportColumn("consultation_ref", "Consultation"),
            ReportColumn("total_ht", "Total HT", "money"),
            ReportColumn("total_ttc", "Total TTC", "money"),
            ReportColumn("devise", "Devise"),
        ],
        filters=["q", "statut", "fournisseur_id", "consultation_id", "date_debut", "date_fin"],
    ),
    "comparaisons": ReportDef(
        key="comparaisons",
        label="Comparaisons",
        description="Tableaux comparatifs et fournisseurs retenus.",
        icon="compare_arrows",
        columns=[
            ReportColumn("id", "ID"),
            ReportColumn("reference", "Référence"),
            ReportColumn("statut", "Statut"),
            ReportColumn("consultation_ref", "Consultation"),
            ReportColumn("fournisseur_retenu", "Retenu"),
            ReportColumn("motif_choix", "Motif"),
        ],
        filters=["q", "statut", "consultation_id"],
    ),
    "bons": ReportDef(
        key="bons",
        label="Bons de commande",
        description="BC, montants, visas et livraisons prévues.",
        icon="receipt_long",
        columns=[
            ReportColumn("id", "ID"),
            ReportColumn("reference", "Référence"),
            ReportColumn("date_bc", "Date BC", "date"),
            ReportColumn("statut", "Statut"),
            ReportColumn("fournisseur", "Fournisseur"),
            ReportColumn("departement", "Département"),
            ReportColumn("projet", "Projet"),
            ReportColumn("total_ht", "Total HT", "money"),
            ReportColumn("total_tva", "Total TVA", "money"),
            ReportColumn("total_ttc", "Total TTC", "money"),
            ReportColumn("date_livraison_prevue", "Livraison prévue", "date"),
        ],
        filters=[
            "q",
            "statut",
            "fournisseur_id",
            "agence_id",
            "date_debut",
            "date_fin",
            "departement",
        ],
    ),
    "livraisons": ReportDef(
        key="livraisons",
        label="Livraisons / BL",
        description="Bons de livraison liés aux commandes.",
        icon="local_shipping",
        columns=[
            ReportColumn("id", "ID"),
            ReportColumn("reference", "Référence"),
            ReportColumn("bon_reference", "BC"),
            ReportColumn("date_bl", "Date BL", "date"),
            ReportColumn("date_livraison", "Livraison", "date"),
            ReportColumn("statut", "Statut"),
            ReportColumn("fournisseur", "Fournisseur"),
            ReportColumn("agence", "Agence"),
            ReportColumn("transporteur", "Transporteur"),
        ],
        filters=["q", "statut", "fournisseur_id", "agence_id", "bon_id", "date_debut", "date_fin"],
    ),
    "receptions": ReportDef(
        key="receptions",
        label="Réceptions",
        description="Réceptions quantitatives sur BC.",
        icon="inventory",
        columns=[
            ReportColumn("id", "ID"),
            ReportColumn("reference", "Référence"),
            ReportColumn("bon_reference", "BC"),
            ReportColumn("date_reception", "Date", "date"),
            ReportColumn("statut", "Statut"),
            ReportColumn("agence", "Agence"),
            ReportColumn("observation", "Observation"),
        ],
        filters=["q", "statut", "agence_id", "bon_id", "date_debut", "date_fin"],
    ),
    "factures": ReportDef(
        key="factures",
        label="Factures",
        description="Factures fournisseurs et contrôle 3 voies.",
        icon="receipt",
        columns=[
            ReportColumn("id", "ID"),
            ReportColumn("reference", "Référence"),
            ReportColumn("numero_fournisseur", "N° fournisseur"),
            ReportColumn("date_facture", "Date", "date"),
            ReportColumn("date_echeance", "Échéance", "date"),
            ReportColumn("statut", "Statut"),
            ReportColumn("fournisseur", "Fournisseur"),
            ReportColumn("bon_reference", "BC"),
            ReportColumn("montant_ht", "HT", "money"),
            ReportColumn("montant_tva", "TVA", "money"),
            ReportColumn("montant_ttc", "TTC", "money"),
            ReportColumn("ecart_quantite", "Écart qty", "bool"),
            ReportColumn("ecart_montant", "Écart mt", "bool"),
        ],
        filters=["q", "statut", "fournisseur_id", "bon_id", "date_debut", "date_fin"],
    ),
    "paiements": ReportDef(
        key="paiements",
        label="Paiements",
        description="Suivi des paiements à régler et payés.",
        icon="payments",
        columns=[
            ReportColumn("id", "ID"),
            ReportColumn("reference", "Référence"),
            ReportColumn("facture_ref", "Facture"),
            ReportColumn("statut", "Statut"),
            ReportColumn("montant", "Montant", "money"),
            ReportColumn("date_echeance", "Échéance", "date"),
            ReportColumn("date_paiement", "Date paiement", "date"),
            ReportColumn("mode_paiement", "Mode"),
            ReportColumn("fournisseur", "Fournisseur"),
        ],
        filters=["q", "statut", "fournisseur_id", "date_debut", "date_fin"],
    ),
    "echeances": ReportDef(
        key="echeances",
        label="Échéances & Anomalies",
        description="Alertes d'échéances, livraisons et écarts 3 voies.",
        icon="notification_important",
        columns=[
            ReportColumn("id", "ID"),
            ReportColumn("type", "Type"),
            ReportColumn("reference", "Référence"),
            ReportColumn("message", "Message"),
            ReportColumn("priorite", "Priorité"),
            ReportColumn("date_echeance", "Échéance", "date"),
        ],
        filters=["q", "priorite", "type"],
        csv_enabled=True,
        group="pilotage",
    ),
}


def catalog() -> list[dict[str, Any]]:
    items = []
    for r in REPORTS.values():
        items.append(
            {
                "key": r.key,
                "label": r.label,
                "description": r.description,
                "icon": r.icon,
                "group": r.group,
                "csv_enabled": r.csv_enabled,
                "filters": r.filters,
                "columns": [
                    {"key": c.key, "label": c.label, "kind": c.kind} for c in r.columns if c.key != "id"
                ],
            }
        )
    items.extend(
        [
            {
                "key": "analyses",
                "label": "Analyses",
                "description": "Volumes, montants et répartition par statut / fournisseur.",
                "icon": "insights",
                "group": "analyses",
                "csv_enabled": True,
                "filters": ["date_debut", "date_fin", "agence_id"],
                "columns": [],
            },
            {
                "key": "personnalise",
                "label": "Rapports personnalisés",
                "description": "Choisir type, colonnes, filtres, tri puis exporter.",
                "icon": "tune",
                "group": "analyses",
                "csv_enabled": True,
                "filters": [],
                "columns": [],
            },
        ]
    )
    return items


def _fmt_date(v: date | datetime | None) -> str | None:
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.astimezone(_TZ).strftime("%d/%m/%Y")
    return v.strftime("%d/%m/%Y")


def _money(v: Any) -> float | None:
    if v is None:
        return None
    return float(Decimal(v))


def _bool_label(v: Any) -> str:
    return "Oui" if v else "Non"


class MgAchatsReportingService:
    def __init__(self, db: AsyncSession):
        self.db = db
        self._agence_cache: dict[uuid.UUID, str] = {}
        self._frs_cache: dict[uuid.UUID, str] = {}

    async def _agence_label(self, agence_id: uuid.UUID | None) -> str | None:
        if not agence_id:
            return None
        if agence_id in self._agence_cache:
            return self._agence_cache[agence_id]
        ag = await self.db.get(Agence, agence_id)
        label = ag.libelle if ag else None
        if label:
            self._agence_cache[agence_id] = label
        return label

    async def _frs_label(self, frs_id: uuid.UUID | None) -> str | None:
        if not frs_id:
            return None
        if frs_id in self._frs_cache:
            return self._frs_cache[frs_id]
        fr = await self.db.get(Fournisseur, frs_id)
        label = fr.raison_sociale if fr else None
        if label:
            self._frs_cache[frs_id] = label
        return label

    def _parse_filters(self, raw: dict[str, Any] | None) -> dict[str, Any]:
        f = dict(raw or {})
        out: dict[str, Any] = {}
        for key in (
            "q",
            "statut",
            "priorite",
            "type_achat",
            "type_fournisseur",
            "ville",
            "pays",
            "departement",
            "type",
            "sort_by",
            "sort_dir",
        ):
            val = f.get(key)
            if val is not None and str(val).strip() != "":
                out[key] = str(val).strip()
        for key in ("agence_id", "fournisseur_id", "consultation_id", "bon_id", "facture_id"):
            val = f.get(key)
            if val:
                try:
                    out[key] = uuid.UUID(str(val))
                except ValueError:
                    raise HTTPException(status.HTTP_400_BAD_REQUEST, detail=f"{key} invalide")
        for key in ("date_debut", "date_fin"):
            val = f.get(key)
            if val:
                if isinstance(val, date) and not isinstance(val, datetime):
                    out[key] = val
                else:
                    try:
                        out[key] = date.fromisoformat(str(val)[:10])
                    except ValueError:
                        raise HTTPException(status.HTTP_400_BAD_REQUEST, detail=f"{key} invalide")
        return out

    def _filters_label(self, filters: dict[str, Any]) -> str:
        parts: list[str] = [DEPT_LABEL]
        period = format_period_label(filters.get("date_debut"), filters.get("date_fin"))
        if period:
            parts.append(period)
        mapping = {
            "statut": "Statut",
            "priorite": "Priorité",
            "type_achat": "Type",
            "type_fournisseur": "Type frs",
            "ville": "Ville",
            "pays": "Pays",
            "departement": "Département",
            "type": "Type alerte",
            "q": "Recherche",
        }
        for k, label in mapping.items():
            if k in filters:
                parts.append(f"{label}={filters[k]}")
        if "agence_id" in filters:
            parts.append(f"Agence={filters['agence_id']}")
        if "fournisseur_id" in filters:
            parts.append(f"Fournisseur={filters['fournisseur_id']}")
        return " · ".join(parts)

    async def preview(
        self,
        report_key: str,
        *,
        filters: dict[str, Any] | None = None,
        page: int = 1,
        size: int = 50,
        columns: list[str] | None = None,
    ) -> dict[str, Any]:
        if report_key == "analyses":
            return await self.analyses(filters=filters)
        if report_key == "personnalise":
            raise HTTPException(
                status.HTTP_400_BAD_REQUEST,
                detail="Utilisez /rapports/personnalise/preview",
            )
        rdef = REPORTS.get(report_key)
        if not rdef:
            raise HTTPException(status.HTTP_404_NOT_FOUND, detail="Rapport inconnu")
        f = self._parse_filters(filters)
        rows, total, totals = await self._fetch_rows(report_key, f, ids=None, limit=None)
        # sort
        sort_by = f.get("sort_by")
        sort_dir = (f.get("sort_dir") or "desc").lower()
        if sort_by:
            reverse = sort_dir != "asc"
            rows.sort(key=lambda x: (x.get(sort_by) is None, x.get(sort_by)), reverse=reverse)
        page, size = max(1, page), max(1, min(size, 200))
        start = (page - 1) * size
        page_rows = rows[start : start + size]
        col_keys = self._resolve_columns(rdef, columns)
        return {
            "report_key": report_key,
            "label": rdef.label,
            "total": total,
            "page": page,
            "size": size,
            "columns": [
                {"key": c.key, "label": c.label, "kind": c.kind}
                for c in rdef.columns
                if c.key in col_keys and c.key != "id"
            ],
            "rows": [{k: row.get(k) for k in col_keys} for row in page_rows],
            "totals": totals,
            "filters_label": self._filters_label(f),
            "empty_message": "Aucune donnée ne correspond aux critères sélectionnés.",
        }

    async def analyses(self, *, filters: dict[str, Any] | None = None) -> dict[str, Any]:
        f = self._parse_filters(filters)
        date_debut = f.get("date_debut")
        date_fin = f.get("date_fin")

        async def _count(model, date_col, extra=None):
            filters_ = [model.deleted_at.is_(None)]
            if date_debut is not None:
                filters_.append(date_col >= date_debut)
            if date_fin is not None:
                filters_.append(date_col <= date_fin)
            if extra:
                filters_.extend(extra)
            return int(await self.db.scalar(select(func.count()).select_from(model).where(*filters_)) or 0)

        async def _sum(model, amount_col, date_col):
            filters_ = [model.deleted_at.is_(None)]
            if date_debut is not None:
                filters_.append(date_col >= date_debut)
            if date_fin is not None:
                filters_.append(date_col <= date_fin)
            return float(
                await self.db.scalar(
                    select(func.coalesce(func.sum(amount_col), 0)).where(*filters_)
                )
                or 0
            )

        by_statut_bc = (
            await self.db.execute(
                select(MgBonCommande.statut, func.count())
                .where(MgBonCommande.deleted_at.is_(None))
                .group_by(MgBonCommande.statut)
            )
        ).all()

        return {
            "report_key": "analyses",
            "label": "Analyses Achats",
            "filters_label": self._filters_label(f),
            "kpis": {
                "nb_demandes": await _count(MgAchatDemande, MgAchatDemande.date_demande),
                "nb_consultations": await _count(
                    MgAchatConsultation, MgAchatConsultation.date_consultation
                ),
                "nb_devis": await _count(MgAchatDevis, MgAchatDevis.date_devis),
                "nb_bons": await _count(MgBonCommande, MgBonCommande.date_bc),
                "nb_receptions": await _count(MgAchatReception, MgAchatReception.date_reception),
                "nb_factures": await _count(MgAchatFacture, MgAchatFacture.date_facture),
                "montant_bc": await _sum(MgBonCommande, MgBonCommande.total_ttc, MgBonCommande.date_bc),
                "montant_factures": await _sum(
                    MgAchatFacture, MgAchatFacture.montant_ttc, MgAchatFacture.date_facture
                ),
                "montant_paiements": await _sum(
                    MgAchatPaiement, MgAchatPaiement.montant, MgAchatPaiement.date_echeance
                ),
            },
            "bons_par_statut": [{"statut": s, "count": int(n)} for s, n in by_statut_bc],
            "empty_message": "Aucune donnée ne correspond aux critères sélectionnés.",
        }

    def _resolve_columns(self, rdef: ReportDef, columns: list[str] | None) -> list[str]:
        allowed = [c.key for c in rdef.columns]
        if not columns:
            return [k for k in allowed if k != "id"] + (["id"] if "id" in allowed else [])
        # always keep id for selection
        picked = [c for c in columns if c in allowed and c != "id"]
        if not picked:
            picked = [k for k in allowed if k != "id"]
        return ["id"] + picked if "id" in allowed else picked

    async def export(
        self,
        report_key: str,
        *,
        fmt: str,
        scope: str,
        ids: list[uuid.UUID] | None,
        filters: dict[str, Any] | None,
        columns: list[str] | None,
        user: User,
    ) -> tuple[bytes, str, str, int]:
        """Returns content, media_type, filename, row_count."""
        if report_key == "analyses":
            return await self._export_analyses(fmt=fmt, filters=filters, user=user)
        rdef = REPORTS.get(report_key)
        if not rdef:
            raise HTTPException(status.HTTP_404_NOT_FOUND, detail="Rapport inconnu")
        if fmt == "csv" and not rdef.csv_enabled:
            raise HTTPException(status.HTTP_400_BAD_REQUEST, detail="CSV non disponible pour ce rapport")
        f = self._parse_filters(filters)
        id_set = ids if scope == "selection" else None
        if scope == "selection" and not id_set:
            raise HTTPException(
                status.HTTP_400_BAD_REQUEST,
                detail="Aucune ligne sélectionnée pour l'export",
            )
        rows, total, totals = await self._fetch_rows(
            report_key, f, ids=id_set, limit=EXPORT_MAX_ROWS
        )
        if scope == "selection" and id_set:
            id_str = {str(i) for i in id_set}
            rows = [r for r in rows if str(r.get("id")) in id_str]
            total = len(rows)
        if total == 0:
            raise HTTPException(
                status.HTTP_404_NOT_FOUND,
                detail="Aucune donnée ne correspond aux critères sélectionnés.",
            )
        if total > EXPORT_MAX_ROWS:
            rows = rows[:EXPORT_MAX_ROWS]
        sort_by = f.get("sort_by")
        if sort_by:
            reverse = (f.get("sort_dir") or "desc").lower() != "asc"
            rows.sort(key=lambda x: (x.get(sort_by) is None, x.get(sort_by)), reverse=reverse)

        col_keys = [k for k in self._resolve_columns(rdef, columns) if k != "id"]
        headers = [next(c.label for c in rdef.columns if c.key == k) for k in col_keys]
        kinds = {c.key: c.kind for c in rdef.columns}
        data_rows: list[list[Any]] = []
        for row in rows:
            line = []
            for k in col_keys:
                v = row.get(k)
                kind = kinds.get(k, "text")
                if kind == "bool":
                    line.append(_bool_label(v))
                elif kind == "date":
                    line.append(v if isinstance(v, str) else _fmt_date(v) or "")
                elif kind in ("money", "number") and v is not None:
                    line.append(float(v))
                else:
                    line.append("" if v is None else v)
            data_rows.append(line)

        when = export_now()
        user_label = user.full_name or user.email
        subtitle = (
            f"{self._filters_label(f)} · Généré par {user_label} · "
            f"{format_export_datetime(when)} · Scope={scope}"
        )
        title = f"BEA-DIGITAL — {rdef.label}"
        stamp = when.strftime("%Y%m%d-%H%M")
        base_name = f"achats-{report_key}-{stamp}"

        if fmt == "csv":
            buf = io.StringIO()
            writer = csv.writer(buf, delimiter=";")
            writer.writerow(headers)
            for line in data_rows:
                writer.writerow(line)
            content = buf.getvalue().encode("utf-8-sig")
            return content, "text/csv; charset=utf-8", f"{base_name}.csv", len(data_rows)

        if fmt == "pdf":
            content = build_styled_pdf(
                report_title=title,
                headers=headers,
                rows=data_rows,
                subtitle=subtitle,
                exported_at=when,
                landscape_mode=True,
            )
            return content, "application/pdf", f"{base_name}.pdf", len(data_rows)

        # xlsx multi-feuilles
        content = self._build_xlsx(
            report_title=title,
            sheet_title=rdef.label[:31],
            headers=headers,
            rows=data_rows,
            subtitle=subtitle,
            totals=totals,
            col_keys=col_keys,
            kinds=kinds,
            exported_at=when,
        )
        return (
            content,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            f"{base_name}.xlsx",
            len(data_rows),
        )

    async def _export_analyses(
        self, *, fmt: str, filters: dict[str, Any] | None, user: User
    ) -> tuple[bytes, str, str, int]:
        data = await self.analyses(filters=filters)
        kpis = data["kpis"]
        headers = ["Indicateur", "Valeur"]
        rows = [[k, v] for k, v in kpis.items()]
        for item in data.get("bons_par_statut") or []:
            rows.append([f"BC statut {item['statut']}", item["count"]])
        if not rows:
            raise HTTPException(
                status.HTTP_404_NOT_FOUND,
                detail="Aucune donnée ne correspond aux critères sélectionnés.",
            )
        when = export_now()
        title = "BEA-DIGITAL — Analyses Achats"
        subtitle = f"{data['filters_label']} · Généré par {user.full_name or user.email}"
        stamp = when.strftime("%Y%m%d-%H%M")
        if fmt == "csv":
            buf = io.StringIO()
            w = csv.writer(buf, delimiter=";")
            w.writerow(headers)
            w.writerows(rows)
            return buf.getvalue().encode("utf-8-sig"), "text/csv; charset=utf-8", f"achats-analyses-{stamp}.csv", len(rows)
        if fmt == "pdf":
            content = build_styled_pdf(
                report_title=title, headers=headers, rows=rows, subtitle=subtitle, exported_at=when
            )
            return content, "application/pdf", f"achats-analyses-{stamp}.pdf", len(rows)
        content = build_styled_workbook(
            sheet_title="Analyses",
            report_title=title,
            headers=headers,
            rows=rows,
            subtitle=subtitle,
            exported_at=when,
        )
        return (
            content,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            f"achats-analyses-{stamp}.xlsx",
            len(rows),
        )

    def _build_xlsx(
        self,
        *,
        report_title: str,
        sheet_title: str,
        headers: Sequence[str],
        rows: Sequence[Sequence[Any]],
        subtitle: str,
        totals: dict[str, Any],
        col_keys: list[str],
        kinds: dict[str, str],
        exported_at: datetime,
    ) -> bytes:
        # Feuille Données via utilitaire existant
        data_bytes = build_styled_workbook(
            sheet_title="Donnees",
            report_title=report_title,
            headers=headers,
            rows=rows,
            subtitle=subtitle,
            exported_at=exported_at,
        )
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data_bytes))
        # Feuille Synthèse
        ws = wb.create_sheet("Synthese", 0)
        ws["A1"] = BANQUE_EL_AMANA["raison_sociale"]
        ws["A1"].font = Font(bold=True, color="1E3A5F", size=12)
        ws["A2"] = report_title
        ws["A2"].font = Font(bold=True, size=14)
        ws["A3"] = subtitle
        ws["A3"].font = Font(italic=True, color="64748B", size=10)
        ws["A5"] = "Indicateur"
        ws["B5"] = "Valeur"
        for cell in (ws["A5"], ws["B5"]):
            cell.fill = PatternFill("solid", fgColor="1E3A5F")
            cell.font = Font(bold=True, color="FFFFFF")
        ws["A6"] = "Nombre de lignes"
        ws["B6"] = len(rows)
        row_i = 7
        for k, v in (totals or {}).items():
            ws.cell(row=row_i, column=1, value=str(k))
            ws.cell(row=row_i, column=2, value=float(v) if isinstance(v, (int, float, Decimal)) else v)
            row_i += 1
        ws["A" + str(row_i + 1)] = DEPT_LABEL
        ws.column_dimensions["A"].width = 36
        ws.column_dimensions["B"].width = 18
        # Filtres auto déjà sur Données
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    async def _fetch_rows(
        self,
        report_key: str,
        filters: dict[str, Any],
        *,
        ids: list[uuid.UUID] | None,
        limit: int | None,
    ) -> tuple[list[dict[str, Any]], int, dict[str, Any]]:
        handlers: dict[str, Callable] = {
            "demandes": self._rows_demandes,
            "fournisseurs": self._rows_fournisseurs,
            "consultations": self._rows_consultations,
            "devis": self._rows_devis,
            "comparaisons": self._rows_comparaisons,
            "bons": self._rows_bons,
            "livraisons": self._rows_livraisons,
            "receptions": self._rows_receptions,
            "factures": self._rows_factures,
            "paiements": self._rows_paiements,
            "echeances": self._rows_echeances,
        }
        handler = handlers.get(report_key)
        if not handler:
            raise HTTPException(status.HTTP_404_NOT_FOUND, detail="Rapport inconnu")
        return await handler(filters, ids=ids, limit=limit)

    async def _rows_demandes(self, f, *, ids, limit):
        filters = [MgAchatDemande.deleted_at.is_(None)]
        if ids:
            filters.append(MgAchatDemande.id.in_(ids))
        if f.get("statut"):
            filters.append(MgAchatDemande.statut == f["statut"])
        if f.get("priorite"):
            filters.append(MgAchatDemande.priorite == f["priorite"])
        if f.get("type_achat"):
            filters.append(MgAchatDemande.type_achat == f["type_achat"])
        if f.get("agence_id"):
            filters.append(MgAchatDemande.agence_id == f["agence_id"])
        if f.get("date_debut"):
            filters.append(MgAchatDemande.date_demande >= f["date_debut"])
        if f.get("date_fin"):
            filters.append(MgAchatDemande.date_demande <= f["date_fin"])
        if f.get("q"):
            like = f"%{f['q']}%"
            filters.append(
                or_(
                    MgAchatDemande.reference.ilike(like),
                    MgAchatDemande.demandeur_nom.ilike(like),
                    MgAchatDemande.projet.ilike(like),
                    MgAchatDemande.motif.ilike(like),
                )
            )
        total = int(
            await self.db.scalar(select(func.count()).select_from(MgAchatDemande).where(*filters))
            or 0
        )
        stmt = (
            select(MgAchatDemande)
            .options(selectinload(MgAchatDemande.lignes))
            .where(*filters)
            .order_by(MgAchatDemande.date_demande.desc())
        )
        if limit:
            stmt = stmt.limit(limit)
        rows = list((await self.db.execute(stmt)).scalars().all())
        out = []
        montant_total = Decimal("0")
        for r in rows:
            mt = sum((ln.montant_estime or Decimal("0")) for ln in (r.lignes or []))
            montant_total += mt
            out.append(
                {
                    "id": str(r.id),
                    "reference": r.reference,
                    "date_demande": _fmt_date(r.date_demande),
                    "statut": r.statut,
                    "priorite": r.priorite,
                    "type_achat": r.type_achat,
                    "demandeur_nom": r.demandeur_nom,
                    "agence": await self._agence_label(r.agence_id),
                    "projet": r.projet,
                    "motif": r.motif,
                    "montant_estime": _money(mt),
                }
            )
        return out, total, {"montant_estime": float(montant_total)}

    async def _rows_fournisseurs(self, f, *, ids, limit):
        from app.services.mg_achats_service import MgAchatsService

        svc = MgAchatsService(self.db)
        # Réutilise le listing métier + stats
        items, total = await svc.list_fournisseurs(
            q=f.get("q"),
            statut=f.get("statut"),
            type_fournisseur=f.get("type_fournisseur"),
            ville=f.get("ville"),
            pays=f.get("pays"),
            actifs_seulement=f.get("statut") != "INACTIF",
            page=1,
            size=limit or EXPORT_MAX_ROWS,
        )
        if ids:
            id_str = {str(i) for i in ids}
            items = [x for x in items if str(x.get("id")) in id_str]
            total = len(items)
        out = []
        for x in items:
            out.append(
                {
                    "id": str(x["id"]),
                    "code": x.get("code"),
                    "raison_sociale": x.get("raison_sociale"),
                    "type_fournisseur": x.get("type_fournisseur"),
                    "ville": x.get("ville"),
                    "pays": x.get("pays"),
                    "telephone": x.get("telephone"),
                    "email": x.get("email"),
                    "nif": x.get("nif"),
                    "statut": "ACTIF" if x.get("is_active", True) else "INACTIF",
                    "nb_devis": x.get("nb_devis") or 0,
                    "nb_bons": x.get("nb_bons") or 0,
                    "nb_factures": x.get("nb_factures") or 0,
                }
            )
        return out, total, {}

    async def _rows_consultations(self, f, *, ids, limit):
        filters = [MgAchatConsultation.deleted_at.is_(None)]
        if ids:
            filters.append(MgAchatConsultation.id.in_(ids))
        if f.get("statut"):
            filters.append(MgAchatConsultation.statut == f["statut"])
        if f.get("agence_id"):
            filters.append(MgAchatConsultation.agence_id == f["agence_id"])
        if f.get("date_debut"):
            filters.append(MgAchatConsultation.date_consultation >= f["date_debut"])
        if f.get("date_fin"):
            filters.append(MgAchatConsultation.date_consultation <= f["date_fin"])
        if f.get("q"):
            like = f"%{f['q']}%"
            filters.append(
                or_(
                    MgAchatConsultation.reference.ilike(like),
                    MgAchatConsultation.objet.ilike(like),
                )
            )
        total = int(
            await self.db.scalar(
                select(func.count()).select_from(MgAchatConsultation).where(*filters)
            )
            or 0
        )
        stmt = (
            select(MgAchatConsultation)
            .options(selectinload(MgAchatConsultation.fournisseurs))
            .where(*filters)
            .order_by(MgAchatConsultation.date_consultation.desc())
        )
        if limit:
            stmt = stmt.limit(limit)
        rows = list((await self.db.execute(stmt)).scalars().all())
        out = []
        for r in rows:
            out.append(
                {
                    "id": str(r.id),
                    "reference": r.reference,
                    "date_consultation": _fmt_date(r.date_consultation),
                    "date_limite": _fmt_date(r.date_limite),
                    "statut": r.statut,
                    "objet": r.objet,
                    "agence": await self._agence_label(r.agence_id),
                    "nb_fournisseurs": len(r.fournisseurs or []),
                }
            )
        return out, total, {}

    async def _rows_devis(self, f, *, ids, limit):
        filters = [MgAchatDevis.deleted_at.is_(None)]
        if ids:
            filters.append(MgAchatDevis.id.in_(ids))
        if f.get("statut"):
            filters.append(MgAchatDevis.statut == f["statut"])
        if f.get("fournisseur_id"):
            filters.append(MgAchatDevis.fournisseur_id == f["fournisseur_id"])
        if f.get("consultation_id"):
            filters.append(MgAchatDevis.consultation_id == f["consultation_id"])
        if f.get("date_debut"):
            filters.append(MgAchatDevis.date_devis >= f["date_debut"])
        if f.get("date_fin"):
            filters.append(MgAchatDevis.date_devis <= f["date_fin"])
        if f.get("q"):
            filters.append(MgAchatDevis.reference.ilike(f"%{f['q']}%"))
        total = int(
            await self.db.scalar(select(func.count()).select_from(MgAchatDevis).where(*filters))
            or 0
        )
        stmt = select(MgAchatDevis).where(*filters).order_by(MgAchatDevis.date_devis.desc())
        if limit:
            stmt = stmt.limit(limit)
        rows = list((await self.db.execute(stmt)).scalars().all())
        cons_ids = {r.consultation_id for r in rows if r.consultation_id}
        cons_map: dict[uuid.UUID, str] = {}
        if cons_ids:
            for cid, ref in (
                await self.db.execute(
                    select(MgAchatConsultation.id, MgAchatConsultation.reference).where(
                        MgAchatConsultation.id.in_(cons_ids)
                    )
                )
            ).all():
                cons_map[cid] = ref
        out = []
        tot_ht = Decimal("0")
        tot_ttc = Decimal("0")
        for r in rows:
            tot_ht += r.montant_ht or Decimal("0")
            tot_ttc += r.montant_ttc or Decimal("0")
            out.append(
                {
                    "id": str(r.id),
                    "reference": r.reference,
                    "date_devis": _fmt_date(r.date_devis),
                    "date_validite": _fmt_date(r.date_validite),
                    "statut": r.statut,
                    "fournisseur": await self._frs_label(r.fournisseur_id),
                    "consultation_ref": cons_map.get(r.consultation_id) if r.consultation_id else None,
                    "total_ht": _money(r.montant_ht),
                    "total_ttc": _money(r.montant_ttc),
                    "devise": r.devise,
                }
            )
        return out, total, {"total_ht": float(tot_ht), "total_ttc": float(tot_ttc)}

    async def _rows_comparaisons(self, f, *, ids, limit):
        filters = [MgAchatComparaison.deleted_at.is_(None)]
        if ids:
            filters.append(MgAchatComparaison.id.in_(ids))
        if f.get("statut"):
            filters.append(MgAchatComparaison.statut == f["statut"])
        if f.get("consultation_id"):
            filters.append(MgAchatComparaison.consultation_id == f["consultation_id"])
        if f.get("q"):
            filters.append(MgAchatComparaison.reference.ilike(f"%{f['q']}%"))
        total = int(
            await self.db.scalar(
                select(func.count()).select_from(MgAchatComparaison).where(*filters)
            )
            or 0
        )
        stmt = (
            select(MgAchatComparaison)
            .where(*filters)
            .order_by(MgAchatComparaison.created_at.desc())
        )
        if limit:
            stmt = stmt.limit(limit)
        rows = list((await self.db.execute(stmt)).scalars().all())
        cons_ids = {r.consultation_id for r in rows}
        cons_map = {}
        if cons_ids:
            for cid, ref in (
                await self.db.execute(
                    select(MgAchatConsultation.id, MgAchatConsultation.reference).where(
                        MgAchatConsultation.id.in_(cons_ids)
                    )
                )
            ).all():
                cons_map[cid] = ref
        out = []
        for r in rows:
            out.append(
                {
                    "id": str(r.id),
                    "reference": r.reference,
                    "statut": r.statut,
                    "consultation_ref": cons_map.get(r.consultation_id),
                    "fournisseur_retenu": await self._frs_label(r.fournisseur_retenu_id),
                    "motif_choix": r.motif_choix,
                }
            )
        return out, total, {}

    async def _rows_bons(self, f, *, ids, limit):
        filters = [MgBonCommande.deleted_at.is_(None)]
        if ids:
            filters.append(MgBonCommande.id.in_(ids))
        if f.get("statut"):
            filters.append(MgBonCommande.statut == f["statut"])
        if f.get("fournisseur_id"):
            filters.append(MgBonCommande.fournisseur_id == f["fournisseur_id"])
        if f.get("agence_id"):
            filters.append(
                or_(
                    MgBonCommande.agence_facturation_id == f["agence_id"],
                    MgBonCommande.agence_livraison_id == f["agence_id"],
                )
            )
        if f.get("departement"):
            filters.append(MgBonCommande.departement.ilike(f"%{f['departement']}%"))
        if f.get("date_debut"):
            filters.append(MgBonCommande.date_bc >= f["date_debut"])
        if f.get("date_fin"):
            filters.append(MgBonCommande.date_bc <= f["date_fin"])
        if f.get("q"):
            like = f"%{f['q']}%"
            filters.append(
                or_(
                    MgBonCommande.reference.ilike(like),
                    MgBonCommande.fournisseur_raison_sociale.ilike(like),
                    MgBonCommande.projet.ilike(like),
                )
            )
        total = int(
            await self.db.scalar(select(func.count()).select_from(MgBonCommande).where(*filters))
            or 0
        )
        stmt = select(MgBonCommande).where(*filters).order_by(MgBonCommande.date_bc.desc())
        if limit:
            stmt = stmt.limit(limit)
        rows = list((await self.db.execute(stmt)).scalars().all())
        out = []
        tot_ht = tot_tva = tot_ttc = Decimal("0")
        for r in rows:
            tot_ht += r.total_ht or Decimal("0")
            tot_tva += r.total_tva or Decimal("0")
            tot_ttc += r.total_ttc or Decimal("0")
            out.append(
                {
                    "id": str(r.id),
                    "reference": r.reference,
                    "date_bc": _fmt_date(r.date_bc),
                    "statut": r.statut,
                    "fournisseur": r.fournisseur_raison_sociale
                    or await self._frs_label(r.fournisseur_id),
                    "departement": r.departement,
                    "projet": r.projet,
                    "total_ht": _money(r.total_ht),
                    "total_tva": _money(r.total_tva),
                    "total_ttc": _money(r.total_ttc),
                    "date_livraison_prevue": _fmt_date(r.date_livraison_prevue),
                }
            )
        return out, total, {
            "total_ht": float(tot_ht),
            "total_tva": float(tot_tva),
            "total_ttc": float(tot_ttc),
        }

    async def _rows_livraisons(self, f, *, ids, limit):
        filters = [MgAchatBl.deleted_at.is_(None)]
        if ids:
            filters.append(MgAchatBl.id.in_(ids))
        if f.get("statut"):
            filters.append(MgAchatBl.statut == f["statut"])
        if f.get("fournisseur_id"):
            filters.append(MgAchatBl.fournisseur_id == f["fournisseur_id"])
        if f.get("agence_id"):
            filters.append(MgAchatBl.agence_id == f["agence_id"])
        if f.get("bon_id"):
            filters.append(MgAchatBl.bon_id == f["bon_id"])
        if f.get("date_debut"):
            filters.append(MgAchatBl.date_bl >= f["date_debut"])
        if f.get("date_fin"):
            filters.append(MgAchatBl.date_bl <= f["date_fin"])
        if f.get("q"):
            filters.append(MgAchatBl.reference.ilike(f"%{f['q']}%"))
        total = int(
            await self.db.scalar(select(func.count()).select_from(MgAchatBl).where(*filters)) or 0
        )
        stmt = select(MgAchatBl).where(*filters).order_by(MgAchatBl.date_bl.desc())
        if limit:
            stmt = stmt.limit(limit)
        rows = list((await self.db.execute(stmt)).scalars().all())
        bon_ids = {r.bon_id for r in rows}
        bon_map = {}
        if bon_ids:
            for bid, ref in (
                await self.db.execute(
                    select(MgBonCommande.id, MgBonCommande.reference).where(
                        MgBonCommande.id.in_(bon_ids)
                    )
                )
            ).all():
                bon_map[bid] = ref
        out = []
        for r in rows:
            out.append(
                {
                    "id": str(r.id),
                    "reference": r.reference,
                    "bon_reference": bon_map.get(r.bon_id),
                    "date_bl": _fmt_date(r.date_bl),
                    "date_livraison": _fmt_date(r.date_livraison),
                    "statut": r.statut,
                    "fournisseur": await self._frs_label(r.fournisseur_id),
                    "agence": await self._agence_label(r.agence_id),
                    "transporteur": r.transporteur,
                }
            )
        return out, total, {}

    async def _rows_receptions(self, f, *, ids, limit):
        filters = [MgAchatReception.deleted_at.is_(None)]
        if ids:
            filters.append(MgAchatReception.id.in_(ids))
        if f.get("statut"):
            filters.append(MgAchatReception.statut == f["statut"])
        if f.get("agence_id"):
            filters.append(MgAchatReception.agence_id == f["agence_id"])
        if f.get("bon_id"):
            filters.append(MgAchatReception.bon_id == f["bon_id"])
        if f.get("date_debut"):
            filters.append(MgAchatReception.date_reception >= f["date_debut"])
        if f.get("date_fin"):
            filters.append(MgAchatReception.date_reception <= f["date_fin"])
        if f.get("q"):
            filters.append(MgAchatReception.reference.ilike(f"%{f['q']}%"))
        total = int(
            await self.db.scalar(
                select(func.count()).select_from(MgAchatReception).where(*filters)
            )
            or 0
        )
        stmt = (
            select(MgAchatReception)
            .where(*filters)
            .order_by(MgAchatReception.date_reception.desc())
        )
        if limit:
            stmt = stmt.limit(limit)
        rows = list((await self.db.execute(stmt)).scalars().all())
        bon_ids = {r.bon_id for r in rows}
        bon_map = {}
        if bon_ids:
            for bid, ref in (
                await self.db.execute(
                    select(MgBonCommande.id, MgBonCommande.reference).where(
                        MgBonCommande.id.in_(bon_ids)
                    )
                )
            ).all():
                bon_map[bid] = ref
        out = []
        for r in rows:
            out.append(
                {
                    "id": str(r.id),
                    "reference": r.reference,
                    "bon_reference": bon_map.get(r.bon_id),
                    "date_reception": _fmt_date(r.date_reception),
                    "statut": r.statut,
                    "agence": await self._agence_label(r.agence_id),
                    "observation": r.observation,
                }
            )
        return out, total, {}

    async def _rows_factures(self, f, *, ids, limit):
        filters = [MgAchatFacture.deleted_at.is_(None)]
        if ids:
            filters.append(MgAchatFacture.id.in_(ids))
        if f.get("statut"):
            filters.append(MgAchatFacture.statut == f["statut"])
        if f.get("fournisseur_id"):
            filters.append(MgAchatFacture.fournisseur_id == f["fournisseur_id"])
        if f.get("bon_id"):
            filters.append(MgAchatFacture.bon_id == f["bon_id"])
        if f.get("date_debut"):
            filters.append(MgAchatFacture.date_facture >= f["date_debut"])
        if f.get("date_fin"):
            filters.append(MgAchatFacture.date_facture <= f["date_fin"])
        if f.get("q"):
            like = f"%{f['q']}%"
            filters.append(
                or_(
                    MgAchatFacture.reference.ilike(like),
                    MgAchatFacture.numero_fournisseur.ilike(like),
                )
            )
        total = int(
            await self.db.scalar(select(func.count()).select_from(MgAchatFacture).where(*filters))
            or 0
        )
        stmt = select(MgAchatFacture).where(*filters).order_by(MgAchatFacture.date_facture.desc())
        if limit:
            stmt = stmt.limit(limit)
        rows = list((await self.db.execute(stmt)).scalars().all())
        bon_ids = {r.bon_id for r in rows}
        bon_map = {}
        if bon_ids:
            for bid, ref in (
                await self.db.execute(
                    select(MgBonCommande.id, MgBonCommande.reference).where(
                        MgBonCommande.id.in_(bon_ids)
                    )
                )
            ).all():
                bon_map[bid] = ref
        out = []
        tot = Decimal("0")
        for r in rows:
            tot += r.montant_ttc or Decimal("0")
            out.append(
                {
                    "id": str(r.id),
                    "reference": r.reference,
                    "numero_fournisseur": r.numero_fournisseur,
                    "date_facture": _fmt_date(r.date_facture),
                    "date_echeance": _fmt_date(r.date_echeance),
                    "statut": r.statut,
                    "fournisseur": await self._frs_label(r.fournisseur_id),
                    "bon_reference": bon_map.get(r.bon_id),
                    "montant_ht": _money(r.montant_ht),
                    "montant_tva": _money(r.montant_tva),
                    "montant_ttc": _money(r.montant_ttc),
                    "ecart_quantite": bool(r.ecart_quantite),
                    "ecart_montant": bool(r.ecart_montant),
                }
            )
        return out, total, {"montant_ttc": float(tot)}

    async def _rows_paiements(self, f, *, ids, limit):
        filters = [MgAchatPaiement.deleted_at.is_(None)]
        if ids:
            filters.append(MgAchatPaiement.id.in_(ids))
        if f.get("statut"):
            filters.append(MgAchatPaiement.statut == f["statut"])
        if f.get("fournisseur_id"):
            filters.append(MgAchatPaiement.fournisseur_id == f["fournisseur_id"])
        if f.get("date_debut"):
            filters.append(MgAchatPaiement.date_echeance >= f["date_debut"])
        if f.get("date_fin"):
            filters.append(MgAchatPaiement.date_echeance <= f["date_fin"])
        if f.get("q"):
            filters.append(MgAchatPaiement.reference.ilike(f"%{f['q']}%"))
        total = int(
            await self.db.scalar(select(func.count()).select_from(MgAchatPaiement).where(*filters))
            or 0
        )
        stmt = (
            select(MgAchatPaiement)
            .where(*filters)
            .order_by(MgAchatPaiement.created_at.desc())
        )
        if limit:
            stmt = stmt.limit(limit)
        rows = list((await self.db.execute(stmt)).scalars().all())
        fac_ids = {r.facture_id for r in rows}
        fac_map = {}
        if fac_ids:
            for fid, ref in (
                await self.db.execute(
                    select(MgAchatFacture.id, MgAchatFacture.reference).where(
                        MgAchatFacture.id.in_(fac_ids)
                    )
                )
            ).all():
                fac_map[fid] = ref
        out = []
        tot = Decimal("0")
        for r in rows:
            tot += r.montant or Decimal("0")
            out.append(
                {
                    "id": str(r.id),
                    "reference": r.reference,
                    "facture_ref": fac_map.get(r.facture_id),
                    "statut": r.statut,
                    "montant": _money(r.montant),
                    "date_echeance": _fmt_date(r.date_echeance),
                    "date_paiement": _fmt_date(r.date_paiement),
                    "mode_paiement": r.mode_paiement,
                    "fournisseur": await self._frs_label(r.fournisseur_id),
                }
            )
        return out, total, {"montant": float(tot)}

    async def _rows_echeances(self, f, *, ids, limit):
        from app.services.mg_achats_service import MgAchatsService

        alertes = await MgAchatsService(self.db).list_alertes()
        out = []
        for a in alertes:
            row = {
                "id": str(a.get("entity_id")),
                "type": a.get("type"),
                "reference": a.get("reference"),
                "message": a.get("message"),
                "priorite": a.get("priorite"),
                "date_echeance": _fmt_date(a.get("date_echeance")),
            }
            if ids and uuid.UUID(str(a.get("entity_id"))) not in set(ids):
                continue
            if f.get("priorite") and row["priorite"] != f["priorite"]:
                continue
            if f.get("type") and row["type"] != f["type"]:
                continue
            if f.get("q"):
                q = f["q"].lower()
                blob = f"{row['reference']} {row['message']} {row['type']}".lower()
                if q not in blob:
                    continue
            out.append(row)
        if limit:
            out = out[:limit]
        return out, len(out), {}

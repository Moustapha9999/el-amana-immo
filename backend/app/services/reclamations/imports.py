"""Import Excel des opérations d'un dossier : modèle unique, analyse (aperçu), confirmation.

Une ligne = une opération (client, compte, référence, date, montant, montant contesté, motif).
L'analyse ne modifie pas le dossier : le fichier est conservé en GED (traçabilité), les lignes
contrôlées sont stockées sur l'import avec leurs erreurs / avertissements, les groupes par client
puis par compte et les totaux. La confirmation rejoue les contrôles sur l'état courant du dossier
(un autre utilisateur a pu le modifier) et n'écrit rien tant qu'une ligne est en erreur.
"""

from __future__ import annotations

import hashlib
import io
import re
import unicodedata
import uuid
from collections import defaultdict
from pathlib import Path
from typing import Any

from fastapi import UploadFile
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models import Reclamation, ReclamationDocument, ReclamationImport
from app.models.reclamations import MODES_IMPORT
from app.services.ged_service import GedService
from app.services.reclamations import constantes as C
from app.services.reclamations import operations as ops
from app.services.reclamations.acces import ReclamationsScope
from app.services.reclamations.dossiers import _evenement, charger, peut_modifier, voit_tous_clients
from app.services.reclamations.regles import Exigences, controles_dossier, exigences, nb_justificatifs
from app.services.reclamations.outils import (
    auditer,
    conflit,
    erreur,
    iso,
    maintenant,
    noms_utilisateurs,
    texte_long,
)

# (champ interne, en-tête du modèle, libellé, règle). Les en-têtes reprennent la nomenclature validée.
COLONNES: list[tuple[str, str, str, str]] = [
    ("client_racine", "client_reference", "Identifiant client",
     "Racine client (6 chiffres, zéros conservés). Vide : titulaire du compte, sinon client du dossier."),
    ("compte_reference", "compte_reference", "Référence du compte",
     "Numéro de compte (11 chiffres, zéros conservés) ou RIB. Vide : compte « à renseigner »."),
    ("reference", "reference_operation", "Référence de l'opération",
     "Une même référence peut revenir dans un autre contexte (avertissement)."),
    ("date_operation", "date_operation", "Date de l'opération", "JJ/MM/AAAA, jamais dans le futur."),
    ("nature_mouvement", "nature_mouvement", "Nature du mouvement", "DEBIT ou CREDIT (D / C acceptés)."),
    ("devise", "devise", "Devise", "Code ISO à 3 lettres parmi les devises autorisées. Jamais déduite par le motif "
                                   "si celui-ci l'exige."),
    ("montant", "montant_initial", "Montant initial", "Strictement positif, deux décimales au plus."),
    ("montant_conteste", "montant_a_extourner", "Montant à extourner",
     "Strictement positif, jamais supérieur au montant initial."),
    ("motif_detail", "motif_ligne", "Motif de la ligne", "Précision propre à cette opération."),
    ("commentaire", "commentaire", "Commentaire", "Libre."),
    ("montant_attendu", "montant_attendu", "Montant attendu", "Motif « montant incorrect » : montant qui aurait dû "
                                                              "être comptabilisé."),
    ("compte_attendu", "compte_attendu", "Compte attendu", "Motif « mauvais compte » : compte qui aurait dû être "
                                                           "mouvementé."),
    ("role_operation", "role_operation", "Rôle de l'opération", "Motif « opération en double » : INITIALE ou "
                                                                "DOUBLON."),
]
COLONNES_OPTIONNELLES = ("montant_attendu", "compte_attendu", "role_operation")

# En-têtes reconnus (normalisés) : nomenclature du modèle + libellés français des anciens fichiers.
_ALIAS: dict[str, tuple[str, ...]] = {
    "client_racine": ("client reference", "identifiant client", "racine client", "racine", "client", "id client"),
    "compte_reference": ("compte reference", "reference du compte", "numero de compte", "compte", "rib",
                         "ref compte"),
    "reference": ("reference operation", "reference de l operation", "ref operation", "reference"),
    "date_operation": ("date operation", "date de l operation", "date"),
    "nature_mouvement": ("nature mouvement", "nature du mouvement", "nature", "sens"),
    "devise": ("devise", "monnaie"),
    "montant": ("montant initial", "montant de l operation", "montant operation", "montant"),
    "montant_conteste": ("montant a extourner", "montant conteste"),
    "motif_detail": ("motif ligne", "motif de la ligne", "motif detaille", "motif"),
    "commentaire": ("commentaire", "observation"),
    "montant_attendu": ("montant attendu",),
    "compte_attendu": ("compte attendu",),
    "role_operation": ("role operation", "role de l operation", "role"),
}
FEUILLE = "Opérations"
LIGNES_ENTETE_MAX = 10
EXTENSIONS = (".xlsx",)
TAILLE_MAX_IMPORT = 5 * 1024 * 1024
FORMATS = {"client_racine": "@", "compte_reference": "@", "reference": "@", "compte_attendu": "@",
           "date_operation": "DD/MM/YYYY", "montant": "#,##0.00", "montant_conteste": "#,##0.00",
           "montant_attendu": "#,##0.00"}


def _cle(texte: Any) -> str:
    t = unicodedata.normalize("NFKD", str(texte or "")).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", " ", t).strip()


# ------------------------------------------------------------------ modèle
def colonnes_du_modele(ex: Exigences | None = None) -> list[tuple[str, str, str, str, bool]]:
    """Colonnes du modèle pour un motif : colonnes de base + colonnes spécifiques exigées ; (…, obligatoire)."""
    ex = ex or Exigences()
    return [(champ, code, libelle, regle, ex.requis(champ)) for champ, code, libelle, regle in COLONNES
            if champ not in COLONNES_OPTIONNELLES or ex.requis(champ)]


def modele(ex: Exigences | None = None, motif: str | None = None) -> bytes:
    """Classeur vierge : feuille « Opérations » (en-têtes seuls) + feuille « Consignes ».

    Avec un motif, les colonnes obligatoires sont signalées (rouge) et la consigne du motif est reprise."""
    ex = ex or Exigences()
    colonnes = colonnes_du_modele(ex)
    wb = Workbook()
    ws = wb.active
    ws.title = FEUILLE
    entete = PatternFill("solid", fgColor="1F4E78")
    obligatoire = PatternFill("solid", fgColor="9C0006")
    for i, (champ, code, _libelle, _regle, requis) in enumerate(colonnes, start=1):
        c = ws.cell(row=1, column=i, value=code)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = obligatoire if requis else entete
        c.alignment = Alignment(vertical="center", wrap_text=True)
        lettre = get_column_letter(i)
        ws.column_dimensions[lettre].width = 40 if champ in ("motif_detail", "commentaire") else 22
        fmt = FORMATS.get(champ)
        if fmt:
            for ligne in range(2, C.MAX_LIGNES_IMPORT + 2):
                ws.cell(row=ligne, column=i).number_format = fmt
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 32

    aide = wb.create_sheet("Consignes")
    aide.column_dimensions["A"].width = 28
    aide.column_dimensions["B"].width = 90
    aide.append(["Import des opérations d'une réclamation"])
    aide["A1"].font = Font(bold=True, size=13)
    aide.append([])
    if motif:
        aide.append(["Motif", motif])
        if ex.regle.get("consigne"):
            aide.append(["Consigne", ex.regle["consigne"]])
        aide.append(["Colonnes obligatoires", ", ".join(code for _c, code, _l, _r, req in colonnes if req) or "—"])
        aide.append(["Devises autorisées", ", ".join(sorted(ex.devises))])
        if ex.statut_validation != "VALIDEE":
            aide.append(["Statut de la règle", "Proposée — en attente de validation par la banque"])
        aide.append([])
    for texte in (
        "Une ligne par opération, à partir de la ligne 2 de la feuille « Opérations ». Ne pas modifier les en-têtes.",
        f"{C.MAX_LIGNES_IMPORT} opérations au plus par fichier.",
        "Cas normal : un client, un ou plusieurs comptes. Plusieurs clients dans un même dossier : cas "
        "exceptionnel, autorisé par la banque et justifié lors de la confirmation.",
        "Client ou compte inconnu de la base clientèle : accepté, mais « à vérifier ». Information absente : "
        "« à renseigner ». Rien n'est considéré comme vérifié par défaut.",
        "Une ligne strictement identique à une autre (ou à une opération déjà au dossier) est refusée. Même "
        "compte, même date et même montant sous une autre référence : doublon potentiel signalé, jamais supprimé.",
        "Saisissez les comptes, racines et références en texte : les zéros initiaux doivent être conservés.",
        "Les montants de devises différentes ne sont jamais additionnés : les totaux sont présentés par devise.",
        "Rien n'est enregistré avant la confirmation : l'aperçu présente les erreurs, les avertissements et les totaux.",
        "Une seule ligne en erreur bloque tout l'import (pas d'import partiel) : corrigez le fichier, puis relancez.",
        f"Taille maximale du fichier : {TAILLE_MAX_IMPORT // (1024 * 1024)} Mo.",
    ):
        aide.append(["", texte])
    aide.append([])
    aide.append(["Colonne", "Libellé", "Règle"])
    for col in (1, 2, 3):
        aide.cell(row=aide.max_row, column=col).font = Font(bold=True)
    aide.column_dimensions["C"].width = 80
    for _champ, code, libelle, regle, requis in colonnes:
        aide.append([code, libelle + (" (obligatoire)" if requis else ""), regle])
    for row in aide.iter_rows(min_row=1, max_row=aide.max_row):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------ lecture
def lire_classeur(contenu: bytes) -> list[tuple[int, dict]]:
    """(n° de ligne Excel, valeurs brutes) des lignes non vides de la feuille des opérations."""
    try:
        wb = load_workbook(io.BytesIO(contenu), read_only=True, data_only=True)
    except Exception:
        raise erreur("Fichier Excel illisible : utilisez le modèle (.xlsx)", "FICHIER_ILLISIBLE") from None
    try:
        ws = wb[FEUILLE] if FEUILLE in wb.sheetnames else wb.worksheets[0]
        lignes = ws.iter_rows(values_only=True)
        colonnes: dict[str, int] | None = None
        numero = 0
        for numero, valeurs in enumerate(lignes, start=1):
            colonnes = _entetes(valeurs)
            if colonnes or numero >= LIGNES_ENTETE_MAX:
                break
        if not colonnes or "reference" not in colonnes or "montant" not in colonnes:
            raise erreur("En-têtes introuvables : la feuille doit contenir au moins « reference_operation » "
                         "et « montant_initial » (téléchargez le modèle)", "ENTETES_MANQUANTES")
        resultat: list[tuple[int, dict]] = []
        for rang, valeurs in enumerate(lignes, start=numero + 1):
            brut = {champ: valeurs[i] if i < len(valeurs) else None for champ, i in colonnes.items()}
            if all(v is None or (isinstance(v, str) and not v.strip()) for v in brut.values()):
                continue
            resultat.append((rang, brut))
            if len(resultat) > C.MAX_LIGNES_IMPORT:
                raise erreur(f"{C.MAX_LIGNES_IMPORT} opérations au plus par fichier", "TROP_D_OPERATIONS")
    finally:
        wb.close()
    if not resultat:
        raise erreur("Le fichier ne contient aucune opération", "FICHIER_VIDE")
    return resultat


def _entetes(valeurs: tuple) -> dict[str, int]:
    trouves: dict[str, int] = {}
    cles = [_cle(v) for v in valeurs]
    for champ, alias in _ALIAS.items():
        for a in alias:
            if a in cles and cles.index(a) not in trouves.values():
                trouves[champ] = cles.index(a)
                break
    return trouves if len(trouves) >= 2 else {}


# ------------------------------------------------------------------ analyse
def _autoriser(s: ReclamationsScope, r) -> None:
    if not peut_modifier(s, r):
        if r.statut in C.STATUTS_TERMINAUX or r.archivee_le:
            raise conflit("Dossier terminé : il n'est plus modifiable", "DOSSIER_TERMINE")
        s.exiger("reclamations.update")
    if not voit_tous_clients(s, r):
        raise erreur("Ce dossier concerne des clients que vous ne pouvez pas consulter", "PERMISSION_DENIED", 403)


async def _controler(db: AsyncSession, r, lignes: list[ops.Ligne], mode: str,
                     ex: Exigences | None = None) -> list[ops.Ligne]:
    """Contrôles complets sur l'état courant ; renvoie les opérations du dossier conservées."""
    if mode == "AJOUT":
        actuelles = await ops.lister(db, r.id)
        par_p = {p.id: p for p in await ops.participants(db, r.id)}
        par_c = {c.id: c for c in await ops.comptes(db, r.id)}
        existantes = ops.lignes_existantes(actuelles, par_p, par_c)
        rattachement = await ops.comptes_du_dossier(db, r.id)
    else:
        existantes = []
        rattachement = {r.compte_reference: r.client_racine} if r.compte_reference else {}
    await ops.resoudre(db, lignes, r.client_racine, rattachement, ex)
    ops.controler(lignes, existantes)
    return existantes


def _synthese(r, lignes: list[ops.Ligne], conservees: list[ops.Ligne], autorise: bool, ex: Exigences,
              nb_pieces: int = 0) -> dict:
    """Groupes client → compte, totaux par devise du fichier et du dossier après import."""
    valides = [lg for lg in lignes if not lg.erreurs]
    groupes: dict[str | None, dict[str | None, list[ops.Ligne]]] = defaultdict(lambda: defaultdict(list))
    for lg in lignes:
        groupes[lg.client_racine][lg.compte_reference].append(lg)

    def totaux(liste: list[ops.Ligne]) -> dict:
        ok = [lg for lg in liste if not lg.erreurs]
        montant, devise = ops.montant_unique(ok)
        conteste, _ = ops.montant_unique(ok, ops.conteste)
        return {"nb": len(liste), "nb_erreurs": len(liste) - len(ok),
                "montant": ops.montant_str(montant), "montant_conteste": ops.montant_str(conteste),
                "devise": devise, "par_devise": ops.par_devise(ok)}

    clients = []
    for racine, par_compte in sorted(groupes.items(), key=lambda x: (x[0] != r.client_racine, x[0] or "~")):
        toutes = [lg for liste in par_compte.values() for lg in liste]
        premiere = next((lg for lg in toutes if lg.client_nom), toutes[0])
        clients.append({
            "racine": racine, "nom": premiere.client_nom if premiere.client_nom else (
                r.client_nom if racine and racine == r.client_racine else None),
            "principal": racine == r.client_racine, "statut": premiere.statut_client or "A_RENSEIGNER",
            **totaux(toutes),
            "comptes": [{"compte_reference": compte, "statut": liste[0].statut_compte or "A_RENSEIGNER",
                         **totaux(liste)}
                        for compte, liste in sorted(par_compte.items(), key=lambda x: x[0] or "~")],
        })
    racines = {lg.client_racine for lg in valides + conservees if lg.client_racine}
    if r.client_racine:
        racines.add(r.client_racine)
    multi = len(racines) > 1
    apres = conservees + valides
    return {
        "lignes": [lg.apercu() for lg in lignes],
        "clients": clients,
        "totaux": totaux(lignes),
        "nb_clients": len(racines), "multi_clients": multi, "regroupement_autorise": autorise,
        "justification_requise": multi and not r.justification_regroupement,
        "justification_actuelle": r.justification_regroupement,
        "apres_import": totaux(apres),
        "nb_conservees": len(conservees),
        "recapitulatif": ops.recapitulatif(lignes, ex, nb_pieces=nb_pieces, racine_principale=r.client_racine,
                                           manques=controles_dossier(ex, apres, nb_pieces)),
        "colonnes": [{"champ": champ, "code": code, "libelle": libelle, "obligatoire": requis}
                     for champ, code, libelle, _regle, requis in colonnes_du_modele(ex)],
    }


async def _deja_importe(db: AsyncSession, empreinte: str, rid: uuid.UUID) -> tuple[bool, list[str]]:
    """Même fichier (empreinte SHA-256) déjà confirmé : sur ce dossier ? sur quels autres dossiers ?"""
    rows = (await db.execute(
        select(ReclamationImport.reclamation_id, Reclamation.reference)
        .join(Reclamation, Reclamation.id == ReclamationImport.reclamation_id)
        .where(ReclamationImport.empreinte == empreinte, ReclamationImport.statut == "CONFIRME"))).all()
    return any(i == rid for i, _ref in rows), sorted({ref for i, ref in rows if i != rid})


async def analyser(db: AsyncSession, s: ReclamationsScope, rid: uuid.UUID, fichier: UploadFile, mode: str,
                   *, forcer: bool = False) -> dict:
    """``forcer`` : réanalyser un fichier déjà importé sur ce dossier (choix explicite, tracé)."""
    if mode not in MODES_IMPORT:
        raise erreur("Mode d'import inconnu", "MODE_INVALIDE")
    nom = Path(fichier.filename or "").name
    if Path(nom).suffix.lower() not in EXTENSIONS:
        raise erreur("Format attendu : classeur Excel .xlsx (téléchargez le modèle)", "FORMAT_IMPORT")
    r = await charger(db, s, rid, verrou=True)
    _autoriser(s, r)
    contenu = await fichier.read(TAILLE_MAX_IMPORT + 1)
    if len(contenu) > TAILLE_MAX_IMPORT:
        raise erreur(f"Fichier trop volumineux ({TAILLE_MAX_IMPORT // (1024 * 1024)} Mo au plus)", "FICHIER_TROP_GROS",
                     413)
    empreinte = hashlib.sha256(contenu).hexdigest()
    ici, ailleurs = await _deja_importe(db, empreinte, r.id)
    if ici and not forcer:
        raise conflit("Ce fichier a déjà été importé sur ce dossier. Confirmez explicitement pour l'analyser à "
                      "nouveau.", "FICHIER_DEJA_IMPORTE")
    ex = await exigences(db, r.motif_id)
    lignes = [ops.lire(brut, rang, ex, devise_defaut=C.DEVISE) for rang, brut in lire_classeur(contenu)]
    conservees = await _controler(db, r, lignes, mode, ex)
    resultat = _synthese(r, lignes, conservees, await ops.regroupement_autorise(db), ex,
                         await nb_justificatifs(db, r.id))
    resultat["fichier_deja_importe"] = {"ce_dossier": ici, "autres_dossiers": ailleurs, "force": bool(ici and forcer)}
    await fichier.seek(0)

    for ancien in (await db.scalars(select(ReclamationImport).where(
            ReclamationImport.reclamation_id == r.id, ReclamationImport.statut == "ANALYSE"))).all():
        ancien.statut = "ABANDONNE"
    doc = await GedService(db).upload(file=fichier, espace_code=C.ESPACE_CODE, module_code=C.MODULE_CODE,
                                      entity=C.GED_ENTITE, entity_id=str(r.id), uploaded_by_id=s.user_id)
    doc.doc_type = "RECLAMATION_IMPORT_OPERATIONS"
    doc.title = f"{C.CATEGORIE_DOCUMENT_LIBELLES['IMPORT_OPERATIONS']} — {r.reference}"[:255]
    doc.reference = r.reference
    doc.date_document = maintenant().date()
    doc.security_level = "restricted"
    doc.ocr_status = "skipped"
    db.add(ReclamationDocument(document_id=doc.id, reclamation_id=r.id, categorie="IMPORT_OPERATIONS",
                               created_by_id=s.user_id))
    nb_err = sum(1 for lg in lignes if lg.erreurs)
    nb_av = sum(1 for lg in lignes if not lg.erreurs and lg.avertissements)
    imp = ReclamationImport(
        reclamation_id=r.id, document_id=doc.id, nom_fichier=nom[:255], empreinte=empreinte,
        statut="ANALYSE", mode=mode, nb_lignes=len(lignes), nb_valides=len(lignes) - nb_err, nb_erreurs=nb_err,
        nb_avertissements=nb_av, resultat=resultat, created_by_id=s.user_id)
    db.add(imp)
    await db.flush()
    db.add(_evenement(r, s, "IMPORT_ANALYSE", statut_avant=r.statut, etape_avant=r.etape,
                      donnees={"fichier": nom, "mode": mode, "lignes": len(lignes), "erreurs": nb_err,
                               "avertissements": nb_av}))
    await db.flush()
    await auditer(db, s, "reclamations.operations.import.analyse", C.GED_ENTITE, r.id, after={
        "import_id": str(imp.id), "document_id": str(doc.id), "fichier": nom, "empreinte": imp.empreinte,
        "mode": mode, "lignes": len(lignes), "erreurs": nb_err, "avertissements": nb_av,
        **({"reimport_force": True} if ici else {}), **({"deja_importe_sur": ailleurs} if ailleurs else {})})
    return await _dict(db, imp, r)


async def _import(db: AsyncSession, rid: uuid.UUID, iid: uuid.UUID) -> ReclamationImport:
    imp = await db.get(ReclamationImport, iid)
    if imp is None or imp.reclamation_id != rid:
        raise erreur("Import introuvable", "NOT_FOUND", 404)
    return imp


async def _dict(db: AsyncSession, imp: ReclamationImport, r) -> dict:
    noms = await noms_utilisateurs(db, {imp.created_by_id, imp.confirme_par_id} - {None})
    return {
        "id": str(imp.id), "reclamation_id": str(r.id), "revision": r.revision,
        "nom_fichier": imp.nom_fichier, "statut": imp.statut, "mode": imp.mode,
        "nb_lignes": imp.nb_lignes, "nb_valides": imp.nb_valides, "nb_erreurs": imp.nb_erreurs,
        "nb_avertissements": imp.nb_avertissements,
        "document_id": str(imp.document_id) if imp.document_id else None,
        "created_at": iso(imp.created_at), "par": noms.get(imp.created_by_id) if imp.created_by_id else None,
        "confirme_le": iso(imp.confirme_le),
        "confirme_par": noms.get(imp.confirme_par_id) if imp.confirme_par_id else None,
        **(imp.resultat or {}),
    }


async def lire(db: AsyncSession, s: ReclamationsScope, rid: uuid.UUID, iid: uuid.UUID) -> dict:
    r = await charger(db, s, rid)
    if not voit_tous_clients(s, r):
        raise erreur("Ce dossier concerne des clients que vous ne pouvez pas consulter", "PERMISSION_DENIED", 403)
    return await _dict(db, await _import(db, r.id, iid), r)


async def rapport_erreurs(db: AsyncSession, s: ReclamationsScope, rid: uuid.UUID, iid: uuid.UUID) -> bytes:
    """Rapport Excel de l'analyse : une ligne par erreur ou avertissement (ligne Excel, colonne, message)."""
    r = await charger(db, s, rid)
    if not voit_tous_clients(s, r):
        raise erreur("Ce dossier concerne des clients que vous ne pouvez pas consulter", "PERMISSION_DENIED", 403)
    imp = await _import(db, r.id, iid)
    codes = {champ: code for champ, code, _l, _r in COLONNES}
    wb = Workbook()
    ws = wb.active
    ws.title = "Rapport"
    ws.append([f"Import « {imp.nom_fichier} » — dossier {r.reference}"])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append([f"{imp.nb_lignes} ligne(s) · {imp.nb_erreurs} en erreur · {imp.nb_avertissements} avec "
               "avertissement · une seule erreur bloque tout l'import"])
    ws.append([])
    ws.append(["Ligne Excel", "Statut", "Colonne", "Code", "Message", "Référence de l'opération"])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
    for lg in (imp.resultat or {}).get("lignes", []):
        detail = lg.get("erreurs_detail") or [{"code": None, "message": m, "champ": None}
                                              for m in lg.get("erreurs") or []]
        for e in detail:
            ws.append([lg.get("ligne"), "Erreur", codes.get(e.get("champ") or "", e.get("champ")), e.get("code"),
                       e.get("message"), lg.get("reference")])
        for message in lg.get("avertissements") or []:
            ws.append([lg.get("ligne"), "Avertissement", None, None, message, lg.get("reference")])
    for lettre, largeur in zip("ABCDEF", (12, 16, 24, 30, 80, 28), strict=True):
        ws.column_dimensions[lettre].width = largeur
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def confirmer(db: AsyncSession, s: ReclamationsScope, rid: uuid.UUID, iid: uuid.UUID, data: dict) -> None:
    r = await charger(db, s, rid, verrou=True)
    _autoriser(s, r)
    imp = await _import(db, r.id, iid)
    if imp.statut != "ANALYSE":
        raise conflit("Cet import a déjà été confirmé ou abandonné", "IMPORT_DEJA_TRAITE")
    if data["revision"] != r.revision:
        raise conflit("Le dossier a été modifié depuis l'analyse : relancez l'import", "CONFLIT_REVISION",
                      revision=r.revision)
    ex = await exigences(db, r.motif_id)
    lignes = [ops.lire(b, int(b.get("ligne") or 0), ex, devise_defaut=C.DEVISE)
              for b in (imp.resultat or {}).get("lignes", [])]
    await _controler(db, r, lignes, imp.mode or "AJOUT", ex)
    en_erreur = [lg for lg in lignes if lg.erreurs]
    if en_erreur:
        lg = en_erreur[0]
        raise erreur(f"{len(en_erreur)} ligne(s) en erreur (ligne {lg.rang} : {lg.erreurs[0][1]}) : corrigez le "
                     "fichier et relancez l'import", "IMPORT_ERREURS")
    if not lignes:
        raise erreur("Aucune opération à importer", "FICHIER_VIDE")
    justification = texte_long(data.get("justification_regroupement"))
    prevus = {lg.client_racine for lg in lignes if lg.client_racine} | ({r.client_racine} if r.client_racine else set())
    if imp.mode != "REMPLACEMENT":
        prevus |= {p.client_racine for p in await ops.participants(db, r.id) if p.client_racine}
    await ops.verifier_regroupement(db, r, justification, prevus)
    actuelles = await ops.lister(db, r.id)
    avant = {"nb": len(actuelles), "par_devise": ops.par_devise(actuelles)}
    if imp.mode == "REMPLACEMENT":
        for o in actuelles:
            await db.delete(o)
        await db.flush()
    await ops.ecrire(db, r, lignes, s.user_id, import_id=imp.id)
    await ops.nettoyer(db, r)
    multi = await ops.verifier_regroupement(db, r, justification)
    imp.statut, imp.confirme_le, imp.confirme_par_id = "CONFIRME", maintenant(), s.user_id
    r.revision += 1
    r.updated_by_id = s.user_id
    apres_ops = await ops.lister(db, r.id)
    apres = {"nb": len(apres_ops), "par_devise": ops.par_devise(apres_ops)}
    db.add(_evenement(r, s, "IMPORT_CONFIRME", statut_avant=r.statut, etape_avant=r.etape,
                      donnees={"fichier": imp.nom_fichier, "mode": imp.mode, "importees": len(lignes),
                               "avant": avant, "apres": apres, **({"multi_clients": True} if multi else {})}))
    await db.flush()
    await auditer(db, s, "reclamations.operations.import.confirm", C.GED_ENTITE, r.id, before=avant, after={
        "import_id": str(imp.id), "fichier": imp.nom_fichier, "mode": imp.mode, "importees": len(lignes), **apres,
        **({"justification_regroupement": r.justification_regroupement} if multi else {})})


async def abandonner(db: AsyncSession, s: ReclamationsScope, rid: uuid.UUID, iid: uuid.UUID) -> dict:
    r = await charger(db, s, rid, verrou=True)
    _autoriser(s, r)
    imp = await _import(db, r.id, iid)
    if imp.statut != "ANALYSE":
        raise conflit("Cet import a déjà été confirmé ou abandonné", "IMPORT_DEJA_TRAITE")
    imp.statut = "ABANDONNE"
    db.add(_evenement(r, s, "IMPORT_ABANDONNE", statut_avant=r.statut, etape_avant=r.etape,
                      donnees={"fichier": imp.nom_fichier}))
    await db.flush()
    await auditer(db, s, "reclamations.operations.import.abandon", C.GED_ENTITE, r.id,
                  after={"import_id": str(imp.id), "fichier": imp.nom_fichier})
    return await _dict(db, imp, r)

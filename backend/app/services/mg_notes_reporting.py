"""Reporting Notes de frais — JSON / CSV / PDF / Excel."""

from __future__ import annotations

import csv
import io
from datetime import date
from decimal import Decimal
from typing import Any
from uuid import UUID

from fastapi import HTTPException, status
from fastapi.responses import Response
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.auth import User
from app.services.mg_notes_service import MgNotesService
from app.services.reporting_export import build_styled_pdf, export_now


REPORTS = {
    "periode": "Notes de frais par période",
    "agence": "Notes par intitulé",
    "demandeur": "Notes par demandeur",
    "categorie": "Dépenses par catégorie",
    "paiement": "Rapport des paiements",
    "attente": "Notes en attente",
    "annuel": "Rapport annuel",
}


class MgNotesReporting:
    def __init__(self, db: AsyncSession):
        self.db = db
        self.svc = MgNotesService(db)

    async def export(
        self,
        user: User,
        report_key: str,
        *,
        fmt: str = "json",
        annee: int | None = None,
        mois: int | None = None,
        agence_id: UUID | None = None,
        statut: str | None = None,
        date_debut: date | None = None,
        date_fin: date | None = None,
    ):
        if report_key not in REPORTS:
            raise HTTPException(status.HTTP_404_NOT_FOUND, detail="Rapport inconnu")

        notes, _ = await self.svc.list_notes(
            user,
            statut=statut,
            agence_id=agence_id,
            date_debut=date_debut,
            date_fin=date_fin,
            page=1,
            size=200,
        )
        if annee:
            notes = [n for n in notes if n.date_demande.year == annee]
        if mois:
            notes = [n for n in notes if n.date_demande.month == mois]

        if report_key == "attente":
            notes = [
                n
                for n in notes
                if n.statut
                in {
                    "SOUMIS",
                    "EN_CONTROLE",
                    "CORRECTION_REQUISE",
                    "VISA_MG",
                    "VISA_DR",
                    "VALIDEE",
                    "MISE_EN_PAIEMENT",
                    "PARTIELLEMENT_PAYEE",
                }
            ]
        elif report_key == "paiement":
            notes = [
                n
                for n in notes
                if n.statut
                in {"VALIDEE", "MISE_EN_PAIEMENT", "PARTIELLEMENT_PAYEE", "PAYEE", "CLOTUREE", "ARCHIVEE"}
            ]

        if report_key == "categorie":
            headers = ["Catégorie", "Nb lignes", "Montant"]
            buckets: dict[str, list[Any]] = {}
            for n in notes:
                for lig in n.lignes:
                    key = lig.categorie_libelle_snapshot or "Sans catégorie"
                    if key not in buckets:
                        buckets[key] = [0, Decimal("0")]
                    buckets[key][0] += 1
                    buckets[key][1] += lig.montant or Decimal("0")
            rows = [[k, v[0], float(v[1])] for k, v in sorted(buckets.items())]
        elif report_key == "agence":
            headers = ["Intitulé", "Nb notes", "Montant total", "Montant payé"]
            buckets = {}
            for n in notes:
                key = n.agence_libelle_snapshot or "—"
                if key not in buckets:
                    buckets[key] = [0, Decimal("0"), Decimal("0")]
                buckets[key][0] += 1
                buckets[key][1] += n.total_mru or Decimal("0")
                buckets[key][2] += n.montant_paye or Decimal("0")
            rows = [[k, v[0], float(v[1]), float(v[2])] for k, v in sorted(buckets.items())]
        elif report_key == "demandeur":
            headers = ["Demandeur", "Nb notes", "Montant total"]
            buckets = {}
            for n in notes:
                key = n.demandeur_nom or "—"
                if key not in buckets:
                    buckets[key] = [0, Decimal("0")]
                buckets[key][0] += 1
                buckets[key][1] += n.total_mru or Decimal("0")
            rows = [[k, v[0], float(v[1])] for k, v in sorted(buckets.items())]
        elif report_key == "annuel":
            year = annee or date.today().year
            headers = ["Mois", "Nb notes", "Montant", "Payé", "Reste"]
            months = {m: [0, Decimal("0"), Decimal("0")] for m in range(1, 13)}
            for n in notes:
                if n.date_demande.year != year:
                    continue
                m = n.date_demande.month
                months[m][0] += 1
                months[m][1] += n.total_mru or Decimal("0")
                months[m][2] += n.montant_paye or Decimal("0")
            noms = [
                "",
                "Janvier",
                "Février",
                "Mars",
                "Avril",
                "Mai",
                "Juin",
                "Juillet",
                "Août",
                "Septembre",
                "Octobre",
                "Novembre",
                "Décembre",
            ]
            rows = [
                [
                    noms[m],
                    months[m][0],
                    float(months[m][1]),
                    float(months[m][2]),
                    float(months[m][1] - months[m][2]),
                ]
                for m in range(1, 13)
            ]
        else:
            headers = [
                "Référence",
                "Date",
                "Demandeur",
                "Intitulé",
                "Département",
                "Montant",
                "Payé",
                "Reste",
                "Statut",
            ]
            rows = [
                [
                    n.reference,
                    str(n.date_demande),
                    n.demandeur_nom or "",
                    n.agence_libelle_snapshot or "",
                    n.departement or "",
                    float(n.total_mru or 0),
                    float(n.montant_paye or 0),
                    float((n.total_mru or 0) - (n.montant_paye or 0)),
                    n.statut,
                ]
                for n in notes
            ]

        if fmt == "json":
            return {
                "key": report_key,
                "title": REPORTS[report_key],
                "headers": headers,
                "rows": rows,
                "count": len(rows),
            }

        if not rows:
            raise HTTPException(
                status.HTTP_404_NOT_FOUND,
                detail="Aucune donnée ne correspond aux critères sélectionnés.",
            )

        when = export_now()
        title = f"BEA-DIGITAL — {REPORTS[report_key]}"
        subtitle = f"Généré par {user.full_name or user.email}"
        stamp = when.strftime("%Y%m%d-%H%M")
        base = f"notes-{report_key}-{stamp}"

        if fmt == "csv":
            buf = io.StringIO()
            w = csv.writer(buf, delimiter=";")
            w.writerow(headers)
            w.writerows(rows)
            return Response(
                content=buf.getvalue().encode("utf-8-sig"),
                media_type="text/csv; charset=utf-8",
                headers={"Content-Disposition": f'attachment; filename="{base}.csv"'},
            )
        if fmt == "pdf":
            content = build_styled_pdf(
                report_title=title,
                headers=headers,
                rows=rows,
                subtitle=subtitle,
                exported_at=when,
                landscape_mode=True,
            )
            return Response(
                content=content,
                media_type="application/pdf",
                headers={"Content-Disposition": f'attachment; filename="{base}.pdf"'},
            )

        # xlsx — synthèse + notes + dépenses
        from openpyxl import Workbook
        from openpyxl.styles import Font

        wb = Workbook()
        ws = wb.active
        ws.title = "Synthese"
        ws["A1"] = title
        ws["A1"].font = Font(bold=True, size=14)
        ws["A2"] = subtitle
        ws["A3"] = f"Lignes : {len(rows)}"
        for i, h in enumerate(headers, 1):
            ws.cell(4, i, h).font = Font(bold=True)
        for r_i, row in enumerate(rows, 5):
            for c_i, val in enumerate(row, 1):
                ws.cell(r_i, c_i, val)

        ws2 = wb.create_sheet("Notes")
        nh = ["Référence", "Date", "Demandeur", "Intitulé", "Montant", "Payé", "Statut"]
        for i, h in enumerate(nh, 1):
            ws2.cell(1, i, h).font = Font(bold=True)
        for r_i, n in enumerate(notes, 2):
            ws2.cell(r_i, 1, n.reference)
            ws2.cell(r_i, 2, str(n.date_demande))
            ws2.cell(r_i, 3, n.demandeur_nom or "")
            ws2.cell(r_i, 4, n.agence_libelle_snapshot or "")
            ws2.cell(r_i, 5, float(n.total_mru or 0))
            ws2.cell(r_i, 6, float(n.montant_paye or 0))
            ws2.cell(r_i, 7, n.statut)

        ws3 = wb.create_sheet("Depenses")
        dh = ["Référence", "Date", "Catégorie", "Description", "Motif", "Montant", "Règlement"]
        for i, h in enumerate(dh, 1):
            ws3.cell(1, i, h).font = Font(bold=True)
        r = 2
        for n in notes:
            for lig in n.lignes:
                ws3.cell(r, 1, n.reference)
                ws3.cell(r, 2, str(lig.date_depense))
                ws3.cell(r, 3, lig.categorie_libelle_snapshot or "")
                ws3.cell(r, 4, lig.description)
                ws3.cell(r, 5, lig.motif or "")
                ws3.cell(r, 6, float(lig.montant or 0))
                ws3.cell(r, 7, lig.mode_reglement or "")
                r += 1

        buf = io.BytesIO()
        wb.save(buf)
        return Response(
            content=buf.getvalue(),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="{base}.xlsx"'},
        )

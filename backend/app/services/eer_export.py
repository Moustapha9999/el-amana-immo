"""Exports EER (PDF / Excel) : liste des dossiers, synthèse de conformité, fiche dossier.

La synthèse reprend la Feuil2 du suivi Excel (par agence, par profil, par état du compte)
sur la référence choisie : ``excel`` (S / T historiques) ou ``bea`` (décision BEA-DIGITAL).
Mise en forme commune BEA DIGITAL (``reporting_export``).
"""

from __future__ import annotations

from collections.abc import Sequence
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from typing import Any, Literal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import mm
from reportlab.platypus import Image, KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.services.reporting_export import (
    _COLOR_BORDER,
    _COLOR_HEADER_FG,
    _COLOR_META,
    _COLOR_NAVY,
    _COLOR_TITLE,
    _COLOR_ZEBRA,
    _THIN,
    _bank_line,
    _pdf_footer,
    _pdf_styles,
    build_styled_pdf,
    build_styled_workbook,
    export_now,
    format_export_datetime,
    resolve_bea_logo_path,
)

Reference = Literal["excel", "bea"]
LIBELLE_REFERENCE = {"excel": "Référence Excel historique (S / T)", "bea": "Décision BEA-DIGITAL"}
LIBELLE_CLASSEMENT = {"CONFORME": "Conforme", "NON_CONFORME": "Non conforme", "NON_EVALUE": "Non évalué"}


def libelle(code: Any) -> str:
    if code is None or code == "":
        return "—"
    if isinstance(code, bool):
        return "Oui" if code else "Non"
    texte = str(code).replace("_", " ").lower()
    return texte[:1].upper() + texte[1:]


def _d(v: date | datetime | None) -> str:
    if v is None:
        return "—"
    return v.strftime("%d/%m/%Y %H:%M") if isinstance(v, datetime) else v.strftime("%d/%m/%Y")


def _taux(v: Decimal | float | None) -> str:
    return "—" if v is None else f"{Decimal(str(v)):.2f} %".replace(".", ",")


@dataclass
class Section:
    titre: str
    entetes: Sequence[str]
    lignes: list[list[Any]] = field(default_factory=list)
    alignements: Sequence[str] | None = None
    total: list[Any] | None = None
    largeurs: Sequence[float] | None = None


# --- Liste des dossiers --------------------------------------------------------------------

ENTETES_LISTE = ["Référence", "Date EER", "Agence", "Client", "Type client", "Profil", "Statut", "Risque",
                 "Analyste", "Conformité Excel", "Décision BEA", "Version"]


def ligne_liste(i: dict[str, Any]) -> list[Any]:
    return [
        i["reference"], _d(i["date_eer"]), i.get("agence_code") or "—", i.get("client_nom") or "—",
        libelle(i.get("type_client")), libelle(i.get("profil")), libelle(i.get("statut")), libelle(i.get("risque")),
        i.get("analyste_nom") or "—", LIBELLE_CLASSEMENT.get(str(i.get("conformite_excel")), "—"),
        libelle(i.get("decision")), i.get("version_courante"),
    ]


def liste_excel(items: list[dict[str, Any]], sous_titre: str | None) -> bytes:
    return build_styled_workbook(sheet_title="Dossiers EER", report_title="Entrées en relation — dossiers",
                                 headers=ENTETES_LISTE, rows=[ligne_liste(i) for i in items], subtitle=sous_titre)


def liste_pdf(items: list[dict[str, Any]], sous_titre: str | None) -> bytes:
    return build_styled_pdf(report_title="Entrées en relation — dossiers", headers=ENTETES_LISTE,
                            rows=[ligne_liste(i) for i in items], subtitle=sous_titre,
                            col_widths=[30, 20, 16, 42, 26, 26, 24, 16, 30, 24, 22, 12])


# --- Synthèse de conformité (Feuil2) ----------------------------------------------------------

def _ligne_conformite(nom: str, bloc: dict[str, Any], reference: Reference) -> list[Any]:
    c = bloc[reference]
    return [nom, c["conformes"], c["non_conformes"], c["non_evalues"], bloc["total"], _taux(c["taux"])]


ENTETES_CONFORMITE = ["Conformes", "Non conformes", "Non évalués", "Total", "Taux de conformité"]
ALIGN_CONFORMITE = ["left", "right", "right", "right", "right", "right"]


def sections_synthese(global_: dict, agences: dict, profils: dict, etats: dict,
                      reference: Reference) -> list[Section]:
    g, flux = global_, global_["flux"]
    ref = g[reference]
    return [
        Section("Synthèse", ["Indicateur", "Valeur"], [
            ["Dossiers évalués (hors abandons)", g["total"]],
            ["Conformes", ref["conformes"]],
            ["Non conformes", ref["non_conformes"]],
            ["Non évalués", ref["non_evalues"]],
            ["Taux de conformité", _taux(ref["taux"])],
            ["Écarts référence Excel / décision BEA-DIGITAL", g["divergences"]],
            ["Dossiers reçus", flux["recus"]],
            ["En cours", flux["en_cours"]],
            ["À compléter", flux["a_completer"]],
            ["Abandonnés", flux["abandonnes"]],
            ["Taux d'abandon", _taux(flux["taux_abandon"])],
        ], ["left", "right"], largeurs=[120, 50]),
        Section("Conformité par agence", ["Agence", *ENTETES_CONFORMITE],
                [_ligne_conformite(f"{a['code']} — {a['libelle']}", a, reference) for a in agences["lignes"]],
                ALIGN_CONFORMITE, _ligne_conformite("Total", agences["total"], reference)),
        Section("Conformité par profil", ["Profil", *ENTETES_CONFORMITE],
                [_ligne_conformite(p["libelle"], p, reference) for p in profils["lignes"]],
                ALIGN_CONFORMITE, _ligne_conformite("Total", profils["total"], reference)),
        Section("État du compte", ["État", "Nombre", "Part"],
                [[e["libelle"], e["nombre"], _taux(e["pourcentage"])] for e in etats["lignes"]],
                ["left", "right", "right"],
                ["Total", etats["total"], "100,00 %" if etats["total"] else "—"], largeurs=[90, 40, 40]),
    ]


# --- Fiche dossier ----------------------------------------------------------------------------

def sections_fiche(d: dict[str, Any], parties: list[dict], checklist: list[dict], anomalies: list[dict],
                   visas: list[dict], documents: list[dict]) -> list[Section]:
    sections = [
        Section("Dossier", ["Rubrique", "Valeur"], [
            ["Référence", d["reference"]], ["Agence", d["agence"]], ["Date EER", _d(d["date_eer"])],
            ["Opération", libelle(d["operation_type"])], ["Type client", libelle(d["type_client"])],
            ["Profil", libelle(d["profil"])], ["Statut", libelle(d["statut"])], ["Étape", libelle(d["etape"])],
            ["Risque LBC-FT", libelle(d["risque"])], ["PPE", libelle(d["ppe"])], ["FATCA", libelle(d["fatca"])],
            ["Avis Conformité requis", libelle(d["avis_requis"])], ["Version", d["version_courante"]],
            ["Conformité physique", libelle(d["conformite_physique"])],
            ["Conformité système", libelle(d["conformite_systeme"])],
            ["Référence Excel", LIBELLE_CLASSEMENT.get(d["conformite_excel"], "—")],
            ["Décision BEA-DIGITAL", libelle(d["decision"])],
            ["Analyste", d.get("analyste") or "—"], ["Soumis le", _d(d["soumis_le"])], ["Validé le", _d(d["valide_le"])],
        ], ["left", "left"], largeurs=[60, 120]),
        Section("Parties", ["Rôle", "Nom", "Nature", "PPE", "FATCA", "Risque"],
                [[libelle(p["role"]), p["nom"], libelle(p["nature"]), libelle(p["ppe"]), libelle(p["fatca"]),
                  libelle(p["risque"])] for p in parties]),
        Section("Checklist", ["Élément", "Catégorie", "Obligatoire", "Présence", "Statut", "Motif"],
                [[i["libelle"], libelle(i["categorie"]), libelle(i["obligatoire"]), libelle(i["presence"]),
                  libelle(i["statut"]), i.get("motif") or ""] for i in checklist],
                largeurs=[70, 28, 20, 20, 26, 60]),
    ]
    if anomalies:
        sections.append(Section("Anomalies", ["Type", "Gravité", "Statut", "Description", "Échéance"],
                                [[libelle(a["type_code"]), libelle(a["gravite"]), libelle(a["statut"]),
                                  a["description"], _d(a["echeance_regularisation"])] for a in anomalies],
                                largeurs=[36, 22, 24, 100, 22]))
    if visas:
        sections.append(Section("Avis Conformité KYC", ["Fonction", "Avis", "Version", "Date", "Commentaire"],
                                [[v["fonction"], libelle(v["avis"]), v["version"], _d(v["vise_le"]),
                                  v.get("commentaire") or ""] for v in visas]))
    if documents:
        sections.append(Section("Pièces GED", ["Type", "Fichier", "Éléments", "Déposé le"],
                                [[doc["type"], doc["filename"], ", ".join(doc["elements"]) or "—",
                                  _d(doc["created_at"])] for doc in documents]))
    return sections


# --- Rendu multi-sections ---------------------------------------------------------------------

def sections_excel(titre: str, sous_titre: str | None, sections: list[Section], feuille: str) -> bytes:
    when = export_now()
    wb = Workbook()
    ws = wb.active
    ws.title = feuille[:31]
    n_cols = max(len(s.entetes) for s in sections)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n_cols)
    ws["A1"].value = _bank_line()
    ws["A1"].font = Font(name="Calibri", size=11, bold=True, color=_COLOR_NAVY)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=n_cols)
    ws["A2"].value = titre
    ws["A2"].font = Font(name="Calibri", size=14, bold=True, color=_COLOR_TITLE)
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=n_cols)
    ws["A3"].value = "  ·  ".join(p for p in (f"Exporté le {format_export_datetime(when)}", sous_titre) if p)
    ws["A3"].font = Font(name="Calibri", size=10, italic=True, color=_COLOR_META)

    entete_fill = PatternFill("solid", fgColor=_COLOR_NAVY)
    zebra = PatternFill("solid", fgColor=_COLOR_ZEBRA)
    total_fill = PatternFill("solid", fgColor="E2E8F0")
    ligne = 5
    largeurs = [12] * (n_cols + 1)
    for s in sections:
        ws.cell(row=ligne, column=1, value=s.titre).font = Font(name="Calibri", size=12, bold=True, color=_COLOR_NAVY)
        ligne += 1
        for c, h in enumerate(s.entetes, start=1):
            cell = ws.cell(row=ligne, column=c, value=h)
            cell.fill, cell.border = entete_fill, _THIN
            cell.font = Font(name="Calibri", size=10, bold=True, color=_COLOR_HEADER_FG)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            largeurs[c] = max(largeurs[c], min(48, len(str(h)) + 2))
        for r, valeurs in enumerate([*s.lignes, *([s.total] if s.total else [])]):
            ligne += 1
            est_total = s.total is not None and r == len(s.lignes)
            for c, v in enumerate(valeurs, start=1):
                cell = ws.cell(row=ligne, column=c, value=v)
                cell.border = _THIN
                cell.font = Font(name="Calibri", size=10, bold=est_total, color=_COLOR_TITLE)
                align = (s.alignements or [])[c - 1] if s.alignements and c <= len(s.alignements) else "left"
                cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
                if est_total:
                    cell.fill = total_fill
                elif r % 2 == 1:
                    cell.fill = zebra
                largeurs[c] = max(largeurs[c], min(60, len(str(v)) + 2))
        ligne += 2
    for c in range(1, n_cols + 1):
        ws.column_dimensions[get_column_letter(c)].width = largeurs[c]
    ws.oddFooter.center.text = f"Exporté le {format_export_datetime(when)} — Page &P / &N"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def sections_pdf(titre: str, sous_titre: str | None, sections: list[Section], *, paysage: bool = False) -> bytes:
    when = export_now()
    exporte = f"Exporté le {format_export_datetime(when)}"
    styles = _pdf_styles()
    pagesize = landscape(A4) if paysage else A4
    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=pagesize, title=titre, leftMargin=12 * mm, rightMargin=12 * mm,
                            topMargin=16 * mm, bottomMargin=18 * mm)
    utile = pagesize[0] - 24 * mm
    story: list = []
    logo = resolve_bea_logo_path()
    if logo is not None:
        try:
            story += [Image(str(logo), width=42 * mm, height=14 * mm), Spacer(1, 2 * mm)]
        except Exception:
            pass
    story += [Paragraph(_bank_line(), styles["bank"]), Paragraph(_esc(titre), styles["title"]),
              Paragraph("  ·  ".join(_esc(p) for p in (exporte, sous_titre) if p), styles["meta"])]
    style_par = {"left": styles["cell"], "right": styles["cell_right"], "center": styles["cell_center"]}
    for s in sections:
        n = len(s.entetes)
        aligns = list(s.alignements or ["left"] * n)
        data = [[Paragraph(_esc(h), styles["header"]) for h in s.entetes]]
        for valeurs in [*s.lignes, *([s.total] if s.total else [])]:
            data.append([Paragraph(_esc(v) or "—", style_par.get(aligns[i] if i < len(aligns) else "left"))
                         for i, v in enumerate(valeurs)])
        if len(data) == 1:
            data.append([Paragraph("Aucune donnée", styles["cell"])] + [""] * (n - 1))
        largeurs = list(s.largeurs) if s.largeurs else [1] * n
        largeurs = [w * utile / sum(largeurs) for w in largeurs]
        if s.largeurs and n == 2:
            largeurs = [w * 0.7 for w in largeurs]
        table = Table(data, repeatRows=1, colWidths=largeurs, hAlign="LEFT")
        regles = [
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{_COLOR_NAVY}")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor(f"#{_COLOR_BORDER}")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor(f"#{_COLOR_ZEBRA}")]),
            ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]
        if s.total:
            regles.append(("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#E2E8F0")))
        table.setStyle(TableStyle(regles))
        story += [KeepTogether([Paragraph(_esc(s.titre), styles["bank"]), Spacer(1, 2)]), table, Spacer(1, 6 * mm)]

    def _page(canvas, document):
        _pdf_footer(canvas, document, exported_label=exporte, report_title=titre)

    doc.build(story, onFirstPage=_page, onLaterPages=_page)
    return buf.getvalue()


def _esc(v: Any) -> str:
    if v is None:
        return ""
    return str(v).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

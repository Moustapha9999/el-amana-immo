"""Exports Excel / PDF avec en-tête Banque El Amana et horodatage."""

from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from typing import Any, Sequence
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.data.el_amana_referentiel import BANQUE_EL_AMANA
from app.models import EcritureComptable, Immobilisation

_TZ = ZoneInfo("Africa/Nouakchott")
_EXCEL_MONTANT_FORMAT = "#,##0.00"

# Palette alignée UI (navy / brand)
_COLOR_NAVY = "1E3A5F"
_COLOR_BRAND = "2874A6"
_COLOR_HEADER_FG = "FFFFFF"
_COLOR_ZEBRA = "F8FAFC"
_COLOR_META = "64748B"
_COLOR_TITLE = "0F172A"
_COLOR_BORDER = "CBD5E1"

_THIN = Border(
    left=Side(style="thin", color=_COLOR_BORDER),
    right=Side(style="thin", color=_COLOR_BORDER),
    top=Side(style="thin", color=_COLOR_BORDER),
    bottom=Side(style="thin", color=_COLOR_BORDER),
)


def export_now() -> datetime:
    return datetime.now(_TZ)


def format_export_datetime(when: datetime | None = None) -> str:
    dt = when or export_now()
    return dt.strftime("%d/%m/%Y à %H:%M")


def format_montant(value: Any) -> str:
    """Montant avec exactement 2 décimales (séparateur français)."""
    if value is None:
        return "—"
    try:
        n = float(value)
    except (TypeError, ValueError):
        return str(value)
    formatted = f"{n:,.2f}"
    return formatted.replace(",", "X").replace(".", ",").replace("X", " ")


def _is_montant_number(value: Any) -> bool:
    """True pour les montants numériques (float/Decimal ; int entiers exclus — années, compteurs)."""
    if isinstance(value, bool):
        return False
    return isinstance(value, (float, Decimal))


def format_period_label(date_debut: date | None, date_fin: date | None) -> str | None:
    if date_debut and date_fin:
        return f"Période du {date_debut.strftime('%d/%m/%Y')} au {date_fin.strftime('%d/%m/%Y')}"
    if date_debut:
        return f"À partir du {date_debut.strftime('%d/%m/%Y')}"
    if date_fin:
        return f"Jusqu’au {date_fin.strftime('%d/%m/%Y')}"
    return None


def _bank_line() -> str:
    return (
        f"{BANQUE_EL_AMANA['raison_sociale']} ({BANQUE_EL_AMANA['sigle']}) — "
        f"Code banque {BANQUE_EL_AMANA['code_banque']} — SWIFT {BANQUE_EL_AMANA['code_swift']}"
    )


def _autosize_columns(ws, min_width: int = 10, max_width: int = 42) -> None:
    for col_idx in range(1, ws.max_column + 1):
        letter = get_column_letter(col_idx)
        max_len = min_width
        for cell in ws[letter]:
            if cell.value is None:
                continue
            max_len = max(max_len, min(max_width, len(str(cell.value)) + 2))
        ws.column_dimensions[letter].width = max_len


def build_styled_workbook(
    *,
    sheet_title: str,
    report_title: str,
    headers: Sequence[str],
    rows: Sequence[Sequence[Any]],
    subtitle: str | None = None,
    exported_at: datetime | None = None,
) -> bytes:
    """Classeur Excel avec en-tête banque, titre, date/heure d'export et tableau stylé."""
    when = exported_at or export_now()
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]

    n_cols = len(headers)
    last_col = get_column_letter(n_cols)

    # Ligne 1 — Banque
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n_cols)
    c1 = ws["A1"]
    c1.value = _bank_line()
    c1.font = Font(name="Calibri", size=11, bold=True, color=_COLOR_NAVY)
    c1.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 20

    # Ligne 2 — Titre du rapport
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=n_cols)
    c2 = ws["A2"]
    c2.value = report_title
    c2.font = Font(name="Calibri", size=14, bold=True, color=_COLOR_TITLE)
    c2.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[2].height = 22

    # Ligne 3 — Métadonnées (date/heure + sous-titre)
    meta_parts = [f"Exporté le {format_export_datetime(when)}"]
    if subtitle:
        meta_parts.append(subtitle)
    meta_parts.append(f"{len(rows)} ligne(s)")
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=n_cols)
    c3 = ws["A3"]
    c3.value = "  ·  ".join(meta_parts)
    c3.font = Font(name="Calibri", size=10, italic=True, color=_COLOR_META)
    c3.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[3].height = 18

    # Ligne 4 — vide
    ws.row_dimensions[4].height = 8

    # Ligne 5 — en-têtes colonnes
    header_row = 5
    header_fill = PatternFill("solid", fgColor=_COLOR_NAVY)
    header_font = Font(name="Calibri", size=10, bold=True, color=_COLOR_HEADER_FG)
    for col_idx, label in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=label)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _THIN
    ws.row_dimensions[header_row].height = 22

    # Données
    zebra = PatternFill("solid", fgColor=_COLOR_ZEBRA)
    data_font = Font(name="Calibri", size=10, color=_COLOR_TITLE)
    for r_idx, row in enumerate(rows):
        excel_row = header_row + 1 + r_idx
        for c_idx, value in enumerate(row, start=1):
            cell_value = float(value) if _is_montant_number(value) else value
            cell = ws.cell(row=excel_row, column=c_idx, value=cell_value)
            cell.font = data_font
            cell.border = _THIN
            if _is_montant_number(value):
                cell.number_format = _EXCEL_MONTANT_FORMAT
                cell.alignment = Alignment(horizontal="right", vertical="center", wrap_text=False)
            else:
                cell.alignment = Alignment(vertical="center", wrap_text=False)
            if r_idx % 2 == 1:
                cell.fill = zebra

    # Bandeau décoratif haut + accent
    for col_idx in range(1, n_cols + 1):
        top_cell = ws.cell(row=1, column=col_idx)
        top_cell.fill = PatternFill("solid", fgColor="EEF6FC")
        top_cell.border = Border(bottom=Side(style="medium", color=_COLOR_BRAND))

    ws.freeze_panes = "A6"
    ws.auto_filter.ref = f"A{header_row}:{last_col}{header_row + max(len(rows), 1)}"
    _autosize_columns(ws)

    ws.print_title_rows = f"1:{header_row}"
    ws.oddHeader.center.text = report_title
    ws.oddFooter.center.text = f"Exporté le {format_export_datetime(when)} — Page &P / &N"
    ws.oddFooter.left.text = BANQUE_EL_AMANA["raison_sociale"]

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _pdf_styles():
    from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT

    base = getSampleStyleSheet()
    return {
        "bank": ParagraphStyle(
            "BankLine",
            parent=base["Normal"],
            fontName="Helvetica-Bold",
            fontSize=10,
            textColor=colors.HexColor(f"#{_COLOR_NAVY}"),
            spaceAfter=2,
        ),
        "title": ParagraphStyle(
            "ReportTitle",
            parent=base["Title"],
            fontName="Helvetica-Bold",
            fontSize=14,
            textColor=colors.HexColor(f"#{_COLOR_TITLE}"),
            spaceBefore=4,
            spaceAfter=4,
            alignment=TA_LEFT,
        ),
        "meta": ParagraphStyle(
            "ReportMeta",
            parent=base["Normal"],
            fontName="Helvetica-Oblique",
            fontSize=9,
            textColor=colors.HexColor(f"#{_COLOR_META}"),
            spaceAfter=10,
        ),
        "header": ParagraphStyle(
            "HeaderCell",
            parent=base["Normal"],
            fontName="Helvetica-Bold",
            fontSize=8,
            leading=10,
            textColor=colors.whitesmoke,
            alignment=TA_CENTER,
        ),
        "cell": ParagraphStyle(
            "CellWrap",
            parent=base["Normal"],
            fontName="Helvetica",
            fontSize=7,
            leading=9,
            textColor=colors.HexColor(f"#{_COLOR_TITLE}"),
            alignment=TA_LEFT,
        ),
        "cell_center": ParagraphStyle(
            "CellCenter",
            parent=base["Normal"],
            fontName="Helvetica",
            fontSize=7,
            leading=9,
            textColor=colors.HexColor(f"#{_COLOR_TITLE}"),
            alignment=TA_CENTER,
        ),
        "cell_right": ParagraphStyle(
            "CellRight",
            parent=base["Normal"],
            fontName="Helvetica",
            fontSize=7,
            leading=9,
            textColor=colors.HexColor(f"#{_COLOR_TITLE}"),
            alignment=TA_RIGHT,
        ),
    }


def _pdf_footer(canvas, doc, *, exported_label: str, report_title: str) -> None:
    canvas.saveState()
    page_w, _ = canvas._pagesize
    canvas.setStrokeColor(colors.HexColor(f"#{_COLOR_BRAND}"))
    canvas.setLineWidth(1.2)
    canvas.line(12 * mm, 12 * mm, page_w - 12 * mm, 12 * mm)
    canvas.setFont("Helvetica", 8)
    canvas.setFillColor(colors.HexColor(f"#{_COLOR_META}"))
    canvas.drawString(12 * mm, 7 * mm, BANQUE_EL_AMANA["raison_sociale"])
    canvas.drawCentredString(page_w / 2, 7 * mm, exported_label)
    canvas.drawRightString(page_w - 12 * mm, 7 * mm, f"Page {doc.page}")
    # Bandeau haut
    canvas.setFillColor(colors.HexColor(f"#{_COLOR_NAVY}"))
    canvas.rect(0, canvas._pagesize[1] - 8 * mm, page_w, 8 * mm, fill=1, stroke=0)
    canvas.setFillColor(colors.white)
    canvas.setFont("Helvetica-Bold", 8)
    canvas.drawString(12 * mm, canvas._pagesize[1] - 5.2 * mm, "Banque El Amana — Immobilisations")
    canvas.drawRightString(page_w - 12 * mm, canvas._pagesize[1] - 5.2 * mm, report_title[:48])
    canvas.restoreState()


def build_styled_pdf(
    *,
    report_title: str,
    headers: Sequence[str],
    rows: Sequence[Sequence[Any]],
    subtitle: str | None = None,
    exported_at: datetime | None = None,
    landscape_mode: bool = True,
    col_widths: Sequence[float] | None = None,
    col_aligns: Sequence[str] | None = None,
) -> bytes:
    when = exported_at or export_now()
    exported_label = f"Exporté le {format_export_datetime(when)}"
    styles = _pdf_styles()
    buf = BytesIO()
    pagesize = landscape(A4) if landscape_mode else A4
    left_m = 12 * mm
    right_m = 12 * mm
    doc = SimpleDocTemplate(
        buf,
        pagesize=pagesize,
        title=report_title,
        leftMargin=left_m,
        rightMargin=right_m,
        topMargin=16 * mm,
        bottomMargin=18 * mm,
    )
    usable_width = pagesize[0] - left_m - right_m

    meta_parts = [exported_label]
    if subtitle:
        meta_parts.append(subtitle)
    meta_parts.append(f"{len(rows)} ligne(s)")

    story: list = [
        Paragraph(_bank_line(), styles["bank"]),
        Paragraph(report_title, styles["title"]),
        Paragraph("  ·  ".join(meta_parts), styles["meta"]),
        Spacer(1, 4),
    ]

    n_cols = len(headers)
    aligns = list(col_aligns) if col_aligns else ["left"] * n_cols
    style_map = {
        "left": styles["cell"],
        "center": styles["cell_center"],
        "right": styles["cell_right"],
    }

    def _cell(value: Any, align: str) -> Paragraph:
        if value is None:
            text = ""
        elif _is_montant_number(value):
            text = format_montant(value)
        else:
            text = str(value)
        # Échapper pour ReportLab Paragraph
        text = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        return Paragraph(text or "—", style_map.get(align, styles["cell"]))

    data = [[Paragraph(h, styles["header"]) for h in headers]]
    for row in rows:
        data.append([_cell(v, aligns[i] if i < len(aligns) else "left") for i, v in enumerate(row)])

    if col_widths:
        widths = list(col_widths)
        total = sum(widths)
        if total > 0 and abs(total - usable_width) > 0.5:
            # Normaliser pour occuper toute la largeur utile
            widths = [w * usable_width / total for w in widths]
    else:
        widths = [usable_width / n_cols] * n_cols

    table = Table(data, repeatRows=1, colWidths=widths, hAlign="LEFT")
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{_COLOR_NAVY}")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor(f"#{_COLOR_BORDER}")),
                ("BOX", (0, 0), (-1, -1), 0.6, colors.HexColor(f"#{_COLOR_NAVY}")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor(f"#{_COLOR_ZEBRA}")]),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    story.append(table)

    def _on_page(canvas, document):
        _pdf_footer(canvas, document, exported_label=exported_label, report_title=report_title)

    doc.build(story, onFirstPage=_on_page, onLaterPages=_on_page)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Exports métier
# ---------------------------------------------------------------------------


def ecritures_to_excel(
    rows: list[EcritureComptable],
    *,
    subtitle: str | None = None,
) -> bytes:
    headers = [
        "Date",
        "Journal",
        "Libellé",
        "Débit",
        "Crédit",
        "Montant",
        "Référence",
        "Auto",
        "Validée",
    ]
    data = [
        [
            row.date_ecriture.strftime("%d/%m/%Y"),
            row.journal_code,
            row.libelle,
            row.compte_debit,
            row.compte_credit,
            float(row.montant),
            row.reference or "",
            "Oui" if row.generee_auto else "Non",
            "Oui" if row.validee else "Non",
        ]
        for row in rows
    ]
    return build_styled_workbook(
        sheet_title="Ecritures",
        report_title="Journal des écritures comptables",
        headers=headers,
        rows=data,
        subtitle=subtitle,
    )


def ecritures_to_pdf(
    rows: list[EcritureComptable],
    *,
    titre: str = "Journal des écritures comptables",
    subtitle: str | None = None,
) -> bytes:
    headers = ["Date", "Jnl", "Libellé", "Débit", "Crédit", "Montant", "Référence"]
    data = [
        [
            row.date_ecriture.strftime("%d/%m/%Y"),
            row.journal_code,
            row.libelle or "",
            row.compte_debit,
            row.compte_credit,
            format_montant(row.montant),
            row.reference or "",
        ]
        for row in rows
    ]
    # Proportions (normalisées sur la largeur utile paysage)
    widths = [24 * mm, 14 * mm, 72 * mm, 24 * mm, 24 * mm, 28 * mm, 58 * mm]
    return build_styled_pdf(
        report_title=titre,
        headers=headers,
        rows=data,
        subtitle=subtitle,
        landscape_mode=True,
        col_widths=widths,
        col_aligns=["center", "center", "left", "center", "center", "right", "left"],
    )


def audit_logs_to_excel(rows: list, *, subtitle: str | None = None) -> bytes:
    headers = ["Date", "Utilisateur", "Action", "Entité", "ID entité", "IP"]
    data = []
    for row in rows:
        email = row.user.email if getattr(row, "user", None) else ""
        created = row.created_at
        if created is not None:
            if created.tzinfo is None:
                created = created.replace(tzinfo=_TZ)
            else:
                created = created.astimezone(_TZ)
            date_str = created.strftime("%d/%m/%Y %H:%M")
        else:
            date_str = ""
        data.append(
            [
                date_str,
                email,
                row.action,
                row.entity,
                row.entity_id or "",
                row.ip_address or "",
            ]
        )
    return build_styled_workbook(
        sheet_title="Audit",
        report_title="Journal d’audit",
        headers=headers,
        rows=data,
        subtitle=subtitle,
    )


_STATUT_IMMO_LABELS = {
    "brouillon": "Brouillon",
    "en_cours_acquisition": "En cours d'acquisition",
    "en_service": "En service",
    "suspendue": "Suspendue",
    "cedee": "Cédée",
    "mise_au_rebut": "Mise au rebut",
    "transferee": "Transférée",
    "reclassee": "Reclassée",
    "archivee": "Archivée",
    "cession": "Cession",
    "rebut": "Rebut",
    "en_cours": "En cours",
    "sortie": "Sortie",
}


def _immobilisation_statut_label(statut: Any) -> str:
    code = statut.value if hasattr(statut, "value") else str(statut or "")
    return _STATUT_IMMO_LABELS.get(code, code or "—")


def _immobilisation_export_row(row: Immobilisation) -> list[Any]:
    categorie = getattr(row, "categorie", None)
    famille = getattr(categorie, "famille", None) if categorie is not None else None
    return [
        row.code_inventaire or "",
        row.designation or "",
        famille or "—",
        float(row.valeur_brute),
        _immobilisation_statut_label(row.statut),
        row.compte_immobilisation or "",
        row.date_acquisition.strftime("%d/%m/%Y") if row.date_acquisition else "",
    ]


def immobilisations_to_excel(
    rows: list[Immobilisation],
    *,
    subtitle: str | None = None,
    report_title: str = "Inventaire des immobilisations",
    sheet_title: str = "Immobilisations",
) -> bytes:
    headers = [
        "Code inventaire",
        "Désignation",
        "Type",
        "Valeur brute",
        "Statut",
        "Compte immo",
        "Date acquisition",
    ]
    data = [_immobilisation_export_row(row) for row in rows]
    return build_styled_workbook(
        sheet_title=sheet_title,
        report_title=report_title,
        headers=headers,
        rows=data,
        subtitle=subtitle,
    )


def immobilisations_to_pdf(
    rows: list[Immobilisation],
    *,
    subtitle: str | None = None,
    report_title: str = "Inventaire des immobilisations",
) -> bytes:
    headers = [
        "Code inventaire",
        "Désignation",
        "Type",
        "Valeur brute",
        "Statut",
        "Compte immo",
        "Date acq.",
    ]
    data = [_immobilisation_export_row(row) for row in rows]
    return build_styled_pdf(
        report_title=report_title,
        headers=headers,
        rows=data,
        subtitle=subtitle,
        landscape_mode=True,
        col_aligns=["left", "left", "left", "right", "left", "left", "center"],
    )


def cession_fiche_to_excel(detail: dict[str, Any]) -> bytes:
    """Fiche cession — Excel (libellé / valeur)."""
    rows = [
        ["Référence cession", detail.get("reference") or "—"],
        ["Date de cession", detail.get("date_cession_fmt") or "—"],
        ["Code inventaire", detail.get("code_inventaire") or "—"],
        ["Désignation", detail.get("designation") or "—"],
        ["Date d'acquisition", detail.get("date_acquisition_fmt") or "—"],
        ["Compte immobilisation", detail.get("compte_immobilisation") or "—"],
        ["Valeur brute (VB)", detail.get("valeur_brute")],
        ["VNC à la cession", detail.get("vnc")],
        ["Prix de cession (PC)", detail.get("prix_cession")],
        ["Résultat (PC − VNC)", detail.get("resultat")],
        ["Plus-value", detail.get("plus_value")],
        ["Moins-value", detail.get("moins_value")],
        ["Cas", detail.get("cas_label") or "—"],
        ["Observations", detail.get("observations") or "—"],
    ]
    return build_styled_workbook(
        sheet_title="Fiche cession",
        report_title="Fiche de cession d'immobilisation",
        headers=["Libellé", "Valeur"],
        rows=rows,
        subtitle=detail.get("subtitle"),
    )


def cession_fiche_to_pdf(detail: dict[str, Any]) -> bytes:
    """Fiche cession — PDF."""
    rows = [
        ["Référence cession", detail.get("reference") or "—"],
        ["Date de cession", detail.get("date_cession_fmt") or "—"],
        ["Code inventaire", detail.get("code_inventaire") or "—"],
        ["Désignation", detail.get("designation") or "—"],
        ["Date d'acquisition", detail.get("date_acquisition_fmt") or "—"],
        ["Compte immobilisation", detail.get("compte_immobilisation") or "—"],
        ["Valeur brute (VB)", detail.get("valeur_brute")],
        ["VNC à la cession", detail.get("vnc")],
        ["Prix de cession (PC)", detail.get("prix_cession")],
        ["Résultat (PC − VNC)", detail.get("resultat")],
        ["Plus-value", detail.get("plus_value")],
        ["Moins-value", detail.get("moins_value")],
        ["Cas", detail.get("cas_label") or "—"],
        ["Observations", detail.get("observations") or "—"],
    ]
    return build_styled_pdf(
        report_title="Fiche de cession d'immobilisation",
        headers=["Libellé", "Valeur"],
        rows=rows,
        subtitle=detail.get("subtitle"),
        landscape_mode=False,
        col_widths=None,
        col_aligns=["left", "left"],
    )


def _cession_list_rows(rows: list[tuple[Any, Any]]) -> list[list[Any]]:
    data: list[list[Any]] = []
    for cession, immo in rows:
        plus = cession.plus_value or 0
        moins = cession.moins_value or 0
        if plus and float(plus) > 0:
            resultat = float(plus)
        elif moins and float(moins) > 0:
            resultat = -float(moins)
        else:
            resultat = 0.0
        data.append(
            [
                cession.date_cession.strftime("%d/%m/%Y") if cession.date_cession else "",
                immo.code_inventaire if immo else "",
                cession.reference or "",
                immo.designation if immo else "",
                float(cession.prix_cession or 0),
                float(cession.vnc or 0),
                resultat,
                float(plus or 0),
                float(moins or 0),
            ]
        )
    return data


def cessions_to_excel(
    rows: list[tuple[Any, Any]],
    *,
    subtitle: str | None = None,
) -> bytes:
    headers = [
        "Date",
        "Code inventaire",
        "Référence",
        "Désignation",
        "Prix de cession",
        "VNC",
        "Résultat",
        "Plus-value",
        "Moins-value",
    ]
    return build_styled_workbook(
        sheet_title="Cessions",
        report_title="Liste des cessions",
        headers=headers,
        rows=_cession_list_rows(rows),
        subtitle=subtitle,
    )


def cessions_to_pdf(
    rows: list[tuple[Any, Any]],
    *,
    subtitle: str | None = None,
) -> bytes:
    headers = ["Date", "Code", "Référence", "Désignation", "Prix", "VNC", "Résultat"]
    data = []
    for row in _cession_list_rows(rows):
        data.append(
            [
                row[0],
                row[1],
                row[2],
                row[3],
                format_montant(row[4]),
                format_montant(row[5]),
                format_montant(row[6]),
            ]
        )
    return build_styled_pdf(
        report_title="Liste des cessions",
        headers=headers,
        rows=data,
        subtitle=subtitle,
        landscape_mode=True,
        col_aligns=["center", "left", "left", "left", "right", "right", "right"],
    )


def rebut_fiche_to_excel(detail: dict[str, Any]) -> bytes:
    """Fiche mise au rebut — Excel (libellé / valeur)."""
    rows = [
        ["Date de mise au rebut", detail.get("date_rebut_fmt") or "—"],
        ["Code inventaire", detail.get("code_inventaire") or "—"],
        ["Désignation", detail.get("designation") or "—"],
        ["Date d'acquisition", detail.get("date_acquisition_fmt") or "—"],
        ["Compte immobilisation", detail.get("compte_immobilisation") or "—"],
        ["Valeur brute (VB)", detail.get("valeur_brute")],
        ["Cumul amortissements", detail.get("cumul_amortissement")],
        ["VNC sortie (perte)", detail.get("vnc")],
        ["Motif", detail.get("motif") or "—"],
        ["Statut immobilisation", detail.get("statut_immobilisation") or "—"],
    ]
    return build_styled_workbook(
        sheet_title="Fiche rebut",
        report_title="Fiche de mise au rebut d'immobilisation",
        headers=["Libellé", "Valeur"],
        rows=rows,
        subtitle=detail.get("subtitle"),
    )


def rebut_fiche_to_pdf(detail: dict[str, Any]) -> bytes:
    """Fiche mise au rebut — PDF."""
    rows = [
        ["Date de mise au rebut", detail.get("date_rebut_fmt") or "—"],
        ["Code inventaire", detail.get("code_inventaire") or "—"],
        ["Désignation", detail.get("designation") or "—"],
        ["Date d'acquisition", detail.get("date_acquisition_fmt") or "—"],
        ["Compte immobilisation", detail.get("compte_immobilisation") or "—"],
        ["Valeur brute (VB)", detail.get("valeur_brute")],
        ["Cumul amortissements", detail.get("cumul_amortissement")],
        ["VNC sortie (perte)", detail.get("vnc")],
        ["Motif", detail.get("motif") or "—"],
        ["Statut immobilisation", detail.get("statut_immobilisation") or "—"],
    ]
    return build_styled_pdf(
        report_title="Fiche de mise au rebut d'immobilisation",
        headers=["Libellé", "Valeur"],
        rows=rows,
        subtitle=detail.get("subtitle"),
        landscape_mode=False,
        col_widths=None,
        col_aligns=["left", "left"],
    )


def _rebut_list_rows(rows: list[tuple[Any, Any]]) -> list[list[Any]]:
    data: list[list[Any]] = []
    for rebut, immo in rows:
        data.append(
            [
                rebut.date_rebut.strftime("%d/%m/%Y") if rebut.date_rebut else "",
                immo.code_inventaire if immo else "",
                immo.designation if immo else "",
                float(rebut.vnc or 0),
                rebut.motif or "",
            ]
        )
    return data


def rebuts_to_excel(
    rows: list[tuple[Any, Any]],
    *,
    subtitle: str | None = None,
) -> bytes:
    headers = ["Date", "Code inventaire", "Désignation", "VNC sortie", "Motif"]
    return build_styled_workbook(
        sheet_title="Rebuts",
        report_title="Liste des mises au rebut",
        headers=headers,
        rows=_rebut_list_rows(rows),
        subtitle=subtitle,
    )


def rebuts_to_pdf(
    rows: list[tuple[Any, Any]],
    *,
    subtitle: str | None = None,
) -> bytes:
    headers = ["Date", "Code", "Désignation", "VNC sortie", "Motif"]
    data = []
    for row in _rebut_list_rows(rows):
        data.append([row[0], row[1], row[2], format_montant(row[3]), row[4]])
    return build_styled_pdf(
        report_title="Liste des mises au rebut",
        headers=headers,
        rows=data,
        subtitle=subtitle,
        landscape_mode=True,
        col_aligns=["center", "left", "left", "right", "left"],
    )


def reevaluation_fiche_to_excel(detail: dict[str, Any]) -> bytes:
    """Fiche réévaluation — Excel (libellé / valeur)."""
    rows = [
        ["Date de réévaluation", detail.get("date_reevaluation_fmt") or "—"],
        ["Code inventaire", detail.get("code_inventaire") or "—"],
        ["Désignation", detail.get("designation") or "—"],
        ["Date d'acquisition", detail.get("date_acquisition_fmt") or "—"],
        ["Compte immobilisation", detail.get("compte_immobilisation") or "—"],
        ["Valeur brute actuelle", detail.get("valeur_brute_actuelle")],
        ["Ancienne VBC", detail.get("ancienne_valeur")],
        ["Nouvelle VBC", detail.get("nouvelle_valeur")],
        ["Écart", detail.get("ecart")],
        ["Sens", detail.get("sens_label") or "—"],
        ["Justificatif", detail.get("justificatif") or "—"],
        ["Statut immobilisation", detail.get("statut_immobilisation") or "—"],
    ]
    return build_styled_workbook(
        sheet_title="Fiche réévaluation",
        report_title="Fiche de réévaluation d'immobilisation",
        headers=["Libellé", "Valeur"],
        rows=rows,
        subtitle=detail.get("subtitle"),
    )


def reevaluation_fiche_to_pdf(detail: dict[str, Any]) -> bytes:
    """Fiche réévaluation — PDF."""
    rows = [
        ["Date de réévaluation", detail.get("date_reevaluation_fmt") or "—"],
        ["Code inventaire", detail.get("code_inventaire") or "—"],
        ["Désignation", detail.get("designation") or "—"],
        ["Date d'acquisition", detail.get("date_acquisition_fmt") or "—"],
        ["Compte immobilisation", detail.get("compte_immobilisation") or "—"],
        ["Valeur brute actuelle", detail.get("valeur_brute_actuelle")],
        ["Ancienne VBC", detail.get("ancienne_valeur")],
        ["Nouvelle VBC", detail.get("nouvelle_valeur")],
        ["Écart", detail.get("ecart")],
        ["Sens", detail.get("sens_label") or "—"],
        ["Justificatif", detail.get("justificatif") or "—"],
        ["Statut immobilisation", detail.get("statut_immobilisation") or "—"],
    ]
    return build_styled_pdf(
        report_title="Fiche de réévaluation d'immobilisation",
        headers=["Libellé", "Valeur"],
        rows=rows,
        subtitle=detail.get("subtitle"),
        landscape_mode=False,
        col_widths=None,
        col_aligns=["left", "left"],
    )


def _reevaluation_list_rows(rows: list[tuple[Any, Any]]) -> list[list[Any]]:
    data: list[list[Any]] = []
    for reev, immo in rows:
        ancienne = float(reev.ancienne_valeur or 0)
        nouvelle = float(reev.nouvelle_valeur or 0)
        data.append(
            [
                reev.date_reevaluation.strftime("%d/%m/%Y") if reev.date_reevaluation else "",
                immo.code_inventaire if immo else "",
                immo.designation if immo else "",
                ancienne,
                nouvelle,
                nouvelle - ancienne,
                reev.justificatif or "",
            ]
        )
    return data


def reevaluations_to_excel(
    rows: list[tuple[Any, Any]],
    *,
    subtitle: str | None = None,
) -> bytes:
    headers = [
        "Date",
        "Code inventaire",
        "Désignation",
        "Ancienne valeur",
        "Nouvelle valeur",
        "Écart",
        "Justificatif",
    ]
    return build_styled_workbook(
        sheet_title="Réévaluations",
        report_title="Liste des réévaluations",
        headers=headers,
        rows=_reevaluation_list_rows(rows),
        subtitle=subtitle,
    )


def reevaluations_to_pdf(
    rows: list[tuple[Any, Any]],
    *,
    subtitle: str | None = None,
) -> bytes:
    headers = ["Date", "Code", "Désignation", "Ancienne", "Nouvelle", "Écart", "Justificatif"]
    data = []
    for row in _reevaluation_list_rows(rows):
        data.append(
            [
                row[0],
                row[1],
                row[2],
                format_montant(row[3]),
                format_montant(row[4]),
                format_montant(row[5]),
                row[6],
            ]
        )
    return build_styled_pdf(
        report_title="Liste des réévaluations",
        headers=headers,
        rows=data,
        subtitle=subtitle,
        landscape_mode=True,
        col_aligns=["center", "left", "left", "right", "right", "right", "left"],
    )


def _format_periode_export(periode: str) -> str:
    import re

    q = re.match(r"^(\d{4})-Q([1-4])$", periode.strip(), re.I)
    if q:
        return f"Trimestre {q.group(2)} — {q.group(1)}"
    m = re.match(r"^(\d{4})-(\d{2})$", periode.strip())
    if m:
        labels = {
            "01": "Janvier",
            "02": "Février",
            "03": "Mars",
            "04": "Avril",
            "05": "Mai",
            "06": "Juin",
            "07": "Juillet",
            "08": "Août",
            "09": "Septembre",
            "10": "Octobre",
            "11": "Novembre",
            "12": "Décembre",
        }
        return f"{labels.get(m.group(2), m.group(2))} {m.group(1)}"
    return periode


def amortissement_fiche_to_excel(
    *,
    meta: dict[str, Any],
    lignes: Sequence[Sequence[Any]],
) -> bytes:
    headers = ["Période", "Amort. période", "Cumul amort.", "VNC", "Statut"]
    return build_styled_workbook(
        sheet_title="Amortissements",
        report_title="Fiche d'amortissement",
        headers=headers,
        rows=lignes,
        subtitle=meta.get("subtitle"),
    )


def amortissement_fiche_to_pdf(
    *,
    meta: dict[str, Any],
    lignes: Sequence[Sequence[Any]],
) -> bytes:
    headers = ["Période", "Amort. période", "Cumul amort.", "VNC", "Statut"]
    return build_styled_pdf(
        report_title="Fiche d'amortissement",
        headers=headers,
        rows=lignes,
        subtitle=meta.get("subtitle"),
        landscape_mode=True,
        col_aligns=["left", "right", "right", "right", "center"],
    )


def ecriture_fiche_to_excel(detail: dict[str, Any]) -> bytes:
    rows = [
        ["Date", detail.get("date_ecriture_fmt") or "—"],
        ["Journal", detail.get("journal_code") or "—"],
        ["Libellé", detail.get("libelle") or "—"],
        ["Compte débit", detail.get("compte_debit") or "—"],
        ["Compte crédit", detail.get("compte_credit") or "—"],
        ["Montant", detail.get("montant")],
        ["Référence", detail.get("reference") or "—"],
        ["Origine", detail.get("origine") or "—"],
        ["Validée", detail.get("validee") or "—"],
        ["Type de mouvement", detail.get("type_mouvement") or "—"],
        ["Code inventaire", detail.get("code_inventaire") or "—"],
        ["Désignation", detail.get("designation") or "—"],
    ]
    return build_styled_workbook(
        sheet_title="Fiche écriture",
        report_title="Fiche d'écriture comptable",
        headers=["Libellé", "Valeur"],
        rows=rows,
        subtitle=detail.get("subtitle"),
    )


def ecriture_fiche_to_pdf(detail: dict[str, Any]) -> bytes:
    rows = [
        ["Date", detail.get("date_ecriture_fmt") or "—"],
        ["Journal", detail.get("journal_code") or "—"],
        ["Libellé", detail.get("libelle") or "—"],
        ["Compte débit", detail.get("compte_debit") or "—"],
        ["Compte crédit", detail.get("compte_credit") or "—"],
        ["Montant", detail.get("montant")],
        ["Référence", detail.get("reference") or "—"],
        ["Origine", detail.get("origine") or "—"],
        ["Validée", detail.get("validee") or "—"],
        ["Type de mouvement", detail.get("type_mouvement") or "—"],
        ["Code inventaire", detail.get("code_inventaire") or "—"],
        ["Désignation", detail.get("designation") or "—"],
    ]
    return build_styled_pdf(
        report_title="Fiche d'écriture comptable",
        headers=["Libellé", "Valeur"],
        rows=rows,
        subtitle=detail.get("subtitle"),
        landscape_mode=False,
        col_aligns=["left", "left"],
    )


def recap_amortissement_to_excel(payload: dict[str, Any]) -> bytes:
    """Récapitulatif tableau d'amortissement au 31/12/N."""
    annee = payload.get("annee")
    headers = [
        "Compte",
        "Intitulé",
        f"VB 31/12/{annee}",
        "Compte amort.",
        f"Amorts cumulés {int(annee) - 1}",
        f"Cessions {annee}",
        f"Dotations {annee}",
        f"Amorts cumulés {annee}",
        f"VNC 31/12/{annee}",
    ]
    rows: list[list[Any]] = []
    for line in payload.get("lignes") or []:
        rows.append(
            [
                line.get("compte_immobilisation") or "",
                line.get("intitule") or "",
                line.get("valeur_brute"),
                line.get("compte_amortissement") or "—",
                line.get("amorts_cumules_n1"),
                line.get("cessions_annee"),
                line.get("dotations_annee"),
                line.get("amorts_cumules_n"),
                line.get("vnc"),
            ]
        )
    totaux = payload.get("totaux") or {}
    rows.append(
        [
            "",
            "TOTAL",
            totaux.get("valeur_brute"),
            "",
            totaux.get("amorts_cumules_n1"),
            totaux.get("cessions_annee"),
            totaux.get("dotations_annee"),
            totaux.get("amorts_cumules_n"),
            totaux.get("vnc"),
        ]
    )
    return build_styled_workbook(
        sheet_title="Récap amortissement",
        report_title="Récapitulatif tableau d'amortissement",
        headers=headers,
        rows=rows,
        subtitle=payload.get("subtitle"),
    )


def recap_amortissement_to_pdf(payload: dict[str, Any]) -> bytes:
    annee = payload.get("annee")
    headers = [
        "Compte",
        "Intitulé",
        f"VB {annee}",
        "Cpt amort.",
        f"Cumul {int(annee) - 1}",
        "Cessions",
        "Dotations",
        f"Cumul {annee}",
        f"VNC {annee}",
    ]
    rows: list[list[Any]] = []
    for line in payload.get("lignes") or []:
        rows.append(
            [
                line.get("compte_immobilisation") or "",
                line.get("intitule") or "",
                line.get("valeur_brute"),
                line.get("compte_amortissement") or "—",
                line.get("amorts_cumules_n1"),
                line.get("cessions_annee"),
                line.get("dotations_annee"),
                line.get("amorts_cumules_n"),
                line.get("vnc"),
            ]
        )
    totaux = payload.get("totaux") or {}
    rows.append(
        [
            "",
            "TOTAL",
            totaux.get("valeur_brute"),
            "",
            totaux.get("amorts_cumules_n1"),
            totaux.get("cessions_annee"),
            totaux.get("dotations_annee"),
            totaux.get("amorts_cumules_n"),
            totaux.get("vnc"),
        ]
    )
    return build_styled_pdf(
        report_title="Récapitulatif tableau d'amortissement",
        headers=headers,
        rows=rows,
        subtitle=payload.get("subtitle"),
        landscape_mode=True,
        col_aligns=["left", "left", "right", "left", "right", "right", "right", "right", "right"],
    )


def recap_amortissement_detail_to_excel(payload: dict[str, Any]) -> bytes:
    """Détail des dotations comptabilisées par immobilisation."""
    annee = payload.get("annee")
    headers = [
        "Code",
        "Désignation",
        "Compte",
        f"Dotation comptabilisée {annee}",
        f"Cumul amort. {annee}",
        f"VNC 31/12/{annee}",
    ]
    rows: list[list[Any]] = []
    for line in payload.get("details") or []:
        rows.append(
            [
                line.get("code_inventaire") or "",
                line.get("designation") or "",
                line.get("compte_immobilisation") or "",
                line.get("dotations_annee"),
                line.get("amorts_cumules_n"),
                line.get("vnc"),
            ]
        )
    totaux = payload.get("totaux") or {}
    rows.append(
        [
            "",
            "TOTAL",
            "",
            totaux.get("dotations_annee"),
            totaux.get("amorts_cumules_n"),
            totaux.get("vnc"),
        ]
    )
    return build_styled_workbook(
        sheet_title="Détail dotations",
        report_title="Détail des dotations d'amortissement par immobilisation",
        headers=headers,
        rows=rows,
        subtitle=payload.get("subtitle"),
    )


def recap_amortissement_detail_to_pdf(payload: dict[str, Any]) -> bytes:
    annee = payload.get("annee")
    headers = [
        "Code",
        "Désignation",
        "Compte",
        f"Dotation {annee}",
        f"Cumul {annee}",
        f"VNC {annee}",
    ]
    rows: list[list[Any]] = []
    for line in payload.get("details") or []:
        rows.append(
            [
                line.get("code_inventaire") or "",
                line.get("designation") or "",
                line.get("compte_immobilisation") or "",
                line.get("dotations_annee"),
                line.get("amorts_cumules_n"),
                line.get("vnc"),
            ]
        )
    totaux = payload.get("totaux") or {}
    rows.append(
        [
            "",
            "TOTAL",
            "",
            totaux.get("dotations_annee"),
            totaux.get("amorts_cumules_n"),
            totaux.get("vnc"),
        ]
    )
    return build_styled_pdf(
        report_title="Détail des dotations d'amortissement par immobilisation",
        headers=headers,
        rows=rows,
        subtitle=payload.get("subtitle"),
        landscape_mode=True,
        col_aligns=["left", "left", "left", "right", "right", "right"],
    )


def _comptes_nature_headers(annee: int) -> list[str]:
    return [
        "Compte",
        "Date",
        "Qté",
        "Désignation",
        "Valeur d'acquisition MRU",
        "Taux",
        f"Amt cumulés fin ex. préc. ({annee - 1})",
        f"Dotation {annee}",
        f"Montant amt fin exercice {annee}",
        "Valeur nette comptable",
        "Agence",
    ]


def _comptes_nature_row(line: dict[str, Any], *, compte: str = "") -> list[Any]:
    taux = line.get("taux")
    return [
        compte,
        line.get("date_acquisition_fmt") or "",
        line.get("quantite"),
        line.get("designation") or "",
        line.get("valeur_acquisition"),
        float(taux) if taux is not None else "",
        line.get("amorts_cumules_n1"),
        line.get("dotations_annee"),
        line.get("amorts_cumules_n"),
        line.get("vnc"),
        line.get("agence") or "",
    ]


def comptes_par_nature_to_excel(payload: dict[str, Any]) -> bytes:
    annee = int(payload.get("annee"))
    headers = _comptes_nature_headers(annee)
    rows: list[list[Any]] = []
    for groupe in payload.get("groupes") or []:
        compte = groupe.get("compte_immobilisation") or ""
        intitule = groupe.get("intitule") or ""
        rows.append([f"{compte} — {intitule}", "", "", "", "", "", "", "", "", "", ""])
        for line in groupe.get("lignes") or []:
            rows.append(_comptes_nature_row(line, compte=compte))
        tot = groupe.get("totaux") or {}
        rows.append(
            [
                "",
                "",
                tot.get("quantite"),
                tot.get("designation") or "Total",
                tot.get("valeur_acquisition"),
                "",
                tot.get("amorts_cumules_n1"),
                tot.get("dotations_annee"),
                tot.get("amorts_cumules_n"),
                tot.get("vnc"),
                "",
            ]
        )
        rows.append(["", "", "", "", "", "", "", "", "", "", ""])
    totaux = payload.get("totaux") or {}
    rows.append(
        [
            "",
            "",
            totaux.get("quantite"),
            totaux.get("designation") or "Total général",
            totaux.get("valeur_acquisition"),
            "",
            totaux.get("amorts_cumules_n1"),
            totaux.get("dotations_annee"),
            totaux.get("amorts_cumules_n"),
            totaux.get("vnc"),
            "",
        ]
    )
    return build_styled_workbook(
        sheet_title="Comptes par nature",
        report_title="Comptes d'immobilisation par nature",
        headers=headers,
        rows=rows,
        subtitle=payload.get("subtitle"),
    )


def comptes_par_nature_to_pdf(payload: dict[str, Any]) -> bytes:
    annee = int(payload.get("annee"))
    headers = [
        "Compte",
        "Date",
        "Qté",
        "Désignation",
        "V. acq.",
        "Taux",
        f"Cumul {annee - 1}",
        f"Dot. {annee}",
        f"Amt {annee}",
        "VNC",
        "Agence",
    ]
    rows: list[list[Any]] = []
    for groupe in payload.get("groupes") or []:
        compte = groupe.get("compte_immobilisation") or ""
        for line in groupe.get("lignes") or []:
            rows.append(_comptes_nature_row(line, compte=compte))
        tot = groupe.get("totaux") or {}
        rows.append(
            [
                "",
                "",
                tot.get("quantite"),
                tot.get("designation") or "Total",
                tot.get("valeur_acquisition"),
                "",
                tot.get("amorts_cumules_n1"),
                tot.get("dotations_annee"),
                tot.get("amorts_cumules_n"),
                tot.get("vnc"),
                "",
            ]
        )
    totaux = payload.get("totaux") or {}
    rows.append(
        [
            "",
            "",
            totaux.get("quantite"),
            totaux.get("designation") or "Total général",
            totaux.get("valeur_acquisition"),
            "",
            totaux.get("amorts_cumules_n1"),
            totaux.get("dotations_annee"),
            totaux.get("amorts_cumules_n"),
            totaux.get("vnc"),
            "",
        ]
    )
    return build_styled_pdf(
        report_title="Comptes d'immobilisation par nature",
        headers=headers,
        rows=rows,
        subtitle=payload.get("subtitle"),
        landscape_mode=True,
        col_aligns=[
            "left",
            "center",
            "center",
            "left",
            "right",
            "right",
            "right",
            "right",
            "right",
            "right",
            "left",
        ],
    )


def ventilation_amortissements_agence_to_excel(payload: dict[str, Any]) -> bytes:
    """Ventilation compte 68 — dotations par agence (détail + sous-totaux)."""
    headers = [
        "Agence",
        "Désignation",
        "Code",
        "Date acquisition",
        "Valeur brute",
        "Taux %",
        "Amort. cumulé",
        "Dotation période",
        "VNC",
    ]
    rows: list[list[Any]] = []
    for groupe in payload.get("groupes") or []:
        agence = groupe.get("agence_libelle") or "Sans agence"
        for line in groupe.get("lignes") or []:
            rows.append(
                [
                    agence,
                    line.get("designation") or "",
                    line.get("code_inventaire") or "",
                    line.get("date_acquisition_fmt") or "",
                    line.get("valeur_brute"),
                    line.get("taux"),
                    line.get("amortissement_cumule"),
                    line.get("dotation_periode"),
                    line.get("vnc"),
                ]
            )
        rows.append(
            [
                f"Sous-total 68 — {agence}",
                "",
                "",
                "",
                "",
                "",
                "",
                groupe.get("total_dotations"),
                "",
            ]
        )
    rows.append(
        [
            "TOTAL 68 — Toutes agences",
            "",
            "",
            "",
            "",
            "",
            "",
            payload.get("total_dotations"),
            "",
        ]
    )
    return build_styled_workbook(
        sheet_title="Amort. par agence",
        report_title="Ventilation des amortissements par agence (compte 68)",
        headers=headers,
        rows=rows,
        subtitle=payload.get("subtitle"),
    )


def ventilation_amortissements_agence_to_pdf(payload: dict[str, Any]) -> bytes:
    headers = [
        "Agence",
        "Désignation",
        "Date acq.",
        "VB",
        "Taux",
        "Cumul",
        "Dotation",
        "VNC",
    ]
    rows: list[list[Any]] = []
    for groupe in payload.get("groupes") or []:
        agence = groupe.get("agence_libelle") or "Sans agence"
        for line in groupe.get("lignes") or []:
            rows.append(
                [
                    agence,
                    line.get("designation") or "",
                    line.get("date_acquisition_fmt") or "",
                    line.get("valeur_brute"),
                    line.get("taux"),
                    line.get("amortissement_cumule"),
                    line.get("dotation_periode"),
                    line.get("vnc"),
                ]
            )
        rows.append(
            [
                f"Sous-total — {agence}",
                "",
                "",
                "",
                "",
                "",
                groupe.get("total_dotations"),
                "",
            ]
        )
    rows.append(
        [
            "TOTAL 68",
            "",
            "",
            "",
            "",
            "",
            payload.get("total_dotations"),
            "",
        ]
    )
    return build_styled_pdf(
        report_title="Ventilation des amortissements par agence (compte 68)",
        headers=headers,
        rows=rows,
        subtitle=payload.get("subtitle"),
        landscape_mode=True,
        col_aligns=["left", "left", "left", "right", "right", "right", "right", "right"],
    )


def _soldes_148_68_headers(annee: int) -> list[str]:
    return [
        "Nature",
        "Compte immo",
        "Compte 148",
        "Valeur brute",
        f"Solde 148-{annee - 1}",
        f"Dotation 68 {annee}",
        f"Solde 148 fin {annee}",
        "Compte 68",
        "VNC",
        "Biens",
    ]


def _soldes_148_68_row(line: dict[str, Any]) -> list[Any]:
    return [
        line.get("nature") or "",
        line.get("compte_immobilisation") or "",
        line.get("compte_amortissement") or "—",
        line.get("valeur_brute"),
        line.get("solde_148_n1"),
        line.get("solde_68"),
        line.get("solde_148"),
        line.get("compte_dotation") or "—",
        line.get("vnc"),
        line.get("nb_biens"),
    ]


def soldes_148_68_to_excel(payload: dict[str, Any]) -> bytes:
    """Consultation comptes 142/148/68 — synthèse par nature ou détail."""
    annee = int(payload.get("annee"))
    famille = str(payload.get("famille_compte") or "148")
    numero = str(payload.get("compte_numero") or famille)
    intitule = str(payload.get("compte_intitule") or "")
    report_title = f"Compte {numero} — {intitule}" if intitule else "Soldes 142 / 148 / 68"
    vue = str(payload.get("vue") or "detail")
    totaux = payload.get("totaux") or {}

    if vue == "synthese":
        headers = _soldes_148_68_headers(annee)
        rows: list[list[Any]] = [_soldes_148_68_row(line) for line in payload.get("lignes") or []]
        rows.append(
            [
                "TOTAL",
                "",
                "",
                totaux.get("valeur_brute"),
                totaux.get("solde_148_n1"),
                totaux.get("solde_68"),
                totaux.get("solde_148"),
                "68 global",
                totaux.get("vnc"),
                totaux.get("nb_biens"),
            ]
        )
        return build_styled_workbook(
            sheet_title=f"Soldes {numero}",
            report_title=report_title,
            headers=headers,
            rows=rows,
            subtitle=payload.get("subtitle"),
        )

    headers = [
        "Date",
        "Référence",
        "Désignation",
        "Catégorie",
        "Valeur brute",
        "Dotation",
        "Amort. cumulé",
        "VNC",
        "Agence",
        "Exercice",
    ]
    rows = []
    for line in payload.get("detail") or []:
        rows.append(
            [
                line.get("date") or "",
                line.get("reference") or "",
                line.get("designation") or "",
                line.get("categorie") or "",
                line.get("valeur_brute"),
                line.get("dotation"),
                line.get("amortissement_cumule"),
                line.get("vnc"),
                line.get("agence") or "",
                line.get("exercice"),
            ]
        )
    rows.append(
        [
            "TOTAL",
            "",
            "",
            "",
            totaux.get("valeur_brute"),
            totaux.get("dotation"),
            totaux.get("amortissement_cumule"),
            totaux.get("vnc"),
            "",
            "",
        ]
    )
    return build_styled_workbook(
        sheet_title=f"Détail {numero}",
        report_title=report_title,
        headers=headers,
        rows=rows,
        subtitle=payload.get("subtitle"),
    )


def soldes_148_68_to_pdf(payload: dict[str, Any]) -> bytes:
    """Consultation comptes 142/148/68 — PDF."""
    annee = int(payload.get("annee"))
    famille = str(payload.get("famille_compte") or "148")
    numero = str(payload.get("compte_numero") or famille)
    intitule = str(payload.get("compte_intitule") or "")
    report_title = f"Compte {numero} — {intitule}" if intitule else "Soldes 142 / 148 / 68"
    vue = str(payload.get("vue") or "detail")
    totaux = payload.get("totaux") or {}

    if vue == "synthese":
        headers = [
            "Nature",
            "Cpt immo",
            "Cpt 148",
            "VB",
            f"148-{annee - 1}",
            f"68 {annee}",
            f"148 fin {annee}",
            "Cpt 68",
            "VNC",
            "Biens",
        ]
        rows: list[list[Any]] = [_soldes_148_68_row(line) for line in payload.get("lignes") or []]
        rows.append(
            [
                "TOTAL",
                "",
                "",
                totaux.get("valeur_brute"),
                totaux.get("solde_148_n1"),
                totaux.get("solde_68"),
                totaux.get("solde_148"),
                "68 global",
                totaux.get("vnc"),
                totaux.get("nb_biens"),
            ]
        )
        return build_styled_pdf(
            report_title=report_title,
            headers=headers,
            rows=rows,
            subtitle=payload.get("subtitle"),
            landscape_mode=True,
            col_aligns=[
                "left",
                "left",
                "left",
                "right",
                "right",
                "right",
                "right",
                "left",
                "right",
                "center",
            ],
        )

    headers = [
        "Date",
        "Réf.",
        "Désignation",
        "Catégorie",
        "VB",
        "Dotation",
        "Amt cumulé",
        "VNC",
        "Agence",
        "Exercice",
    ]
    rows = []
    for line in payload.get("detail") or []:
        rows.append(
            [
                line.get("date") or "",
                line.get("reference") or "",
                line.get("designation") or "",
                line.get("categorie") or "",
                line.get("valeur_brute"),
                line.get("dotation"),
                line.get("amortissement_cumule"),
                line.get("vnc"),
                line.get("agence") or "",
                line.get("exercice"),
            ]
        )
    rows.append(
        [
            "TOTAL",
            "",
            "",
            "",
            totaux.get("valeur_brute"),
            totaux.get("dotation"),
            totaux.get("amortissement_cumule"),
            totaux.get("vnc"),
            "",
            "",
        ]
    )
    return build_styled_pdf(
        report_title=report_title,
        headers=headers,
        rows=rows,
        subtitle=payload.get("subtitle"),
        landscape_mode=True,
        col_aligns=[
            "center",
            "left",
            "left",
            "left",
            "right",
            "right",
            "right",
            "right",
            "left",
            "center",
        ],
    )

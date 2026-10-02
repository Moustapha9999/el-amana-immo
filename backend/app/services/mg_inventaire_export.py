"""Exports d'un inventaire mensuel : Excel (Synthèse / Détail / Écarts) et PDF avec bloc de validation."""

from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import mm
from reportlab.platypus import (
    Image,
    KeepTogether,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from app.services.reporting_export import (
    _COLOR_BORDER,
    _COLOR_BRAND,
    _COLOR_META,
    _COLOR_NAVY,
    _COLOR_TITLE,
    _COLOR_ZEBRA,
    _THIN,
    _bank_line,
    _pdf_footer,
    _pdf_styles,
    export_now,
    format_export_datetime,
    resolve_bea_logo_path,
)

STATUT_LABELS = {
    "BROUILLON": "Brouillon",
    "EN_COURS": "En cours",
    "A_CONTROLER": "À contrôler",
    "VALIDE": "Validé",
    "AJUSTE": "Ajusté",
    "ARCHIVE": "Archivé",
    "ANNULE": "Annulé",
}
LIGNE_LABELS = {
    "NON_COMPTE": "Non compté",
    "CONFORME": "Conforme",
    "ECART_NEGATIF": "Écart négatif",
    "ECART_POSITIF": "Écart positif",
    "EXCLU": "Exclu",
}
DETAIL_HEADERS = [
    "Code", "Désignation", "Famille", "Unité", "Stock théorique", "Stock physique",
    "Écart", "Écart %", "Statut", "Compté par", "Compté le", "Observation",
]
ECART_HEADERS = ["Code", "Désignation", "Famille", "Théorique", "Physique", "Écart", "Écart %", "Sens", "Observation"]


def _num(v: Any) -> int | float | None:
    if v is None:
        return None
    if isinstance(v, Decimal):
        return int(v) if v == v.to_integral_value() else float(v)
    return v


def _d(v: date | datetime | None, with_time: bool = False) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%d/%m/%Y %H:%M" if with_time else "%d/%m/%Y")
    return v.strftime("%d/%m/%Y")


def _pct(v: float | None) -> str:
    if v is None:
        return "n/d"
    return f"{v:+.2f} %".replace(".", ",")


def synthese_rows(inv: dict) -> list[tuple[str, Any]]:
    s = inv["stats"]
    return [
        ("Référence", inv["reference"]),
        ("Libellé", inv.get("libelle") or ""),
        ("Période", inv.get("periode_libelle") or ""),
        ("Date d'inventaire", _d(inv.get("date_debut"))),
        ("Agence", inv.get("agence_libelle") or "Toutes agences"),
        ("Famille", inv.get("famille_libelle") or "Toutes familles"),
        ("Responsable", inv.get("responsable_nom") or ""),
        ("Statut", STATUT_LABELS.get(inv["statut"], inv["statut"])),
        ("Source", "Import Excel" if inv.get("source") == "IMPORT_EXCEL" else "Saisie"),
        ("Photo du stock théorique", _d(inv.get("snapshot_at"), True)),
        ("Articles inventoriés", s["total"]),
        ("Articles à compter", s["a_compter"]),
        ("Articles comptés", s["comptes"]),
        ("Articles non comptés", s["non_comptes"]),
        ("Articles exclus", s["exclus"]),
        ("Progression", f"{s['progression']:.1f} %".replace(".", ",")),
        ("Articles sans écart", s["sans_ecart"]),
        ("Écarts négatifs (manquants)", s["ecarts_negatifs"]),
        ("Écarts positifs (excédents)", s["ecarts_positifs"]),
        ("Total théorique (comptés)", _num(s["total_theorique"])),
        ("Total physique (comptés)", _num(s["total_physique"])),
        ("Écart net", _num(s["ecart_net"])),
        ("Validé le", _d(inv.get("valide_at"), True)),
        ("Validé par", inv.get("valide_by_nom") or ""),
        ("Validation forcée", "Oui" if inv.get("validation_forcee") else "Non"),
        ("Ajustements générés le", _d(inv.get("ajustements_at"), True)),
        ("Observations", inv.get("observation") or ""),
    ] + _rapprochement_rows(inv)


def _rapprochement_rows(inv: dict) -> list[tuple[str, Any]]:
    s = inv["stats"]
    if not s.get("rapprochement"):
        return []
    meta = (inv.get("import_meta") or {}).get("rapprochement") or {}
    return [
        ("Référence officielle", meta.get("source_officielle") or ""),
        ("Plan de rapprochement", meta.get("fichier") or ""),
        ("Comptage de référence le", meta.get("date_comptage") or ""),
        ("Stock système avant rapprochement", _num(s["total_systeme"])),
        ("Ajustements de stock", _num(s["ajustement_net"])),
        ("Stock actuel retenu (référence)", _num(s["total_retenu"])),
        ("Stock physique compté", _num(s["total_physique"])),
        ("Écart physique restant à régulariser", _num(s["ecart_a_regulariser"])),
    ]


def _statut(lg: dict) -> str:
    if lg.get("ancienne_agence"):
        return "Ancienne agence"
    label = LIGNE_LABELS.get(lg["statut_ligne"], lg["statut_ligne"])
    return f"{label} — à régulariser" if lg.get("a_regulariser") else label


def detail_row(lg: dict) -> list[Any]:
    return [
        lg["article_code"],
        lg["article_designation"],
        lg.get("famille_libelle") or "",
        lg.get("unite") or "",
        _num(lg.get("theorique_reference", lg["stock_theorique"])),
        _num(lg["stock_physique"]),
        _num(lg["ecart"]),
        _pct(lg["ecart_pourcentage"]) if lg["ecart"] is not None else "",
        _statut(lg),
        lg.get("compte_par_nom") or "",
        _d(lg.get("compte_at"), True),
        lg.get("observation") or "",
    ]


def ecart_row(lg: dict) -> list[Any]:
    return [
        lg["article_code"],
        lg["article_designation"],
        lg.get("famille_libelle") or "",
        _num(lg.get("theorique_reference", lg["stock_theorique"])),
        _num(lg["stock_physique"]),
        _num(lg["ecart"]),
        _pct(lg["ecart_pourcentage"]),
        ("Manquant" if lg["ecart"] < 0 else "Excédent") + (" — à régulariser" if lg.get("a_regulariser") else ""),
        lg.get("observation") or "",
    ]


def lignes_ecart(lignes: list[dict]) -> list[dict]:
    return sorted(
        (lg for lg in lignes if lg["statut_ligne"] in {"ECART_NEGATIF", "ECART_POSITIF"}),
        key=lambda lg: (-abs(lg["ecart"]), lg["article_code"] or ""),
    )


def _titre(inv: dict) -> str:
    return f"Inventaire {inv['reference']} — {inv.get('periode_libelle') or _d(inv.get('date_debut'))}"


def _entete_feuille(ws, titre: str, n_cols: int, when: datetime) -> int:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n_cols)
    ws["A1"].value = _bank_line()
    ws["A1"].font = Font(name="Calibri", size=11, bold=True, color=_COLOR_NAVY)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=n_cols)
    ws["A2"].value = titre
    ws["A2"].font = Font(name="Calibri", size=14, bold=True, color=_COLOR_TITLE)
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=n_cols)
    ws["A3"].value = f"Exporté le {format_export_datetime(when)}"
    ws["A3"].font = Font(name="Calibri", size=10, italic=True, color=_COLOR_META)
    for c in range(1, n_cols + 1):
        ws.cell(row=1, column=c).fill = PatternFill("solid", fgColor="EEF6FC")
    return 5


def _tableau(ws, start: int, headers: list[str], rows: list[list[Any]], widths: list[int]) -> None:
    head_fill = PatternFill("solid", fgColor=_COLOR_NAVY)
    zebra = PatternFill("solid", fgColor=_COLOR_ZEBRA)
    neg = Font(name="Calibri", size=10, color="B91C1C", bold=True)
    pos = Font(name="Calibri", size=10, color="B45309", bold=True)
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=start, column=c, value=h)
        cell.fill = head_fill
        cell.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _THIN
    ecart_col = headers.index("Écart") + 1 if "Écart" in headers else None
    for r, row in enumerate(rows, start=1):
        for c, v in enumerate(row, start=1):
            cell = ws.cell(row=start + r, column=c, value=v)
            cell.border = _THIN
            cell.font = Font(name="Calibri", size=10, color=_COLOR_TITLE)
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                cell.alignment = Alignment(horizontal="right")
                cell.number_format = "#,##0;-#,##0;0"
            if c == ecart_col and isinstance(v, (int, float)) and v:
                cell.font = neg if v < 0 else pos
            if r % 2 == 0:
                cell.fill = zebra
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=start + 1, column=1)
    if rows:
        ws.auto_filter.ref = f"A{start}:{get_column_letter(len(headers))}{start + len(rows)}"


def inventaire_to_excel(inv: dict, lignes: list[dict]) -> bytes:
    when = export_now()
    titre = _titre(inv)
    wb = Workbook()

    ws = wb.active
    ws.title = "Synthèse"
    row = _entete_feuille(ws, titre, 2, when)
    for label, value in synthese_rows(inv):
        a = ws.cell(row=row, column=1, value=label)
        a.font = Font(name="Calibri", size=10, bold=True, color=_COLOR_NAVY)
        a.border = _THIN
        b = ws.cell(row=row, column=2, value=value)
        b.border = _THIN
        b.alignment = Alignment(horizontal="left", wrap_text=True)
        row += 1
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 60

    ws = wb.create_sheet("Détail")
    start = _entete_feuille(ws, f"{titre} — détail", len(DETAIL_HEADERS), when)
    _tableau(ws, start, DETAIL_HEADERS, [detail_row(lg) for lg in lignes],
             [11, 42, 22, 9, 14, 14, 10, 10, 14, 22, 17, 36])

    ws = wb.create_sheet("Écarts")
    start = _entete_feuille(ws, f"{titre} — écarts", len(ECART_HEADERS), when)
    _tableau(ws, start, ECART_HEADERS, [ecart_row(lg) for lg in lignes_ecart(lignes)],
             [11, 42, 22, 12, 12, 10, 10, 11, 36])

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _esc(v: Any) -> str:
    text = "" if v is None else str(v)
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def inventaire_to_pdf(inv: dict, lignes: list[dict]) -> bytes:
    when = export_now()
    exported_label = f"Exporté le {format_export_datetime(when)}"
    titre = _titre(inv)
    styles = _pdf_styles()
    buf = BytesIO()
    pagesize = landscape(A4)
    doc = SimpleDocTemplate(
        buf, pagesize=pagesize, title=titre,
        leftMargin=12 * mm, rightMargin=12 * mm, topMargin=16 * mm, bottomMargin=18 * mm,
    )
    largeur = pagesize[0] - 24 * mm
    grille = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{_COLOR_NAVY}")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor(f"#{_COLOR_BORDER}")),
        ("BOX", (0, 0), (-1, -1), 0.6, colors.HexColor(f"#{_COLOR_NAVY}")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor(f"#{_COLOR_ZEBRA}")]),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]

    def section(texte: str) -> Paragraph:
        return Paragraph(f"<font color='#{_COLOR_BRAND}'>{_esc(texte)}</font>", styles["title"])

    def tableau(headers: list[str], rows: list[list[Any]], widths: list[float], aligns: list[str]) -> Table:
        style_map = {"l": styles["cell"], "c": styles["cell_center"], "r": styles["cell_right"]}
        data = [[Paragraph(_esc(h), styles["header"]) for h in headers]]
        for row in rows:
            data.append([
                Paragraph(_esc(v) if v not in (None, "") else "—", style_map[aligns[i]]) for i, v in enumerate(row)
            ])
        total = sum(widths)
        t = Table(data, repeatRows=1, colWidths=[w * largeur / total for w in widths], hAlign="LEFT")
        t.setStyle(TableStyle(grille))
        return t

    story: list = []
    logo = resolve_bea_logo_path()
    if logo is not None:
        try:
            story += [Image(str(logo), width=42 * mm, height=14 * mm), Spacer(1, 2 * mm)]
        except Exception:
            pass
    story += [
        Paragraph(_esc(_bank_line()), styles["bank"]),
        Paragraph(_esc(titre), styles["title"]),
        Paragraph(_esc(exported_label), styles["meta"]),
        section("1. Synthèse"),
    ]
    synth = synthese_rows(inv)
    moitie = (len(synth) + 1) // 2
    gauche, droite = synth[:moitie], synth[moitie:]
    data = []
    for i in range(moitie):
        lg = gauche[i]
        rd = droite[i] if i < len(droite) else ("", "")
        data.append([
            Paragraph(f"<b>{_esc(lg[0])}</b>", styles["cell"]), Paragraph(_esc(lg[1]), styles["cell"]),
            Paragraph(f"<b>{_esc(rd[0])}</b>", styles["cell"]), Paragraph(_esc(rd[1]), styles["cell"]),
        ])
    t = Table(data, colWidths=[largeur * 0.2, largeur * 0.3, largeur * 0.2, largeur * 0.3], hAlign="LEFT")
    t.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor(f"#{_COLOR_BORDER}")),
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#EEF6FC")),
        ("BACKGROUND", (2, 0), (2, -1), colors.HexColor("#EEF6FC")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))
    story += [t, PageBreak(), section("2. Détail des articles")]
    story.append(tableau(
        DETAIL_HEADERS[:9] + ["Observation"],
        [detail_row(lg)[:9] + [detail_row(lg)[11]] for lg in lignes],
        [10, 34, 16, 7, 10, 10, 8, 9, 11, 25],
        ["l", "l", "l", "c", "r", "r", "r", "r", "c", "l"],
    ))
    ecarts = lignes_ecart(lignes)
    story += [PageBreak(), section("3. Écarts")]
    if ecarts:
        story.append(tableau(
            ECART_HEADERS, [ecart_row(lg) for lg in ecarts],
            [10, 34, 16, 10, 10, 8, 9, 10, 25],
            ["l", "l", "l", "r", "r", "r", "r", "c", "l"],
        ))
    else:
        story.append(Paragraph("Aucun écart constaté sur les articles comptés.", styles["cell"]))

    signatures = [
        ["Établi par (responsable inventaire)", "Contrôlé par", "Validé par"],
        [
            inv.get("responsable_nom") or "",
            "",
            inv.get("valide_by_nom") or "",
        ],
        [
            "Date : " + _d(inv.get("date_debut")),
            "Date :",
            "Date : " + _d(inv.get("valide_at")),
        ],
        ["Signature :\n\n\n", "Signature :\n\n\n", "Signature :\n\n\n"],
    ]
    sig = Table(
        [[Paragraph(_esc(c).replace("\n", "<br/>"), styles["header" if r == 0 else "cell"]) for c in row]
         for r, row in enumerate(signatures)],
        colWidths=[largeur / 3] * 3,
        rowHeights=[None, None, None, 28 * mm],
    )
    sig.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{_COLOR_NAVY}")),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor(f"#{_COLOR_BORDER}")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    mention = (
        "Validation forcée avec des articles non comptés." if inv.get("validation_forcee") else
        "Les écarts ci-dessus sont constatés contradictoirement ; les ajustements de stock sont générés "
        "après validation."
    )
    story += [Spacer(1, 6 * mm), KeepTogether([
        section("4. Validation"),
        Paragraph(_esc(mention), styles["meta"]),
        sig,
    ])]

    def _on_page(canvas, document):
        _pdf_footer(canvas, document, exported_label=exported_label, report_title=titre)

    doc.build(story, onFirstPage=_on_page, onLaterPages=_on_page)
    return buf.getvalue()

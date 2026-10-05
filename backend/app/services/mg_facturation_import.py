"""Import Excel des références fournisseurs (ex. SOMELEC) → points de facturation.

Excel → lecture / mapping → prévisualisation → contrôle doublons → confirmation → points.
Aucun montant n'est importé : les montants viennent des factures réelles.
"""

from __future__ import annotations

import difflib
import io
import re
import uuid
from dataclasses import dataclass, field
from datetime import date

from fastapi import UploadFile
from sqlalchemy import select

from app.core.exceptions import AppError
from app.models.auth import Agence, User
from app.models.mg_ops import MgPointFacturation
from app.schemas.mg_facturation import ImportChoix, PointFacturationCreate
from app.services.mg_facturation_service import (
    PERIODICITES_POINT,
    TYPES_POINT,
    MgFacturationService,
    decouper_reference,
    mois_suivant,
    normaliser_reference,
    sans_accents,
)

MAX_IMPORT_BYTES = 5 * 1024 * 1024
RE_LABEL_REF = re.compile(r"^\s*r[eé]f[eé]?r[ae]n?ces?\b", re.IGNORECASE)
RE_LABEL_MONTANT = re.compile(r"montant", re.IGNORECASE)
MOTS_VIDES = {"agence", "bureau", "siege", "de", "des", "du", "la", "le", "les", "l", "d", "et", "pdv", "pvd", "amanty", "bea"}


@dataclass
class LigneImport:
    reference: str
    normalisee: str
    compteur: str | None
    nom_source: str
    nom: str
    type_point: str
    feuille: str
    cellule: str
    statut: str = "NOUVEAU"
    agence_id: str | None = None
    agence_libelle: str | None = None
    score: float | None = None
    point_existant: dict | None = None

    def as_dict(self) -> dict:
        return {
            "reference": self.reference,
            "normalisee": self.normalisee,
            "compteur": self.compteur,
            "nom_source": self.nom_source,
            "nom": self.nom,
            "type_point": self.type_point,
            "feuille": self.feuille,
            "cellule": self.cellule,
            "statut": self.statut,
            "agence_id": self.agence_id,
            "agence_libelle": self.agence_libelle,
            "score": self.score,
            "point_existant": self.point_existant,
        }


@dataclass
class Analyse:
    lignes: list[LigneImport] = field(default_factory=list)
    feuilles: list[str] = field(default_factory=list)
    montants_detectes: int = 0
    avertissements: list[str] = field(default_factory=list)


def _texte(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r"\s+", " ", str(value)).strip()


def _cellule(row: int, col: int) -> str:
    lettres = ""
    n = col + 1
    while n:
        n, r = divmod(n - 1, 26)
        lettres = chr(65 + r) + lettres
    return f"{lettres}{row + 1}"


def ressemble_reference(value) -> bool:
    """Référence compteur / abonnement : au moins 9 chiffres, éventuellement espacés ou entre parenthèses."""
    if value is None or isinstance(value, bool):
        return False
    if isinstance(value, (int, float)):
        return len(str(int(value))) >= 9
    text = _texte(value)
    if not text or re.search(r"[a-zA-Z]{4,}", text.split("(")[0] if "(" in text else text):
        return False
    _, digits, _ = decouper_reference(text)
    return len(digits) >= 9


def deduire_type(nom: str, contexte: str = "") -> str:
    n = sans_accents(nom)
    c = sans_accents(contexte)
    if "siege" in n:
        return "SIEGE"
    if re.search(r"\bp[dv][dv]\b", n) or re.search(r"\bp[dv][dv]\b", c):
        return "PDV"
    if "agence" in n or "bureau" in n:
        return "AGENCE"
    return "AUTRE"


def nom_point(nom: str, type_point: str) -> str:
    """« PVD CARREFOUR 24 » / « SKY RIM » (section PDV AMANTY) → « PDV Amanty Carrefour 24 »."""
    propre = _texte(nom)
    if type_point != "PDV":
        return propre[:1].upper() + propre[1:] if propre else propre
    reste = re.sub(r"^\s*(p[dv][dv]\s*)?(amanty\s*)?", "", propre, flags=re.IGNORECASE).strip(" -")
    return f"PDV Amanty {reste.title()}".strip() if reste else "PDV Amanty"


def _tokens(text: str) -> set[str]:
    return {t for t in re.split(r"[^a-z0-9]+", sans_accents(text)) if t and t not in MOTS_VIDES}


def suggerer_agence(nom: str, agences: list[tuple[str, str]]) -> tuple[str | None, str | None, float | None]:
    """Meilleure agence (id, libellé, score 0–1) par recouvrement de mots + similarité."""
    cible = _tokens(nom)
    if not cible:
        return None, None, None
    meilleur: tuple[str | None, str | None, float] = (None, None, 0.0)
    for aid, libelle in agences:
        mots = _tokens(libelle)
        if not mots:
            continue
        recouvrement = len(cible & mots) / max(len(cible), len(mots))
        ratio = difflib.SequenceMatcher(None, " ".join(sorted(cible)), " ".join(sorted(mots))).ratio()
        score = max(recouvrement, ratio * 0.9)
        if score > meilleur[2]:
            meilleur = (aid, libelle, score)
    if meilleur[2] < 0.6:
        return None, None, round(meilleur[2], 2) if meilleur[2] else None
    return meilleur[0], meilleur[1], round(meilleur[2], 2)


def analyser_classeur(content: bytes) -> Analyse:
    """Extrait les couples (nom de site, référence) de toutes les feuilles."""
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:  # fichier corrompu ou non Excel
        raise AppError("Fichier Excel illisible (.xlsx attendu)", code="IMPORT_FICHIER_INVALIDE") from exc

    analyse = Analyse()
    for ws in wb.worksheets:
        grille = [list(r) for r in ws.iter_rows(values_only=True)]
        if not grille:
            continue
        analyse.feuilles.append(ws.title)
        utilisees: set[tuple[int, int]] = set()
        largeur = max(len(r) for r in grille)
        for r in grille:
            r.extend([None] * (largeur - len(r)))

        def valeur(i: int, j: int):
            return grille[i][j] if 0 <= i < len(grille) and 0 <= j < largeur else None

        # 1) Disposition « ligne » : libellé Référence, noms sur la ligne d'en-tête au-dessus.
        for i, row in enumerate(grille):
            for j, cell in enumerate(row):
                texte = _texte(cell)
                if texte and RE_LABEL_MONTANT.search(texte) and not ressemble_reference(cell):
                    analyse.montants_detectes += sum(
                        1 for v in row[j + 1:] if isinstance(v, (int, float)) and not isinstance(v, bool) and v
                    )
                    utilisees.update((i, k) for k in range(j, largeur))
                    break
                if not texte or not RE_LABEL_REF.match(texte):
                    continue
                for k in range(j + 1, largeur):
                    ref = row[k]
                    if not ressemble_reference(ref):
                        continue
                    nom = ""
                    for up in range(i - 1, max(-1, i - 4), -1):
                        candidat = _texte(valeur(up, k))
                        if candidat and not ressemble_reference(valeur(up, k)):
                            nom = candidat
                            utilisees.add((up, k))
                            break
                    _ajouter(analyse, ws.title, i, k, nom, ref, "")
                    utilisees.add((i, k))
                utilisees.add((i, j))
                break

        # 2) Disposition « paires » : nom à gauche de la référence, section éventuelle au-dessus.
        contexte = ""
        for i, row in enumerate(grille):
            textes = [(j, _texte(c)) for j, c in enumerate(row) if _texte(c) and (i, j) not in utilisees]
            refs = [j for j, c in enumerate(row) if (i, j) not in utilisees and ressemble_reference(c)]
            if textes and not refs and len(textes) == 1:
                contexte = textes[0][1]
                continue
            for k in refs:
                nom = ""
                for left in range(k - 1, -1, -1):
                    candidat = _texte(row[left])
                    if candidat and not ressemble_reference(row[left]):
                        nom = candidat
                        break
                if not nom:
                    for up in range(i - 1, max(-1, i - 3), -1):
                        candidat = _texte(valeur(up, k))
                        if candidat and not ressemble_reference(valeur(up, k)):
                            nom = candidat
                            break
                _ajouter(analyse, ws.title, i, k, nom, row[k], contexte)
                utilisees.add((i, k))
    wb.close()

    vues: dict[str, LigneImport] = {}
    for ligne in analyse.lignes:
        premier = vues.get(ligne.normalisee)
        if premier is None:
            vues[ligne.normalisee] = ligne
        else:
            ligne.statut = "DOUBLON_FICHIER"
            ligne.point_existant = {"code": None, "nom": premier.nom, "cellule": f"{premier.feuille}!{premier.cellule}"}
    if not analyse.lignes:
        analyse.avertissements.append("Aucune référence fournisseur détectée dans le fichier.")
    return analyse


def _ajouter(analyse: Analyse, feuille: str, i: int, k: int, nom: str, ref, contexte: str) -> None:
    affichee, _digits, compteur = decouper_reference(ref)
    normalisee = normaliser_reference(ref)
    if not normalisee:
        return
    type_point = deduire_type(nom, contexte)
    analyse.lignes.append(
        LigneImport(
            reference=affichee,
            normalisee=normalisee,
            compteur=compteur,
            nom_source=nom or "(sans nom)",
            nom=nom_point(nom, type_point) if nom else f"Point {normalisee}",
            type_point=type_point,
            feuille=feuille,
            cellule=_cellule(i, k),
        )
    )


class MgFacturationImport:
    def __init__(self, svc: MgFacturationService):
        self.svc = svc
        self.db = svc.db

    async def _lire(self, file: UploadFile) -> bytes:
        name = (file.filename or "").lower()
        if not name.endswith((".xlsx", ".xlsm")):
            raise AppError("Format .xlsx attendu", code="IMPORT_FICHIER_INVALIDE")
        content = await file.read()
        if not content:
            raise AppError("Fichier vide", code="FICHIER_VIDE")
        if len(content) > MAX_IMPORT_BYTES:
            raise AppError("Fichier trop volumineux (max 5 Mo)", status_code=413, code="FICHIER_TROP_VOLUMINEUX")
        return content

    async def _analyser(self, file: UploadFile, fournisseur_id: uuid.UUID) -> tuple[Analyse, str]:
        fournisseur = await self.svc._fournisseur(fournisseur_id)
        analyse = analyser_classeur(await self._lire(file))
        ag_stmt = select(Agence.id, Agence.libelle).where(Agence.deleted_at.is_(None), Agence.is_active.is_(True))
        if self.svc.scope_agence:
            ag_stmt = ag_stmt.where(Agence.id == self.svc.scope_agence)
        agences = [(str(a), l) for a, l in (await self.db.execute(ag_stmt)).all()]
        existants = {
            p.reference_normalisee: p
            for p in (
                await self.db.execute(
                    select(MgPointFacturation).where(
                        MgPointFacturation.fournisseur_id == fournisseur.id,
                        MgPointFacturation.deleted_at.is_(None),
                    )
                )
            ).scalars().all()
        }
        for ligne in analyse.lignes:
            if ligne.type_point in {"AGENCE", "SIEGE"}:
                ligne.agence_id, ligne.agence_libelle, ligne.score = suggerer_agence(ligne.nom_source, agences)
            existant = existants.get(ligne.normalisee)
            if existant and ligne.statut == "NOUVEAU":
                ligne.statut = "EXISTANT"
                ligne.point_existant = {"code": existant.code, "nom": existant.nom, "id": str(existant.id)}
        return analyse, fournisseur.raison_sociale

    def _resume(self, analyse: Analyse) -> dict:
        uniques = [l for l in analyse.lignes if l.statut != "DOUBLON_FICHIER"]
        return {
            "total_detectees": len(analyse.lignes),
            "references_uniques": len(uniques),
            "agences": sum(1 for l in uniques if l.type_point == "AGENCE"),
            "sieges": sum(1 for l in uniques if l.type_point == "SIEGE"),
            "pdv": sum(1 for l in uniques if l.type_point == "PDV"),
            "autres": sum(1 for l in uniques if l.type_point == "AUTRE"),
            "nouveaux": sum(1 for l in uniques if l.statut == "NOUVEAU"),
            "existants": sum(1 for l in uniques if l.statut == "EXISTANT"),
            "doublons_fichier": sum(1 for l in analyse.lignes if l.statut == "DOUBLON_FICHIER"),
            "agences_rapprochees": sum(1 for l in uniques if l.agence_id),
            "montants_detectes": analyse.montants_detectes,
            "montants_importes": 0,
        }

    async def preview(self, file: UploadFile, fournisseur_id: uuid.UUID) -> dict:
        analyse, fournisseur = await self._analyser(file, fournisseur_id)
        avertissements = list(analyse.avertissements)
        if analyse.montants_detectes:
            avertissements.append(
                f"{analyse.montants_detectes} montant(s) détecté(s) : non importés. Les montants proviennent "
                "uniquement des factures réelles saisies ensuite."
            )
        else:
            avertissements.append("Aucun montant dans le fichier : seuls les points de facturation seront créés.")
        return {
            "fichier": file.filename,
            "fournisseur_id": str(fournisseur_id),
            "fournisseur": fournisseur,
            "feuilles": analyse.feuilles,
            "lignes": [l.as_dict() for l in analyse.lignes],
            "resume": self._resume(analyse),
            "avertissements": avertissements,
        }

    async def confirm(
        self,
        file: UploadFile,
        fournisseur_id: uuid.UUID,
        user: User,
        *,
        choix: list[ImportChoix],
        suivi_depuis: date | None,
        periodicite: str,
        type_facture: str | None,
    ) -> dict:
        periodicite = (periodicite or "MENSUEL").upper()
        if periodicite not in PERIODICITES_POINT:
            raise AppError("Périodicité inconnue", code="PERIODICITE_INVALIDE")
        if suivi_depuis is None:
            today = date.today()
            y, m = mois_suivant(today.year, today.month, -1)
            suivi_depuis = date(y, m, 1)
        analyse, fournisseur = await self._analyser(file, fournisseur_id)
        par_ref = {c.normalisee: c for c in choix}
        crees: list[dict] = []
        ignores = 0
        for ligne in analyse.lignes:
            if ligne.statut != "NOUVEAU":
                continue
            c = par_ref.get(ligne.normalisee)
            if c is not None and not c.inclure:
                ignores += 1
                continue
            type_point = ((c.type_point if c and c.type_point else ligne.type_point) or "AUTRE").upper()
            if type_point not in TYPES_POINT:
                raise AppError(f"Type de point inconnu pour {ligne.reference}", code="POINT_TYPE_INVALIDE")
            agence_id = c.agence_id if c and "agence_id" in c.model_fields_set else (
                uuid.UUID(ligne.agence_id) if ligne.agence_id else None
            )
            nom = (c.nom.strip() if c and c.nom and c.nom.strip() else ligne.nom)[:255]
            point = await self.svc.create_point(
                PointFacturationCreate(
                    type_point=type_point,
                    nom=nom,
                    fournisseur_id=fournisseur_id,
                    reference_fournisseur=ligne.reference[:80],
                    compteur=ligne.compteur,
                    agence_id=agence_id,
                    type_facture=type_facture,
                    periodicite=periodicite,
                    date_debut=suivi_depuis,
                    description=f"Import Excel « {file.filename} » — {ligne.feuille}!{ligne.cellule} ({ligne.nom_source})",
                ),
                user,
                commit=False,
                origine=f"import {file.filename}",
            )
            crees.append({"id": str(point.id), "code": point.code, "nom": point.nom, "reference": point.reference_fournisseur})
        resume = self._resume(analyse)
        if crees:
            await self.svc._audit(
                user, "factures.points.import", None, entity="point_facturation",
                after={"fichier": file.filename, "fournisseur": fournisseur, "crees": len(crees), "ignores": ignores,
                       "existants": resume["existants"], "doublons_fichier": resume["doublons_fichier"]},
            )
        await self.db.commit()
        return {
            "crees": crees,
            "nb_crees": len(crees),
            "ignores": ignores,
            "existants": resume["existants"],
            "doublons_fichier": resume["doublons_fichier"],
            "suivi_depuis": suivi_depuis.isoformat(),
        }

"""Extraction texte — pdfplumber (PDF natif) + Tesseract (images / PDF scannés)."""

from __future__ import annotations

import logging
from pathlib import Path

logger = logging.getLogger(__name__)

_IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".tif", ".tiff", ".bmp"}
_MIN_NATIVE_CHARS = 40


def extract_text_from_file(path: Path, *, lang: str = "fra+eng") -> str:
    """Retourne le texte extrait. Lève RuntimeError si échec technique."""
    if not path.is_file():
        raise RuntimeError(f"Fichier introuvable: {path}")

    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _extract_pdf(path, lang=lang)
    if suffix in _IMAGE_SUFFIXES:
        return _ocr_image(path, lang=lang)
    if suffix in {".txt", ".csv"}:
        return path.read_text(encoding="utf-8", errors="ignore")
    if suffix == ".docx":
        return "\n".join(docx_paragraphs(path))
    if suffix in {".xlsx", ".xls"}:
        return _classeur_texte(path)
    # .doc (Word 97-2003 binaire) : pas d'extracteur sans dépendance — texte vide (statut done côté worker)
    logger.info("OCR non applicable pour suffixe %s (%s)", suffix, path.name)
    return ""


_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def docx_paragraphs(path: Path, *, limite: int | None = None) -> list[str]:
    """Paragraphes d'un .docx (corps + tableaux) via la bibliothèque standard."""
    import zipfile
    from xml.etree import ElementTree

    with zipfile.ZipFile(path) as z:
        xml = z.read("word/document.xml")
    racine = ElementTree.fromstring(xml)
    paragraphes: list[str] = []
    for p in racine.iter(f"{_W}p"):
        morceaux = []
        for n in p.iter():
            if n.tag == f"{_W}t" and n.text:
                morceaux.append(n.text)
            elif n.tag == f"{_W}tab":
                morceaux.append("\t")
        texte = "".join(morceaux).strip()
        if texte:
            paragraphes.append(texte)
            if limite and len(paragraphes) >= limite:
                break
    return paragraphes


def _classeur_texte(path: Path) -> str:
    lignes: list[str] = []
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
                    if cells:
                        lignes.append(" ".join(cells))
        finally:
            wb.close()
    else:
        import xlrd

        book = xlrd.open_workbook(path)
        for sheet in book.sheets():
            for r in range(sheet.nrows):
                cells = [str(v).strip() for v in sheet.row_values(r) if str(v).strip()]
                if cells:
                    lignes.append(" ".join(cells))
    return "\n".join(lignes)


def _extract_pdf(path: Path, *, lang: str) -> str:
    native = _pdf_native_text(path)
    if len(native.strip()) >= _MIN_NATIVE_CHARS:
        return native.strip()
    # PDF scanné : OCR page par page si pdf2image/poppler disponibles
    ocr_pages = _pdf_ocr_pages(path, lang=lang)
    if ocr_pages.strip():
        return ocr_pages.strip()
    if native.strip():
        return native.strip()
    raise RuntimeError(
        "PDF sans texte extractible (natif insuffisant et OCR pages indisponible)"
    )


def _pdf_native_text(path: Path) -> str:
    try:
        import pdfplumber
    except ImportError as exc:
        raise RuntimeError("pdfplumber non installé") from exc

    parts: list[str] = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            text = page.extract_text() or ""
            if text.strip():
                parts.append(text)
    return "\n\n".join(parts)


def _pdf_ocr_pages(path: Path, *, lang: str) -> str:
    try:
        from pdf2image import convert_from_path
    except ImportError:
        logger.warning("pdf2image absent — OCR PDF scanné indisponible")
        return ""

    try:
        images = convert_from_path(str(path), dpi=200)
    except Exception:
        logger.exception("Conversion PDF→images échouée pour %s", path)
        return ""

    parts: list[str] = []
    for img in images:
        try:
            parts.append(_ocr_pil(img, lang=lang))
        except Exception:
            logger.exception("OCR page PDF échouée")
    return "\n\n".join(p for p in parts if p.strip())


def _ocr_image(path: Path, *, lang: str) -> str:
    try:
        from PIL import Image
    except ImportError as exc:
        raise RuntimeError("Pillow non installé") from exc
    with Image.open(path) as img:
        return _ocr_pil(img, lang=lang)


def _ocr_pil(image, *, lang: str) -> str:
    try:
        import pytesseract
    except ImportError as exc:
        raise RuntimeError("pytesseract non installé") from exc
    try:
        text = pytesseract.image_to_string(image, lang=lang)
    except pytesseract.TesseractNotFoundError as exc:
        raise RuntimeError("Binaire tesseract introuvable sur le système") from exc
    except Exception as exc:
        # Fallback anglais seul si pack fra manquant
        try:
            text = pytesseract.image_to_string(image, lang="eng")
        except Exception as exc2:
            raise RuntimeError(f"Tesseract a échoué: {exc2}") from exc
    return (text or "").strip()

"""Cycle de vie des contrats MG — sur mg_contrats et les référentiels existants."""

from __future__ import annotations

import csv
import io
import uuid
from datetime import date, timedelta
from decimal import Decimal

from fastapi import HTTPException, status
from sqlalchemy.exc import IntegrityError
from fastapi.responses import Response
from sqlalchemy import func, or_, select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.models.auth import Agence, User
from app.models.mg_ops import (
    MgContrat,
    MgContratEcheance,
    MgContratHistorique,
    MgContratPaiement,
    MgContratParametre,
    MgContratType,
)
from app.models.organisation import Fournisseur
from app.schemas.mg_ops import ContratCreate, ContratUpdate
from app.models.enums import TypeNotification
from app.services.audit_helpers import record_audit
from app.services.notification_service import NotificationService
from app.services.reporting_export import build_styled_pdf, export_now

LOCKED = {"ARCHIVE", "ANNULE"}
TRANSITIONS: dict[str, tuple[set[str], str]] = {
    "soumettre": ({"BROUILLON", "EN_PREPARATION"}, "EN_VALIDATION"),
    "valider": ({"EN_VALIDATION"}, "ACTIF"),
    "rejeter": ({"EN_VALIDATION"}, "REJETE"),
    "suspendre": ({"ACTIF"}, "SUSPENDU"),
    "reprendre": ({"SUSPENDU"}, "ACTIF"),
    "expirer": ({"ACTIF", "SUSPENDU"}, "EXPIRE"),
    "annuler": ({"BROUILLON", "EN_PREPARATION", "EN_VALIDATION", "ACTIF", "SUSPENDU", "REJETE"}, "ANNULE"),
    "archiver": ({"ACTIF", "EXPIRE", "SUSPENDU", "REJETE"}, "ARCHIVE"),
}
DEFAULT_TYPES = [
    ("FOURNITURE", "Fourniture"),
    ("MAINTENANCE", "Maintenance"),
    ("LOCATION", "Location"),
    ("ASSURANCE", "Assurance"),
    ("TRANSPORT", "Transport"),
    ("TELECOM", "Télécommunication"),
    ("NETTOYAGE", "Nettoyage"),
    ("SECURITE", "Sécurité"),
    ("PRESTATION", "Prestation"),
    ("AUTRE", "Autre"),
]
DEFAULT_PARAMS = [
    ("contrats.prefixe", "CTR", "Préfixe des références"),
    ("contrats.alerte_urgent", "7", "Jours — niveau urgent"),
    ("contrats.alerte_attention", "30", "Jours — niveau attention"),
    ("contrats.alerte_info", "90", "Jours — niveau info"),
    ("contrats.taux_tva", "0", "Taux de TVA par défaut (%)"),
]


def paiement_statut(
    *,
    date_prevue: date,
    date_reelle: date | None,
    montant_prevu: Decimal,
    montant_paye: Decimal,
    today: date | None = None,
) -> str:
    today = today or date.today()
    prevu = montant_prevu or Decimal("0")
    paye = montant_paye or Decimal("0")
    if paye <= 0 and date_reelle is None:
        return "EN_RETARD" if date_prevue < today else "A_VENIR"
    if prevu > 0 and paye + Decimal("0.001") < prevu:
        return "PARTIELLEMENT_PAYE"
    return "PAYE"


def niveau_alerte(jours: int, urgent: int, attention: int) -> str:
    if jours < 0:
        return "CRITIQUE"
    if jours <= urgent:
        return "URGENT"
    if jours <= attention:
        return "ATTENTION"
    return "INFO"


class MgContratsService:
    def __init__(self, db: AsyncSession):
        self.db = db

    async def ensure_defaults(self) -> None:
        ntypes = await self.db.scalar(select(func.count()).select_from(MgContratType))
        if not ntypes:
            for code, libelle in DEFAULT_TYPES:
                self.db.add(MgContratType(code=code, libelle=libelle, actif=True))
        nparams = await self.db.scalar(select(func.count()).select_from(MgContratParametre))
        if not nparams:
            for cle, valeur, libelle in DEFAULT_PARAMS:
                self.db.add(MgContratParametre(cle=cle, valeur=valeur, libelle=libelle))
        await self.db.commit()

    async def _param(self, cle: str, default: str) -> str:
        row = await self.db.scalar(select(MgContratParametre).where(MgContratParametre.cle == cle))
        return row.valeur if row and row.valeur != "" else default

    async def _next_ref(self) -> str:
        prefix = await self._param("contrats.prefixe", "CTR")
        year = date.today().year
        like = f"{prefix}-{year}-%"
        count = await self.db.scalar(
            select(func.count()).select_from(MgContrat).where(MgContrat.reference.ilike(like))
        )
        return f"{prefix}-{year}-{int(count or 0) + 1:04d}"

    def _hist(self, contrat: MgContrat, action: str, user: User, frm: str | None, to: str | None, commentaire: str | None = None) -> None:
        contrat.historique.append(
            MgContratHistorique(
                action=action,
                from_statut=frm,
                to_statut=to,
                user_id=user.id,
                user_nom=user.full_name,
                commentaire=commentaire,
            )
        )

    async def _audit(self, user: User, action: str, contrat_id, before=None, after=None) -> None:
        await record_audit(
            self.db,
            user=user,
            action=action,
            entity="contrat",
            entity_id=str(contrat_id),
            before=before,
            after=after,
            espace_code="moyens-generaux",
            module_code="contrats-echeances",
        )

    def _check_dates(self, debut: date | None, fin: date | None) -> None:
        if debut and fin and fin < debut:
            raise HTTPException(status.HTTP_400_BAD_REQUEST, detail="La date de fin ne peut pas précéder la date de début")

    def _montants(self, ht: Decimal | None, taux: Decimal | None, montant: Decimal | None) -> tuple[Decimal | None, Decimal | None, Decimal | None]:
        if ht is None:
            return None, taux, montant
        rate = taux if taux is not None else Decimal("0")
        if rate < 0 or ht < 0:
            raise HTTPException(status.HTTP_400_BAD_REQUEST, detail="Montant ou TVA invalide")
        ttc = (ht * (Decimal("1") + rate / Decimal("100"))).quantize(Decimal("0.01"))
        return ht, rate, ttc

    async def _apply_refs(self, contrat: MgContrat, *, agence_id, fournisseur_id, responsable_id) -> None:
        if agence_id:
            ag = await self.db.get(Agence, agence_id)
            if not ag or ag.deleted_at is not None:
                raise HTTPException(status.HTTP_400_BAD_REQUEST, detail="Agence introuvable")
            contrat.agence_id = ag.id
            contrat.agence_libelle_snapshot = ag.libelle
        if fournisseur_id:
            fr = await self.db.get(Fournisseur, fournisseur_id)
            if not fr or fr.deleted_at is not None:
                raise HTTPException(status.HTTP_400_BAD_REQUEST, detail="Fournisseur introuvable")
            contrat.fournisseur_id = fr.id
            contrat.fournisseur_snapshot = fr.raison_sociale
        if responsable_id:
            user = await self.db.get(User, responsable_id)
            if not user or user.deleted_at is not None:
                raise HTTPException(status.HTTP_400_BAD_REQUEST, detail="Responsable introuvable")
            contrat.responsable_id = user.id
            contrat.responsable_nom = user.full_name

    async def _lock(self, contrat_id: uuid.UUID) -> MgContrat:
        contrat = await self.db.scalar(
            select(MgContrat)
            .options(
                selectinload(MgContrat.echeances),
                selectinload(MgContrat.paiements),
                selectinload(MgContrat.historique),
            )
            .where(MgContrat.id == contrat_id, MgContrat.deleted_at.is_(None))
        )
        if not contrat:
            raise HTTPException(status.HTTP_404_NOT_FOUND, detail="Contrat introuvable")
        return contrat

    def _refresh_echeance(self, contrat: MgContrat) -> None:
        dates = [e.date_prevue for e in contrat.echeances if e.statut not in {"FAITE", "ANNULEE"}]
        if contrat.date_fin:
            dates.append(contrat.date_fin)
        futur = [d for d in dates if d >= date.today()]
        contrat.prochain_echeance = min(futur) if futur else (min(dates) if dates else contrat.date_fin)

    async def list_contrats(
        self,
        *,
        q: str | None = None,
        statut: str | None = None,
        agence_id: uuid.UUID | None = None,
        horizon: str | None = None,
        renouveles: bool = False,
    ) -> list[MgContrat]:
        stmt = select(MgContrat).where(MgContrat.deleted_at.is_(None))
        if renouveles:
            stmt = stmt.where(MgContrat.contrat_precedent_id.is_not(None))
        if statut:
            stmt = stmt.where(MgContrat.statut == statut)
        if agence_id:
            stmt = stmt.where(MgContrat.agence_id == agence_id)
        if q and q.strip():
            like = f"%{q.strip()}%"
            stmt = stmt.where(
                or_(
                    MgContrat.reference.ilike(like),
                    MgContrat.titre.ilike(like),
                    MgContrat.numero_contrat.ilike(like),
                    MgContrat.fournisseur_snapshot.ilike(like),
                    MgContrat.agence_libelle_snapshot.ilike(like),
                    MgContrat.responsable_nom.ilike(like),
                )
            )
        rows = list((await self.db.execute(stmt.order_by(MgContrat.date_debut.desc()))).scalars().all())
        if not horizon:
            return rows
        today = date.today()
        kept = []
        for c in rows:
            fin = c.prochain_echeance or c.date_fin
            if fin is None:
                continue
            delta = (fin - today).days
            if horizon == "retard" and delta < 0:
                kept.append(c)
            elif horizon == "7" and 0 <= delta <= 7:
                kept.append(c)
            elif horizon == "30" and 0 <= delta <= 30:
                kept.append(c)
            elif horizon == "60" and 0 <= delta <= 60:
                kept.append(c)
            elif horizon == "90" and 0 <= delta <= 90:
                kept.append(c)
        return kept

    async def get_contrat(self, contrat_id: uuid.UUID) -> MgContrat:
        return await self._lock(contrat_id)

    async def create_contrat(self, data: ContratCreate, user: User) -> MgContrat:
        await self.ensure_defaults()
        self._check_dates(data.date_debut, data.date_fin)
        ht, taux, ttc = self._montants(data.montant_ht, data.taux_tva, data.montant)
        contrat = MgContrat(
            reference=await self._next_ref(),
            titre=data.titre.strip(),
            numero_contrat=(data.numero_contrat or "").strip() or None,
            description=data.description,
            type_contrat=(data.type_contrat or "AUTRE").strip().upper(),
            date_signature=data.date_signature,
            date_debut=data.date_debut,
            date_fin=data.date_fin,
            montant_ht=ht,
            taux_tva=taux,
            montant=ttc if ttc is not None else data.montant,
            devise=(data.devise or "MRU").upper(),
            periodicite=data.periodicite or "ANNUEL",
            prochain_echeance=data.prochain_echeance or data.date_fin,
            alerte_jours=data.alerte_jours,
            mode_paiement=data.mode_paiement,
            observation=data.observation,
            fournisseur_snapshot=data.fournisseur_snapshot,
            statut="BROUILLON",
        )
        await self._apply_refs(
            contrat,
            agence_id=data.agence_id,
            fournisseur_id=data.fournisseur_id,
            responsable_id=data.responsable_id,
        )
        self._hist(contrat, "creer", user, None, "BROUILLON")
        self.db.add(contrat)
        await self.db.commit()
        await self._audit(user, "contrats.create", contrat.id, after={"ref": contrat.reference})
        await self.db.commit()
        return await self.get_contrat(contrat.id)

    async def update_contrat(self, contrat_id: uuid.UUID, data: ContratUpdate, user: User) -> MgContrat:
        contrat = await self._lock(contrat_id)
        debut = data.date_debut or contrat.date_debut
        fin = data.date_fin if data.date_fin is not None else contrat.date_fin
        self._check_dates(debut, fin)
        for field in (
            "titre",
            "numero_contrat",
            "description",
            "date_signature",
            "date_debut",
            "date_fin",
            "periodicite",
            "prochain_echeance",
            "alerte_jours",
            "mode_paiement",
            "observation",
            "devise",
            "type_contrat",
            "fournisseur_snapshot",
        ):
            val = getattr(data, field)
            if val is not None:
                setattr(contrat, field, val.strip().upper() if field == "type_contrat" and isinstance(val, str) else val)
        if data.montant_ht is not None or data.taux_tva is not None:
            ht, taux, ttc = self._montants(
                data.montant_ht if data.montant_ht is not None else contrat.montant_ht,
                data.taux_tva if data.taux_tva is not None else contrat.taux_tva,
                contrat.montant,
            )
            contrat.montant_ht, contrat.taux_tva, contrat.montant = ht, taux, ttc
        elif data.montant is not None:
            contrat.montant = data.montant
        await self._apply_refs(
            contrat,
            agence_id=data.agence_id,
            fournisseur_id=data.fournisseur_id,
            responsable_id=data.responsable_id,
        )
        self._refresh_echeance(contrat)
        self._hist(contrat, "modifier", user, contrat.statut, contrat.statut)
        await self.db.commit()
        await self._audit(user, "contrats.update", contrat.id, after={"ref": contrat.reference})
        await self.db.commit()
        return await self.get_contrat(contrat.id)

    async def soft_delete(self, contrat_id: uuid.UUID, user: User) -> None:
        contrat = await self._lock(contrat_id)
        from datetime import datetime, timezone

        ancien = contrat.statut
        contrat.deleted_at = datetime.now(timezone.utc)
        self._hist(contrat, "supprimer", user, ancien, None)
        await self.db.commit()
        await self._audit(user, "contrats.delete", contrat.id, before={"statut": ancien, "ref": contrat.reference})
        await self.db.commit()

    async def transition(self, contrat_id: uuid.UUID, action: str, user: User, commentaire: str | None = None) -> MgContrat:
        key = action.strip().lower()
        if key not in TRANSITIONS:
            raise HTTPException(status.HTTP_400_BAD_REQUEST, detail="Action invalide")
        allowed, new_statut = TRANSITIONS[key]
        contrat = await self._lock(contrat_id)
        if contrat.statut == new_statut:
            return contrat
        if contrat.statut not in allowed:
            raise HTTPException(status.HTTP_400_BAD_REQUEST, detail=f"Transition impossible depuis {contrat.statut}")
        if key == "rejeter" and not (commentaire or "").strip():
            raise HTTPException(status.HTTP_400_BAD_REQUEST, detail="Motif obligatoire")
        old = contrat.statut
        contrat.statut = new_statut
        self._hist(contrat, key, user, old, new_statut, commentaire)
        if key in {"soumettre", "valider", "rejeter"} and contrat.responsable_id and contrat.responsable_id != user.id:
            await NotificationService(self.db).create(
                user_id=contrat.responsable_id,
                type_notification=TypeNotification.SYSTEME,
                titre=f"Contrat {contrat.reference}",
                message=f"Le contrat {contrat.reference} est passé à {new_statut}. Consulter la fiche.",
                entity="contrat",
                entity_id=str(contrat.id),
                espace_code="moyens-generaux",
                module_code="contrats-echeances",
                event_type=f"contrats.{key}",
                emetteur_type="utilisateur",
                emetteur_label=user.full_name or "Système",
                actor_user_id=user.id,
            )
        await self.db.commit()
        await self._audit(user, f"contrats.{key}", contrat.id, before={"statut": old}, after={"statut": new_statut})
        await self.db.commit()
        return await self.get_contrat(contrat.id)

    async def renouveler(self, contrat_id: uuid.UUID, user: User) -> MgContrat:
        src = await self._lock(contrat_id)
        if src.statut not in {"ACTIF", "EXPIRE"}:
            raise HTTPException(status.HTTP_400_BAD_REQUEST, detail="Renouvellement possible pour un contrat actif ou expiré")
        child = await self.db.scalar(
            select(MgContrat.id).where(
                MgContrat.contrat_precedent_id == src.id,
                MgContrat.deleted_at.is_(None),
                MgContrat.statut.notin_(["ANNULE"]),
            )
        )
        if child:
            raise HTTPException(status.HTTP_400_BAD_REQUEST, detail="Un renouvellement existe déjà pour ce contrat")
        debut = (src.date_fin + timedelta(days=1)) if src.date_fin else date.today()
        fin = None
        if src.date_debut and src.date_fin:
            fin = debut + (src.date_fin - src.date_debut)
        nouveau = MgContrat(
            reference=await self._next_ref(),
            titre=src.titre,
            numero_contrat=src.numero_contrat,
            description=src.description,
            type_contrat=src.type_contrat,
            fournisseur_id=src.fournisseur_id,
            fournisseur_snapshot=src.fournisseur_snapshot,
            agence_id=src.agence_id,
            agence_libelle_snapshot=src.agence_libelle_snapshot,
            responsable_id=src.responsable_id,
            responsable_nom=src.responsable_nom,
            date_debut=debut,
            date_fin=fin,
            date_signature=None,
            montant=src.montant,
            montant_ht=src.montant_ht,
            taux_tva=src.taux_tva,
            devise=src.devise or "MRU",
            periodicite=src.periodicite,
            prochain_echeance=fin,
            alerte_jours=src.alerte_jours,
            mode_paiement=src.mode_paiement,
            contrat_precedent_id=src.id,
            statut="BROUILLON",
        )
        self._hist(nouveau, "renouveler", user, None, "BROUILLON", f"Issu de {src.reference}")
        self._hist(src, "renouveler", user, src.statut, src.statut, None)
        self.db.add(nouveau)
        try:
            await self.db.commit()
        except IntegrityError:
            await self.db.rollback()
            raise HTTPException(status.HTTP_400_BAD_REQUEST, detail="Un renouvellement existe déjà pour ce contrat")
        await self._audit(user, "contrats.renouveler", nouveau.id, after={"precedent": src.reference, "ref": nouveau.reference})
        await self.db.commit()
        return await self.get_contrat(nouveau.id)

    async def add_echeance(self, contrat_id: uuid.UUID, data: dict, user: User) -> MgContrat:
        contrat = await self._lock(contrat_id)
        contrat.echeances.append(
            MgContratEcheance(
                type_echeance=(data.get("type_echeance") or "AUTRE").upper(),
                date_prevue=data["date_prevue"],
                date_reelle=data.get("date_reelle"),
                montant=data.get("montant"),
                responsable_nom=data.get("responsable_nom") or contrat.responsable_nom,
                statut=data.get("statut") or "A_VENIR",
                commentaire=data.get("commentaire"),
            )
        )
        self._refresh_echeance(contrat)
        self._hist(contrat, "echeance", user, contrat.statut, contrat.statut, data.get("type_echeance"))
        await self.db.commit()
        await self._audit(user, "contrats.echeance", contrat.id)
        await self.db.commit()
        return await self.get_contrat(contrat.id)

    async def update_echeance(self, echeance_id: uuid.UUID, data: dict, user: User) -> dict:
        echeance = await self.db.get(MgContratEcheance, echeance_id)
        if not echeance:
            raise HTTPException(status.HTTP_404_NOT_FOUND, detail="Échéance introuvable")
        contrat = await self._lock(echeance.contrat_id)
        if data.get("type_echeance"):
            echeance.type_echeance = str(data["type_echeance"]).upper()
        if data.get("date_prevue"):
            echeance.date_prevue = data["date_prevue"]
        if "date_reelle" in data:
            echeance.date_reelle = data.get("date_reelle")
        if "montant" in data:
            echeance.montant = data.get("montant")
        if data.get("statut"):
            echeance.statut = str(data["statut"]).upper()
        if "commentaire" in data:
            echeance.commentaire = data.get("commentaire")
        if data.get("responsable_nom") is not None:
            echeance.responsable_nom = data.get("responsable_nom")
        self._refresh_echeance(contrat)
        self._hist(contrat, "echeance_modifier", user, contrat.statut, contrat.statut, echeance.type_echeance)
        await self.db.commit()
        await self._audit(user, "contrats.echeance.update", contrat.id, after={"echeance": str(echeance.id)})
        await self.db.commit()
        return {"id": str(echeance.id), "contrat_id": str(contrat.id)}

    async def delete_echeance(self, echeance_id: uuid.UUID, user: User) -> None:
        echeance = await self.db.get(MgContratEcheance, echeance_id)
        if not echeance:
            raise HTTPException(status.HTTP_404_NOT_FOUND, detail="Échéance introuvable")
        contrat = await self._lock(echeance.contrat_id)
        if echeance in contrat.echeances:
            contrat.echeances.remove(echeance)
        await self.db.delete(echeance)
        await self.db.flush()
        self._refresh_echeance(contrat)
        self._hist(contrat, "echeance_supprimer", user, contrat.statut, contrat.statut)
        await self.db.commit()
        await self._audit(user, "contrats.echeance.delete", contrat.id, before={"echeance": str(echeance_id)})
        await self.db.commit()

    async def add_paiement(self, contrat_id: uuid.UUID, data: dict, user: User) -> MgContrat:
        contrat = await self._lock(contrat_id)
        if contrat.statut in LOCKED:
            raise HTTPException(status.HTTP_400_BAD_REQUEST, detail="Contrat non modifiable")
        ref = (data.get("reference") or "").strip()
        if ref and any((p.reference or "").strip() == ref for p in contrat.paiements):
            raise HTTPException(status.HTTP_400_BAD_REQUEST, detail="Cette référence de paiement existe déjà")
        prevu = Decimal(str(data.get("montant_prevu") or 0))
        paye = Decimal(str(data.get("montant_paye") or 0))
        if prevu < 0 or paye < 0:
            raise HTTPException(status.HTTP_400_BAD_REQUEST, detail="Montant invalide")
        date_prevue = data["date_prevue"]
        statut = paiement_statut(
            date_prevue=date_prevue,
            date_reelle=data.get("date_reelle"),
            montant_prevu=prevu,
            montant_paye=paye,
        )
        contrat.paiements.append(
            MgContratPaiement(
                reference=ref or None,
                date_prevue=date_prevue,
                date_reelle=data.get("date_reelle"),
                montant_prevu=prevu,
                montant_paye=paye,
                devise=contrat.devise or "MRU",
                statut=statut,
                mode=data.get("mode") or contrat.mode_paiement,
                commentaire=data.get("commentaire"),
                echeance_id=data.get("echeance_id"),
            )
        )
        self._hist(contrat, "paiement", user, contrat.statut, contrat.statut, ref or None)
        try:
            await self.db.commit()
        except IntegrityError:
            await self.db.rollback()
            raise HTTPException(status.HTTP_400_BAD_REQUEST, detail="Cette référence de paiement existe déjà")
        await self._audit(user, "contrats.paiement", contrat.id, after={"statut": statut, "paye": str(paye)})
        await self.db.commit()
        return await self.get_contrat(contrat.id)

    async def alertes(self) -> list[dict]:
        await self.ensure_defaults()
        urgent = int(await self._param("contrats.alerte_urgent", "7"))
        attention = int(await self._param("contrats.alerte_attention", "30"))
        info = int(await self._param("contrats.alerte_info", "90"))
        rows = await self.list_contrats()
        today = date.today()
        out = []
        for c in rows:
            if c.statut in {"ANNULE", "ARCHIVE", "BROUILLON"}:
                continue
            fin = c.date_fin or c.prochain_echeance
            if fin is not None:
                jours = (fin - today).days
                if jours <= info:
                    kind = "expiration" if c.date_fin else "echeance"
                    out.append(
                        {
                            "type": kind,
                            "contrat_id": str(c.id),
                            "reference": c.reference,
                            "titre": c.titre,
                            "fournisseur": c.fournisseur_snapshot,
                            "agence": c.agence_libelle_snapshot,
                            "echeance": fin.isoformat(),
                            "jours": jours,
                            "niveau": niveau_alerte(jours, urgent, attention),
                            "statut": c.statut,
                        }
                    )
            if not c.responsable_id and c.statut == "ACTIF":
                out.append(
                    {
                        "type": "sans_responsable",
                        "contrat_id": str(c.id),
                        "reference": c.reference,
                        "titre": c.titre,
                        "fournisseur": c.fournisseur_snapshot,
                        "agence": c.agence_libelle_snapshot,
                        "echeance": None,
                        "jours": None,
                        "niveau": "ATTENTION",
                        "statut": c.statut,
                    }
                )
            if not c.fournisseur_id and not (c.fournisseur_snapshot or "").strip() and c.statut == "ACTIF":
                out.append(
                    {
                        "type": "sans_fournisseur",
                        "contrat_id": str(c.id),
                        "reference": c.reference,
                        "titre": c.titre,
                        "fournisseur": None,
                        "agence": c.agence_libelle_snapshot,
                        "echeance": None,
                        "jours": None,
                        "niveau": "ATTENTION",
                        "statut": c.statut,
                    }
                )
        for p in (
            await self.db.execute(
                select(MgContratPaiement, MgContrat)
                .join(MgContrat, MgContrat.id == MgContratPaiement.contrat_id)
                .where(
                    MgContrat.deleted_at.is_(None),
                    MgContrat.statut.in_(["ACTIF", "SUSPENDU", "EXPIRE"]),
                    MgContratPaiement.statut.in_(["A_VENIR", "EN_RETARD", "PARTIELLEMENT_PAYE"]),
                )
            )
        ).all():
            pay, contrat = p
            jours = (pay.date_prevue - today).days
            if pay.statut == "EN_RETARD" or jours <= attention:
                out.append(
                    {
                        "type": "paiement",
                        "contrat_id": str(contrat.id),
                        "reference": contrat.reference,
                        "titre": contrat.titre,
                        "fournisseur": contrat.fournisseur_snapshot,
                        "agence": contrat.agence_libelle_snapshot,
                        "echeance": pay.date_prevue.isoformat(),
                        "jours": jours,
                        "niveau": niveau_alerte(jours, urgent, attention),
                        "statut": pay.statut,
                    }
                )
        return out

    async def dashboard(self) -> dict:
        rows = await self.list_contrats()
        today = date.today()

        def count(statut: str) -> int:
            return sum(1 for c in rows if c.statut == statut)

        montant = sum((c.montant or 0) for c in rows if c.statut == "ACTIF")
        a_venir = Decimal("0")
        en_retard = Decimal("0")
        pays = (
            await self.db.execute(
                select(MgContratPaiement).join(MgContrat, MgContrat.id == MgContratPaiement.contrat_id).where(MgContrat.deleted_at.is_(None))
            )
        ).scalars().all()
        for p in pays:
            reste = (p.montant_prevu or 0) - (p.montant_paye or 0)
            if p.statut == "EN_RETARD":
                en_retard += reste
            elif p.statut in {"A_VENIR", "PARTIELLEMENT_PAYE"} and p.date_prevue >= today:
                a_venir += reste
        expires_dates = sum(
            1
            for c in rows
            if c.statut == "ACTIF" and c.date_fin and c.date_fin < today
        )
        echeance_30 = sum(
            1
            for c in rows
            if c.statut == "ACTIF"
            and (c.date_fin or c.prochain_echeance)
            and 0 <= ((c.date_fin or c.prochain_echeance) - today).days <= 30
        )
        return {
            "total": len(rows),
            "actifs": count("ACTIF"),
            "brouillons": count("BROUILLON"),
            "en_validation": count("EN_VALIDATION"),
            "suspendus": count("SUSPENDU"),
            "expires": count("EXPIRE"),
            "date_depassee": expires_dates,
            "echeance_30": echeance_30,
            "archives": count("ARCHIVE"),
            "annules": count("ANNULE"),
            "montant_actifs": float(montant),
            "paiements_a_venir": float(a_venir),
            "paiements_en_retard": float(en_retard),
            "alertes": len(await self.alertes()) if rows else 0,
        }

    async def list_responsables(self, q: str | None = None) -> list[dict]:
        stmt = select(User).where(User.deleted_at.is_(None), User.is_active.is_(True))
        if q and q.strip():
            like = f"%{q.strip()}%"
            stmt = stmt.where(or_(User.full_name.ilike(like), User.email.ilike(like)))
        rows = list((await self.db.execute(stmt.order_by(User.full_name).limit(200))).scalars().all())
        return [{"id": str(u.id), "full_name": u.full_name, "email": u.email} for u in rows]

    async def list_echeances(self, horizon: str | None = None) -> list[dict]:
        rows = (
            await self.db.execute(
                select(MgContratEcheance, MgContrat)
                .join(MgContrat, MgContrat.id == MgContratEcheance.contrat_id)
                .where(MgContrat.deleted_at.is_(None))
                .order_by(MgContratEcheance.date_prevue)
            )
        ).all()
        today = date.today()
        out = []
        for echeance, contrat in rows:
            jours = (echeance.date_prevue - today).days
            if horizon == "retard" and jours >= 0:
                continue
            if horizon in {"7", "30", "60", "90"} and not (0 <= jours <= int(horizon)):
                continue
            out.append(
                {
                    "id": str(echeance.id),
                    "contrat_id": str(contrat.id),
                    "reference": contrat.reference,
                    "titre": contrat.titre,
                    "type_echeance": echeance.type_echeance,
                    "date_prevue": echeance.date_prevue.isoformat(),
                    "date_reelle": echeance.date_reelle.isoformat() if echeance.date_reelle else None,
                    "jours": jours,
                    "statut": echeance.statut,
                    "commentaire": echeance.commentaire,
                    "montant": float(echeance.montant) if echeance.montant is not None else None,
                    "agence": contrat.agence_libelle_snapshot,
                    "fournisseur": contrat.fournisseur_snapshot,
                }
            )
        return out

    async def list_paiements(self, statut: str | None = None) -> list[dict]:
        stmt = (
            select(MgContratPaiement, MgContrat)
            .join(MgContrat, MgContrat.id == MgContratPaiement.contrat_id)
            .where(MgContrat.deleted_at.is_(None))
            .order_by(MgContratPaiement.date_prevue)
        )
        if statut:
            stmt = stmt.where(MgContratPaiement.statut == statut)
        rows = (await self.db.execute(stmt)).all()
        return [
            {
                "id": str(pay.id),
                "contrat_id": str(contrat.id),
                "reference": contrat.reference,
                "titre": contrat.titre,
                "paiement_ref": pay.reference,
                "date_prevue": pay.date_prevue.isoformat(),
                "date_reelle": pay.date_reelle.isoformat() if pay.date_reelle else None,
                "montant_prevu": float(pay.montant_prevu or 0),
                "montant_paye": float(pay.montant_paye or 0),
                "reste": float((pay.montant_prevu or 0) - (pay.montant_paye or 0)),
                "statut": pay.statut,
                "devise": pay.devise,
                "fournisseur": contrat.fournisseur_snapshot,
                "agence": contrat.agence_libelle_snapshot,
            }
            for pay, contrat in rows
        ]

    async def list_agences(self) -> list[Agence]:
        return list(
            (
                await self.db.execute(
                    select(Agence).where(Agence.is_active.is_(True), Agence.deleted_at.is_(None)).order_by(Agence.libelle)
                )
            ).scalars().all()
        )

    async def list_fournisseurs(self, q: str | None = None) -> list[dict]:
        stmt = select(Fournisseur).where(Fournisseur.deleted_at.is_(None), Fournisseur.is_active.is_(True))
        if q and q.strip():
            like = f"%{q.strip()}%"
            stmt = stmt.where(or_(Fournisseur.raison_sociale.ilike(like), Fournisseur.code.ilike(like)))
        rows = list((await self.db.execute(stmt.order_by(Fournisseur.raison_sociale).limit(200))).scalars().all())
        return [
            {
                "id": str(f.id),
                "code": f.code,
                "raison_sociale": f.raison_sociale,
                "telephone": f.telephone,
                "email": f.email,
                "adresse": f.adresse,
            }
            for f in rows
        ]

    async def list_types(self) -> list[MgContratType]:
        await self.ensure_defaults()
        return list((await self.db.execute(select(MgContratType).order_by(MgContratType.libelle))).scalars().all())

    async def list_params(self) -> list[MgContratParametre]:
        await self.ensure_defaults()
        return list((await self.db.execute(select(MgContratParametre).order_by(MgContratParametre.cle))).scalars().all())

    async def set_param(self, cle: str, valeur: str, libelle: str | None = None) -> MgContratParametre:
        row = await self.db.scalar(select(MgContratParametre).where(MgContratParametre.cle == cle))
        if not row:
            raise HTTPException(status.HTTP_404_NOT_FOUND, detail="Paramètre introuvable")
        row.valeur = valeur
        if libelle is not None:
            row.libelle = libelle.strip() or row.libelle
        await self.db.commit()
        await self.db.refresh(row)
        return row

    async def create_param(self, cle: str, valeur: str, libelle: str | None) -> MgContratParametre:
        key = cle.strip()
        if not key:
            raise HTTPException(status.HTTP_400_BAD_REQUEST, detail="Clé obligatoire")
        exists = await self.db.scalar(select(MgContratParametre.id).where(MgContratParametre.cle == key))
        if exists:
            raise HTTPException(status.HTTP_400_BAD_REQUEST, detail="Ce paramètre existe déjà")
        row = MgContratParametre(cle=key, valeur=valeur, libelle=(libelle or "").strip() or None)
        self.db.add(row)
        try:
            await self.db.commit()
        except IntegrityError:
            await self.db.rollback()
            raise HTTPException(status.HTTP_400_BAD_REQUEST, detail="Ce paramètre existe déjà")
        await self.db.refresh(row)
        return row

    async def delete_param(self, cle: str) -> None:
        row = await self.db.scalar(select(MgContratParametre).where(MgContratParametre.cle == cle))
        if not row:
            raise HTTPException(status.HTTP_404_NOT_FOUND, detail="Paramètre introuvable")
        await self.db.delete(row)
        await self.db.commit()

    async def create_type(self, code: str, libelle: str) -> MgContratType:
        key = code.strip().upper().replace(" ", "_")
        label = libelle.strip()
        if not key or not label:
            raise HTTPException(status.HTTP_400_BAD_REQUEST, detail="Code et libellé obligatoires")
        row = MgContratType(code=key, libelle=label, actif=True)
        self.db.add(row)
        try:
            await self.db.commit()
        except IntegrityError:
            await self.db.rollback()
            raise HTTPException(status.HTTP_400_BAD_REQUEST, detail="Ce type existe déjà")
        await self.db.refresh(row)
        return row

    async def update_type(self, type_id: uuid.UUID, *, libelle: str | None, actif: bool | None) -> MgContratType:
        row = await self.db.get(MgContratType, type_id)
        if not row:
            raise HTTPException(status.HTTP_404_NOT_FOUND, detail="Type introuvable")
        if libelle is not None and libelle.strip():
            row.libelle = libelle.strip()
        if actif is not None:
            row.actif = actif
        await self.db.commit()
        await self.db.refresh(row)
        return row

    async def delete_type(self, type_id: uuid.UUID) -> dict:
        row = await self.db.get(MgContratType, type_id)
        if not row:
            raise HTTPException(status.HTTP_404_NOT_FOUND, detail="Type introuvable")
        used = await self.db.scalar(
            select(func.count()).select_from(MgContrat).where(
                MgContrat.deleted_at.is_(None),
                MgContrat.type_contrat == row.code,
            )
        )
        if used:
            row.actif = False
            await self.db.commit()
            return {"ok": True, "desactive": True}
        await self.db.delete(row)
        await self.db.commit()
        return {"ok": True, "desactive": False}

    async def export(self, user: User, report_key: str, fmt: str):
        rows_src = await self.list_contrats()
        if report_key == "actifs":
            rows_src = [c for c in rows_src if c.statut == "ACTIF"]
        elif report_key == "expires":
            today = date.today()
            rows_src = [c for c in rows_src if c.statut == "EXPIRE" or (c.date_fin and c.date_fin < today and c.statut == "ACTIF")]
        elif report_key == "echeances":
            rows_src = await self.list_contrats(horizon="90")
        headers = ["Référence", "Objet", "Type", "Fournisseur", "Agence", "Début", "Fin", "Montant", "Statut"]
        rows = [
            [
                c.reference,
                c.titre,
                c.type_contrat,
                c.fournisseur_snapshot or "",
                c.agence_libelle_snapshot or "",
                str(c.date_debut),
                str(c.date_fin or ""),
                float(c.montant or 0),
                c.statut,
            ]
            for c in rows_src
        ]
        if fmt == "json":
            return {"key": report_key, "title": f"Contrats — {report_key}", "headers": headers, "rows": rows, "count": len(rows)}
        if not rows:
            raise HTTPException(status.HTTP_404_NOT_FOUND, detail="Aucune donnée pour ce rapport")
        await self._audit(user, "contrats.export", None, after={"rapport": report_key, "format": fmt})
        await self.db.commit()
        when = export_now()
        title = f"BEA-DIGITAL — Contrats ({report_key})"
        if fmt == "csv":
            buf = io.StringIO()
            w = csv.writer(buf, delimiter=";")
            w.writerow(headers)
            w.writerows(rows)
            return Response(
                content=buf.getvalue().encode("utf-8-sig"),
                media_type="text/csv; charset=utf-8",
                headers={"Content-Disposition": f'attachment; filename="contrats-{report_key}.csv"'},
            )
        if fmt == "pdf":
            content = build_styled_pdf(
                report_title=title,
                headers=headers,
                rows=rows,
                subtitle=f"Généré par {user.full_name or user.email}",
                exported_at=when,
                landscape_mode=True,
            )
            return Response(
                content=content,
                media_type="application/pdf",
                headers={"Content-Disposition": f'attachment; filename="contrats-{report_key}.pdf"'},
            )
        from openpyxl import Workbook
        from openpyxl.styles import Font

        wb = Workbook()
        ws = wb.active
        ws.title = "Contrats"
        ws["A1"] = title
        ws["A1"].font = Font(bold=True, size=14)
        for i, h in enumerate(headers, 1):
            ws.cell(3, i, h).font = Font(bold=True)
        for r_i, row in enumerate(rows, 4):
            for c_i, val in enumerate(row, 1):
                ws.cell(r_i, c_i, val)
        out = io.BytesIO()
        wb.save(out)
        return Response(
            content=out.getvalue(),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="contrats-{report_key}.xlsx"'},
        )

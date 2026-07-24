from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.core.exceptions import ValidationError
from app.models import Agence, CategorieImmobilisation, Immobilisation
from app.schemas.immobilisation import ImmobilisationCreate
from app.services.immobilisation_defaults import load_categorie, prepare_create, validate_immobilisation
from app.services.immobilisation_service import ImmobilisationService


def immobilisations_import_template_bytes() -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Import"
    ws.append(
        [
            "code_inventaire",
            "designation",
            "categorie_code",
            "date_acquisition",
            "date_comptabilisation",
            "valeur_brute",
            "taux",
            "agence_code",
            "valeur_residuelle",
            "numero_facture",
        ]
    )
    ws.append(
        ["IMMO-001", "Exemple matériel", "TY-142041", "2026-01-15", "2026-01-15", "150000", "20", "", "0", ""]
    )
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _cell_date(value) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str) and value.strip():
        return date.fromisoformat(value.strip()[:10])
    raise ValidationError(f"Date invalide : {value}")


def _cell_decimal(value) -> Decimal:
    try:
        return Decimal(str(value).replace(",", ".")).quantize(Decimal("0.01"))
    except (InvalidOperation, ValueError) as exc:
        raise ValidationError(f"Montant invalide : {value}") from exc


class ImmobilisationImportService:
    def __init__(self, db: AsyncSession):
        self.db = db

    async def import_from_xlsx(self, content: bytes) -> tuple[int, list[str]]:
        wb = load_workbook(BytesIO(content), read_only=True)
        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)
        header = [str(h).strip().lower() if h else "" for h in next(rows_iter)]
        required = {"code_inventaire", "designation", "categorie_code", "date_acquisition", "valeur_brute"}
        if not required.issubset(set(header)):
            raise ValidationError(f"Colonnes requises : {', '.join(sorted(required))}")

        idx = {name: header.index(name) for name in header if name}

        agences = {
            a.code: a.id
            for a in (
                await self.db.execute(select(Agence).where(Agence.deleted_at.is_(None)))
            ).scalars().all()
        }
        categories = {
            c.code: c
            for c in (
                await self.db.execute(
                    select(CategorieImmobilisation).where(CategorieImmobilisation.deleted_at.is_(None))
                )
            ).scalars().all()
        }

        created = 0
        errors: list[str] = []
        line_no = 1
        service = ImmobilisationService(self.db)

        for raw in rows_iter:
            line_no += 1
            if not raw or all(c is None or str(c).strip() == "" for c in raw):
                continue
            try:
                code = str(raw[idx["code_inventaire"]]).strip()
                designation = str(raw[idx["designation"]]).strip()
                cat_code = str(raw[idx["categorie_code"]]).strip()
                if not code or not designation or not cat_code:
                    raise ValidationError("code, désignation et categorie_code obligatoires")
                categorie = categories.get(cat_code)
                if categorie is None:
                    raise ValidationError(f"Catégorie inconnue : {cat_code}")

                agence_id = None
                if "agence_code" in idx and raw[idx["agence_code"]]:
                    ac = str(raw[idx["agence_code"]]).strip()
                    if ac and ac not in agences:
                        raise ValidationError(f"Agence inconnue : {ac}")
                    agence_id = agences.get(ac) if ac else None

                date_acq = _cell_date(raw[idx["date_acquisition"]])
                if "date_comptabilisation" in idx and raw[idx["date_comptabilisation"]]:
                    date_compta = _cell_date(raw[idx["date_comptabilisation"]])
                else:
                    date_compta = date_acq

                taux = None
                if "taux" in idx and raw[idx["taux"]] is not None and str(raw[idx["taux"]]).strip() != "":
                    taux = _cell_decimal(raw[idx["taux"]])

                numero_facture = None
                if "numero_facture" in idx and raw[idx["numero_facture"]]:
                    numero_facture = str(raw[idx["numero_facture"]]).strip() or None

                residuelle = Decimal("0")
                if "valeur_residuelle" in idx and raw[idx["valeur_residuelle"]] is not None:
                    residuelle = _cell_decimal(raw[idx["valeur_residuelle"]])

                payload = ImmobilisationCreate(
                    code_inventaire=code,
                    designation=designation,
                    categorie_id=categorie.id,
                    agence_id=agence_id,
                    date_acquisition=date_acq,
                    date_comptabilisation=date_compta,
                    valeur_brute=_cell_decimal(raw[idx["valeur_brute"]]),
                    valeur_residuelle=residuelle,
                    taux=taux,
                    numero_facture=numero_facture,
                )
                data = prepare_create(payload, categorie)
                item = Immobilisation(**data)
                from app.services.immobilisation_defaults import apply_categorie_defaults

                apply_categorie_defaults(
                    item,
                    categorie,
                    override_comptes=True,
                    preserve_taux=taux is not None,
                )
                validate_immobilisation(item, categorie)
                item.qr_code_data = f"IMMO:{item.code_inventaire}"
                item.barcode_data = item.code_inventaire
                await service.repo.add(item)
                created += 1
            except Exception as exc:
                errors.append(f"Ligne {line_no}: {exc}")

        return created, errors

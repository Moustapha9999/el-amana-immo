"""Formation & Sensibilisation — feuille de présence (PDF / Excel) et rapports (PDF / Excel).

Habillage commun BEA DIGITAL (``reporting_export``) : logo, ligne banque, pied de page horodaté.
La feuille de présence est un document imprimable : colonne « Signature » large et lignes
hautes pour signer à la main.
"""

from __future__ import annotations

from datetime import date
from io import BytesIO
from typing import Any, Sequence

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XlImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from reportlab.graphics.charts.barcharts import HorizontalBarChart, VerticalBarChart
from reportlab.graphics.charts.legends import Legend
from reportlab.graphics.charts.piecharts import Pie
from reportlab.graphics.shapes import Drawing, String
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (
    CondPageBreak,
    Image,
    KeepTogether,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from app.services.reporting_export import (
    _COLOR_BORDER,
    _COLOR_BRAND,
    _COLOR_META,
    _COLOR_NAVY,
    _COLOR_TITLE,
    _COLOR_ZEBRA,
    _bank_line,
    _pdf_footer,
    _pdf_styles,
    build_styled_workbook_multi,
    export_now,
    format_export_datetime,
    resolve_bea_logo_path,
)

NAVY = colors.HexColor(f"#{_COLOR_NAVY}")
BRAND = colors.HexColor(f"#{_COLOR_BRAND}")
BORDER = colors.HexColor(f"#{_COLOR_BORDER}")
ZEBRA = colors.HexColor(f"#{_COLOR_ZEBRA}")
META = colors.HexColor(f"#{_COLOR_META}")
OK = colors.HexColor("#0f766e")
KO = colors.HexColor("#c2410c")
GRIS = colors.HexColor("#94a3b8")
PALETTE = [colors.HexColor(c) for c in ("#1a5278", "#0f766e", "#c2410c", "#7c3aed", "#0891b2",
                                        "#b91c1c", "#64748b", "#ca8a04")]


def fr_date(iso: str | date | None) -> str:
    if not iso:
        return "—"
    d = iso if isinstance(iso, date) else date.fromisoformat(str(iso)[:10])
    return d.strftime("%d/%m/%Y")


def fr_taux(v: float | None) -> str:
    return "—" if v is None else f"{v:.1f} %".replace(".", ",")


def _esc(s: Any) -> str:
    return str(s if s is not None else "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _doc(buf: BytesIO, titre: str) -> SimpleDocTemplate:
    return SimpleDocTemplate(buf, pagesize=A4, title=titre, leftMargin=14 * mm, rightMargin=14 * mm,
                             topMargin=16 * mm, bottomMargin=18 * mm)


def _logo(story: list, largeur: float = 42) -> None:
    path = resolve_bea_logo_path()
    if path is not None:
        try:
            story.append(Image(str(path), width=largeur * mm, height=largeur / 3 * mm))
        except Exception:
            pass


# --------------------------------------------------------------- feuille de présence
def feuille_presence_pdf(session: dict) -> bytes:
    when = export_now()
    titre = f"Feuille de présence — {session['reference']}"
    buf = BytesIO()
    doc = _doc(buf, titre)
    st = _pdf_styles()
    centre = ParagraphStyle("c", parent=st["title"], alignment=TA_CENTER)
    banque = ParagraphStyle("b", parent=centre, fontSize=17, textColor=NAVY, spaceAfter=0)
    sous = ParagraphStyle("s", parent=centre, fontSize=12, textColor=BRAND, spaceBefore=0)
    feuille = ParagraphStyle("f", parent=centre, fontSize=11, textColor=colors.HexColor(f"#{_COLOR_TITLE}"),
                             spaceBefore=2, spaceAfter=8)
    lab = ParagraphStyle("l", parent=st["cell"], fontName="Helvetica-Bold", fontSize=9.5, textColor=NAVY)
    val = ParagraphStyle("v", parent=st["cell"], fontSize=9.5, leading=12)
    head = ParagraphStyle("h", parent=st["header"], fontSize=9.5, leading=12)
    cell = ParagraphStyle("ce", parent=st["cell"], fontSize=9.5, leading=12)

    story: list = []
    _logo(story)
    story += [
        Paragraph("BANQUE EL AMANA", banque),
        Paragraph("FORMATION &amp; SENSIBILISATION", sous),
        Paragraph("FEUILLE DE PRÉSENCE", feuille),
    ]
    infos = [
        ("THÈME", session.get("theme_libelle") or "—"),
        ("DATE", fr_date(session.get("date_session"))),
        ("LIEU", (session.get("lieu") or {}).get("libelle") or "—"),
        ("FORMATEUR", session.get("formateur_libelle") or "—"),
    ]
    if session.get("intitule"):
        infos.insert(1, ("INTITULÉ", session["intitule"]))
    info_tbl = Table([[Paragraph(k, lab), Paragraph(_esc(v), val)] for k, v in infos],
                     colWidths=[32 * mm, doc.width - 32 * mm])
    info_tbl.setStyle(TableStyle([
        ("BOX", (0, 0), (-1, -1), 0.8, NAVY),
        ("LINEBELOW", (0, 0), (-1, -2), 0.3, BORDER),
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#EEF6FC")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    story += [info_tbl, Spacer(1, 6 * mm)]

    participants = session.get("participants") or []
    data = [[Paragraph("N°", head), Paragraph("Nom et prénom", head), Paragraph("Signature", head)]]
    for i, p in enumerate(participants, 1):
        data.append([Paragraph(str(i), ParagraphStyle("n", parent=cell, alignment=TA_CENTER)),
                     Paragraph(_esc(p.get("nom_complet")), cell), ""])
    vides = 0 if participants else 15
    for i in range(len(participants) + 1, len(participants) + 1 + vides):
        data.append([Paragraph(str(i), ParagraphStyle("n", parent=cell, alignment=TA_CENTER)), "", ""])
    largeurs = [14 * mm, 88 * mm, doc.width - 102 * mm]
    tbl = Table(data, colWidths=largeurs, repeatRows=1,
                rowHeights=[9 * mm] + [12 * mm] * (len(data) - 1))
    tbl.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), NAVY),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#94A3B8")),
        ("BOX", (0, 0), (-1, -1), 0.9, NAVY),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, ZEBRA]),
    ]))
    story.append(tbl)
    pied = Table(
        [[Paragraph(f"Participants attendus : <b>{len(participants)}</b>", val),
          Paragraph("Signature du formateur", lab)],
         ["", ""]],
        colWidths=[doc.width / 2, doc.width / 2], rowHeights=[8 * mm, 20 * mm])
    pied.setStyle(TableStyle([
        ("BOX", (1, 0), (1, 1), 0.6, NAVY),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (1, 0), (1, 1), 6),
    ]))
    story += [Spacer(1, 7 * mm), KeepTogether(pied)]
    label = f"{session['reference']} · Imprimé le {format_export_datetime(when)}"

    def _page(c, d):
        _pdf_footer(c, d, exported_label=label, report_title="Feuille de présence")

    doc.build(story, onFirstPage=_page, onLaterPages=_page)
    return buf.getvalue()


def feuille_presence_excel(session: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Feuille de présence"
    thin = Side(style="thin", color="94A3B8")
    bord = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws.column_dimensions["A"].width = 7
    ws.column_dimensions["B"].width = 46
    ws.column_dimensions["C"].width = 40
    logo = resolve_bea_logo_path()
    if logo is not None:
        try:
            img = XlImage(str(logo))
            img.width, img.height = 96, 32
            ws.add_image(img, "A1")
        except Exception:
            pass
    ws.row_dimensions[1].height = 30
    lignes = [
        (2, "BANQUE EL AMANA", 16, _COLOR_NAVY),
        (3, "FORMATION & SENSIBILISATION", 13, _COLOR_BRAND),
        (4, "FEUILLE DE PRÉSENCE", 12, _COLOR_TITLE),
    ]
    for r, texte, taille, couleur in lignes:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
        c = ws.cell(row=r, column=1, value=texte)
        c.font = Font(name="Calibri", size=taille, bold=True, color=couleur)
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[r].height = 22
    infos = [
        ("THÈME", session.get("theme_libelle") or "—"),
        ("DATE", fr_date(session.get("date_session"))),
        ("LIEU", (session.get("lieu") or {}).get("libelle") or "—"),
        ("FORMATEUR", session.get("formateur_libelle") or "—"),
    ]
    r = 6
    for k, v in infos:
        a = ws.cell(row=r, column=1, value=k)
        a.font = Font(bold=True, color=_COLOR_NAVY)
        a.fill = PatternFill("solid", fgColor="EEF6FC")
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        b = ws.cell(row=r, column=2, value=v)
        b.alignment = Alignment(wrap_text=True, vertical="center")
        for col in (1, 2, 3):
            ws.cell(row=r, column=col).border = bord
        ws.row_dimensions[r].height = 20
        r += 1
    r += 1
    for col, h in enumerate(("N°", "Nom et prénom", "Signature"), 1):
        c = ws.cell(row=r, column=col, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=_COLOR_NAVY)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = bord
    ws.row_dimensions[r].height = 24
    entete = r
    participants = session.get("participants") or []
    noms = [p.get("nom_complet") for p in participants] or [None] * 15
    for i, nom in enumerate(noms, 1):
        r += 1
        ws.cell(row=r, column=1, value=i).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=r, column=2, value=nom).alignment = Alignment(vertical="center")
        for col in (1, 2, 3):
            ws.cell(row=r, column=col).border = bord
        ws.row_dimensions[r].height = 34
    r += 2
    ws.cell(row=r, column=1, value=f"Participants attendus : {len(participants)}").font = Font(italic=True,
                                                                                              color=_COLOR_META)
    ws.cell(row=r, column=3, value="Signature du formateur :").font = Font(bold=True, color=_COLOR_NAVY)
    ws.print_title_rows = f"{entete}:{entete}"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddFooter.center.text = f"{session['reference']} — Page &P / &N"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------------ rapports
def libelle_filtres(data: dict, noms: dict[str, str]) -> str:
    f = data.get("filtres") or {}
    parts = []
    if f.get("annee"):
        parts.append(f"Année {f['annee']}")
    if f.get("date_debut") or f.get("date_fin"):
        parts.append(f"Du {fr_date(f.get('date_debut')) if f.get('date_debut') else '…'} "
                     f"au {fr_date(f.get('date_fin')) if f.get('date_fin') else '…'}")
    for k, lab in (("theme_id", "Thème"), ("formateur_id", "Formateur"), ("lieu_id", "Lieu"),
                   ("entite_id", "Entité"), ("perimetre_id", "Périmètre"), ("employe_id", "Employé"),
                   ("fonction_id", "Fonction")):
        if f.get(k):
            parts.append(f"{lab} : {noms.get(str(f[k]), '—')}")
    if f.get("presence"):
        parts.append({"PRESENT": "Présents", "ABSENT": "Absents", "NON_SAISI": "Présence non saisie"}.get(
            f["presence"], f["presence"]))
    return " · ".join(parts) or "Toutes les formations (hors annulées)"


def _rows_dim(lignes: list[dict]) -> list[list]:
    return [[l["libelle"], l["formations"], l["participants"], l["presents"], l["absents"], fr_taux(l["taux"])]
            for l in lignes]


_H_DIM = ["Libellé", "Formations", "Participations", "Présents", "Absents", "Taux de présence"]
_H_FORM = ["Référence", "Date", "Thème", "Lieu", "Formateur", "Statut", "Participants", "Présents", "Absents", "Taux"]
_H_PART = ["Date", "Référence", "Thème", "Lieu", "Nom et prénom", "Fonction", "Entité", "Périmètre", "Présence"]


def _rows_form(data: dict) -> list[list]:
    return [[f["reference"], fr_date(f["date"]), f["theme"], f["lieu"], f["formateur"], f["statut_libelle"],
             f["participants"], f["presents"], f["absents"], fr_taux(f["taux"])] for f in data["formations"]]


def _rows_part(lignes: list[dict]) -> list[list]:
    return [[fr_date(p["date"]), p["reference"], p["theme"], p["lieu"], p["nom_complet"], p["fonction"] or "—",
             p["entite"] or "—", p["perimetre"] or "—", p["presence_libelle"]] for p in lignes]


def _synthese_rows(data: dict, filtres: str) -> list[list]:
    k = data["kpis"]
    return [
        ["Filtres", filtres],
        ["Formations", k["formations"]],
        ["Formations réalisées", k["formations_realisees"]],
        ["Formations à venir", k["formations_a_venir"]],
        ["Participations", k["participations"]],
        ["Présents", k["presents"]],
        ["Absents", k["absents"]],
        ["Présences non saisies", k["non_saisis"]],
        ["Taux de présence", fr_taux(k["taux_presence"])],
        ["Employés concernés", k["employes_concernes"]],
        ["Employés formés (au moins une présence)", k["employes_formes"]],
        ["Employés actifs au référentiel", k["employes_actifs"]],
        ["Thèmes couverts", k["themes_couverts"]],
    ] + [[f"Année {a['libelle']}", f"{a['formations']} formation(s) · {a['participants']} participation(s) · "
                                     f"taux {fr_taux(a['taux'])}"] for a in data["par_annee"]]


def rapport_excel(data: dict, filtres: str) -> bytes:
    parts = data["participations"]
    return build_styled_workbook_multi(
        report_title="Rapport Formation & Sensibilisation",
        subtitle=filtres,
        sheets=[
            ("Synthèse", "Synthèse", ["Indicateur", "Valeur"], _synthese_rows(data, filtres)),
            ("Formations", "Formations", _H_FORM, _rows_form(data)),
            ("Participants", "Participants", _H_PART, _rows_part(parts)),
            ("Présences", "Présences", _H_PART, _rows_part([p for p in parts if p["presence"] == "PRESENT"])),
            ("Absences", "Absences", _H_PART, _rows_part([p for p in parts if p["presence"] == "ABSENT"])),
            ("Par_Thème", "Par thème", _H_DIM, _rows_dim(data["par_theme"])),
            ("Par_Entité", "Par entité", _H_DIM, _rows_dim(data["par_entite"])),
            ("Par_Périmètre", "Par périmètre", _H_DIM, _rows_dim(data["par_perimetre"])),
        ],
    )


def _graph_titre(d: Drawing, texte: str, largeur: float, hauteur: float) -> None:
    d.add(String(0, hauteur - 12, texte, fontName="Helvetica-Bold", fontSize=10, fillColor=NAVY))


def _graph_annees(data: dict, largeur: float) -> Drawing | None:
    annees = data["par_annee"]
    if not annees:
        return None
    h = 62 * mm
    d = Drawing(largeur, h)
    _graph_titre(d, "Formations et présences par année", largeur, h)
    ch = VerticalBarChart()
    ch.x, ch.y, ch.width, ch.height = 28, 30, largeur - 140, h - 58
    ch.data = [[a["presents"] for a in annees], [a["absents"] for a in annees]]
    ch.categoryAxis.categoryNames = [a["libelle"] for a in annees]
    ch.categoryAxis.labels.fontSize = 8
    ch.valueAxis.labels.fontSize = 7
    ch.valueAxis.valueMin = 0
    ch.categoryAxis.style = "stacked"
    ch.bars[0].fillColor, ch.bars[1].fillColor = OK, KO
    ch.bars.strokeColor = None
    ch.barSpacing = 2
    d.add(ch)
    lg = Legend()
    lg.x, lg.y = largeur - 100, h - 34
    lg.fontSize = 8
    lg.alignment = "right"
    lg.colorNamePairs = [(OK, "Présents"), (KO, "Absents")]
    d.add(lg)
    for i, a in enumerate(annees[-6:]):
        d.add(String(largeur - 100, h - 64 - i * 11, f"{a['libelle']} : {a['formations']} formation(s)",
                     fontName="Helvetica", fontSize=7.5, fillColor=META))
    return d


def _graph_hbar(lignes: list[dict], titre: str, largeur: float, couleur, top: int = 10) -> Drawing | None:
    lignes = [l for l in lignes if l["participants"] or l["formations"]][:top]
    if not lignes:
        return None
    h = max(40 * mm, (len(lignes) * 7 + 22) * mm / 1.6)
    d = Drawing(largeur, h)
    _graph_titre(d, titre, largeur, h)
    ch = HorizontalBarChart()
    ch.x, ch.y, ch.width, ch.height = 150, 8, largeur - 190, h - 34
    lignes = list(reversed(lignes))
    ch.data = [[l["participants"] for l in lignes]]
    ch.categoryAxis.categoryNames = [l["libelle"][:34] for l in lignes]
    ch.categoryAxis.labels.fontSize = 7.5
    ch.categoryAxis.labels.boxAnchor = "e"
    ch.valueAxis.labels.fontSize = 7
    ch.valueAxis.valueMin = 0
    ch.bars[0].fillColor = couleur
    ch.bars.strokeColor = None
    ch.barLabelFormat = "%d"
    ch.barLabels.fontSize = 7
    ch.barLabels.nudge = 8
    d.add(ch)
    return d


def _graph_presence(data: dict, largeur: float) -> Drawing | None:
    vals = [p for p in data["presence"] if p["valeur"]]
    if not vals:
        return None
    h = 55 * mm
    d = Drawing(largeur, h)
    _graph_titre(d, "Répartition présence / absence", largeur, h)
    pie = Pie()
    pie.x, pie.y, pie.width, pie.height = 20, 8, h - 34, h - 34
    pie.data = [p["valeur"] for p in vals]
    pie.labels = None
    pie.simpleLabels = 1
    couleurs = {"PRESENT": OK, "ABSENT": KO, "NON_SAISI": GRIS}
    for i, p in enumerate(vals):
        pie.slices[i].fillColor = couleurs[p["cle"]]
        pie.slices[i].strokeColor = colors.white
        pie.slices[i].strokeWidth = 1.5
    d.add(pie)
    total = sum(p["valeur"] for p in vals)
    lg = Legend()
    lg.x, lg.y = h, h - 30
    lg.fontSize = 9
    lg.colorNamePairs = [(couleurs[p["cle"]], f"{p['libelle']} : {p['valeur']} ({p['valeur'] * 100 / total:.0f} %)")
                         for p in vals]
    d.add(lg)
    return d


def _tableau(headers: Sequence[str], rows: Sequence[Sequence[Any]], largeurs: Sequence[float],
             st: dict, *, num_from: int | None = None) -> Table:
    head = ParagraphStyle("hh", parent=st["header"], fontSize=7.5, leading=9)
    cell = ParagraphStyle("cc", parent=st["cell"], fontSize=7.5, leading=9.5)
    right = ParagraphStyle("cr", parent=st["cell_right"], fontSize=7.5, leading=9.5)
    data = [[Paragraph(h, head) for h in headers]]
    for r in rows:
        data.append([Paragraph(_esc(v if v not in (None, "") else "—"),
                               right if num_from is not None and i >= num_from else cell) for i, v in enumerate(r)])
    if len(data) == 1:
        data.append([Paragraph("Aucune donnée", cell)] + [""] * (len(headers) - 1))
    t = Table(data, colWidths=largeurs, repeatRows=1, hAlign="LEFT")
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), NAVY),
        ("GRID", (0, 0), (-1, -1), 0.3, BORDER),
        ("BOX", (0, 0), (-1, -1), 0.6, NAVY),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, ZEBRA]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    return t


def rapport_pdf(data: dict, filtres: str) -> bytes:
    when = export_now()
    titre = "Rapport Formation & Sensibilisation"
    buf = BytesIO()
    doc = _doc(buf, titre)
    st = _pdf_styles()
    w = doc.width
    h2 = ParagraphStyle("h2", parent=st["title"], fontSize=12, textColor=NAVY, spaceBefore=10, spaceAfter=6)
    txt = ParagraphStyle("t", parent=st["cell"], fontSize=9.5, leading=13.5)
    k = data["kpis"]

    story: list = []
    _logo(story)
    story += [Spacer(1, 2 * mm), Paragraph(_bank_line(), st["bank"]), Paragraph(titre, st["title"]),
              Paragraph(f"Exporté le {format_export_datetime(when)} · {_esc(filtres)}", st["meta"])]

    # Résumé
    meilleur_theme = data["par_theme"][0]["libelle"] if data["par_theme"] else None
    resume = (
        f"Sur le périmètre retenu, <b>{k['formations']}</b> formation(s) ont été recensées "
        f"(<b>{k['formations_realisees']}</b> réalisée(s), {k['formations_a_venir']} à venir), pour "
        f"<b>{k['participations']}</b> participation(s) : <b>{k['presents']}</b> présent(s) et "
        f"<b>{k['absents']}</b> absent(s), soit un taux de présence de <b>{fr_taux(k['taux_presence'])}</b>. "
        f"<b>{k['employes_formes']}</b> employé(s) distinct(s) ont suivi au moins une formation "
        f"sur {k['employes_actifs']} employé(s) actif(s) au référentiel."
    )
    if meilleur_theme:
        resume += f" Thème le plus suivi : <b>{_esc(meilleur_theme)}</b>."
    if k["non_saisis"]:
        resume += f" Attention : {k['non_saisis']} présence(s) restent à saisir."
    story += [Paragraph("Résumé", h2), Paragraph(resume, txt), Spacer(1, 4 * mm)]

    # KPI
    big = ParagraphStyle("k", parent=st["title"], fontSize=17, alignment=TA_CENTER, textColor=NAVY,
                         spaceBefore=0, spaceAfter=0)
    small = ParagraphStyle("ks", parent=st["cell_center"], fontSize=7.5, textColor=META)
    cartes = [("Employés formés", k["employes_formes"]), ("Formations", k["formations"]),
              ("Participations", k["participations"]), ("Présents", k["presents"]),
              ("Absents", k["absents"]), ("Taux de présence", fr_taux(k["taux_presence"]))]
    kt = Table([[Paragraph(str(v), big) for _l, v in cartes], [Paragraph(l, small) for l, _v in cartes]],
               colWidths=[w / len(cartes)] * len(cartes))
    kt.setStyle(TableStyle([
        ("BOX", (0, 0), (-1, -1), 0.6, BORDER),
        ("LINEAFTER", (0, 0), (-2, -1), 0.4, BORDER),
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F5F9FC")),
        ("LINEABOVE", (0, 0), (-1, 0), 2.2, BRAND),
        ("TOPPADDING", (0, 0), (-1, 0), 8), ("BOTTOMPADDING", (0, 1), (-1, 1), 7),
    ]))
    story += [Paragraph("Indicateurs clés", h2), kt]

    # Graphiques
    story.append(Paragraph("Graphiques", h2))
    for g in (_graph_annees(data, w), _graph_presence(data, w),
              _graph_hbar(data["par_theme"], "Participations par thème", w, PALETTE[0]),
              _graph_hbar(data["par_perimetre"], "Participations par périmètre", w, PALETTE[3]),
              _graph_hbar(data["par_entite"], "Participations par entité (top 10)", w, PALETTE[4])):
        if g is not None:
            story += [CondPageBreak(g.height + 6 * mm), g, Spacer(1, 4 * mm)]

    # Statistiques
    dim_w = [w * 0.40] + [w * 0.12] * 5
    for titre_dim, cle in (("Statistiques par thème", "par_theme"), ("Statistiques par entité", "par_entite"),
                           ("Statistiques par périmètre", "par_perimetre")):
        story += [CondPageBreak(40 * mm), Paragraph(titre_dim, h2),
                  _tableau(_H_DIM, _rows_dim(data[cle]), dim_w, st, num_from=1)]

    # Listes
    story += [PageBreak(), Paragraph("Liste des formations", h2),
              _tableau(_H_FORM, _rows_form(data),
                       [w * x for x in (0.11, 0.08, 0.2, 0.09, 0.14, 0.09, 0.08, 0.07, 0.07, 0.07)], st, num_from=6)]
    part_w = [w * x for x in (0.08, 0.1, 0.16, 0.08, 0.17, 0.11, 0.12, 0.1, 0.08)]
    parts = data["participations"]
    for titre_l, lignes in (("Participants", parts),
                            ("Présences", [p for p in parts if p["presence"] == "PRESENT"]),
                            ("Absences", [p for p in parts if p["presence"] == "ABSENT"])):
        story += [CondPageBreak(40 * mm), Paragraph(f"{titre_l} ({len(lignes)})", h2),
                  _tableau(_H_PART, _rows_part(lignes), part_w, st)]

    label = f"Exporté le {format_export_datetime(when)}"

    def _page(c, d):
        _pdf_footer(c, d, exported_label=label, report_title=titre)

    doc.build(story, onFirstPage=_page, onLaterPages=_page)
    return buf.getvalue()

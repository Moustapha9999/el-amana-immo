"""Import d'un inventaire Excel (fiches de stock) : analyse, mapping, anomalies, création de campagne.

L'import ne crée jamais d'article : chaque ligne Excel est rapprochée d'un article existant
(code, puis référence, puis désignation unique). Les lignes non reconnues ou en doublon
bloquent l'import tant qu'elles ne sont pas résolues (article choisi ou ligne ignorée).
"""

from __future__ import annotations

import hashlib
import re
import unicodedata
import uuid
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from io import BytesIO

from fastapi import Request, status
from openpyxl import load_workbook
from sqlalchemy import select

from app.core.exceptions import AppError
from app.models.mg_stock import MgArticle, MgInventaireLigne
from app.schemas.mg_stock import InventaireCreate, InventaireImportOptions
from app.schemas.nombres import as_qty
from app.services.mg_inventaire_service import (
    PERM_SAISIE,
    MgInventaireService,
    fin_de_journee,
    periode_label,
)
from app.services.mg_stock_periodes import nature_ecart

MAX_LIGNES = 5000
MAX_OCTETS = 10 * 1024 * 1024

# Champ cible → libellés d'en-tête reconnus (comparés normalisés : minuscules, sans accents).
CHAMPS: dict[str, tuple[str, ...]] = {
    "code": ("ref", "ref.", "reference", "code", "code article", "ref article"),
    "designation": ("article", "designation", "libelle", "article (libelle de la fiche)", "libelle article"),
    "famille": ("categorie", "famille", "sous-famille"),
    "stock_theorique": ("stock final theorique", "stock theorique", "theorique", "stock systeme", "stock final"),
    "stock_physique": (
        "stock physique constate",
        "stock physique",
        "inventaire physique",
        "physique",
        "quantite physique",
        "quantite comptee",
    ),
    "statut_agence": ("statut agence",),
    "verifie": ("verifie (ok)", "verifie", "controle"),
    "observation": ("observations", "observation", "commentaire", "commentaires"),
    "ecart": ("ecart (physique - theorique)", "ecart"),
    "stock_initial": ("stock initial",),
    "entrees": ("entrees",),
    "sorties": ("sorties",),
}
CHAMPS_LIBELLES = {
    "code": "Code / Réf. article",
    "designation": "Désignation",
    "famille": "Famille / Catégorie",
    "stock_theorique": "Stock théorique (fichier)",
    "stock_physique": "Stock physique",
    "statut_agence": "Statut agence",
    "verifie": "Vérifié",
    "observation": "Observations",
    "ecart": "Écart (fichier, recalculé)",
    "stock_initial": "Stock initial",
    "entrees": "Entrées",
    "sorties": "Sorties",
}
OBLIGATOIRES = ("stock_physique",)
BLOQUANTS = frozenset({"INCONNU", "DOUBLON", "QTE_INVALIDE"})


def normaliser(texte) -> str:
    if texte is None:
        return ""
    s = unicodedata.normalize("NFKD", str(texte)).encode("ascii", "ignore").decode("ascii")
    s = s.replace("−", "-").replace("–", "-").lower().strip()
    return re.sub(r"\s+", " ", s)


def _match_champ(entete: str) -> tuple[str | None, bool]:
    """(champ, exact) pour un en-tête normalisé."""
    for champ, synonymes in CHAMPS.items():
        if entete in synonymes:
            return champ, True
    for champ, synonymes in CHAMPS.items():
        if any(entete.startswith(s + " ") or entete.startswith(s + "(") for s in synonymes if len(s) > 3):
            return champ, False
    return None, False


def _quantite(valeur) -> tuple[Decimal | None, str | None]:
    if valeur is None or (isinstance(valeur, str) and not valeur.strip()):
        return None, None
    if isinstance(valeur, bool):
        return None, f"valeur « {valeur} » non numérique"
    try:
        q = Decimal(str(valeur).replace(" ", "").replace(",", "."))
    except (InvalidOperation, ValueError):
        return None, f"valeur « {valeur} » non numérique"
    if q < 0:
        return None, f"quantité négative ({valeur})"
    if q != q.to_integral_value():
        return None, f"quantité non entière ({valeur})"
    return q, None


def _detecter_entete(ws) -> tuple[int | None, dict[int, str], list[str]]:
    """Ligne d'en-tête = première ligne (sur 20) reconnaissant au moins 3 champs dont un stock."""
    for r_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=min(ws.max_row, 20), values_only=True), start=1):
        champs: dict[int, str] = {}
        for c_idx, val in enumerate(row):
            champ, _ = _match_champ(normaliser(val))
            if champ:
                champs[c_idx] = champ
        if len(champs) >= 3 and {"stock_physique", "stock_theorique"} & set(champs.values()):
            entetes = [str(v).strip() if v is not None else "" for v in row]
            return r_idx, champs, entetes
    return None, {}, []


class InventaireImportService:
    def __init__(self, inventaires: MgInventaireService):
        self.inv = inventaires
        self.db = inventaires.db

    def _classeur(self, contenu: bytes):
        if len(contenu) > MAX_OCTETS:
            raise AppError("Fichier trop volumineux (10 Mo maximum).", status.HTTP_413_REQUEST_ENTITY_TOO_LARGE)
        try:
            return load_workbook(BytesIO(contenu), data_only=True, read_only=False)
        except Exception as exc:
            raise AppError(
                "Fichier illisible : un classeur Excel (.xlsx) est attendu.",
                status.HTTP_422_UNPROCESSABLE_ENTITY,
                code="IMPORT_FICHIER_INVALIDE",
            ) from exc

    async def analyser(self, contenu: bytes, nom_fichier: str, options: InventaireImportOptions) -> dict:
        wb = self._classeur(contenu)
        feuilles = [ws.title for ws in wb.worksheets]
        ws = None
        entete_row, colonnes, entetes = None, {}, []
        if options.feuille:
            if options.feuille not in feuilles:
                raise AppError(f"Feuille « {options.feuille} » absente du classeur.", status.HTTP_422_UNPROCESSABLE_ENTITY)
            ws = wb[options.feuille]
            entete_row, colonnes, entetes = _detecter_entete(ws)
        else:
            for candidate in wb.worksheets:
                entete_row, colonnes, entetes = _detecter_entete(candidate)
                if entete_row:
                    ws = candidate
                    break
        if ws is None or entete_row is None:
            raise AppError(
                "Aucune feuille ne contient d'en-tête d'inventaire reconnu (Référence, Désignation, Stock physique…).",
                status.HTTP_422_UNPROCESSABLE_ENTITY,
                code="IMPORT_ENTETE_INTROUVABLE",
            )

        # Mapping : détection automatique puis surcharges explicites (champ → libellé d'en-tête).
        mapping: dict[str, int] = {}
        ambigues: list[dict] = []
        for c_idx, champ in colonnes.items():
            if champ in mapping:
                ambigues.append(
                    {
                        "champ": champ,
                        "libelle": CHAMPS_LIBELLES[champ],
                        "colonnes": [entetes[mapping[champ]], entetes[c_idx]],
                        "retenue": entetes[mapping[champ]],
                    }
                )
                continue
            mapping[champ] = c_idx
        for champ, libelle in (options.mapping or {}).items():
            if champ not in CHAMPS:
                continue
            if not libelle:
                mapping.pop(champ, None)
                continue
            cible = normaliser(libelle)
            idx = next((i for i, e in enumerate(entetes) if normaliser(e) == cible), None)
            if idx is None:
                raise AppError(f"Colonne « {libelle} » introuvable pour {CHAMPS_LIBELLES[champ]}.",
                               status.HTTP_422_UNPROCESSABLE_ENTITY)
            mapping[champ] = idx
        manquants = [CHAMPS_LIBELLES[c] for c in OBLIGATOIRES if c not in mapping]
        if "code" not in mapping and "designation" not in mapping:
            manquants.append("Code ou Désignation")
        colonnes_ignorees = [
            e for i, e in enumerate(entetes) if e and i not in mapping.values()
        ]

        articles = list(
            (await self.db.execute(select(MgArticle).where(MgArticle.deleted_at.is_(None)))).scalars().all()
        )
        par_code = {normaliser(a.code): a for a in articles}
        par_ref: dict[str, MgArticle] = {}
        for a in articles:
            if a.reference:
                par_ref.setdefault(normaliser(a.reference), a)
        par_designation: dict[str, list[MgArticle]] = {}
        for a in articles:
            par_designation.setdefault(normaliser(a.designation), []).append(a)
        par_id = {str(a.id): a for a in articles}
        stocks = await self.inv.stock_a_date(articles, fin_de_journee(options.date_inventaire))

        def cell(row, champ):
            idx = mapping.get(champ)
            return row[idx] if idx is not None and idx < len(row) else None

        lignes: list[dict] = []
        vus: dict[uuid.UUID, int] = {}
        codes_vus: dict[str, int] = {}
        for r_idx, row in enumerate(ws.iter_rows(min_row=entete_row + 1, values_only=True), start=entete_row + 1):
            code = str(cell(row, "code")).strip() if cell(row, "code") not in (None, "") else None
            designation = str(cell(row, "designation")).strip() if cell(row, "designation") not in (None, "") else None
            if not code and not designation:
                continue
            if len(lignes) >= MAX_LIGNES:
                raise AppError(f"Plus de {MAX_LIGNES} lignes : découpez le fichier.", status.HTTP_422_UNPROCESSABLE_ENTITY)
            anomalies: list[str] = []
            avertissements: list[str] = []
            theo_src, err_theo = _quantite(cell(row, "stock_theorique"))
            physique, err_phys = _quantite(cell(row, "stock_physique"))
            statut_agence = normaliser(cell(row, "statut_agence"))
            verifie = normaliser(cell(row, "verifie"))
            exclu = "ancien" in statut_agence or verifie.startswith("non") or "(x)" in verifie
            for err, lib in ((err_theo, "théorique"), (err_phys, "physique")):
                if err:
                    anomalies.append(f"Stock {lib} : {err}")

            article, methode = None, None
            resolution = (options.resolutions or {}).get(str(r_idx))
            if resolution and resolution != "IGNORER":
                article, methode = par_id.get(resolution), "RESOLUTION"
                if article is None:
                    anomalies.append("Article choisi introuvable")
            elif code and normaliser(code) in par_code:
                article, methode = par_code[normaliser(code)], "CODE"
            elif code and normaliser(code) in par_ref:
                article, methode = par_ref[normaliser(code)], "REFERENCE"
            elif designation and len(par_designation.get(normaliser(designation), [])) == 1:
                article, methode = par_designation[normaliser(designation)][0], "DESIGNATION"
                avertissements.append("Reconnu par la désignation : vérifier la correspondance")
            if not code:
                avertissements.append("Article sans code dans le fichier")
            if article is not None and designation and normaliser(designation) != normaliser(article.designation):
                avertissements.append(f"Libellé différent du référentiel : « {article.designation} »")

            if code:
                cle = normaliser(code)
                if cle in codes_vus and resolution != "IGNORER":
                    anomalies.append(f"Code en doublon (déjà ligne {codes_vus[cle]})")
                codes_vus.setdefault(cle, r_idx)
            if article is not None and resolution != "IGNORER":
                if article.id in vus and not any("doublon" in a for a in anomalies):
                    anomalies.append(f"Article en doublon (déjà ligne {vus[article.id]})")
                vus.setdefault(article.id, r_idx)

            theo_systeme = stocks.get(article.id) if article is not None else None
            if resolution == "IGNORER":
                statut = "IGNORE"
            elif any("doublon" in a.lower() for a in anomalies):
                statut = "DOUBLON"
            elif article is None:
                statut = "INCONNU"
                anomalies.append("Article non reconnu dans le référentiel")
            elif err_theo or err_phys:
                statut = "QTE_INVALIDE"
            elif exclu:
                statut = "EXCLU"
            elif physique is None:
                statut = "NON_COMPTE"
                avertissements.append("Stock physique vide : l'article restera à compter")
            else:
                statut = "OK"
            if article is not None and (not article.stockable or not article.is_active):
                avertissements.append("Article inactif ou non stockable : ligne non importée")
                if statut in {"OK", "NON_COMPTE", "EXCLU"}:
                    statut = "HORS_PERIMETRE"
            if (
                options.agence_id
                and article is not None
                and article.agence_id not in (None, options.agence_id)
                and statut in {"OK", "NON_COMPTE", "EXCLU"}
            ):
                statut = "HORS_PERIMETRE"
                avertissements.append("Article d'une autre agence que le périmètre choisi")

            lignes.append(
                {
                    "ligne": r_idx,
                    "code": code,
                    "designation": designation,
                    "famille": str(cell(row, "famille")).strip() if cell(row, "famille") else None,
                    "stock_theorique_fichier": as_qty(theo_src) if theo_src is not None else None,
                    "stock_physique": as_qty(physique) if physique is not None else None,
                    "stock_theorique_systeme": as_qty(theo_systeme) if theo_systeme is not None else None,
                    "ecart": as_qty(physique - theo_systeme)
                    if physique is not None and theo_systeme is not None and not exclu
                    else None,
                    "observation": str(cell(row, "observation")).strip() if cell(row, "observation") else None,
                    "exclu": exclu,
                    "statut": statut,
                    "methode": methode,
                    "article_id": str(article.id) if article else None,
                    "article_code": article.code if article else None,
                    "article_designation": article.designation if article else None,
                    "anomalies": anomalies,
                    "avertissements": avertissements,
                }
            )

        def nb(*statuts):
            return sum(1 for lg in lignes if lg["statut"] in statuts)

        retenues = [lg for lg in lignes if lg["statut"] in {"OK", "NON_COMPTE", "EXCLU"}]
        articles_fichier = {lg["article_id"] for lg in retenues}
        perimetre = await self.inv.articles_perimetre(
            agence_id=options.agence_id, famille_id=None, fin=fin_de_journee(options.date_inventaire)
        )
        absents = [a for a in perimetre if str(a.id) not in articles_fichier]
        ecarts = [lg for lg in retenues if lg["statut"] == "OK" and lg["ecart"]]
        divergences = [
            lg for lg in retenues
            if lg["stock_theorique_fichier"] is not None
            and lg["stock_theorique_systeme"] is not None
            and lg["stock_theorique_fichier"] != lg["stock_theorique_systeme"]
        ]
        resume = {
            "lignes_detectees": len(lignes),
            "reconnus": len(retenues) + nb("HORS_PERIMETRE"),
            "a_importer": len(retenues),
            "comptes": nb("OK"),
            "non_comptes": nb("NON_COMPTE"),
            "exclus": nb("EXCLU"),
            "inconnus": nb("INCONNU"),
            "doublons": nb("DOUBLON"),
            "quantites_invalides": nb("QTE_INVALIDE"),
            "ignores": nb("IGNORE"),
            "hors_perimetre": nb("HORS_PERIMETRE"),
            "sans_code": sum(1 for lg in lignes if not lg["code"]),
            "reconnus_par_designation": sum(1 for lg in lignes if lg["methode"] == "DESIGNATION"),
            "ecarts_physique_systeme": len(ecarts),
            "divergences_theorique": len(divergences),
            "articles_perimetre_absents": len(absents),
        }
        anomalies_globales: list[str] = []
        if manquants:
            anomalies_globales.append("Colonnes obligatoires non mappées : " + ", ".join(manquants))
        for label, key in (
            ("article(s) non reconnu(s)", "inconnus"),
            ("doublon(s)", "doublons"),
            ("quantité(s) invalide(s)", "quantites_invalides"),
        ):
            if resume[key]:
                anomalies_globales.append(f"{resume[key]} {label} à résoudre")
        if not retenues:
            anomalies_globales.append("Aucune ligne importable")
        avertissements_globaux: list[str] = []
        if ambigues:
            avertissements_globaux.append(f"{len(ambigues)} colonne(s) ambiguë(s) : vérifier le mapping")
        if resume["divergences_theorique"]:
            avertissements_globaux.append(
                f"{resume['divergences_theorique']} article(s) dont le théorique du fichier diffère du stock système "
                "à la date : le théorique de l'inventaire reste le stock système (figé), celui du fichier est conservé "
                "pour le rapprochement."
            )
        if resume["exclus"]:
            avertissements_globaux.append(
                f"{resume['exclus']} article(s) marqué(s) « ancienne agence / X » : importés comme exclus du comptage "
                "(aucun ajustement)."
            )
        if absents:
            avertissements_globaux.append(
                f"{len(absents)} article(s) du périmètre absent(s) du fichier : ils resteront « non comptés »."
            )
        return {
            "fichier": nom_fichier,
            "sha256": hashlib.sha256(contenu).hexdigest(),
            "feuilles": feuilles,
            "feuille": ws.title,
            "ligne_entete": entete_row,
            "entetes": [e for e in entetes if e],
            "mapping": {champ: entetes[idx] for champ, idx in mapping.items()},
            "champs": [{"champ": c, "libelle": CHAMPS_LIBELLES[c], "obligatoire": c in OBLIGATOIRES} for c in CHAMPS],
            "colonnes_ambigues": ambigues,
            "colonnes_ignorees": colonnes_ignorees,
            "periode": periode_label(
                options.annee or options.date_inventaire.year, options.mois or options.date_inventaire.month
            ),
            "resume": resume,
            "anomalies": anomalies_globales,
            "avertissements": avertissements_globaux,
            "bloquant": bool(anomalies_globales),
            "articles_absents": [{"id": str(a.id), "code": a.code, "designation": a.designation} for a in absents[:200]],
            "lignes": lignes,
        }

    async def importer(self, contenu: bytes, nom_fichier: str, options: InventaireImportOptions, *,
                       request: Request | None = None):
        self.inv.exiger(PERM_SAISIE)
        analyse = await self.analyser(contenu, nom_fichier, options)
        if analyse["bloquant"]:
            raise AppError(
                "Import bloqué : " + " ; ".join(analyse["anomalies"]),
                status.HTTP_422_UNPROCESSABLE_ENTITY,
                code="IMPORT_ANOMALIES",
            )
        inv = await self.inv.creer(
            InventaireCreate(
                libelle=f"Inventaire {analyse['periode']}",
                date_debut=options.date_inventaire,
                annee=options.annee,
                mois=options.mois,
                agence_id=options.agence_id,
                responsable_nom=options.responsable_nom,
                observation=options.observation,
            ),
            request=request,
            source="IMPORT_EXCEL",
            commit=False,
        )
        lignes_db = {
            lg.article_id: lg
            for lg in (
                await self.db.execute(select(MgInventaireLigne).where(MgInventaireLigne.inventaire_id == inv.id))
            ).scalars().all()
        }
        now = datetime.now(timezone.utc)
        user_id = self.inv.user.id if self.inv.user else None
        comptes = 0
        for lg in analyse["lignes"]:
            if lg["statut"] not in {"OK", "NON_COMPTE", "EXCLU"}:
                continue
            ligne = lignes_db.get(uuid.UUID(lg["article_id"]))
            if ligne is None:
                continue
            if lg["stock_theorique_fichier"] is not None:
                ligne.stock_theorique_source = Decimal(lg["stock_theorique_fichier"])
            if lg["observation"]:
                ligne.observation = lg["observation"][:255]
            if lg["statut"] == "EXCLU":
                ligne.statut_comptage = "EXCLU"
                continue
            if lg["statut"] == "OK":
                physique = Decimal(lg["stock_physique"])
                ligne.stock_physique = physique
                ligne.ecart = physique - Decimal(ligne.stock_theorique or 0)
                ligne.nature_ecart = nature_ecart(ligne.ecart)
                ligne.statut_comptage = "COMPTE"
                ligne.compte_par = user_id
                ligne.compte_at = now
                comptes += 1
        if comptes:
            inv.statut = "EN_COURS"
        inv.import_meta = {
            "fichier": nom_fichier,
            "sha256": analyse["sha256"],
            "feuille": analyse["feuille"],
            "ligne_entete": analyse["ligne_entete"],
            "mapping": analyse["mapping"],
            "resume": analyse["resume"],
            "resolutions": options.resolutions,
            "importe_at": now.isoformat(),
        }
        await self.db.flush()
        await self.inv._audit(
            "import_excel", "mg_inventaire", inv.id, request=request,
            after={"reference": inv.reference, "fichier": nom_fichier, "sha256": analyse["sha256"],
                   **analyse["resume"]},
        )
        await self.db.commit()
        return inv, analyse

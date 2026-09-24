"""Import contrôlé du référentiel articles (USB / Excel).

Architecture uniquement : aucune écriture en base tant que les fichiers
ne sont pas fournis et qu'un opérateur n'a pas validé le mapping.

Phases :
  1. Analyse (feuilles, colonnes, types, doublons)
  2. Rapport de mapping Excel → table/colonne
  3. Prévisualisation
  4. Validation
  5. Import transactionnel (jamais d'écrasement silencieux)
"""

from __future__ import annotations

from io import BytesIO
from typing import Any

from fastapi import HTTPException, status

TARGET_ARTICLE_COLUMNS = {
    "code": "mg_articles.code",
    "reference": "mg_articles.reference",
    "designation": "mg_articles.designation",
    "famille": "mg_article_familles.libelle",
    "sous_famille": "mg_articles.sous_famille",
    "uom": "mg_articles.uom",
    "stockable": "mg_articles.stockable",
    "stock_min": "mg_articles.stock_min",
    "stock_max": "mg_articles.stock_max",
    "emplacement": "mg_articles.emplacement",
    "fournisseur_habituel": "mg_articles.fournisseur_habituel",
}

ALIASES = {
    "code article": "code",
    "code": "code",
    "ref": "reference",
    "référence": "reference",
    "reference": "reference",
    "désignation": "designation",
    "designation": "designation",
    "libelle": "designation",
    "libellé": "designation",
    "famille": "famille",
    "sous famille": "sous_famille",
    "sous-famille": "sous_famille",
    "unité": "uom",
    "unite": "uom",
    "uom": "uom",
    "stockable": "stockable",
    "seuil": "stock_min",
    "stock min": "stock_min",
    "stock max": "stock_max",
    "emplacement": "emplacement",
    "fournisseur": "fournisseur_habituel",
}


def protocol() -> dict[str, Any]:
    return {
        "status": "en_attente_fichiers",
        "message": (
            "Aucun fichier USB n'a été déposé. L'import automatique est interdit. "
            "Lorsque les fichiers seront fournis : analyse → mapping → prévisualisation "
            "→ validation → import transactionnel."
        ),
        "phases": [
            "analyse",
            "mapping",
            "preview",
            "validation",
            "import_transactionnel",
        ],
        "tables_cibles": ["mg_article_familles", "mg_articles"],
        "colonnes_cibles": TARGET_ARTICLE_COLUMNS,
        "regles": [
            "Ne jamais écraser la production sans validation.",
            "Détecter doublons (code) et articles similaires (désignation).",
            "Les articles avec historique restent ACTIF/INACTIF — pas de suppression.",
            "Les stocks éventuels du fichier ne créent pas d'ENTREE artificielle de report.",
        ],
    }


def analyze_workbook(content: bytes, filename: str) -> dict[str, Any]:
    """Phase 1 — lecture seule, aucune écriture."""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise HTTPException(
            status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="openpyxl indisponible",
        ) from exc
    try:
        wb = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST,
            detail=f"Fichier illisible : {exc}",
        ) from exc

    sheets = []
    mapping: list[dict[str, str]] = []
    for name in wb.sheetnames:
        ws = wb[name]
        rows = list(ws.iter_rows(max_row=6, values_only=True))
        headers = [str(c).strip() if c is not None else "" for c in (rows[0] if rows else [])]
        sample = []
        for raw in rows[1:6]:
            sample.append([("" if c is None else str(c)) for c in raw])
        col_map = []
        for h in headers:
            key = ALIASES.get(h.lower().strip())
            col_map.append(
                {
                    "excel": h,
                    "cible": TARGET_ARTICLE_COLUMNS.get(key, "") if key else "",
                    "cle": key or "",
                }
            )
            if key:
                mapping.append(
                    {"feuille": name, "excel": h, "table_colonne": TARGET_ARTICLE_COLUMNS[key]}
                )
        sheets.append(
            {
                "nom": name,
                "colonnes": headers,
                "echantillon": sample,
                "mapping_propose": col_map,
            }
        )
    return {
        "filename": filename,
        "phase": "analyse",
        "ecriture": False,
        "feuilles": sheets,
        "mapping": mapping,
        "incoherences": [
            "Doublons et unités à contrôler en phase 3 (prévisualisation).",
        ],
        "message": "Analyse uniquement — aucun article n'a été créé ni modifié.",
    }

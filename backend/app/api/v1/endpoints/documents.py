"""Document Service — ingestion, recherche, OCR retry, versions, téléchargement."""

from __future__ import annotations

from datetime import date
from pathlib import Path
from uuid import UUID, uuid4

from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, Request, UploadFile, status
from fastapi.responses import FileResponse, Response
from sqlalchemy.ext.asyncio import AsyncSession

from app.api.deps import require_permission
from app.core.exceptions import AppError, raise_http_from_app
from app.db.session import get_db
from app.models import User
from app.schemas.documents import (
    DocumentAuditEventOut,
    DocumentMetadataIn,
    DocumentListOut,
    DocumentOut,
    SoftDeleteIn,
    document_out,
)
from app.services.audit_helpers import record_audit
from app.services.document_ingest_service import DocumentIngestService
from app.services.document_query_service import DocumentQueryService
from app.services.document_reporting_service import DocumentReportingService
from app.services.ged_service import GedService

router = APIRouter(prefix="/documents", tags=["documents"])


def _cell(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def spreadsheet_preview(path: Path, filename: str) -> dict:
    """Aperçu tabulaire des classeurs. PDF et images restent côté navigateur."""
    lower = filename.lower()
    max_rows, max_cols = 120, 16
    rows: list[list[str]] = []
    title = filename
    truncated = False

    if lower.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            ws = wb.active
            title = ws.title or filename
            for index, row in enumerate(ws.iter_rows(values_only=True)):
                if index >= max_rows:
                    truncated = True
                    break
                cells = [_cell(c) for c in list(row)[:max_cols]]
                if any(cells):
                    rows.append(cells)
        finally:
            wb.close()
    elif lower.endswith(".xls"):
        import xlrd

        book = xlrd.open_workbook(path)
        sheet = book.sheet_by_index(0)
        title = sheet.name or filename
        limit = min(sheet.nrows, max_rows)
        truncated = sheet.nrows > max_rows
        for r in range(limit):
            cells = [_cell(sheet.cell_value(r, c)) for c in range(min(sheet.ncols, max_cols))]
            if any(cells):
                rows.append(cells)
    elif lower.endswith(".csv"):
        import csv

        title = filename
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            for index, row in enumerate(csv.reader(handle, delimiter=";")):
                if index >= max_rows:
                    truncated = True
                    break
                cells = [_cell(c) for c in row[:max_cols]]
                if any(cells):
                    rows.append(cells)
    elif lower.endswith(".docx"):
        from app.services.ocr_extract import docx_paragraphs

        try:
            paragraphes = docx_paragraphs(path, limite=400)
        except Exception:
            return {"kind": "none", "title": filename, "rows": [], "truncated": False}
        return {
            "kind": "text",
            "title": filename,
            "rows": [[p] for p in paragraphes],
            "truncated": len(paragraphes) >= 400,
        }
    else:
        return {"kind": "none", "title": filename, "rows": [], "truncated": False}

    return {"kind": "table", "title": title, "rows": rows, "truncated": truncated}


@router.post("", response_model=DocumentOut, status_code=status.HTTP_201_CREATED)
async def upload_document(
    request: Request,
    file: UploadFile = File(...),
    espace_code: str = Form(...),
    module_code: str = Form(...),
    entity: str = Form("manuel"),
    entity_id: str | None = Form(None),
    doc_type: str | None = Form(None),
    title: str | None = Form(None),
    description: str | None = Form(None),
    reference: str | None = Form(None),
    date_document: date | None = Form(None),
    agence_id: UUID | None = Form(None),
    department_id: UUID | None = Form(None),
    fournisseur_id: UUID | None = Form(None),
    security_level: str = Form("internal"),
    user: User = Depends(require_permission("ged.write")),
    db: AsyncSession = Depends(get_db),
):
    """Flux 1 — upload manuel + OCR asynchrone."""
    try:
        eid = (entity_id or "").strip() or str(uuid4())
        query = DocumentQueryService(db)
        allowed = await query.user_espace_codes(user)
        espace = espace_code.strip().lower()
        if not user.is_superuser and espace not in allowed:
            raise HTTPException(status.HTTP_403_FORBIDDEN, detail="Espace non autorisé")

        row = await DocumentIngestService(db).ingest_document(
            file=file,
            espace_code=espace,
            module_code=module_code,
            entity=entity,
            entity_id=eid,
            uploaded_by_id=user.id,
            title=title,
            description=description,
            doc_type=doc_type,
            reference=reference,
            date_document=date_document,
            agence_id=agence_id,
            department_id=department_id,
            fournisseur_id=fournisseur_id,
            security_level=security_level,
        )
        await record_audit(
            db,
            user=user,
            action="document_ingest",
            entity="ged_document",
            entity_id=str(row.id),
            after={
                "espace_code": row.espace_code,
                "module_code": row.module_code,
                "filename": row.filename,
                "ocr_status": row.ocr_status,
            },
            request=request,
            espace_code=row.espace_code,
            module_code="documents",
        )
        return document_out(row)
    except AppError as exc:
        raise_http_from_app(exc)


@router.post(
    "/from-operation",
    response_model=DocumentOut,
    status_code=status.HTTP_201_CREATED,
)
async def archive_from_operation(
    request: Request,
    file: UploadFile = File(...),
    espace_code: str = Form(...),
    module_code: str = Form(...),
    source_type: str = Form(..., description="entity métier (= entity)"),
    source_id: str = Form(..., description="id opération (= entity_id)"),
    doc_type: str | None = Form(None),
    title: str | None = Form(None),
    description: str | None = Form(None),
    reference: str | None = Form(None),
    date_document: date | None = Form(None),
    agence_id: UUID | None = Form(None),
    department_id: UUID | None = Form(None),
    fournisseur_id: UUID | None = Form(None),
    security_level: str = Form("internal"),
    user: User = Depends(require_permission("ged.write")),
    db: AsyncSession = Depends(get_db),
):
    """Flux 2 — archivage depuis une opération métier."""
    try:
        query = DocumentQueryService(db)
        allowed = await query.user_espace_codes(user)
        espace = espace_code.strip().lower()
        if not user.is_superuser and espace not in allowed:
            raise HTTPException(status.HTTP_403_FORBIDDEN, detail="Espace non autorisé")

        row = await DocumentIngestService(db).ingest_document(
            file=file,
            espace_code=espace,
            module_code=module_code,
            entity=source_type.strip(),
            entity_id=source_id.strip(),
            uploaded_by_id=user.id,
            title=title,
            description=description,
            doc_type=doc_type,
            reference=reference,
            date_document=date_document,
            agence_id=agence_id,
            department_id=department_id,
            fournisseur_id=fournisseur_id,
            security_level=security_level,
        )
        await record_audit(
            db,
            user=user,
            action="document_archive_operation",
            entity="ged_document",
            entity_id=str(row.id),
            after={
                "source_type": row.entity,
                "source_id": row.entity_id,
                "module_code": row.module_code,
                "ocr_status": row.ocr_status,
            },
            request=request,
            espace_code=row.espace_code,
            module_code=row.module_code,
        )
        return document_out(row)
    except AppError as exc:
        raise_http_from_app(exc)


@router.get("/search", response_model=DocumentListOut)
async def search_documents(
    request: Request,
    espace_code: str | None = Query(None),
    module_code: str | None = Query(None),
    doc_type: str | None = Query(None),
    agence_id: UUID | None = Query(None),
    fournisseur_id: UUID | None = Query(None),
    department_id: UUID | None = Query(None),
    ocr_status: str | None = Query(None),
    security_level: str | None = Query(None),
    date_debut: date | None = Query(None),
    date_fin: date | None = Query(None),
    q: str | None = Query(None),
    search_ocr: bool = Query(False),
    page: int = Query(1, ge=1),
    size: int = Query(50, ge=1, le=200),
    user: User = Depends(require_permission("ged.read")),
    db: AsyncSession = Depends(get_db),
):
    try:
        rows, total, meta = await DocumentQueryService(db).search(
            user=user,
            espace_code=espace_code,
            module_code=module_code,
            doc_type=doc_type,
            agence_id=agence_id,
            fournisseur_id=fournisseur_id,
            department_id=department_id,
            ocr_status=ocr_status,
            security_level=security_level,
            date_debut=date_debut,
            date_fin=date_fin,
            q=q,
            search_ocr=search_ocr,
            page=page,
            size=size,
            general=espace_code is None,
        )
        await record_audit(
            db,
            user=user,
            action="document_search",
            entity="ged_document",
            entity_id=None,
            after={"q": q, "search_ocr": search_ocr, "total": total},
            request=request,
            module_code="documents",
        )
        return DocumentListOut(
            items=[document_out(r) for r in rows],
            total=total,
            page=page,
            size=size,
            ocr_pending_hint=bool(meta.get("ocr_pending_hint")),
        )
    except AppError as exc:
        raise_http_from_app(exc)


@router.get("/rapports/export")
async def export_documents_report(
    request: Request,
    format: str = Query("csv", alias="format"),
    report_key: str = Query("documents"),
    espace_code: str | None = Query(None),
    module_code: str | None = Query(None),
    doc_type: str | None = Query(None),
    ocr_status: str | None = Query(None),
    date_debut: date | None = Query(None),
    date_fin: date | None = Query(None),
    q: str | None = Query(None),
    general: bool = Query(False),
    user: User = Depends(require_permission("ged.export", "ged.read", "mg.archives.export")),
    db: AsyncSession = Depends(get_db),
):
    try:
        content, media, filename = await DocumentReportingService(db).export(
            user=user,
            fmt=format,
            espace_code=espace_code,
            module_code=module_code,
            doc_type=doc_type,
            ocr_status=ocr_status,
            date_debut=date_debut,
            date_fin=date_fin,
            q=q,
            general=general or espace_code is None,
            report_key=report_key,
        )
        await record_audit(
            db,
            user=user,
            action="document_export",
            entity="ged_document",
            entity_id=None,
            after={"format": format, "report_key": report_key},
            request=request,
            module_code="documents",
        )
        return Response(
            content=content,
            media_type=media,
            headers={"Content-Disposition": f'attachment; filename="{filename}"'},
        )
    except AppError as exc:
        raise_http_from_app(exc)


@router.get("/{document_id}", response_model=DocumentOut)
async def get_document(
    document_id: UUID,
    request: Request,
    user: User = Depends(require_permission("ged.read")),
    db: AsyncSession = Depends(get_db),
):
    try:
        row = await DocumentQueryService(db).get_accessible(document_id, user)
        await record_audit(
            db,
            user=user,
            action="document_view",
            entity="ged_document",
            entity_id=str(row.id),
            request=request,
            espace_code=row.espace_code,
            module_code="documents",
        )
        return document_out(row, include_ocr_text=True)
    except AppError as exc:
        raise_http_from_app(exc)


@router.patch("/{document_id}", response_model=DocumentOut)
async def update_document_metadata(
    document_id: UUID,
    body: DocumentMetadataIn,
    request: Request,
    user: User = Depends(require_permission("ged.write")),
    db: AsyncSession = Depends(get_db),
):
    """Modifie les métadonnées GED. Ne modifie pas l'objet métier source."""
    try:
        row = await DocumentQueryService(db).get_accessible(document_id, user)
        before = {
            "title": row.title,
            "doc_type": row.doc_type,
            "reference": row.reference,
        }
        row = await DocumentIngestService(db).update_metadata(row, body)
        await record_audit(
            db,
            user=user,
            action="document_metadata_update",
            entity="ged_document",
            entity_id=str(row.id),
            before=before,
            after={
                "title": row.title,
                "doc_type": row.doc_type,
                "reference": row.reference,
                "archived_at": row.archived_at.isoformat() if row.archived_at else None,
            },
            request=request,
            espace_code=row.espace_code,
            module_code="documents",
        )
        return document_out(row)
    except AppError as exc:
        raise_http_from_app(exc)


@router.get("/{document_id}/versions", response_model=list[DocumentOut])
async def list_versions(
    document_id: UUID,
    request: Request,
    user: User = Depends(require_permission("ged.read")),
    db: AsyncSession = Depends(get_db),
):
    try:
        rows = await DocumentQueryService(db).list_versions(document_id, user)
        await record_audit(
            db,
            user=user,
            action="document_versions_list",
            entity="ged_document",
            entity_id=str(document_id),
            request=request,
            module_code="documents",
        )
        return [document_out(r) for r in rows]
    except AppError as exc:
        raise_http_from_app(exc)


@router.post(
    "/{document_id}/versions",
    response_model=DocumentOut,
    status_code=status.HTTP_201_CREATED,
)
async def create_version(
    document_id: UUID,
    request: Request,
    file: UploadFile = File(...),
    version_comment: str | None = Form(None),
    user: User = Depends(require_permission("ged.write")),
    db: AsyncSession = Depends(get_db),
):
    try:
        await DocumentQueryService(db).get_accessible(document_id, user)
        row = await DocumentIngestService(db).create_version(
            parent_id=document_id,
            file=file,
            uploaded_by_id=user.id,
            version_comment=version_comment,
        )
        await record_audit(
            db,
            user=user,
            action="document_version_create",
            entity="ged_document",
            entity_id=str(row.id),
            after={
                "parent_document_id": str(document_id),
                "version": row.version,
                "version_comment": row.version_comment,
            },
            request=request,
            espace_code=row.espace_code,
            module_code="documents",
        )
        return document_out(row)
    except AppError as exc:
        raise_http_from_app(exc)


@router.get("/{document_id}/relations", response_model=list[DocumentOut])
async def list_relations(
    document_id: UUID,
    request: Request,
    user: User = Depends(require_permission("ged.read")),
    db: AsyncSession = Depends(get_db),
):
    try:
        rows = await DocumentQueryService(db).list_relations(document_id, user)
        await record_audit(
            db,
            user=user,
            action="document_relations_view",
            entity="ged_document",
            entity_id=str(document_id),
            request=request,
            module_code="documents",
        )
        return [document_out(r) for r in rows]
    except AppError as exc:
        raise_http_from_app(exc)


@router.get("/{document_id}/audit", response_model=list[DocumentAuditEventOut])
async def list_document_audit(
    document_id: UUID,
    request: Request,
    limit: int = Query(100, ge=1, le=500),
    user: User = Depends(require_permission("ged.read", "mg.archives.view", "archives.general.view")),
    db: AsyncSession = Depends(get_db),
):
    try:
        rows = await DocumentQueryService(db).list_audit(document_id, user, limit=limit)
        await record_audit(
            db,
            user=user,
            action="document_audit_view",
            entity="ged_document",
            entity_id=str(document_id),
            request=request,
            module_code="documents",
        )
        return [
            DocumentAuditEventOut(
                id=r.id,
                action=r.action,
                created_at=r.created_at,
                user_id=r.user_id,
                after=r.after_data,
                before=r.before_data,
            )
            for r in rows
        ]
    except AppError as exc:
        raise_http_from_app(exc)


@router.post("/{document_id}/retry-ocr", response_model=DocumentOut)
async def retry_ocr(
    document_id: UUID,
    request: Request,
    user: User = Depends(
        require_permission("ged.write", "archives.general.ocr.retry", "mg.archives.update")
    ),
    db: AsyncSession = Depends(get_db),
):
    try:
        row = await DocumentQueryService(db).get_accessible(document_id, user)
        row = await DocumentIngestService(db).retry_ocr(row.id)
        await record_audit(
            db,
            user=user,
            action="ocr_retry",
            entity="ged_document",
            entity_id=str(row.id),
            after={"ocr_status": row.ocr_status},
            request=request,
            espace_code=row.espace_code,
            module_code="documents",
        )
        return document_out(row)
    except AppError as exc:
        raise_http_from_app(exc)


@router.get("/{document_id}/preview")
async def preview_document(
    document_id: UUID,
    user: User = Depends(require_permission("ged.read", "archives.general.view", "mg.archives.view")),
    db: AsyncSession = Depends(get_db),
):
    """Contenu lisible d'un classeur (Excel / CSV)."""
    try:
        row = await DocumentQueryService(db).get_accessible(document_id, user)
        path = GedService(db).absolute_path(row.stored_path)
        if not path.is_file():
            raise HTTPException(status_code=404, detail="Fichier introuvable sur le serveur")
        return spreadsheet_preview(path, row.filename)
    except AppError as exc:
        raise_http_from_app(exc)


@router.get("/{document_id}/download")
async def download_document(
    document_id: UUID,
    request: Request,
    user: User = Depends(require_permission("ged.download", "ged.read", "mg.archives.download", "archives.general.download")),
    db: AsyncSession = Depends(get_db),
):
    try:
        row = await DocumentQueryService(db).get_accessible(document_id, user)
        path = GedService(db).absolute_path(row.stored_path)
        if not path.exists():
            raise HTTPException(status_code=404, detail="Fichier introuvable sur le serveur")
        await record_audit(
            db,
            user=user,
            action="document_download",
            entity="ged_document",
            entity_id=str(row.id),
            after={"filename": row.filename},
            request=request,
            espace_code=row.espace_code,
            module_code="documents",
        )
        return FileResponse(
            path,
            filename=row.filename,
            media_type=row.mime_type or "application/octet-stream",
        )
    except AppError as exc:
        raise_http_from_app(exc)


@router.delete("/{document_id}", response_model=DocumentOut)
async def soft_delete_document(
    document_id: UUID,
    request: Request,
    body: SoftDeleteIn | None = None,
    user: User = Depends(require_permission("ged.write")),
    db: AsyncSession = Depends(get_db),
):
    try:
        await DocumentQueryService(db).get_accessible(document_id, user)
        reason = body.reason if body else None
        row = await DocumentIngestService(db).soft_delete(
            document_id, user_id=user.id, reason=reason
        )
        await record_audit(
            db,
            user=user,
            action="document_delete",
            entity="ged_document",
            entity_id=str(row.id),
            after={"reason": reason},
            request=request,
            espace_code=row.espace_code,
            module_code="documents",
        )
        return document_out(row)
    except AppError as exc:
        raise_http_from_app(exc)


@router.post("/{document_id}/restore", response_model=DocumentOut)
async def restore_document(
    document_id: UUID,
    request: Request,
    user: User = Depends(require_permission("ged.write", "mg.archives.restore")),
    db: AsyncSession = Depends(get_db),
):
    try:
        await DocumentQueryService(db).get_accessible(
            document_id, user, include_deleted=True
        )
        row = await DocumentIngestService(db).restore(document_id)
        await record_audit(
            db,
            user=user,
            action="document_restore",
            entity="ged_document",
            entity_id=str(row.id),
            request=request,
            espace_code=row.espace_code,
            module_code="documents",
        )
        return document_out(row)
    except AppError as exc:
        raise_http_from_app(exc)

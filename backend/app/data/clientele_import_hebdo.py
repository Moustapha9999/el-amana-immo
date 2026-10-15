"""Format d'import hebdomadaire de classification clients.

Les listes observées dans le classeur source ne sont **pas** le référentiel maître
BEA DIGITAL : elles servent uniquement à détecter les consignes, à contrôler la
cohérence interne (code / libellé / famille BCM) et à signaler une valeur hors
liste observée. Le moteur continue de lire les référentiels versionnés.
"""

from __future__ import annotations

import json
import re
from functools import lru_cache
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from app.services.clientele.consolidation import normaliser_entete

_LISTES = Path(__file__).with_name("clientele_import_hebdo_listes.json")

COLONNES_REQUETE = ("CODE AGENCE", "NOM AGENCE", "RACINE CLIENT", "NOM CLIENT")
COLONNES_METIER = (
    "DATE EER", "PROFIL", "SOUS PROFIL", "Catégorie juridique",
    "Code Secteur BCM", "Libellé secteur BCM", "Famille BCM",
    "NIF", "NNI", "Nationnalité", "Résidence",
    "Statut compte", "Statut Gel",
    "Date mise à jour client", "Date classe risque LBC FT",
    "Classe risque LBC FT", "Motif de risque LBC FT", "Analyste KYC",
)
COLONNES_COMPLEMENT = ("DATE_CREATION_ENTREPRISE", "ENTREPRISE_OUVERTURE_RECENTE")
COLONNES_CLIENTS = COLONNES_REQUETE + COLONNES_METIER + COLONNES_COMPLEMENT
COLONNES_EXPORT = COLONNES_REQUETE + COLONNES_METIER

RESULTATS_MOTEUR = (
    ("SCORE_CALCULE", "Score additif du moteur SCORE (CDC)."),
    ("NIVEAU_PROPOSE_MOTEUR", "Niveau proposé : FAIBLE / MOYEN / ELEVE / INTERDIT / NON CLASSÉ."),
    ("STATUT_EVALUATION", "EVALUE, NON_DISPONIBLE, A_ARBITRER, A_VERIFIER…"),
    ("MOTIF_MOTEUR", "Motif généré à partir des critères réellement utilisés."),
    ("CLASSE_RISQUE_RETENUE", "Niveau enregistré après décision (moteur, fichier ou manuel)."),
    ("MODE_DECISION", "MOTEUR, FICHIER ou MANUEL."),
    ("DATE_CALCUL", "Date/heure du calcul dans BEA DIGITAL."),
    ("VERSION_MOTEUR", "Version du moteur et des règles."),
)

FEUILLES_CLIENTS = frozenset({"CLIENTS"})
FEUILLES_TECHNIQUES = frozenset({
    "REFERENTIELS", "SECTEURS BCM", "SECTEURS_BCM", "NOTICE IMPORT", "NOTICE_IMPORT",
    "DICTIONNAIRE", "LISEZ MOI", "LISEZ_MOI", "RESULTATS MOTEUR", "RESULTATS_MOTEUR",
    "LIRE", "NOTES",
})
INTERDITS_DATE_OUVERTURE = frozenset({
    "DATE EER", "DATE D EER", "DATE ENTREE EN RELATION", "DATE D ENTREE EN RELATION",
})
_CONSIGNES = frozenset({
    "REQUETE", "LISTE INTERNE", "LISTE BCM", "LISTE PAYS",
    "LIBRE FORMAT DATE", "FORMAT DATE",
})
_MAX_NIF, _MAX_NNI = 9, 10


def _norm(v: Any) -> str:
    return " ".join(re.sub(r"[^A-Z0-9]+", " ", normaliser_entete(str(v or ""))).split())


@lru_cache(maxsize=1)
def listes_observees() -> dict:
    return json.loads(_LISTES.read_text(encoding="utf-8"))


def est_feuille_technique(titre: str) -> bool:
    return _norm(titre) in {_norm(x) for x in FEUILLES_TECHNIQUES}


def est_feuille_clients(titre: str) -> bool:
    return _norm(titre) in {_norm(x) for x in FEUILLES_CLIENTS}


def est_valeur_consigne(v: Any) -> bool:
    n = _norm(v)
    return bool(n) and (n in _CONSIGNES or n.endswith("CARACTERES MAX") or n.endswith("CARACTERE MAX"))


def est_ligne_consigne(row: tuple) -> bool:
    valeurs = [c for c in row if c is not None and str(c).strip() != ""]
    return bool(valeurs) and all(est_valeur_consigne(c) for c in valeurs)


def date_ouverture_interdite(entete: str | None) -> bool:
    return bool(entete) and _norm(entete) in {_norm(x) for x in INTERDITS_DATE_OUVERTURE}


def _cell(donnees: dict, *aliases: str) -> str | None:
    index = {_norm(k): k for k in donnees}
    for a in aliases:
        k = index.get(_norm(a))
        if k is not None:
            v = donnees.get(k)
            if v is not None and str(v).strip() != "":
                return str(v).strip()
    return None


def valeur_source(donnees: dict | None, *aliases: str) -> str | None:
    """Valeur brute (non vide) d'une colonne du fichier, en-tête comparé sans casse ni accents."""
    return _cell(donnees or {}, *aliases)


def _index_secteurs() -> tuple[dict[str, tuple[str, str]], dict[str, str]]:
    par_code: dict[str, tuple[str, str]] = {}
    par_lib: dict[str, str] = {}
    for s in listes_observees().get("secteurs") or []:
        code = str(s.get("code") or "").strip()
        lib, fam = _norm(s.get("libelle")), _norm(s.get("famille"))
        if code:
            par_code[code] = (lib, fam)
            if lib:
                par_lib[lib] = code
    return par_code, par_lib


def _dans(valeur: str | None, cle: str) -> bool:
    if not valeur:
        return True
    attendus = {_norm(x) for x in (listes_observees().get(cle) or [])}
    return not attendus or _norm(valeur) in attendus


def valider_ligne_hebdo(donnees: dict) -> list[str]:
    """Contrôles de format et de cohérence interne. N'invente rien, ne tronque pas."""
    anomalies: list[str] = []
    nif = _cell(donnees, "NIF")
    if nif and len(nif) > _MAX_NIF:
        anomalies.append(f"NIF trop long ({len(nif)} caractères, {_MAX_NIF} max selon le format source) : non tronqué")
    nni = _cell(donnees, "NNI")
    if nni and len(nni) > _MAX_NNI:
        anomalies.append(f"NNI trop long ({len(nni)} caractères, {_MAX_NNI} max selon le format source) : non tronqué")
    date_eer = _cell(donnees, "DATE EER", "DATE D EER")
    if date_eer:
        anomalies.extend(_ctrl_date("DATE EER", date_eer))
    d_classe = _cell(donnees, "DATE CLASSE RISQUE LBC FT", "DATE CLASSE RISQUE")
    if d_classe:
        anomalies.extend(_ctrl_date("Date classe risque LBC FT", d_classe))

    code = _cell(donnees, "CODE SECTEUR BCM", "CODE SECTEUR")
    lib = _cell(donnees, "LIBELLE SECTEUR BCM", "LIBELLE SECTEUR")
    fam = _cell(donnees, "FAMILLE BCM", "FAMILLE SECTEUR BCM")
    if code or lib or fam:
        par_code, par_lib = _index_secteurs()
        code_n = str(int(float(code))) if code and re.fullmatch(r"\d+(\.0+)?", code) else (code or "")
        if code_n and code_n not in par_code:
            anomalies.append(f"Code secteur BCM « {code_n} » hors liste observée (référentiel maître non substitué)")
        elif code_n:
            lib_att, fam_att = par_code[code_n]
            if lib and _norm(lib) != lib_att:
                anomalies.append(f"Libellé secteur BCM incohérent avec le code {code_n}")
            if fam and _norm(fam) != fam_att:
                anomalies.append(f"Famille BCM incohérente avec le code {code_n}")
        elif lib and _norm(lib) not in par_lib:
            anomalies.append("Libellé secteur BCM hors liste observée")

    if not _dans(_cell(donnees, "PROFIL"), "profils"):
        anomalies.append("PROFIL hors liste observée")
    if not _dans(_cell(donnees, "SOUS PROFIL", "SOUS_PROFIL"), "sous_profils"):
        anomalies.append("SOUS PROFIL hors liste observée")
    if not _dans(_cell(donnees, "CATEGORIE JURIDIQUE"), "categories_juridiques"):
        anomalies.append("Catégorie juridique hors liste observée (liste source hétérogène)")
    if not _dans(_cell(donnees, "STATUT COMPTE"), "statuts_compte"):
        anomalies.append("Statut compte hors liste observée — sémantique non assimilée automatiquement")
    if not _dans(_cell(donnees, "STATUT GEL"), "statuts_gel"):
        anomalies.append("Statut Gel hors liste observée — distinct du statut compte")
    if _cell(donnees, "STATUT COMPTE") or _cell(donnees, "STATUT GEL"):
        if not _cell(donnees, "COMPTE", "RIB", "NUMERO COMPTE"):
            anomalies.append("Statut compte / gel sans COMPTE ni RIB : aucune agrégation multi-comptes n'est déduite")

    nat = _cell(donnees, "NATIONNALITE", "NATIONALITE")
    res = _cell(donnees, "RESIDENCE")
    if nat and _norm(nat) == "LISTE PAYS":
        anomalies.append("Nationnalité : valeur de consigne, pas un pays")
    if res and _norm(res) == "LISTE PAYS":
        anomalies.append("Résidence : valeur de consigne, pas un pays")
    return anomalies


def _ctrl_date(libelle: str, brut: str) -> list[str]:
    if est_valeur_consigne(brut):
        return [f"{libelle} : valeur de consigne, pas une date"]
    return []


def colonnes_export(type_lot: str, entetes: list[str] | None) -> list[str]:
    """En-têtes du fichier exporté : 22 colonnes du modèle pour une liste hebdomadaire."""
    if type_lot != "LISTE":
        exclus = {_norm(x) for x in COLONNES_COMPLEMENT}
        return [e for e in (entetes or []) if e and _norm(e) not in exclus]
    return list(COLONNES_EXPORT)


def valeur_export(donnees: dict | None, canon: str) -> str | None:
    if not donnees:
        return None
    index = {_norm(k): k for k in donnees}
    k = index.get(_norm(canon))
    if k is None:
        return None
    v = donnees.get(k)
    if v is None or str(v).strip() == "":
        return None
    return str(v).strip()


def ligne_export_hebdo(donnees: dict | None, *, racine: str | None = None,
                       nom: str | None = None) -> list[str | None]:
    """Une ligne CLIENTS : uniquement les 22 colonnes source, sans métadonnées."""
    row: list[str | None] = []
    for canon in COLONNES_EXPORT:
        v = valeur_export(donnees, canon)
        if not v and _norm(canon) == "RACINE CLIENT":
            v = racine
        if not v and _norm(canon) == "NOM CLIENT":
            v = nom
        row.append(v)
    return row


def construire_modele() -> bytes:
    """Classeur propre : CLIENTS / REFERENTIELS / SECTEURS_BCM / NOTICE_IMPORT / LISEZ_MOI."""
    listes = listes_observees()
    wb = Workbook()

    lisez = wb.active
    lisez.title = "LISEZ_MOI"
    lisez.append(["BEA DIGITAL — Modèle d'import hebdomadaire de classification clients"])
    lisez["A1"].font = Font(bold=True, size=14, color="1A5278")
    notices = [
        ("Objectif", "Chaque semaine, déposer l'extraction et les compléments métier. "
         "BEA DIGITAL rapproche les racines, évalue le moteur SCORE et préserve les décisions humaines."),
        ("Feuille à remplir", "CLIENTS uniquement. Une ligne par RACINE CLIENT (pas une ligne par compte)."),
        ("Clé", "RACINE CLIENT = CLIENT ORION, exactement 6 chiffres, texte, zéros initiaux conservés."),
        ("Comptes", "Pas de COMPTE / RIB dans ce format. Statut compte et Statut Gel ne sont pas agrégés automatiquement."),
        ("Classification source", "Classe / motif / date de classe = données source, jamais un remplacement silencieux du moteur."),
        ("Entreprise récente", "Ancienneté > 366 j → 100 ; ≤ 366 j → 10 000. DATE EER n'est pas une date de création."),
        ("Priorité", "Le moteur propose. La Conformité décide. Un import ultérieur n'écrase jamais une décision manuelle."),
        ("Référentiels", "Les listes de ce classeur sont observées dans le fichier source. "
         "Le référentiel maître reste celui de BEA DIGITAL."),
        ("Pays", "Aucune liste de pays n'était fournie dans le classeur source. Ne pas inventer."),
        ("Résultats moteur", "Ne pas remplir SCORE / niveau proposé : générés par BEA DIGITAL à l'export."),
        ("Version", "1.1 — adapté au format hebdomadaire réel (22 colonnes + 2 complémentaires)."),
    ]
    lisez.append(["Rubrique", "Consignes"])
    for a, b in notices:
        lisez.append([a, b])
    lisez.column_dimensions["A"].width = 28
    lisez.column_dimensions["B"].width = 110
    lisez.row_dimensions[1].height = 22

    clients = wb.create_sheet("CLIENTS")
    clients.append(list(COLONNES_CLIENTS))
    for cell in clients[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1A5278")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    clients.row_dimensions[1].height = 32
    clients.freeze_panes = "A2"
    clients.auto_filter.ref = f"A1:{get_column_letter(len(COLONNES_CLIENTS))}1"
    for i, nom in enumerate(COLONNES_CLIENTS, start=1):
        clients.column_dimensions[get_column_letter(i)].width = max(16, min(34, len(nom) + 4))
    clients.column_dimensions["C"].number_format = "@"
    clients.column_dimensions["L"].number_format = "@"
    clients.column_dimensions["M"].number_format = "@"
    oui = DataValidation(type="list", formula1='"Oui,Non"', allow_blank=True)
    clients.add_data_validation(oui)
    oui.add("X2:X5001")

    ref = wb.create_sheet("REFERENTIELS")
    entetes_ref = (
        "PROFIL", "SOUS_PROFIL", "CATEGORIE_JURIDIQUE", "STATUT_COMPTE", "STATUT_GEL",
        "CLASSE_RISQUE_LBC_FT", "MOTIF_RISQUE_LBC_FT", "ANALYSTE_KYC", "PAYS",
    )
    ref.append(entetes_ref)
    ref.append(["Listes observées dans le fichier source — à valider métier. Pas le référentiel maître."]
               + [""] * 7 + ["Absentes du fichier source — référentiel pays BEA DIGITAL."])
    cles = ("profils", "sous_profils", "categories_juridiques", "statuts_compte", "statuts_gel",
            "classes_risque", "motifs", "analystes")
    colonnes = [listes.get(k) or [] for k in cles] + [[]]
    for i in range(max(len(c) for c in colonnes)):
        ref.append([(c[i] if i < len(c) else None) for c in colonnes])
    for cell in ref[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0F3A57")
    for i in range(1, 10):
        ref.column_dimensions[get_column_letter(i)].width = 28

    sec = wb.create_sheet("SECTEURS_BCM")
    sec.append(["CODE_SECTEUR_BCM", "LIBELLE_SECTEUR_BCM", "FAMILLE_BCM", "STATUT"])
    sec.append(["Observé dans le fichier source", "", "",
                "À valider avec le référentiel BCM officiel — pas le maître BEA DIGITAL"])
    for s in listes.get("secteurs") or []:
        sec.append([s["code"], s["libelle"], s["famille"], "OBSERVÉ — À VALIDER"])
    for cell in sec[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0F3A57")
    sec.column_dimensions["A"].width = 20
    sec.column_dimensions["B"].width = 70
    sec.column_dimensions["C"].width = 28
    sec.column_dimensions["D"].width = 36

    notice = wb.create_sheet("NOTICE_IMPORT")
    notice.append(["CHAMP", "ORIGINE", "FORMAT", "OBLIGATOIRE", "RÈGLE"])
    for row in _NOTICE:
        notice.append(list(row))
    for cell in notice[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1A5278")
    notice.column_dimensions["A"].width = 32
    notice.column_dimensions["B"].width = 28
    notice.column_dimensions["C"].width = 28
    notice.column_dimensions["D"].width = 22
    notice.column_dimensions["E"].width = 78

    res = wb.create_sheet("RESULTATS_MOTEUR")
    res.append(["Ne pas remplir — champs générés par BEA DIGITAL à l'export du lot."])
    res.append(["CHAMP", "DESCRIPTION"])
    for a, b in RESULTATS_MOTEUR:
        res.append([a, b])
    res.column_dimensions["A"].width = 32
    res.column_dimensions["B"].width = 80
    for sheet in wb.worksheets:
        if sheet.title != "CLIENTS":
            sheet.protection.sheet = True
            sheet.protection.enable()
            for row in sheet.iter_rows():
                for cell in row:
                    cell.protection = Protection(locked=True)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


_NOTICE = (
    ("CODE AGENCE", "Requête ORION", "Texte", "Oui", "Contrôler dans le référentiel agences."),
    ("NOM AGENCE", "Requête ORION", "Texte", "Oui", "Vérifier avec CODE AGENCE."),
    ("RACINE CLIENT", "Requête ORION", "Texte, 6 chiffres", "Oui", "Clé de rapprochement ; jamais un entier."),
    ("NOM CLIENT", "Requête ORION", "Texte", "Oui", "Comparé au référentiel ; aucun écrasement silencieux."),
    ("DATE EER", "Source métier", "Date JJ/MM/AAAA", "À confirmer", "Jamais une date de création d'entreprise."),
    ("PROFIL", "Référentiel métier", "Valeur contrôlée", "À confirmer", "Liste observée ; maître = BEA DIGITAL."),
    ("SOUS PROFIL", "Référentiel métier", "Valeur contrôlée", "À confirmer", "Liste observée."),
    ("Catégorie juridique", "Référentiel métier", "Valeur contrôlée", "À confirmer", "Liste source hétérogène."),
    ("Code Secteur BCM", "Référentiel BCM", "Code", "Selon activité", "Cohérence code / libellé / famille."),
    ("Libellé secteur BCM", "Référentiel BCM", "Texte", "Selon activité", "Doit correspondre au code."),
    ("Famille BCM", "Référentiel BCM", "Texte", "Selon activité", "Doit correspondre au code."),
    ("NIF", "Source client", "Texte, max 9", "À confirmer", "Non tronqué. Anomalie si plus long."),
    ("NNI", "Source client", "Texte, max 10", "À confirmer", "Non tronqué. Anomalie si plus long."),
    ("Nationnalité", "Référentiel pays", "Texte", "À confirmer", "Orthographe source conservée. Liste pays absente du fichier."),
    ("Résidence", "Référentiel pays", "Texte", "À confirmer", "Distincte de la nationalité."),
    ("Statut compte", "ORION consolidé", "Valeur contrôlée", "À confirmer", "Ne pas assimiler fermé / bloqué / clôturé."),
    ("Statut Gel", "Métier / ORION", "Valeur contrôlée", "À confirmer", "Distinct du statut compte."),
    ("Date mise à jour client", "Source métier", "Date JJ/MM/AAAA", "Selon dispos.", "Date métier ≠ date de calcul."),
    ("Date classe risque LBC FT", "Source métier", "Date JJ/MM/AAAA", "Selon dispos.", "Date de la classe source."),
    ("Classe risque LBC FT", "Fichier / métier", "Classe source", "Selon dispos.", "Distincte du niveau calculé."),
    ("Motif de risque LBC FT", "Fichier / métier", "Texte", "Selon dispos.", "Distinct du motif moteur."),
    ("Analyste KYC", "Fichier / métier", "Nom", "Selon dispos.", "Distinct de l'utilisateur importateur."),
    ("DATE_CREATION_ENTREPRISE", "Source fiable", "Date JJ/MM/AAAA", "Si critère applicable", "Ne pas substituer DATE EER."),
    ("ENTREPRISE_OUVERTURE_RECENTE", "Métier ou calcul", "Oui / Non", "Facultatif si date", "Comparé au calcul ; contradiction signalée."),
)

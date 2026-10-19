"""Déclaration SIC sur base : référentiel versionné, compléments KYC, cycle complet, historique, non-modification.

N'ajoute et ne supprime que les racines 997001 à 997003, les versions de référentiel « test-sic » et la
correspondance d'agence créée par le test. La version active éventuelle est restaurée en fin de test.
"""

from __future__ import annotations

import io
import uuid
from datetime import date

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import select, text
from sqlalchemy.exc import DBAPIError
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine
from sqlalchemy.pool import NullPool

from app.core.config import get_settings
from app.core.exceptions import AppError
from app.data.clientele_sic import COLONNES
from app.data.plateforme_catalogue import FUNCTIONAL_PERMISSIONS, ROLE_PERMISSIONS
from app.models import Agence, User
from app.services.clientele.bcm_referentiel import ClienteleBcmReferentielService
from app.services.clientele.sic import ClienteleSicService
from app.services.clientele.service import Ctx

_IN = "('997001','997002','997003')"
REF_DATE = date(2002, 2, 2)
LIBELLE = "test-sic"


@pytest.fixture
async def db():
    engine = create_async_engine(get_settings().database_url, poolclass=NullPool)
    try:
        async with engine.connect() as conn:
            pret = await conn.scalar(text("SELECT to_regclass('public.clientele_declarations_sic') IS NOT NULL"))
    except Exception:
        pret = False
    if not pret:
        await engine.dispose()
        pytest.skip("Tables SIC absentes (migration 20261019_clientele_13)")
    session = async_sessionmaker(engine, expire_on_commit=False)()
    active = await session.scalar(text(
        "SELECT id FROM clientele_bcm_ref_versions WHERE statut = 'ACTIVE' AND libelle NOT LIKE :l"),
        {"l": f"{LIBELLE}%"})
    agence = await session.scalar(select(Agence.id).where(Agence.deleted_at.is_(None)).limit(1))
    if not agence:
        await engine.dispose()
        pytest.skip("aucune agence en base de test")
    mapping_existant = await session.scalar(text(
        "SELECT 1 FROM clientele_sic_agences WHERE agence_id = :a"), {"a": agence})
    await _nettoyer(session, None, None)
    await session.commit()
    session.info.update(agence=agence, active=active)
    try:
        yield session
    finally:
        await session.rollback()
        await _nettoyer(session, active, None if mapping_existant else agence)
        await session.commit()
        await session.close()
        await engine.dispose()


async def _nettoyer(session, active, agence_creee) -> None:
    await session.execute(text("DELETE FROM clientele_declarations_sic WHERE date_reference = :d"), {"d": REF_DATE})
    await session.execute(text("SET LOCAL session_replication_role = replica"))
    await session.execute(text(
        "DELETE FROM clientele_bcm_ref_codes WHERE version_id IN "
        "(SELECT id FROM clientele_bcm_ref_versions WHERE libelle LIKE :l)"), {"l": f"{LIBELLE}%"})
    await session.execute(text("DELETE FROM clientele_bcm_ref_versions WHERE libelle LIKE :l"), {"l": f"{LIBELLE}%"})
    await session.execute(text("SET LOCAL session_replication_role = origin"))
    if active:
        await session.execute(text(
            "UPDATE clientele_bcm_ref_versions SET statut = 'ACTIVE', archivee_le = NULL WHERE id = :i"), {"i": active})
    if agence_creee:
        await session.execute(text("DELETE FROM clientele_sic_agences WHERE agence_id = :a"), {"a": agence_creee})
    await session.execute(text(f"DELETE FROM clientele_kyc WHERE racine_client IN {_IN}"))
    await session.execute(text(f"DELETE FROM clientele_comptes WHERE racine_client IN {_IN}"))
    await session.execute(text(f"DELETE FROM clientele_clients WHERE racine_client IN {_IN}"))


async def _ctx(db) -> Ctx:
    user = await db.scalar(select(User).limit(1))
    if not user:
        pytest.skip("aucun utilisateur en base de test")
    return Ctx(user=user, permissions={"clientele.bcm.view", "clientele.bcm.prepare", "clientele.bcm.valider",
                                       "clientele.bcm.referentiel", "clientele.sic.depot", "clientele.view"})


async def _client(db, racine: str, *, agent: str, nom: str, prenoms: str | None, nni: str | None = None,
                  nif: str | None = None, categorie: str | None = None, secteur: str | None = None) -> None:
    await db.execute(text("""
        INSERT INTO clientele_clients (id, racine_client, raison_sociale, prenoms, statut_resident, nni, nif,
                                       categorie_juridique, agent_economique, secteur_activite,
                                       premiere_extraction, date_extraction, created_at, updated_at)
        VALUES (gen_random_uuid(), :r, :nom, :p, 'R', :nni, :nif, :cat, :ag, :sect, '2026-10-01', '2026-10-01',
                now(), now())
    """), {"r": racine, "nom": nom, "p": prenoms, "nni": nni, "nif": nif, "cat": categorie, "ag": agent,
           "sect": secteur})
    compte = f"{racine}00001"
    await db.execute(text("""
        INSERT INTO clientele_comptes (id, compte, rib, racine_client, agence_id, etat_compte, date_ouverture, devise,
                                       premiere_extraction, date_extraction, created_at, updated_at)
        VALUES (gen_random_uuid(), :c, :rib, :r, :a, 'OUVERT', '2001-05-04', 'MRU', '2026-10-01', '2026-10-01',
                now(), now())
    """), {"c": compte, "rib": f"0000000000{compte}00", "r": racine, "a": db.info["agence"]})


def _referentiel(doublon: bool = False, suffixe: str = "") -> bytes:
    wb = Workbook()
    wb.active.title = "pp"
    ws = wb.create_sheet("ref")
    ws.append(["Code", "Secteur d'activité", "Catégorie"])
    ws.append([101, "Agriculture", "Secteur primaire"])
    ws.append([205, f"Commerce{suffixe}", "Secteur tertiaire"])
    ws.append([1101, "Salarié", "Personnne physique"])
    if doublon:
        ws.append([101, "Doublon", "Secteur primaire"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_permissions_sic_catalogue():
    codes = {c for c, _l, _m in FUNCTIONAL_PERMISSIONS}
    assert {"clientele.bcm.referentiel", "clientele.sic.depot"} <= codes
    assert "clientele.sic.depot" in ROLE_PERMISSIONS["clientele.admin"]
    assert "clientele.sic.depot" not in ROLE_PERMISSIONS["clientele.gestionnaire"]
    assert "clientele.bcm.referentiel" not in ROLE_PERMISSIONS["clientele.lecteur"]


async def test_referentiel_versionne(db):
    c = await _ctx(db)
    ref = ClienteleBcmReferentielService(db, c)
    fichier_doublon = _referentiel(doublon=True)
    v_doublon = await ref.importer(fichier_doublon, "ref.xlsx", f"{LIBELLE} doublon", None)
    assert v_doublon["statut"] == "BROUILLON" and v_doublon["doublons"] == 1
    with pytest.raises(AppError, match="dupliqué"):
        await ref.activer(uuid.UUID(v_doublon["id"]), "test")
    with pytest.raises(AppError, match="déjà importé"):
        await ref.importer(fichier_doublon, "ref.xlsx", f"{LIBELLE} bis", None)

    v1 = await ref.importer(_referentiel(), "ref.xlsx", f"{LIBELLE} v1", None)
    assert (v1["nb_codes"], v1["nb_pp"], v1["nb_pm"]) == (3, 1, 2)
    await ref.activer(uuid.UUID(v1["id"]), "test")
    v2 = await ref.importer(_referentiel(suffixe=" et distribution"), "ref2.xlsx", f"{LIBELLE} v2", None)
    diff = await ref.comparer(uuid.UUID(v1["id"]), uuid.UUID(v2["id"]))
    assert [m["code"] for m in diff["modifies"]] == ["205"] and not diff["ajoutes"] and not diff["retires"]
    await ref.activer(uuid.UUID(v2["id"]), "test")
    versions = {v["id"]: v["statut"] for v in await ref.versions()}
    assert versions[v1["id"]] == "ARCHIVEE" and versions[v2["id"]] == "ACTIVE"
    with pytest.raises(AppError):
        await ref.supprimer(uuid.UUID(v1["id"]))

    codes = await ref.codes(None, None, "PP", 1, 50)
    assert [x["code"] for x in codes["items"]] == ["1101"]

    async with db.bind.connect() as conn:
        with pytest.raises(DBAPIError, match="ni modifiable"):
            await conn.execute(text("UPDATE clientele_bcm_ref_codes SET libelle = 'x' WHERE version_id = :v"),
                               {"v": uuid.UUID(v1["id"])})
        await conn.rollback()


async def test_cycle_sic_complet_sans_modifier_les_clients(db):
    c = await _ctx(db)
    ref = ClienteleBcmReferentielService(db, c)
    v = await ref.importer(_referentiel(), "ref.xlsx", f"{LIBELLE} cycle", None)
    await ref.activer(uuid.UUID(v["id"]), "test")
    await ref.definir_agence(db.info["agence"], "Tevragh Zeina", None)
    await _client(db, "997001", agent="PARTICULIERS", nom="SIDI MOHAMED AHMED", prenoms="[IDP] 1 AHMED",
                  nni="0012345678", categorie="PERSONNE PHYSIQUE")
    await _client(db, "997002", agent="PARTICULIERS", nom="MARIEM", prenoms="[IDP] 2 MARIEM", nni=None)
    await _client(db, "997003", agent="AUTRES SOCIETES", nom="SOCIETE ALPHA", prenoms=None, nif="00098765",
                  categorie="SARL", secteur="DIVERS")
    await db.commit()
    await ref.definir_complement("997001", {"genre": "F", "telephone": "22 33 44 55", "code_csp_bcm": "1101"},
                                 "test")
    with pytest.raises(AppError, match="hors catégorie"):
        await ref.definir_complement("997003", {"code_secteur_bcm": "1101"}, "test")
    await ref.definir_complement("997003", {"code_secteur_bcm": "205"}, "test")
    avant = (await db.execute(text(
        f"SELECT racine_client, raison_sociale, prenoms, nni, nif, categorie_juridique, updated_at "
        f"FROM clientele_clients WHERE racine_client IN {_IN} ORDER BY 1"))).all()
    classes_avant = await db.scalar(text(
        f"SELECT count(*) FROM clientele_classif_historique WHERE racine_client IN {_IN}"))

    sic = ClienteleSicService(db, c)
    pp = await sic.creer("PP", REF_DATE, None, prefixe="997")
    assert pp["nb_candidats"] == 2 and pp["nb_lignes"] == 1 and pp["nb_exclus"] == 1
    assert pp["statut"] == "A_CORRIGER" and pp["ref_version_id"] == v["id"]
    exclus = await sic.lignes(uuid.UUID(pp["id"]), "EXCLUE", None, None, None, 1, 50)
    assert [x["racine"] for x in exclus["items"]] == ["997002"]
    assert {a["champ"] for a in exclus["items"][0]["anomalies"]} >= {"NNI", "Genre", "NumTel", "Nom"}

    with pytest.raises(AppError, match="commentaire"):
        await sic.valider(uuid.UUID(pp["id"]), None)
    with pytest.raises(AppError):
        await sic.deposer(uuid.UUID(pp["id"]), date.today(), "ACCEPTE", None)
    pp = await sic.valider(uuid.UUID(pp["id"]), "Exclusion de 997002 acceptée (test)")
    assert pp["statut"] == "VALIDEE_EXPORT"
    pp = await sic.generer(uuid.UUID(pp["id"]))
    assert pp["statut"] == "GENEREE" and pp["controle_fichier"]["ok"]
    contenu, nom = await sic.fichier(uuid.UUID(pp["id"]))
    assert nom == "SIC_PP_BEA_20020202.xlsx"
    ws = load_workbook(io.BytesIO(contenu))["pp"]
    assert [x.value for x in ws[1]] == list(COLONNES["PP"])
    assert [x.value for x in ws[2]] == ["997001", "MR", "0012345678", "AHMED", "SIDI MOHAMED", "F", "1101",
                                        "22334455", "04/05/2001", "Tevragh Zeina"]
    with pytest.raises(AppError):
        await sic.controler(uuid.UUID(pp["id"]))
    rapport, _ = await sic.rapport(uuid.UUID(pp["id"]))
    assert load_workbook(io.BytesIO(rapport))["Anomalies"]["A2"].value == "997002"
    pp = await sic.deposer(uuid.UUID(pp["id"]), date.today(), "ACCEPTE", None)
    assert pp["statut"] == "DEPOSEE" and pp["depot_resultat"] == "ACCEPTE"

    pm = await sic.creer("PM", REF_DATE, None, prefixe="997")
    assert pm["nb_candidats"] == 1 and pm["nb_lignes"] == 1
    [ligne] = (await sic.lignes(uuid.UUID(pm["id"]), None, None, None, None, 1, 10))["items"]
    assert ligne["valeurs"][:4] == ["997003", "00098765", "SOCIETE ALPHA", "205"]

    historique = [h for h in await sic.historique() if h["id"] in (pp["id"], pm["id"])]
    assert {h["type"] for h in historique} == {"Déclaration SIC PP", "Déclaration SIC PM"}
    assert all(h["ref_version"] == f"{LIBELLE} cycle" for h in historique)

    apres = (await db.execute(text(
        f"SELECT racine_client, raison_sociale, prenoms, nni, nif, categorie_juridique, updated_at "
        f"FROM clientele_clients WHERE racine_client IN {_IN} ORDER BY 1"))).all()
    assert apres == avant
    assert await db.scalar(text(
        f"SELECT count(*) FROM clientele_classif_historique WHERE racine_client IN {_IN}")) == classes_avant
    with pytest.raises(AppError):
        await ref.supprimer(uuid.UUID(v["id"]))

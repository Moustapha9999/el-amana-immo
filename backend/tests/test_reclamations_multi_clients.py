# ruff: noqa: F811
"""Réclamations — dossier multi-comptes / multi-clients et import Excel des opérations.

Clients et comptes réels de la base clientèle (transaction annulée : rien n'est laissé en base).
"""

from __future__ import annotations

import uuid
from decimal import Decimal
from io import BytesIO

import pytest
from fastapi import HTTPException, UploadFile
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select

from app.core.exceptions import AppError
from app.models import ClienteleCompte, GedDocument, ReclamationParametre
from app.services.ged_service import GedService
from app.services.reclamations import constantes as C
from app.services.reclamations import dossiers, imports
from app.services.reclamations.acces import ReclamationsScope
from tests.test_reclamations_dossiers import _code, _saisie, _sr
from tests.test_reclamations_fondations import _unite, _user, db  # noqa: F401


async def _refus(coro, attendu: str) -> None:
    with pytest.raises((AppError, HTTPException)) as e:
        await coro
    assert _code(e.value) == attendu, e.value


async def _clients(db) -> tuple[str, list[str], str, str]:
    """Client A (2 comptes) et client B (1 compte)."""
    a = await db.scalar(select(ClienteleCompte.racine_client).group_by(ClienteleCompte.racine_client)
                        .having(func.count() >= 2).order_by(ClienteleCompte.racine_client).limit(1))
    if a is None:
        pytest.skip("base clientèle vide")
    comptes_a = list((await db.scalars(select(ClienteleCompte.compte).where(
        ClienteleCompte.racine_client == a).order_by(ClienteleCompte.compte).limit(2))).all())
    b, compte_b = (await db.execute(select(ClienteleCompte.racine_client, ClienteleCompte.compte).where(
        ClienteleCompte.racine_client != a).order_by(ClienteleCompte.compte).limit(1))).one()
    return a, comptes_a, b, compte_b


def _op(ref: str, montant: str, compte: str | None = None, racine: str | None = None, **extra) -> dict:
    return {"reference": ref, "montant": Decimal(montant), "compte_reference": compte, "client_racine": racine,
            **extra}


async def test_un_client_plusieurs_comptes(db):
    s = await _sr(db)
    unite = await _unite(db)
    a, (c1, c2), _b, _cb = await _clients(db)
    d = await dossiers.creer(db, s, await _saisie(db, unite.id, client_racine=a, operations=[
        _op("R1", "100", c1), _op("R2", "50", c1, montant_conteste=Decimal("20")), _op("R3", "30", c2)]))
    assert d["nb_clients"] == 1 and len(d["participants"]) == 1
    assert {c["compte_reference"]: c["nb_operations"] for c in d["comptes"]} == {c1: 2, c2: 1}
    assert Decimal(d["montant_total"]) == Decimal("180.00")
    assert Decimal(d["montant_conteste_total"]) == Decimal("150.00")
    assert all(c["statut"] == "VERIFIE" for c in d["comptes"])


async def test_plusieurs_clients_justification_et_masquage(db):
    s = await _sr(db)
    unite = await _unite(db)
    a, (c1, _c2), b, cb = await _clients(db)
    d = await dossiers.creer(db, s, await _saisie(db, unite.id, client_racine=a, operations=[_op("R1", "100", c1)]))
    rid = uuid.UUID(d["id"])
    lignes = [_op("R1", "100", c1), _op("R9", "70", cb)]
    # Compte de B : rattaché à B (titulaire) → 2 clients → justification obligatoire.
    await _refus(dossiers.remplacer_operations(db, s, rid, {"revision": d["revision"], "operations": lignes}),
                 "JUSTIFICATION_REGROUPEMENT_OBLIGATOIRE")
    # Compte de B déclaré sous le client A : incohérent.
    await _refus(dossiers.remplacer_operations(db, s, rid, {"revision": d["revision"], "operations": [
        _op("R9", "70", cb, a)]}), "COMPTE_AUTRE_CLIENT")
    d = await dossiers.remplacer_operations(db, s, rid, {"revision": d["revision"], "operations": lignes,
                                                        "justification_regroupement": "Même fraude, deux clients"})
    assert d["nb_clients"] == 2 and d["justification_regroupement"] == "Même fraude, deux clients"
    associe = next(p for p in d["participants"] if p["role"] == "ASSOCIE")
    assert associe["racine"] == b and not associe["masque"]

    # Accès par le seul périmètre organisationnel : client associé et ses comptes masqués.
    lecteur = ReclamationsScope(user=await _user(db), permissions=frozenset({"reclamations.view"}),
                                unites_perimetre_org=frozenset({unite.id}))
    vue = await dossiers.detail(db, lecteur, await dossiers.charger(db, lecteur, rid))
    masque = next(p for p in vue["participants"] if p["role"] == "ASSOCIE")
    assert masque["masque"] and masque["racine"] != b and vue["clients_masques"]
    assert all(o["compte_reference"] != cb for o in vue["operations"])
    assert not dossiers.actions(lecteur, await dossiers.charger(db, lecteur, rid))["modifier_operations"]

    # Regroupement interdit par la banque : refusé.
    p = await db.get(ReclamationParametre, C.PARAM_MULTI_CLIENTS)
    if p is None:
        pytest.skip("paramètre absent")
    p.valeur = False
    await db.flush()
    await _refus(dossiers.remplacer_operations(db, s, rid, {"revision": d["revision"], "operations": [
        *lignes, _op("R10", "5", cb)]}), "REGROUPEMENT_NON_AUTORISE")


# En-têtes de l'ancien modèle (libellés français) : toujours reconnus.
ANCIENS_ENTETES = ["Identifiant client", "Référence du compte", "Référence de l'opération", "Date de l'opération",
                   "Montant de l'opération", "Montant contesté", "Motif détaillé"]


def _classeur(lignes: list[list]) -> UploadFile:
    wb = Workbook()
    ws = wb.active
    ws.title = "Opérations"
    ws.append(ANCIENS_ENTETES)
    for ligne in lignes:
        ws.append(ligne)
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return UploadFile(file=buf, filename="operations.xlsx")


async def test_modele_import():
    wb = load_workbook(BytesIO(imports.modele()))
    assert wb.sheetnames == ["Opérations", "Consignes"]
    assert [c.value for c in wb["Opérations"][1]] == [code for _c, code, *_ in imports.colonnes_du_modele()]
    assert wb["Opérations"].max_row >= 1 and wb["Opérations"]["A2"].value is None


async def test_import_analyse_puis_confirmation(db):
    s = await _sr(db)
    unite = await _unite(db)
    a, (c1, c2), b, cb = await _clients(db)
    d = await dossiers.creer(db, s, await _saisie(db, unite.id, client_racine=a))
    rid = uuid.UUID(d["id"])
    fichiers: list[str] = []
    try:
        res = await imports.analyser(db, s, rid, _classeur([
            [a, c1, "OP1", "01/10/2026", 100, None, "Débit non autorisé"],
            [None, c1, "OP2", None, "1 250,50", 250, None],
            [None, c2, "OP1", None, 100, None, None],          # même référence, autre compte : avertissement
            [None, c2, "OP1", None, 100, None, None],          # ligne identique : erreur
            [None, None, None, None, 10, None, None],          # référence absente : erreur
            [b, cb, "OPB", None, 40, 50, None],                # contesté > montant : erreur
        ]), "AJOUT")
        doc = await db.get(GedDocument, uuid.UUID(res["document_id"]))
        fichiers.append(doc.stored_path)
        assert (res["nb_lignes"], res["nb_erreurs"]) == (6, 3)
        statuts = [x["statut"] for x in res["lignes"]]
        assert statuts[3] == "ERREUR" and statuts[4] == "ERREUR" and statuts[5] == "ERREUR"
        assert res["lignes"][2]["avertissements"]
        assert res["clients"][0]["principal"] and res["clients"][0]["racine"] == a
        await _refus(imports.confirmer(db, s, rid, uuid.UUID(res["id"]), {"revision": res["revision"]}),
                     "IMPORT_ERREURS")
        await imports.abandonner(db, s, rid, uuid.UUID(res["id"]))

        res = await imports.analyser(db, s, rid, _classeur([
            [a, c1, "OP1", "01/10/2026", 100, None, None],
            [None, c2, "OP2", None, "1 250,50", 250, None],
            [b, cb, "OPB", None, 40, None, None],
        ]), "AJOUT")
        doc = await db.get(GedDocument, uuid.UUID(res["document_id"]))
        fichiers.append(doc.stored_path)
        assert res["nb_erreurs"] == 0 and res["multi_clients"] and res["justification_requise"]
        assert Decimal(res["apres_import"]["montant"]) == Decimal("1390.50")
        iid = uuid.UUID(res["id"])
        await _refus(imports.confirmer(db, s, rid, iid, {"revision": res["revision"]}),
                     "JUSTIFICATION_REGROUPEMENT_OBLIGATOIRE")
        await imports.confirmer(db, s, rid, iid, {"revision": res["revision"],
                                                  "justification_regroupement": "Opérations liées"})
        fiche = await dossiers.detail(db, s, await dossiers.charger(db, s, rid))
        assert fiche["nb_operations"] == 3 and fiche["nb_clients"] == 2
        assert {o["import_id"] for o in fiche["operations"]} == {str(iid)}
        assert {i["id"]: i["statut"] for i in fiche["imports"]}[str(iid)] == "CONFIRME"
        await _refus(imports.confirmer(db, s, rid, iid, {"revision": fiche["revision"]}), "IMPORT_DEJA_TRAITE")
    finally:
        ged = GedService(db)
        for chemin in fichiers:
            ged.absolute_path(chemin).unlink(missing_ok=True)

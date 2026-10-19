"""Export Excel BCM au format du modèle officiel (sans base de données)."""

from __future__ import annotations

from datetime import datetime
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.data.clientele_bcm import tableaux
from app.services.clientele.declaration_modele import (
    A_CONFIGURER, BLOCS, FEUILLE, NON_CALCULABLE, SOURCE_ABSENTE, classeur_modele, positions,
)

MODELE = Path(__file__).resolve().parents[2] / "Déclaration mensuelle - BEA - août - 2026.xlsx"
META = {"periode": "Août 2026", "statut": "CALCULEE", "genere_le": datetime(2026, 10, 19, 11, 40),
        "genere_par": "Analyste Test", "moteur_version": "m1", "grille_version": "g1"}


def _cellules(valeurs: dict[tuple[str, str, str], int | float | None],
              a_configurer: set[tuple[str, str]] = frozenset()) -> dict:
    """Cellules au format de ``ClienteleDeclarationService._assembler`` pour la grille courante."""
    out = []
    for tab in tableaux():
        lignes = []
        for ligne in tab["lignes"]:
            cells = {}
            for col in tab["colonnes"]:
                cle = (tab["code"], ligne["code"], col["code"])
                absente = ligne["statut"] == "A_CONFIGURER"
                ac = (tab["code"], ligne["code"]) in a_configurer or absente
                cells[col["code"]] = {"code": ".".join(cle), "statut": "A_CONFIGURER" if ac else ligne["statut"],
                                      "valeur": None if ac else valeurs.get(cle, 0), "reserve": ligne["reserve"],
                                      "fiabilite": "SOURCE_ABSENTE" if absente else None}
            lignes.append({"code": ligne["code"], "cellules": cells})
        out.append({"code": tab["code"], "lignes": lignes})
    return {"tableaux": out}


def _ws(contenu: bytes, feuille: str = FEUILLE):
    return load_workbook(BytesIO(contenu))[feuille]


def test_blocs_couvrent_toute_la_grille():
    grille = {(t["code"], ligne["code"]) for t in tableaux() for ligne in t["lignes"]}
    modele = {(b["tableau"], code) for b in BLOCS for code, _l in b["lignes"]}
    assert grille == modele
    for b in BLOCS:
        tab = next(t for t in tableaux() if t["code"] == b["tableau"])
        assert b["codes"] == tuple(c["code"] for c in tab["colonnes"])


@pytest.mark.skipif(not MODELE.exists(), reason="Modèle Excel BCM absent à la racine du dépôt")
def test_structure_identique_au_modele():
    ref = load_workbook(MODELE).worksheets[0]
    ws = _ws(classeur_modele(_cellules({}), [], META))
    assert ws.title == ref.title
    assert {str(r) for r in ws.merged_cells.ranges} == {str(r) for r in ref.merged_cells.ranges}
    chiffrees = positions()
    for ligne in ref.iter_rows(min_row=1, max_row=ref.max_row):
        for c in ligne:
            if c.coordinate in chiffrees:
                assert ws[c.coordinate].value is not None, c.coordinate
                continue
            attendu = c.value.strip() if isinstance(c.value, str) else c.value
            obtenu = ws[c.coordinate].value
            assert (obtenu.strip() if isinstance(obtenu, str) else obtenu) == attendu, c.coordinate
            assert bool(ws[c.coordinate].font.bold) == bool(c.font.bold), c.coordinate
    cellules_ref = {c.coordinate for row in ref.iter_rows() for c in row
                    if c.value is not None and c.column > 1 and not c.font.bold}
    assert cellules_ref == set(chiffrees)
    for col in "ABCDE":
        assert ws.column_dimensions[col].width == pytest.approx(ref.column_dimensions[col].width, abs=0.1)


def test_valeurs_calculees_aux_bonnes_positions():
    valeurs = {("T1A", "stock_m1", "pp"): 34426, ("T1A", "stock_m1", "pm"): 5692,
               ("T1A", "stock_m1", "total"): 40118, ("T1A", "maj_mois", "total"): 292,
               ("T1B", "stock_m1", "eleve"): 678, ("T4", "eleve", "valeur"): 680,
               ("T4", "eleve", "ratio"): 1.68, ("T4", "total", "valeur"): 40410, ("T4", "total", "ratio"): 100.0}
    ws = _ws(classeur_modele(_cellules(valeurs), [], META))
    assert (ws["B5"].value, ws["C5"].value, ws["E5"].value) == (34426, 5692, 40118)
    assert ws["E7"].value == 292 and ws["B12"].value == 678
    assert (ws["B41"].value, ws["C41"].value) == (680, 1.68)
    assert (ws["B44"].value, ws["C44"].value) == (40410, 100.0)
    assert ws["C41"].number_format == "0.00" and ws["B5"].number_format == "#,##0"


def test_a_configurer_jamais_un_nombre():
    ws = _ws(classeur_modele(_cellules({("T1A", "cible", "total"): 589, ("T2", "extraites", "valeur"): 2492},
                                       {("T1A", "cible"), ("T1A", "maj_5ans"), ("T1B", "cible")}), [], META))
    for ref in ("B6", "E6", "B8", "B13"):
        assert ws[ref].value == A_CONFIGURER, ref
    for ref in ("B21", "B24", "B25", "B31", "B35"):
        assert ws[ref].value == SOURCE_ABSENTE, ref


def test_mois_sans_activite_zeros_et_ratio_non_calculable():
    cellules = _cellules({})
    for tab in cellules["tableaux"]:
        if tab["code"] == "T4":
            for ligne in tab["lignes"]:
                ligne["cellules"]["ratio"]["valeur"] = None
    ws = _ws(classeur_modele(cellules, [], META))
    assert (ws["B5"].value, ws["E7"].value, ws["B44"].value) == (0, 0, 0)
    assert ws["C44"].value == NON_CALCULABLE


def test_feuille_controle_tracabilite():
    controles = [{"code": "t4.somme", "etat": "AVEC_RESERVE", "libelle": "T4", "bloquant": False}]
    meta = {**META, "regles": [{"code": "actif.valeurs_gel", "statut": "A_VALIDER", "libelle": "Gel actif",
                                "valeur": ["ACTIF"], "origine": "Proposition"}],
            "exceptions": {"forme": {"EXC_FORME_VIDE": 12}, "risque_fin": {"INTERDIT": 3}, "doublons": 0}}
    ws = _ws(classeur_modele(_cellules({}, {("T1A", "cible")}), controles, meta), "Contrôle")
    textes = [c.value for row in ws.iter_rows() for c in row if c.value is not None]
    assert "Août 2026" in textes and "Analyste Test" in textes and "19/10/2026 11:40" in textes
    assert any(isinstance(t, str) and "Aucune transmission automatique" in t for t in textes)
    assert "B6" in textes and "B21" in textes and "t4.somme" in textes
    assert "actif.valeurs_gel" in textes and "À valider" in textes and "ACTIF" in textes
    assert "Forme juridique : EXC_FORME_VIDE" in textes and "Risque fin de mois : INTERDIT" in textes

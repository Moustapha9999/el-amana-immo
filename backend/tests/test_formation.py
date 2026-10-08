"""Formation & Sensibilisation — normalisation, import Excel, exports, règles de session.

Les tests base tournent dans une transaction externe annulée à la fin : aucune donnée laissée.
Ignorés si la migration ``20261007_formation_module`` n'est pas appliquée.
"""

from __future__ import annotations

import uuid
from datetime import date, datetime, timedelta
from io import BytesIO

import pytest
from fastapi import HTTPException, UploadFile
from openpyxl import Workbook, load_workbook
from sqlalchemy import select, text
from sqlalchemy.ext.asyncio import AsyncSession, create_async_engine
from sqlalchemy.pool import NullPool

from app.core.config import get_settings
from app.core.exceptions import AppError
from app.data.module_backup_scopes import ESPACE_MODULES, MODULE_BACKUP_SCOPES
from app.data.plateforme_catalogue import FUNCTIONAL_PERMISSIONS, PLATEFORME_MODULES, ROLE_PERMISSIONS
from app.models import AuditLog, FormationReferentiel, GedDocument, User
from app.services import formation_documents, formation_export
from app.services.ged_service import GedService
from app.services.formation_import import SEP_FORMATEURS, SEP_THEMES, FormationImportService, _detecter_colonnes
from app.services.formation_reporting import Filtres, FormationReporting
from app.services.formation_service import Ctx, FormationService, cle, cle_identite


# ----------------------------------------------------------------------- unitaires
def test_cle_et_identite():
    assert cle("  Procédure  EER ") == "procedure eer"
    assert cle("LBC-FT") == cle("LBC FT") == "lbc ft"
    assert cle_identite("Mohamed", "Ahmed") == cle_identite("AHMED Mohamed") == "ahmed mohamed"
    assert cle_identite("Ould Cheikh", "Sidi") == cle_identite("Sidi OULD CHEIKH")


def test_catalogue_formation():
    module = next(m for m in PLATEFORME_MODULES if m["code"] == "formation")
    assert module["domaine_code"] == "conformite-securite-financiere"
    assert module["entry_path"].startswith("/formation/")
    perms = {c for c, _l, m in FUNCTIONAL_PERMISSIONS if m == "formation"}
    assert {"formation.view", "formation.attendance.manage", "formation.import.execute",
            "formation.reporting.export"} <= perms
    for role in ("formation.lecteur", "formation.gestionnaire", "formation.admin"):
        assert set(ROLE_PERMISSIONS[role]) <= perms
    assert "formation.attendance.manage" not in ROLE_PERMISSIONS["formation.lecteur"]
    assert "formation.references.manage" not in ROLE_PERMISSIONS["formation.gestionnaire"]
    assert "formation" in ESPACE_MODULES["audit-controle-conformite"]
    assert "formation_participants" in MODULE_BACKUP_SCOPES["formation"]["exclusive_tables"]


def test_separateurs_multi_valeurs():
    assert SEP_THEMES.split("LBC FT/Procédure ouverture de compte/projet fiabilisation") == [
        "LBC FT", "Procédure ouverture de compte", "projet fiabilisation"]
    assert [t for t in SEP_THEMES.split("Procédure EER, Fiabilisation, Remediation") if t] == [
        "Procédure EER", "Fiabilisation", "Remediation"]
    assert SEP_FORMATEURS.split("Abass Ngam / Fatimata Thiam") == ["Abass Ngam", "Fatimata Thiam"]


def test_detection_colonnes_excel_historique():
    entetes = ["Date", "Thème", "Lieu", "Formateur", "Nom du participant", "Fonction", "Entité", "Périmètre"]
    cols = _detecter_colonnes(entetes)
    assert cols == {"date": 0, "theme": 1, "lieu": 2, "formateur": 3, "participant": 4, "fonction": 5,
                    "entite": 6, "perimetre": 7}
    assert "nom" not in cols


def _session_exemple(n: int = 3) -> dict:
    return {
        "reference": "FOR-2026-0001", "intitule": None, "date_session": "2026-03-12",
        "theme_libelle": "LBC FT, FATCA", "formateur_libelle": "Abass Ngam / Yero Dieng",
        "lieu": {"libelle": "Nouakchott"},
        "participants": [{"nom_complet": f"Employé {i}"} for i in range(n)],
    }


def test_feuille_presence_pdf_et_excel():
    pdf = formation_export.feuille_presence_pdf(_session_exemple())
    assert pdf.startswith(b"%PDF") and len(pdf) > 1500
    assert formation_export.feuille_presence_pdf(_session_exemple(0)).startswith(b"%PDF")
    wb = load_workbook(BytesIO(formation_export.feuille_presence_excel(_session_exemple())))
    valeurs = [c for row in wb.active.iter_rows(values_only=True) for c in row if c]
    for attendu in ("BANQUE EL AMANA", "FORMATION & SENSIBILISATION", "Nom et prénom", "Signature", "Employé 2"):
        assert attendu in valeurs


def _rapport_exemple() -> dict:
    ligne = {"libelle": "LBC FT", "formations": 2, "participants": 5, "presents": 4, "absents": 1, "employes": 5,
             "taux": 80.0}
    part = {"session_id": "x", "reference": "FOR-2026-0001", "date": "2026-03-12", "theme": "LBC FT",
            "lieu": "Nouakchott", "formateur": "A", "employe_id": "e", "nom": "Ba", "prenom": "Amadou",
            "nom_complet": "Amadou Ba", "fonction": "Caissier", "entite": "AGENCE KSAR",
            "perimetre": "DIRECTION COMMERCIALE", "presence": "PRESENT", "presence_libelle": "Présent"}
    return {
        "filtres": {"annee": 2026}, "kpis": {
            "formations": 2, "formations_realisees": 2, "formations_a_venir": 0, "participations": 5,
            "presents": 4, "absents": 1, "non_saisis": 0, "taux_presence": 80.0, "employes_concernes": 5,
            "employes_formes": 4, "employes_actifs": 10, "themes_couverts": 1},
        "par_annee": [{**ligne, "libelle": "2026"}], "par_mois": [], "par_theme": [ligne],
        "par_entite": [{**ligne, "libelle": "AGENCE KSAR"}], "par_perimetre": [{**ligne, "libelle": "DC"}],
        "par_lieu": [], "par_formateur": [], "par_fonction": [],
        "presence": [{"libelle": "Présents", "valeur": 4, "cle": "PRESENT"},
                     {"libelle": "Absents", "valeur": 1, "cle": "ABSENT"},
                     {"libelle": "Non saisis", "valeur": 0, "cle": "NON_SAISI"}],
        "formations": [{"id": "x", "reference": "FOR-2026-0001", "date": "2026-03-12", "intitule": None,
                        "theme": "LBC FT", "lieu": "Nouakchott", "formateur": "A", "statut": "CLOTUREE",
                        "statut_libelle": "Clôturée", "participants": 5, "presents": 4, "absents": 1,
                        "non_saisis": 0, "taux": 80.0}],
        "participations": [part, {**part, "presence": "ABSENT", "presence_libelle": "Absent"}],
        "participations_total": 2,
    }


def test_rapports_pdf_et_excel():
    data = _rapport_exemple()
    assert formation_export.rapport_pdf(data, "Année 2026").startswith(b"%PDF")
    wb = load_workbook(BytesIO(formation_export.rapport_excel(data, "Année 2026")))
    assert wb.sheetnames == ["Synthèse", "Formations", "Participants", "Présences", "Absences",
                             "Par_Thème", "Par_Entité", "Par_Périmètre"]


# ----------------------------------------------------------------------- base
@pytest.fixture
async def db():
    engine = create_async_engine(get_settings().database_url, poolclass=NullPool)
    try:
        async with engine.connect() as conn:
            pret = await conn.scalar(text("SELECT to_regclass('public.formation_sessions') IS NOT NULL"))
    except Exception:
        pret = False
    if not pret:
        await engine.dispose()
        pytest.skip("Migration 20261007_formation_module non appliquée")
    conn = await engine.connect()
    trans = await conn.begin()
    session = AsyncSession(bind=conn, join_transaction_mode="create_savepoint", expire_on_commit=False)
    try:
        yield session
    finally:
        await session.close()
        await trans.rollback()
        await conn.close()
        await engine.dispose()


async def _ctx(db: AsyncSession, *perms: str) -> Ctx:
    user = await db.scalar(select(User).where(User.is_superuser.is_(True), User.deleted_at.is_(None)).limit(1))
    if user is None:
        pytest.skip("Aucun super-utilisateur")
    if perms:
        # Utilisateur non superuser simulé : permissions explicites seulement.
        clone = User(id=user.id, email=user.email, full_name=user.full_name, hashed_password="x", is_superuser=False)
        return Ctx(user=clone, permissions=set(perms))
    return Ctx(user=user, permissions={"*"})


async def _ref(db: AsyncSession, domaine: str) -> str:
    return str(await db.scalar(select(FormationReferentiel.id).where(
        FormationReferentiel.domaine == domaine, FormationReferentiel.actif.is_(True)).limit(1)))


async def _employe(svc: FormationService, nom: str) -> dict:
    entite = (await svc.lister_entites())[0]
    return await svc.creer_employe({"nom": nom, "prenom": "Test", "entite_id": entite["id"]}, forcer=True)


@pytest.mark.asyncio
async def test_cycle_session_presences_cloture(db):
    svc = FormationService(db, await _ctx(db))
    suffixe = uuid.uuid4().hex[:6].upper()
    e1, e2 = await _employe(svc, f"ZZA{suffixe}"), await _employe(svc, f"ZZB{suffixe}")
    assert e1["perimetre"], "le périmètre est déduit de l'entité"

    s = await svc.creer_session({
        "date_session": date.today() - timedelta(days=1), "theme_ids": [await _ref(db, "THEME")],
        "lieu_id": await _ref(db, "LIEU"), "formateur_ids": [await _ref(db, "FORMATEUR")],
        "employe_ids": [e1["id"], e2["id"]],
    })
    assert s["reference"].startswith("FOR-") and s["statut"] == "PLANIFIEE"
    assert all(p["presence"] is None for p in s["participants"]), "personne n'est présent par défaut"
    assert s["a_saisir"] and not s["actions"]["cloturer"]

    p1, p2 = (p["id"] for p in s["participants"])
    with pytest.raises(HTTPException) as conflit:
        await svc.saisir_presences(uuid.UUID(s["id"]), {p1: "PRESENT"}, revision=s["revision"] + 5)
    assert conflit.value.status_code == 409

    s = await svc.saisir_presences(uuid.UUID(s["id"]), {p1: "PRESENT"}, revision=s["revision"])
    assert s["statut"] == "PLANIFIEE" and s["stats"]["non_saisis"] == 1
    s = await svc.saisir_presences(uuid.UUID(s["id"]), {p2: "ABSENT"}, revision=s["revision"])
    assert s["statut"] == "REALISEE" and s["stats"]["taux_presence"] == 50.0

    with pytest.raises(AppError) as motif:
        await svc.saisir_presences(uuid.UUID(s["id"]), {p2: "PRESENT"}, revision=s["revision"])
    assert motif.value.code == "MOTIF_OBLIGATOIRE"

    assert not s["actions"]["annuler"], "annulation interdite après saisie des présences"
    s = await svc.changer_statut(uuid.UUID(s["id"]), "cloturer", s["revision"], None)
    assert s["statut"] == "CLOTUREE" and not s["actions"]["modifier"]
    with pytest.raises(AppError):
        await svc.modifier_session(uuid.UUID(s["id"]), {"revision": s["revision"], "intitule": "x"})
    gestionnaire = FormationService(db, await _ctx(db, "formation.view", "formation.update", "formation.close"))
    with pytest.raises(HTTPException):
        await gestionnaire.supprimer_session(uuid.UUID(s["id"]), "test")
    assert not (await gestionnaire.detail_session(uuid.UUID(s["id"])))["actions"]["supprimer"]

    fiche = await svc.fiche_employe(uuid.UUID(e1["id"]))
    assert fiche["stats"]["presents"] == 1 and fiche["historique"][0]["presence"] == "PRESENT"

    rep = await FormationReporting(db).reporting(Filtres(employe_id=uuid.UUID(e2["id"])))
    assert rep["kpis"]["absents"] == 1 and rep["kpis"]["formations"] == 1

    actions = set((await db.scalars(select(AuditLog.action).where(AuditLog.entity_id == s["id"]))).all())
    assert {"formation.session.create", "formation.presence.update", "formation.session.close"} <= actions


@pytest.mark.asyncio
async def test_presence_future_et_annulation(db):
    svc = FormationService(db, await _ctx(db))
    e = await _employe(svc, f"ZZF{uuid.uuid4().hex[:6].upper()}")
    s = await svc.creer_session({
        "date_session": date.today() + timedelta(days=10), "theme_ids": [await _ref(db, "THEME")],
        "lieu_id": await _ref(db, "LIEU"), "formateur_ids": [await _ref(db, "FORMATEUR")], "employe_ids": [e["id"]],
    })
    with pytest.raises(AppError) as futur:
        await svc.saisir_presences(uuid.UUID(s["id"]), {s["participants"][0]["id"]: "PRESENT"}, s["revision"])
    assert futur.value.code == "FORMATION_FUTURE"
    with pytest.raises(AppError):
        await svc.changer_statut(uuid.UUID(s["id"]), "annuler", s["revision"], None)
    s = await svc.changer_statut(uuid.UUID(s["id"]), "annuler", s["revision"], "Formateur indisponible")
    assert s["statut"] == "ANNULEE"
    s = await svc.changer_statut(uuid.UUID(s["id"]), "retablir", s["revision"], None)
    assert s["statut"] == "PLANIFIEE"


@pytest.mark.asyncio
async def test_feuille_signee_ged(db, monkeypatch):
    monkeypatch.setattr(formation_documents, "enqueue_ocr", lambda _id: None)
    svc = FormationService(db, await _ctx(db))
    e = await _employe(svc, f"ZZG{uuid.uuid4().hex[:6].upper()}")
    base = {"theme_ids": [await _ref(db, "THEME")], "lieu_id": await _ref(db, "LIEU"),
            "formateur_ids": [await _ref(db, "FORMATEUR")], "employe_ids": [e["id"]]}
    futur = await svc.creer_session({"date_session": date.today() + timedelta(days=5), **base})
    pdf = b"%PDF-1.4\n1 0 obj<<>>endobj\ntrailer<<>>\n%%EOF\n"
    with pytest.raises(AppError) as avant:
        await formation_documents.deposer(svc, uuid.UUID(futur["id"]), UploadFile(BytesIO(pdf), filename="f.pdf"))
    assert avant.value.code == "DEPOT_INTERDIT"

    s = await svc.creer_session({"date_session": date.today() - timedelta(days=1), **base})
    sid = uuid.UUID(s["id"])
    lecteur = FormationService(db, await _ctx(db, "formation.view"))
    with pytest.raises(HTTPException):
        await formation_documents.deposer(lecteur, sid, UploadFile(BytesIO(pdf), filename="f.pdf"))

    chemins = []
    try:
        r = await formation_documents.deposer(svc, sid, UploadFile(BytesIO(pdf), filename="feuille signée.pdf"))
        assert len(r["items"]) == 1 and r["actions"]["retirer"]
        doc = await db.get(GedDocument, uuid.UUID(r["items"][0]["id"]))
        chemins.append(doc.stored_path)
        assert (doc.module_code, doc.entity, doc.entity_id) == ("formation", "formation_session", s["id"])
        assert doc.doc_type == "FEUILLE_PRESENCE_SIGNEE" and doc.reference == s["reference"]

        liste = await svc.lister_sessions(feuille="AVEC", q=s["reference"])
        assert [x["id"] for x in liste["items"]] == [s["id"]] and liste["items"][0]["feuilles_signees"] == 1
        assert not (await svc.lister_sessions(feuille="SANS", q=s["reference"]))["items"]

        _, chemin = await formation_documents.fichier(svc, sid, doc.id)
        assert chemin.read_bytes() == pdf

        p = s["participants"][0]["id"]
        s = await svc.saisir_presences(sid, {p: "PRESENT"}, s["revision"])
        s = await svc.changer_statut(sid, "cloturer", s["revision"], None)
        gestionnaire = FormationService(db, await _ctx(db, "formation.view", "formation.attendance.manage"))
        with pytest.raises(AppError) as fige:
            await formation_documents.retirer(gestionnaire, sid, doc.id, "erreur")
        assert fige.value.code == "RETRAIT_INTERDIT"
        s = await svc.changer_statut(sid, "rouvrir", s["revision"], "Mauvais scan")
        with pytest.raises(AppError):
            await formation_documents.retirer(svc, sid, doc.id, " ")
        r = await formation_documents.retirer(svc, sid, doc.id, "Mauvais scan")
        assert r["items"] == []

        actions = set((await db.scalars(select(AuditLog.action).where(AuditLog.entity_id == s["id"]))).all())
        assert {"formation.document.upload", "formation.document.download", "formation.document.delete"} <= actions
    finally:
        for rel in chemins:
            GedService(db).absolute_path(rel).unlink(missing_ok=True)


@pytest.mark.asyncio
async def test_doublon_employe_et_permissions(db):
    svc = FormationService(db, await _ctx(db))
    nom = f"ZZD{uuid.uuid4().hex[:6].upper()}"
    await svc.creer_employe({"nom": nom, "prenom": "Aminata"})
    with pytest.raises(HTTPException) as doublon:
        await svc.creer_employe({"nom": "Aminata", "prenom": nom})
    assert doublon.value.status_code == 409 and doublon.value.detail["code"] == "DOUBLON_EMPLOYE"

    lecteur = FormationService(db, await _ctx(db, "formation.view", "formation.employees.view"))
    with pytest.raises(HTTPException) as refus:
        await lecteur.creer_employe({"nom": "Interdit"})
    assert refus.value.status_code == 403


@pytest.mark.asyncio
async def test_admin_supprime_tout(db):
    admin = FormationService(db, await _ctx(db))
    gestionnaire = FormationService(db, await _ctx(
        db, "formation.view", "formation.update", "formation.employees.manage", "formation.references.manage"))
    suffixe = uuid.uuid4().hex[:6].upper()
    e1, e2 = await _employe(admin, f"ZZS{suffixe}"), await _employe(admin, f"ZZT{suffixe}")
    theme = await admin.creer_referentiel("THEME", f"Thème jetable {suffixe}")
    fonction = await admin.creer_referentiel("FONCTION", f"Fonction jetable {suffixe}")
    await admin.modifier_employe(uuid.UUID(e2["id"]), {"fonction_id": fonction["id"]})
    base = {"theme_ids": [await _ref(db, "THEME"), theme["id"]], "lieu_id": await _ref(db, "LIEU"),
            "formateur_ids": [await _ref(db, "FORMATEUR")]}
    s = await admin.creer_session({"date_session": date.today() - timedelta(days=2), **base,
                                   "employe_ids": [e1["id"], e2["id"]]})
    sid = uuid.UUID(s["id"])
    p1, p2 = (p["id"] for p in s["participants"])
    s = await admin.saisir_presences(sid, {p1: "PRESENT", p2: "ABSENT"}, s["revision"])
    s = await admin.changer_statut(sid, "cloturer", s["revision"], None)
    assert s["actions"]["supprimer"] and s["actions"]["retirer_participants"]

    # Valeurs utilisées : refusées au gestionnaire, détachées puis supprimées par l'administrateur.
    with pytest.raises(HTTPException) as utilise:
        await gestionnaire.supprimer_referentiel(uuid.UUID(theme["id"]), forcer=True)
    assert utilise.value.detail["code"] == "REFERENTIEL_UTILISE"
    await admin.supprimer_referentiel(uuid.UUID(theme["id"]), forcer=True)
    await admin.supprimer_referentiel(uuid.UUID(fonction["id"]), forcer=True)
    s = await admin.detail_session(sid)
    assert theme["id"] not in {t["id"] for t in s["themes"]} and len(s["themes"]) == 1
    with pytest.raises(HTTPException) as lieu:
        await admin.supprimer_referentiel(uuid.UUID(base["lieu_id"]), forcer=True)
    assert lieu.value.detail["code"] == "REFERENTIEL_OBLIGATOIRE"

    # Participant retiré d'une formation clôturée, puis employé supprimé avec ses participations.
    s = await admin.retirer_participant(sid, uuid.UUID(p2), s["revision"], "Erreur de saisie")
    assert s["stats"]["participants"] == 1
    with pytest.raises(HTTPException):
        await gestionnaire.supprimer_employe(uuid.UUID(e1["id"]), "Doublon")
    with pytest.raises(AppError):
        await admin.supprimer_employe(uuid.UUID(e1["id"]), " ")
    assert await admin.supprimer_employe(uuid.UUID(e1["id"]), "Doublon") == e1["nom_complet"]
    with pytest.raises(AppError):
        await admin.fiche_employe(uuid.UUID(e1["id"]))
    s = await admin.detail_session(sid)
    assert s["stats"]["participants"] == 0

    # Formation clôturée supprimée définitivement.
    assert await admin.supprimer_session(sid, "Formation de test") == s["reference"]
    with pytest.raises(AppError):
        await admin.detail_session(sid)
    actions = set((await db.scalars(select(AuditLog.action).where(AuditLog.entity_id.in_(
        [s["id"], e1["id"], theme["id"]])))).all())
    assert {"formation.session.delete", "formation.employe.delete", "formation.referentiel.delete"} <= actions


@pytest.mark.asyncio
async def test_admin_supprime_entite_utilisee(db):
    admin = FormationService(db, await _ctx(db))
    perimetre = await _ref(db, "PERIMETRE")
    entite = await admin.creer_entite({"libelle": f"Entité jetable {uuid.uuid4().hex[:6]}", "perimetre_id": perimetre})
    e = await admin.creer_employe({"nom": f"ZZE{uuid.uuid4().hex[:6].upper()}", "entite_id": entite["id"]}, forcer=True)
    gestionnaire = FormationService(db, await _ctx(db, "formation.view", "formation.references.manage"))
    with pytest.raises(HTTPException):
        await gestionnaire.supprimer_entite(uuid.UUID(entite["id"]), forcer=True)
    await admin.supprimer_entite(uuid.UUID(entite["id"]), forcer=True)
    assert (await admin.fiche_employe(uuid.UUID(e["id"])))["entite_id"] is None


def _classeur_historique(suffixe: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Feuil1"
    ws.append(["Date mise à jour", datetime(2024, 11, 20)])
    ws.append([])
    ws.append(["Date", "Thème", "Lieu", "Formateur", "Nom du participant", "Fonction", "Entité", "Périmètre"])
    ws.append([datetime(2023, 5, 4), "LBC FT/Procédure ouverture de compte", "Nouakchott",
               "Abass Ngam / Fatimata Thiam", f"Mariem Mint {suffixe}", "Caissière", "Agence Ksar", "#REF!"])
    ws.append([datetime(2023, 5, 4), "LBC FT/Procédure ouverture de compte", "Nouakchott",
               "Abass Ngam / Fatimata Thiam", f"Sidi {suffixe} Ahmed", "Chef d'agence Ksar", "AGENCE KSAR", "Néant"])
    ws.append([datetime(2023, 5, 4), "LBC FT/Procédure ouverture de compte", "Nouakchott",
               "Abass Ngam / Fatimata Thiam", f"MINT {suffixe} Mariem", "Caissière", "Agence Ksar", ""])
    ws.append([None, "FATCA", "Rosso", "Yero Dieng", f"Sans Date {suffixe}", "Caissier", "XXX", ""])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.mark.asyncio
async def test_import_excel_analyse_et_confirmation(db):
    ctx = await _ctx(db)
    imp = FormationImportService(db, ctx)
    suffixe = uuid.uuid4().hex[:6].upper()
    res = await imp.analyser(_classeur_historique(suffixe), "Suivi des formations.xlsx")
    a = res["analyse"]
    assert a["ligne_entete"] == 3 and a["stats"]["lignes_ignorees"] == 1
    assert a["stats"]["sessions"] == 1
    # « Mariem Mint X » et « MINT X Mariem » = même personne (ordre des mots indifférent).
    assert a["stats"]["personnes"] == 2
    themes = {v["brut"]: v for v in a["valeurs"]["THEME"]}
    assert themes["LBC FT"]["action"] == "MAPPER" and themes["LBC FT"]["score"] == 1.0
    assert any(x["niveau"] == "bloquant" for x in a["anomalies"])
    fonctions = {v["brut"]: v for v in a["valeurs"]["FONCTION"]}
    assert fonctions["Chef d'agence Ksar"]["cible"] == "Chef d'agence"

    personnes = {p["cle"]: {"action": "NOUVEAU", "nom": p["brut"].split()[-1], "prenom": " ".join(p["brut"].split()[:-1])}
                 for p in a["personnes"]}
    with pytest.raises(AppError) as verif:
        await imp.confirmer(uuid.UUID(res["id"]), {"personnes": personnes})
    assert verif.value.code == "VERIFICATION_NOMS_REQUISE"

    fait = await imp.confirmer(uuid.UUID(res["id"]), {"personnes": personnes, "verification_noms": True})
    assert fait["statut"] == "IMPORTE"
    assert fait["resultat"]["sessions_creees"] == 1 and fait["resultat"]["participations_creees"] == 2

    svc = FormationService(db, ctx)
    liste = await svc.lister_sessions(q="lbc", statut="CLOTUREE", annee=2023, taille=200)
    session = next(s for s in liste["items"] if s["source"] == "IMPORT" and s["date_session"] == "2023-05-04")
    detail = await svc.detail_session(uuid.UUID(session["id"]))
    assert {p["presence"] for p in detail["participants"]} == {"PRESENT"}
    assert len(detail["formateurs"]) == 2 and len(detail["themes"]) == 2

    # Ré-import du même fichier : signalé, puis idempotent si forcé.
    res2 = await imp.analyser(_classeur_historique(suffixe), "Suivi des formations.xlsx")
    assert res2["analyse"]["deja_importe"]
    with pytest.raises(HTTPException):
        await imp.confirmer(uuid.UUID(res2["id"]), {"verification_noms": True})
    fait2 = await imp.confirmer(uuid.UUID(res2["id"]), {"verification_noms": True, "forcer": True})
    assert fait2["resultat"].get("participations_creees", 0) == 0

    with pytest.raises(HTTPException):
        await FormationImportService(db, await _ctx(db, "formation.import.execute")).supprimer(uuid.UUID(res2["id"]))
    await imp.supprimer(uuid.UUID(res2["id"]))
    with pytest.raises(AppError):
        await imp.detail(uuid.UUID(res2["id"]))
    assert (await svc.detail_session(uuid.UUID(session["id"])))["source"] == "IMPORT"

"""Consolidation ORION sans base : une ligne = un compte, un client par racine.

Le test sur le fichier réel est sauté tant que ``CLIENTELE_ETAT_COMPTE`` ne pointe pas vers
l'extrait « Etat Compte BEA » (données nominatives : jamais versionné).
"""

from __future__ import annotations

import os
from datetime import date, datetime

import pytest

from app.services.clientele.consolidation import LigneInvalide, cle_rib, consolider, lire_ligne


def rib_pour(compte: str, guichet: str = "00001") -> str:
    return f"00007{guichet}{compte}{cle_rib('00007', guichet, compte):02d}"


def brute(racine: object, compte: str, rib: str | None = None, **autres) -> dict:
    valeurs = {
        "RIB": rib or rib_pour(compte),
        "COMPTE": compte,
        "CODE_AGENCE_COMPTE": "00001",
        "AGENCE_COMPTE": "AGENCE CENTRALE PARTICULIERS",
        "CLIENT": racine,
        "NCG": "210000",
        "RUBRIQUE_COMPTABLE": "C.ORD.CLIENTELE RESIDENT",
        "RAISON_SOCIAL": f"CLIENT {racine}",
        "DATE_NAISSANCE": "31/12/1966",
        "NATIONALITE": "MAURITANIE",
        "LISTE_IDP_PRENOMS": "PRENOM",
        "ETAT_COMPTE": "Ouvert",
        "AGENT_ECONOMIQUE": "PARTICULIERS",
        "SITUATION_JURIDIQUE": "(Autre) personne physique",
        "CATEGORIE_JURIDIQUE": "PERSONNE PHYSIQUE",
        "R/N - Statut résident": "R",
        "TYPE_IDENTIFIANT": "NNI",
        "IDENTIFIANT_LISTE": "1234567890",
        "CONFORMITE_COMPTE": "COMPTES CONFORMES",
        "DATOUV": 46296,
        "DDC": None,
        "DDD": datetime(2026, 9, 1),
        "DEVISE": "MRU",
    }
    valeurs.update(autres)
    return valeurs


def consolider_brutes(*lignes: dict):
    return consolider(lire_ligne(b, i) for i, b in enumerate(lignes, 2))


def test_rib_exemple_metier():
    assert rib_pour("00000100001") == "00007000010000010000159"


def test_un_client_trois_comptes_puis_deuxieme_client():
    resultat = consolider_brutes(
        brute("000001", "00000100001"),
        brute("000001", "00000100002"),
        brute("000001", "00000100003"),
    )
    assert (resultat.nb_clients, resultat.nb_comptes, resultat.nb_rib) == (1, 3, 3)
    assert {c.racine_client for c in resultat.comptes.values()} == {"000001"}

    resultat = consolider_brutes(
        brute("000001", "00000100001"),
        brute("000001", "00000100002"),
        brute("000001", "00000100003"),
        brute("000002", "00000200001"),
    )
    assert (resultat.nb_clients, resultat.nb_comptes, resultat.nb_rib) == (2, 4, 4)
    assert not resultat.anomalies


def test_zeros_initiaux_conserves():
    ligne = lire_ligne(brute("000001", "00000100001"), 2)
    assert ligne.client.racine_client == "000001"
    assert ligne.compte.compte == "00000100001"
    assert ligne.compte.rib.startswith("00007")


@pytest.mark.parametrize("racine", [1, 1.0])
def test_racine_numerique_refusee(racine):
    with pytest.raises(LigneInvalide, match="zéros initiaux perdus"):
        lire_ligne(brute(racine, "00000100001"), 2)


@pytest.mark.parametrize("racine", ["00001", "0000001", "00000A", "", None])
def test_racine_longueur_ou_format_invalide(racine):
    with pytest.raises(LigneInvalide, match="CLIENT"):
        lire_ligne(brute(racine, "00000100001"), 2)


def test_rib_qui_ne_contient_pas_le_compte():
    with pytest.raises(LigneInvalide, match="ne contient pas le COMPTE"):
        lire_ligne(brute("000001", "00000100001", rib=rib_pour("00000100002")), 2)


def test_guichet_du_rib_different_de_l_agence():
    with pytest.raises(LigneInvalide, match="guichet"):
        lire_ligne(brute("000001", "00000100001", rib=rib_pour("00000100001", "08001")), 2)


def test_donnees_client_incoherentes_bloquent():
    resultat = consolider_brutes(
        brute("000001", "00000100001"),
        brute("000001", "00000100002", RAISON_SOCIAL="AUTRE NOM"),
    )
    assert resultat.nb_clients == 1
    assert resultat.bloquante
    assert resultat.anomalies[0].code == "CLIENT_INCOHERENT"
    assert "raison_sociale" in resultat.anomalies[0].message


def test_attributs_de_compte_peuvent_varier():
    resultat = consolider_brutes(
        brute("000001", "00000100001", DEVISE="MRU", ETAT_COMPTE="Ouvert", NCG="210000"),
        brute("000001", "00000100002", DEVISE="EUR", ETAT_COMPTE="Fermé ", NCG="240100"),
    )
    assert not resultat.bloquante
    assert resultat.comptes["00000100002"].etat_compte == "FERME"


def test_compte_duplique_bloque():
    resultat = consolider_brutes(brute("000001", "00000100001"), brute("000001", "00000100001"))
    assert resultat.nb_comptes == 1
    assert [a.code for a in resultat.anomalies] == ["COMPTE_DUPLIQUE"]


def test_cle_rib_invalide_signalee_sans_bloquer():
    rib = rib_pour("00000100001")
    faux = rib[:21] + f"{(int(rib[21:]) + 1) % 100:02d}"
    resultat = consolider_brutes(brute("000001", "00000100001", rib=faux))
    assert [(a.code, a.bloquante) for a in resultat.anomalies] == [("CLE_RIB_INVALIDE", False)]


def test_racine_en_position_1_du_compte_acceptee():
    resultat = consolider_brutes(brute("000026", "00000260007"))
    assert not resultat.anomalies


def test_racine_absente_du_compte_signalee():
    resultat = consolider_brutes(brute("000002", "00000100009"))
    assert [(a.code, a.bloquante) for a in resultat.anomalies] == [("RACINE_HORS_COMPTE", False)]


def test_conversions_dates_et_identifiants():
    ligne = lire_ligne(brute("000001", "00000100001", DATOUV=35365, DDC=46296), 2)
    assert ligne.compte.date_ouverture == date(1996, 10, 27)
    assert ligne.compte.ddc == date(2026, 10, 1)
    assert ligne.compte.ddd == date(2026, 9, 1)
    assert ligne.client.date_naissance == date(1966, 12, 31)
    assert ligne.client.nni == "1234567890" and ligne.client.nif is None

    joint = lire_ligne(brute("000001", "00000100001", DATE_NAISSANCE="01/01/1970, 02/02/1980",
                             IDENTIFIANT_LISTE="1111111111, 2222222222"), 2)
    assert joint.client.date_naissance is None
    assert joint.client.date_naissance_orion == "01/01/1970, 02/02/1980"
    assert joint.client.nni is None

    double = lire_ligne(brute("000001", "00000100001", IDENTIFIANT_LISTE="AB 1234567/1234567890"), 2)
    assert double.client.nni is None
    assert double.client.identifiant_orion == "AB 1234567/1234567890"


@pytest.mark.skipif(not os.environ.get("CLIENTELE_ETAT_COMPTE"), reason="CLIENTELE_ETAT_COMPTE non défini")
def test_fichier_reel_etat_compte():
    from openpyxl import load_workbook

    feuille = load_workbook(os.environ["CLIENTELE_ETAT_COMPTE"], read_only=True, data_only=True).worksheets[0]
    lignes = feuille.iter_rows(values_only=True)
    entete = next(lignes)
    resultat = consolider(
        lire_ligne(dict(zip(entete, valeurs)), numero) for numero, valeurs in enumerate(lignes, 2) if valeurs[0]
    )
    assert not resultat.bloquante, resultat.anomalies[:5]
    assert resultat.nb_comptes == resultat.nb_rib == resultat.nb_lignes
    assert resultat.nb_clients == len({c.racine_client for c in resultat.comptes.values()})
    assert resultat.nb_clients < resultat.nb_comptes

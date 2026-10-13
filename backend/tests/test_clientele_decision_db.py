"""Politique de décision de classification sur PostgreSQL (directive du 10/10/2026).

Exécution uniquement sur une copie dont le nom se termine par ``_test`` (ex. bea_digital_clientele_test),
migration ``20261010_clientele_08`` appliquée. Chaque test est annulé à la fin.
"""

from __future__ import annotations

from datetime import date, datetime
from io import BytesIO
from uuid import UUID

import pytest
from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select, text
from sqlalchemy.engine import make_url
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine
from sqlalchemy.pool import NullPool

from app.core.config import get_settings
from app.core.exceptions import AppError
from app.models import (
    AuditLog,
    ClienteleClassifHistorique,
    ClienteleClassifReexamen,
    ClienteleClassification,
    ClienteleClient,
    User,
)
from app.services.clientele.classification import ClienteleClassificationService, autorisation_classement_masse
from app.services.clientele.lots import ClienteleLotService
from app.services.clientele.service import Ctx

EXTRACTION = date(2026, 10, 7)
RACINES = ("990001", "990002", "990003", "990004", "990005", "990006")
BASE = {"agent_economique": "AUTRES SOCIETES", "statut_resident": "R", "nationalite": "MAURITANIE"}
PERMS_GESTION = {"clientele.view", "clientele.classif.view", "clientele.classif.execute", "clientele.import.execute"}
PERMS_ADMIN = PERMS_GESTION | {"clientele.classif.force", "clientele.classif.arbitrer", "clientele.admin"}


@pytest.fixture
async def db(tmp_path, monkeypatch):
    url = get_settings().database_url
    if not (make_url(url).database or "").endswith("_test"):
        pytest.skip("Tests réservés à une copie *_test de la base (jamais la base réelle)")
    engine = create_async_engine(url, poolclass=NullPool)
    try:
        async with engine.connect() as conn:
            pret = await conn.scalar(text("SELECT to_regclass('public.clientele_classif_lots') IS NOT NULL"))
    except Exception:
        pret = False
    if not pret:
        await engine.dispose()
        pytest.skip("Migration 20261010_clientele_08 non appliquée")
    monkeypatch.setattr(ClienteleLotService, "_dossier", lambda self: tmp_path)
    session = async_sessionmaker(engine, expire_on_commit=False)()
    clients = {
        "990001": BASE, "990002": BASE, "990003": {**BASE, "statut_resident": "N"}, "990004": BASE,
        "990005": {**BASE, "secteur_activite": "Casino"}, "990006": {},
    }
    for r, attrs in clients.items():
        session.add(ClienteleClient(racine_client=r, raison_sociale=f"TEST {r}", premiere_extraction=EXTRACTION,
                                    date_extraction=EXTRACTION, **attrs))
    await session.flush()
    try:
        yield session
    finally:
        await session.rollback()
        await session.close()
        await engine.dispose()


async def _ctx(db, perms: set[str]) -> Ctx:
    user = await db.scalar(select(User).where(User.is_superuser.is_(False)).limit(1))
    if user is None:
        pytest.skip("Aucun utilisateur non superadmin dans la copie")
    return Ctx(user=user, permissions=set(perms))


ENTETES = ["CLIENT", "RAISON SOCIALE", "AGENCE", "NCG", "DEVISE", "ETAT", "SOLDE", "Motif de risque",
           "Classe risque LBC/FT", "Entreprise ouverture récente", "Date ouverture", "Date de mise à jour",
           "Analyste KYC", "Colonne 14 libre"]


def _fichier() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(ENTETES)

    def ligne(racine, classe, ouverture=None, recente=None, motif=None):
        ws.append([racine, "X", "00001", "210000", "MRU", "Ouvert", 0, motif, classe, recente, ouverture,
                   datetime(2026, 10, 8), "A. SALL", f"libre {racine}"])

    ligne(990001, "Moyen", datetime(2025, 10, 6), motif="Société récente")   # 366 j → 10 000
    ligne(990002, "Faible", datetime(2025, 10, 5))                           # 367 j → 100
    ligne(990003, "Faible")                                                  # non-résident : ÉLEVÉ
    ligne(990004, "Faible")                                                  # décision manuelle existante
    ligne(990005, "Interdit")
    ligne(990006, "Moyen")                                                   # moteur NON_CLASSE
    ligne(990099, "Faible")                                                  # inconnue
    ligne(990001, "Moyen")                                                   # doublon
    ligne("ABC", "Faible")                                                   # invalide
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def _manuel_990004(db, ctx) -> None:
    await ClienteleClassificationService(db, ctx).decider_manuel("990004", "MOYEN", "Revue Conformité")


async def _lot_evalue(db, ctx, contenu: bytes | None = None,
                      type_lot: str = "LISTE") -> tuple[ClienteleLotService, dict]:
    svc = ClienteleLotService(db, ctx)
    lot = await svc.analyser(contenu or _fichier(), "hebdo.xlsx", type_lot=type_lot,
                             periode_reference=date(2026, 10, 8))
    lot = await svc.evaluer(UUID(lot["id"]))
    return svc, lot


def _par_racine(lot: dict) -> dict[str, dict]:
    return {x["racine"]: x for x in lot["lignes"] if x["statut"] == "OK"}


async def test_controles_mapping_et_colonnes(db):
    ctx = await _ctx(db, PERMS_GESTION)
    svc, lot = await _lot_evalue(db, ctx)
    s = lot["stats"]
    assert (s["racines_ok"], s["inconnues"], s["doublons"], s["invalides"]) == (6, 1, 1, 1)
    assert s["avec_date_ouverture"] == 2 and s["avec_analyste"] == 6 and s["avec_date_maj"] == 6
    assert lot["colonne_entreprise_recente"] == "Entreprise ouverture récente"
    assert lot["colonne_date_ouverture"] == "Date ouverture" and lot["colonne_analyste"] == "Analyste KYC"
    assert "Colonne 14 libre" in lot["colonnes_affichees"]
    lignes = _par_racine(lot)
    assert lignes["990001"]["colonnes"]["Colonne 14 libre"] == "libre 990001"
    assert await db.scalar(select(func.count()).select_from(ClienteleClient)
                           .where(ClienteleClient.racine_client == "990099")) == 0
    export = load_workbook(BytesIO(await svc.exporter(UUID(lot["id"]))), read_only=True)
    assert export.sheetnames == ["CLIENTS"]
    source = {str(c) for row in export["CLIENTS"].iter_rows(max_row=12, values_only=True) for c in row if c}
    assert "CODE AGENCE" in source and "RACINE CLIENT" in source and "Analyste KYC" in source
    assert "DATE_CREATION_ENTREPRISE" not in source
    assert "Colonne 14 libre" not in source
    assert "Classe source" not in source
    assert "990001" in source


async def test_seuil_366_et_coherence_individuel_lot(db):
    ctx = await _ctx(db, PERMS_GESTION)
    _, lot = await _lot_evalue(db, ctx)
    lignes = _par_racine(lot)
    assert lignes["990001"]["moteur"]["score"] == 10_200 and lignes["990001"]["moteur"]["niveau"] == "MOYEN"
    assert lignes["990002"]["moteur"]["score"] == 300 and lignes["990002"]["moteur"]["niveau"] == "FAIBLE"
    individuel = await ClienteleClassificationService(db, ctx).evaluer_racine("990001", persister=False)
    assert individuel["score_total"] == 10_200 and individuel["date_reference"] == EXTRACTION.isoformat()
    assert individuel["source_date_reference"] == "EXTRACTION_ORION"
    er = next(lg for lg in individuel["lignes"] if lg["critere"] == "ENTREPRISE_RECENTE")
    assert er["source_donnee"] == "FICHIER" and er["poids"] == 10_000


async def test_fichier_divergence_manuel_dates_analyste(db):
    ctx = await _ctx(db, PERMS_GESTION)
    await _manuel_990004(db, ctx)
    contenu = _fichier()
    svc_liste, liste = await _lot_evalue(db, ctx, contenu)
    with pytest.raises(AppError) as refus:
        await svc_liste.classer(UUID(liste["id"]), {"origine": "FICHIER", "motif": "Revue hebdo"})
    assert refus.value.code == "DECISION_PAR_LIGNE"
    svc, lot = await _lot_evalue(db, ctx, contenu, "SITUATION")
    lignes = _par_racine(lot)
    assert lignes["990003"]["divergence"] is True
    assert lignes["990006"]["decision_detail"]["comparaison"] == "NON_COMPARABLE"
    res = await svc.classer(UUID(lot["id"]), {"origine": "FICHIER", "motif": "Revue hebdo"})
    decisions = {x["racine"]: x["decision"] for x in res["lignes"] if x["statut"] == "OK"}
    assert decisions["990001"] == "ECRIT"
    assert decisions["990003"] == "DIVERGENCE_A_TRAITER"
    assert decisions["990004"] == "MANUEL_PRESERVE"
    c1 = await db.get(ClienteleClassification, "990001")
    assert (c1.niveau, c1.mode_decision, c1.classe_source_fichier, c1.niveau_propose) == ("MOYEN", "FICHIER",
                                                                                          "MOYEN", "MOYEN")
    assert c1.score_calcule == 10_200 and c1.analyste_kyc == "A. SALL" and c1.date_maj_metier == date(2026, 10, 8)
    assert c1.motif_source == "Société récente" and c1.ligne_source == 2 and c1.date_reference == EXTRACTION
    assert c1.version_moteur and c1.version_regles and c1.evaluation_id
    assert await db.get(ClienteleClassification, "990003") is None
    c4 = await db.get(ClienteleClassification, "990004")
    assert (c4.niveau, c4.source) == ("MOYEN", "MANUEL")
    reex = (await db.scalars(select(ClienteleClassifReexamen).where(
        ClienteleClassifReexamen.racine_client == "990004"))).all()
    assert len(reex) == 1 and reex[0].niveau_suggere == "FAIBLE" and reex[0].source_suggestion == "FICHIER"
    h1 = (await db.scalars(select(ClienteleClassifHistorique).where(
        ClienteleClassifHistorique.racine_client == "990001"))).all()
    assert len(h1) == 1 and h1[0].analyste_kyc == "A. SALL" and h1[0].mode_decision == "FICHIER"

    svc2, lot2 = await _lot_evalue(db, ctx, contenu, "SITUATION")
    assert lot2["doublon_de"] in (liste["id"], lot["id"])
    res2 = await svc2.classer(UUID(lot2["id"]), {"origine": "FICHIER", "motif": "Ré-import"})
    assert {x["racine"]: x["decision"] for x in res2["lignes"] if x["statut"] == "OK"}["990001"] == "INCHANGE"
    assert await db.scalar(select(func.count()).select_from(ClienteleClassifHistorique)
                           .where(ClienteleClassifHistorique.racine_client == "990001")) == 1
    assert await db.scalar(select(func.count()).select_from(ClienteleClassifReexamen)
                           .where(ClienteleClassifReexamen.racine_client == "990004")) == 1

    ligne3 = lignes["990003"]["ligne"]
    await svc2.decider_divergences(UUID(lot["id"]), {"decisions": [
        {"ligne": ligne3, "choix": "MOTEUR", "motif": "Non-résident confirmé"}]})
    c3 = await db.get(ClienteleClassification, "990003")
    assert (c3.niveau, c3.mode_decision, c3.classe_source_fichier) == ("ELEVE", "MANUEL", "FAIBLE")
    assert c3.motif_decision_manuelle == "Non-résident confirmé" and c3.niveau_propose == "ELEVE"


async def test_moteur_non_admissible_et_manuel_preserve(db, masse_autorisee):
    ctx = await _ctx(db, PERMS_GESTION)
    await _manuel_990004(db, ctx)
    svc, lot = await _lot_evalue(db, ctx)
    res = await svc.classer(UUID(lot["id"]), {"origine": "MOTEUR", "motif": "Moteur"})
    decisions = {x["racine"]: x["decision"] for x in res["lignes"] if x["statut"] == "OK"}
    assert decisions["990006"] == "NON_ADMISSIBLE"
    assert decisions["990004"] in ("MANUEL_PRESERVE",)
    assert decisions["990005"] == "ECRIT"
    assert await db.get(ClienteleClassification, "990006") is None
    c5 = await db.get(ClienteleClassification, "990005")
    assert (c5.niveau, c5.mode_decision) == ("INTERDIT", "MOTEUR")


async def test_forcage_permission_confirmation_audit(db):
    gest = await _ctx(db, PERMS_GESTION)
    await _manuel_990004(db, gest)
    svc, lot = await _lot_evalue(db, gest, type_lot="SITUATION")
    lid = UUID(lot["id"])
    with pytest.raises(HTTPException) as e:
        await svc.classer(lid, {"origine": "FICHIER", "motif": "Forçage", "forcer": True, "confirmation": True})
    assert e.value.status_code == 403
    admin = await _ctx(db, PERMS_ADMIN)
    svc_admin = ClienteleLotService(db, admin)
    with pytest.raises(AppError) as e2:
        await svc_admin.classer(lid, {"origine": "FICHIER", "motif": "Forçage", "forcer": True})
    assert e2.value.code == "CONFIRMATION_REQUISE"
    await svc_admin.classer(lid, {"origine": "FICHIER", "motif": "Forçage validé", "forcer": True,
                                  "confirmation": True})
    c4 = await db.get(ClienteleClassification, "990004")
    assert (c4.niveau, c4.mode_decision, c4.force) == ("FAIBLE", "FICHIER", True)
    assert await db.scalar(select(func.count()).select_from(AuditLog).where(
        AuditLog.action == "clientele.classif.forcage", AuditLog.entity_id == "990004")) == 1


async def test_blocage_interdit_non_contournable(db):
    gest = await _ctx(db, PERMS_GESTION)
    classif = ClienteleClassificationService(db, gest)
    with pytest.raises(AppError) as e:
        await classif.decider_manuel("990005", "FAIBLE", "Erreur de secteur")
    assert e.value.code == "BLOCAGE_INTERDIT"
    with pytest.raises(HTTPException):
        await classif.decider_manuel("990005", "FAIBLE", "Erreur de secteur", forcer=True, confirmation=True)
    await classif.decider_manuel("990005", "INTERDIT", "Confirmé")
    admin = ClienteleClassificationService(db, await _ctx(db, PERMS_ADMIN))
    await admin.decider_manuel("990005", "MOYEN", "Secteur ORION erroné, justificatif RCCM", forcer=True,
                               confirmation=True)
    c5 = await db.get(ClienteleClassification, "990005")
    assert (c5.niveau, c5.blocage_leve, c5.force, c5.niveau_propose) == ("MOYEN", True, True, "INTERDIT")
    assert await db.scalar(select(func.count()).select_from(ClienteleClassifHistorique)
                           .where(ClienteleClassifHistorique.racine_client == "990005")) == 2


async def test_fiche_classer_par_moteur(db):
    assert autorisation_classement_masse()["actif"] is False
    ctx = await _ctx(db, PERMS_GESTION)
    classif = ClienteleClassificationService(db, ctx)
    r6 = await classif.classer_par_moteur("990006", {"motif": "Évaluation fiche"})
    assert r6["resultat"] == "NON_ADMISSIBLE" and r6["classification"] is None
    r2 = await classif.classer_par_moteur("990002", {"motif": "Évaluation fiche"})
    assert r2["resultat"] == "ECRIT"
    c = r2["classification"]
    assert c["mode_decision"] == "MOTEUR" and c["niveau"] == c["niveau_propose"] == "FAIBLE"
    assert c["score_calcule"] == 200 and c["date_reference"] == EXTRACTION.isoformat()
    again = await classif.classer_par_moteur("990002", {"motif": "Évaluation fiche"})
    assert again["resultat"] == "INCHANGE"
    with pytest.raises(AppError):
        await classif.classer_par_moteur("990002", {"motif": ""})


async def test_classement_masse_moteur(db, masse_autorisee):
    ctx = await _ctx(db, PERMS_GESTION)
    classif = ClienteleClassificationService(db, ctx)
    await _manuel_990004(db, ctx)
    await classif._enregistrer("990003", "FAIBLE", source="EXCEL", version_id=None, motifs=[],
                               motif_risque=None, motif_classement="Fichier",
                               trace={"mode_decision": "FICHIER"})
    with pytest.raises(AppError):
        await classif.classer_masse_moteur({"motif": " "})
    with pytest.raises(AppError):
        await classif.classer_masse_moteur({"motif": "x", "profil": "ZZ"})
    q, filtres = classif._perimetre_masse({"profil": "pm", "code_agence": "00001"})
    assert filtres == {"profil": "PM", "code_agence": "00001"}
    assert await db.scalar(select(func.count()).select_from(q.subquery())) >= 0

    async def tout(motif: str) -> tuple[int, dict]:
        cumul: dict[str, int] = {}
        p = await classif.classer_masse_moteur({"motif": motif, "apres": "989999", "taille": 4})
        traites = p["traites"]
        while True:
            for k, v in p["bilan"].items():
                cumul[k] = cumul.get(k, 0) + v
            if not p["prochain"] or p["prochain"] >= "990006":
                return traites, cumul
            p = await classif.classer_masse_moteur({"motif": motif, "apres": p["prochain"], "taille": 4})
            traites += p["traites"]

    traites, b = await tout("Classification initiale")
    mes = {r: await db.get(ClienteleClassification, r) for r in RACINES}
    assert traites >= 6
    assert mes["990001"].mode_decision == "MOTEUR" and mes["990001"].evaluation_id
    assert (mes["990004"].niveau, mes["990004"].source) == ("MOYEN", "MANUEL")
    assert (mes["990003"].niveau, mes["990003"].source) == ("FAIBLE", "EXCEL")
    assert mes["990005"].niveau == "INTERDIT" and mes["990006"] is None
    assert b["MANUEL_PRESERVE"] >= 1 and b["FICHIER_PRESERVE"] >= 1 and b["NON_ADMISSIBLE"] >= 1
    reex3 = (await db.scalars(select(ClienteleClassifReexamen).where(
        ClienteleClassifReexamen.racine_client == "990003"))).all()
    assert len(reex3) == 1 and reex3[0].niveau_suggere == "ELEVE"

    _, b2 = await tout("Relance")
    assert b2.get("ECRIT", 0) == 0
    assert await db.scalar(select(func.count()).select_from(ClienteleClassifReexamen)
                           .where(ClienteleClassifReexamen.racine_client == "990003")) == 1


async def test_saisie_date_ouverture_apres_import(db):
    ctx = await _ctx(db, PERMS_GESTION)
    svc, lot = await _lot_evalue(db, ctx)
    lid = UUID(lot["id"])
    ligne3 = _par_racine(lot)["990003"]["ligne"]
    x = await svc.saisir_date_ouverture(lid, {"ligne": ligne3, "date_ouverture": "2026-03-01"})
    assert x["source_date_ouverture"] == "SAISIE" and x["moteur"] is None
    with pytest.raises(AppError):
        await svc.saisir_date_ouverture(lid, {"ligne": ligne3, "date_ouverture": "pas une date"})
    lot = await svc.evaluer(lid)
    ev = await ClienteleClassificationService(db, ctx).evaluer_racine("990003", persister=False)
    er = next(lg for lg in ev["lignes"] if lg["critere"] == "ENTREPRISE_RECENTE")
    assert er["source_donnee"] == "SAISIE" and er["poids"] == 10_000
    assert _par_racine(lot)["990003"]["moteur"]["score"] == ev["score_total"]

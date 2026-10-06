"""Facturation Fournisseurs sur PostgreSQL — page Fournisseurs, file de contrôle, historique, exports.

Même principe que ``test_mg_facturation_profils`` : transaction externe annulée à la fin.
"""

from __future__ import annotations

import uuid
from datetime import date
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.core.exceptions import AppError
from app.schemas.mg_facturation import FactureCreate, NouveauFournisseurIn, ProfilCreate
from app.services.mg_facturation_analytics import MgFacturationAnalytics
from tests.test_mg_facturation_profils import _contexte, _joindre_scan, _numero, db  # noqa: F401


async def test_nouveau_fournisseur_et_profil(db):  # noqa: F811
    user, _profils, svc = await _contexte(db)
    nom = f"OPERATEUR TEST {uuid.uuid4().hex[:6].upper()}"
    p = await svc.create_profil(
        ProfilCreate(
            code=f"T_{uuid.uuid4().hex[:6]}", libelle="Opérateur test", type_facture="internet",
            nouveau_fournisseur=NouveauFournisseurIn(raison_sociale=f"  {nom} "),
            champs={"numero_fournisseur": "obligatoire", "montant_ht": "masque", "montant_tva": "masque"},
            libelles={"montant_ttc": "Net à payer"},
        ),
        user,
    )
    assert p["fournisseur"] == nom and p["taux_tva"] is None and p["type_facture"] == "INTERNET"

    with pytest.raises(AppError) as err:
        await svc.create_profil(
            ProfilCreate(code=f"T_{uuid.uuid4().hex[:6]}", libelle="Doublon",
                         nouveau_fournisseur=NouveauFournisseurIn(raison_sociale=nom.lower())),
            user,
        )
    assert err.value.code == "FOURNISSEUR_EXISTANT"
    with pytest.raises(AppError) as err:
        await svc.create_profil(ProfilCreate(code=f"T_{uuid.uuid4().hex[:6]}", libelle="Sans fournisseur"), user)
    assert err.value.code == "PROFIL_FOURNISSEUR_REQUIS"

    f = await svc.create_facture(
        FactureCreate(profil_id=uuid.UUID(p["id"]), date_facture=date.today(), numero_fournisseur=_numero(),
                      montant_ttc=Decimal("750"), annee=date.today().year, mois=date.today().month),
        user,
    )
    stats = await svc.stats_fournisseurs(date.today().year)
    item = next(i for i in stats["items"] if i["profil_id"] == p["id"])
    assert item["nb"] == 1 and item["montant"] == 750.0 and item["reste"] == 750.0
    assert item["a_valider"] == 1 and item["mensuel"][date.today().month - 1] == 750.0
    assert f["fournisseur"] == nom


async def test_file_controles_et_validation_en_lot(db):  # noqa: F811
    user, profils, svc = await _contexte(db)
    ok = await svc.create_facture(
        FactureCreate(profil_id=profils["MATTEL_SMS"].id, date_facture=date.today(), numero_fournisseur=_numero(),
                      montant_ht=Decimal("1000"), montant_tva=Decimal("180"),
                      date_echeance=date.today()),
        user,
    )
    incoherente = await svc.create_facture(
        FactureCreate(profil_id=profils["MATTEL_SMS"].id, date_facture=date.today(), numero_fournisseur=_numero(),
                      montant_ht=Decimal("1000"), montant_tva=Decimal("50")),
        user,
    )
    sans_piece = await svc.create_facture(
        FactureCreate(profil_id=profils["MATTEL_SMS"].id, date_facture=date.today(), numero_fournisseur=_numero(),
                      montant_ht=Decimal("1000"), montant_tva=Decimal("180")),
        user,
    )
    _joindre_scan(db, uuid.UUID(ok["id"]), user)
    _joindre_scan(db, uuid.UUID(incoherente["id"]), user)
    file = await svc.file_controles({})
    par_id = {i["id"]: i for i in file["items"]}
    assert par_id[ok["id"]]["niveau_controle"] == "ok" and par_id[ok["id"]]["nb_documents"] == 1
    assert par_id[incoherente["id"]]["niveau_controle"] == "attention"
    assert "tva_incoherente" in {c["code"] for c in par_id[incoherente["id"]]["controles"]}
    assert par_id[sans_piece["id"]]["niveau_controle"] == "bloquant"
    assert file["compteurs"]["a_valider"] >= 3

    ids = [uuid.UUID(x["id"]) for x in (ok, incoherente, sans_piece)] + [uuid.uuid4()]
    res = await svc.transition_lot(ids, "valider", user, None)
    assert res["succes"] == 2 and res["echecs"] == 2
    assert {i["id"] for i in (await svc.file_controles({}))["items"]} & {ok["id"], incoherente["id"]} == set()

    journal = await svc.journal({"q": ok["reference"]})
    actions = {i["action"] for i in journal["items"]}
    assert "VALIDER" in actions
    assert all(i["reference"] == ok["reference"] for i in journal["items"])


async def test_exports_csv_et_classeur(db):  # noqa: F811
    user, profils, svc = await _contexte(db)
    await svc.create_facture(
        FactureCreate(profil_id=profils["SNDE"].id, date_facture=date.today(), numero_fournisseur=_numero(),
                      montant_ttc=Decimal("1234.5"), annee=date.today().year, mois=date.today().month),
        user,
    )
    analytics = MgFacturationAnalytics(svc)
    filtres = {"annee": date.today().year}
    csv = await analytics.export(user, "fournisseurs_mois", "csv", filtres)
    texte = csv.body.decode("utf-8")
    assert texte.startswith("\ufeffFournisseur;Janv") or texte.startswith("\ufeffFournisseur;")
    assert "SNDE" in texte

    xlsx = await analytics.export(user, "classeur", "xlsx", filtres)
    wb = load_workbook(BytesIO(xlsx.body))
    assert {"Synthèse", "Fournisseurs", "Fournisseurs mois", "Registre"} <= set(wb.sheetnames)

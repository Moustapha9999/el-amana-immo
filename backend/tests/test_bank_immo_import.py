"""Tests du parser / helpers d'import tableau banque."""

from datetime import date
from decimal import Decimal
from io import BytesIO

import xlrd
from openpyxl import Workbook

from app.services.bank_immo_import import (
    compute_dotation_exercice,
    parse_amount,
    parse_bank_workbook,
    parse_date,
    parse_taux,
    resolve_categorie_code,
)
from app.services.recap_amortissement import _is_import_banque, _mouvements_immo
from app.models.immobilisation import Immobilisation
from app.models.enums import StatutImmobilisation


def test_resolve_categorie_code():
    assert resolve_categorie_code("aai") == "TY-142010"
    assert resolve_categorie_code("matinfo ") == "TY-142041"
    assert resolve_categorie_code("constructé") == "TY-142020"
    assert resolve_categorie_code("FRIAS emmission EMPRT") == "TY-147030"
    assert resolve_categorie_code("recapAmort2022") is None


def test_parse_amount_and_taux():
    assert parse_amount("5 606 827,70") == Decimal("5606827.70")
    assert parse_amount(-168400) == Decimal("-168400.00")
    assert parse_amount("-") is None
    assert parse_taux("10%") == Decimal("10.00")
    assert parse_taux(0.1) == Decimal("10.00")
    assert parse_taux(10) == Decimal("10.00")


def test_parse_date_typo_and_serial():
    assert parse_date("01/01/20120") == date(2020, 1, 1)
    assert parse_date("06/01/2026") == date(2026, 1, 6)
    assert parse_date("02/05/212") == date(2012, 5, 2)  # typo banque 212 → 2012
    serial = xlrd.xldate.xldate_from_date_tuple((2007, 1, 2), 0)
    assert parse_date(serial, datemode=0) == date(2007, 1, 2)


def _build_aai_fixture_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "aai"
    ws.append(["BEA"])
    ws.append(["TABLEAU D'AMORTISSEMENT AAI"])
    ws.append(["COMPTE N°: 142010"])
    ws.append(
        [
            "Date",
            "Qté",
            "Désignation",
            "d'Acquisition MRU",
            "Taux",
            "Fin Exr.Précé",
            "Exer. En. C",
            "Fin Exercice",
            "Comptable",
            "AGENCE",
        ]
    )
    ws.append(["01/01/03", None, "REPORT 2003", 5606827.70, "10%", 5606827.70, None, 5606827.70, 0, None])
    ws.append(
        ["02/01/07", 1, "Construction chambre forte NDB", 106000.00, 0.1, 106000.00, None, 106000.00, 0, None]
    )
    ws.append(
        [
            "06/01/2026",
            None,
            "RGLT FACT ETS KERIM",
            203500.00,
            "10%",
            None,
            9892.36,
            9892.36,
            193607.64,
            "SIEGE CENTRAL",
        ]
    )
    # Reports de solde annuels — doivent être ignorés
    ws.append(
        ["01/01/10", None, "Report 01/01/2010", 7200000.00, "10%", 7200000.00, None, 7200000.00, 0, None]
    )
    ws.append(
        ["01/01/22", None, "Report de solde 31/12/2021", 9000000.00, None, 9000000.00, None, 9000000.00, 0, None]
    )
    # Variante orthographe banque « Rapport exercice »
    ws.append(
        ["01/01/25", None, "Rapport exercice 2024", 5916327.70, None, 5712827.70, None, 5712827.70, 0, None]
    )
    ws.append(
        ["30/06/26", None, "Solde au 30/06/2026", 5916327.70, None, 5712827.70, 9892.36, 5722720.06, 193607.64, None]
    )
    other = wb.create_sheet("recapAmort2022")
    other.append(["RECAP"])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_parse_bank_workbook_aai_fixture():
    content = _build_aai_fixture_xlsx()
    rows = parse_bank_workbook(content, "fixture.xlsx")
    assert len(rows) == 3
    assert rows[0].is_report is True
    assert rows[0].valeur_brute == Decimal("5606827.70")
    assert rows[0].categorie_code == "TY-142010"

    assert rows[1].designation.startswith("Construction")
    assert rows[1].taux == Decimal("10.00")
    assert rows[1].amt_n1 == Decimal("106000.00")
    assert rows[1].dotation == Decimal("0.00")

    assert rows[2].date_acquisition == date(2026, 1, 6)
    assert rows[2].dotation == Decimal("9892.36")
    assert rows[2].vnc == Decimal("193607.64")
    assert rows[2].agence_label == "SIEGE CENTRAL"

    vb = sum((r.valeur_brute for r in rows), Decimal("0"))
    assert vb == Decimal("5916327.70")


def _build_logiciel_fixture_xlsx() -> bytes:
    """Feuille type Logiciel 147530 : pas de colonnes Dotation / Fin exercice / VNC."""
    wb = Workbook()
    ws = wb.active
    ws.title = "logiciel"
    ws.append(["BEA"])
    ws.append(["TABLEAU D'AMORTISSEMENT LOGICIEL"])
    ws.append(["COMPTE N°: 1475300008  Compte Amort : N° 148700"])
    ws.append(["Date", "Qté", "Désignation", "d'acquisition", "Taux", "Fin Exr.Précé"])
    ws.append(["0101/05", None, "REPORT 2004", 1588323.26, "10%", 1588323.26])
    ws.append(["13/04/15", 1, "Licence CBS 30 000 Euros", 1035240.00, "10%", 1009359.00])
    # Lignes de total / report annuels — à ignorer
    ws.append(["31/12/16", None, "Total au 31/12/2016", 2623563.26, None, 2597682.26])
    ws.append(["01/01/17", None, "Report 01/01/2017", 2623563.26, None, 2597682.26])
    ws.append(["06/10/25", None, "RGLT FACT NSERVICES", 637200.00, "10%", None])
    ws.append(["15/01/26", None, "RGLT FACT EUT N° F128726", 363803.18, "10%", None])
    ws.append(["30/06/26", None, "Solde au 30/06/2026", 3624566.44, None, 2597682.26])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_parse_bank_workbook_logiciel_sans_colonnes_dotation():
    rows = parse_bank_workbook(_build_logiciel_fixture_xlsx(), "fixture.xlsx")
    assert len(rows) == 4
    assert all(r.calc_dotation is True for r in rows)
    assert all(r.categorie_code == "TY-147530" for r in rows)
    # Au parsing, pas de dotation banque : cumul fin = cumul N-1
    assert all(r.dotation == Decimal("0.00") for r in rows)
    assert all(r.amt_fin == r.amt_n1 for r in rows)

    vb = sum((r.valeur_brute for r in rows), Decimal("0"))
    n1 = sum((r.amt_n1 for r in rows), Decimal("0"))
    assert vb == Decimal("3624566.44")
    assert n1 == Decimal("2597682.26")


def _build_frais_emprunt_fixture_xlsx() -> bytes:
    """Feuille FRIAS emmission EMPRT : Solde de report avec libellé en colonne Qté."""
    wb = Workbook()
    ws = wb.active
    ws.title = "FRIAS emmission EMPRT"
    ws.append(["BEA - DEPARTEMENT FINANCE ET COMPTABILITE"])
    ws.append(["TABLEAU D'AMORTISSEMENT FRAIS IMMOBILISES"])
    ws.append(["COMPTE N°: 147030"])
    # Pas d'en-tête Qté / Désignation (mise en page banque réelle)
    ws.append(
        [
            "Date",
            None,
            None,
            "Valeur acquisition",
            "taux",
            "Amt Cumulés debut",
            "Dotation Exercice",
            "Amt fin exercice",
            "Valeur Net",
            "AGENCE",
        ]
    )
    ws.append(["15/09/25", None, "FRAIS EMRT", 15917280.00, "33%", None, None, None, 15917280.00, None])
    ws.append(["30/12/25", None, "FRAIS EMRT IFC", 10138800.00, None, None, None, None, 10138800.00, None])
    ws.append(["29/09/25", None, "Solde au 29/09/2025", 26056080.00, None, None, None, None, 26056080.00, None])
    # Report de solde mal aligné : texte Solde en Qté, désignation vide, "Date" en date
    ws.append(["Date ", "Solde au 01/01/2025", None, 26056080.00, None, None, None, None, 26056080.00, None])
    ws.append(
        [
            "05/01/26",
            None,
            "RGLT FRAIS MISSION IFC",
            17426250.00,
            "33%",
            None,
            2839549.35,
            2839549.35,
            14586700.65,
            "SIEGE CENTRAL",
        ]
    )
    ws.append(
        [
            "30/06/26",
            None,
            "Solde au 29/06/2026",
            43482330.00,
            None,
            None,
            2839549.35,
            2839549.35,
            40642780.65,
            None,
        ]
    )
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_parse_bank_workbook_frais_emprunt_ignore_solde_en_qte():
    """Le Solde reporté (libellé hors Désignation) ne doit pas doublonner la VB."""
    rows = parse_bank_workbook(_build_frais_emprunt_fixture_xlsx(), "fixture.xlsx")
    assert all(r.categorie_code == "TY-147030" for r in rows)
    assert len(rows) == 3
    assert [r.designation for r in rows] == [
        "FRAIS EMRT",
        "FRAIS EMRT IFC",
        "RGLT FRAIS MISSION IFC",
    ]
    assert not any("sans libellé" in r.designation.lower() for r in rows)
    vb = sum((r.valeur_brute for r in rows), Decimal("0"))
    assert vb == Decimal("43482330.00")
    assert rows[2].dotation == Decimal("2839549.35")
    assert rows[2].vnc == Decimal("14586700.65")


def _build_matinfo_opening_st_fixture_xlsx() -> bytes:
    """Matériel informatique : S/T d'ouverture 2005 puis acquisitions 2006."""
    wb = Workbook()
    ws = wb.active
    ws.title = "matinfo"
    ws.append(["TABLEAU D'AMORTISSEMENT MATERIEL INFORMATIQUE"])
    ws.append(
        [
            "Date",
            "Qté",
            "Désignations",
            None,
            "Valeur d'Acquisit°",
            "Taux",
            "Fin. Ex. Précédent",
            "Fin. Ex; en cours",
            "Amrt Fin . Ex",
            "Net",
        ]
    )
    # Stock d'ouverture (à conserver comme REPORT)
    ws.append(["31/12/05", "Q", "S/T", None, 2627450.00, None, 2627450.00, 0, 2627450.00, 0])
    ws.append(["23/01/06", 1, "Portable Toshiba", None, 73000.00, "20%", 73000.00, None, 73000.00, 0])
    ws.append(["07/02/06", 1, "Imprimante HP", None, 15000.00, "20%", 15000.00, None, 15000.00, 0])
    ws.append(["30/12/06", None, "S/T", None, 2715450.00, None, 2715450.00, 0, 2715450.00, 0])
    # Report d'ouverture année suivante — à ignorer
    ws.append(["30/06/06", "Q", "S/T", None, 2715450.00, None, 2715450.00, 0, 2715450.00, 0])
    ws.append(["02/01/07", 1, "Guichet Automatique", None, 100000.00, "20%", 100000.00, None, 100000.00, 0])
    # Total fin d'année sans libellé — à ignorer
    ws.append(["31/12/08", None, None, None, 2815450.00, None, 2815450.00, 0, 2815450.00, 0])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_parse_matinfo_keeps_opening_st_skips_blank_total():
    rows = parse_bank_workbook(_build_matinfo_opening_st_fixture_xlsx(), "fixture.xlsx")
    assert all(r.categorie_code == "TY-142041" for r in rows)
    assert rows[0].is_report is True
    assert rows[0].designation == "REPORT 2005"
    assert rows[0].valeur_brute == Decimal("2627450.00")
    assert [r.designation for r in rows] == [
        "REPORT 2005",
        "Portable Toshiba",
        "Imprimante HP",
        "Guichet Automatique",
    ]
    vb = sum((r.valeur_brute for r in rows), Decimal("0"))
    assert vb == Decimal("2815450.00")
    assert not any("sans libellé" in r.designation.lower() for r in rows)


def _header_matex() -> list:
    return [
        "Date",
        "Qté",
        "Désignation",
        "d'acquisition",
        "Taux",
        "Fin Exr.Précé",
        "Exer. En. C",
        "Fin Exercice",
        "Comptable",
    ]


def _build_matexhisto_deux_blocs_xlsx() -> bytes:
    """MatexHisto : 2 tableaux (comptes 1420970002 & 1420970015), chacun avec son REPORT."""
    wb = Workbook()
    ws = wb.active
    ws.title = "matexhisto"
    ws.append(["TABLEAU D'AMORTISSEMENT MAT EXPLOI HISTO"])
    ws.append(["COMPTE N°: 1420970002 & 1420970015"])
    # Bloc 1 — compte récent
    ws.append(_header_matex())
    ws.append(
        ["01/01/16", None, "REPORT 2015", 1697736.00, "10%", 1697736.31, None, 1697736.31, -0.31]
    )
    ws.append(
        ["18/08/22", 1, "Coffre fort", 42500.00, "10%", 10093.75, 4250.00, 14343.75, 28156.25]
    )
    ws.append(
        [
            "31/12/24",
            None,
            "Solde au 31 decembre 2023",
            1740236.00,
            None,
            1707830.06,
            4250.00,
            1712080.06,
            28155.94,
        ]
    )
    # Bloc 2 — historique (REPORT 2003 daté 01/01/15 comme sur le fichier banque)
    ws.append([])
    ws.append(_header_matex())
    ws.append(
        ["01/01/15", None, "REPORT 2003", 1337171.90, None, 1337171.90, None, 1337171.90, None]
    )
    ws.append(
        ["24/11/05", 1, "coffre fort", 46800.00, "10%", 46800.00, None, 46800.00, 0]
    )
    ws.append(
        [
            "31/12/23",
            None,
            "Solde au 31/12/2023",
            1383971.90,
            None,
            1383971.90,
            0,
            1383971.90,
            0,
        ]
    )
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_matexhisto_deux_blocs_garde_les_deux_reports():
    """Chaque bloc d'en-tête conserve son REPORT YYYY (pas un seul pour toute la feuille)."""
    rows = parse_bank_workbook(_build_matexhisto_deux_blocs_xlsx(), "matex.xlsx")
    assert all(r.categorie_code == "TY-142097" for r in rows)
    reports = [r for r in rows if r.is_report]
    assert len(reports) == 2
    assert {r.designation for r in reports} == {"REPORT 2015", "REPORT 2003"}
    r2003 = next(r for r in reports if r.designation == "REPORT 2003")
    assert r2003.valeur_brute == Decimal("1337171.90")
    # Date banque 01/01/15 → ancrée sur 2003 (année du libellé)
    assert r2003.date_acquisition == date(2003, 1, 1)
    vb = sum((r.valeur_brute for r in rows), Decimal("0"))
    assert vb == Decimal("3124207.90")


def test_compute_dotation_exercice_2026():
    dix = Decimal("10")
    # Bien totalement amorti → dotation nulle
    assert compute_dotation_exercice(
        Decimal("1588323.26"), Decimal("1588323.26"), dix, date(2004, 1, 1)
    ) == Decimal("0.00")
    # Acquisition < 2026 : 12 mois pleins, plafonnés au restant à amortir
    assert compute_dotation_exercice(
        Decimal("637200.00"), Decimal("0.00"), dix, date(2025, 10, 6)
    ) == Decimal("63720.00")
    assert compute_dotation_exercice(
        Decimal("1035240.00"), Decimal("1009359.00"), dix, date(2015, 4, 13)
    ) == Decimal("25881.00")
    # Acquisition 2026 : Excel DAYS360(15/01, 30/06) = 165 jours
    assert compute_dotation_exercice(
        Decimal("363803.18"), Decimal("0.00"), dix, date(2026, 1, 15)
    ) == Decimal("16674.31")
    # Acquisition postérieure à l'arrêté → aucune dotation
    assert compute_dotation_exercice(
        Decimal("100000.00"), Decimal("0.00"), dix, date(2026, 8, 10)
    ) == Decimal("0.00")
    # VB négative / reclassement / nivellement : pas de nouvelle dotation
    assert compute_dotation_exercice(
        Decimal("-2168127.00"),
        Decimal("-216812.70"),
        dix,
        date(2005, 11, 24),
        designation="Reclassement Financement BID",
    ) == Decimal("0.00")
    assert compute_dotation_exercice(
        Decimal("7953216.59"),
        Decimal("0.00"),
        dix,
        date(2025, 1, 14),
        designation="NIVELLEMENT SOLDE COMPTE 147530/47 VERS COMPTE 147530/50",
    ) == Decimal("0.00")


def test_mouvements_prefer_bank_seed():
    immo = Immobilisation(
        code_inventaire="142010-TEST",
        designation="Test banque",
        quantite=1,
        date_acquisition=date(2020, 1, 1),
        valeur_brute=Decimal("100000.00"),
        taux=Decimal("10"),
        statut=StatutImmobilisation.EN_SERVICE,
        compte_immobilisation="142010",
        metadata_json={"source": "import_banque"},
    )
    assert _is_import_banque(immo) is True
    mvts = _mouvements_immo(
        immo,
        2026,
        montants_dotation_db=[Decimal("5000.00")],
        cumul_n1_db=Decimal("40000.00"),
        cumul_fin_n_db=Decimal("45000.00"),
        vnc_fin_n_db=Decimal("55000.00"),
    )
    assert mvts is not None
    assert mvts["amorts_cumules_n1"] == Decimal("40000.00")
    assert mvts["dotations_annee"] == Decimal("5000.00")
    assert mvts["amorts_cumules_n"] == Decimal("45000.00")
    assert mvts["vnc"] == Decimal("55000.00")


def test_mouvements_stock_excel_2025_keeps_amt_n1_dotation():
    """Coffre 17/09/2015 : Excel amt_n1=8325 + dot=675 → fin=9000 (écart type 675)."""
    immo = Immobilisation(
        code_inventaire="MatExp-2015-002",
        designation="Coffre fort",
        quantite=1,
        date_acquisition=date(2015, 9, 17),
        valeur_brute=Decimal("9000.00"),
        taux=Decimal("10"),
        statut=StatutImmobilisation.EN_SERVICE,
        compte_immobilisation="142097",
        metadata_json={
            "source": "import_banque",
            "bank": {
                "amt_n1": "8325.00",
                "dotation": "675.00",
                "amt_fin": "9000.00",
                "vnc": "0.00",
                "seed_mode": "stock_ouverture",
            },
        },
    )
    # Année du Solde Excel : coller les colonnes banque
    mvts_2025 = _mouvements_immo(immo, 2025, cumul_n1_db=Decimal("9000.00"))
    assert mvts_2025 is not None
    assert mvts_2025["amorts_cumules_n1"] == Decimal("8325.00")
    assert mvts_2025["dotations_annee"] == Decimal("675.00")
    assert mvts_2025["amorts_cumules_n"] == Decimal("9000.00")
    assert mvts_2025["vnc"] == Decimal("0.00")

    # 2026 : ouverture = fin Excel ; pas de rejouer la dotation 2025
    mvts_2026 = _mouvements_immo(
        immo,
        2026,
        montants_dotation_db=[],
        cumul_n1_db=Decimal("9000.00"),
    )
    assert mvts_2026 is not None
    assert mvts_2026["amorts_cumules_n1"] == Decimal("9000.00")
    assert mvts_2026["dotations_annee"] == Decimal("0.00")
    assert mvts_2026["amorts_cumules_n"] == Decimal("9000.00")

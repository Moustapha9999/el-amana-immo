"""Déclaration SIC : colonnes, référentiel, découpage des noms, anomalies, fichier Excel. Tests purs."""

from __future__ import annotations

import io
from datetime import date, datetime

import pytest
from openpyxl import load_workbook

from app.data.clientele_sic import (
    COLONNES,
    PAR_CODE_SIC,
    decouper_nom,
    lire_referentiel,
    valider_valeur_sic,
)
from app.services.clientele.sic_calcul import (
    A_CORRIGER,
    AVERTISSEMENT,
    BLOQUANTE,
    CandidatSic,
    ContexteSic,
    construire,
    ecrire_xlsx,
    formater_date,
    rapport_anomalies_xlsx,
    verifier_xlsx,
)

AGENCE = "11111111-1111-1111-1111-111111111111"
CODES = {"101": "PM", "205": "PM", "1101": "PP", "1102": "PP"}


def _ctx(**kw) -> ContexteSic:
    regles = {r["code"]: r["valeur"] for r in PAR_CODE_SIC.values()}
    regles.update(kw.pop("regles", {}))
    return ContexteSic(regles=regles, regles_bcm=kw.pop("regles_bcm", {}), codes_ref=dict(CODES),
                       agences=kw.pop("agences", {AGENCE: "Tevragh Zeina"}),
                       secteurs=kw.pop("secteurs", {"COMMERCE": ("205", "CONFIRMEE"), "BTP": ("101", "A_VALIDER")}))


def _pp(**kw) -> CandidatSic:
    base = dict(racine="000123", profil="PP", categorie="PERSONNE PHYSIQUE", statut_resident="R",
                nni="0012345678", nom_complet="SIDI MOHAMED AHMED", prenoms="[IDP] 0043639 AHMED",
                date_entree=date(2015, 3, 7), agence_id=AGENCE, agence_code="00001", nb_comptes=1, nb_agences=1,
                eer_sexe="M", eer_telephone="22 33 44 55", kyc_csp="1101")
    base.update(kw)
    return CandidatSic(**base)


def _pm(**kw) -> CandidatSic:
    base = dict(racine="004567", profil="PM", categorie="SARL", nif="00098765", nom_complet="SOCIETE ALPHA",
                secteur_orion="Commerce", date_entree=date(2019, 12, 31), agence_id=AGENCE, agence_code="00001",
                nb_comptes=2, nb_agences=1)
    base.update(kw)
    return CandidatSic(**base)


def _gravites(ligne: dict, champ: str | None = None) -> set[str]:
    return {a["gravite"] for a in ligne["anomalies"] if champ is None or a["champ"] == champ}


# --------------------------------------------------------------------------- structure
def test_colonnes_modele_exactes():
    assert COLONNES["PP"] == ("CodeClient", "StatutResidentiel", "NNI", "Prenom", "Nom", "Genre",
                              "CodeCategorieSocioprofessionnelle", "NumTel", "DateEntreeEnRelation",
                              "MoughataaAgence")
    assert COLONNES["PM"] == ("CodeClient", "NIF", "RaisonSociale", "CodeSecteurActivite", "DateEntreeEnRelation",
                              "MoughataaAgence")


def test_referentiel_lecture_et_anomalies():
    lignes = [("Code", "Secteur d'activité", "Catégorie"),
              (101.0, "Agriculture", "Secteur primaire"),
              (1101, "Salarié", "Personnne physique"),
              (1109, "Autre", "Autre PP"),
              ("101", "Doublon", "Secteur primaire"),
              ("ABC", "Invalide", "X"),
              (205, "Commerce", None)]
    codes, anomalies = lire_referentiel(lignes)
    usages = {c["code"]: c["usage"] for c in codes}
    assert usages["101"] == "PM" and usages["1101"] == "PP" and usages["1109"] == "PP"
    motifs = " | ".join(a["motif"] for a in anomalies)
    assert "dupliqué" in motifs and "non numérique" in motifs and "Catégorie absente" in motifs


def test_referentiel_entetes_absents():
    codes, anomalies = lire_referentiel([("A", "B"), (1, 2)])
    assert codes == [] and anomalies


@pytest.mark.parametrize("complet,prenoms,attendu", [
    ("SIDI MOHAMED AHMED", "[IDP] 0043639 AHMED", ("SIDI MOHAMED", "AHMED", None)),
    ("FATIMETOU MINT ELY", "[IDM] 12 MINT ELY", ("FATIMETOU", "MINT ELY", None)),
])
def test_decoupage_nom(complet, prenoms, attendu):
    assert decouper_nom(complet, prenoms) == attendu


@pytest.mark.parametrize("complet,prenoms", [
    ("SIDI MOHAMED AHMED VAL", "[IDP] 1 AHMED"),   # nom tronqué / ordre différent
    ("AHMED", "[IDP] 1 AHMED"),                     # prénom = nom complet
    ("SIDI AHMED", "[IDP] 0043639"),                # prénom absent
])
def test_decoupage_nom_jamais_devine(complet, prenoms):
    nom, _prenom, motif = decouper_nom(complet, prenoms)
    assert nom is None and motif


def test_decoupage_prenom_en_tete_sur_option_seulement():
    assert decouper_nom("AHMED SIDI MATALLA", "[IDP] 1 AHMED")[0] is None
    assert decouper_nom("AHMED SIDI MATALLA", "[IDP] 1 AHMED", en_tete=True) == ("SIDI MATALLA", "AHMED", None)
    [lg] = construire([_pp(nom_complet="AHMED SIDI MATALLA")], "PP",
                      _ctx(regles={"sic.nom_prenom": "PRENOM_EN_FIN_OU_EN_TETE"}))
    assert lg["statut"] == "EXPORTEE" and lg["valeurs"][3:5] == ["AHMED", "SIDI MATALLA"]


def test_regles_sic_validation():
    assert valider_valeur_sic("sic.format_date", "DATE_EXCEL") == "DATE_EXCEL"
    with pytest.raises(ValueError):
        valider_valeur_sic("sic.format_date", "AAAA")
    with pytest.raises(ValueError):
        valider_valeur_sic("sic.genre", {"M": "X"})
    with pytest.raises(ValueError):
        valider_valeur_sic("sic.telephone_format", "[")


# --------------------------------------------------------------------------- lignes PP
def test_pp_complet_exporte_identifiants_texte():
    [lg] = construire([_pp()], "PP", _ctx())
    assert lg["statut"] == "EXPORTEE" and lg["anomalies"] == []
    assert lg["valeurs"] == ["000123", "MR", "0012345678", "AHMED", "SIDI MOHAMED", "H", "1101", "22334455",
                             "07/03/2015", "Tevragh Zeina"]


def test_pp_champs_absents_bloquants_jamais_inventes():
    [lg] = construire([_pp(eer_sexe=None, eer_telephone=None, kyc_csp=None, statut_resident=None, nni=None)],
                      "PP", _ctx())
    assert lg["statut"] == "EXPORTEE" and lg["gravite_max"] == BLOQUANTE
    for champ in ("Genre", "NumTel", "CodeCategorieSocioprofessionnelle", "StatutResidentiel", "NNI"):
        assert BLOQUANTE in _gravites(lg, champ), champ
    assert lg["valeurs"][5] is None and lg["valeurs"][6] is None


def test_pp_csp_inconnue_ou_reservee_pm():
    [inconnu] = construire([_pp(kyc_csp="1999")], "PP", _ctx())
    [pm] = construire([_pp(kyc_csp="101")], "PP", _ctx())
    assert BLOQUANTE in _gravites(inconnu, "CodeCategorieSocioprofessionnelle")
    assert BLOQUANTE in _gravites(pm, "CodeCategorieSocioprofessionnelle")


def test_pp_kyc_complete_eer_absent_et_divergence():
    [kyc] = construire([_pp(eer_sexe=None, eer_telephone=None, kyc_genre="F", kyc_telephone="+22244556677")],
                       "PP", _ctx())
    assert kyc["statut"] == "EXPORTEE" and kyc["valeurs"][5] == "F" and kyc["valeurs"][7] == "+22244556677"
    [div] = construire([_pp(kyc_genre="F")], "PP", _ctx())
    assert div["valeurs"][5] == "H" and A_CORRIGER in _gravites(div, "Genre")


def test_pp_nni_hors_format_exporte_a_corriger_et_tel_invalide_bloque():
    [lg] = construire([_pp(nni="123456")], "PP", _ctx())
    assert lg["statut"] == "EXPORTEE" and lg["valeurs"][2] == "123456" and A_CORRIGER in _gravites(lg, "NNI")
    [tel] = construire([_pp(eer_telephone="123")], "PP", _ctx())
    assert tel["statut"] == "EXPORTEE" and tel["valeurs"][7] is None and BLOQUANTE in _gravites(tel, "NumTel")


def test_pp_nom_non_decoupable_bloque():
    [lg] = construire([_pp(nom_complet="SIDI MOHAMED AHMED OULD", prenoms="[IDP] 1 AHMED")], "PP", _ctx())
    assert lg["gravite_max"] == BLOQUANTE and BLOQUANTE in _gravites(lg, "Nom")


def test_profil_non_identifie_et_confusion_pp_pm():
    [ni] = construire([_pp(profil="NON_IDENTIFIE")], "PP", _ctx())
    assert ni["gravite_max"] == BLOQUANTE and BLOQUANTE in _gravites(ni, "Profil")
    formes = {"forme.correspondance": {"PP": ["PERSONNE PHYSIQUE"], "PM": ["SARL"], "CJ": []},
              "forme.repli_profil": False}
    [conf] = construire([_pp(categorie="SARL")], "PP", _ctx(regles_bcm=formes))
    assert conf["gravite_max"] == BLOQUANTE and BLOQUANTE in _gravites(conf, "Profil")


def test_agence_sans_correspondance_et_multi_agences():
    [sans] = construire([_pp()], "PP", _ctx(agences={}))
    assert sans["statut"] == "EXPORTEE" and sans["valeurs"][9] is None and BLOQUANTE in _gravites(sans, "MoughataaAgence")
    [multi] = construire([_pp(nb_agences=2, nb_comptes=3)], "PP", _ctx())
    assert multi["statut"] == "EXPORTEE" and _gravites(multi) == {AVERTISSEMENT}


def test_doublons_racine_et_identifiant():
    lignes = construire([_pp(), _pp()], "PP", _ctx())
    assert all(lg["gravite_max"] == BLOQUANTE for lg in lignes)
    partage = construire([_pp(), _pp(racine="000124")], "PP", _ctx())
    assert all(A_CORRIGER in _gravites(lg, "NNI") for lg in partage)


def test_date_absente_bloquante_et_formats():
    [lg] = construire([_pp(date_entree=None)], "PP", _ctx())
    assert BLOQUANTE in _gravites(lg, "DateEntreeEnRelation")
    d = date(2024, 8, 5)
    assert formater_date(d, "TEXTE_JJ_MM_AAAA") == "05/08/2024"
    assert formater_date(d, "TEXTE_MM_JJ_AAAA_HEURE") == "08/05/2024 00:00:00"
    assert formater_date(d, "DATE_EXCEL") == datetime(2024, 8, 5)


# --------------------------------------------------------------------------- lignes PM
def test_pm_complet_et_secteurs():
    [ok] = construire([_pm()], "PM", _ctx())
    assert ok["statut"] == "EXPORTEE"
    assert ok["valeurs"] == ["004567", "00098765", "SOCIETE ALPHA", "205", "31/12/2019", "Tevragh Zeina"]
    [a_valider] = construire([_pm(secteur_orion="BTP")], "PM", _ctx())
    assert a_valider["statut"] == "EXPORTEE" and A_CORRIGER in _gravites(a_valider, "CodeSecteurActivite")
    for lib in ("SANS", "DIVERS", "", "Pêche artisanale"):
        [lg] = construire([_pm(secteur_orion=lib)], "PM", _ctx())
        assert lg["valeurs"][3] is None and BLOQUANTE in _gravites(lg, "CodeSecteurActivite"), lib
    [kyc] = construire([_pm(secteur_orion="SANS", kyc_secteur="101")], "PM", _ctx())
    assert kyc["statut"] == "EXPORTEE" and kyc["valeurs"][3] == "101"


def test_pm_nif_absent_nni_jamais_substitue():
    [lg] = construire([_pm(nif=None, nni="0012345678")], "PM", _ctx())
    assert lg["statut"] == "EXPORTEE" and lg["valeurs"][1] is None and BLOQUANTE in _gravites(lg, "NIF")


# --------------------------------------------------------------------------- fichier Excel
@pytest.mark.parametrize("fmt", ["TEXTE_JJ_MM_AAAA", "TEXTE_MM_JJ_AAAA_HEURE", "DATE_EXCEL"])
def test_xlsx_relu_conforme(fmt):
    ctx = _ctx(regles={"sic.format_date": fmt})
    lignes = construire([_pp(), _pp(racine="000124", nni="0012345679")], "PP", ctx)
    valeurs = [lg["valeurs"] for lg in lignes]
    contenu = ecrire_xlsx("PP", valeurs, fmt)
    verif = verifier_xlsx(contenu, "PP", valeurs, fmt, CODES, ctx.regles)
    assert verif["ok"], [c for c in verif["controles"] if not c["ok"]]
    ws = load_workbook(io.BytesIO(contenu))["pp"]
    assert [c.value for c in ws[1]] == list(COLONNES["PP"])
    assert ws["A2"].value == "000123" and ws["A2"].data_type == "s" and ws["C2"].value == "0012345678"
    assert ws.max_row == 3


def test_xlsx_incomplet_accepte_cellules_vides_signalees():
    ctx = _ctx(agences={})
    lignes = construire([_pp(eer_sexe=None), _pp(racine="000124", nni="0012345679")], "PP", ctx)
    valeurs = [lg["valeurs"] for lg in lignes]
    contenu = ecrire_xlsx("PP", valeurs, "TEXTE_JJ_MM_AAAA")
    verif = verifier_xlsx(contenu, "PP", valeurs, "TEXTE_JJ_MM_AAAA", CODES, ctx.regles)
    assert verif["ok"]
    vides = next(c for c in verif["controles"] if c["code"] == "obligatoires")
    assert not vides["ok"] and not vides["bloquant"] and "3 cellule(s) vide(s) sur 2 ligne(s)" in vides["detail"]
    assert load_workbook(io.BytesIO(contenu))["pp"].max_row == 3


def test_xlsx_controle_detecte_ecart():
    ctx = _ctx()
    valeurs = [construire([_pm()], "PM", ctx)[0]["valeurs"]]
    contenu = ecrire_xlsx("PM", valeurs, "TEXTE_JJ_MM_AAAA")
    verif = verifier_xlsx(contenu, "PM", valeurs + [valeurs[0]], "TEXTE_JJ_MM_AAAA", CODES, ctx.regles)
    assert not verif["ok"]
    assert {c["code"] for c in verif["controles"] if not c["ok"]} >= {"nb_lignes", "contenu"}


def test_rapport_anomalies():
    lignes = construire([_pp(eer_sexe=None), _pp(racine="000124", nni="0012345679")], "PP", _ctx())
    contenu = rapport_anomalies_xlsx([("Déclaration SIC", "PP")], lignes)
    wb = load_workbook(io.BytesIO(contenu))
    assert wb.sheetnames == ["Synthèse", "Anomalies"]
    ws = wb["Anomalies"]
    assert ws["A2"].value == "000123" and ws["C2"].value == BLOQUANTE and ws.max_row == 2

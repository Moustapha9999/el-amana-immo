# ruff: noqa: F811
"""Import Excel de la liste officielle directions / départements / services dans l'organigramme CORE.

Transaction externe annulée : rien n'est laissé en base.
"""

from __future__ import annotations

import uuid
from io import BytesIO

import pytest
from fastapi import UploadFile
from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from app.core.exceptions import AppError
from app.models import AuditLog
from app.models.organisation import OrgEntite
from app.services import org_import
from tests.test_org_referentiel import _user, db  # noqa: F401


def _fichier(lignes: list[list], entetes: list[str] | None = None) -> UploadFile:
    wb = Workbook()
    ws = wb.active
    ws.title = org_import.FEUILLE
    ws.append(entetes or ["direction", "departement", "service", "type_service", "code", "sigle"])
    for ligne in lignes:
        ws.append(ligne)
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return UploadFile(file=buf, filename="organigramme.xlsx")


async def _direction(db) -> OrgEntite:
    d = await db.scalar(select(OrgEntite).where(OrgEntite.type_entite == "DIRECTION", OrgEntite.is_active,
                                                OrgEntite.deleted_at.is_(None)).order_by(OrgEntite.code).limit(1))
    if d is None:
        pytest.skip("organigramme vide")
    return d


def test_modele():
    wb = load_workbook(BytesIO(org_import.modele()))
    assert wb.sheetnames == [org_import.FEUILLE, "Consignes"]
    assert [c.value for c in wb[org_import.FEUILLE][1]] == [c for c, _ in org_import.COLONNES]


async def test_apercu_puis_application_idempotente(db):
    actor = await _user(db)
    d = await _direction(db)
    dep, svc = f"Département test {uuid.uuid4().hex[:6]}", "Service des litiges"
    lignes = [[d.libelle, dep, None, None, None, "DTI"],
              [d.libelle.upper(), dep, svc, None, None, None],
              [d.libelle, dep, "Cellule veille", "CELLULE", None, None]]
    ap = await org_import.apercu(db, _fichier(lignes))
    assert ap["applicable"] and ap["entites"]["a_creer"] == 3 and ap["entites"]["existantes"] == 1
    assert {lg["statut"] for lg in ap["lignes"]} == {"CREER"}
    # Rien n'est écrit par l'aperçu.
    assert await db.scalar(select(OrgEntite.id).where(OrgEntite.libelle == dep)) is None

    res = await org_import.appliquer(db, actor, _fichier(lignes), ap["empreinte"])
    assert len(res["creees"]) == 3
    cree = await db.scalar(select(OrgEntite).where(OrgEntite.libelle == dep))
    assert cree.parent_id == d.id and cree.type_entite == "DEPARTEMENT" and cree.sigle == "DTI"
    assert cree.source == org_import.SOURCE and not cree.a_confirmer
    cellule = await db.scalar(select(OrgEntite).where(OrgEntite.parent_id == cree.id,
                                                      OrgEntite.libelle == "Cellule veille"))
    assert cellule.type_entite == "CELLULE"
    assert await db.scalar(select(AuditLog.id).where(AuditLog.action == "org.import.apply",
                                                     AuditLog.user_id == actor.id))

    # Rejouer le même fichier : tout existe, rien n'est recréé.
    ap2 = await org_import.apercu(db, _fichier(lignes))
    assert ap2["entites"]["a_creer"] == 0 and {lg["statut"] for lg in ap2["lignes"]} == {"EXISTANT"}
    assert (await org_import.appliquer(db, actor, _fichier(lignes), None))["creees"] == []


async def test_homonymes_sous_parents_differents(db):
    actor = await _user(db)
    d = await _direction(db)
    a, b = f"Dép A {uuid.uuid4().hex[:6]}", f"Dép B {uuid.uuid4().hex[:6]}"
    lignes = [[d.libelle, a, "Service courrier", None, None, None],
              [d.libelle, b, "Service courrier", None, None, None]]
    await org_import.appliquer(db, actor, _fichier(lignes), None)
    services = (await db.scalars(select(OrgEntite).where(OrgEntite.libelle == "Service courrier",
                                                         OrgEntite.source == org_import.SOURCE))).all()
    assert len({s.parent_id for s in services}) == 2


async def test_bloquants_aucune_ecriture(db):
    actor = await _user(db)
    d = await _direction(db)
    dep = f"Dép {uuid.uuid4().hex[:6]}"
    lignes = [[d.libelle, dep, None, None, None, None],
              [None, f"Inconnu {uuid.uuid4().hex[:6]}", "Svc", None, None, None],   # département sans direction
              [d.libelle, dep, "Svc", "BUREAU", None, None]]                         # type invalide
    ap = await org_import.apercu(db, _fichier(lignes))
    assert not ap["applicable"] and ap["nb_lignes_bloquantes"] == 2
    with pytest.raises(AppError) as e:
        await org_import.appliquer(db, actor, _fichier(lignes), None)
    assert e.value.code == "IMPORT_BLOQUE"
    assert await db.scalar(select(OrgEntite.id).where(OrgEntite.libelle == dep)) is None


async def test_entite_desactivee_et_absentes(db):
    d = await _direction(db)
    dep = OrgEntite(code=f"T{uuid.uuid4().hex[:8]}".upper(), libelle=f"Dép inactif {uuid.uuid4().hex[:6]}",
                    type_entite="DEPARTEMENT", parent_id=d.id, is_active=False)
    db.add(dep)
    await db.flush()
    ap = await org_import.apercu(db, _fichier([[d.libelle, dep.libelle, None, None, None, None]]))
    assert ap["lignes"][0]["statut"] == "ERREUR" and "désactivée" in ap["lignes"][0]["message"]
    # Les entités actives non citées sont listées, jamais supprimées.
    assert ap["absentes_du_fichier"] and all(x["id"] != str(d.id) for x in ap["absentes_du_fichier"])


async def test_fichier_modifie_et_format(db):
    actor = await _user(db)
    d = await _direction(db)
    with pytest.raises(AppError) as e:
        await org_import.appliquer(db, actor, _fichier([[d.libelle, None, None, None, None, None]]), "0" * 64)
    assert e.value.code == "FICHIER_MODIFIE"
    with pytest.raises(AppError) as e:
        await org_import.apercu(db, UploadFile(file=BytesIO(b"x"), filename="liste.csv"))
    assert e.value.code == "FORMAT_IMPORT"
    with pytest.raises(AppError) as e:
        await org_import.apercu(db, _fichier([["x"]], entetes=["libelle"]))
    assert e.value.code == "ENTETES_MANQUANTES"

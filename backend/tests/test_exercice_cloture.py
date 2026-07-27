"""Tests clôture d'exercice — snapshot archives + ouverture 148 sans report 68."""

from datetime import date
from decimal import Decimal
from types import SimpleNamespace
from unittest.mock import AsyncMock, MagicMock
from uuid import uuid4

import pytest
from openpyxl import load_workbook

from app.core.exceptions import ValidationError
from app.services.exercice_cloture_service import (
    KIND_CLOTURE,
    ExerciceClotureService,
    build_cloture_xlsx,
    periode_ouverture,
)
from app.services.recap_amortissement import _mouvements_immo


def test_periode_ouverture():
    assert periode_ouverture(2026) == "2026-12"
    assert periode_ouverture(2025) == "2025-12"


def test_build_cloture_xlsx_contains_snapshot_columns():
    content = build_cloture_xlsx(
        annee=2026,
        nature_code="TY-142010",
        lignes=[
            {
                "date_acquisition": date(2026, 1, 6),
                "quantite": 1,
                "designation": "RGLT FACT ETS KERIM",
                "valeur_brute": Decimal("203500.00"),
                "taux": Decimal("10"),
                "amt_n1": Decimal("0"),
                "dotation": Decimal("10175.00"),
                "amt_fin": Decimal("10175.00"),
                "vnc": Decimal("193325.00"),
                "agence_label": "NDB",
            }
        ],
    )
    wb = load_workbook(filename=__import__("io").BytesIO(content))
    ws = wb.active
    assert "CLÔTURE 31/12/2026" in str(ws["A2"].value)
    # Header row 4, data row 5
    assert ws.cell(5, 3).value == "RGLT FACT ETS KERIM"
    assert float(ws.cell(5, 4).value) == 203500.0
    assert float(ws.cell(5, 7).value) == 10175.0  # Exer. En. C = dotation N


def test_mouvements_snapshot_amounts_coherent():
    """amt_n1 / dotation / amt_fin alignés sur _mouvements_immo (source snapshot)."""
    immo = SimpleNamespace(
        date_acquisition=date(2024, 1, 1),
        date_fin=None,
        valeur_brute=Decimal("100000.00"),
        valeur_residuelle=Decimal("0"),
        taux=Decimal("10.0000"),
        duree_annees=10,
        duree_mois=0,
        compte_immobilisation="142010",
        compte_amortissement="148211",
        metadata_json=None,
    )
    mvts = _mouvements_immo(
        immo,
        2026,
        montants_dotation_db=[Decimal("2500"), Decimal("2500"), Decimal("2500"), Decimal("2500")],
        cumul_n1_db=Decimal("20000.00"),
        cumul_fin_n_db=Decimal("30000.00"),
        vnc_fin_n_db=Decimal("70000.00"),
    )
    assert mvts is not None
    assert mvts["amorts_cumules_n1"] == Decimal("20000.00")
    assert mvts["dotations_annee"] == Decimal("10000.00")
    assert mvts["amorts_cumules_n"] == Decimal("30000.00")
    assert mvts["vnc"] == Decimal("70000.00")


@pytest.mark.asyncio
async def test_seed_ouverture_creates_periode_with_zero_montant():
    """Ouverture {N}-12 : cumul 148 reporté, montant 0 → pas d'impact 68."""
    db = MagicMock()
    svc = ExerciceClotureService(db)
    immo = SimpleNamespace(id=uuid4())
    mvts = {
        "amorts_cumules_n": Decimal("45000.00"),
        "vnc": Decimal("55000.00"),
    }

    empty = MagicMock()
    empty.scalar_one_or_none.return_value = None
    db.execute = AsyncMock(return_value=empty)
    db.add = MagicMock()

    n = await svc._seed_ouvertures_n1(2026, [(immo, mvts)])
    assert n == 1
    db.add.assert_called_once()
    amort = db.add.call_args[0][0]
    assert amort.periode == "2026-12"
    assert amort.montant == Decimal("0.00")
    assert amort.cumul == Decimal("45000.00")
    assert amort.vnc == Decimal("55000.00")
    assert amort.valide is True


@pytest.mark.asyncio
async def test_seed_ouverture_updates_existing_forces_montant_zero():
    """Report 148 : même si un montant N-1 traînait (ancien import), on le remet à 0."""
    db = MagicMock()
    svc = ExerciceClotureService(db)
    immo = SimpleNamespace(id=uuid4())
    existing = MagicMock()
    existing.montant = Decimal("5606827.70")  # ancien seed import banque
    existing.cumul = Decimal("0")
    existing.vnc = Decimal("0")

    found = MagicMock()
    found.scalar_one_or_none.return_value = existing
    db.execute = AsyncMock(return_value=found)

    await svc._seed_ouvertures_n1(
        2026,
        [(immo, {"amorts_cumules_n": Decimal("50000"), "vnc": Decimal("150000")})],
    )
    assert existing.montant == Decimal("0.00")
    assert existing.cumul == Decimal("50000.00")
    assert existing.vnc == Decimal("150000.00")
    assert existing.valide is True
    db.add.assert_not_called()


def test_bank_snapshot_amounts_keeps_amt_n1():
    from app.services.exercice_cloture_service import _bank_snapshot_amounts

    immo = SimpleNamespace(
        valeur_brute=Decimal("100000.00"),
        metadata_json={
            "source": "import_banque",
            "bank": {
                "amt_n1": "80000.00",
                "dotation": "5000.00",
                "amt_fin": "85000.00",
                "vnc": "15000.00",
            },
        },
    )
    mvts = _bank_snapshot_amounts(immo)
    assert mvts is not None
    assert mvts["amorts_cumules_n1"] == Decimal("80000.00")
    assert mvts["dotations_annee"] == Decimal("5000.00")
    assert mvts["amorts_cumules_n"] == Decimal("85000.00")
    assert mvts["vnc"] == Decimal("15000.00")
    # Stock historique : ouverture N+1 = amt_fin
    assert mvts["cumul_ouverture_n1"] == Decimal("85000.00")
    assert mvts["vnc_ouverture_n1"] == Decimal("15000.00")


def test_bank_snapshot_arrete_courant_uses_amt_n1_for_opening():
    from app.services.exercice_cloture_service import _bank_snapshot_amounts

    immo = SimpleNamespace(
        valeur_brute=Decimal("100000.00"),
        metadata_json={
            "source": "import_banque",
            "bank": {
                "amt_n1": "80000.00",
                "dotation": "5000.00",
                "amt_fin": "85000.00",
                "vnc": "15000.00",
                "seed_mode": "arrete_courant",
            },
        },
    )
    mvts = _bank_snapshot_amounts(immo)
    assert mvts is not None
    assert mvts["cumul_ouverture_n1"] == Decimal("80000.00")
    assert mvts["vnc_ouverture_n1"] == Decimal("20000.00")


@pytest.mark.asyncio
async def test_cloture_idempotente_sans_force():
    db = MagicMock()
    svc = ExerciceClotureService(db)
    dossier = MagicMock()
    dossier.annee = 2026
    dossier.fichiers = [MagicMock(kind=KIND_CLOTURE)]
    svc._get_or_create_dossier = AsyncMock(return_value=dossier)

    with pytest.raises(ValidationError, match="clôture système existe déjà"):
        await svc.cloturer(annee=2026, force=False)

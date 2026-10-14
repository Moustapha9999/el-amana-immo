# ruff: noqa: F811
"""Réclamations — circuit 4 étapes, motif « Autre », opérations multiples, indicateurs.

Scénarios de recette (transaction annulée, aucune donnée laissée en base) :
1. motif / sous-motif du référentiel ; 2. motif « Autre » ; 3. sous-motif « Autre » ;
4. une opération ; 5. plusieurs opérations et total ; 6. modification selon droits et état ;
7. circuit complet jusqu'à la clôture SR ; 8. un service ne clôture pas ;
9. indicateurs sans double comptage ; 10. données existantes conservées après migration.
"""

from __future__ import annotations

import importlib.util
import pathlib
import uuid
from datetime import UTC, datetime, timedelta
from decimal import Decimal
from io import BytesIO

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy import select, text

from app.core.exceptions import AppError
from app.models import ReclamationReferentiel, ReclamationUnite
from app.services.reclamations import circuit, dossiers, reponses, reporting
from app.services.reclamations import constantes as C
from app.services.reclamations.acces import ReclamationsScope, permissions_effectives
from tests.test_reclamations_dossiers import _agent, _classement, _code, _saisie, _sr
from tests.test_reclamations_fondations import _unite, _user, db  # noqa: F401


async def _agent_service(db, unite_id, *extra: str) -> ReclamationsScope:
    return ReclamationsScope(
        user=await _user(db),
        permissions=permissions_effectives({"reclamations.view", "reclamations.treat", *extra}),
        unites=frozenset({unite_id}),
    )


async def _sr_unite(db) -> ReclamationUnite:
    u = await db.scalar(select(ReclamationUnite).where(ReclamationUnite.est_service_reclamations.is_(True)))
    return u or await _unite(db)


async def _canal_reponse(db) -> uuid.UUID:
    return await db.scalar(select(ReclamationReferentiel.id).where(
        ReclamationReferentiel.domaine == "CANAL_REPONSE", ReclamationReferentiel.actif.is_(True)))


async def _refus(coro, attendu: str) -> None:
    with pytest.raises((AppError, HTTPException)) as e:
        await coro
    assert _code(e.value) == attendu, e.value


def _ops(*montants: str) -> list[dict]:
    return [{"reference": f"OP-{i}", "montant": Decimal(m)} for i, m in enumerate(montants, start=1)]


# ------------------------------------------------------------- 1-3. classement
async def test_motif_sous_motif_referentiel(db):
    sr = await _sr(db)
    unite = await _unite(db)
    t, m, sm = await _classement(db, sr)
    d = await dossiers.creer(db, sr, await _saisie(db, unite.id, motif_id=uuid.UUID(m["id"]),
                                                   sous_motif_id=uuid.UUID(sm["id"])))
    assert (d["type_ref"]["id"], d["motif"], d["sous_motif"]) == (t["id"], "Motif test", "Sous-motif test")
    # Sous-motif incompatible avec le motif : refusé ; changer de motif impose de vider le sous-motif.
    _t2, m2, _sm2 = await _classement(db, sr)
    rid = uuid.UUID(d["id"])
    await _refus(dossiers.modifier(db, sr, rid, {"revision": d["revision"], "motif_id": uuid.UUID(m2["id"])}),
                 "CLASSEMENT_INCOHERENT")
    d = await dossiers.modifier(db, sr, rid, {"revision": d["revision"], "motif_id": uuid.UUID(m2["id"]),
                                              "type_id": uuid.UUID(_t2["id"]), "effacer": ["sous_motif_id"]})
    assert d["motif_ref"]["id"] == m2["id"] and d["sous_motif_ref"] is None


async def test_motif_autre_saisie_libre(db):
    sr = await _sr(db)
    unite = await _unite(db)
    t, m, _sm = await _classement(db, sr)
    d = await dossiers.creer(db, sr, await _saisie(db, unite.id, type_id=uuid.UUID(t["id"]),
                                                   motif_autre="  Litige commerçant  "))
    assert d["motif_autre"] == "Litige commerçant" and d["motif_ref"] is None
    assert d["motif"] == "Autre : Litige commerçant"
    await _refus(dossiers.creer(db, sr, await _saisie(db, unite.id, motif_id=uuid.UUID(m["id"]),
                                                      motif_autre="Doublon")), "CLASSEMENT_INCOHERENT")
    # Le filtre « Autre » retrouve les saisies libres.
    assert any(x["id"] == d["id"] for x in (await dossiers.lister(db, sr, {"motif_id": "AUTRE"}))["items"])
    # Choisir ensuite un motif du référentiel efface la saisie libre.
    d = await dossiers.modifier(db, sr, uuid.UUID(d["id"]), {"revision": d["revision"],
                                                             "motif_id": uuid.UUID(m["id"])})
    assert d["motif_autre"] is None and d["motif_ref"]["id"] == m["id"]


async def test_sous_motif_autre(db):
    sr = await _sr(db)
    unite = await _unite(db)
    _t, m, sm = await _classement(db, sr)
    d = await dossiers.creer(db, sr, await _saisie(db, unite.id, motif_id=uuid.UUID(m["id"]),
                                                   sous_motif_autre="Retrait GAB partiel"))
    assert d["sous_motif"] == "Autre : Retrait GAB partiel" and d["sous_motif_ref"] is None
    d2 = await dossiers.creer(db, sr, await _saisie(db, unite.id, motif_autre="Nouveau produit",
                                                    sous_motif_autre="Souscription"))
    assert (d2["motif"], d2["sous_motif"]) == ("Autre : Nouveau produit", "Autre : Souscription")
    await _refus(dossiers.creer(db, sr, await _saisie(db, unite.id, sous_motif_autre="Sans motif")),
                 "CLASSEMENT_INCOHERENT")
    await _refus(dossiers.creer(db, sr, await _saisie(db, unite.id, motif_autre="Libre",
                                                      sous_motif_id=uuid.UUID(sm["id"]))),
                 "CLASSEMENT_INCOHERENT")


# ------------------------------------------------------------- 4-6. opérations
async def test_une_operation(db):
    sr = await _sr(db)
    unite = await _unite(db)
    d = await dossiers.creer(db, sr, await _saisie(db, unite.id, operations=_ops("15209.5")))
    assert d["nb_operations"] == 1 and Decimal(d["montant_total"]) == Decimal("15209.50")
    assert d["operations"][0] == {**d["operations"][0], "reference": "OP-1", "montant": "15209.50",
                                  "devise": "MRU"}


async def test_plusieurs_operations_et_total(db):
    sr = await _sr(db)
    unite = await _unite(db)
    d = await dossiers.creer(db, sr, await _saisie(db, unite.id, operations=_ops("100", "200.25", "0.75")))
    assert d["nb_operations"] == 3 and Decimal(d["montant_total"]) == Decimal("301.00")
    assert [o["reference"] for o in d["operations"]] == ["OP-1", "OP-2", "OP-3"]
    ligne = (await dossiers.lister(db, sr, {"q": "OP-2"}))["items"]
    assert any(x["id"] == d["id"] and x["nb_operations"] == 3 for x in ligne)

    doublon = [{"reference": "A1", "montant": Decimal("1")}, {"reference": " a1 ", "montant": Decimal("2")}]
    await _refus(dossiers.creer(db, sr, await _saisie(db, unite.id, operations=doublon)), "OPERATION_DOUBLON")
    await _refus(dossiers.creer(db, sr, await _saisie(db, unite.id, operations=_ops("0"))), "MONTANT_INVALIDE")
    await _refus(dossiers.creer(db, sr, await _saisie(db, unite.id, operations=_ops("-5"))), "MONTANT_INVALIDE")
    await _refus(dossiers.creer(db, sr, await _saisie(db, unite.id, operations=_ops("1.005"))), "MONTANT_INVALIDE")
    await _refus(dossiers.creer(db, sr, await _saisie(db, unite.id, operations=[{"reference": " ", "montant": 3}])),
                 "OPERATION_REFERENCE_OBLIGATOIRE")


async def test_modification_operations_selon_droits_et_etat(db):
    sr = await _sr(db)
    unite = await _unite(db)
    point = await _agent(db, unites={unite.id})
    d = await dossiers.creer(db, point, await _saisie(db, unite.id, operations=_ops("10")))
    rid = uuid.UUID(d["id"])
    assert d["actions"]["modifier_operations"]
    d = await dossiers.remplacer_operations(db, point, rid, {"revision": d["revision"],
                                                              "operations": _ops("10", "20")})
    assert d["nb_operations"] == 2 and Decimal(d["montant_total"]) == Decimal("30")
    await _refus(dossiers.remplacer_operations(db, point, rid, {"revision": 1, "operations": []}),
                 "CONFLIT_REVISION")
    chrono = await dossiers.chronologie(db, sr, rid)
    assert chrono[-1]["type"] == "OPERATIONS" and chrono[-1]["donnees"]["avant"]["nb"] == 1

    # Qualifié par le SR : l'auteur (point de réception) ne modifie plus ; le SR (droit de modification) oui.
    t, _m, _sm = await _classement(db, sr)
    d = await circuit.qualifier(db, sr, rid, {"revision": d["revision"], "type_id": uuid.UUID(t["id"])})
    assert d["statut"] == "A_AFFECTER" and not (await circuit.fiche(
        db, point, await dossiers.charger(db, point, rid)))["actions"]["modifier_operations"]
    with pytest.raises(HTTPException) as e:
        await dossiers.remplacer_operations(db, point, rid, {"revision": d["revision"], "operations": []})
    assert e.value.status_code == 403
    d = await dossiers.remplacer_operations(db, sr, rid, {"revision": d["revision"], "operations": _ops("5")})
    assert d["nb_operations"] == 1

    r = await dossiers.charger(db, sr, rid)
    r.statut, r.etape, r.date_cloture = "CLOTUREE", "CLOTURE", datetime.now(UTC)
    await db.flush()
    await _refus(dossiers.remplacer_operations(db, sr, rid, {"revision": r.revision, "operations": []}),
                 "DOSSIER_TERMINE")


# ------------------------------------------------------------------ 7-8. circuit
async def test_circuit_complet_jusqua_cloture_sr(db):
    sr = await _sr(db)
    unite_sr = await _sr_unite(db)
    service = await _unite(db)
    t, m, _sm = await _classement(db, sr)
    d = await dossiers.creer(db, sr, await _saisie(db, unite_sr.id, unite_concernee_id=service.id,
                                                   operations=_ops("2500")))
    rid = uuid.UUID(d["id"])
    assert (d["statut"], d["etape"]) == ("A_QUALIFIER", "RECEPTION")
    assert d["unite_concernee_ref"]["id"] == str(service.id)

    d = await circuit.qualifier(db, sr, rid, {"revision": d["revision"], "type_id": uuid.UUID(t["id"]),
                                              "motif_id": uuid.UUID(m["id"])})
    assert (d["statut"], d["etape"]) == ("A_AFFECTER", "AFFECTATION") and d["actions"]["affecter"]

    d = await circuit.affecter(db, sr, rid, {"revision": d["revision"], "unite_id": service.id,
                                             "instructions": "Vérifier le GAB"})
    assert (d["statut"], d["etape"]) == ("EN_TRAITEMENT", "TRAITEMENT")

    agent = await _agent_service(db, service.id)
    assert (await circuit.fiche(db, agent, await dossiers.charger(db, agent, rid)))["actions"]["retourner"]
    d = await circuit.retourner(db, agent, rid, {
        "revision": d["revision"], "analyse": "Transaction non aboutie côté GAB.",
        "conclusion": "Remboursement recommandé.", "decision_proposee": "FAVORABLE"})
    assert (d["statut"], d["etape"]) == ("RETOUR_RECU", "CLOTURE")
    assert d["affectation_courante"]["recu_le"] is not None

    d = await circuit.evaluer(db, sr, rid, {"revision": d["revision"], "evaluation": "SUFFISANT"})
    assert (d["statut"], d["etape"]) == ("A_CLOTURER", "CLOTURE") and d["actions"]["cloturer"]
    assert "qualifier" not in d["actions"] and "enregistrer_reponse" not in d["actions"]

    d = await reponses.enregistrer_envoi(db, sr, rid, {
        "revision": d["revision"], "contenu": "Nous procédons au remboursement.", "resultat": "FAVORABLE",
        "canal_envoi_id": await _canal_reponse(db), "destinataire": "Client Test"})
    assert d["statut"] == "A_CLOTURER" and d["reponse_courante"]["statut"] == "ENVOYEE"
    assert d["actions"]["cloturer"]
    d = await circuit.cloturer(db, sr, rid, {"revision": d["revision"], "motif_resolution": "Remboursé"})
    assert (d["statut"], d["resultat"]) == ("CLOTUREE", "FAVORABLE") and d["cloture_par"]

    # Réouverture : l'historique conserve la clôture précédente.
    d = await circuit.rouvrir(db, sr, rid, {"revision": d["revision"], "motif": "Client conteste le montant"})
    assert (d["statut"], d["etape"]) == ("REOUVERTE", "AFFECTATION") and d["nb_reouvertures"] == 1
    types = [e["type"] for e in await dossiers.chronologie(db, sr, rid)]
    assert types[-3:] == ["REPONSE_ENVOYEE", "CLOTURE", "REOUVERTURE"]
    assert {"QUALIFICATION", "AFFECTATION", "RETOUR_SERVICE", "EVALUATION"} <= set(types)
    reouv = (await dossiers.chronologie(db, sr, rid))[-1]
    assert reouv["donnees"]["cloture_precedente"]["statut"] == "CLOTUREE" and reouv["auteur"]


async def test_service_ne_cloture_pas(db):
    sr = await _sr(db)
    service = await _unite(db)
    t, _m, _ = await _classement(db, sr)
    d = await dossiers.creer(db, sr, await _saisie(db, service.id, type_id=uuid.UUID(t["id"])))
    rid = uuid.UUID(d["id"])
    d = await circuit.affecter(db, sr, rid, {"revision": d["revision"], "unite_id": service.id})
    agent = await _agent_service(db, service.id, "reclamations.close", "reclamations.response.send")
    d = await circuit.retourner(db, agent, rid, {"revision": d["revision"], "analyse": "OK", "conclusion": "OK"})
    d = await circuit.evaluer(db, sr, rid, {"revision": d["revision"], "evaluation": "SUFFISANT"})
    d = await reponses.enregistrer_envoi(db, sr, rid, {"revision": d["revision"], "contenu": "Réponse",
                                                       "canal_envoi_id": await _canal_reponse(db),
                                                       "resultat": "DEFAVORABLE"})
    vue_service = await circuit.fiche(db, agent, await dossiers.charger(db, agent, rid))
    assert not vue_service["actions"]["cloturer"]
    await _refus(circuit.cloturer(db, agent, rid, {"revision": d["revision"], "resultat": "DEFAVORABLE"}),
                 "CLOTURE_RESERVEE_SR")
    assert (await circuit.cloturer(db, sr, rid, {"revision": d["revision"]}))["statut"] == "CLOTUREE"


async def test_renvoi_si_retour_insuffisant(db):
    sr = await _sr(db)
    service = await _unite(db)
    t, _m, _ = await _classement(db, sr)
    d = await dossiers.creer(db, sr, await _saisie(db, service.id, type_id=uuid.UUID(t["id"])))
    rid = uuid.UUID(d["id"])
    d = await circuit.affecter(db, sr, rid, {"revision": d["revision"], "unite_id": service.id})
    agent = await _agent_service(db, service.id)
    d = await circuit.retourner(db, agent, rid, {"revision": d["revision"], "analyse": "Incomplet",
                                                 "conclusion": "Manque le ticket"})
    await _refus(circuit.evaluer(db, sr, rid, {"revision": d["revision"], "evaluation": "INSUFFISANT"}),
                 "MOTIF_OBLIGATOIRE")
    d = await circuit.evaluer(db, sr, rid, {"revision": d["revision"], "evaluation": "INSUFFISANT",
                                            "motif": "Joindre le ticket GAB"})
    assert (d["statut"], d["etape"]) == ("EN_TRAITEMENT", "TRAITEMENT")
    assert d["nb_renvois"] == 1 and len(d["affectations"]) == 2


# --------------------------------------------------------------- 9. indicateurs
async def test_indicateurs_sans_double_comptage(db):
    sr = await _sr(db)
    unite = await _unite(db)
    t, m, _ = await _classement(db, sr)
    marqueur = f"KPI-{uuid.uuid4().hex[:8]}"
    base = {"objet": f"Contestation {marqueur}", "motif_id": uuid.UUID(m["id"])}
    await dossiers.creer(db, sr, await _saisie(db, unite.id, operations=_ops("100", "200", "300.50"), **base))
    await dossiers.creer(db, sr, await _saisie(db, unite.id, operations=_ops("50"), **base))
    await dossiers.creer(db, sr, await _saisie(db, unite.id, objet=f"Accueil {marqueur}",
                                               motif_autre="Attente"))
    k = await reporting.tableau_de_bord(db, sr, {"q": marqueur})
    assert k["total"] == 3 and k["nb_operations"] == 4 and k["dossiers_avec_operations"] == 2
    assert Decimal(k["montant_total"]) == Decimal("650.50")
    motifs = {x["libelle"]: x for x in k["motifs"]}
    assert motifs["Motif test"]["n"] == 2 and motifs["Motif test"]["nb_operations"] == 4
    assert Decimal(motifs["Motif test"]["montant"]) == Decimal("650.50")
    assert motifs["Autre (saisie libre)"]["n"] == 1
    assert sum(x["n"] for x in k["canaux"]) == 3
    assert sum(x["n"] for x in k["statuts"]) == 3 and sum(x["n"] for x in k["etapes"]) == 3
    assert len(k["tendance"]) == 12 and sum(m["recues"] for m in k["tendance"]) == 3
    assert sum(t["n"] for t in k["anciennete"]) == k["ouvertes"] == 3
    assert k["performance"]["cloturees"] == 0 and k["performance"]["delai_moyen_jours"] is None
    assert {x["etape"] for x in (await reporting.circuit(db, sr))["colonnes"]} == set(C.ETAPE_LIBELLES)

    classeur = load_workbook(BytesIO(await reporting.exporter(db, sr, {"q": marqueur})), read_only=True)
    assert classeur.sheetnames == ["Réclamations", "Opérations"]
    valeurs = [row for row in classeur["Réclamations"].iter_rows(values_only=True)]
    total = next(row for row in valeurs if row[0] == "Total")
    assert total[16] == 4 and Decimal(str(total[17])) == Decimal("650.50")
    pdf = await reporting.exporter(db, sr, {"q": marqueur}, format="pdf")
    assert pdf.startswith(b"%PDF")


# ------------------------------------------------------- 10. reprise des données
async def test_donnees_existantes_conservees_apres_migration(db):
    """La migration 02 reprend chaque statut historique et chaque opération déjà saisie."""
    chemin = pathlib.Path(__file__).parents[1] / "alembic/versions/20261014_reclamations_02_circuit_operations.py"
    spec = importlib.util.spec_from_file_location("migration_rc02", chemin)
    mig = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mig)
    assert set(mig.ANCIENS_STATUTS) - {"ATTENTE_INFO", "SUSPENDUE"} == set(mig.CORRESPONDANCE)
    assert set(mig.CORRESPONDANCE.values()) <= set(C.STATUT_LIBELLES)
    assert all(C.ETAPE_DU_STATUT[s] in C.ETAPE_LIBELLES for s in C.STATUT_LIBELLES)

    version = await db.scalar(text("SELECT version_num FROM alembic_version"))
    # Révisions préfixées par la date : toute révision ultérieure (autre module compris) inclut la 02.
    assert version >= mig.revision
    perdues = await db.scalar(text(
        "SELECT count(*) FROM reclamations r WHERE (NULLIF(btrim(r.operation_reference), '') IS NOT NULL "
        "OR r.montant_conteste > 0) AND NOT EXISTS (SELECT 1 FROM reclamation_operations o "
        "WHERE o.reclamation_id = r.id)"))
    assert perdues == 0
    hors_circuit = await db.scalar(text(
        "SELECT count(*) FROM reclamations WHERE statut <> ALL(:s) OR etape <> ALL(:e)"),
        {"s": list(C.STATUT_LIBELLES), "e": list(C.ETAPE_LIBELLES)})
    assert hors_circuit == 0
    assert await db.scalar(text("SELECT count(*) FROM reclamations WHERE date_reclamation IS NULL")) == 0
    # Les anciens codes restent lisibles dans l'historique.
    assert dossiers.libelle_statut("REPONSE_A_PREPARER") == "Réponse à préparer"


async def test_aucune_echeance_calculee_a_la_creation(db):
    """« Délais & alertes » retiré : la création et la qualification ne posent plus d'échéance."""
    sr = await _sr(db)
    unite_sr = await _sr_unite(db)
    t, _m, _ = await _classement(db, sr)
    d = await dossiers.creer(db, sr, await _saisie(
        db, unite_sr.id, type_id=uuid.UUID(t["id"]),
        date_premiere_reception=datetime.now(UTC) - timedelta(days=1)))
    assert not {"echeance", "etat_delai", "situation_delai", "en_retard", "regle_delai"} & d.keys()
    r = await dossiers.charger(db, sr, uuid.UUID(d["id"]))
    assert (r.echeance, r.etat_delai, r.regle_delai_id) == (None, None, None)

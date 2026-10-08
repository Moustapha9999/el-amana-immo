"""Phases 4–6 : rapprochement d'imports, classification configurable, filtrage / alertes."""

from __future__ import annotations

import uuid
from datetime import date
from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import select, text
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine
from sqlalchemy.pool import NullPool

from app.core.config import get_settings
from app.data.plateforme_catalogue import FUNCTIONAL_PERMISSIONS, ROLE_PERMISSIONS
from app.models import (
    ClienteleClassification,
    ClienteleClient,
    ClienteleCompte,
    User,
)
from app.services.clientele.classification import ClienteleClassificationService
from app.services.clientele.filtrage import ClienteleFiltrageService
from app.services.clientele.import_service import ClienteleImportService
from app.services.clientele.rapprochement import ClienteleRapprochementService
from app.services.clientele.service import Ctx
from tests.test_clientele_consolidation import brute
from tests.test_clientele_import import classeur

EXTRACTION_A = date(2026, 9, 30)
EXTRACTION_B = date(2026, 10, 31)
RACINES = ("000001", "000002", "000003")


@pytest.fixture
async def db():
    engine = create_async_engine(get_settings().database_url, poolclass=NullPool)
    try:
        async with engine.connect() as conn:
            pret = await conn.scalar(text(
                "SELECT to_regclass('public.clientele_rapprochements') IS NOT NULL"))
    except Exception:
        pret = False
    if not pret:
        await engine.dispose()
        pytest.skip("Tables phase 4 absentes (migration 20261008_clientele_03)")
    session = async_sessionmaker(engine, expire_on_commit=False)()
    await _nettoyer(session)
    await session.commit()
    try:
        yield session
    finally:
        await session.rollback()
        await _nettoyer(session)
        await session.commit()
        await session.close()
        await engine.dispose()


_IN = "('000001','000002','000003')"


async def _nettoyer(session) -> None:
    await session.execute(text(f"""
        DELETE FROM clientele_alerte_justificatifs WHERE alerte_id IN (
            SELECT id FROM clientele_alertes WHERE racine_client IN {_IN})
    """))
    await session.execute(text(f"""
        DELETE FROM clientele_alerte_evenements WHERE alerte_id IN (
            SELECT id FROM clientele_alertes WHERE racine_client IN {_IN})
    """))
    await session.execute(text(
        f"DELETE FROM clientele_filtrage_empreintes WHERE racine_client IN {_IN}"))
    await session.execute(text(f"DELETE FROM clientele_alertes WHERE racine_client IN {_IN}"))
    await session.execute(text(
        f"DELETE FROM clientele_classif_historique WHERE racine_client IN {_IN}"))
    await session.execute(text(
        f"DELETE FROM clientele_classifications WHERE racine_client IN {_IN}"))
    await session.execute(text("DELETE FROM clientele_classif_regles WHERE motif LIKE 'TEST %'"))
    await session.execute(text("DELETE FROM clientele_rapprochement_ecarts"))
    await session.execute(text("DELETE FROM clientele_rapprochements"))
    await session.execute(text("DELETE FROM clientele_import_comptes"))
    await session.execute(text("DELETE FROM clientele_import_clients"))
    await session.execute(text("DELETE FROM clientele_import_anomalies"))
    await session.execute(text("DELETE FROM clientele_import_lignes"))
    await session.execute(text("DELETE FROM clientele_imports"))
    await session.execute(text(f"DELETE FROM clientele_comptes WHERE racine_client IN {_IN}"))
    await session.execute(text(f"DELETE FROM clientele_clients WHERE racine_client IN {_IN}"))
    await session.execute(text(
        "DELETE FROM clientele_filtrage_entrees WHERE raison_sociale LIKE 'TEST %' "
        "OR nom LIKE 'TEST %'"))


async def _ctx(db) -> Ctx:
    user = await db.scalar(select(User).limit(1))
    if not user:
        pytest.skip("aucun utilisateur en base de test")
    return Ctx(user=user, permissions={"clientele.admin"})


async def _importer(db, ctx, nom: str, extraction: date, *lignes: dict) -> uuid.UUID:
    svc = ClienteleImportService(db, ctx)
    a = await svc.analyser(classeur(*lignes), nom)
    c = await svc.confirmer(uuid.UUID(a["id"]), {"date_extraction": extraction.isoformat()})
    return uuid.UUID(c["id"])


def test_catalogue_phases_456():
    perms = {c for c, _l, m in FUNCTIONAL_PERMISSIONS if m == "clientele"}
    assert {"clientele.rapprochement.execute", "clientele.classif.execute",
            "clientele.classif.admin", "clientele.filtrage.decide"} <= perms
    assert "clientele.classif.admin" not in ROLE_PERMISSIONS["clientele.gestionnaire"]
    assert "clientele.classif.view" in ROLE_PERMISSIONS["clientele.lecteur"]
    assert "clientele.filtrage.decide" not in ROLE_PERMISSIONS["clientele.lecteur"]


@pytest.mark.asyncio
async def test_rapprochement_septembre_octobre(db):
    ctx = await _ctx(db)
    id_a = await _importer(
        db, ctx, "sept.xlsx", EXTRACTION_A,
        brute("000001", "00000100001"),
        brute("000001", "00000100002"),
        brute("000002", "00000200001"),
    )
    id_b = await _importer(
        db, ctx, "oct.xlsx", EXTRACTION_B,
        brute("000001", "00000100001", **{"R/N - Statut résident": "N"}),
        brute("000001", "00000100002", ETAT_COMPTE="Fermé", **{"R/N - Statut résident": "N"}),
        brute("000001", "00000100003", **{"R/N - Statut résident": "N"}),
        brute("000003", "00000300001"),
    )
    n_clients = await db.scalar(select(ClienteleClient).where(
        ClienteleClient.racine_client == "000002"))
    assert n_clients is not None  # 000002 absent de l'extraction B mais toujours en base

    svc = ClienteleRapprochementService(db, ctx)
    rap = await svc.calculer(id_a, id_b)
    s = rap["synthese"]
    assert s["clients_nouveaux"] == 1  # 000003
    assert s["clients_absents"] == 1   # 000002
    assert s["comptes_nouveaux"] == 2  # 00000100003 + 00000300001
    assert s["comptes_absents"] == 1   # 00000200001
    assert s["comptes_fermes"] == 1
    assert s["clients_modifies"] >= 1  # résidence
    assert "ABSENT_EXTRACTION n'est pas une suppression" in s["regle"]

    still = await db.scalar(select(ClienteleCompte).where(ClienteleCompte.compte == "00000200001"))
    assert still is not None

    ecarts = await svc.ecarts(uuid.UUID(rap["id"]), categorie="ABSENT_EXTRACTION")
    racines_abs = {e["racine_client"] for e in ecarts["items"]}
    assert "000002" in racines_abs

    xlsx = await svc.exporter_excel(uuid.UUID(rap["id"]))
    wb = load_workbook(BytesIO(xlsx))
    assert "Synthèse" in wb.sheetnames
    pdf = await svc.exporter_pdf(uuid.UUID(rap["id"]))
    assert pdf.startswith(b"%PDF")


@pytest.mark.asyncio
async def test_classification_moteur_configurable_et_manuel(db):
    ctx = await _ctx(db)
    await _importer(db, ctx, "c.xlsx", EXTRACTION_B, brute("000001", "00000100001",
                                                           **{"R/N - Statut résident": "N"}))
    svc = ClienteleClassificationService(db, ctx)
    ref = await svc.referentiel()
    assert {n["code"] for n in ref["niveaux"]} == {"FAIBLE", "MOYEN", "ELEVE", "INTERDIT"}
    version = next(v for v in ref["versions"] if v["statut"] == "ACTIVE")
    critere = next(c for c in ref["criteres"] if c["code"] == "RESIDENCE")

    with pytest.raises(Exception) as vide:
        await svc.appliquer_moteur()
    assert "AUCUNE_REGLE" in str(vide.value) or "règle" in str(vide.value).lower()

    await svc.creer_regle(uuid.UUID(version["id"]), {
        "critere_id": critere["id"], "niveau_cible": "ELEVE", "operateur": "EGAL",
        "champ_source": "statut_resident", "portee": "CLIENT", "valeur": "N",
        "motif": "TEST non-résident", "priorite": 10, "poids": 5, "actif": True,
    })
    bilan = await svc.appliquer_moteur()
    assert bilan["modifies"] >= 1
    courante = await svc.courante("000001")
    assert courante["classification"]["niveau"] == "ELEVE"
    assert "non-résident" in (courante["classification"]["motif_classement"] or "").lower() or (
        courante["classification"]["motifs"])

    await svc.modifier_individuel("000001", {
        "niveau": "MOYEN", "motif_classement": "TEST ajustement analyste", "motif_risque": "revue",
    })
    assert (await svc.courante("000001"))["classification"]["niveau"] == "MOYEN"
    assert (await svc.courante("000001"))["classification"]["source"] == "MANUEL"

    await svc.appliquer_moteur()
    assert (await svc.courante("000001"))["classification"]["source"] == "MANUEL"
    assert (await svc.courante("000001"))["classification"]["niveau"] == "MOYEN"

    await svc.appliquer_moteur(forcer=True)
    assert (await svc.courante("000001"))["classification"]["niveau"] == "ELEVE"
    hist = courante["historique"] if False else (await svc.historique("000001"))
    assert len(hist) >= 3
    with pytest.raises(Exception):
        await svc.modifier_individuel("000001", {
            "niveau": "FAIBLE", "motif_classement": "x", "racine_client": "000099",
        })


@pytest.mark.asyncio
async def test_classification_excel_et_racine_immutable(db):
    ctx = await _ctx(db)
    await _importer(db, ctx, "e.xlsx", EXTRACTION_B, brute("000001", "00000100001"))
    svc = ClienteleClassificationService(db, ctx)
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(["RACINE_CLIENT", "NIVEAU", "MOTIF_RISQUE", "MOTIF_CLASSEMENT"])
    ws.append(["000001", "INTERDIT", "liste", "TEST excel"])
    buf = BytesIO()
    wb.save(buf)
    analyse = await svc.analyser_excel(buf.getvalue())
    assert analyse["nb_valides"] == 1
    await svc.confirmer_excel(analyse["lignes"])
    assert (await db.get(ClienteleClassification, "000001")).niveau == "INTERDIT"
    assert (await db.scalar(select(ClienteleClient.racine_client)
                            .where(ClienteleClient.racine_client == "000001"))) == "000001"


@pytest.mark.asyncio
async def test_filtrage_potentiel_workflow_faux_positif(db):
    ctx = await _ctx(db)
    await _importer(
        db, ctx, "f.xlsx", EXTRACTION_B,
        brute("000001", "00000100001", LISTE_INTERDICTION="OFAC-TEST", RAISON_SOCIAL="CLIENT UNIQUE TEST"),
    )
    svc = ClienteleFiltrageService(db, ctx)
    scan = await svc.scanner()
    assert scan["alertes_crees"] >= 1
    page = await svc.lister()
    alerte = next(i for i in page["items"] if i["motif"] == "LISTE_INTERDICTION")
    assert alerte["statut"] == "NOUVELLE"
    assert alerte["correspondance"]["potentielle"] is True

    d = await svc.decider(uuid.UUID(alerte["id"]), {"statut": "A_ANALYSER"})
    d = await svc.decider(uuid.UUID(alerte["id"]), {"statut": "EN_INVESTIGATION"})
    d = await svc.decider(uuid.UUID(alerte["id"]), {
        "statut": "FAUX_POSITIF", "commentaire": "homonyme contrôlé",
    })
    assert d["statut"] == "FAUX_POSITIF"
    await svc.decider(uuid.UUID(alerte["id"]), {"statut": "CLOTUREE", "commentaire": "clôturé"})

    scan2 = await svc.scanner()
    assert scan2["alertes_crees"] >= 1
    page2 = await svc.lister(statut="NOUVELLE")
    nouvelle = next(i for i in page2["items"] if i["motif"] == "LISTE_INTERDICTION")
    assert nouvelle["precedent_faux_positif"] is True
    assert nouvelle["id"] != alerte["id"]
    assert nouvelle["statut"] != "CONFIRMEE"

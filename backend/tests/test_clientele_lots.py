from io import BytesIO

from openpyxl import Workbook, load_workbook

from app.data.clientele_sources import libelle_sources
from app.services.clientele.lots import _detecter, normaliser_niveau, normaliser_racine


def test_racine_zeros_restaures():
    assert normaliser_racine(1234) == "001234"
    assert normaliser_racine(1234.0) == "001234"
    assert normaliser_racine(" 059904 ") == "059904"


def test_racine_invalide():
    assert normaliser_racine(None) is None
    assert normaliser_racine("") is None
    assert normaliser_racine("12A456") is None
    assert normaliser_racine(1234567) is None
    assert normaliser_racine(12.5) is None


def test_niveau_normalise():
    assert normaliser_niveau("Élevé") == "ELEVE"
    assert normaliser_niveau("risque moyen") == "MOYEN"
    assert normaliser_niveau("FAIBLE") == "FAIBLE"
    assert normaliser_niveau("Interdit") == "INTERDIT"
    assert normaliser_niveau("") is None
    assert normaliser_niveau("???") is None


def _classeur(lignes: list[list]) -> object:
    wb = Workbook()
    ws = wb.active
    ws.title = "Feuil1"
    for r in lignes:
        ws.append(r)
    buf = BytesIO()
    wb.save(buf)
    return load_workbook(BytesIO(buf.getvalue()), read_only=True, data_only=True)


def test_detection_entete_decalee():
    wb = _classeur([
        ["Situation des comptes PP et PM"],
        [],
        ["CLIENT", "RAISON SOCIALE", "CLASSE RISQUE LBC FT"],
        [1234, "ALPHA", "Moyen"],
    ])
    feuille, ligne, entetes = _detecter(wb)
    assert (feuille, ligne) == ("Feuil1", 3)
    assert entetes[:3] == ["CLIENT", "RAISON SOCIALE", "CLASSE RISQUE LBC FT"]


def test_colonne_sans_titre_selectionnable():
    wb = _classeur([[None, "NOM"], [1234, "ALPHA"]])
    _, _, entetes = _detecter(wb)
    assert entetes[0] == "(colonne A)"


def test_libelles_sources():
    assert libelle_sources("Matrice Faible / V4 100.") == (
        "Matrice des risques LBC/FT BEA Faible / Fiche de scoring 100.")
    assert libelle_sources("Conflit V1/V4") == "Conflit Matrice des risques LBC/FT BEA / Fiche de scoring"
    assert libelle_sources("MATRICE_V1 SCORING_V4 niveau_matrice") == "MATRICE_V1 SCORING_V4 niveau_matrice"
    assert libelle_sources("Matrice des risques LBC/FT BEA") == "Matrice des risques LBC/FT BEA"

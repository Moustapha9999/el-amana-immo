# ruff: noqa: F811
"""Réclamations — règles par motif, multi-devises, doublons, import, rattachement, régularisation.

Scénarios de recette (transaction annulée ; fichiers GED supprimés en fin de test) :
1. règle normalisée / refusée ; 2. champs obligatoires du motif ; 3. nature D/C incohérente ;
4. extourne partielle ; 5. montant incorrect ; 6. mauvais compte ; 7. devises (liste, défaut, refus) ;
8. doublon exact (erreur) et potentiel (avertissement) ; 9. zéros de tête conservés ;
10. totaux par devise sans mélange ; 11. récapitulatif avant transmission ; 12. vérification à blanc ;
13. transmission bloquée si dossier incomplet ; 14. rattachement département → service ;
15. régularisation à effectuer puis confirmée (preuve, droit) ; 16. décision défavorable ;
17. pièce par opération ; 18. import : réimport, limite de taille, rapport d'erreurs ;
19. modèle d'import par motif ; 20. filtres ; 21. tableau de bord par devise ; 22. paramétrage des devises.
"""

from __future__ import annotations

import uuid
from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO

import pytest
from fastapi import HTTPException, UploadFile
from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from app.core.exceptions import AppError
from app.models import GedDocument, ReclamationReferentiel
from app.models.organisation import OrgEntite
from app.services.ged_service import GedService
from app.services.reclamations import circuit, documents, dossiers, imports, regles, reporting
from app.services.reclamations import constantes as C
from app.services.reclamations import operations as ops
from app.services.reclamations.acces import ReclamationsScope, permissions_effectives
from app.services.reclamations.regles import Exigences
from tests.test_reclamations_dossiers import PDF, _classement, _code, _saisie, _sr
from tests.test_reclamations_fondations import _unite, _user, db  # noqa: F401

EXTOURNE_DEBIT = {"champs_requis": ["reference", "date_operation", "compte_reference", "montant", "devise",
                                    "montant_conteste"],
                  "operations_min": 1, "nature_attendue": "DEBIT", "regularisation": True}
DEVISES = frozenset({"MRU", "USD", "EUR"})


def _ex(**regle) -> Exigences:
    return Exigences(regle=regles.normaliser_regle(regle), devises=DEVISES)


def _codes(lg: ops.Ligne) -> set[str]:
    return {c for c, _m in lg.erreurs}


async def _refus(coro, attendu: str) -> None:
    with pytest.raises((AppError, HTTPException)) as e:
        await coro
    assert _code(e.value) == attendu, e.value


async def _motif_regle(db, s, regle: dict) -> tuple[uuid.UUID, uuid.UUID]:
    t, m, _sm = await _classement(db, s)
    await regles.enregistrer(db, s, uuid.UUID(m["id"]), regle)
    return uuid.UUID(t["id"]), uuid.UUID(m["id"])


def _purger(db, rid: uuid.UUID):
    async def _f():
        ged = GedService(db)
        for doc in (await db.scalars(select(GedDocument).where(GedDocument.entity_id == str(rid)))).all():
            ged.absolute_path(doc.stored_path).unlink(missing_ok=True)
    return _f


# ------------------------------------------------------------------ 1-11 : sans base
def test_regle_normalisee_et_refusee():
    with pytest.raises(AppError) as e:
        regles.normaliser_regle({"champs_requis": ["iban"]})
    assert e.value.code == "REGLE_INVALIDE"
    with pytest.raises(AppError):
        regles.normaliser_regle({"nature_attendue": "AUTRE"})
    assert regles.normaliser_regle({"controle": "DOUBLE"})["operations_min"] == 2
    assert regles.normaliser_regle({"cle_inconnue": 1})["champs_requis"] == ["reference", "montant"]


def test_champs_obligatoires_du_motif():
    lg = ops.lire({"reference": "R1", "montant": "100"}, 1, _ex(**EXTOURNE_DEBIT))
    champs = {ch for ch in lg.erreurs_champs}
    assert {"date_operation", "compte_reference", "devise", "montant_conteste"} <= champs
    assert "DEVISE_OBLIGATOIRE" in _codes(lg)
    ok = ops.lire({"reference": "R1", "montant": "100", "montant_conteste": "100", "devise": "MRU",
                   "date_operation": "01/10/2026", "compte_reference": "00012345678"}, 1, _ex(**EXTOURNE_DEBIT))
    assert ok.statut == ops.OK


def test_nature_incoherente():
    lg = ops.lire({"reference": "R1", "montant": "10", "devise": "MRU", "nature_mouvement": "C"}, 1,
                  _ex(nature_attendue="DEBIT"))
    assert "NATURE_INCOHERENTE" in _codes(lg)
    assert ops.lire({"reference": "R1", "montant": "10", "devise": "MRU", "nature_mouvement": "D"}, 1,
                    _ex(nature_attendue="DEBIT")).nature_mouvement == "DEBIT"
    assert "NATURE_INVALIDE" in _codes(ops.lire({"reference": "R", "montant": "1", "devise": "MRU",
                                                 "nature_mouvement": "X"}, 1, _ex()))


def test_extourne_partielle():
    ex = _ex(controle="PARTIELLE", champs_requis=["reference", "montant", "montant_conteste", "devise"])
    assert "EXTOURNE_NON_PARTIELLE" in _codes(ops.lire(
        {"reference": "R", "montant": "100", "montant_conteste": "100", "devise": "MRU"}, 1, ex))
    ok = ops.lire({"reference": "R", "montant": "100", "montant_conteste": "40", "devise": "MRU"}, 1, ex)
    assert ok.statut == ops.AVERTISSEMENT and "ORION" in ok.avertissements[0]
    assert "MONTANT_CONTESTE_SUPERIEUR" in _codes(ops.lire(
        {"reference": "R", "montant": "100", "montant_conteste": "150", "devise": "MRU"}, 1, ex))


def test_montant_incorrect():
    ex = _ex(controle="MONTANT_INCORRECT", champs_requis=["reference", "montant", "montant_attendu", "devise"])
    base = {"reference": "R", "montant": "100", "devise": "MRU"}
    assert "MONTANT_ATTENDU_IDENTIQUE" in _codes(ops.lire({**base, "montant_attendu": "100"}, 1, ex))
    assert "ECART_INCOHERENT" in _codes(ops.lire({**base, "montant_attendu": "80", "montant_conteste": "30"}, 1, ex))
    assert ops.lire({**base, "montant_attendu": "80", "montant_conteste": "20"}, 1, ex).statut == ops.OK
    assert "CHAMP_REQUIS" in _codes(ops.lire(base, 1, ex))


def test_mauvais_compte():
    ex = _ex(controle="MAUVAIS_COMPTE", champs_requis=["reference", "compte_reference", "compte_attendu",
                                                       "montant", "devise"])
    base = {"reference": "R", "montant": "5", "devise": "MRU", "compte_reference": "00011111111"}
    assert "COMPTE_ATTENDU_IDENTIQUE" in _codes(ops.lire({**base, "compte_attendu": "00011111111"}, 1, ex))
    assert ops.lire({**base, "compte_attendu": "00022222222"}, 1, ex).statut == ops.OK


def test_devises():
    ex = _ex()
    assert "DEVISE_INVALIDE" in _codes(ops.lire({"reference": "R", "montant": "1", "devise": "XOF"}, 1, ex))
    assert ops.lire({"reference": "R", "montant": "1", "devise": "usd"}, 1, ex).devise == "USD"
    # Saisie / import : MRU retenue (avertissement) si le motif n'exige pas la devise…
    defaut = ops.lire({"reference": "R", "montant": "1"}, 1, ex, devise_defaut=C.DEVISE)
    assert defaut.devise == C.DEVISE and defaut.statut == ops.AVERTISSEMENT
    # … mais jamais quand le motif l'exige.
    exige = ops.lire({"reference": "R", "montant": "1"}, 1, _ex(champs_requis=["reference", "devise"]),
                     devise_defaut=C.DEVISE)
    assert exige.devise is None and "DEVISE_OBLIGATOIRE" in _codes(exige)


def test_doublons_exacts_et_potentiels():
    ex = _ex()
    brut = {"reference": "R1", "montant": "100", "devise": "MRU", "compte_reference": "00012345678",
            "date_operation": "01/10/2026"}
    lignes = [ops.lire(brut, 1, ex), ops.lire(brut, 2, ex), ops.lire({**brut, "reference": "R2"}, 3, ex),
              ops.lire({**brut, "compte_reference": "00099999999", "montant": "7"}, 4, ex)]
    ops.controler(lignes)
    assert lignes[0].statut == ops.OK
    assert "OPERATION_DOUBLON" in _codes(lignes[1])
    assert lignes[2].statut == ops.AVERTISSEMENT and "doublon potentiel" in lignes[2].avertissements[0]
    assert "déjà utilisée" in lignes[3].avertissements[0]
    # Rien n'est supprimé : toutes les lignes restent présentes.
    assert len(lignes) == 4


def test_zeros_de_tete_conserves():
    lg = ops.lire({"reference": "000123", "montant": 1, "devise": "MRU", "compte_reference": 1234,
                   "client_racine": 42}, 1, _ex())
    assert lg.reference == "000123" and lg.compte_reference == "00000001234" and lg.client_racine == "000042"
    assert ops.compte_texte("00001 00001 00012345678 12") == "00012345678"


def test_totaux_par_devise_sans_melange():
    ex = _ex()
    lignes = [ops.lire({"reference": f"R{i}", "montant": m, "devise": d}, i, ex)
              for i, (m, d) in enumerate([("100", "MRU"), ("50", "USD"), ("25.50", "MRU")], start=1)]
    totaux = {x["devise"]: x for x in ops.par_devise(lignes)}
    assert ops.par_devise(lignes)[0]["devise"] == "MRU"
    assert Decimal(totaux["MRU"]["montant"]) == Decimal("125.50") and totaux["MRU"]["nb"] == 2
    assert Decimal(totaux["USD"]["montant"]) == Decimal("50.00")
    assert ops.montant_unique(lignes) == (None, None)
    assert ops.montant_unique(lignes[::2]) == (Decimal("125.50"), "MRU")


def test_recapitulatif_avant_transmission():
    ex = _ex(controle="DOUBLE", champs_requis=["reference", "montant", "devise", "role_operation"],
             justificatif_requis=True)
    lignes = [ops.lire({"reference": "R1", "montant": "10", "devise": "MRU", "role_operation": "INITIALE"}, 1, ex),
              ops.lire({"reference": "R2", "montant": "10", "devise": "EUR", "role_operation": "INITIALE"}, 2, ex),
              ops.lire({"reference": "R3", "montant": "10", "devise": "MRU"}, 3, ex)]
    rec = ops.recapitulatif(lignes, ex)
    assert rec["nb_operations"] == 3 and rec["nb_lignes_en_erreur"] == 1 and rec["devises_multiples"]
    assert rec["montant_total"] is None
    assert {x["devise"]: x["montant"] for x in rec["par_devise"]} == {"MRU": "10.00", "EUR": "10.00"}
    manques = " ".join(rec["a_completer_avant_transmission"])
    assert "initiale" in manques and "justificatif" in manques


# ------------------------------------------------------------------ 12-22 : base
async def test_verification_a_blanc(db):
    s = await _sr(db)
    _t, motif = await _motif_regle(db, s, EXTOURNE_DEBIT)
    res = await dossiers.verifier_operations(db, s, {"motif_id": motif, "operations": [
        {"reference": "R1", "montant": "100", "devise": "MRU", "nature_mouvement": "CREDIT"}]})
    detail = res["lignes"][0]["erreurs_detail"]
    assert {d["champ"] for d in detail} >= {"nature_mouvement", "date_operation", "montant_conteste"}
    assert res["recapitulatif"]["nb_erreurs"] >= 3 and res["recapitulatif"]["regle"]["nature_attendue"] == "DEBIT"


async def test_transmission_bloquee_si_incomplet(db):
    s = await _sr(db)
    unite = await _unite(db)
    regle = {"champs_requis": ["reference", "montant", "devise", "role_operation"], "controle": "DOUBLE",
             "operations_min": 2}
    t, motif = await _motif_regle(db, s, regle)
    # Une saisie non conforme à la règle est refusée dès l'enregistrement.
    await _refus(dossiers.creer(db, s, await _saisie(db, unite.id, type_id=t, motif_id=motif, operations=[
        {"reference": "R1", "montant": Decimal("10"), "devise": "MRU"}])), "CHAMP_REQUIS")
    # Enregistré avec une seule opération (incomplet), il ne peut pas être transmis.
    d = await dossiers.creer(db, s, await _saisie(db, unite.id, type_id=t, motif_id=motif, operations=[
        {"reference": "R1", "montant": Decimal("10"), "devise": "MRU", "role_operation": "INITIALE"}]))
    assert d["a_completer_avant_transmission"]
    await _refus(circuit.traiter_interne(db, s, uuid.UUID(d["id"]), {"revision": d["revision"], "motif": "x"}),
                 "DOSSIER_INCOMPLET")


async def _org_departement_service(db) -> tuple[OrgEntite, OrgEntite, OrgEntite]:
    arbre = await dossiers.org.charger_arbre(db)
    svc = dep = None
    for e in arbre.entites.values():
        if e.type_entite in C.TYPES_ORG_SERVICE and e.is_active:
            deps = [a for a in arbre.chemin(e.id)[:-1] if a.type_entite == "DEPARTEMENT"]
            if deps:
                svc, dep = e, deps[-1]
                break
    if svc is None:
        pytest.skip("organigramme sans service rattaché à un département")
    autre = next((e for e in arbre.entites.values() if e.type_entite == "DEPARTEMENT" and e.is_active
                  and e.id not in {a.id for a in arbre.chemin(svc.id)}), None)
    if autre is None:
        pytest.skip("un seul département")
    return dep, svc, autre


async def test_rattachement_departement_service(db):
    s = await _sr(db)
    unite = await _unite(db)
    dep, svc, autre = await _org_departement_service(db)
    await _refus(dossiers.creer(db, s, await _saisie(db, unite.id, org_departement_id=autre.id,
                                                     org_service_id=svc.id)), "SERVICE_HORS_DEPARTEMENT")
    point = await db.scalar(select(ReclamationReferentiel.id).where(
        ReclamationReferentiel.domaine == "POINT_RECEPTION", ReclamationReferentiel.code == C.POINT_SERVICE_CONCERNE))
    if point is not None:
        await _refus(dossiers.creer(db, s, await _saisie(db, unite.id, point_reception_id=point)),
                     "SERVICE_OBLIGATOIRE")
    # Service seul : département déduit de la hiérarchie réelle.
    d = await dossiers.creer(db, s, await _saisie(db, unite.id, org_service_id=svc.id))
    assert d["service_ref"]["id"] == str(svc.id) and d["departement_ref"]["id"] == str(dep.id)
    # Filtre par département : sous-arbre complet.
    arbre = await dossiers.org.charger_arbre(db)
    liste = await dossiers.lister(db, s, {"org_departement_ids": arbre.descendants(dep.id)})
    assert d["id"] in {x["id"] for x in liste["items"]}
    liste = await dossiers.lister(db, s, {"org_service_ids": arbre.descendants(autre.id)})
    assert d["id"] not in {x["id"] for x in liste["items"]}


async def _dossier_a_regulariser(db, s, unite) -> dict:
    t, motif = await _motif_regle(db, s, {**EXTOURNE_DEBIT, "champs_requis": ["reference", "montant", "devise"]})
    d = await dossiers.creer(db, s, await _saisie(db, unite.id, type_id=t, motif_id=motif, operations=[
        {"reference": "EXT-1", "montant": Decimal("300"), "montant_conteste": Decimal("300"), "devise": "MRU",
         "nature_mouvement": "DEBIT"}]))
    rid = uuid.UUID(d["id"])
    d = await circuit.traiter_interne(db, s, rid, {"revision": d["revision"], "motif": "Débit erroné confirmé"})
    return d


async def test_regularisation_a_effectuer_puis_confirmee(db):
    s = await _sr(db)
    unite = await _unite(db)
    d = await _dossier_a_regulariser(db, s, unite)
    rid = uuid.UUID(d["id"])
    try:
        d = await circuit.cloturer(db, s, rid, {"revision": d["revision"], "resultat": "FAVORABLE"})
        # Décision favorable ≠ extourne exécutée.
        assert d["statut"] == "CLOTUREE" and d["regularisation"]["statut"] == "A_EFFECTUER"
        assert d["actions"]["confirmer_regularisation"]
        await _refus(circuit.archiver(db, s, rid, {"revision": d["revision"]}), "REGULARISATION_EN_ATTENTE")
        data = {"revision": d["revision"], "effectuee_le": date.today(), "reference": "ECR-2026-001"}
        await _refus(circuit.confirmer_regularisation(db, s, rid, data), "PREUVE_OBLIGATOIRE")
        await _refus(circuit.confirmer_regularisation(db, s, rid, {**data, "effectuee_le": date.today()
                                                                    + timedelta(days=1)}), "DATE_FUTURE")
        await _refus(circuit.confirmer_regularisation(db, s, rid, {**data, "reference": " "}),
                     "REFERENCE_OBLIGATOIRE")
        sans_droit = ReclamationsScope(user=await _user(db), membre_service_central=True,
                                       permissions=permissions_effectives({"reclamations.view", "reclamations.close",
                                                                           "reclamations.scope.all"})
                                       - {"reclamations.regularisation.confirm"})
        with pytest.raises(HTTPException) as e:
            await circuit.confirmer_regularisation(db, sans_droit, rid, data)
        assert e.value.status_code == 403
        await documents.deposer(db, s, rid, UploadFile(file=BytesIO(PDF), filename="extourne.pdf"),
                                "PREUVE_REGULARISATION", "Avis d'extourne", operation_reference="EXT-1")
        d = await dossiers.detail(db, s, await dossiers.charger(db, s, rid))
        d = await circuit.confirmer_regularisation(db, s, rid, {**data, "revision": d["revision"],
                                                                "commentaire": "Extourne passée le jour même"})
        reg = d["regularisation"]
        assert reg["statut"] == "EFFECTUEE" and reg["reference"] == "ECR-2026-001" and reg["agent"]
        assert d["statut"] == "CLOTUREE"
        evts = {e["type_evenement"] for e in d.get("chronologie", d.get("evenements", []))}
        if evts:
            assert {"REGULARISATION_A_EFFECTUER", "REGULARISATION_CONFIRMEE"} <= evts
        assert (await circuit.archiver(db, s, rid, {"revision": d["revision"]}))["archivee_le"]
    finally:
        await _purger(db, rid)()


async def test_decision_defavorable_sans_regularisation(db):
    s = await _sr(db)
    unite = await _unite(db)
    d = await _dossier_a_regulariser(db, s, unite)
    d = await circuit.cloturer(db, s, uuid.UUID(d["id"]), {"revision": d["revision"], "resultat": "DEFAVORABLE"})
    assert d["regularisation"] is None and not d["actions"]["confirmer_regularisation"]


async def test_piece_par_operation(db):
    s = await _sr(db)
    unite = await _unite(db)
    d = await dossiers.creer(db, s, await _saisie(db, unite.id, operations=[
        {"reference": "OP-77", "montant": Decimal("10"), "devise": "MRU"}]))
    rid = uuid.UUID(d["id"])
    try:
        await _refus(documents.deposer(db, s, rid, UploadFile(file=BytesIO(PDF), filename="a.pdf"), "JUSTIFICATIF",
                                       None, operation_reference="OP-INCONNUE"), "OPERATION_INCONNUE")
        await documents.deposer(db, s, rid, UploadFile(file=BytesIO(PDF), filename="a.pdf"), "JUSTIFICATIF",
                                None, operation_reference="op-77")
        items = (await documents.lister(db, s, rid))["items"]
        assert [x["operation_reference"] for x in items if x["categorie"] == "JUSTIFICATIF"] == ["op-77"]
    finally:
        await _purger(db, rid)()


def _classeur(lignes: list[list]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Opérations"
    ws.append(["reference_operation", "date_operation", "nature_mouvement", "devise", "montant_initial",
               "montant_a_extourner", "commentaire"])
    for ligne in lignes:
        ws.append(ligne)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _upload(contenu: bytes) -> UploadFile:
    return UploadFile(file=BytesIO(contenu), filename="operations.xlsx")


async def test_import_reimport_taille_et_rapport(db, monkeypatch):
    s = await _sr(db)
    unite = await _unite(db)
    d = await dossiers.creer(db, s, await _saisie(db, unite.id))
    rid = uuid.UUID(d["id"])
    try:
        mauvais = _classeur([["0001", "01/10/2026", "D", "MRU", 100, 150, None],
                             ["0002", None, "X", "XOF", 20, None, None]])
        res = await imports.analyser(db, s, rid, _upload(mauvais), "AJOUT")
        assert res["nb_erreurs"] == 2 and [x["statut"] for x in res["lignes"]] == ["ERREUR", "ERREUR"]
        assert res["lignes"][0]["reference"] == "0001"
        rapport = load_workbook(BytesIO(await imports.rapport_erreurs(db, s, rid, uuid.UUID(res["id"]))))
        cellules = " ".join(str(c.value) for row in rapport.active.iter_rows() for c in row if c.value)
        assert "devise" in cellules.lower() and "nature" in cellules.lower()
        await imports.abandonner(db, s, rid, uuid.UUID(res["id"]))

        bon = _classeur([["0001", "01/10/2026", "D", "MRU", 100, 40, None],
                         ["0002", "02/10/2026", "C", "USD", 20, None, "Frais"]])
        res = await imports.analyser(db, s, rid, _upload(bon), "AJOUT")
        assert res["nb_erreurs"] == 0
        await imports.confirmer(db, s, rid, uuid.UUID(res["id"]), {"revision": res["revision"]})
        fiche = await dossiers.detail(db, s, await dossiers.charger(db, s, rid))
        assert {x["devise"] for x in fiche["par_devise"]} == {"MRU", "USD"}
        assert {o["reference"] for o in fiche["operations"]} == {"0001", "0002"}
        # Réimport accidentel du même fichier : refusé sauf choix explicite.
        await _refus(imports.analyser(db, s, rid, _upload(bon), "AJOUT"), "FICHIER_DEJA_IMPORTE")
        res = await imports.analyser(db, s, rid, _upload(bon), "AJOUT", forcer=True)
        assert res["fichier_deja_importe"] and res["nb_erreurs"] == 2      # doublons exacts signalés
        await imports.abandonner(db, s, rid, uuid.UUID(res["id"]))

        monkeypatch.setattr(imports, "TAILLE_MAX_IMPORT", 100)
        with pytest.raises((AppError, HTTPException)) as e:
            await imports.analyser(db, s, rid, _upload(bon), "AJOUT", forcer=True)
        assert getattr(e.value, "status_code", None) == 413
    finally:
        await _purger(db, rid)()


async def test_modele_import_par_motif(db):
    s = await _sr(db)
    _t, motif = await _motif_regle(db, s, EXTOURNE_DEBIT)
    ex = await regles.exigences(db, motif)
    requis = {code for _c, code, _l, _r, req in imports.colonnes_du_modele(ex) if req}
    assert {"reference_operation", "date_operation", "compte_reference", "devise", "montant_initial",
            "montant_a_extourner"} <= requis
    wb = load_workbook(BytesIO(imports.modele(ex)))
    assert [c.value for c in wb["Opérations"][1]][:10] == [
        "client_reference", "compte_reference", "reference_operation", "date_operation", "nature_mouvement",
        "devise", "montant_initial", "montant_a_extourner", "motif_ligne", "commentaire"]


async def test_filtres_et_tableau_de_bord_par_devise(db):
    s = await _sr(db)
    unite = await _unite(db)
    ref = f"FLT-{uuid.uuid4().hex[:6]}"
    d = await dossiers.creer(db, s, await _saisie(db, unite.id, operations=[
        {"reference": ref, "montant": Decimal("10"), "devise": "MRU"},
        {"reference": f"{ref}-B", "montant": Decimal("5"), "devise": "USD"}]))
    trouves = await dossiers.lister(db, s, {"operation_reference": ref})
    assert [x["id"] for x in trouves["items"]] == [d["id"]]
    ligne = trouves["items"][0]
    assert ligne["devises_multiples"] and ligne["montant_total"] is None
    assert d["id"] in {x["id"] for x in (await dossiers.lister(db, s, {"devise": "USD"}))["items"]}
    tb = await reporting.tableau_de_bord(db, s, {"operation_reference": ref})
    par = {x["devise"]: Decimal(x["montant"]) for x in tb["montants_par_devise"]}
    assert par == {"MRU": Decimal("10.00"), "USD": Decimal("5.00")}


async def test_parametrage_devises_et_validation_regle(db):
    s = await _sr(db)
    await _refus(regles.modifier_devises(db, s, ["USD"]), "DEVISE_REFERENCE")
    await _refus(regles.modifier_devises(db, s, ["MRU", "DOLLAR"]), "DEVISE_INVALIDE")
    assert await regles.modifier_devises(db, s, ["mru", "usd", "EUR", "USD"]) == ["MRU", "USD", "EUR"]
    assert await regles.devises_autorisees(db) == frozenset({"MRU", "USD", "EUR"})
    _t, motif = await _motif_regle(db, s, EXTOURNE_DEBIT)
    assert (await regles.valider(db, s, motif, "VALIDEE"))["statut_validation"] == "VALIDEE"
    # Toute modification repasse la règle en « proposée ».
    r = await regles.enregistrer(db, s, motif, {**EXTOURNE_DEBIT, "justificatif_requis": True})
    assert r["statut_validation"] == "PROPOSEE"
    lecteur = ReclamationsScope(user=await _user(db), permissions=frozenset({"reclamations.view"}))
    with pytest.raises(HTTPException) as e:
        await regles.valider(db, lecteur, motif, "VALIDEE")
    assert e.value.status_code == 403

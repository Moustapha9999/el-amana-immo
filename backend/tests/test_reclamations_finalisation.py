# ruff: noqa: F811
"""Réclamations — finalisation : délais centralisés, indicateurs, comparaison, référentiels, permissions.

Recette 1–13 du prompt de finalisation (transaction annulée, aucune donnée laissée en base) :
1. cycle de vie d'une règle ; 2. règle en attente jamais réglementaire, conflits, sélection ;
3. quatre délais (événements présents / absents / incohérents) ; 4. échéance proche, retard, alertes
sans doublon ; 5. réouverture (simple et multiple) ; 6. performance des services ; 7. filtre agence ;
8. KPI cliquables = listes ; 9. comparaison A/B ; 10. KPI = exports ; 11. opérations multiples sans
double comptage ; 12. permissions retirées ; 13. vérification « Marketing et communication ».
"""

from __future__ import annotations

import pathlib
import uuid
from datetime import UTC, date, datetime, timedelta
from decimal import Decimal
from io import BytesIO

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy import func, select, text

from app.core.exceptions import AppError
from app.data.plateforme_catalogue import FUNCTIONAL_PERMISSIONS
from app.models import Agence, Notification, Reclamation, ReclamationEvenement, ReclamationRegleDelai, ReclamationUnite
from app.services.reclamations import circuit, delais, dossiers, indicateurs, organisation, reporting
from app.services.reclamations import echeances as E
from app.services.reclamations.acces import ReclamationsScope, permissions_effectives
from tests.test_reclamations_circuit import _agent_service, _ops, _refus
from tests.test_reclamations_dossiers import _agent, _classement, _saisie, _sr
from tests.test_reclamations_fondations import _unite, _user, db  # noqa: F401

NOW = datetime.now(UTC)
AUJOURDHUI = NOW.date()
PERIODE = {"kpi_debut": AUJOURDHUI - timedelta(days=30), "kpi_fin": AUJOURDHUI}
RETIREES = {"reclamations.suspend", "reclamations.response.prepare", "reclamations.response.validate",
            "reclamations.audit.view"}


def _marqueur() -> str:
    return f"FIN-{uuid.uuid4().hex[:8]}"


async def _dossier(db, sr, unite, marqueur: str, **extra) -> Reclamation:
    d = await dossiers.creer(db, sr, await _saisie(db, unite.id, objet=f"Dossier {marqueur}", **extra))
    return await dossiers.charger(db, sr, uuid.UUID(d["id"]))


async def _regle(db, sr, type_id, *, valider: bool = True, **extra) -> dict:
    data = {"libelle": f"Règle {uuid.uuid4().hex[:5]}", "type_id": type_id, "duree": 10,
            "date_effet": date(2020, 1, 1), **extra}
    g = await delais.creer_regle(db, sr, data)
    if valider:
        g = await delais.valider_regle(db, sr, uuid.UUID(g["id"]), {"statut_validation": "VALIDEE"})
    return g


async def _total(db, s, **f) -> int:
    return (await dossiers.lister(db, s, f, taille=100))["total"]


# ---------------------------------------------------------------- 1. cycle de vie
async def test_regle_creer_modifier_activer_valider_versionner(db):
    sr = await _sr(db)
    unite = await _unite(db)
    t, _m, _ = await _classement(db, sr)
    tid = uuid.UUID(t["id"])
    g = await _regle(db, sr, tid, valider=False, seuil_alerte=2, destinataires=["SR", "GESTIONNAIRE"])
    rid = uuid.UUID(g["id"])
    assert (g["statut_validation"], g["etat_effectif"], g["reglementaire"]) == (
        "A_VALIDER", "EN_ATTENTE_VALIDATION", False)
    assert g["destinataires"] == ["GESTIONNAIRE", "SR"] and g["point_arrivee"] == "CLOTURE"

    g = await delais.modifier_regle(db, sr, rid, {"duree": 7})
    assert g["duree"] == 7 and g["statut_validation"] == "A_VALIDER"
    assert (await delais.modifier_regle(db, sr, rid, {"actif": False}))["etat_effectif"] == "INACTIVE"
    assert (await delais.modifier_regle(db, sr, rid, {"actif": True}))["etat_effectif"] == "EN_ATTENTE_VALIDATION"
    await _refus(delais.modifier_regle(db, sr, rid, {"seuil_alerte": 7}), "SEUIL_INVALIDE")

    g = await delais.valider_regle(db, sr, rid, {"statut_validation": "VALIDEE"})
    assert (g["etat_effectif"], g["reglementaire"]) == ("APPLIQUEE", True)
    # Paramètre de calcul modifié sur une règle validée non utilisée : retour en validation.
    g = await delais.modifier_regle(db, sr, rid, {"duree": 10})
    assert g["statut_validation"] == "A_VALIDER"
    g = await delais.valider_regle(db, sr, rid, {"statut_validation": "VALIDEE"})

    r = await _dossier(db, sr, unite, _marqueur(), type_id=tid, date_premiere_reception=NOW - timedelta(days=3))
    echeance_initiale = r.echeance
    assert r.etat_delai == "CALCULE" and r.regle_delai_id == rid
    # Règle appliquée : paramètres de calcul figés, libellé modifiable.
    await _refus(delais.modifier_regle(db, sr, rid, {"duree": 20}), "REGLE_UTILISEE")
    assert (await delais.modifier_regle(db, sr, rid, {"libelle": "Délai type test"}))["libelle"] == "Délai type test"

    # Nouvelle version à date d'effet : l'ancienne est bornée à la validation, sans recalcul du passé.
    v = await delais.versionner_regle(db, sr, rid, {"date_effet": AUJOURDHUI, "duree": 5})
    assert (v["remplace_id"], v["statut_validation"], v["duree"], v["type_id"]) == (str(rid), "A_VALIDER", 5, str(tid))
    assert v["seuil_alerte"] == 2
    await _refus(delais.versionner_regle(db, sr, rid, {"date_effet": AUJOURDHUI, "duree": 4}), "VERSION_EXISTANTE")
    await _refus(delais.versionner_regle(db, sr, uuid.UUID(v["id"]), {"date_effet": date(2019, 1, 1)}),
                 "DATE_EFFET_INVALIDE")
    await delais.valider_regle(db, sr, uuid.UUID(v["id"]), {"statut_validation": "VALIDEE"})
    ancienne = await db.get(ReclamationRegleDelai, rid)
    assert ancienne.date_fin == AUJOURDHUI - timedelta(days=1)
    await db.refresh(r)
    assert r.echeance == echeance_initiale and r.regle_delai_id == rid

    nouveau = await _dossier(db, sr, unite, _marqueur(), type_id=tid, date_premiere_reception=NOW - timedelta(minutes=1))
    assert nouveau.regle_delai_id == uuid.UUID(v["id"])

    # Version suivante : seuil d'alerte vidé ; la fin n'est pas reprise de la version remplacée.
    v2 = await delais.versionner_regle(db, sr, uuid.UUID(v["id"]), {
        "date_effet": AUJOURDHUI + timedelta(days=30), "detacher_seuil": True})
    assert (v2["seuil_alerte"], v2["date_fin"], v2["duree"]) == (None, None, 5)

    histo = await delais.historique_regle(db, rid)
    actions = {e["action"] for e in histo["evenements"]}
    assert {"reclamations.delai.create", "reclamations.delai.update", "reclamations.delai.validate",
            "reclamations.delai.version"} <= actions
    assert [x["id"] for x in histo["versions"]] == [v2["id"], v["id"], str(rid)]
    regles = {x["id"]: x for x in await delais.lister_regles(db)}
    assert regles[str(rid)]["remplacee_par"] == v["id"] and regles[str(rid)]["nb_dossiers"] == 1


# ------------------------------------------- 2. règle en attente, conflits, sélection
def _regle_memoire(**champs) -> ReclamationRegleDelai:
    base = {"id": uuid.uuid4(), "libelle": "R", "type_id": None, "motif_id": None, "sous_motif_id": None,
            "priorite": None, "duree": 10, "unite_duree": "JOURS_CALENDAIRES", "point_depart": "PREMIERE_RECEPTION",
            "point_arrivee": "CLOTURE", "seuil_alerte": None, "date_effet": date(2020, 1, 1), "date_fin": None,
            "statut_validation": "VALIDEE", "actif": True}
    return ReclamationRegleDelai(**{**base, **champs})


def test_selection_specificite_attente_conflit_depart():
    t, m, sm = uuid.uuid4(), uuid.uuid4(), uuid.uuid4()
    r = Reclamation(type_id=t, motif_id=m, sous_motif_id=sm, priorite="NORMALE",
                    date_premiere_reception=datetime(2026, 10, 1, 9, tzinfo=UTC), statut="A_QUALIFIER")
    generale = _regle_memoire(duree=30)
    par_motif = _regle_memoire(type_id=t, motif_id=m, duree=15)
    par_sous_motif = _regle_memoire(type_id=t, motif_id=m, sous_motif_id=sm, duree=5)
    autre_motif = _regle_memoire(motif_id=uuid.uuid4(), duree=1)
    assert E.selectionner([generale, par_motif, autre_motif], r).regle is par_motif
    assert E.selectionner([generale, par_motif, par_sous_motif], r).regle is par_sous_motif
    # Une règle plus spécifique mais non validée n'est jamais appliquée ni présentée comme telle.
    attente = _regle_memoire(type_id=t, motif_id=m, sous_motif_id=sm, duree=2, statut_validation="A_VALIDER")
    sel = E.selectionner([generale, attente], r)
    assert sel.regle is generale and sel.en_attente == [attente]
    assert E.selectionner([attente], r).etat == "EN_ATTENTE_VALIDATION"
    assert E.selectionner([], r).etat == "NON_CONFIGURE"
    hors_periode = _regle_memoire(date_effet=date(2026, 11, 1))
    assert E.selectionner([hors_periode], r).etat == "NON_CONFIGURE"
    # Deux règles validées de même spécificité, paramètres différents : conflit, aucune échéance.
    conflit = E.selectionner([par_motif, _regle_memoire(type_id=t, motif_id=m, duree=20)], r)
    assert conflit.etat == "CONFLIT" and conflit.regle is None and len(conflit.conflit) == 2
    E.poser(r, conflit, set())
    assert r.echeance is None and r.etat_delai == "CONFLIT"
    assert E.situation(r, NOW) == "NON_CALCULABLE"
    # Départ « réception SR » non survenu : pas d'échéance inventée.
    sr = _regle_memoire(point_depart="RECEPTION_SR")
    E.poser(r, E.selectionner([sr], r), set())
    assert r.etat_delai == "SANS_DEPART" and r.echeance is None


def test_calendrier_jours_ouvres_heures_et_seuil():
    vendredi = datetime(2026, 10, 9, 10, tzinfo=UTC)
    ferie = {date(2026, 10, 13)}
    assert E.ajouter_jours(vendredi, 2, "JOURS_OUVRES", ferie).date() == date(2026, 10, 14)
    assert E.ajouter_jours(vendredi, 2, "JOURS_CALENDAIRES", ferie) == vendredi + timedelta(days=2)
    assert E.ajouter_jours(vendredi, 36, "HEURES", ferie) == vendredi + timedelta(hours=36)
    echeance = datetime(2026, 10, 14, 23, 59, 59, tzinfo=UTC)
    assert E.retrancher(echeance, 2, "JOURS_OUVRES", ferie).date() == date(2026, 10, 9)
    assert E.retrancher(echeance, 5, "HEURES", ferie) == echeance - timedelta(hours=5)


async def test_regle_en_attente_et_conflit_en_base(db):
    sr = await _sr(db)
    unite = await _unite(db)
    t, _m, _ = await _classement(db, sr)
    tid = uuid.UUID(t["id"])
    attente = await _regle(db, sr, tid, valider=False)
    r = await _dossier(db, sr, unite, _marqueur(), type_id=tid)
    ligne = dossiers.ligne(r, await dossiers.Contexte.charger(db))
    assert ligne["echeance"] is None and ligne["etat_delai"] == "EN_ATTENTE_VALIDATION"
    assert ligne["etat_delai_libelle"] == "Délai en attente de validation"
    assert ligne["situation_delai"] == "NON_CALCULABLE"
    await delais.valider_regle(db, sr, uuid.UUID(attente["id"]), {"statut_validation": "VALIDEE"})
    concurrente = await _regle(db, sr, tid, valider=False, duree=20)
    await _refus(delais.valider_regle(db, sr, uuid.UUID(concurrente["id"]), {"statut_validation": "VALIDEE"}),
                 "REGLE_CONFLIT")
    # Application explicite aux dossiers sans délai : tracée dans l'historique, rien d'autre n'est recalculé.
    res = await delais.appliquer_sans_delai(db, sr)
    assert res["calcules"] >= 1
    await db.refresh(r)
    assert r.etat_delai == "CALCULE" and r.echeance is not None
    assert "DELAI_CALCULE" in [e["type"] for e in await dossiers.chronologie(db, sr, r.id)]


# ------------------------------------------------------------------ 3. quatre délais
async def test_quatre_delais_evenements_presents_absents_incoherents(db):
    sr = await _sr(db)
    unite = await _unite(db)
    m = _marqueur()
    d1 = await _dossier(db, sr, unite, m, date_premiere_reception=NOW - timedelta(days=5))
    d1.date_reception_sr = d1.date_premiere_reception + timedelta(hours=4)
    d1.date_retour_service = NOW - timedelta(days=3)
    d1.date_reponse_client = NOW - timedelta(days=2)
    d1.date_cloture = d1.date_premiere_reception + timedelta(days=4)
    d1.statut, d1.etape = "CLOTUREE", "CLOTURE"
    d2 = await _dossier(db, sr, unite, m, date_premiere_reception=NOW - timedelta(days=4))
    d2.date_reception_sr = None
    d3 = await _dossier(db, sr, unite, m, date_premiere_reception=NOW - timedelta(days=6))
    d3.date_reception_sr = d3.date_premiere_reception + timedelta(hours=2)
    d3.date_retour_service = NOW - timedelta(days=1)
    d3.date_reponse_client = NOW - timedelta(days=2)
    d3.statut, d3.etape = "A_CLOTURER", "CLOTURE"
    await db.flush()

    p = indicateurs.Periode.de(PERIODE["kpi_debut"], PERIODE["kpi_fin"])
    res = await indicateurs.pilotage(db, sr, {"q": m}, p, lambda _i: None)
    k = {x["code"]: x for x in res["delais"]}
    assert k["prise_en_charge"]["termines"] == {"n": 2, "moyenne_heures": 3.0, "moyenne_jours": 0.12}
    assert k["prise_en_charge"]["en_cours"]["n"] == 1
    assert k["reponse_client"]["termines"]["n"] == 1 and k["reponse_client"]["termines"]["moyenne_heures"] == 24.0
    assert k["reponse_client"]["anomalies"] == {"n": 1, "references": [d3.reference]}
    assert k["cloture"]["termines"]["n"] == 1 and k["cloture"]["termines"]["moyenne_heures"] == 96.0
    assert k["cloture"]["en_cours"]["n"] == 2
    assert k["traitement_service"]["termines"] == {"n": 0, "moyenne_heures": None, "moyenne_jours": None}
    # Chaque compte = total de la liste filtrée par l'indicateur.
    for code, attendu in (("delai_prise_en_charge_termines", 2), ("delai_prise_en_charge_en_cours", 1),
                          ("delai_reponse_client_anomalies", 1), ("delai_cloture_en_cours", 2),
                          ("delai_cloture_termines", 1)):
        assert await _total(db, sr, q=m, indicateur=code, **PERIODE) == attendu, code
    vide = await indicateurs.pilotage(db, sr, {"q": m}, indicateurs.Periode.de(date(2001, 1, 1), date(2001, 1, 31)),
                                      lambda _i: None)
    assert all(d["termines"]["moyenne_heures"] is None for d in vide["delais"])
    assert vide["reouverture"]["taux"] is None


# -------------------------------------------- 4. échéance proche, retard, alertes
async def test_echeance_proche_retard_et_alertes_sans_doublon(db):
    sr = await _sr(db)
    unite = await _unite(db)
    t, _m, _ = await _classement(db, sr)
    tid = uuid.UUID(t["id"])
    await _regle(db, sr, tid, seuil_alerte=3, destinataires=["GESTIONNAIRE"])
    m = _marqueur()
    proche = await _dossier(db, sr, unite, m, type_id=tid, date_premiere_reception=NOW - timedelta(days=8))
    retard = await _dossier(db, sr, unite, m, type_id=tid, date_premiere_reception=NOW - timedelta(days=12))
    ok = await _dossier(db, sr, unite, m, type_id=tid, date_premiere_reception=NOW - timedelta(days=1))
    for r in (proche, retard, ok):
        r.gestionnaire_id = sr.user_id
    await db.flush()
    assert [E.situation(r, NOW) for r in (proche, retard, ok)] == ["ECHEANCE_PROCHE", "EN_RETARD", "DANS_DELAI"]
    assert await _total(db, sr, q=m, situation_delai="ECHEANCE_PROCHE") == 1
    assert await _total(db, sr, q=m, en_retard=True) == 1
    al = await delais.alertes(db, sr)
    assert retard.reference in {x["reference"] for x in al["retards"]}
    assert proche.reference in {x["reference"] for x in al["alertes"]}

    async def notifs() -> int:
        return await db.scalar(select(func.count()).select_from(Notification).where(
            Notification.entity_id.in_([str(proche.id), str(retard.id), str(ok.id)]))) or 0

    await delais.envoyer_alertes(db)
    assert await notifs() == 2
    await delais.envoyer_alertes(db)
    assert await notifs() == 2
    niveaux = dict((await db.execute(select(Reclamation.id, Reclamation.niveau_alerte).where(
        Reclamation.id.in_([proche.id, retard.id, ok.id])))).all())
    assert niveaux == {proche.id: "PROCHE", retard.id: "RETARD", ok.id: None}


# ------------------------------------------------------------------- 5. réouverture
async def test_reouverture_simple_et_multiple(db):
    sr = await _sr(db)
    unite = await _unite(db)
    m = _marqueur()
    x = await _dossier(db, sr, unite, m)
    y = await _dossier(db, sr, unite, m)
    z = await _dossier(db, sr, unite, m)

    def ev(r, type_evenement, jours):
        db.add(ReclamationEvenement(reclamation_id=r.id, type_evenement=type_evenement,
                                    date_effet=NOW - timedelta(days=jours), auteur_id=sr.user_id, donnees={}))

    for type_evenement, jours in (("CLOTURE", 9), ("REOUVERTURE", 8), ("CLOTURE", 6), ("REOUVERTURE", 5)):
        ev(x, type_evenement, jours)
    ev(y, "CLOTURE", 4)
    ev(z, "CLOTURE", 90)
    ev(z, "REOUVERTURE", 3)
    await db.flush()
    p = indicateurs.Periode.de(PERIODE["kpi_debut"], PERIODE["kpi_fin"])
    ro = (await indicateurs.pilotage(db, sr, {"q": m}, p, lambda _i: None))["reouverture"]
    assert (ro["cloturees"], ro["reouvertes"], ro["taux"]) == (2, 1, 50.0)
    assert (ro["nb_reouvertures"], ro["dossiers_reouverts"], ro["dossiers_plusieurs_reouvertures"]) == (3, 2, 1)
    assert await _total(db, sr, q=m, indicateur="reouvertes", **PERIODE) == 1
    assert await _total(db, sr, q=m, indicateur="clotures", **PERIODE) == 2


# -------------------------------------- 6. services ; 11. opérations sans double comptage
async def test_performance_services_et_operations_multiples(db):
    sr = await _sr(db)
    service = await _unite(db)
    t, _m, _ = await _classement(db, sr)
    m = _marqueur()
    ids = []
    for montants in (("100", "200", "300"), ("50",)):
        d = await dossiers.creer(db, sr, await _saisie(db, service.id, objet=f"Dossier {m}", type_id=uuid.UUID(t["id"]),
                                                       operations=_ops(*montants)))
        d = await circuit.affecter(db, sr, uuid.UUID(d["id"]), {"revision": d["revision"], "unite_id": service.id})
        ids.append((uuid.UUID(d["id"]), d["revision"]))
    agent = await _agent_service(db, service.id)
    await circuit.retourner(db, agent, ids[0][0], {"revision": ids[0][1], "analyse": "Vu", "conclusion": "OK"})

    p = indicateurs.Periode.de(PERIODE["kpi_debut"], PERIODE["kpi_fin"])
    res = await indicateurs.pilotage(db, sr, {"q": m}, p, lambda _i: "Service test")
    (ligne,) = [s for s in res["services"] if s["unite_id"] == str(service.id)]
    assert (ligne["affectes"], ligne["en_cours"], ligne["retournes"], ligne["en_retard"]) == (2, 1, 1, 0)
    assert ligne["delai_moyen_heures"] is not None and ligne["age_moyen_heures"] is not None
    ts = next(d for d in res["delais"] if d["code"] == "traitement_service")
    assert ts["termines"]["n"] == 1 and ts["en_cours"]["n"] == 1 and ts["unite"] == "fiche de traitement"
    for code, attendu in (("service_affectes", 2), ("service_en_cours", 1), ("service_retournes", 1)):
        assert await _total(db, sr, q=m, indicateur=code, kpi_unite_id=service.id, **PERIODE) == attendu, code

    tb = await reporting.tableau_de_bord(db, sr, {"q": m})
    assert (tb["total"], tb["nb_operations"], tb["dossiers_avec_operations"]) == (2, 4, 2)
    assert Decimal(tb["montant_total"]) == Decimal("650")
    liste = await dossiers.lister(db, sr, {"q": m, "indicateur": "avec_operations"})
    assert liste["total"] == 2 and liste["totaux"]["nb_operations"] == 4
    assert Decimal(liste["totaux"]["montant_total"]) == Decimal("650")


# -------------------------------------------------------------- 7. filtre agence
async def test_filtre_agence_sur_tous_les_indicateurs(db):
    sr = await _sr(db)
    unite = await _unite(db)
    a1, a2 = (await db.scalars(select(Agence.id).order_by(Agence.code).limit(2))).all()
    m = _marqueur()
    await _dossier(db, sr, unite, m, client_agence_id=a1, operations=_ops("10"))
    await _dossier(db, sr, unite, m, client_agence_id=a2)
    await _dossier(db, sr, unite, m)
    for agence, attendu in ((a1, 1), (a2, 1), (dossiers.AUCUNE, 1), (None, 3)):
        f = {"q": m, "client_agence_id": agence}
        tb = await reporting.tableau_de_bord(db, sr, f)
        assert tb["total"] == attendu and await _total(db, sr, **f) == attendu, agence
        assert sum(x["n"] for x in tb["agences"]) == attendu
        en_cours = next(d for d in tb["pilotage"]["delais"] if d["code"] == "cloture")["en_cours"]["n"]
        assert en_cours == attendu
    tb = await reporting.tableau_de_bord(db, sr, {"q": m, "client_agence_id": a1})
    assert tb["nb_operations"] == 1 and tb["filtres"]["client_agence_id"] == str(a1)
    assert {x["libelle"] for x in (await reporting.tableau_de_bord(db, sr, {"q": m}))["agences"]} >= {"Sans agence"}


# ------------------------------------------------------- 8. KPI cliquables = listes
async def test_kpi_cliquables_correspondent_aux_listes(db):
    sr = await _sr(db)
    unite = await _unite(db)
    m = _marqueur()
    await _dossier(db, sr, unite, m, date_premiere_reception=NOW - timedelta(days=3), operations=_ops("12.5", "7.5"))
    await _dossier(db, sr, unite, m, date_premiere_reception=NOW - timedelta(days=45))
    vieux = await _dossier(db, sr, unite, m, date_premiere_reception=NOW - timedelta(days=10))
    vieux.date_cloture, vieux.statut, vieux.etape = NOW - timedelta(days=2), "CLOTUREE", "CLOTURE"
    await db.flush()
    debut = AUJOURDHUI - timedelta(days=90)
    tb = await reporting.tableau_de_bord(db, sr, {"q": m, "date_debut": debut})
    liens = tb["liens"]

    def params(cle: str) -> dict:
        out = {}
        for k, v in liens[cle].items():
            if v is None:
                continue
            out[k] = date.fromisoformat(v) if k.endswith(("debut", "fin")) else (v == "true" if v in ("true",) else v)
        return out

    assert tb["volume_30j"] == await _total(db, sr, q=m, **params("volume_30j")) == 2
    assert tb["dossiers_avec_operations"] == await _total(db, sr, q=m, **params("operations")) == 1
    liste_ops = await dossiers.lister(db, sr, {"q": m, **params("operations")})
    assert Decimal(liste_ops["totaux"]["montant_total"]) == Decimal(tb["montant_total"]) == Decimal("20")
    assert liste_ops["totaux"]["nb_operations"] == tb["nb_operations"] == 2
    assert tb["performance"]["delai_cloture_n"] == await _total(db, sr, q=m, **params("delai_moyen")) == 1
    assert tb["performance"]["delai_cloture_jours"] == 8.0
    assert tb["ouvertes"] == await _total(db, sr, q=m, **params("ouvertes")) == 2
    assert tb["retards"] == await _total(db, sr, q=m, **params("retards"))
    assert tb["total"] == await _total(db, sr, q=m, **params("total")) == 3


# ---------------------------------------------------------- 9. comparaison A/B
def test_comparer_variation_et_sens():
    a = {"recues": 0, "clotures": 10, "taux_reouverture": 20.0, "delai_cloture": 48.0, "montant": "100"}
    b = {"recues": 5, "clotures": 15, "taux_reouverture": 10.0, "delai_cloture": 60.0, "montant": "150"}
    lignes = {x["code"]: x for x in indicateurs.comparer(a, b)}
    assert lignes["recues"]["ecart"] == 5 and lignes["recues"]["variation"] is None
    assert lignes["recues"]["favorable"] is None
    assert (lignes["clotures"]["variation"], lignes["clotures"]["favorable"]) == (50.0, True)
    assert lignes["taux_reouverture"]["favorable"] is True
    assert (lignes["delai_cloture"]["ecart"], lignes["delai_cloture"]["favorable"]) == (12.0, False)
    assert lignes["montant"]["variation"] == 50.0 and lignes["montant"]["favorable"] is None
    assert lignes["taux_respect"]["ecart"] is None and lignes["taux_respect"]["variation"] is None


async def test_comparaison_deux_periodes_et_exports(db):
    sr = await _sr(db)
    unite = await _unite(db)
    m = _marqueur()
    await _dossier(db, sr, unite, m, date_premiere_reception=NOW - timedelta(days=40), operations=_ops("10"))
    await _dossier(db, sr, unite, m, date_premiere_reception=NOW - timedelta(days=5), operations=_ops("30", "5"))
    await _dossier(db, sr, unite, m, date_premiere_reception=NOW - timedelta(days=3))
    pa = indicateurs.Periode.de(AUJOURDHUI - timedelta(days=60), AUJOURDHUI - timedelta(days=31))
    pb = indicateurs.Periode.de(AUJOURDHUI - timedelta(days=30), AUJOURDHUI)
    comp = await indicateurs.comparaison(db, sr, {"q": m}, pa, pb)
    lignes = {x["code"]: x for x in comp["lignes"]}
    assert (lignes["recues"]["a"], lignes["recues"]["b"], lignes["recues"]["variation"]) == (1, 2, 100.0)
    assert (lignes["nb_operations"]["a"], lignes["nb_operations"]["b"]) == (1, 2)
    assert Decimal(lignes["montant"]["b"]) == Decimal("35")
    assert comp["avertissements"] and "durées différentes" in comp["avertissements"][0]
    assert await _total(db, sr, q=m, date_debut=pb.jour_debut, date_fin=pb.jour_fin) == lignes["recues"]["b"]
    with pytest.raises(AppError):
        indicateurs.Periode.de(AUJOURDHUI, AUJOURDHUI - timedelta(days=1))
    with pytest.raises(AppError):
        await indicateurs.comparaison(db, sr, {}, indicateurs.Periode.de(None, AUJOURDHUI), pb)

    classeur = load_workbook(BytesIO(await reporting.exporter_comparaison(db, sr, {"q": m}, pa, pb)), read_only=True)
    assert classeur.sheetnames == ["Comparaison", "Méthode"]
    rangees = list(classeur["Comparaison"].iter_rows(values_only=True))
    assert any(r[0] == "Réclamations reçues" and r[1] == "1" and r[2] == "2" for r in rangees)
    assert (await reporting.exporter_comparaison(db, sr, {"q": m}, pa, pb, format="pdf")).startswith(b"%PDF")


# ------------------------------------------------------------- 10. KPI = exports
async def test_exports_coherents_avec_les_indicateurs(db):
    sr = await _sr(db)
    unite = await _unite(db)
    m = _marqueur()
    await _dossier(db, sr, unite, m, operations=_ops("40", "2"))
    await _dossier(db, sr, unite, m)
    tb = await reporting.tableau_de_bord(db, sr, {"q": m})
    classeur = load_workbook(BytesIO(await reporting.exporter_pilotage(db, sr, {"q": m})), read_only=True)
    assert classeur.sheetnames == ["Synthèse", "Délais", "Services", "Réouverture", "Agences"]
    synthese = {r[1]: r[2] for r in classeur["Synthèse"].iter_rows(values_only=True) if r and r[1]}
    assert synthese["Réclamations (période, date de réception)"] == tb["total"] == 2
    assert synthese["Opérations contestées"] == tb["nb_operations"] == 2
    assert Decimal(str(synthese[f"Montant contesté ({'MRU'})"])) == Decimal("42")
    assert (await reporting.exporter_pilotage(db, sr, {"q": m}, format="pdf")).startswith(b"%PDF")
    registre = load_workbook(BytesIO(await reporting.exporter(db, sr, {"q": m, "indicateur": "avec_operations"})),
                             read_only=True)
    lignes = [r for r in registre["Réclamations"].iter_rows(values_only=True) if r and str(r[0]).startswith("REC-")]
    assert len(lignes) == await _total(db, sr, q=m, indicateur="avec_operations") == 1


# ---------------------------------------------- 12. permissions ; référentiels
async def test_permissions_obsoletes_retirees(db):
    assert not RETIREES & {c for c, _l, _m in FUNCTIONAL_PERMISSIONS}
    restantes = (await db.scalars(text("SELECT code FROM permissions WHERE code = ANY(:c)"),
                                  {"c": list(RETIREES)})).all()
    assert restantes == []
    # Les étapes du circuit gardent leurs permissions effectives (réponse, clôture, réouverture).
    assert {"reclamations.response.send", "reclamations.close", "reclamations.reopen"} <= permissions_effectives(
        {"reclamations.admin"})
    # La traçabilité technique reste : journal d'audit et historique des dossiers non touchés.
    assert await db.scalar(text("SELECT to_regclass('public.reclamation_evenements') IS NOT NULL"))


async def test_referentiels_proposes_puis_valides(db):
    sr = await _sr(db)
    t, m, sm = await _classement(db, sr)
    assert (t["statut_validation"], t["a_valider"]) == ("PROPOSEE", True)
    await _refus(organisation.valider_referentiel(db, sr, uuid.UUID(sm["id"]), {"statut_validation": "VALIDEE"}),
                 "PARENT_NON_VALIDE")
    for x in (t, m, sm):
        v = await organisation.valider_referentiel(db, sr, uuid.UUID(x["id"]), {"statut_validation": "VALIDEE"})
        assert v["statut_validation"] == "VALIDEE" and v["valide_le"]
    await _refus(organisation.valider_referentiel(db, sr, uuid.UUID(m["id"]), {"statut_validation": "PROPOSEE"}),
                 "ENFANTS_VALIDES")
    # Libellé modifié : la valeur redevient une proposition, l'historique des dossiers est conservé.
    v = await organisation.modifier_referentiel(db, sr, uuid.UUID(sm["id"]), {"libelle": "Sous-motif renommé"})
    assert v["statut_validation"] == "PROPOSEE"
    lecteur = await _agent(db, perms=("reclamations.view", "reclamations.references.manage"))
    with pytest.raises(HTTPException) as e:
        await organisation.valider_referentiel(db, lecteur, uuid.UUID(t["id"]), {"statut_validation": "PROPOSEE"})
    assert e.value.status_code == 403
    canal = next(r for r in await organisation.lister_referentiels(db, "CANAL"))
    assert canal["a_valider"] is False and canal["statut_validation"] is None


# --------------------------------------- 13. Marketing ; enregistrement d'appel
async def test_marketing_communication_conforme_a_la_source(db):
    """Vérification sans déplacement : rattachement identique à la liste « entité » du module Formation."""
    source = pathlib.Path(__file__).parents[1] / "alembic/versions/20261007_formation_module.py"
    assert '("MARKETING ET COMMUNICATION", "MARKETING & COMMUNICATION", NKC)' in source.read_text(encoding="utf-8")
    service = await db.scalar(select(ReclamationUnite).where(ReclamationUnite.libelle == "MARKETING ET COMMUNICATION"))
    if service is None:
        pytest.skip("Référentiel des unités non initialisé")
    parent = await db.get(ReclamationUnite, service.parent_id)
    assert (service.type_unite, parent.type_unite, parent.libelle) == (
        "SERVICE", "DEPARTEMENT", "MARKETING & COMMUNICATION")


async def test_enregistrement_appel_desactive_par_defaut(db):
    assert await dossiers.enregistrement_autorise(db) is False


# --------------------------------------------------------------- périmètre
async def test_indicateurs_restent_dans_le_perimetre(db):
    sr = await _sr(db)
    unite, autre = await _unite(db), await _unite(db)
    m = _marqueur()
    await _dossier(db, sr, unite, m)
    await _dossier(db, sr, autre, m)
    agent = ReclamationsScope(user=await _user(db), permissions=permissions_effectives(
        {"reclamations.view", "reclamations.reporting.view"}), unites=frozenset({unite.id}))
    tb = await reporting.tableau_de_bord(db, agent, {"q": m})
    assert tb["total"] == 1 == await _total(db, agent, q=m)
    assert next(d for d in tb["pilotage"]["delais"] if d["code"] == "cloture")["en_cours"]["n"] == 1

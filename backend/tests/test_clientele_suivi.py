"""Suivi de la classification PP / PM : pourcentages, séparation PP / PM, exports, lots par sélection."""

from __future__ import annotations

from datetime import date
from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import select, text
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine
from sqlalchemy.pool import NullPool

from app.core.config import get_settings
from app.core.exceptions import AppError
from app.models import ClienteleClassification, User
from app.services.clientele.consolidation import consolider, lire_ligne
from app.services.clientele.lots import ClienteleLotService
from app.services.clientele.persistance import enregistrer
from app.services.clientele.service import Ctx
from app.services.clientele.situation import ClienteleSituationService
from app.services.clientele.suivi import ClienteleSuiviService, pourcentage, pourcentage_fr
from tests.test_clientele_consolidation import brute

EXTRACTION = date(2026, 10, 8)
RACINES = ("000001", "000002", "000003", "000004")
PERMS = {"clientele.view", "clientele.classif.view", "clientele.classif.execute", "clientele.export"}


def test_pourcentage_borne_et_population_vide():
    assert pourcentage(0, 0) is None
    assert pourcentage_fr(None) == "Non applicable"
    assert pourcentage(1, 3) == 33.3
    assert pourcentage(5, 5) == 100.0
    assert pourcentage(7, 5) == 100.0
    assert pourcentage(0, 5) == 0.0
    assert pourcentage_fr(82.5) == "82,5 %"


@pytest.fixture
async def db():
    engine = create_async_engine(get_settings().database_url, poolclass=NullPool)
    try:
        async with engine.connect() as conn:
            pret = await conn.scalar(text("SELECT to_regclass('public.clientele_situation') IS NOT NULL"))
    except Exception:
        pret = False
    if not pret:
        await engine.dispose()
        pytest.skip("Vue clientele_situation absente")
    session = async_sessionmaker(engine, expire_on_commit=False)()
    liste = ",".join(f"'{r}'" for r in RACINES)
    for table in ("clientele_classif_historique", "clientele_classifications", "clientele_classif_evaluations",
                  "clientele_comptes", "clientele_clients"):
        await session.execute(text(f"DELETE FROM {table} WHERE racine_client IN ({liste})"))
    try:
        yield session
    finally:
        await session.rollback()
        await session.close()
        await engine.dispose()


async def _ctx(db) -> Ctx:
    user = await db.scalar(select(User).limit(1))
    if not user:
        pytest.skip("aucun utilisateur")
    return Ctx(user=user, permissions=set(PERMS))


async def _seed(db) -> None:
    pm = {"AGENT_ECONOMIQUE": "AUTRES SOCIETES", "CATEGORIE_JURIDIQUE": "SOCIETE COMMERCIALE",
          "TYPE_IDENTIFIANT": "NIF"}
    await enregistrer(db, consolider(lire_ligne(b, i) for i, b in enumerate((
        brute("000001", "00000100001"),
        brute("000002", "00000200001"),
        brute("000003", "00000300001", IDENTIFIANT_LISTE="00012345", **pm),
        brute("000004", "00000400001", **pm),
    ), 2)), EXTRACTION)
    await db.flush()


async def test_progression_pp_pm_separees_et_coherentes(db):
    ctx = await _ctx(db)
    svc = ClienteleSuiviService(db, ctx)
    avant = await svc.progression()
    await _seed(db)
    db.add(ClienteleClassification(racine_client="000001", niveau="MOYEN", source="MANUEL", motifs=[]))
    db.add(ClienteleClassification(racine_client="000003", niveau="FAIBLE", source="MOTEUR", motifs=[]))
    await db.flush()
    apres = await svc.progression()

    for cle in ("pp", "pm"):
        assert apres[cle]["total"] - avant[cle]["total"] == 2
        assert apres[cle]["classes"] - avant[cle]["classes"] == 1
        assert apres[cle]["non_classes"] - avant[cle]["non_classes"] == 1
        b = apres[cle]
        assert b["classes"] + b["non_classes"] == b["total"]
        assert b["pourcentage"] is None or 0 <= b["pourcentage"] <= 100
    g = apres["global"]
    assert g["total"] == apres["pp"]["total"] + apres["pm"]["total"]
    assert g["classes"] == apres["pp"]["classes"] + apres["pm"]["classes"]
    assert apres["pp"]["manuels"] - avant["pp"]["manuels"] == 1


async def test_listes_et_exports_jamais_melanges(db):
    ctx = await _ctx(db)
    await _seed(db)
    db.add(ClienteleClassification(racine_client="000001", niveau="ELEVE", source="MANUEL", motifs=[]))
    await db.flush()
    svc = ClienteleSuiviService(db, ctx)

    pp = await svc.lister(profil="PP", q="00000")
    assert {x["racine_client"] for x in pp["items"]} >= {"000001", "000002"}
    assert all(x["profil"] == "PP" for x in pp["items"])
    classes = await svc.lister(profil="PP", racine="000001", statut="CLASSE")
    assert classes["total"] == 1 and classes["items"][0]["statut_classification"] == "CLASSE"
    assert (await svc.lister(profil="PP", racine="000002", statut="CLASSE"))["total"] == 0
    assert (await svc.lister(profil="PM", racine="000001"))["total"] == 0

    apercu = await svc.apercu_export(profil="PP", racine="000001")
    assert apercu["filtre"] == 1 and apercu["tout"] >= 2

    contenu, nom = await svc.exporter(profil="PP", format="xlsx", portee="filtre", racine="000001")
    assert nom.startswith("classification-pp-filtre-") and nom.endswith(".xlsx")
    ws = load_workbook(BytesIO(contenu)).worksheets[0]
    valeurs = [[c.value for c in row] for row in ws.iter_rows(min_row=6)]
    assert len([v for v in valeurs if any(v)]) == apercu["filtre"]
    assert valeurs[0][0] == "000001"
    assert ws.freeze_panes == "A6" and ws.auto_filter.ref

    pm_xlsx, _ = await svc.exporter(profil="PM", format="xlsx", portee="filtre", racine="000003")
    ws = load_workbook(BytesIO(pm_xlsx)).worksheets[0]
    ligne = [c.value for c in ws[6]]
    assert ligne[0] == "000003" and ligne[1] == "00012345"

    pdf, nom_pdf = await svc.exporter(profil="PP", format="pdf", portee="filtre", racine="000001")
    assert pdf[:4] == b"%PDF" and nom_pdf.endswith(".pdf")

    situation = ClienteleSituationService(db, ctx)
    contenu, _ = await situation.exporter(profil="PM", format="xlsx", portee="filtre", racine="000004")
    ws = load_workbook(BytesIO(contenu)).worksheets[0]
    assert ws.cell(row=6, column=1).value == "000004"


async def test_pdf_multipage_entetes_repetes(db):
    from app.services.reporting_export import build_styled_pdf

    rows = [[f"{i:06d}", "NOM TRÈS LONG AVEC ACCENTS ÉÈÀ " * 2, "MOYEN"] for i in range(400)]
    pdf = build_styled_pdf(report_title="Test", headers=["Racine", "Nom", "Classe"], rows=rows)
    assert pdf.count(b"/Type /Page") - pdf.count(b"/Type /Pages") > 1


async def test_lot_par_selection_refuse_le_melange(db):
    ctx = await _ctx(db)
    await _seed(db)
    lots = ClienteleLotService(db, ctx)
    with pytest.raises(AppError) as err:
        await lots.depuis_selection(profil="PP", racines=["000001", "000003"], analyste_kyc=None)
    assert err.value.code == "PROFIL_MELANGE"
    lot = await lots.depuis_selection(profil="PP", racines=["1", "000002", "000001"], analyste_kyc=None)
    assert lot["profil"] == "PP" and lot["racines"] == 2

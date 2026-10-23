# ruff: noqa: F811
"""Réclamations — décisions de finalisation (octobre 2026).

Recette (transaction annulée, données fictives) :
- état « Suspendue » : transitions, motif, étape conservée, reprise, blocage du circuit, permissions,
  filtres, liste, tableau de bord, exports ; distinct de clôturé / archivé / supprimé ; aucun délai ;
- règles par motif validées par le porteur du projet (valeurs inchangées, appliquées, pas une
  validation Conformité) ; cohérence avec la nomenclature des 26 types ;
- unités Marketing : deux libellés distincts, sélectionnables (organigramme CORE) ;
- plus aucune mention fonctionnelle des délais dans le catalogue.
"""

from __future__ import annotations

import uuid
from io import BytesIO

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy import select, text
from sqlalchemy.exc import IntegrityError

from app.api.v1.endpoints.reclamations import organigramme_concerne
from app.data.plateforme_catalogue import FUNCTIONAL_PERMISSIONS, PLATEFORME_MODULES, RBAC_ROLES
from app.data.reclamations_nomenclature import NOMENCLATURE
from app.data.reclamations_regles import REGLES_EXTOURNES, TYPE_EXTOURNES
from app.models import OrgEntite, Reclamation, ReclamationEvenement, ReclamationReferentiel, ReclamationRegleMotif
from app.services.reclamations import circuit, dossiers, regles, reporting
from app.services.reclamations import constantes as C
from app.services.reclamations.acces import ReclamationsScope, permissions_effectives
from tests.test_reclamations_circuit import _refus
from tests.test_reclamations_dossiers import _classement, _saisie, _sr
from tests.test_reclamations_fondations import _unite, _user, db  # noqa: F401

MOTIF = {"motif": "Piece complementaire attendue du client (test)"}


def _m() -> str:
    return f"SUSP-{uuid.uuid4().hex[:8]}"


async def _dossier(db, sr, unite, marqueur: str, **extra) -> dict:
    return await dossiers.creer(db, sr, await _saisie(db, unite.id, objet=f"Dossier {marqueur}", **extra))


async def _en_traitement(db, sr, unite, marqueur: str) -> dict:
    t, _mo, _sm = await _classement(db, sr)
    d = await _dossier(db, sr, unite, marqueur, type_id=uuid.UUID(t["id"]))
    return await circuit.affecter(db, sr, uuid.UUID(d["id"]), {"revision": d["revision"], "unite_id": unite.id})


async def _suspendre(db, s, d: dict, **extra) -> dict:
    return await circuit.suspendre(db, s, uuid.UUID(d["id"]), {"revision": d["revision"], **MOTIF, **extra})


async def _evenements(db, rid: str) -> list[str]:
    return list((await db.scalars(select(ReclamationEvenement.type_evenement).where(
        ReclamationEvenement.reclamation_id == uuid.UUID(rid)).order_by(ReclamationEvenement.date_effet))).all())


# ------------------------------------------------------------------ transitions
async def test_suspension_et_reprise_conservent_etape_et_statut(db):
    sr = await _sr(db)
    unite = await _unite(db)
    d = await _en_traitement(db, sr, unite, _m())
    assert (d["statut"], d["etape"]) == ("EN_TRAITEMENT", "TRAITEMENT")
    assert d["actions"]["suspendre"] and not d["actions"]["reprendre"]

    s = await _suspendre(db, sr, d)
    assert (s["statut"], s["statut_libelle"], s["etape"]) == ("SUSPENDUE", C.STATUT_LIBELLES["SUSPENDUE"], "TRAITEMENT")
    assert (s["groupe"], s["termine"]) == ("EN_COURS", False)
    assert s["suspension"]["statut_avant"] == "EN_TRAITEMENT"
    assert s["suspension"]["motif"] == MOTIF["motif"]
    assert s["suspension"]["depuis"]
    assert s["suspension"]["statut_avant_libelle"] == C.STATUT_LIBELLES["EN_TRAITEMENT"]
    assert s["date_cloture"] is None and s["archivee_le"] is None and s["supprimee_le"] is None
    r = await db.get(Reclamation, uuid.UUID(s["id"]))
    assert r.echeance is None and r.date_cloture is None

    a = s["actions"]
    assert a["reprendre"] and not a["suspendre"]
    assert not any(a[k] for k in ("affecter", "traiter_interne", "evaluer", "cloturer", "refuser",
                                  "annuler", "rouvrir", "archiver"))
    assert a["documents_deposer"] and a["documents_voir"] and a["modifier"]

    p = await circuit.reprendre(db, sr, uuid.UUID(s["id"]), {"revision": s["revision"], "motif": "Piece recue"})
    assert (p["statut"], p["etape"], p["suspension"]) == ("EN_TRAITEMENT", "TRAITEMENT", None)
    assert p["actions"]["suspendre"]
    await db.refresh(r)
    assert (r.statut_avant_suspension, r.motif_suspension, r.date_suspension) == (None, None, None)
    ev = await _evenements(db, p["id"])
    assert ev[-2:] == ["SUSPENSION", "REPRISE"]


async def test_suspension_depuis_reception_puis_circuit_normal(db):
    sr = await _sr(db)
    unite = await _unite(db)
    d = await _dossier(db, sr, unite, _m())
    s = await _suspendre(db, sr, d)
    assert (s["statut"], s["etape"], s["suspension"]["statut_avant"]) == ("SUSPENDUE", "RECEPTION", "A_QUALIFIER")
    rid = uuid.UUID(s["id"])
    for action, corps in (("qualifier", {}), ("affecter", {"unite_id": unite.id}), ("refuser", MOTIF),
                          ("annuler", MOTIF), ("traiter_interne", MOTIF), ("cloturer", {"resultat": "FAVORABLE"}),
                          ("suspendre", MOTIF)):
        await _refus(circuit.executer(db, sr, rid, action, {"revision": s["revision"], **corps}), "DOSSIER_SUSPENDU")
    p = await circuit.reprendre(db, sr, rid, {"revision": s["revision"]})
    assert p["statut"] == "A_QUALIFIER"
    rej = await circuit.refuser(db, sr, rid, {"revision": p["revision"], "motif": "Hors perimetre"})
    assert rej["statut"] == "REJETEE"


async def test_suspension_motif_obligatoire_et_etats_refuses(db):
    sr = await _sr(db)
    unite = await _unite(db)
    d = await _dossier(db, sr, unite, _m())
    await _refus(circuit.suspendre(db, sr, uuid.UUID(d["id"]), {"revision": d["revision"]}), "MOTIF_OBLIGATOIRE")
    await _refus(circuit.suspendre(db, sr, uuid.UUID(d["id"]), {"revision": d["revision"] + 7, **MOTIF}),
                 "CONFLIT_REVISION")
    await _refus(circuit.reprendre(db, sr, uuid.UUID(d["id"]), {"revision": d["revision"]}), "ACTION_INVALIDE")
    rej = await circuit.refuser(db, sr, uuid.UUID(d["id"]), {"revision": d["revision"], "motif": "Hors perimetre"})
    assert rej["actions"]["suspendre"] is False
    await _refus(_suspendre(db, sr, rej), "DOSSIER_TERMINE")
    arch = await circuit.archiver(db, sr, uuid.UUID(rej["id"]), {"revision": rej["revision"]})
    await _refus(_suspendre(db, sr, arch), "DOSSIER_ARCHIVE")
    await _refus(circuit.reprendre(db, sr, uuid.UUID(arch["id"]), {"revision": arch["revision"]}), "DOSSIER_ARCHIVE")


async def test_suspension_reservee_au_service_central_habilite(db):
    sr = await _sr(db)
    unite = await _unite(db)
    d = await _dossier(db, sr, unite, _m())
    rid = uuid.UUID(d["id"])
    sans_droit = ReclamationsScope(user=await _user(db), permissions=permissions_effectives(
        {"reclamations.view", "reclamations.scope.all", "reclamations.treat"}))
    with pytest.raises(HTTPException) as e:
        await _suspendre(db, sans_droit, d)
    assert e.value.status_code == 403
    hors_sr = ReclamationsScope(user=await _user(db), permissions=permissions_effectives(
        {"reclamations.view", "reclamations.scope.all", "reclamations.assign"}))
    await _refus(_suspendre(db, hors_sr, d), "SUSPENSION_RESERVEE_SR")
    vue = await circuit.fiche(db, hors_sr, await dossiers.charger(db, hors_sr, rid))
    assert vue["actions"]["suspendre"] is False
    agent_sr = ReclamationsScope(user=await _user(db), permissions=permissions_effectives(
        {"reclamations.view", "reclamations.scope.all", "reclamations.assign"}), membre_service_central=True)
    s = await _suspendre(db, agent_sr, d)
    assert s["statut"] == "SUSPENDUE"
    await _refus(circuit.reprendre(db, hors_sr, rid, {"revision": s["revision"]}), "SUSPENSION_RESERVEE_SR")
    assert (await circuit.reprendre(db, agent_sr, rid, {"revision": s["revision"]}))["statut"] == "A_QUALIFIER"


async def test_contrainte_base_statut_avant_suspension(db):
    sr = await _sr(db)
    unite = await _unite(db)
    d = await _dossier(db, sr, unite, _m())
    rid = uuid.UUID(d["id"])
    with pytest.raises(IntegrityError):
        async with db.begin_nested():
            await db.execute(text(
                "UPDATE reclamations SET statut = 'SUSPENDUE', statut_avant_suspension = NULL WHERE id = :id"
            ), {"id": rid})


# ------------------------------------------------- filtres, liste, KPI, exports
async def test_suspendue_dans_filtres_tableau_de_bord_et_exports(db):
    sr = await _sr(db)
    unite = await _unite(db)
    m = _m()
    s = await _suspendre(db, sr, await _dossier(db, sr, unite, m))
    await _dossier(db, sr, unite, m)

    async def total(**f) -> int:
        return (await dossiers.lister(db, sr, {"q": m, **f}, taille=100))["total"]

    assert await total(statut=["SUSPENDUE"]) == 1
    assert await total(groupe=["EN_COURS"]) == 2 and await total(ouvertes=True) == 2
    ligne = next(x for x in (await dossiers.lister(db, sr, {"q": m, "statut": ["SUSPENDUE"]}))["items"])
    assert (ligne["id"], ligne["statut_libelle"]) == (s["id"], C.STATUT_LIBELLES["SUSPENDUE"])

    tb = await reporting.tableau_de_bord(db, sr, {"q": m})
    assert (tb["ouvertes"], tb["suspendues"]) == (2, 1)
    assert {x["code"]: x["n"] for x in tb["statuts"]}["SUSPENDUE"] == 1
    lien = tb["liens"]["suspendues"]
    assert lien["statut"] == "SUSPENDUE" and await total(statut=[lien["statut"]]) == tb["suspendues"]

    classeur = load_workbook(BytesIO(await reporting.exporter(db, sr, {"q": m, "statut": ["SUSPENDUE"]})),
                             read_only=True)
    feuille = classeur.worksheets[0]
    lignes = [r for r in feuille.iter_rows(values_only=True) if r and str(r[0]).startswith("REC-")]
    assert len(lignes) == 1
    assert C.STATUT_LIBELLES["SUSPENDUE"] in lignes[0]


def test_statut_suspendu_distinct_et_sans_delai():
    assert C.STATUT_LIBELLES["SUSPENDUE"] == "Suspendue"
    assert "SUSPENDUE" not in C.STATUT_LIBELLES_HISTORIQUES
    assert C.STATUT_SUSPENDU not in C.STATUTS_TERMINAUX
    assert C.STATUT_GROUPES[C.STATUT_SUSPENDU] == "EN_COURS"
    assert C.STATUTS_SUSPENDABLES.isdisjoint(C.STATUTS_TERMINAUX | {C.STATUT_SUSPENDU})
    assert {"suspendre", "reprendre"} <= circuit.ACTIONS.keys()
    assert "reclamations.suspend" not in {c for c, _l, _m in FUNCTIONAL_PERMISSIONS}
    assert len([c for c, _l, _m in FUNCTIONAL_PERMISSIONS if c.startswith("reclamations.")]) == 28


# ------------------------------------------------------------ règles par motif
async def test_regles_extournes_validees_par_le_porteur_valeurs_inchangees(db):
    rows = (await db.execute(select(ReclamationRegleMotif, ReclamationReferentiel.libelle).join(
        ReclamationReferentiel, ReclamationReferentiel.id == ReclamationRegleMotif.motif_id))).all()
    if not rows:
        pytest.skip("Regles par motif non initialisees")
    par_libelle = {lib: r for r, lib in rows}
    assert set(REGLES_EXTOURNES) <= set(par_libelle)
    for libelle, source in REGLES_EXTOURNES.items():
        r = par_libelle[libelle]
        assert r.regles == regles.normaliser_regle(source), libelle
        if r.validation_origine == "PORTEUR_PROJET":
            assert (r.statut_validation, r.valide_par_id) == ("VALIDEE", None) and r.valide_le
        ex = await regles.exigences(db, r.motif_id)
        d = ex.to_dict()
        assert ex.regle == regles.normaliser_regle(source)
        assert d["statut_validation"] == r.statut_validation
    porteur = [r for r, _l in rows if r.validation_origine == "PORTEUR_PROJET"]
    assert porteur, "migration 20261023_reclamations_11 non appliquee"
    libelle = C.ORIGINE_VALIDATION_REGLE_LIBELLES["PORTEUR_PROJET"]
    assert "porteur du projet" in libelle
    assert "hors validation" in libelle


async def test_regle_modifiee_redevient_proposee_puis_validee_par_habilite(db):
    sr = await _sr(db)
    _t, mo, _sm = await _classement(db, sr)
    motif = uuid.UUID(mo["id"])
    r = await regles.enregistrer(db, sr, motif, {"champs_requis": ["reference", "montant", "devise"]})
    assert (r["statut_validation"], r["validation_origine"]) == ("PROPOSEE", None)
    v = await regles.valider(db, sr, motif, "VALIDEE")
    assert v["validation_origine"] == "HABILITE"
    assert v["statut_validation_libelle"] == C.VALIDATION_REGLE_LIBELLES["VALIDEE"]
    row = await db.scalar(select(ReclamationRegleMotif).where(ReclamationRegleMotif.motif_id == motif))
    row.validation_origine, row.valide_par_id = "PORTEUR_PROJET", None
    await db.flush()
    assert (await regles.exigences(db, motif)).to_dict()["validation_origine"] == "PORTEUR_PROJET"
    m2 = await regles.enregistrer(db, sr, motif, {"champs_requis": ["reference", "montant"]})
    assert (m2["statut_validation"], m2["validation_origine"]) == ("PROPOSEE", None)


def test_nomenclature_26_types_coherente_avec_les_regles():
    types = [t for _f, t, _n, _m in NOMENCLATURE]
    assert len(types) == 26 == len(set(types))
    extournes = next(m for _f, t, _n, m in NOMENCLATURE if t == TYPE_EXTOURNES)
    assert set(REGLES_EXTOURNES) == set(extournes)
    assert all(isinstance(m, str) for _f, _t, _n, motifs in NOMENCLATURE for m in motifs)


# ------------------------------------------------------------------ Marketing
async def test_unites_marketing_distinctes_et_selectionnables(db):
    dep = await db.scalar(select(OrgEntite).where(OrgEntite.code == "MKT-COM", OrgEntite.deleted_at.is_(None)))
    svc = await db.scalar(select(OrgEntite).where(OrgEntite.code == "MKT-COM-SVC", OrgEntite.deleted_at.is_(None)))
    if dep is None or svc is None:
        pytest.skip("migration 20261023_reclamations_11 non appliquee")
    assert (dep.libelle, dep.type_entite) == ("MARKETING & COMMUNICATION", "DEPARTEMENT")
    assert (svc.libelle, svc.type_entite, svc.parent_id) == ("MARKETING ET COMMUNICATION", "SERVICE", dep.id)
    assert dep.a_confirmer and svc.a_confirmer and dep.id != svc.id

    ctx = await dossiers.Contexte.charger(db)
    org = organigramme_concerne(ctx)
    assert str(dep.id) in {x["id"] for x in org["departements"]}
    service = next(x for x in org["services"] if x["id"] == str(svc.id))
    assert service["departement_id"] == str(dep.id) and service["a_confirmer"]

    sr = await _sr(db)
    unite = await _unite(db)
    m = _m()
    d = await _dossier(db, sr, unite, m, org_service_id=svc.id)
    assert (d["departement_concerne"], d["service_concerne"]) == (
        "MARKETING & COMMUNICATION", "MARKETING ET COMMUNICATION")
    assert (await dossiers.lister(db, sr, {"q": m, "org_service_ids": [svc.id]}))["total"] == 1
    assert (await dossiers.lister(db, sr, {"q": m, "org_departement_ids": [dep.id]}))["total"] == 1
    tb = await reporting.tableau_de_bord(db, sr, {"q": m})
    assert {x["libelle"] for x in tb["departements"]} == {"MARKETING & COMMUNICATION"}


# -------------------------------------------------------- délais & alertes
async def test_plus_de_mention_fonctionnelle_des_delais(db):
    module = next(x for x in PLATEFORME_MODULES if x["code"] == "reclamations")
    assert "delai" not in module["description"].lower().replace("é", "e")
    admin = next(r for r in RBAC_ROLES if r[0] == "reclamations.admin")
    assert "delai" not in admin[2].lower().replace("é", "e")
    en_base = await db.scalar(text("SELECT description FROM roles WHERE code = 'reclamations.admin'"))
    if en_base is not None:
        assert "delai" not in en_base.lower().replace("é", "e")
    assert not any("delai" in c for c, _l, _m in FUNCTIONAL_PERMISSIONS)

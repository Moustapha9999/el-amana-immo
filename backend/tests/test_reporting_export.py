from io import BytesIO

from openpyxl import load_workbook

from app.services.reporting_export import (
    build_styled_pdf,
    build_styled_workbook,
    ecritures_to_excel,
    format_export_datetime,
    format_period_label,
)
from datetime import date


def test_ecritures_excel_empty():
    data = ecritures_to_excel([])
    assert data[:2] == b"PK"
    wb = load_workbook(BytesIO(data))
    ws = wb.active
    assert "Banque El Amana" in str(ws["A1"].value)
    assert "Exporté le" in str(ws["A3"].value)


def test_styled_workbook_contains_timestamp_and_headers():
    data = build_styled_workbook(
        sheet_title="Test",
        report_title="Rapport test",
        headers=["A", "B"],
        rows=[["1", "2"]],
    )
    wb = load_workbook(BytesIO(data))
    ws = wb.active
    assert ws["A2"].value == "Rapport test"
    assert "Exporté le" in str(ws["A3"].value)
    assert ws["A5"].value == "A"
    assert ws["B6"].value == "2"


def test_styled_pdf_bytes():
    data = build_styled_pdf(
        report_title="Rapport PDF",
        headers=["Col"],
        rows=[["valeur"]],
        landscape_mode=False,
    )
    assert data[:4] == b"%PDF"


def test_format_helpers():
    assert "à" in format_export_datetime()
    assert format_period_label(date(2026, 1, 1), date(2026, 1, 31)).startswith("Période")


def test_cessions_list_export_empty():
    from app.services.reporting_export import cessions_to_excel, cessions_to_pdf

    assert cessions_to_excel([])[:2] == b"PK"
    assert cessions_to_pdf([])[:4] == b"%PDF"


def test_rebuts_list_export_empty():
    from app.services.reporting_export import rebuts_to_excel, rebuts_to_pdf

    assert rebuts_to_excel([])[:2] == b"PK"
    assert rebuts_to_pdf([])[:4] == b"%PDF"


def test_reevaluations_list_export_empty():
    from app.services.reporting_export import reevaluations_to_excel, reevaluations_to_pdf

    assert reevaluations_to_excel([])[:2] == b"PK"
    assert reevaluations_to_pdf([])[:4] == b"%PDF"

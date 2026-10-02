"""Inventaire mensuel Stock & Fournitures — écarts, saisie, workflow, import Excel, exports (sans DB)."""

from datetime import date
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock, MagicMock, patch
from uuid import uuid4

import pytest
from openpyxl import Workbook, load_workbook

from app.core.exceptions import AppError
from app.schemas.mg_stock import InventaireImportOptions, InventaireLigneSaisie
from app.services import mg_inventaire_export as export
from app.services.mg_inventaire_import import (
    InventaireImportService,
    _detecter_entete,
    _match_champ,
    _quantite,
    normaliser,
)
from app.services.mg_inventaire_service import (
    MSG_VERROU,
    PERM_AJUSTEMENT,
    PERM_GESTION,
    PERM_SAISIE,
    PERM_VALIDATION,
    MgInventaireService,
    ajustement_ligne,
    ecart_a_regulariser,
    ecart_pourcentage,
    periode_label,
    statut_ligne,
    theorique_reference,
)
from app.services.mg_inventaire_reference import InventaireReferenceService

TOUTES = {PERM_SAISIE, PERM_VALIDATION, PERM_AJUSTEMENT, PERM_GESTION}


def _stats(**kw):
    base = {
        "total": 10, "a_compter": 10, "comptes": 10, "non_comptes": 0, "exclus": 0,
        "sans_ecart": 9, "ecarts_negatifs": 1, "ecarts_positifs": 0,
        "total_theorique": Decimal("100"), "total_physique": Decimal("99"),
        "ecart_net": Decimal("-1"), "progression": 100.0,
    }
    base.update(kw)
    return base


def _inv(statut="EN_COURS", **kw):
    data = dict(
        id=uuid4(), reference="INV-2026-09-001", statut=statut, date_debut=date(2026, 9, 30),
        date_fin=None, agence_id=None, famille_id=None, periode_id=None, ajustements_at=None,
        validation_forcee=False, valide_at=None, valide_by=None, motif_annulation=None, import_meta=None,
    )
    data.update(kw)
    return SimpleNamespace(**data)


def _svc(perms=TOUTES, inv=None, stats=None):
    db = MagicMock()
    db.flush = AsyncMock()
    db.commit = AsyncMock()
    user = SimpleNamespace(id=uuid4(), full_name="Testeur")
    svc = MgInventaireService(db, user, set(perms))
    svc._audit = AsyncMock()
    if inv is not None:
        svc.get = AsyncMock(return_value=inv)
        svc.stats = AsyncMock(return_value={inv.id: stats or _stats()})
        svc.nb_ajustements = AsyncMock(return_value=0)
        svc.mouvements_depuis = AsyncMock(return_value=0)
    return svc


# --------------------------------------------------------------------------- écarts


@pytest.mark.parametrize(
    "theo, phys, ecart, statut, pct",
    [
        (10, 10, 0, "CONFORME", 0.0),
        (10, 8, -2, "ECART_NEGATIF", -20.0),
        (10, 12, 2, "ECART_POSITIF", 20.0),
        (0, 5, 5, "ECART_POSITIF", None),
        (5, 0, -5, "ECART_NEGATIF", -100.0),
        (0, 0, 0, "CONFORME", 0.0),
    ],
)
def test_ecart_physique_moins_theorique(theo, phys, ecart, statut, pct):
    e = Decimal(phys) - Decimal(theo)
    assert e == ecart
    assert statut_ligne("COMPTE", e) == statut
    assert ecart_pourcentage(Decimal(theo), e) == pct


def test_statut_ligne_non_compte_et_exclu():
    assert statut_ligne("NON_COMPTE", None) == "NON_COMPTE"
    assert statut_ligne("EXCLU", Decimal("3")) == "EXCLU"
    assert ecart_pourcentage(Decimal("4"), None) is None


def test_periode_label():
    assert periode_label(2026, 9) == "Septembre 2026"
    assert periode_label(None, 9) is None


# --------------------------------------------------------------------------- saisie


def _ligne(statut="NON_COMPTE", theo=Decimal("10"), phys=None):
    return SimpleNamespace(
        id=uuid4(), article_id=uuid4(), stock_theorique=theo, stock_physique=phys, ecart=None,
        nature_ecart=None, statut_comptage=statut, observation=None, compte_par=None, compte_at=None,
        updated_at=None, stock_cible=None, stock_theorique_source=None, donnees_source=None,
    )


def _saisie_svc(inv, ligne, perms=TOUTES):
    svc = _svc(perms, inv)
    svc.db.scalar = AsyncMock(return_value=ligne)
    svc.db.get = AsyncMock(return_value=SimpleNamespace(code="B0028"))
    svc.ligne = AsyncMock(return_value={"id": ligne.id})
    return svc


@pytest.mark.asyncio
@pytest.mark.parametrize("theo, phys", [(0, 5), (5, 0), (0, 0), (84, 85)])
async def test_saisie_calcule_ecart_et_marque_compte(theo, phys):
    inv = _inv("EN_COURS")
    ligne = _ligne(theo=Decimal(theo))
    svc = _saisie_svc(inv, ligne)
    out = await svc.saisir_ligne(inv.id, ligne.id, InventaireLigneSaisie(stock_physique=phys))
    assert ligne.statut_comptage == "COMPTE"
    assert ligne.stock_physique == phys
    assert ligne.ecart == Decimal(phys) - Decimal(theo)
    assert out["statut"] == "EN_COURS"
    action = svc._audit.await_args_list[0].args[0]
    assert action == "saisie_physique"


@pytest.mark.asyncio
async def test_premiere_saisie_demarre_le_brouillon():
    inv = _inv("BROUILLON")
    ligne = _ligne()
    svc = _saisie_svc(inv, ligne)
    out = await svc.saisir_ligne(inv.id, ligne.id, InventaireLigneSaisie(stock_physique=3))
    assert inv.statut == "EN_COURS" and out["statut"] == "EN_COURS"
    assert [c.args[0] for c in svc._audit.await_args_list] == ["saisie_physique", "demarrer"]


@pytest.mark.asyncio
async def test_effacer_remet_non_compte():
    inv = _inv()
    ligne = _ligne("COMPTE", phys=Decimal("4"))
    svc = _saisie_svc(inv, ligne)
    await svc.saisir_ligne(inv.id, ligne.id, InventaireLigneSaisie(effacer=True))
    assert ligne.statut_comptage == "NON_COMPTE" and ligne.stock_physique is None and ligne.ecart is None


@pytest.mark.asyncio
async def test_exclure_puis_saisie_refusee():
    inv = _inv()
    ligne = _ligne()
    svc = _saisie_svc(inv, ligne)
    await svc.saisir_ligne(inv.id, ligne.id, InventaireLigneSaisie(exclure=True))
    assert ligne.statut_comptage == "EXCLU"
    with pytest.raises(AppError) as exc:
        await svc.saisir_ligne(inv.id, ligne.id, InventaireLigneSaisie(stock_physique=2))
    assert exc.value.code == "LIGNE_EXCLUE"


def test_quantite_negative_refusee_par_schema():
    with pytest.raises(ValueError):
        InventaireLigneSaisie(stock_physique=-1)


@pytest.mark.asyncio
@pytest.mark.parametrize("statut", ["VALIDE", "AJUSTE", "ARCHIVE", "ANNULE"])
async def test_saisie_verrouillee_apres_validation(statut):
    inv = _inv(statut)
    ligne = _ligne()
    svc = _saisie_svc(inv, ligne)
    with pytest.raises(AppError) as exc:
        await svc.saisir_ligne(inv.id, ligne.id, InventaireLigneSaisie(stock_physique=1))
    assert exc.value.message == MSG_VERROU


@pytest.mark.asyncio
async def test_saisie_sans_permission():
    inv = _inv()
    svc = _saisie_svc(inv, _ligne(), perms={"mg.stock.view"})
    with pytest.raises(AppError) as exc:
        await svc.saisir_ligne(inv.id, uuid4(), InventaireLigneSaisie(stock_physique=1))
    assert exc.value.status_code == 403


# --------------------------------------------------------------------------- workflow


@pytest.mark.asyncio
async def test_soumettre_sans_comptage_refuse():
    inv = _inv("EN_COURS")
    svc = _svc(inv=inv, stats=_stats(comptes=0, non_comptes=10))
    with pytest.raises(AppError):
        await svc.transition(inv.id, "soumettre")


@pytest.mark.asyncio
async def test_valider_avec_non_comptes_refuse_sans_forcage():
    inv = _inv("A_CONTROLER")
    svc = _svc(inv=inv, stats=_stats(comptes=8, non_comptes=2))
    with pytest.raises(AppError) as exc:
        await svc.transition(inv.id, "valider")
    assert exc.value.code == "INVENTAIRE_INCOMPLET"


@pytest.mark.asyncio
async def test_forcer_requiert_permission_gestion_et_motif():
    inv = _inv("A_CONTROLER")
    svc = _svc({PERM_VALIDATION}, inv=inv, stats=_stats(comptes=8, non_comptes=2))
    with pytest.raises(AppError) as exc:
        await svc.transition(inv.id, "valider", forcer=True, motif="ok")
    assert exc.value.status_code == 403
    svc = _svc(inv=inv, stats=_stats(comptes=8, non_comptes=2))
    with pytest.raises(AppError):
        await svc.transition(inv.id, "valider", forcer=True, motif="  ")
    await svc.transition(inv.id, "valider", forcer=True, motif="Articles introuvables")
    assert inv.statut == "VALIDE" and inv.validation_forcee is True


@pytest.mark.asyncio
async def test_valider_complet():
    inv = _inv("A_CONTROLER")
    svc = _svc(inv=inv)
    await svc.transition(inv.id, "valider")
    assert inv.statut == "VALIDE" and inv.valide_at is not None
    svc._audit.assert_awaited()


@pytest.mark.asyncio
async def test_valider_deux_fois_refuse():
    inv = _inv("VALIDE")
    svc = _svc(inv=inv)
    with pytest.raises(AppError) as exc:
        await svc.transition(inv.id, "valider")
    assert exc.value.code == "INVENTAIRE_DEJA_VALIDE"


@pytest.mark.asyncio
async def test_ajustements_generes_une_seule_fois():
    inv = _inv("VALIDE")
    svc = _svc(inv=inv)
    svc._generer_ajustements = AsyncMock(return_value={"ajustements": 1, "ecart_net": 1})
    await svc.transition(inv.id, "generer_ajustements")
    assert inv.statut == "AJUSTE" and inv.ajustements_at is not None
    with pytest.raises(AppError) as exc:
        await svc.transition(inv.id, "generer_ajustements")
    assert exc.value.message == "Les ajustements de cet inventaire ont déjà été appliqués."
    svc._generer_ajustements.assert_awaited_once()


@pytest.mark.asyncio
async def test_ajustements_requierent_permission():
    inv = _inv("VALIDE")
    svc = _svc({PERM_VALIDATION, PERM_SAISIE}, inv=inv)
    with pytest.raises(AppError) as exc:
        await svc.transition(inv.id, "generer_ajustements")
    assert exc.value.status_code == 403


@pytest.mark.asyncio
async def test_ajustements_avant_validation_refuses():
    inv = _inv("A_CONTROLER")
    svc = _svc(inv=inv)
    with pytest.raises(AppError):
        await svc.transition(inv.id, "generer_ajustements")


@pytest.mark.asyncio
async def test_annuler_motif_obligatoire_et_bloque_si_ajustements():
    inv = _inv("EN_COURS")
    svc = _svc(inv=inv)
    with pytest.raises(AppError):
        await svc.transition(inv.id, "annuler")
    svc.nb_ajustements = AsyncMock(return_value=3)
    with pytest.raises(AppError):
        await svc.transition(inv.id, "annuler", motif="Erreur")
    svc.nb_ajustements = AsyncMock(return_value=0)
    await svc.transition(inv.id, "annuler", motif="Erreur de périmètre")
    assert inv.statut == "ANNULE" and inv.motif_annulation == "Erreur de périmètre"


@pytest.mark.asyncio
async def test_archiver_apres_ajustement():
    inv = _inv("AJUSTE")
    svc = _svc(inv=inv)
    await svc.transition(inv.id, "archiver")
    assert inv.statut == "ARCHIVE"


# --------------------------------------------------------------------------- rapprochement banque


def test_b0028_ajustement_vers_stock_retenu_et_ecart_a_regulariser():
    ligne = _ligne("COMPTE", theo=Decimal("29"), phys=Decimal("85"))
    ligne.stock_theorique_source = Decimal("84")
    ligne.stock_cible = Decimal("84")
    assert theorique_reference(ligne) == 84
    assert ajustement_ligne(ligne) == 55
    assert ecart_a_regulariser(ligne) == 1


def test_ligne_sans_cible_ajuste_au_physique():
    ligne = _ligne("COMPTE", theo=Decimal("2415"), phys=Decimal("1915"))
    ligne.stock_theorique_source = Decimal("1900")
    assert theorique_reference(ligne) == 2415
    assert ajustement_ligne(ligne) == -500
    assert ecart_a_regulariser(ligne) is None
    assert ajustement_ligne(_ligne("EXCLU", theo=Decimal("21"))) is None


@pytest.mark.asyncio
async def test_saisie_ligne_rapprochee_ecart_vs_theorique_banque():
    inv = _inv("EN_COURS")
    ligne = _ligne(theo=Decimal("29"))
    ligne.stock_theorique_source = Decimal("84")
    ligne.stock_cible = Decimal("84")
    svc = _saisie_svc(inv, ligne)
    await svc.saisir_ligne(inv.id, ligne.id, InventaireLigneSaisie(stock_physique=85))
    assert ligne.ecart == 1


@pytest.mark.asyncio
async def test_archivage_bloque_si_controle_rapprochement_en_echec():
    inv = _inv("AJUSTE", import_meta={"rapprochement": {"mouvements": {}}})
    svc = _svc(inv=inv)
    svc.controles = AsyncMock(return_value=[
        {"code": "stock_final", "libelle": "Stock final BEA DIGITAL", "attendu": 8930, "obtenu": 9175, "ok": False},
    ])
    with pytest.raises(AppError) as exc:
        await svc.transition(inv.id, "archiver")
    assert exc.value.code == "RAPPROCHEMENT_CONTROLES" and inv.statut == "AJUSTE"
    svc.controles = AsyncMock(return_value=[{"code": "stock_final", "libelle": "x", "attendu": 1, "obtenu": 1, "ok": True}])
    await svc.transition(inv.id, "archiver")
    assert inv.statut == "ARCHIVE"


@pytest.mark.asyncio
async def test_ajustements_refuses_si_mouvements_deja_existants():
    inv = _inv("VALIDE")
    svc = _svc(inv=inv)
    svc.nb_ajustements = AsyncMock(return_value=17)
    svc._generer_ajustements = AsyncMock()
    with pytest.raises(AppError) as exc:
        await svc.transition(inv.id, "generer_ajustements")
    assert exc.value.code == "INVENTAIRE_DEJA_AJUSTE"
    svc._generer_ajustements.assert_not_awaited()


REF_ENTETES = [
    "N°", "Fiche", "Article (libellé de la fiche)", "Réf.", "Catégorie", "Statut agence", "Stock initial",
    "Entrées", "Sorties", "Stock final théorique", "Vérifié (OK)", "Stock physique constaté",
    "Écart (physique − théorique)", "Statut inventaire", "Stock actuel agence", "Consommation (sorties)",
    "Alerte stock", "Observations",
]


def _reference_xlsx(*, plan_b0028=55, controle_variation=54, ancien_dans_plan=False, inventaire="INV-2026-09-004"):
    wb = Workbook()
    wb.active.title = "LisezMoi"
    ws = wb.create_sheet("Inventaire_Reference")
    ws.append(REF_ENTETES)
    ws.append([1, 1, "Agrafeuse 24/6", "B0001", "Petit matériel", "Agence actuelle", 22, 0, 2, 20, "OK", 20, 0,
               "Conforme", 20, 2, "OK", None])
    ws.append([2, 1, "Rame de Papiers", "B0028", "Papeterie", "Agence actuelle", 100, 0, 16, 84, "OK", 85, 1,
               "Écart en plus", 84, 16, "OK", None])
    ws.append([3, 1, "Cartouche 12 A", "C0002", "Consommables", "Ancienne agence", 21, 0, 0, 21, "Non (X)", None,
               None, None, None, 0, None, None])
    ws = wb.create_sheet("Mouvements_A_Appliquer")
    ws.append(["Réf.", "Article", "Stock BEA actuel", "Stock final théorique banque", "Stock actuel agence banque",
               "Stock physique banque", "Écart physique vs théorique banque", "Ajustement à appliquer",
               "Type mouvement", "Source"])
    ws.append(["B0001", "Agrafeuse 24/6", 21, 20, 20, 20, 0, -1, "AJUSTEMENT", "Inventaire bancaire Septembre 2026"])
    ws.append(["B0028", "Rame de Papiers", 29, 84, 84, 85, 1, plan_b0028, "AJUSTEMENT",
               "Inventaire bancaire Septembre 2026"])
    if ancien_dans_plan:
        ws.append(["C0002", "Cartouche 12 A", 21, 21, 0, 0, 0, -21, "AJUSTEMENT", "x"])
    ws.append(["TOTAL", None, None, 104, 104, 105, 1, 54, None, None])
    ws = wb.create_sheet("Ecarts_Banque")
    ws.append(["Réf.", "Article", "Stock final théorique banque", "Stock physique", "Écart", "Statut", "Décision"])
    ws.append(["B0028", "Rame de Papiers", 84, 85, 1, "Écart en plus", "À régulariser séparément"])
    ws = wb.create_sheet("Controle")
    ws.append(["Contrôle", "Résultat", "Attendu / interprétation"])
    for libelle, valeur in [
        ("Références bancaires", 3), ("Articles agence actuelle", 2), ("Ancienne agence", 1),
        ("Erreurs formule banque", 0), ("Stock actuel agence banque", 104),
        ("Stock physique agence actuelle", 105), ("Écart physique non régularisé", 1),
        ("Ajustement total vers stock actuel banque", controle_variation), ("Nombre de mouvements à appliquer", 2),
    ]:
        ws.append([libelle, valeur, str(valeur)])
    ws = wb.create_sheet("Meta")
    ws.append(["Clé", "Valeur"])
    for cle, valeur in [("Source autoritaire", "Inventaire_Stocks.xlsx"), ("Inventaire", inventaire),
                        ("Date inventaire", "30/09/2026"), ("Comptage bancaire", "02/10/2026")]:
        ws.append([cle, valeur])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _reference_svc():
    inv = _inv("EN_COURS", reference="INV-2026-09-004")
    svc = _svc(inv=inv)
    rows = []
    for code, systeme in (("B0001", 21), ("B0028", 29), ("C0002", 21)):
        ligne = _ligne("COMPTE", theo=Decimal(systeme))
        rows.append((ligne, SimpleNamespace(id=ligne.article_id, code=code, stock_actuel=Decimal(systeme))))
    result = MagicMock()
    result.all.return_value = rows
    svc.db.execute = AsyncMock(return_value=result)
    return inv, InventaireReferenceService(svc)


@pytest.mark.asyncio
async def test_reference_banque_conforme():
    inv, ref = _reference_svc()
    analyse = await ref.analyser(inv.id, _reference_xlsx(), "ref.xlsx")
    assert analyse["anomalies"] == [] and not analyse["bloquant"]
    assert analyse["mouvements"] == {"B0001": -1, "B0028": 55}
    c = analyse["calcul"]
    assert (c["references"], c["agence_actuelle"], c["ancienne_agence"]) == (3, 2, 1)
    assert (c["stock_systeme"], c["stock_final"], c["stock_physique"], c["ecart_restant"]) == (50, 104, 105, 1)
    assert analyse["registre_ecarts"] == [{"code": "B0028", "ecart": 1}]
    assert analyse["date_comptage"] == "02/10/2026"
    ancien = next(lg for lg in analyse["lignes"] if lg["code"] == "C0002")
    assert ancien["ancienne_agence"] and ancien["ajustement"] is None


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "kwargs, extrait",
    [
        ({"plan_b0028": 56, "controle_variation": 55}, "B0028"),
        ({"controle_variation": -854}, "variation"),
        ({"ancien_dans_plan": True, "controle_variation": 33}, "ancienne agence"),
        ({"inventaire": "INV-2026-08-001"}, "INV-2026-08-001"),
    ],
)
async def test_reference_banque_incoherente_bloque(kwargs, extrait):
    inv, ref = _reference_svc()
    analyse = await ref.analyser(inv.id, _reference_xlsx(**kwargs), "ref.xlsx")
    assert analyse["bloquant"]
    assert any(extrait in a for a in analyse["anomalies"]), analyse["anomalies"]


@pytest.mark.asyncio
async def test_suppression_reservee_brouillon():
    inv = _inv("VALIDE", deleted_at=None, is_active=True)
    svc = _svc(inv=inv)
    with pytest.raises(AppError):
        await svc.supprimer(inv.id)


@pytest.mark.asyncio
async def test_apercu_validation():
    inv = _inv("A_CONTROLER")
    svc = _svc({PERM_VALIDATION}, inv=inv, stats=_stats(comptes=7, non_comptes=3))
    out = await svc.apercu_validation(inv.id)
    assert out["peut_valider"] is False and out["peut_forcer"] is False
    assert "3 article(s) non compté(s)" in out["message"]


def test_permission_gestion_ne_finit_pas_par_admin():
    assert not PERM_GESTION.endswith(".admin")


# --------------------------------------------------------------------------- import Excel


def test_normalisation_et_mapping_entetes():
    assert normaliser("  Stock physique   constaté ") == "stock physique constate"
    assert _match_champ("ref.") == ("code", True)
    assert _match_champ("stock physique constate")[0] == "stock_physique"
    assert _match_champ("stock final theorique")[0] == "stock_theorique"
    assert _match_champ("ecart (physique − theorique)")[0] == "ecart"


@pytest.mark.parametrize(
    "valeur, attendu, erreur",
    [(5, Decimal(5), False), ("12", Decimal(12), False), (None, None, False), ("", None, False),
     (-1, None, True), ("abc", None, True), (2.5, None, True)],
)
def test_quantite(valeur, attendu, erreur):
    q, err = _quantite(valeur)
    assert q == attendu
    assert bool(err) is erreur


def _classeur(rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Inventaire"
    ws.append(["INVENTAIRE SEPTEMBRE 2026"])
    ws.append([])
    ws.append([])
    ws.append(["N°", "Article (libellé de la fiche)", "Réf.", "Catégorie", "Statut agence",
               "Stock final théorique", "Vérifié (OK)", "Stock physique constaté", "Observations"])
    for r in rows:
        ws.append(r)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_detection_entete_ligne_4():
    wb = load_workbook(BytesIO(_classeur([])))
    row, colonnes, _ = _detecter_entete(wb.active)
    assert row == 4
    assert {"code", "designation", "stock_theorique", "stock_physique"} <= set(colonnes.values())


def _art(code, designation, stock=10):
    return SimpleNamespace(
        id=uuid4(), code=code, reference=None, designation=designation, stockable=True,
        is_active=True, agence_id=None, stock_actuel=Decimal(stock),
    )


def _import_svc(articles):
    inv_svc = _svc()
    result = MagicMock()
    result.scalars.return_value.all.return_value = articles
    inv_svc.db.execute = AsyncMock(return_value=result)
    inv_svc.stock_a_date = AsyncMock(return_value={a.id: Decimal(a.stock_actuel) for a in articles})
    inv_svc.articles_perimetre = AsyncMock(return_value=articles)
    return InventaireImportService(inv_svc)


@pytest.mark.asyncio
async def test_analyse_import_detecte_anomalies():
    a1, a2, a3 = _art("B0001", "Stylo bleu", 10), _art("B0028", "Rame papier", 84), _art("C0001", "Encre", 3)
    contenu = _classeur([
        [1, "Stylo bleu", "B0001", "Bureau", "Agence actuelle", 10, "OK", 10, None],
        [2, "Rame papier", "B0028", "Bureau", "Agence actuelle", 84, "OK", 85, "1 de plus"],
        [3, "Encre", "C0001", "Conso", "Ancienne agence", 3, "X", None, None],
        [4, "Inconnu", "Z9999", "Bureau", "Agence actuelle", 1, "OK", 1, None],
        [5, "Stylo bleu", "B0001", "Bureau", "Agence actuelle", 10, "OK", 10, None],
        [6, "Agrafes", "B0050", "Bureau", "Agence actuelle", 2, "OK", -3, None],
    ])
    svc = _import_svc([a1, a2, a3])
    out = await svc.analyser(contenu, "inv.xlsx", InventaireImportOptions(date_inventaire=date(2026, 9, 30)))
    par_ligne = {lg["ligne"]: lg for lg in out["lignes"]}
    assert par_ligne[5]["statut"] == "OK"
    assert par_ligne[6]["statut"] == "OK" and par_ligne[6]["ecart"] == 1
    assert par_ligne[7]["statut"] == "EXCLU"
    assert par_ligne[8]["statut"] == "INCONNU"
    assert par_ligne[9]["statut"] == "DOUBLON"
    assert par_ligne[10]["statut"] in {"INCONNU", "QTE_INVALIDE"}
    assert out["bloquant"] is True
    assert out["resume"]["doublons"] == 1 and out["resume"]["exclus"] == 1


@pytest.mark.asyncio
async def test_analyse_import_resolutions_debloquent():
    a1 = _art("B0001", "Stylo bleu", 10)
    contenu = _classeur([
        [1, "Stylo bleu", "B0001", "Bureau", "Agence actuelle", 10, "OK", 9, None],
        [2, "Stylo bleu (ancien code)", "S-01", "Bureau", "Agence actuelle", 10, "OK", 9, None],
    ])
    svc = _import_svc([a1])
    out = await svc.analyser(
        contenu, "inv.xlsx",
        InventaireImportOptions(date_inventaire=date(2026, 9, 30), resolutions={"6": "IGNORER"}),
    )
    assert out["bloquant"] is False
    assert {lg["ligne"]: lg["statut"] for lg in out["lignes"]} == {5: "OK", 6: "IGNORE"}
    assert out["resume"]["ecarts_physique_systeme"] == 1


@pytest.mark.asyncio
async def test_import_fichier_invalide():
    svc = _import_svc([])
    with pytest.raises(AppError) as exc:
        await svc.analyser(b"pas un excel", "x.xlsx", InventaireImportOptions(date_inventaire=date(2026, 9, 30)))
    assert exc.value.code == "IMPORT_FICHIER_INVALIDE"


# --------------------------------------------------------------------------- exports


def _export_data():
    inv = {
        "reference": "INV-2026-09-001", "libelle": "Inventaire Septembre 2026", "periode_libelle": "Septembre 2026",
        "date_debut": date(2026, 9, 30), "agence_libelle": None, "famille_libelle": None,
        "responsable_nom": "Magasinier", "statut": "VALIDE", "source": "IMPORT_EXCEL", "snapshot_at": None,
        "stats": _stats(), "valide_at": None, "valide_by_nom": "Valideur", "validation_forcee": False,
        "ajustements_at": None, "observation": None,
    }
    base = dict(famille_libelle="Bureau", unite="U", compte_par_nom="Testeur", compte_at=None, observation=None)
    lignes = [
        dict(base, article_code="B0001", article_designation="Stylo", stock_theorique=Decimal(10),
             stock_physique=Decimal(10), ecart=Decimal(0), ecart_pourcentage=0.0, statut_ligne="CONFORME"),
        dict(base, article_code="B0028", article_designation="Rame", stock_theorique=Decimal(84),
             stock_physique=Decimal(85), ecart=Decimal(1), ecart_pourcentage=1.19, statut_ligne="ECART_POSITIF"),
        dict(base, article_code="C0001", article_designation="Encre", stock_theorique=Decimal(3),
             stock_physique=None, ecart=None, ecart_pourcentage=None, statut_ligne="EXCLU"),
    ]
    return inv, lignes


def test_export_excel_trois_feuilles():
    inv, lignes = _export_data()
    wb = load_workbook(BytesIO(export.inventaire_to_excel(inv, lignes)))
    assert wb.sheetnames == ["Synthèse", "Détail", "Écarts"]
    ecarts = [r for r in wb["Écarts"].iter_rows(min_row=6, values_only=True) if r[0]]
    assert len(ecarts) == 1 and ecarts[0][0] == "B0028"


def test_export_pdf():
    inv, lignes = _export_data()
    assert export.inventaire_to_pdf(inv, lignes).startswith(b"%PDF")


@patch("app.services.mg_inventaire_export.export_now")
def test_export_sans_ecart(mock_now):
    from datetime import datetime

    mock_now.return_value = datetime(2026, 10, 2, 10, 0)
    inv, lignes = _export_data()
    assert export.lignes_ecart(lignes[:1]) == []
    assert export.inventaire_to_pdf(inv, lignes[:1]).startswith(b"%PDF")

"""Extrait Matrice V1 + Scoring V4 vers un module Python. Ne décide aucune règle."""
from __future__ import annotations

import json
import re
import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
MATRICE = ROOT / "Matrice des risques LBC FT BEA V1-7-5-25.xlsx"
SCORING = ROOT / "Risque LBC FT Client V4-2026.xlsx"
OUT_PY = ROOT / "backend" / "app" / "data" / "clientele_classif_referentiel.py"
OUT_JSON = ROOT / "backend" / "app" / "data" / "clientele_classif_referentiel.json"


def n(v: object) -> str:
    if v is None:
        return ""
    s = str(v).strip()
    return "" if s.lower() in ("none", "nan") else s


def slug(s: str, maxlen: int = 80) -> str:
    s = s.upper().replace("É", "E").replace("È", "E").replace("Ê", "E").replace("À", "A")
    s = s.replace("Ô", "O").replace("Ù", "U").replace("Ç", "C").replace("Ï", "I")
    s = s.replace("Î", "I").replace("Û", "U").replace("Ä", "A").replace("Ö", "O")
    s = re.sub(r"[^A-Z0-9]+", "_", s).strip("_")
    return (s or "X")[:maxlen]


def norm_niveau(v: object) -> str | None:
    s = n(v).lower().replace("é", "e")
    if not s:
        return None
    if s.startswith("inter") or "inacceptable" in s or "unacceptable" in s:
        return "INTERDIT" if s.startswith("inter") else "ELEVE"
    if "inacceptable" in s:
        return "ELEVE"
    if s.startswith("elev") or s.startswith("high"):
        return "ELEVE"
    if s.startswith("moy") or s.startswith("med"):
        return "MOYEN"
    if s.startswith("faib") or s.startswith("low"):
        return "FAIBLE"
    return None


def poids_vers_niveau(p: int | None) -> str | None:
    if p is None:
        return None
    if p >= 100_000:
        return "ELEVE"
    if p >= 10_000:
        return "MOYEN"
    if p > 0:
        return "FAIBLE"
    return None


def paire_colonnes(ws, row_header: int = 2) -> list[tuple[int, int, str]]:
    headers = list(ws.iter_rows(min_row=row_header, max_row=row_header, max_col=ws.max_column, values_only=True))[0]
    pairs = []
    i = 0
    while i < len(headers):
        h = headers[i]
        if h and i + 1 < len(headers) and headers[i + 1] and "risque" in str(headers[i + 1]).lower():
            pairs.append((i, i + 1, str(h).strip()))
            i += 3
            continue
        i += 1
    return pairs


def extraire_paires(ws, pairs, min_row=3) -> dict[str, list[dict]]:
    out: dict[str, list[dict]] = {}
    for c_val, c_risq, titre in pairs:
        items = []
        for row in ws.iter_rows(min_row=min_row, max_row=ws.max_row, min_col=c_val + 1, max_col=c_risq + 1,
                                values_only=True):
            v, r = n(row[0]), n(row[1])
            if not v and not r:
                continue
            items.append({"libelle": v, "niveau_matrice": norm_niveau(r), "brut_matrice": r})
        out[titre] = items
    return out


def extraire_zone(ws) -> list[dict]:
    out = []
    for row in ws.iter_rows(min_row=3, max_row=ws.max_row, max_col=2, values_only=True):
        pays, niv = n(row[0]), n(row[1])
        if not pays:
            continue
        label = norm_niveau(niv)
        # Inacceptable n'existe pas dans la matrice ; Interdit est explicite.
        if niv.lower().startswith("inter"):
            label = "INTERDIT"
        out.append({"libelle": pays, "niveau_matrice": label, "brut_matrice": niv, "source": "MATRICE_V1"})
    return out


def extraire_produit(ws) -> list[dict]:
    rows = list(ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=5, values_only=True))
    out = []
    if not rows:
        return out
    # paires (A,B) monnaie locale et (D,E) monnaie étrangère
    for row in rows[1:]:
        if n(row[0]):
            out.append({"libelle": n(row[0]), "famille": "MONNAIE_LOCALE",
                        "niveau_matrice": norm_niveau(row[1]), "source": "MATRICE_V1"})
        if n(row[3]):
            out.append({"libelle": n(row[3]), "famille": "MONNAIE_ETRANGERE",
                        "niveau_matrice": norm_niveau(row[4]), "source": "MATRICE_V1"})
    return out


def extraire_canal_matrice(ws) -> list[dict]:
    out = []
    for row in ws.iter_rows(min_row=3, max_row=ws.max_row, max_col=2, values_only=True):
        if n(row[0]):
            out.append({"libelle": n(row[0]), "niveau_matrice": norm_niveau(row[1]), "source": "MATRICE_V1"})
    return out


def extraire_liste_v4(ws) -> dict[str, list[dict]]:
    """Feuille Liste — ignore Pays (K/L) et Nationalité (M/N), lookup cassé."""
    def col(c_lib: int, c_score: int, start: int = 5) -> list[dict]:
        items = []
        for r in range(start, ws.max_row + 1):
            lib, sc = ws.cell(r, c_lib).value, ws.cell(r, c_score).value
            if lib is None:
                continue
            lib_s = n(lib)
            if lib_s.lower() in ("score", "profil", "pays", "nationalité", "nationalite"):
                continue
            try:
                p = int(float(sc)) if sc is not None and not isinstance(sc, str) else None
            except (TypeError, ValueError):
                p = None
            if isinstance(sc, str) and sc.upper() == "SCORE":
                continue
            items.append({"libelle": lib_s, "poids_v4": p, "niveau_v4": poids_vers_niveau(p)})
        return items

    ppe = []
    for r in (1, 2):
        lib, sc = n(ws.cell(r, 9).value), ws.cell(r, 10).value
        if lib:
            p = int(float(sc)) if sc is not None else None
            ppe.append({"libelle": lib, "poids_v4": p, "niveau_v4": poids_vers_niveau(p)})

    return {
        "profil": col(1, 2, 5),
        "personne_physique": col(3, 4, 5),
        "personne_morale_publique": col(5, 6, 5),
        "personne_morale_privee": col(7, 8, 5),
        "personne_morale_association": col(9, 10, 5),
        "ppe_oui_non": ppe,
        "origine_fonds": col(15, 16, 5),
        "secteur": col(17, 18, 5),
        "beneficiaire_effectif": col(19, 20, 5),
        "procuration": col(21, 22, 5),
        "etape_eer": col(24, 25, 5),
        "transaction": col(26, 27, 5),
        "canal": col(28, 29, 5),
        "_ignore": "Colonnes Pays (K/L) et Nationalité (M/N) : lookup cassé, quasi toutes à 100000.",
    }


def extraire_pays_v4(ws) -> list[dict]:
    out = []
    for r in range(4, ws.max_row + 1):
        fr, scp, nat, scn, en, alias, risk = [ws.cell(r, c).value for c in range(1, 8)]
        if not fr and not en:
            continue
        try:
            p_pays = int(float(scp)) if scp is not None else None
        except (TypeError, ValueError):
            p_pays = None
        try:
            p_nat = int(float(scn)) if scn is not None else None
        except (TypeError, ValueError):
            p_nat = None
        aliases = sorted({n(x) for x in (fr, en, alias, nat) if n(x)})
        risk_s = n(risk)
        niveau_label = None
        if "inacceptable" in risk_s.lower() or "unacceptable" in risk_s.lower():
            niveau_label = "INACCEPTABLE_V4"
        else:
            niveau_label = poids_vers_niveau(p_pays)
        out.append({
            "nom_fr": n(fr) or n(en),
            "nom_en": n(en) or None,
            "alias_en_caps": n(alias) or None,
            "nationalite": n(nat) or None,
            "poids_pays_v4": p_pays,
            "poids_nationalite_v4": p_nat,
            "niveau_pays_v4": poids_vers_niveau(p_pays),
            "niveau_nationalite_v4": poids_vers_niveau(p_nat),
            "label_v4": risk_s or None,
            "categorie_v4": niveau_label,
            "aliases": aliases,
            "source": "SCORING_V4_PAYS_ENG_FR",
        })
    return out


def reconcilier_pays(pays_m: list[dict], pays_v4: list[dict]) -> list[dict]:
    def key(s: str) -> str:
        return slug(s)

    idx: dict[str, dict] = {}
    for p in pays_v4:
        for a in p["aliases"] + [p["nom_fr"], p.get("nom_en") or "", p.get("alias_en_caps") or ""]:
            if a:
                idx[key(a)] = p

    fusion = []
    vus: set[int] = set()
    for m in pays_m:
        p4 = idx.get(key(m["libelle"]))
        rec = {
            "code": slug(m["libelle"], 48),
            "nom_fr": m["libelle"],
            "nom_en": p4.get("nom_en") if p4 else None,
            "nationalite": p4.get("nationalite") if p4 else None,
            "aliases": sorted({m["libelle"], *(p4["aliases"] if p4 else [])}),
            "niveau_matrice": m["niveau_matrice"],
            "poids_pays_v4": p4["poids_pays_v4"] if p4 else None,
            "niveau_pays_v4": p4["niveau_pays_v4"] if p4 else None,
            "poids_nationalite_v4": p4["poids_nationalite_v4"] if p4 else None,
            "label_v4": p4.get("label_v4") if p4 else None,
            "categorie_v4": p4.get("categorie_v4") if p4 else None,
            "sources": ["MATRICE_V1"] + (["SCORING_V4"] if p4 else []),
        }
        if p4:
            vus.add(id(p4))
        fusion.append(rec)
    for p in pays_v4:
        if id(p) in vus:
            continue
        fusion.append({
            "code": slug(p["nom_fr"], 48),
            "nom_fr": p["nom_fr"],
            "nom_en": p.get("nom_en"),
            "nationalite": p.get("nationalite"),
            "aliases": p["aliases"],
            "niveau_matrice": None,
            "poids_pays_v4": p["poids_pays_v4"],
            "niveau_pays_v4": p["niveau_pays_v4"],
            "poids_nationalite_v4": p["poids_nationalite_v4"],
            "label_v4": p.get("label_v4"),
            "categorie_v4": p.get("categorie_v4"),
            "sources": ["SCORING_V4"],
        })
    return fusion


def divergences_pays(fusion: list[dict]) -> list[dict]:
    div = []
    for p in fusion:
        nm, nv = p.get("niveau_matrice"), p.get("niveau_pays_v4")
        # V4 n'a pas INTERDIT : Inacceptable/100000 ≡ ELEVE côté score.
        nv_cmp = "ELEVE" if p.get("categorie_v4") == "INACCEPTABLE_V4" else nv
        if nm == "INTERDIT" and nv_cmp == "ELEVE":
            div.append({
                "domaine": "PAYS",
                "cle": p["nom_fr"],
                "source_a": "MATRICE_V1",
                "valeur_a": "INTERDIT",
                "source_b": "SCORING_V4",
                "valeur_b": p.get("label_v4") or "ELEVE (poids 100000)",
                "statut": "A_ARBITRER",
                "note": "La matrice pose INTERDIT ; V4 n'a pas ce niveau (score 100000 / Inacceptable).",
            })
            continue
        if nm and nv_cmp and nm != nv_cmp:
            div.append({
                "domaine": "PAYS", "cle": p["nom_fr"],
                "source_a": "MATRICE_V1", "valeur_a": nm,
                "source_b": "SCORING_V4", "valeur_b": nv_cmp,
                "statut": "A_ARBITRER",
            })
        elif nm and not nv_cmp:
            div.append({
                "domaine": "PAYS", "cle": p["nom_fr"],
                "source_a": "MATRICE_V1", "valeur_a": nm,
                "source_b": "SCORING_V4", "valeur_b": "ABSENT",
                "statut": "A_ARBITRER",
            })
        elif nv_cmp and not nm:
            div.append({
                "domaine": "PAYS", "cle": p["nom_fr"],
                "source_a": "MATRICE_V1", "valeur_a": "ABSENT",
                "source_b": "SCORING_V4", "valeur_b": nv_cmp,
                "statut": "A_ARBITRER",
            })
    return div


def main() -> int:
    if not MATRICE.exists() or not SCORING.exists():
        print("Excel manquants", file=sys.stderr)
        return 1

    wb_m = load_workbook(MATRICE, data_only=True)
    info = {
        "titre": n(wb_m["Info document"]["A6"].value),
        "date_maj": str(wb_m["Info document"]["B9"].value or ""),
        "elaboration": n(wb_m["Info document"]["B10"].value),
        "validation": n(wb_m["Info document"]["B11"].value),
    }
    client_pairs = extraire_paires(wb_m["Client"], paire_colonnes(wb_m["Client"]))
    zone = extraire_zone(wb_m["Zone géographique"])
    produits = extraire_produit(wb_m["Produit - Service - Opération"])
    canaux_m = extraire_canal_matrice(wb_m["Canal de distribution"])

    wb_v = load_workbook(SCORING, data_only=True)
    liste = extraire_liste_v4(wb_v["Liste"])
    pays_v4 = extraire_pays_v4(wb_v["Pays Eng-FR"])

    wb_vf = load_workbook(SCORING, data_only=False)
    ws_s = wb_vf["Scoring"]
    formules = {
        "H11_canal": str(ws_s["H11"].value),
        "H19_profil": str(ws_s["H19"].value),
        "H21_categorie": str(ws_s["H21"].value),
        "H23_age": str(ws_s["H23"].value),
        "H25_nationalite": str(ws_s["H25"].value),
        "H27_pays": str(ws_s["H27"].value),
        "H29_ppe": str(ws_s["H29"].value),
        "H31_procuration": str(ws_s["H31"].value),
        "H33_secteur": str(ws_s["H33"].value),
        "H36_total": str(ws_s["H36"].value),
        "C39_niveau": str(getattr(ws_s["C39"].value, "text", ws_s["C39"].value)),
        "C41_due_diligence": str(ws_s["C41"].value),
        "note_c39": (
            "V4 : score>100000 → ÉLEVÉ ; score>10000 → MOYEN ; score<10000 → FAIBLE ; "
            "score==10000 → « Veuillez compléter tous les champs ! ». "
            "Les VLOOKUP Pays/Nationalité pointent vers Liste K:L et M:N (colonnes corrompues)."
        ),
    }

    pays = reconcilier_pays(zone, pays_v4)
    div_pays = divergences_pays(pays)

    # Formes : PM privée Moyen (matrice) vs poids 100 (V4)
    div_autres = [
        {
            "domaine": "PROFIL",
            "cle": "Personne_Morale_Privée",
            "source_a": "MATRICE_V1",
            "valeur_a": "MOYEN",
            "source_b": "SCORING_V4",
            "valeur_b": "FAIBLE (poids 100, feuille Liste col. A/B)",
            "statut": "A_ARBITRER",
            "note": "La matrice classe le profil PM privée en Moyen ; le scoring V4 lui donne 100.",
        },
        {
            "domaine": "FORME_JURIDIQUE",
            "cle": "SARL / SA / SUARL / SNC / SCS / GIE / professions libérales",
            "source_a": "MATRICE_V1",
            "valeur_a": "MOYEN (toutes les PM privées listées)",
            "source_b": "SCORING_V4",
            "valeur_b": "FAIBLE (poids 100 pour SA, SARL, SUARL, SNC, SCS, SP, SAS, SCA, etc.)",
            "statut": "A_ARBITRER",
        },
        {
            "domaine": "PP_CATEGORIE",
            "cle": "Sans emploi / Entrepreneur",
            "source_a": "MATRICE_V1",
            "valeur_a": "Sans emploi=MOYEN, Entrepreneur=MOYEN",
            "source_b": "SCORING_V4",
            "valeur_b": "1104-Sans emploi=100 FAIBLE, 1105-Entrepreuneur=100 FAIBLE",
            "statut": "A_ARBITRER",
        },
        {
            "domaine": "PP_CATEGORIE",
            "cle": "Non Résident (catégorie PP)",
            "source_a": "MATRICE_V1",
            "valeur_a": "Non Résident = ÉLEVÉ (colonne Personne_Physique)",
            "source_b": "SCORING_V4",
            "valeur_b": "ABSENT de la liste Personne_Physique V4 (1101–1108)",
            "statut": "A_ARBITRER",
            "note": "V4 n'a pas la catégorie PP « Non Résident ». La résidence est un critère séparé dans la matrice.",
        },
        {
            "domaine": "FEUILLE_LISTE",
            "cle": "Colonnes Pays et Nationalité (Liste K:L et M:N)",
            "source_a": "SCORING_V4_LISTE",
            "valeur_a": "Quasi toutes les lignes à 100000 (lookup cassé)",
            "source_b": "SCORING_V4_PAYS_ENG_FR",
            "valeur_b": "Table pays réelle à utiliser",
            "statut": "INVALIDE_SOURCE",
            "note": "Ne jamais charger Liste!K:L ni Liste!M:N.",
        },
        {
            "domaine": "SEUIL_SCORE",
            "cle": "Formule C39 / H36",
            "source_a": "SCORING_V4",
            "valeur_a": ">100000 ÉLEVÉ ; >10000 MOYEN ; <10000 FAIBLE ; ==10000 message d'incomplétude",
            "source_b": "PROPOSITION_BEA_DIGITAL",
            "valeur_b": ">=100000 ÉLEVÉ ; >=10000 MOYEN ; critère évalué sinon FAIBLE ; 0 évalué = NON_CLASSE",
            "statut": "A_ARBITRER",
            "note": "La proposition BEA DIGITAL n'est pas dans les Excel ; elle évite le piège V4 score==10000.",
        },
        {
            "domaine": "PPE_NON",
            "cle": "PPE = Non",
            "source_a": "SCORING_V4",
            "valeur_a": "poids 100 (contribue au total)",
            "source_b": "MATRICE_V1",
            "valeur_b": "Faible",
            "statut": "ALIGNEE_NIVEAU",
            "note": "Niveau Faible aligné, mais V4 additionne 100. Absence de donnée ≠ Non.",
        },
        {
            "domaine": "SANCTIONS",
            "cle": "Filtrage / liste d'interdiction confirmée",
            "source_a": "MATRICE_V1",
            "valeur_a": "Oui = ÉLEVÉ (pas INTERDIT)",
            "source_b": "SCORING_V4",
            "valeur_b": "Critère absent de la fiche Scoring",
            "statut": "A_ARBITRER",
            "note": "INTERDIT vs ÉLEVÉ : décision Conformité, pas technique.",
        },
        {
            "domaine": "SECTEUR",
            "cle": "Casino / jeux de hasard / armement / pornographie / alcool",
            "source_a": "MATRICE_V1",
            "valeur_a": "INTERDIT",
            "source_b": "SCORING_V4",
            "valeur_b": "ABSENT de la liste secteurs V4 (2 secteurs à 100000 : ONG et Autres/Divers)",
            "statut": "A_ARBITRER",
        },
        {
            "domaine": "SECTEUR",
            "cle": "913-Autres ou Divers=(AUTRES (NON DÉFINIS))",
            "source_a": "SCORING_V4",
            "valeur_a": "100000 ÉLEVÉ",
            "source_b": "MATRICE_V1",
            "valeur_b": "Plusieurs « Divers … » à Faible ou Élevé selon la ligne, pas un seul code 913",
            "statut": "A_ARBITRER",
        },
        {
            "domaine": "NATIONALITE_V4",
            "cle": "Score nationalité Pays Eng-FR",
            "source_a": "SCORING_V4",
            "valeur_a": "Mauritanie=100 ; toutes les autres nationalités=10000",
            "source_b": "MATRICE_V1",
            "valeur_b": "mauritanienne=Faible ; étrangère=Moyen",
            "statut": "ALIGNEE",
            "note": "Les deux sources s'accordent : étrangère = Moyen. Ne pas utiliser Liste!M:N.",
        },
    ]

    secteurs_m = client_pairs.get("Sous secteur d'activité", [])
    v4_eleve_sect = [x["libelle"] for x in liste.get("secteur", []) if x.get("poids_v4") == 100000]
    interdit_m = [x["libelle"] for x in secteurs_m if x.get("niveau_matrice") == "INTERDIT"]
    eleve_m = [x["libelle"] for x in secteurs_m if x.get("niveau_matrice") == "ELEVE"]

    payload = {
        "meta": {
            "matrice": {
                "fichier": MATRICE.name,
                **info,
            },
            "scoring": {
                "fichier": SCORING.name,
                "version_affichee": "BANQUE EL AMANA V4.26",
            },
            "poids_echelle": {"FAIBLE": 100, "MOYEN": 10000, "ELEVE": 100000},
            "niveaux": ["FAIBLE", "MOYEN", "ELEVE", "INTERDIT"],
            "hierarchie": {"FAIBLE": 1, "MOYEN": 2, "ELEVE": 3, "INTERDIT": 4},
            "avertissement": (
                "Référentiels extraits des Excel BEA. Aucune règle n'est VALIDEE. "
                "Les divergences restent A_ARBITRER. La Conformité tranche avant activation. "
                "Une donnée absente n'est jamais FAIBLE."
            ),
        },
        "formules_v4": formules,
        "matrice_client": client_pairs,
        "pays": pays,
        "produits_matrice": produits,
        "canaux_matrice": canaux_m,
        "listes_v4": {k: v for k, v in liste.items() if k != "_ignore"},
        "liste_v4_avertissement": liste["_ignore"],
        "synthese_pays": {
            "nb_fusion": len(pays),
            "nb_matrice": len(zone),
            "nb_v4": len(pays_v4),
            "niveaux_matrice": dict(Counter(p["niveau_matrice"] for p in zone)),
            "poids_v4": dict(Counter(p["poids_pays_v4"] for p in pays_v4)),
            "interdit_matrice": [p["libelle"] for p in zone if p["niveau_matrice"] == "INTERDIT"],
            "inacceptable_v4": [p["nom_fr"] for p in pays_v4 if p.get("categorie_v4") == "INACCEPTABLE_V4"],
        },
        "synthese_secteurs": {
            "nb_matrice": len(secteurs_m),
            "niveaux_matrice": dict(Counter(x.get("niveau_matrice") for x in secteurs_m)),
            "interdit_matrice": interdit_m,
            "eleve_matrice": eleve_m,
            "nb_v4": len(liste.get("secteur", [])),
            "eleve_v4": v4_eleve_sect,
        },
        "divergences": div_pays + div_autres,
    }

    header = '''"""Référentiels extraits des Excel LBC/FT BEA — proposition, pas validation.

Sources :
- Matrice des risques LBC FT BEA V1-7-5-25.xlsx (validée Abass N'gam, MAJ 2024-12-27)
- Risque LBC FT Client V4-2026.xlsx (fiche Scoring + Liste masquée + Pays Eng-FR)

Aucune valeur n'est une règle métier validée pour BEA DIGITAL.
Les divergences sont A_ARBITRER. Ne pas charger Liste!Pays / Liste!Nationalité.
Fichier généré par backend/scripts/extract_classif_lbc.py — ne pas éditer à la main.
"""
from __future__ import annotations

import json
from pathlib import Path

REFERENTIEL = json.loads(
    (Path(__file__).with_suffix(".json")).read_text(encoding="utf-8"))
'''
    OUT_JSON.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    OUT_PY.write_text(header, encoding="utf-8")
    print(f"écrit {OUT_JSON} et {OUT_PY}")
    print(f"  pays={len(pays)} divergences={len(payload['divergences'])} "
          f"secteurs_m={len(secteurs_m)} secteurs_v4={len(liste.get('secteur', []))}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

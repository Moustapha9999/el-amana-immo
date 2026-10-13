"""Simulation du classement de masse et rapport d'un classement passé — LECTURE SEULE GARANTIE.

Connexions ``app.db.lecture_seule`` : PostgreSQL refuse toute écriture (``default_transaction_read_only``)
et SQLAlchemy bloque tout flush / commit. Aucune évaluation photographiée, classification ni audit.

Chaque classeur contient une feuille « Contrôles » (cohérence synthèse / détails) : si un contrôle échoue,
le script sort en erreur (code 2). Chaque exécution est tracée dans ``journal.jsonl`` (sans donnée
client). Les classeurs contiennent des données clients CONFIDENTIELLES : ``backend/storage/rapports/``
(hors Git, hors image Docker, non servi par HTTP), purgés après ``--conservation-jours`` jours.

  python scripts/simulation_classification.py simuler [--profil PP|PM|NON_IDENTIFIE] [--limite N]
  python scripts/simulation_classification.py rapport-tache <tache_id>
"""
from __future__ import annotations

import argparse
import asyncio
import hashlib
import json
import os
import sys
from collections import Counter, defaultdict
from datetime import datetime, timedelta
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import select
from sqlalchemy.engine import make_url

from app.core.config import get_settings
from app.db.lecture_seule import creer_moteur_lecture_seule, fabrique_lecture_seule
from app.models import (
    ClienteleClassifEvaluation,
    ClienteleClassifEvaluationLigne,
    ClienteleClassifHistorique,
    ClienteleClassifReexamen,
    ClienteleClassifTache,
    ClienteleClassification,
    ClienteleClient,
    ClienteleCompte,
    ClienteleSituation,
    User,
)
from app.services.clientele.classification import (
    SOURCES_FICHIER,
    SOURCES_MANUELLES,
    ClienteleClassificationService,
)
from app.services.clientele.scoring import (
    POIDS_ECHELLE,
    VERSION_MOTEUR,
    VERSION_REGLES,
    codes_liste_interdiction,
    evaluer,
    niveau_depuis_score_cdc,
)
from app.services.clientele.service import Ctx

DOSSIER = Path(__file__).resolve().parents[1] / "storage" / "rapports"
JOURNAL = DOSSIER / "journal.jsonl"
CONSERVATION_JOURS = 90
NIVEAUX = ("FAIBLE", "MOYEN", "ELEVE", "INTERDIT")
RANG = {"FAIBLE": 1, "MOYEN": 2, "ELEVE": 3, "INTERDIT": 4}
PAQUET = 200
AVERTISSEMENT = ("CONFIDENTIEL — données clients (racine, nom, classes de risque LBC/FT). Usage interne "
                 "Conformité / Risques. Ne pas transmettre hors BEA, ne pas déposer sur un partage non protégé. "
                 f"Conservation : {CONSERVATION_JOURS} jours sauf décision contraire de la Conformité.")

# Variantes de règles À L'ÉTUDE, recalculées sur les lignes du moteur de production. Aucune ne modifie
# le moteur : elles chiffrent l'impact avant décision de la Conformité.
SCENARIOS = {
    "S0": f"Règles de production actuelles (moteur {VERSION_MOTEUR}) : liste ORION hors score",
    "S1": "Liste ORION comptée ÉLEVÉ (ancienne règle), tous comptes",
    "S2": "Liste ORION comptée ÉLEVÉ, comptes OUVERTS seulement",
    "S3": "S0 + double comptage association / profil supprimé",
    "S4": "S0 + secteur en conflit au niveau de la Fiche (ex. Divers Commerce Général → FAIBLE)",
}


def _recalcul(lignes, scenario: str, codes_tous: list[str], codes_ouverts: list[str]) -> tuple[str | None, str, int]:
    """(niveau, statut, score) d'une variante, avec les mêmes règles de décision que ``evaluer``."""
    contrib: dict[str, int] = {lg.critere: lg.poids for lg in lignes if lg.contribue_au_score}
    par = {lg.critere: lg for lg in lignes}
    if scenario == "S1" and codes_tous:
        contrib["LISTE_INTERDICTION"] = POIDS_ECHELLE["ELEVE"]
    if scenario == "S2" and codes_ouverts:
        contrib["LISTE_INTERDICTION"] = POIDS_ECHELLE["ELEVE"]
    if scenario == "S3" and contrib.get("ASSOCIATION", 0) >= POIDS_ECHELLE["ELEVE"] \
            and contrib.get("PROFIL", 0) >= POIDS_ECHELLE["ELEVE"]:
        contrib.pop("ASSOCIATION")
    secteur = par.get("SECTEUR")
    if scenario == "S4" and secteur and secteur.divergence and "SECTEUR" in contrib and secteur.niveau_v4 in POIDS_ECHELLE:
        contrib["SECTEUR"] = POIDS_ECHELLE[secteur.niveau_v4]
    score, nb = sum(contrib.values()), len(contrib)
    blocking = any(lg.blocking_propose and lg.etat == "EVALUE" for lg in lignes)
    conflits = any(lg.statut_regle == "A_ARBITRER" and lg.etat == "EVALUE" for lg in lignes)
    if blocking:
        return "INTERDIT", "BLOQUANT", score
    if nb == 0 and not conflits:
        return None, "NON_CLASSE", score
    return niveau_depuis_score_cdc(score, nb), ("A_ARBITRER" if conflits else "EVALUE"), score


def _retenu(actuelle: tuple[str, str] | None, niveau: str | None, statut: str) -> tuple[str | None, str]:
    """Niveau qui serait RETENU et résultat, selon ``ecrire_resultat_moteur`` (classement de masse)."""
    niv_act = actuelle[0] if actuelle else None
    if statut not in ("EVALUE", "BLOQUANT") or niveau not in NIVEAUX:
        return niv_act, "NON_ADMISSIBLE"
    if actuelle and actuelle[1] in SOURCES_MANUELLES:
        return niv_act, "MANUEL_PRESERVE"
    if actuelle and actuelle[1] in SOURCES_FICHIER:
        return niv_act, "FICHIER_PRESERVE"
    if actuelle and actuelle[0] == niveau and actuelle[1] == "MOTEUR":
        return niv_act, "INCHANGE"
    return niveau, "ECRIT"


def _sens(avant: str | None, apres: str | None) -> str:
    if avant == apres:
        return "="
    if avant is None:
        return "nouveau"
    if apres is None:
        return "retiré"
    return "hausse" if RANG[apres] > RANG[avant] else "baisse"


def _feuille(wb: Workbook, titre: str, entetes: list[str], lignes: list[list]) -> None:
    ws = wb.create_sheet(titre[:31])
    ws.append(entetes)
    for c in ws[1]:
        c.font = Font(bold=True)
    for lg in lignes:
        ws.append(lg)
    ws.freeze_panes = "A2"
    for i, h in enumerate(entetes, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = max(12, min(60, len(str(h)) + 4))


def _matrice(compteur: Counter) -> list[list]:
    cles = ["(aucune)", *NIVEAUX]
    return [[a, *[compteur.get((a, b), 0) for b in cles], sum(compteur.get((a, b), 0) for b in cles)] for a in cles]


def _sha256(chemin: Path) -> str:
    return hashlib.sha256(chemin.read_bytes()).hexdigest()


def _classeur(commande: str, parametres: dict) -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.title = f"BEA DIGITAL — {commande}"
    wb.properties.subject = "CONFIDENTIEL — Conformité LBC/FT"
    wb.properties.creator = "BEA DIGITAL (script lecture seule)"
    _feuille(wb, "Lisez-moi", ["Élément", "Valeur"], [
        ["Avertissement", AVERTISSEMENT],
        ["Commande", commande],
        *[[k, v] for k, v in parametres.items()],
        ["Écritures en base", "aucune : connexion PostgreSQL en lecture seule + flush / commit bloqués"],
    ])
    return wb


def _parametres(**extra) -> dict:
    url = make_url(get_settings().database_url)
    return {
        "Date d'exécution": datetime.now().strftime("%d/%m/%Y %H:%M:%S"),
        "Base": f"{url.host}:{url.port}/{url.database}",
        "Moteur / règles": f"{VERSION_MOTEUR} / {VERSION_REGLES}",
        "Empreinte du script (SHA-256)": _sha256(Path(__file__)),
        **extra,
    }


def _enregistrer(wb: Workbook, nom: str, commande: str, parametres: dict, controles: list[list],
                 resume: dict, conservation: int) -> tuple[Path, bool]:
    ok = all(c[1] == "OK" for c in controles)
    _feuille(wb, "Contrôles", ["Contrôle", "Résultat", "Attendu", "Obtenu"], controles)
    DOSSIER.mkdir(parents=True, exist_ok=True)
    if os.name == "posix":
        os.chmod(DOSSIER, 0o700)
    chemin = DOSSIER / nom
    wb.save(chemin)
    if os.name == "posix":
        os.chmod(chemin, 0o600)
    purges = _purger(conservation)
    with JOURNAL.open("a", encoding="utf-8") as j:
        j.write(json.dumps({
            "date": datetime.now().isoformat(timespec="seconds"), "commande": commande,
            "parametres": {k: v for k, v in parametres.items() if k != "Date d'exécution"},
            "fichier": chemin.name, "sha256_classeur": _sha256(chemin), "controles_ok": ok,
            "resume": resume, "purges": purges, "conservation_jours": conservation,
        }, ensure_ascii=False) + "\n")
    return chemin, ok


def _purger(jours: int) -> list[str]:
    limite = datetime.now() - timedelta(days=jours)
    purges = []
    for f in DOSSIER.glob("*.xlsx"):
        if datetime.fromtimestamp(f.stat().st_mtime) < limite:
            f.unlink()
            purges.append(f.name)
    return purges


def _ctl(libelle: str, attendu, obtenu) -> list:
    return [libelle, "OK" if attendu == obtenu else "ÉCART", attendu, obtenu]


async def simuler(profil: str | None, limite: int | None, conservation: int) -> tuple[Path, bool]:
    moteur = creer_moteur_lecture_seule()
    Session = fabrique_lecture_seule(moteur)
    try:
        async with Session() as db:
            user = await db.scalar(select(User).where(User.is_superuser.is_(True)).limit(1))
            q = select(ClienteleClient.racine_client).order_by(ClienteleClient.racine_client)
            if profil:
                q = q.join(ClienteleSituation, ClienteleSituation.racine_client == ClienteleClient.racine_client) \
                     .where(ClienteleSituation.profil_derive == profil)
            racines = list((await db.scalars(q.limit(limite) if limite else q)).all())
            situ = {r: (p or "NON_IDENTIFIE", a or "—") for r, p, a in (await db.execute(
                select(ClienteleSituation.racine_client, ClienteleSituation.profil_derive,
                       ClienteleSituation.code_agence))).all()}
            actuelles = {r: (n, s) for r, n, s in (await db.execute(
                select(ClienteleClassification.racine_client, ClienteleClassification.niveau,
                       ClienteleClassification.source))).all()}
        ctx = Ctx(user=user, permissions={"clientele.classif.view"})

        trans = {s: Counter() for s in SCENARIOS}
        trans_pop = {s: defaultdict(Counter) for s in SCENARIOS}
        resultats = {s: Counter() for s in SCENARIOS}
        agences: dict[str, Counter] = defaultdict(Counter)
        causes = Counter()
        details, doubles = [], []
        provisoires = Counter()
        debut = datetime.now()

        for i in range(0, len(racines), PAQUET):
            async with Session() as db:
                svc = ClienteleClassificationService(db, ctx)
                paquet = racines[i:i + PAQUET]
                clients = {c.racine_client: c for c in (await db.scalars(
                    select(ClienteleClient).where(ClienteleClient.racine_client.in_(paquet)))).all()}
                comptes: dict[str, list] = defaultdict(list)
                for cp in (await db.scalars(select(ClienteleCompte).where(ClienteleCompte.racine_client.in_(paquet)))).all():
                    comptes[cp.racine_client].append(cp)
                for r in paquet:
                    ev = evaluer(await svc._collecter(clients[r], comptes[r], None))
                    codes_tous = codes_liste_interdiction(cp.liste_interdiction for cp in comptes[r])
                    codes_ouverts = codes_liste_interdiction(
                        cp.liste_interdiction for cp in comptes[r] if cp.etat_compte == "OUVERT")
                    pop, agence = situ.get(r, ("NON_IDENTIFIE", "—"))
                    act = actuelles.get(r)
                    avant = act[0] if act else None
                    niv = {}
                    for s in SCENARIOS:
                        n_s, st_s, _ = _recalcul(ev.lignes, s, codes_tous, codes_ouverts)
                        ret, res = _retenu(act, n_s, st_s)
                        niv[s] = ret
                        trans[s][(avant or "(aucune)", ret or "(aucune)")] += 1
                        trans_pop[s][pop][_sens(avant, ret)] += 1
                        resultats[s][res] += 1
                    sens = _sens(avant, niv["S0"])
                    agences[agence][sens] += 1
                    if ev.statut in ("EVALUE", "BLOQUANT") and ev.completude == "INCOMPLETE":
                        provisoires[ev.niveau_final] += 1
                    eleves = {lg.critere for lg in ev.lignes
                              if lg.contribue_au_score and (lg.poids or 0) >= POIDS_ECHELLE["ELEVE"]}
                    if {"ASSOCIATION", "PROFIL"} <= eleves:
                        doubles.append([r, clients[r].raison_sociale, pop, "Association + Profil association",
                                        ev.score_total])
                    if sens != "=":
                        cause = "Liste ORION hors score" if niv["S1"] == avant else "Autre (données ou référentiel)"
                        causes[(cause, sens)] += 1
                        details.append([r, clients[r].raison_sociale, pop, agence, avant, act[1] if act else None,
                                        niv["S0"], sens, ev.score_total, cause, ", ".join(codes_tous),
                                        ", ".join(codes_ouverts), ev.motif_genere])
            if (i // PAQUET) % 25 == 0:
                print(f"{min(i + PAQUET, len(racines))} / {len(racines)} ({datetime.now() - debut})", flush=True)
    finally:
        await moteur.dispose()

    total = len(racines)
    parametres = _parametres(**{"Population": profil or "TOUS", "Limite": limite or "aucune",
                                "Clients simulés": total, "Durée": str(datetime.now() - debut).split(".")[0],
                                **{f"Scénario {k}": v for k, v in SCENARIOS.items()}})
    wb = _classeur("simulation de masse", parametres)
    synth = [["Clients simulés", total], []]
    sommes = {}
    for s, lib in SCENARIOS.items():
        base = Counter()
        for pop_c in trans_pop[s].values():
            base.update(pop_c)
        sommes[s] = base
        synth.append([s, lib])
        synth.append(["", "baisses", base["baisse"], "hausses", base["hausse"], "inchangés", base["="],
                      "nouveaux", base["nouveau"], "retirés", base["retiré"]])
        synth.append(["", *[f"{k} {v}" for k, v in sorted(resultats[s].items())]])
    synth += [[], ["Classés provisoires (donnée essentielle manquante : PPE, filtrage…)"],
              *[[k or "(aucun)", v] for k, v in provisoires.items()]]
    _feuille(wb, "Synthèse", ["Élément", "Valeur", "", "", "", "", "", "", "", "", ""], synth)
    for s in SCENARIOS:
        _feuille(wb, f"{s} classe actuelle→simulée", ["Actuelle \\ Simulée", "(aucune)", *NIVEAUX, "Total"],
                 _matrice(trans[s]))
    _feuille(wb, "Par population", ["Scénario", "Population", "Baisses", "Hausses", "Inchangés", "Nouveaux", "Retirés"],
             [[s, p, c["baisse"], c["hausse"], c["="], c["nouveau"], c["retiré"]]
              for s in SCENARIOS for p, c in sorted(trans_pop[s].items())])
    _feuille(wb, "S0 par agence", ["Agence", "Baisses", "Hausses", "Inchangés", "Nouveaux", "Retirés"],
             [[a, c["baisse"], c["hausse"], c["="], c["nouveau"], c["retiré"]] for a, c in sorted(agences.items())])
    _feuille(wb, "S0 causes", ["Cause", "Sens", "Clients"], [[c, s, n] for (c, s), n in causes.most_common()])
    _feuille(wb, "S0 détail changements",
             ["Racine", "Nom", "Population", "Agence", "Classe actuelle", "Source actuelle", "Classe simulée",
              "Sens", "Score simulé", "Cause", "Codes ORION (tous comptes)", "Codes ORION (comptes ouverts)",
              "Motif moteur simulé"], details)
    _feuille(wb, "Doubles comptages", ["Racine", "Nom", "Population", "Notion comptée deux fois", "Score"], doubles)

    changes_s0 = total - sommes["S0"]["="]
    agences_tot = Counter()
    for c in agences.values():
        agences_tot.update(c)
    pop_detail = Counter((d[2], d[7]) for d in details)
    controles = [_ctl("Racines distinctes simulées", total, len(set(racines)))]
    for s in SCENARIOS:
        controles += [
            _ctl(f"{s} : total de la matrice actuelle→simulée = clients", total, sum(trans[s].values())),
            _ctl(f"{s} : somme par population = clients", total, sum(sommes[s].values())),
            _ctl(f"{s} : somme des résultats moteur = clients", total, sum(resultats[s].values())),
            _ctl(f"{s} : hors diagonale de la matrice = changements",
                 total - sommes[s]["="], sum(v for (a, b), v in trans[s].items() if a != b)),
        ]
    controles += [
        _ctl("S0 : somme par agence = clients", total, sum(agences_tot.values())),
        _ctl("S0 : baisses par agence = baisses de synthèse", sommes["S0"]["baisse"], agences_tot["baisse"]),
        _ctl("S0 : hausses par agence = hausses de synthèse", sommes["S0"]["hausse"], agences_tot["hausse"]),
        _ctl("S0 : lignes de détail = changements de synthèse", changes_s0, len(details)),
        _ctl("S0 : somme des causes = changements", changes_s0, sum(causes.values())),
        *[_ctl(f"S0 : détail {p} / {s} = synthèse par population", trans_pop["S0"][p][s], pop_detail[(p, s)])
          for p in sorted(trans_pop["S0"]) for s in ("baisse", "hausse", "nouveau", "retiré")],
        _ctl("Provisoires ≤ clients", True, sum(provisoires.values()) <= total),
    ]
    resume = {s: {k: v for k, v in sommes[s].items()} for s in SCENARIOS}
    nom = f"simulation-classification-{datetime.now():%Y%m%d-%H%M%S}.xlsx"
    return _enregistrer(wb, nom, "simuler", parametres, controles, resume, conservation)


async def rapport_tache(tache_id: str, conservation: int) -> tuple[Path, bool]:
    moteur = creer_moteur_lecture_seule()
    try:
        async with fabrique_lecture_seule(moteur)() as db:
            t = await db.get(ClienteleClassifTache, tache_id)
            if not t:
                raise SystemExit("Tâche introuvable")
            debut, fin = t.debut or t.created_at, (t.fin or datetime.now(t.created_at.tzinfo)) + timedelta(seconds=5)
            histo = list((await db.scalars(select(ClienteleClassifHistorique).where(
                ClienteleClassifHistorique.created_at.between(debut, fin),
                ClienteleClassifHistorique.motif_classement == t.motif,
                ClienteleClassifHistorique.source == "MOTEUR").order_by(ClienteleClassifHistorique.racine_client))).all())
            racines = [h.racine_client for h in histo]
            noms = dict((await db.execute(select(ClienteleClient.racine_client, ClienteleClient.raison_sociale)
                                          .where(ClienteleClient.racine_client.in_(racines)))).all())
            agences = dict((await db.execute(select(ClienteleSituation.racine_client, ClienteleSituation.code_agence)
                                             .where(ClienteleSituation.racine_client.in_(racines)))).all())
            courantes = {c.racine_client: c for c in (await db.scalars(
                select(ClienteleClassification).where(ClienteleClassification.racine_client.in_(racines)))).all()}
            lignes = []
            stats = Counter()
            for h in histo:
                r = h.racine_client
                precedent = await db.scalar(select(ClienteleClassifHistorique).where(
                    ClienteleClassifHistorique.racine_client == r, ClienteleClassifHistorique.created_at < debut)
                    .order_by(ClienteleClassifHistorique.created_at.desc()).limit(1))
                posterieurs = await db.scalar(select(ClienteleClassifHistorique.id).where(
                    ClienteleClassifHistorique.racine_client == r, ClienteleClassifHistorique.created_at > fin).limit(1))
                reex = await db.scalar(select(ClienteleClassifReexamen.id).where(
                    ClienteleClassifReexamen.racine_client == r, ClienteleClassifReexamen.statut == "OUVERT").limit(1))
                ev_old = await db.scalar(select(ClienteleClassifEvaluation).where(
                    ClienteleClassifEvaluation.racine_client == r, ClienteleClassifEvaluation.created_at < debut)
                    .order_by(ClienteleClassifEvaluation.created_at.desc()).limit(1))
                ev_new = await db.get(ClienteleClassifEvaluation, h.evaluation_id) if h.evaluation_id else None

                async def eleves(ev) -> set[str]:
                    if not ev:
                        return set()
                    return set((await db.scalars(select(ClienteleClassifEvaluationLigne.critere).where(
                        ClienteleClassifEvaluationLigne.evaluation_id == ev.id,
                        ClienteleClassifEvaluationLigne.contribue_au_score.is_(True),
                        ClienteleClassifEvaluationLigne.poids >= POIDS_ECHELLE["ELEVE"]))).all())

                e_old, e_new = await eleves(ev_old), await eleves(ev_new)
                perdus = sorted(e_old - e_new)
                regle = ("Liste ORION suspendue (hors score)" if perdus == ["LISTE_INTERDICTION"]
                         else ("Critères ÉLEVÉ perdus : " + ", ".join(perdus)) if perdus else "Données modifiées")
                c = courantes.get(r)
                obstacles = []
                if not c or c.niveau != h.nouvelle_classe or c.source != "MOTEUR":
                    obstacles.append("classification courante modifiée depuis")
                if posterieurs:
                    obstacles.append("changement postérieur dans l'historique")
                if reex:
                    obstacles.append("réexamen ouvert")
                if not precedent or precedent.nouvelle_classe != h.ancienne_classe:
                    obstacles.append("classe précédente introuvable dans l'historique")
                elif precedent.source in SOURCES_MANUELLES:
                    obstacles.append("classe précédente manuelle")
                restaurable = "Oui" if not obstacles else "À examiner : " + " ; ".join(obstacles)
                stats[("restaurable", restaurable == "Oui")] += 1
                stats[("regle", regle)] += 1
                stats[("transition", f"{h.ancienne_classe} → {h.nouvelle_classe}")] += 1
                lignes.append([
                    r, noms.get(r), agences.get(r), h.ancienne_classe, h.nouvelle_classe,
                    h.created_at.strftime("%d/%m/%Y %H:%M:%S"), regle,
                    ev_old.version_moteur if ev_old else None, ev_old.score_total if ev_old else None,
                    ", ".join(sorted(e_old)), ev_new.score_total if ev_new else None, ", ".join(sorted(e_new)),
                    precedent.source if precedent else None,
                    precedent.created_at.strftime("%d/%m/%Y %H:%M") if precedent else None,
                    (precedent.motif_risque or "")[:500] if precedent else None,
                    (h.motif_risque or "")[:500], restaurable,
                ])
            preserves = list((await db.execute(
                select(ClienteleClassification.racine_client, ClienteleClassification.niveau,
                       ClienteleClassification.source, ClienteleClassification.classifie_le)
                .join(ClienteleSituation, ClienteleSituation.racine_client == ClienteleClassification.racine_client)
                .where(ClienteleSituation.profil_derive == (t.profil or ClienteleSituation.profil_derive),
                       ClienteleClassification.source.in_(SOURCES_MANUELLES + SOURCES_FICHIER)))).all())
            reex_crees = list((await db.execute(
                select(ClienteleClassifReexamen.racine_client, ClienteleClassifReexamen.niveau_actuel,
                       ClienteleClassifReexamen.niveau_suggere, ClienteleClassifReexamen.statut)
                .where(ClienteleClassifReexamen.created_at.between(debut, fin)))).all())
    finally:
        await moteur.dispose()

    parametres = _parametres(**{"Tâche": str(t.id), "Motif de la tâche": t.motif, "Population": t.profil or "TOUS"})
    wb = _classeur("rapport de tâche", parametres)
    synth = [["Début", debut.strftime("%d/%m/%Y %H:%M:%S")], ["Fin", t.fin.strftime("%d/%m/%Y %H:%M:%S") if t.fin else None],
             ["Bilan", str(t.bilan)], ["Classes modifiées", len(histo)], []]
    synth += [[k[1], v] for k, v in stats.items() if k[0] == "transition"] + [[]]
    synth += [[k[1], v] for k, v in stats.items() if k[0] == "regle"] + [[]]
    synth += [["Restaurables sans obstacle", stats[("restaurable", True)]],
              ["À examiner", stats[("restaurable", False)]], [],
              ["Décisions manuelles / fichier de la population (non touchées)", len(preserves)],
              ["Réexamens ouverts pendant la tâche", len(reex_crees)]]
    _feuille(wb, "Synthèse", ["Élément", "Valeur"], synth)
    _feuille(wb, "Clients modifiés", [
        "Racine", "Nom", "Agence", "Ancienne classe", "Nouvelle classe", "Date du changement", "Règle responsable",
        "Moteur précédent", "Score précédent", "Critères ÉLEVÉ précédents", "Nouveau score", "Critères ÉLEVÉ nouveaux",
        "Source de la classe précédente", "Classe précédente posée le", "Motif précédent", "Nouveau motif",
        "Restauration depuis l'historique"], lignes)
    _feuille(wb, "Décisions préservées", ["Racine", "Niveau", "Source", "Classé le"],
             [[r, n, s, d.strftime("%d/%m/%Y %H:%M") if d else None] for r, n, s, d in preserves])
    _feuille(wb, "Réexamens de la tâche", ["Racine", "Niveau actuel", "Niveau suggéré", "Statut"],
             [list(x) for x in reex_crees])
    ecrits = (t.bilan or {}).get("ECRIT", 0)
    controles = [
        _ctl("Lignes de détail = classes écrites par la tâche (bilan ECRIT)", ecrits, len(lignes)),
        _ctl("Racines distinctes = lignes", len(lignes), len({x[0] for x in lignes})),
        _ctl("Somme des transitions = lignes", len(lignes), sum(v for k, v in stats.items() if k[0] == "transition")),
        _ctl("Somme des règles = lignes", len(lignes), sum(v for k, v in stats.items() if k[0] == "regle")),
        _ctl("Restaurables + à examiner = lignes", len(lignes),
             stats[("restaurable", True)] + stats[("restaurable", False)]),
    ]
    resume = {"ecrits": ecrits, "restaurables": stats[("restaurable", True)], "a_examiner": stats[("restaurable", False)]}
    nom = f"rapport-tache-{str(t.id)[:8]}-{datetime.now():%Y%m%d-%H%M%S}.xlsx"
    return _enregistrer(wb, nom, "rapport-tache", parametres, controles, resume, conservation)


def main() -> None:
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--conservation-jours", type=int, default=CONSERVATION_JOURS)
    sp = p.add_subparsers(dest="cmd", required=True)
    s = sp.add_parser("simuler")
    s.add_argument("--profil", choices=("PP", "PM", "NON_IDENTIFIE"))
    s.add_argument("--limite", type=int)
    r = sp.add_parser("rapport-tache")
    r.add_argument("tache_id")
    a = p.parse_args()
    coro = (simuler(a.profil, a.limite, a.conservation_jours) if a.cmd == "simuler"
            else rapport_tache(a.tache_id, a.conservation_jours))
    chemin, ok = asyncio.run(coro)
    print(f"Classeur : {chemin}")
    print("Contrôles : OK" if ok else "Contrôles : ÉCART — voir la feuille « Contrôles »")
    sys.exit(0 if ok else 2)


if __name__ == "__main__":
    main()

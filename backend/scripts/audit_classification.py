"""Audit de la classification LBC/FT avant paramétrage — LECTURE SEULE GARANTIE.

Connexion ``app.db.lecture_seule`` (PostgreSQL ``default_transaction_read_only`` + flush / commit bloqués) :
aucune classification, évaluation, règle ni historique n'est modifié et le moteur n'est pas relancé.
Les impacts sont chiffrés à partir des évaluations déjà photographiées.

Classeur CONFIDENTIEL dans ``backend/storage/rapports/`` (hors Git, hors image, non servi par HTTP) avec
une feuille « Contrôles » ; un écart → code de sortie 2. Exécution tracée dans ``journal.jsonl``.

  python scripts/audit_classification.py --matrice /chemin/Matrice.xlsx [--fiche /chemin/V4.xlsx]
"""
from __future__ import annotations

import argparse
import asyncio
import difflib
import re
import sys
from collections import Counter, defaultdict
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(Path(__file__).resolve().parent))

from openpyxl import load_workbook
from sqlalchemy import text

from app.data.clientele_classif_matrice import (
    ENTREPRISE_RECENTE_JOURS,
    POIDS_ECHELLE,
    VERSION_MOTEUR,
    VERSION_REGLES,
)
from app.data.clientele_import_hebdo import listes_observees
from app.db.lecture_seule import creer_moteur_lecture_seule, fabrique_lecture_seule
from app.services.clientele.classification import autorisation_classement_masse
from app.services.clientele.lots import CORRESPONDANCES_OFFICIELLES, normaliser_niveau
from app.services.clientele.scoring import (
    REFERENTIEL,
    _chercher_pays,
    _est_mauritanie,
    _index,
    niveau_depuis_score_cdc,
    normaliser,
    slug_cle,
)
from simulation_classification import CONSERVATION_JOURS, _classeur, _ctl, _enregistrer, _feuille, _parametres

AUJOURDHUI = date.today().strftime("%d/%m/%Y")
NIVEAUX_PAYS = {"FAIBLE": "FAIBLE", "MOYEN": "MOYEN", "ELEVE": "ELEVE", "INTERDIT": "INTERDIT"}
MAX_NIF, MAX_NNI = 9, 10
STATUT_DIV = {"A_ARBITRER": "À arbitrer", "ALIGNEE": "À analyser", "INVALIDE_SOURCE": "À analyser",
              "VALIDEE": "Validé", "REJETEE": "Rejeté"}


async def _q(db, sql: str, **p) -> list:
    return list((await db.execute(text(sql), p)).mappings().all())


def _niv(score: int, nb: int) -> str | None:
    return niveau_depuis_score_cdc(score, nb)


# --------------------------------------------------------------------------- 0. état général
async def etat_general(db) -> tuple[list[list], dict]:
    v = {}
    v["clients"] = (await _q(db, "SELECT count(*) n FROM clientele_clients"))[0]["n"]
    v["comptes"] = (await _q(db, "SELECT count(*) n FROM clientele_comptes"))[0]["n"]
    cls = await _q(db, """SELECT coalesce(mode_decision, source) mode, source, niveau, version_moteur, count(*) n
                          FROM clientele_classifications GROUP BY 1,2,3,4 ORDER BY 1,3,4""")
    v["classes"] = sum(r["n"] for r in cls)
    v["manuels"] = sum(r["n"] for r in cls if r["mode"] == "MANUEL")
    v["moteur"] = sum(r["n"] for r in cls if r["mode"] == "MOTEUR")
    v["non_classes"] = v["clients"] - v["classes"]
    reex = await _q(db, "SELECT statut, count(*) n FROM clientele_classif_reexamens GROUP BY 1")
    v["reexamens_ouverts"] = sum(r["n"] for r in reex if r["statut"] == "OUVERT")
    div = await _q(db, "SELECT statut, count(*) n FROM clientele_classif_divergences GROUP BY 1")
    v["divergences_a_arbitrer"] = sum(r["n"] for r in div if r["statut"] == "A_ARBITRER")
    val = await _q(db, "SELECT statut, count(*) n, count(*) FILTER (WHERE actif) actifs FROM clientele_classif_valeurs GROUP BY 1")
    v["valeurs_a_arbitrer"] = sum(r["n"] for r in val if r["statut"] == "A_ARBITRER")
    v["valeurs_actives"] = sum(r["actifs"] for r in val)
    vers = await _q(db, "SELECT numero, libelle, statut, mode, (SELECT count(*) FROM clientele_classif_regles r "
                        "WHERE r.version_id = v.id) nb FROM clientele_classif_versions v ORDER BY numero")
    lots = await _q(db, "SELECT type, statut, count(*) n FROM clientele_classif_lots GROUP BY 1,2")
    ev = await _q(db, "SELECT count(*) n FROM clientele_classif_evaluations")
    histo = await _q(db, "SELECT source, count(*) n FROM clientele_classif_historique GROUP BY 1 ORDER BY 2 DESC")
    verrou = autorisation_classement_masse()
    v["verrou_masse_actif"] = bool(verrou.get("actif"))
    lignes = [
        ["Date de l'audit", AUJOURDHUI],
        ["Moteur / règles du code", f"{VERSION_MOTEUR} / {VERSION_REGLES}"],
        ["Clients (racines)", v["clients"]], ["Comptes", v["comptes"]],
        ["Classés", v["classes"]], ["  dont moteur", v["moteur"]], ["  dont décision manuelle", v["manuels"]],
        ["Non classés (aucune classification retenue)", v["non_classes"]],
        ["Réexamens ouverts", v["reexamens_ouverts"]],
        ["Divergences Matrice / Fiche à arbitrer", v["divergences_a_arbitrer"]],
        ["Valeurs du référentiel maître à arbitrer", v["valeurs_a_arbitrer"]],
        ["Valeurs du référentiel maître ACTIVES", v["valeurs_actives"]],
        ["Évaluations photographiées", ev[0]["n"]],
        ["Classement de masse", "OUVERT" if v["verrou_masse_actif"] else "VERROUILLÉ côté serveur (aucune référence de PV)"],
        [], ["Classifications retenues", "Mode", "Source", "Niveau", "Version moteur", "Nombre"],
        *[["", r["mode"], r["source"], r["niveau"], r["version_moteur"], r["n"]] for r in cls],
        [], ["Versions du moteur à règles historique", "Numéro", "Libellé", "Statut", "Mode", "Règles"],
        *[["", r["numero"], r["libelle"], r["statut"], r["mode"], r["nb"]] for r in vers],
        [], ["Lots", "Type", "Statut", "Nombre"], *[["", r["type"], r["statut"], r["n"]] for r in lots],
        [], ["Historique des classes", "Source", "Lignes"], *[["", r["source"], r["n"]] for r in histo],
        [], ["Référentiel maître (valeurs)", "Statut", "Lignes", "Actives"],
        *[["", r["statut"], r["n"], r["actifs"]] for r in val],
    ]
    return lignes, v


# --------------------------------------------------------------------------- évaluations de référence
async def evaluations(db) -> dict[str, dict]:
    """Évaluation de référence par client : celle de la classification retenue, sinon la plus récente."""
    rows = await _q(db, """
        WITH ref AS (
            SELECT c.racine_client, c.evaluation_id id, true classe, c.niveau niveau_retenu,
                   coalesce(c.mode_decision, c.source) mode
            FROM clientele_classifications c
            UNION ALL
            SELECT * FROM (
                SELECT DISTINCT ON (e.racine_client) e.racine_client, e.id, false, NULL::varchar, NULL::varchar
                FROM clientele_classif_evaluations e
                WHERE NOT EXISTS (SELECT 1 FROM clientele_classifications c WHERE c.racine_client = e.racine_client)
                ORDER BY e.racine_client, e.created_at DESC) x
        )
        SELECT ref.*, e.score_total, e.nb_evalues, e.statut, e.niveau_final, e.version_moteur, e.coherence,
               e.created_at, s.profil_derive, s.nom_client, s.code_agence
        FROM ref LEFT JOIN clientele_classif_evaluations e ON e.id = ref.id
        LEFT JOIN clientele_situation s ON s.racine_client = ref.racine_client""")
    ev = {r["racine_client"]: dict(r, lignes={}) for r in rows}
    ids = {r["id"]: r["racine_client"] for r in rows if r["id"]}
    lignes = await _q(db, """
        SELECT l.evaluation_id, l.critere, l.etat, l.valeur, l.poids, l.niveau_matrice, l.niveau_v4,
               l.divergence, l.statut_regle, l.contribue_au_score, l.source_donnee, l.blocking_propose
        FROM clientele_classif_evaluation_lignes l
        WHERE l.evaluation_id = ANY(:ids)""", ids=list(ids))
    for lg in lignes:
        ev[ids[lg["evaluation_id"]]]["lignes"][lg["critere"]] = lg
    return ev


def _impact(e: dict, ajout: int, critere: str) -> str | None:
    """Niveau obtenu si le critère ``critere`` (aujourd'hui exclu) comptait ``ajout`` points."""
    if e.get("score_total") is None:
        return None
    if e["statut"] == "BLOQUANT":
        return "INTERDIT"
    autres_conflits = [c for c, lg in e["lignes"].items()
                       if c != critere and lg["statut_regle"] == "A_ARBITRER" and lg["etat"] == "EVALUE"]
    if autres_conflits:
        return "A_ARBITRER"
    return _niv(e["score_total"] + ajout, e["nb_evalues"] + (1 if ajout > 0 else 0))


# --------------------------------------------------------------------------- 2. date d'ouverture entreprise
async def date_entreprise(db, ev: dict) -> tuple[list[list], list[list], dict]:
    rows = await _q(db, """
        SELECT s.racine_client, s.nom_client, s.profil_derive, cl.categorie_juridique, cl.situation_juridique,
               nullif(trim(cl.rcs), '') rcs, s.date_ouverture date_premier_compte,
               (SELECT max(l.date_ouverture_entreprise) FROM clientele_classif_lot_lignes l
                 WHERE l.racine_client = s.racine_client AND l.statut = 'OK') date_lot,
               (SELECT max(m.date_creation) FROM eer_dossiers d JOIN eer_parties_morales m ON m.partie_id = d.client_partie_id
                 WHERE d.racine_client = s.racine_client AND d.deleted_at IS NULL) date_eer
        FROM clientele_situation s JOIN clientele_clients cl ON cl.racine_client = s.racine_client
        WHERE s.profil_derive <> 'PP'""")
    det, stats = [], Counter()
    for r in rows:
        e = ev.get(r["racine_client"], {})
        lg = e.get("lignes", {}).get("ENTREPRISE_RECENTE") or {}
        source = "SAISIE/FICHIER (lot)" if r["date_lot"] else ("EER" if r["date_eer"] else None)
        etat = lg.get("etat") or "SANS_EVALUATION"
        if source:
            statut = "DATE DISPONIBLE"
        elif r["profil_derive"] == "NON_IDENTIFIE":
            statut = "APPLICABILITÉ À VÉRIFIER (profil non identifié)"
        else:
            statut = "DONNÉE MANQUANTE"
        stats[(r["profil_derive"], statut)] += 1
        stats[("etat", etat)] += 1
        avant = e.get("niveau_retenu") or (e.get("niveau_final") if e.get("statut") in ("EVALUE", "BLOQUANT") else None)
        si_recente = _impact(e, POIDS_ECHELLE["MOYEN"], "ENTREPRISE_RECENTE") if etat != "EVALUE" else None
        si_ancienne = _impact(e, POIDS_ECHELLE["FAIBLE"], "ENTREPRISE_RECENTE") if etat != "EVALUE" else None
        if r["profil_derive"] == "PM" and etat != "EVALUE":
            stats[("recente", avant, si_recente)] += 1
            stats[("ancienne", avant, si_ancienne)] += 1
        det.append([r["racine_client"], r["nom_client"], r["profil_derive"], r["categorie_juridique"],
                    r["situation_juridique"], "Oui" if r["rcs"] else "Non", r["date_lot"], r["date_eer"],
                    source or "—", statut, etat, lg.get("source_donnee"),
                    r["date_premier_compte"], avant, si_recente, si_ancienne])
    synth = [
        ["Règle", f"Entreprise ≤ {ENTREPRISE_RECENTE_JOURS} jours = 10 000 (MOYEN) ; au-delà = 100 (FAIBLE). "
                  "Personnes morales uniquement."],
        ["Sources lues par le moteur (ordre)", "1. date saisie par l'analyste sur une ligne de lot ; "
                                               "2. colonne « Date ouverture » du fichier ; 3. date de création PM de l'EER."],
        ["Date d'ouverture de compte ORION", "JAMAIS utilisée par le moteur (vérifié dans "
                                             "ClienteleClassificationService._collecter). Affichée ici pour information."],
        ["Date RCCM", "Aucun champ « date RCCM » dans ORION État des comptes : seul le numéro RCCM (rcs) existe."],
        [], ["Population", "Statut", "Clients"],
        *[[k[0], k[1], n] for k, n in sorted(stats.items(), key=str) if len(k) == 2 and k[0] in ("PM", "NON_IDENTIFIE")],
        [], ["État du critère dans l'évaluation de référence", "", "Clients"],
        *[["", k[1], n] for k, n in stats.items() if len(k) == 2 and k[0] == "etat"],
        [], ["Impact si une date était fournie (PM sans critère évalué)", "Niveau actuel → niveau obtenu", "Clients"],
        *[[("Entreprise récente (+10 000)" if k[0] == "recente" else "Entreprise ancienne (+100)"),
           f"{k[1] or '(aucun)'} → {k[2] or '(aucun)'}", n]
          for k, n in sorted(((k, n) for k, n in stats.items() if k[0] in ("recente", "ancienne")), key=str)],
    ]
    deux = {k: n for k, n in stats.items() if len(k) == 2}
    resume = {"pm_sans_date": sum(n for (p, s), n in deux.items() if p == "PM" and s == "DONNÉE MANQUANTE"),
              "pm_total": sum(n for (p, s), n in deux.items() if p == "PM"),
              "avec_date": sum(n for (p, s), n in deux.items() if p in ("PM", "NON_IDENTIFIE") and s == "DATE DISPONIBLE")}
    return synth, det, resume


# --------------------------------------------------------------------------- 3. statut compte / gel
async def statuts_comptes(db) -> tuple[list[list], list[list], dict]:
    rows = await _q(db, """
        SELECT racine_client, count(*) nb, count(*) FILTER (WHERE etat_compte = 'OUVERT') ouverts,
               count(*) FILTER (WHERE etat_compte = 'FERME') fermes
        FROM clientele_comptes GROUP BY 1""")
    etats = await _q(db, "SELECT etat_compte, count(*) n FROM clientele_comptes GROUP BY 1")
    lot = await _q(db, """SELECT donnees->>'STATUT COMPTE' sc, donnees->>'STATUT GEL' sg, count(*) n
                          FROM clientele_classif_lot_lignes GROUP BY 1,2""")
    c = Counter()
    mixtes = []
    for r in rows:
        cat = "TOUS OUVERTS" if r["fermes"] == 0 else ("TOUS FERMÉS" if r["ouverts"] == 0 else "MIXTE")
        c[cat] += 1
        c[("multi", r["nb"] > 1)] += 1
        if cat == "MIXTE":
            mixtes.append([r["racine_client"], r["nb"], r["ouverts"], r["fermes"],
                           "OUVERT (vue clientele_situation : ≥ 1 compte ouvert)", "Proposition : Ouvert — à valider"])
    obs = listes_observees()
    synth = [
        ["Champ « Statut du compte » (ORION etat_compte)", ", ".join(f"{r['etat_compte']} = {r['n']}" for r in etats)],
        ["Champ « Statut de gel »", "AUCUNE source : ORION État des comptes ne contient pas de statut de gel. "
                                     "Le fichier hebdomadaire a une colonne « Statut Gel » sans compte ni RIB."],
        ["Valeurs « Statut compte » observées (fichier hebdomadaire)", ", ".join(obs.get("statuts_compte") or [])],
        ["Valeurs « Statut gel » observées (fichier hebdomadaire)", ", ".join(obs.get("statuts_gel") or [])],
        ["Hors liste cible", "« Bloqué » (statut compte) et « Cloturé » (statut gel) ne sont ni Ouvert/Fermé ni "
                             "Actif/Inactif : correspondance à arbitrer, jamais assimilée."],
        ["Lignes de lot (statut compte | statut gel)",
         "; ".join(f"{r['sc'] or '—'} | {r['sg'] or '—'} = {r['n']}" for r in lot)],
        ["Utilisation par le moteur", "Aucune (ni statut compte ni gel n'entrent dans le score). La liste "
                                      "d'interdiction est lue sur TOUS les comptes, ouverts ou fermés."],
        ["Agrégation existante", "Vue clientele_situation : client CLOTURE si tous ses comptes sont FERME, "
                                 "sinon OUVERT. Règle technique, non approuvée formellement."],
        [], ["Clients par composition de comptes", "Clients"],
        ["Tous comptes ouverts", c["TOUS OUVERTS"]], ["Tous comptes fermés", c["TOUS FERMÉS"]],
        ["Mixte (ouverts + fermés)", c["MIXTE"]], ["Clients multi-comptes", c[("multi", True)]],
        ["Clients mono-compte", c[("multi", False)]],
    ]
    return synth, mixtes, {"mixtes": c["MIXTE"], "tous_fermes": c["TOUS FERMÉS"], "multi": c[("multi", True)]}


# --------------------------------------------------------------------------- 4. catégories juridiques
def _rapprocher_forme(brut: str) -> tuple[str, dict | None, list[str]]:
    sv, nv = slug_cle(brut), normaliser(brut)
    formes = _index()["formes"]
    exact = [f for f in formes if normaliser(f.get("libelle")) == nv or slug_cle(f.get("libelle")) == sv]
    if exact:
        return "EXACTE", _fusion(exact), []
    partiel = [f for f in formes if sv and slug_cle(f.get("libelle"))
               and (sv in slug_cle(f.get("libelle")) or slug_cle(f.get("libelle")) in sv)]
    libs = sorted({f.get("libelle") for f in partiel})
    if len(libs) == 1:
        return "PROPOSÉE", _fusion(partiel), libs
    if libs:
        return "AMBIGUË", None, libs
    return "NON RECONNUE", None, []


def _fusion(hits: list[dict]) -> dict:
    return {
        "libelle": next((h.get("libelle") for h in hits if h.get("libelle")), None),
        "profil": next((h.get("profil_matrice") or h.get("profil_v4") for h in hits
                        if h.get("profil_matrice") or h.get("profil_v4")), None),
        "nm": next((h.get("niveau_matrice") for h in hits if h.get("niveau_matrice")), None),
        "pv": next((h.get("poids_v4") for h in hits if h.get("poids_v4") is not None), None),
        "nv": next((h.get("niveau_v4") for h in hits if h.get("niveau_v4")), None),
    }


async def categories(db) -> tuple[list[list], dict]:
    rows = await _q(db, """
        SELECT champ, valeur, profil, count(*) n FROM (
            SELECT 'categorie_juridique' champ, cl.categorie_juridique valeur, s.profil_derive profil
            FROM clientele_clients cl JOIN clientele_situation s USING (racine_client)
            UNION ALL
            SELECT 'situation_juridique', cl.situation_juridique, s.profil_derive
            FROM clientele_clients cl JOIN clientele_situation s USING (racine_client)) x
        GROUP BY 1,2,3""")
    agg: dict[tuple, Counter] = defaultdict(Counter)
    for r in rows:
        agg[(r["champ"], r["valeur"])][r["profil"]] += r["n"]
    for v in listes_observees().get("categories_juridiques") or []:
        agg.setdefault(("fichier_hebdomadaire", v), Counter())
    slugs: dict[str, set] = defaultdict(set)
    for (champ, v) in agg:
        if v:
            slugs[slug_cle(v)].add(v)
    table, stats = [], Counter()
    for (champ, v), prof in sorted(agg.items(), key=lambda kv: (kv[0][0], -sum(kv[1].values()))):
        total = sum(prof.values())
        if not v or not str(v).strip():
            table.append([champ, v, "", "", "", "", "", "", total, ", ".join(f"{k} {n}" for k, n in prof.items()),
                          "VIDE", "Valeur absente : aucun niveau attribué.", "Donnée manquante", AUJOURDHUI])
            stats["VIDE"] += total if champ != "fichier_hebdomadaire" else 0
            continue
        statut, f, cand = _rapprocher_forme(v)
        nm, nv = (f or {}).get("nm"), (f or {}).get("nv")
        if f and nm and nv and nm != nv:
            niveau = f"CONFLIT : Matrice {nm} / Fiche {nv}"
            just = "Les deux sources divergent : exclu de la somme jusqu'à arbitrage Conformité."
        elif f and (nm or nv):
            niveau = nm or nv
            just = f"{'Libellé identique' if statut == 'EXACTE' else 'Rapprochement partiel unique'} au référentiel ({'Matrice' if nm else 'Fiche'})."
        elif f:
            niveau = "—"
            just = "Rapproché d'un libellé sans niveau."
        else:
            niveau = "—"
            just = ("Plusieurs libellés candidats : " + " | ".join(cand)) if cand else "Aucun libellé du référentiel."
        doublon = len(slugs[slug_cle(v)]) > 1
        validation = {"EXACTE": "Correspondance exacte — validation Conformité à confirmer",
                      "PROPOSÉE": "Proposée — à valider", "AMBIGUË": "Ambiguë — à arbitrer",
                      "NON RECONNUE": "Non reconnue — à rattacher manuellement"}[statut]
        table.append([champ, v, normaliser(v), (f or {}).get("libelle") or "", (f or {}).get("profil") or "",
                      nm or "", (f"{nv} ({(f or {}).get('pv')})" if nv else ""), niveau, total,
                      ", ".join(f"{k} {n}" for k, n in prof.items()),
                      statut + (" — DOUBLON après normalisation" if doublon else ""), just, validation, AUJOURDHUI])
        if champ != "fichier_hebdomadaire":
            stats[statut] += total
            stats[("valeurs", statut)] += 1
    return table, stats


# --------------------------------------------------------------------------- 5. NIF / NNI
def _controle_id(v: str | None, maxi: int) -> list[str]:
    if v is None or v == "":
        return ["VIDE"]
    a = []
    if v != v.strip():
        a.append("ESPACES EN BORDURE")
    s = v.strip()
    if not s:
        return ["VIDE (espaces seulement)"]
    if re.search(r"[\s\-./_]", s):
        a.append("SÉPARATEURS INTERNES")
    if re.search(r"[^\d\s\-./_]", s):
        a.append("CARACTÈRES NON NUMÉRIQUES")
    chiffres = re.sub(r"\D", "", s)
    if len(s) > maxi:
        a.append(f"TROP LONG ({len(s)} > {maxi})")
    if a and chiffres and len(chiffres) <= maxi and not re.search(r"[^\d\s\-./_]", s):
        a.append("serait conforme après retrait des séparateurs (NON appliqué)")
    return a or ["CONFORME"]


async def identifiants(db) -> tuple[list[list], list[list], dict]:
    rows = await _q(db, """SELECT cl.racine_client, s.profil_derive, cl.nni, cl.nif, cl.type_identifiant
                          FROM clientele_clients cl JOIN clientele_situation s USING (racine_client)""")
    c, det = Counter(), []
    obligatoire = {("PP", "NNI"), ("PM", "NIF")}
    for r in rows:
        for champ, maxi in (("NNI", MAX_NNI), ("NIF", MAX_NIF)):
            v = r[champ.lower()]
            a = _controle_id(v, maxi)
            statut = a[0] if a == ["CONFORME"] or a[0].startswith("VIDE") else "NON CONFORME"
            if statut.startswith("VIDE"):
                statut = "MANQUANT (obligatoire proposé)" if (r["profil_derive"], champ) in obligatoire else "VIDE (facultatif proposé)"
            c[(r["profil_derive"], champ, statut)] += 1
            if statut == "CONFORME" and v.startswith("0"):
                c[(champ, "zeros")] += 1
            if statut == "NON CONFORME":
                for x in a:
                    c[(champ, "anomalie", x.split(" (")[0])] += 1
                det.append([r["racine_client"], r["profil_derive"], champ, repr(v), len(v), " ; ".join(a),
                            r["type_identifiant"]])
    synth = [
        ["Règle demandée", f"NNI : chiffres uniquement, ≤ {MAX_NNI} ; NIF : chiffres uniquement, ≤ {MAX_NIF}. "
                           "Chaînes de caractères (zéros initiaux conservés)."],
        ["Stockage actuel", "clientele_clients.nni / nif : VARCHAR(40) — chaînes, zéros préservés. Aucune contrainte."],
        ["Obligation (PROPOSITION à valider)", "PP : NNI obligatoire, NIF facultatif ; PM : NIF obligatoire, NNI facultatif ; "
                                               "profil non identifié : non déterminé."],
        ["Contrôle actuel", "Import hebdomadaire : longueur seulement (signalée, jamais tronquée). Import ORION : aucun."],
        [], ["Profil", "Champ", "Statut", "Clients"],
        *[[p, ch, s, n] for (p, ch, s), n in sorted(((k, n) for k, n in c.items() if len(k) == 3 and k[1] in ("NNI", "NIF")), key=str)],
        [], ["Anomalies (non conformes)", "Champ", "Type", "Occurrences"],
        *[["", k[0], k[2], n] for k, n in sorted(c.items(), key=str) if len(k) == 3 and k[1] == "anomalie"],
        [], ["Conformes commençant par 0 (zéro initial à préserver)", "NNI", c[("NNI", "zeros")], "NIF", c[("NIF", "zeros")]],
    ]
    resume = {ch: sum(n for k, n in c.items() if len(k) == 3 and k[1] == ch and k[2] == "NON CONFORME") for ch in ("NNI", "NIF")}
    return synth, det, resume


# --------------------------------------------------------------------------- 6. pays
def _lire_matrice(chemin: Path) -> list[dict]:
    wb = load_workbook(chemin, data_only=True, read_only=True)
    ws = next(w for w in wb.worksheets if slug_cle(w.title).startswith("ZONE"))
    out = []
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        nom, niv = (row + (None, None))[:2]
        if i <= 2 or nom is None and niv is None:
            continue
        out.append({"ligne": i, "brut": "" if nom is None else str(nom), "niveau_brut": "" if niv is None else str(niv)})
    return out


def _lire_fiche(chemin: Path | None) -> dict[str, dict]:
    if not chemin:
        return {}
    wb = load_workbook(chemin, data_only=True, read_only=True)
    ws = next(w for w in wb.worksheets if "ENG" in slug_cle(w.title))
    out = {}
    for row in ws.iter_rows(min_row=4, values_only=True):
        if row and row[0] and len(row) >= 7 and row[5]:
            out[slug_cle(row[5])] = {"fr": row[0], "score": row[1], "label": row[6]}
    return out


async def pays(db, chemin_matrice: Path, chemin_fiche: Path | None) -> tuple[list[list], list[list], list[list], dict]:
    src = _lire_matrice(chemin_matrice)
    fiche = _lire_fiche(chemin_fiche)
    ref_json = {slug_cle(p.get("nom_fr")): p for p in REFERENTIEL["pays"] if "MATRICE_V1" in (p.get("sources") or [])}
    vus: Counter = Counter(slug_cle(r["brut"]) for r in src)
    largeur = Counter(len(r["brut"]) for r in src).most_common(1)[0][0]
    table, stats = [], Counter()
    for r in src:
        propre = " ".join(r["brut"].split())
        k = slug_cle(propre)
        niv = NIVEAUX_PAYS.get(slug_cle(r["niveau_brut"]).replace("ELEVEE", "ELEVE"))
        anomalies = []
        if not niv:
            anomalies.append("NIVEAU NON RECONNU")
        if vus[k] > 1:
            anomalies.append("DOUBLON")
        if len(r["brut"]) == largeur and not r["brut"].endswith(" ") and len(propre) == largeur:
            anomalies.append(f"NOM PROBABLEMENT TRONQUÉ ({largeur} caractères pleins)")
        j = ref_json.get(k)
        if not j:
            anomalies.append("ABSENT DU RÉFÉRENTIEL MOTEUR (JSON)")
        elif j.get("niveau_matrice") != niv:
            anomalies.append(f"NIVEAU DIFFÉRENT DANS LE MOTEUR ({j.get('niveau_matrice')})")
        f = fiche.get(k)
        stats[niv or "?"] += 1
        for a in anomalies:
            stats[("anomalie", a.split(" (")[0])] += 1
        table.append([r["ligne"], r["brut"], r["niveau_brut"], propre, k, niv or "", "",
                      "Non déterminé : aucune table ISO 3166-1 de référence dans le projet",
                      (f or {}).get("fr") or "", (f or {}).get("label") or "", " ; ".join(anomalies) or "OK",
                      "Fourni par la Banque — validation Conformité à confirmer"])

    nat = await _q(db, """SELECT cl.nationalite v, count(*) n,
                                 count(*) FILTER (WHERE c.niveau IS NOT NULL) classes
                          FROM clientele_clients cl LEFT JOIN clientele_classifications c USING (racine_client)
                          GROUP BY 1 ORDER BY 2 DESC""")
    rap = []
    for r in nat:
        v = r["v"]
        if not v or not v.strip():
            rap.append(["Nationalité (ORION)", v, r["n"], "—", "—", "VIDE", "Donnée manquante"])
            stats[("nat", "VIDE")] += r["n"]
            continue
        if _est_mauritanie(v):
            p = next((p for p in REFERENTIEL["pays"] if p.get("code") == "MAURITANIA"), None)
        else:
            p = _chercher_pays(v)
        if not p:
            statut = "NON RAPPROCHÉ"
        elif "MATRICE_V1" in (p.get("sources") or []):
            statut = "RAPPROCHÉ (Matrice)"
        else:
            statut = "RAPPROCHÉ À UNE ENTRÉE FICHE SEULE — niveau Matrice absent (fusion incomplète)"
        stats[("nat", statut.split(" —")[0])] += r["n"]
        niveau = (p or {}).get("niveau_matrice") or ""
        if niveau == "INTERDIT":
            stats[("nat_interdit",)] += r["n"]
        rap.append(["Nationalité (ORION)", v, r["n"], (p or {}).get("nom_fr") or "", niveau, statut,
                    "Le moteur applique MOYEN forfaitaire à toute nationalité étrangère (niveau pays non utilisé)."])
    res = await _q(db, """SELECT p.pays_residence v, count(*) n, count(d.racine_client) liees
                          FROM eer_parties p LEFT JOIN eer_dossiers d ON d.client_partie_id = p.id AND d.deleted_at IS NULL
                          WHERE p.pays_residence IS NOT NULL GROUP BY 1""")
    for r in res:
        p = _chercher_pays(r["v"])
        rap.append(["Pays de résidence (EER)", r["v"], r["n"], (p or {}).get("nom_fr") or "", (p or {}).get("niveau_matrice") or "",
                    "RAPPROCHÉ" if p else "NON RAPPROCHÉ",
                    f"{r['liees']} dossier(s) EER rattaché(s) à une racine client : le moteur ne lit ce pays que via l'EER."])
    doublons_fusion = []
    for p in REFERENTIEL["pays"]:
        if p.get("sources") == ["SCORING_V4"]:
            doublons_fusion.append([p.get("code"), p.get("nom_fr"), p.get("nom_en"), p.get("label_v4"),
                                    "Entrée Fiche non fusionnée avec la Matrice : un client rapproché ici n'a pas de niveau Matrice"])
    stats["lignes_source"] = len(src)
    return table, rap, doublons_fusion, stats


# --------------------------------------------------------------------------- 7. HAUT / FORT / BAS
async def conversions(db) -> tuple[list[list], dict]:
    rows = await _q(db, """SELECT l.classe_source_brut brut, l.classe_source norm, count(*) n,
                                  count(c.racine_client) classifications
                           FROM clientele_classif_lot_lignes l
                           LEFT JOIN clientele_classifications c
                             ON c.lot_id = l.lot_id AND c.ligne_source = l.ligne
                           GROUP BY 1,2 ORDER BY 3 DESC""")
    synonymes = {"HAUT", "FORT", "BAS"}
    out = [
        ["Emplacement", "backend/app/services/clientele/lots.py — normaliser_niveau() (lignes 119-131) : "
                        "HAUT / FORT → ELEVE ; BAS → FAIBLE. Utilisé à l'analyse d'un lot (colonne classe de risque)."],
        ["Contrôle fichier / moteur", "lots.py — CORRESPONDANCES_OFFICIELLES : FAIBLE, MOYEN, ELEVE, INTERDIT "
                                      "uniquement ; HAUT / FORT / BAS → NON_COMPARABLE (aucune décision possible)."],
        ["Valeur brute conservée", "Oui : clientele_classif_lot_lignes.classe_source_brut (valeur normalisée : classe_source)."],
        ["Tests", "backend/tests/test_clientele_lots.py::test_niveau_normalise ne teste pas HAUT / FORT / BAS."],
        [], ["Valeur brute", "Valeur normalisée actuelle", "Lignes de lot", "Classifications issues de ces lignes", "Synonyme non officiel ?"],
        *[[r["brut"], r["norm"], r["n"], r["classifications"],
           "OUI" if slug_cle((r["brut"] or "").replace("risque", "")).strip("_") in synonymes else "non"] for r in rows],
        [], ["Proposition pour la Conformité", "Ne retenir que Faible / Moyen / Élevé / Interdit. HAUT, FORT, BAS → "
                                               "« classe illisible » (anomalie de ligne), valeur brute conservée, aucune "
                                               "classe déduite. À appliquer seulement après validation."],
    ]
    concernes = sum(r["n"] for r in rows if slug_cle((r["brut"] or "").replace("risque", "")).strip("_") in synonymes)
    assert normaliser_niveau("HAUT") == "ELEVE" and "HAUT" not in CORRESPONDANCES_OFFICIELLES
    return out, {"lignes_synonymes": concernes}


# --------------------------------------------------------------------------- 8. divergences
REGLE_MOTEUR = {
    "PAYS": "Pays de résidence : niveau Matrice retenu, Fiche en contrôle (SOURCE_V1). Nationalité : MOYEN forfaitaire.",
    "SECTEUR": "Matrice retenue, conflit signalé dans le motif (CDC §13).",
    "FORME_JURIDIQUE": "Conflit : critère exclu de la somme et évaluation « À arbitrer » (non admissible).",
    "PROFIL": "Conflit : critère exclu de la somme et évaluation « À arbitrer » (non admissible).",
    "PP_CATEGORIE": "Catégorie PP non évaluée séparément (résidence = critère distinct).",
    "PPE_NON": "PPE = Non → 100 (FAIBLE) additionné. Donnée absente ≠ Non.",
    "SANCTIONS": "Alerte confirmée → ÉLEVÉ, statut A_ARBITRER.",
    "SEUIL_SCORE": "Seuils BEA DIGITAL appliqués : ≥ 100 000 ÉLEVÉ ; ≥ 10 000 MOYEN ; sinon FAIBLE ; 0 critère = NON CLASSÉ.",
    "FEUILLE_LISTE": "Colonnes Liste!K:L et M:N jamais chargées ; table Pays Eng-FR utilisée.",
    "NATIONALITE_V4": "Étrangère = MOYEN (sources concordantes).",
}
DECISION_ATTENDUE = {
    "PAYS": "Confirmer le niveau du pays (Matrice seule, Fiche absente ou différente).",
    "SECTEUR": "Choisir le niveau applicable au secteur.",
    "FORME_JURIDIQUE": "PM privée (SARL, SA, SUARL…) : MOYEN (Matrice) ou FAIBLE (Fiche) ?",
    "PROFIL": "Profil PM privée : MOYEN (Matrice) ou FAIBLE (Fiche) ?",
    "PP_CATEGORIE": "Niveau des catégories PP concernées.",
    "PPE_NON": "PPE = Non doit-il additionner 100 points ?",
    "SANCTIONS": "Filtrage confirmé : ÉLEVÉ ou INTERDIT ?",
    "SEUIL_SCORE": "Valider les seuils proposés par BEA DIGITAL.",
    "FEUILLE_LISTE": "Confirmer que la table Pays Eng-FR est la seule source Fiche des pays.",
    "NATIONALITE_V4": "Confirmer étrangère = MOYEN.",
}


async def divergences(db, ev: dict) -> tuple[list[list], dict]:
    rows = await _q(db, """SELECT id, domaine, cle, source_a, valeur_a, source_b, valeur_b, statut, note,
                                  validee_le FROM clientele_classif_divergences ORDER BY domaine, cle""")
    nat_count = Counter()
    for e in ev.values():
        lg = e["lignes"].get("NATIONALITE")
        if lg and lg["valeur"]:
            p = None if _est_mauritanie(lg["valeur"]) else _chercher_pays(lg["valeur"])
            if p:
                nat_count[slug_cle(p.get("nom_fr"))] += 1

    def pm_privee(critere: str) -> tuple[int, Counter, Counter]:
        n, v1, v4 = 0, Counter(), Counter()
        for e in ev.values():
            lg = e["lignes"].get(critere)
            if lg and lg["divergence"] and lg["etat"] == "EVALUE":
                n += 1
                av = e.get("niveau_retenu") or "(aucun)"
                v1[(av, _impact(e, POIDS_ECHELLE["MOYEN"], critere) or "(aucun)")] += 1
                v4[(av, _impact(e, POIDS_ECHELLE["FAIBLE"], critere) or "(aucun)")] += 1
        return n, v1, v4

    secteurs = Counter()
    for e in ev.values():
        lg = e["lignes"].get("SECTEUR")
        if lg and lg["divergence"]:
            secteurs[lg["valeur"]] += 1

    def fmt(c: Counter) -> str:
        return " ; ".join(f"{a}→{b} : {n}" for (a, b), n in sorted(c.items(), key=str) if a != b) or "aucun changement de niveau"

    out = []
    for r in rows:
        dom = r["domaine"]
        concernes, impact, propo = "", "", ""
        if dom == "PAYS":
            n = nat_count.get(slug_cle(r["cle"]), 0)
            concernes = f"{n} client(s) de cette nationalité ; 0 pays de résidence alimenté"
            impact = "Aucun sur les classements actuels (pays de résidence non alimenté ; nationalité = MOYEN forfaitaire)."
            propo = ("Fiche ABSENTE : niveau Matrice seul disponible — pas de conflit de niveau, à confirmer."
                     if r["valeur_b"] == "ABSENT" else
                     "Matrice ABSENTE : aucune valeur Matrice, le niveau Fiche ne doit pas être substitué." if r["valeur_a"] == "ABSENT"
                     else "Niveaux différents : arbitrage Conformité.")
        elif dom in ("FORME_JURIDIQUE", "PROFIL"):
            n, v1, v4 = pm_privee(dom)
            concernes = f"{n} client(s) dont le critère {dom} est en conflit dans l'évaluation de référence"
            impact = f"Si Matrice (MOYEN, +10 000) : {fmt(v1)}. Si Fiche (FAIBLE, +100) : {fmt(v4)}."
            propo = "Aucune proposition technique : choix de politique de risque."
        elif dom == "SECTEUR":
            concernes = "; ".join(f"{k} : {n}" for k, n in secteurs.most_common(5)) or "0 client avec conflit sectoriel évalué"
            impact = "Matrice déjà retenue par le moteur : impact seulement si la Conformité choisit la Fiche."
            propo = "Maintenir la Matrice tant que non arbitré (règle CDC §13 déjà en place)."
        elif dom == "SEUIL_SCORE":
            concernes, impact = "Tous les clients évalués", "Structurant : conditionne tous les niveaux."
            propo = "Seuils BEA DIGITAL évitent l'ambiguïté V4 « score == 10 000 »."
        elif dom == "PPE_NON":
            concernes, impact = "0 (aucun EER rattaché à un client)", "Aucun aujourd'hui."
        elif dom == "SANCTIONS":
            concernes, impact = "0 alerte confirmée", "Aucun aujourd'hui."
        else:
            concernes, impact = "—", "—"
        out.append([str(r["id"])[:8], dom, r["cle"], f"{r['source_a']} : {r['valeur_a']}", f"{r['source_b']} : {r['valeur_b']}",
                    REGLE_MOTEUR.get(dom, "—"), concernes, impact, propo, DECISION_ATTENDUE.get(dom, "—"),
                    STATUT_DIV.get(r["statut"], r["statut"]), r["note"] or ""])
    vals = await _q(db, """SELECT id, dimension, critere, code, libelle, niveau_v1, score_v1, niveau_v4, score_v4,
                                  statut, note FROM clientele_classif_valeurs
                           WHERE statut IN ('A_ARBITRER', 'A_CONFIGURER') ORDER BY statut, critere, code""")
    for v in vals:
        if v["statut"] == "A_ARBITRER":
            crit = "PROFIL" if v["critere"] == "PROFIL" else "FORME_JURIDIQUE"
            n, v1, v4 = pm_privee(crit)
            impact = f"Si Matrice : {fmt(v1)}. Si Fiche : {fmt(v4)}."
            concernes = f"{n} client(s) (critère {crit} en conflit)"
            dec = "MOYEN (Matrice) ou FAIBLE (Fiche) ?"
        else:
            concernes, impact, dec = "0 (critère Canal non alimenté)", "Aucun", "Définir le canal, sa source et son niveau."
        out.append([str(v["id"])[:8], f"VALEUR {v['dimension']}/{v['critere']}", f"{v['code']} — {v['libelle']}",
                    f"Matrice : {v['niveau_v1'] or '—'} ({v['score_v1'] if v['score_v1'] is not None else '—'})",
                    f"Fiche : {v['niveau_v4'] or '—'} ({v['score_v4'] if v['score_v4'] is not None else '—'})",
                    "Exclu de la somme (A_ARBITRER / A_CONFIGURER)", concernes, impact,
                    "Aucune : décision Conformité.", dec,
                    "À arbitrer" if v["statut"] == "A_ARBITRER" else "À analyser (à configurer)", v["note"] or ""])
    return out, {"divergences": len(rows), "a_arbitrer": sum(1 for r in rows if r["statut"] == "A_ARBITRER"),
                 "valeurs": len(vals)}


# --------------------------------------------------------------------------- 10. non classés + réexamens
def _cause(e: dict) -> str:
    st = e.get("statut")
    if not e.get("id"):
        return "Aucune évaluation photographiée"
    if st == "A_ARBITRER":
        conf = [c for c, lg in e["lignes"].items() if lg["statut_regle"] == "A_ARBITRER" and lg["etat"] == "EVALUE"]
        return "Conflit de référentiel non arbitré : " + ", ".join(sorted(conf))
    if st == "NON_CLASSE":
        nd = [c for c, lg in e["lignes"].items() if lg["etat"] in ("NON_DISPONIBLE", "A_VERIFIER")
              and c in ("PROFIL", "FORME_JURIDIQUE", "RESIDENCE", "NATIONALITE", "SECTEUR", "LISTE_INTERDICTION")]
        return "Aucun critère évalué — données manquantes / non reconnues : " + ", ".join(sorted(nd))
    if st in ("EVALUE", "BLOQUANT"):
        return "Évaluation admissible jamais enregistrée (classement de masse verrouillé)"
    return f"Autre ({st})"


async def non_classes(db, ev: dict) -> tuple[list[list], list[list], dict]:
    det, causes = [], Counter()
    for r, e in ev.items():
        if e["classe"]:
            continue
        cause = _cause(e)
        causes[cause.split(" :")[0]] += 1
        lg = e["lignes"]
        det.append([r, e.get("nom_client"), e.get("profil_derive"), e.get("code_agence"), e.get("statut"),
                    e.get("score_total"), e.get("nb_evalues"), e.get("version_moteur"),
                    e["created_at"].strftime("%d/%m/%Y") if e.get("created_at") else None, cause,
                    *[f"{lg[c]['etat']}{' : ' + str(lg[c]['valeur'])[:40] if lg.get(c) and lg[c]['valeur'] else ''}"
                      if lg.get(c) else "—" for c in ("PROFIL", "FORME_JURIDIQUE", "RESIDENCE", "NATIONALITE", "SECTEUR")]])
    clients_sans_eval = await _q(db, """SELECT count(*) n FROM clientele_clients cl WHERE NOT EXISTS
        (SELECT 1 FROM clientele_classifications c WHERE c.racine_client = cl.racine_client)
        AND NOT EXISTS (SELECT 1 FROM clientele_classif_evaluations e WHERE e.racine_client = cl.racine_client)""")
    for _ in range(clients_sans_eval[0]["n"]):
        causes["Aucune évaluation photographiée"] += 1
    reex = await _q(db, """
        SELECT r.id, r.racine_client, s.nom_client, r.origine, r.niveau_actuel, r.source_actuelle, r.niveau_suggere,
               r.source_suggestion, r.motif, r.statut, r.created_at, u.full_name auteur,
               c.niveau niveau_courant, coalesce(c.mode_decision, c.source) mode_courant, c.motif_decision_manuelle,
               cu.full_name decide_par, c.classifie_le,
               (SELECT count(*) FROM clientele_classif_historique h WHERE h.racine_client = r.racine_client) nb_histo
        FROM clientele_classif_reexamens r
        LEFT JOIN clientele_situation s ON s.racine_client = r.racine_client
        LEFT JOIN users u ON u.id = r.created_by_id
        LEFT JOIN clientele_classifications c ON c.racine_client = r.racine_client
        LEFT JOIN users cu ON cu.id = c.classifie_par_id
        ORDER BY r.created_at""")
    reex_l = [[str(x["id"])[:8], x["racine_client"], x["nom_client"], x["origine"], x["statut"],
               x["niveau_actuel"], x["source_actuelle"], x["niveau_suggere"], x["source_suggestion"], x["motif"],
               x["created_at"].strftime("%d/%m/%Y %H:%M"), x["auteur"], x["niveau_courant"], x["mode_courant"],
               x["motif_decision_manuelle"], x["decide_par"],
               x["classifie_le"].strftime("%d/%m/%Y %H:%M") if x["classifie_le"] else None, x["nb_histo"],
               "Traitement par un arbitre (POST /classification/reexamens/{id}/traiter) — aucune fermeture automatique"]
              for x in reex]
    return det, reex_l, {"causes": dict(causes), "reexamens": len(reex),
                         "ouverts": sum(1 for x in reex if x["statut"] == "OUVERT")}


# --------------------------------------------------------------------------- 11. produit / canal
async def produit_canal(db) -> list[list]:
    ref = REFERENTIEL
    ncg = await _q(db, "SELECT count(DISTINCT ncg) n, count(DISTINCT rubrique_comptable) r FROM clientele_comptes")
    rub = await _q(db, "SELECT rubrique_comptable v, count(*) n FROM clientele_comptes GROUP BY 1 ORDER BY 2 DESC LIMIT 15")
    eer = await _q(db, """SELECT count(*) n, count(operation_type) op, count(origine_fonds) orig,
                                 count(racine_client) lies FROM eer_dossiers WHERE deleted_at IS NULL""")
    out = [
        ["Critère moteur", "État dans le moteur", "Source attendue", "Source disponible aujourd'hui", "Clé de rapprochement", "Propriétaire (à confirmer)"],
        ["PRODUIT", "NON_DISPONIBLE", "Matrice feuille « Produit - Service - Opération »", "Aucune (pas de flux opérations)", "racine / compte", "Direction des Opérations / ORION"],
        ["TYPE_OPERATION", "NON_DISPONIBLE", "Flux opérations", "Aucune", "compte + date", "ORION"],
        ["PAYS_PROVENANCE_FONDS / DESTINATION", "NON_DISPONIBLE", "Transferts (SWIFT / RIA…)", "Aucune", "compte + référence", "Monétique / International"],
        ["CANAL", "NON_DISPONIBLE", "Canal d'ouverture / d'opération", "Aucune (absent d'ORION État des comptes)", "compte", "Réseau / Digital"],
        ["ORIGINE_FONDS", "NON_DISPONIBLE", "EER.origine_fonds", f"{eer[0]['orig']} dossier(s) EER renseignés, {eer[0]['lies']} rattaché(s) à une racine", "racine_client EER", "Conformité / EER"],
        [], ["Référentiel Matrice — produits", "Niveau"],
        *[[p.get("libelle"), p.get("niveau_matrice")] for p in ref.get("produits_matrice") or []],
        [], ["Référentiel Matrice — canaux", "Niveau"],
        *[[p.get("libelle"), p.get("niveau_matrice")] for p in ref.get("canaux_matrice") or []],
        [], ["Canaux à configurer (référentiel maître)", "DIGITAL, APPLICATION, WEB, INTERMEDIAIRE, AUTRE — aucun score inventé"],
        [], ["Donnée candidate (NON utilisée)", "Comptes : NCG distincts", ncg[0]["n"], "rubriques comptables distinctes", ncg[0]["r"]],
        ["Rubrique comptable (top 15)", "Comptes"], *[[r["v"], r["n"]] for r in rub],
        [], ["Règle", "Ces critères restent NON_DISPONIBLE (hors score) tant que la source et la règle ne sont pas validées. "
                      "Aucune valeur par défaut."],
    ]
    return out


# --------------------------------------------------------------------------- 12. analystes KYC
async def analystes(db) -> tuple[list[list], dict]:
    lot = await _q(db, "SELECT analyste_kyc v, count(*) n FROM clientele_classif_lot_lignes WHERE analyste_kyc IS NOT NULL GROUP BY 1")
    cls = await _q(db, "SELECT analyste_kyc v, count(*) n FROM clientele_classifications WHERE analyste_kyc IS NOT NULL GROUP BY 1")
    users = await _q(db, "SELECT id, full_name, email, is_active FROM users WHERE deleted_at IS NULL")
    valeurs: dict[str, dict] = defaultdict(lambda: {"lot": 0, "cls": 0, "liste": False})
    for r in lot:
        valeurs[r["v"]]["lot"] += r["n"]
    for r in cls:
        valeurs[r["v"]]["cls"] += r["n"]
    for v in listes_observees().get("analystes") or []:
        valeurs[v]["liste"] = True
    norm = {u["id"]: normaliser(u["full_name"]) for u in users}
    out, stats = [], Counter()
    for v, info in sorted(valeurs.items()):
        nv = normaliser(v)
        if nv in ("NON IDENTIFIE", ""):
            statut, cands = "VALEUR DE CONSIGNE (pas un analyste)", []
        else:
            exact = [u for u in users if norm[u["id"]] == nv]
            tokens = set(nv.split())
            partiel = [u for u in users if u not in exact and tokens and tokens <= set(norm[u["id"]].split())]
            proches = [u for u in users if u not in exact and u not in partiel
                       and difflib.SequenceMatcher(None, nv, norm[u["id"]]).ratio() >= 0.8]
            if len(exact) == 1:
                statut, cands = "NOM IDENTIQUE (pas d'identifiant stable)", exact
            elif exact:
                statut, cands = "HOMONYMES", exact
            elif len(partiel) == 1:
                statut, cands = "PRÉNOM / NOM PARTIEL — candidat unique", partiel
            elif partiel:
                statut, cands = "PRÉNOM / NOM PARTIEL — plusieurs candidats", partiel
            elif proches:
                statut, cands = "NOM SIMILAIRE", proches
            else:
                statut, cands = "AUCUN COMPTE CORRESPONDANT", []
        stats[statut] += 1
        out.append([v, info["lot"], info["cls"], "Oui" if info["liste"] else "Non", statut,
                    " | ".join(f"{u['full_name']} <{u['email']}>{'' if u['is_active'] else ' (inactif)'}" for u in cands),
                    "À valider par le responsable KYC — aucun compte créé, aucun rôle attribué"])
    return out, dict(stats)


# --------------------------------------------------------------------------- orchestration
async def auditer(chemin_matrice: Path, chemin_fiche: Path | None, conservation: int) -> tuple[Path, bool]:
    moteur = creer_moteur_lecture_seule()
    try:
        async with fabrique_lecture_seule(moteur)() as db:
            gen, v = await etat_general(db)
            ev = await evaluations(db)
            de_synth, de_det, de_res = await date_entreprise(db, ev)
            st_synth, st_mixtes, st_res = await statuts_comptes(db)
            cat_table, cat_stats = await categories(db)
            id_synth, id_det, id_res = await identifiants(db)
            p_table, p_rap, p_fusion, p_stats = await pays(db, chemin_matrice, chemin_fiche)
            conv, conv_res = await conversions(db)
            div, div_res = await divergences(db, ev)
            nc_det, reex, nc_res = await non_classes(db, ev)
            pc = await produit_canal(db)
            ana, ana_res = await analystes(db)
    finally:
        await moteur.dispose()

    parametres = _parametres(**{"Matrice (source pays)": chemin_matrice.name,
                                "Fiche de scoring": chemin_fiche.name if chemin_fiche else "—"})
    wb = _classeur("audit classification", parametres)
    _feuille(wb, "0 État général", ["Élément", "Valeur", "", "", "", ""], gen)
    _feuille(wb, "2 Date entreprise — synthèse", ["Élément", "Valeur", "Clients"], de_synth)
    _feuille(wb, "2 Date entreprise — clients", [
        "Racine", "Nom", "Profil", "Catégorie juridique", "Situation juridique", "N° RCCM présent",
        "Date ouverture (lot)", "Date création (EER)", "Source utilisable", "Statut", "État critère moteur",
        "Source critère moteur", "1er compte ORION (NON utilisé)", "Niveau actuel",
        "Niveau si entreprise récente", "Niveau si entreprise ancienne"], de_det)
    _feuille(wb, "3 Statuts comptes", ["Élément", "Valeur"], st_synth)
    _feuille(wb, "3 Clients statuts mixtes", ["Racine", "Comptes", "Ouverts", "Fermés", "Agrégat actuel", "Statut"], st_mixtes)
    _feuille(wb, "4 Catégories juridiques", [
        "Champ source", "Valeur source (originale)", "Valeur normalisée", "Catégorie cible (référentiel)",
        "Profil cible", "Niveau Matrice", "Niveau Fiche (poids)", "Niveau de risque associé", "Clients",
        "Répartition par profil", "Type de correspondance", "Justification", "Statut de validation",
        "Date de modification"], cat_table)
    _feuille(wb, "5 NIF-NNI synthèse", ["Élément", "Valeur", "", ""], id_synth)
    _feuille(wb, "5 NIF-NNI non conformes", ["Racine", "Profil", "Champ", "Valeur brute", "Longueur",
                                             "Anomalies", "Type identifiant ORION"], id_det)
    _feuille(wb, "6 Pays — référentiel source", [
        "Ligne Excel", "Nom d'origine (brut)", "Niveau d'origine", "Nom nettoyé", "Clé normalisée",
        "Niveau normalisé", "Code ISO", "Motif ISO", "Nom Fiche (contrôle)", "Libellé Fiche", "Contrôle", "Statut"], p_table)
    _feuille(wb, "6 Pays — rapprochement clients", ["Champ", "Valeur", "Clients", "Pays rapproché",
                                                   "Niveau Matrice", "Statut", "Observation"], p_rap)
    _feuille(wb, "6 Pays — entrées non fusionnées", ["Code", "Nom", "Nom anglais", "Libellé Fiche", "Observation"], p_fusion)
    _feuille(wb, "7 HAUT-FORT-BAS", ["Élément", "Valeur", "", "", ""], conv)
    _feuille(wb, "8 Divergences et valeurs", [
        "Identifiant", "Domaine", "Règle / clé", "Valeur source A", "Valeur source B", "Règle moteur actuelle",
        "Clients concernés", "Impact potentiel", "Proposition technique", "Décision attendue", "Statut", "Note"], div)
    _feuille(wb, "10 Non classés", ["Racine", "Nom", "Profil", "Agence", "Statut évaluation", "Score", "Critères évalués",
                                    "Version moteur", "Évalué le", "Cause", "Profil (critère)", "Forme juridique",
                                    "Résidence", "Nationalité", "Secteur"], nc_det)
    _feuille(wb, "10 Réexamens", ["Id", "Racine", "Nom", "Origine", "Statut", "Niveau actuel", "Source actuelle",
                                  "Niveau suggéré", "Source suggestion", "Motif", "Ouvert le", "Ouvert par",
                                  "Classe courante", "Mode", "Motif décision manuelle", "Décidé par", "Décidé le",
                                  "Lignes d'historique", "Traitement attendu"], reex)
    _feuille(wb, "11 Produit-Canal", ["Élément", "", "", "", "", ""], pc)
    _feuille(wb, "12 Analystes KYC", ["Valeur du fichier", "Lignes de lot", "Classifications", "Liste observée",
                                      "Statut", "Comptes candidats", "Validation"], ana)

    nc_total = sum(nc_res["causes"].values())
    cat_clients = sum(n for k, n in cat_stats.items() if isinstance(k, str))
    controles = [
        _ctl("Classés + non classés = clients", v["clients"], v["classes"] + v["non_classes"]),
        _ctl("Non classés détaillés (avec ou sans évaluation) = non classés", v["non_classes"], nc_total),
        _ctl("Moteur + manuels = classés", v["classes"], v["moteur"] + v["manuels"]),
        _ctl("Classement de masse verrouillé", False, v["verrou_masse_actif"]),
        _ctl("Lignes pays importées = lignes source", p_stats["lignes_source"], len(p_table)),
        _ctl("Somme des niveaux pays = lignes source", p_stats["lignes_source"],
             sum(n for k, n in p_stats.items() if k in ("FAIBLE", "MOYEN", "ELEVE", "INTERDIT", "?"))),
        _ctl("Catégories : clients comptés deux fois (catégorie + situation)", 2 * v["clients"], cat_clients),
        _ctl("Réexamens détaillés = réexamens", nc_res["reexamens"], len(reex)),
    ]
    resume = {"etat": v, "date_entreprise": de_res, "statuts": st_res, "nif_nni": id_res,
              "pays": {k: n for k, n in p_stats.items() if isinstance(k, str)},
              "conversions": conv_res, "divergences": div_res, "non_classes": nc_res["causes"],
              "analystes": ana_res,
              "categories": {k: n for k, n in cat_stats.items() if isinstance(k, str)}}
    for k, n in sorted(p_stats.items(), key=str):
        print("pays", k, n)
    for k, n in sorted(cat_stats.items(), key=str):
        print("categories", k, n)
    for k, n in nc_res["causes"].items():
        print("non classés", k, n)
    print("resume", resume)
    nom = f"audit-classification-{date.today():%Y%m%d}-{__import__('datetime').datetime.now():%H%M%S}.xlsx"
    return _enregistrer(wb, nom, "audit", parametres, controles, resume, conservation)


def main() -> None:
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--matrice", required=True, type=Path)
    p.add_argument("--fiche", type=Path)
    p.add_argument("--conservation-jours", type=int, default=CONSERVATION_JOURS)
    a = p.parse_args()
    chemin, ok = asyncio.run(auditer(a.matrice, a.fiche, a.conservation_jours))
    print(f"Classeur : {chemin}")
    print("Contrôles : OK" if ok else "Contrôles : ÉCART — voir la feuille « Contrôles »")
    sys.exit(0 if ok else 2)


if __name__ == "__main__":
    main()
